"""PF / ESI coverage as an Excel file and a PDF: how many employees have PF / ESI deductions and how many do not.

    GET /api/pf-esi/coverage/export?format=xlsx|pdf&type=staff|production|all[&branchId=]

The same numbers as GET /api/pf-esi/coverage (both read pf_esi_views.load_coverage, so a file can never disagree with the page):
one section per employee class (Staff, Production) holding a summary block (the counts of pf_esi.coverage_counts) and one row
per employee (pf_esi.coverage_rows). Excel gets one worksheet per class, the PDF one printable section per class.

Permission is the one of /pf-esi/coverage (the `employees` module, see permission_registry.URL_MODULE_MAP); a branch-scoped login
only ever gets its own branch. READ-ONLY apart from one audit-log line per file. Errors use the project's JSON shape
({"error": "..."}).

`format` is the file type here, but Django REST framework reads a `?format=` query parameter as a renderer suffix and answers
404 for one it does not know: this view therefore turns that negotiation off (the only renderer is JSON, used for errors).
"""

from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from rest_framework.decorators import api_view
from rest_framework.negotiation import BaseContentNegotiation
from rest_framework.request import Request
from rest_framework.response import Response
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image as RLImage,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from . import pf_esi
from .audit_utils import log_action
from .auth import require_hr
from .branch_scope import get_branch_scope
from .clock import ist_now, ist_today
from .models import Branch
from .pf_esi_views import CoverageParamError, load_coverage
from .reporting.access import hr_display_name
from .reporting.branding import company
from .reporting.definitions.pf_esi_coverage import CATEGORY_TEXT, CLASS_TITLES, STATUS_TEXT, yes_no
from .reporting.export_pdf import BRAND as PDF_BRAND
from .reporting.export_pdf import BRAND_SOFT as PDF_BRAND_SOFT
from .reporting.export_pdf import GRID as PDF_GRID
from .reporting.export_pdf import INK, MUTED, ZEBRA as PDF_ZEBRA
from .reporting.export_pdf import _canvas_maker, ensure_fonts
from .reporting.export_xlsx import BRAND, BRAND_SOFT, GRID, ZEBRA, safe_cell, sheet_title
from .view_common import error_response

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"
#: The Report Center's ceilings for one file: a bigger company narrows the filters (type / branch) instead of getting a cut-off one.
XLSX_ROW_LIMIT = 30_000
PDF_ROW_LIMIT = 5_000

#: (key of pf_esi.coverage_counts, label, indented): the order reads as the arithmetic
#: total = coEmp + tds + notCoEmp, and coEmp = withAny + withNeither + notEvaluated.
SUMMARY_LINES = (
    ("total", "Active employees", False),
    ("coEmp", "Co Emp (PF and ESI apply)", False),
    ("withPf", "With PF", True),
    ("withEsi", "With ESI", True),
    ("withBoth", "With both PF and ESI", True),
    ("withAny", "With PF and/or ESI", True),
    ("withNeither", "Co Emp but no PF or ESI (no rule matched)", True),
    ("notEvaluated", "Not evaluated yet (production: no payroll generated)", True),
    ("tds", "TDS employees (TDS only)", False),
    ("notCoEmp", "Not Co Emp (no PF or ESI)", False),
)
COLUMNS = (
    ("Emp Code", 12, "left"),
    ("Employee", 28, "left"),
    ("Class", 12, "left"),
    ("Department", 22, "left"),
    ("Branch", 18, "left"),
    ("Category", 13, "center"),
    ("PF", 7, "center"),
    ("PF %", 8, "right"),
    ("ESI", 7, "center"),
    ("ESI %", 8, "right"),
    ("Status", 22, "left"),
    ("TDS %", 8, "right"),
)
NOTES = (
    "Only Co Emp employees (Employees > Co Emp) can have PF or ESI. TDS employees have TDS only, and every other employee has "
    "neither.",
    "Staff are judged on their current salary split by the PF and ESI rules of Settings > PF & ESI (a unit's own rules when it "
    "has them, else the rules for all units). This shows what the next payroll would do, not what a past payroll deducted.",
    "Production workers are judged on the monthly equivalent of their latest generated payroll period; a Co Emp worker with no "
    "payroll yet is 'Not evaluated'.",
    "Only active employees are counted.",
)


class _IgnoreFormatParam(BaseContentNegotiation):
    """See the module notes: `?format=` is the file type here, not a DRF renderer suffix."""

    def select_parser(self, request, parsers):
        return parsers[0] if parsers else None

    def select_renderer(self, request, renderers, format_suffix=None):
        return renderers[0], renderers[0].media_type


# ─────────────────────────────────────────────────────────────────────────────
#  The data of one export
# ─────────────────────────────────────────────────────────────────────────────


def _natural(text) -> list:
    return [(0, int(p), "") if p.isdigit() else (1, 0, p.lower()) for p in re.split(r"(\d+)", str(text or "")) if p]


def _class_of(row: dict) -> str:
    return pf_esi.CLASS_PRODUCTION if row["employmentType"] == "production" else pf_esi.CLASS_STAFF


def _branch_text(request: Request, branch_id: int | None) -> str:
    wanted = branch_id if branch_id is not None else get_branch_scope(request)
    if wanted is None:
        return "All branches"
    name = Branch.objects.filter(pk=wanted).values_list("name", flat=True).first()
    return name or f"Branch {wanted}"


def build_dataset(request: Request) -> dict:
    """{"type", "branch", "sheets": [{"class", "title", "counts", "rows"}]} for the request. Raises CoverageParamError."""
    kind, branch_id, rows = load_coverage(request)
    sheets = []
    for emp_class in (pf_esi.CLASS_STAFF, pf_esi.CLASS_PRODUCTION):
        if kind not in ("all", emp_class):
            continue
        mine = sorted((r for r in rows if _class_of(r) == emp_class), key=lambda r: _natural(r["employeeCode"]))
        sheets.append(
            {
                "class": emp_class,
                "title": CLASS_TITLES[emp_class],
                "counts": pf_esi.coverage_counts(rows, emp_class),
                "rows": mine,
            }
        )
    return {"type": kind, "branch": _branch_text(request, branch_id), "sheets": sheets}


def _filter_text(data: dict) -> str:
    scope = {"all": "Staff and production", "staff": "Staff", "production": "Production"}[data["type"]]
    return f"Active employees   |   Type: {scope}   |   Branch: {data['branch']}"


def _detail_cells(row: dict) -> list:
    """One employee's cells, in COLUMNS order (rates stay numbers, blanks stay None)."""
    return [
        row["employeeCode"],
        row["name"],
        CLASS_TITLES[_class_of(row)],
        row["departmentName"],
        row["branchName"],
        CATEGORY_TEXT[row["category"]],
        yes_no(row["pf"]),
        row["pfRatePercent"],
        yes_no(row["esi"]),
        row["esiRatePercent"],
        STATUS_TEXT[row["status"]],
        row["tdsPercent"],
    ]


# ─────────────────────────────────────────────────────────────────────────────
#  Excel
# ─────────────────────────────────────────────────────────────────────────────


def build_xlsx(data: dict, generated_by: str) -> bytes:
    co = company()
    stamp = f"Generated {ist_now().strftime('%d-%b-%Y %H:%M')} by {generated_by}"
    n_cols = len(COLUMNS)
    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_company = Font(name="Calibri", size=15, bold=True, color=BRAND)
    f_sub = Font(name="Calibri", size=10, color="666666")
    f_title = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    f_head = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    f_body = Font(name="Calibri", size=10)
    f_bold = Font(name="Calibri", size=10, bold=True)
    f_note = Font(name="Calibri", size=9, italic=True, color="666666")
    f_section = Font(name="Calibri", size=11, bold=True, color=BRAND)
    fill_brand = PatternFill("solid", fgColor=BRAND)
    fill_soft = PatternFill("solid", fgColor=BRAND_SOFT)
    fill_zebra = PatternFill("solid", fgColor=ZEBRA)

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "UKTextiles HRMS"
    wb.properties.title = "PF / ESI Coverage"

    def put(ws, row, col, value, font=f_body, fill=None, align="left", wrap=False, fmt=None, box=False):
        cell = ws.cell(row=row, column=col, value=safe_cell(value))
        if isinstance(value, str):
            cell.data_type = "s"  # an employee name that starts with "=" must stay text, never become a formula
        cell.font = font
        if fill is not None:
            cell.fill = fill
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fmt:
            cell.number_format = fmt
        if box:
            cell.border = border
        return cell

    for sheet in data["sheets"]:
        ws = wb.create_sheet(sheet_title(f"{sheet['title']} PF-ESI coverage"))
        ws.sheet_view.showGridLines = False
        for index, (_label, width, _align) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width

        # letterhead
        indent = ""
        if co.logo_png:
            try:
                image = XlImage(io.BytesIO(co.logo_png))
                image.width = image.height = 44
                ws.add_image(image, "A1")
                indent = "      "
            except Exception:  # a logo that cannot be drawn must not stop the file
                pass
        put(ws, 1, 1, indent + co.name, f_company)
        ws.row_dimensions[1].height = 26
        sub = "  |  ".join(b for b in (co.address, co.contact, f"GSTIN {co.gstin}" if co.gstin else "") if b)
        put(ws, 2, 1, indent + (sub or co.tagline), f_sub)
        for col in range(1, n_cols + 1):
            put(ws, 3, col, None, f_title, fill_brand)
        put(ws, 3, 1, f"PF / ESI COVERAGE - {sheet['title'].upper()}", f_title, fill_brand)
        ws.row_dimensions[3].height = 24
        for col in range(1, n_cols + 1):
            put(ws, 4, col, None, f_bold, fill_soft)
        put(ws, 4, 1, _filter_text(data), f_bold, fill_soft)
        put(ws, 5, 1, stamp, f_note)

        # summary block: the label spans A:C, the number sits in D
        row = 7
        put(ws, row, 1, "SUMMARY", f_section)
        row += 1
        for key, label, indented in SUMMARY_LINES:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            put(ws, row, 1, ("    " if indented else "") + label, f_body if indented else f_bold, box=True)
            for col in (2, 3):
                ws.cell(row=row, column=col).border = border
            put(ws, row, 4, sheet["counts"][key], f_bold, align="right", fmt="#,##0", box=True)
            row += 1

        # detail table
        row += 1
        put(ws, row, 1, "EMPLOYEES", f_section)
        row += 1
        header_row = row
        for col, (label, _width, _align) in enumerate(COLUMNS, start=1):
            put(ws, row, col, label, f_head, fill_brand, align="center", wrap=True, box=True)
        ws.row_dimensions[row].height = 28
        if not sheet["rows"]:
            row += 1
            put(ws, row, 1, "No active employees match the selected filters.", f_note)
        for index, record in enumerate(sheet["rows"], start=1):
            row += 1
            fill = fill_zebra if index % 2 == 0 else None
            for col, ((_label, _width, align), value) in enumerate(zip(COLUMNS, _detail_cells(record)), start=1):
                fmt = '0.##"%"' if _label.endswith("%") and value is not None else None
                put(ws, row, col, value, f_body, fill, align=align, fmt=fmt, box=True)
        if sheet["rows"]:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{row}"

        # notes
        row += 2
        for note in NOTES:
            put(ws, row, 1, note, f_note)
            row += 1
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
#  PDF
# ─────────────────────────────────────────────────────────────────────────────

_PDF_WIDTHS_CM = (2.0, 4.6, 2.0, 3.4, 3.0, 2.1, 1.2, 1.4, 1.2, 1.4, 3.4, 1.4)


def _percent_text(value) -> str:
    return "-" if value is None else f"{value:g}%"


def _pdf_cells(record: dict) -> list:
    cells = _detail_cells(record)
    out = []
    for (label, _width, _align), value in zip(COLUMNS, cells):
        if label.endswith("%"):
            out.append(_percent_text(value))
        else:
            out.append("-" if value in (None, "") else str(value))
    return out


def build_pdf(data: dict, generated_by: str) -> bytes:
    regular, bold, _rupee_ok = ensure_fonts()
    co = company()
    page = landscape(A4)
    margin = 1.0 * cm
    avail = page[0] - 2 * margin
    widths = [w / sum(_PDF_WIDTHS_CM) * avail for w in _PDF_WIDTHS_CM]
    size = 7.5

    def style(name, **kw) -> ParagraphStyle:
        base = dict(fontName=regular, fontSize=size, leading=size * 1.25, textColor=INK)
        base.update(kw)
        return ParagraphStyle(name, **base)

    s_head = style("h", fontName=bold, alignment=TA_CENTER, textColor=colors.white)
    s_note = style("n", fontSize=7, leading=9, textColor=MUTED)
    aligns = {"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT}
    cell_styles = [style(f"c{i}", alignment=aligns[a]) for i, (_l, _w, a) in enumerate(COLUMNS)]

    story: list = []

    # letterhead
    left = [
        Paragraph(f"<b>{escape(co.name)}</b>", style("co", fontName=bold, fontSize=14, leading=17, textColor=PDF_BRAND))
    ]
    sub = "  |  ".join(b for b in (co.address, co.contact, f"GSTIN {co.gstin}" if co.gstin else "") if b)
    if sub:
        left.append(Paragraph(escape(sub), style("sub", fontSize=7.5, leading=10, textColor=MUTED)))
    logo = None
    if co.logo_png:
        try:
            logo = RLImage(io.BytesIO(co.logo_png), width=1.3 * cm, height=1.3 * cm, kind="proportional")
        except Exception:
            logo = None
    header = (
        Table([[logo, left]], colWidths=[1.7 * cm, avail - 1.7 * cm])
        if logo is not None
        else Table([[left]], colWidths=[avail])
    )
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(header)
    story.append(Spacer(1, 4))

    def bar(text, background, text_color, font, font_size):
        table = Table(
            [
                [
                    Paragraph(
                        escape(text),
                        style("bar", fontName=font, fontSize=font_size, leading=font_size * 1.3, textColor=text_color),
                    )
                ]
            ],
            colWidths=[avail],
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), background),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        return table

    story.append(bar("PF / ESI COVERAGE", PDF_BRAND, colors.white, bold, 11))
    story.append(bar(_filter_text(data), PDF_BRAND_SOFT, INK, regular, 8))
    story.append(Spacer(1, 6))

    for position, sheet in enumerate(data["sheets"]):
        if position:
            story.append(PageBreak())
        story.append(bar(f"{sheet['title'].upper()} EMPLOYEES", PDF_BRAND_SOFT, PDF_BRAND, bold, 10))
        story.append(Spacer(1, 4))

        # the summary as cards, five to a line
        cards = [
            [
                Paragraph(escape(label), style("kl", fontSize=7, leading=9, textColor=MUTED)),
                Paragraph(
                    f"<b>{sheet['counts'][key]:,}</b>",
                    style("kv", fontName=bold, fontSize=12, leading=15, textColor=PDF_BRAND),
                ),
            ]
            for key, label, _indented in SUMMARY_LINES
        ]
        per_line = 5
        for start in range(0, len(cards), per_line):
            chunk = cards[start : start + per_line]
            chunk += [["", ""]] * (per_line - len(chunk))
            table = Table([chunk], colWidths=[avail / per_line] * per_line)
            table.setStyle(
                TableStyle(
                    [
                        ("BOX", (0, 0), (-1, -1), 0.5, PDF_GRID),
                        ("INNERGRID", (0, 0), (-1, -1), 0.5, PDF_GRID),
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 8))

        if not sheet["rows"]:
            story.append(
                Paragraph(
                    "No active employees match the selected filters.",
                    style("e", fontSize=10, leading=14, textColor=MUTED),
                )
            )
            continue
        grid = [[Paragraph(escape(label), s_head) for label, _w, _a in COLUMNS]]
        for record in sheet["rows"]:
            grid.append([Paragraph(escape(text), cell_styles[i]) for i, text in enumerate(_pdf_cells(record))])
        table = Table(grid, colWidths=widths, repeatRows=1, splitByRow=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), PDF_BRAND),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PDF_ZEBRA]),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.25, PDF_GRID),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(table)

    story.append(Spacer(1, 6))
    for note in NOTES:
        story.append(Paragraph(escape(note), s_note))
        story.append(Spacer(1, 2))

    buffer = io.BytesIO()
    doc = BaseDocTemplate(
        buffer,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=1.5 * cm,
        title="PF / ESI Coverage",
        author=co.name,
    )
    frame = Frame(
        margin,
        1.5 * cm,
        avail,
        page[1] - margin - 1.5 * cm,
        id="body",
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )
    doc.addPageTemplates([PageTemplate(id="coverage", frames=[frame])])
    stamp = f"Generated {ist_now().strftime('%d-%b-%Y %H:%M')} by {generated_by}"
    doc.build(story, canvasmaker=_canvas_maker(stamp, f"{co.name} - PF / ESI Coverage", regular))
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
#  The endpoint
# ─────────────────────────────────────────────────────────────────────────────


def _generated_by(request: Request) -> str:
    return hr_display_name(request)


@api_view(["GET"])
@require_hr
def pf_esi_coverage_export(request: Request) -> Response | HttpResponse:
    fmt = (request.query_params.get("format") or "").strip().lower()
    if fmt not in ("xlsx", "pdf"):
        return error_response("format must be xlsx or pdf")
    try:
        data = build_dataset(request)
    except CoverageParamError as problem:
        return error_response(str(problem))
    count = sum(len(sheet["rows"]) for sheet in data["sheets"])
    limit = XLSX_ROW_LIMIT if fmt == "xlsx" else PDF_ROW_LIMIT
    if count > limit:
        return error_response(
            f"This report has {count:,} employees, more than the {limit:,} a {'Excel' if fmt == 'xlsx' else 'PDF'} file "
            "holds. Choose one employee type or one branch and try again.",
            413,
        )
    by = _generated_by(request)
    content = build_xlsx(data, by) if fmt == "xlsx" else build_pdf(data, by)
    log_action(
        request,
        "export",
        "employees",
        description=f"PF / ESI coverage - {fmt.upper()} - {count} employees - type {data['type']}, {data['branch']}"[
            :480
        ],
    )
    response = HttpResponse(content, content_type=XLSX_MIME if fmt == "xlsx" else PDF_MIME)
    response["Content-Disposition"] = (
        f'attachment; filename="pf_esi_coverage_{data["type"]}_{ist_today().isoformat()}.{fmt}"'
    )
    response["Cache-Control"] = "no-store"
    return response


# `?format=` is the file type here: switch off DRF's renderer-suffix negotiation for this one view (module notes).
pf_esi_coverage_export.cls.content_negotiation_class = _IgnoreFormatParam
