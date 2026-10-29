"""
Report Center - HR Reports: hr-time-card (Time Card, Form 25 B).

One card per employee: In / Lunch Out / Lunch In / Out, the overtime window, hours worked, the shifts present and the OT
shifts of each day, with a printable one-page-per-employee PDF in the layout of the company's printed register.

Every date is in August 2026 (already over: "today" is patched to 2026-09-05). Shift 09:00-18:00, first half ends 13:30,
lunch 60 min. Calendar: Mon 3, Tue 4, Wed 5, Thu 6 (company holiday "Test Holiday", national), Fri 7, Sat 8, Sun 9, Mon 10.

Run via: python manage.py test api.tests_reporting_hr_time_card -v 2
"""

import io
from datetime import date, time
from decimal import Decimal
from unittest import mock

import pdfplumber
from django.apps import apps
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .attendance_final import _holiday_dates_for_month, compute_range_records
from .jwt_utils import sign_token
from .models import (
    AttendanceDayRecord,
    AttendanceLog,
    Branch,
    Department,
    Designation,
    Employee,
    EmployeeShiftAssignment,
    Holiday,
    HRUser,
    LeaveRequest,
    LeaveType,
    OvertimeRecord,
    PayrollSettings,
    ResignationRequest,
    Role,
    ShiftTemplate,
)
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import hr_reports_time_card as TC
from .reporting.export_xlsx import HEADER_ROW

RID = "hr-time-card"
TODAY = date(2026, 9, 5)
PERIOD = {"dateFrom": "2026-08-03", "dateTo": "2026-08-10"}
SRC = "biometric:adms:SN1"
LABELS = [
    "Employee Code",
    "Employee Name",
    "Department",
    "Designation",
    "SNo",
    "Date",
    "In Time",
    "Lunch Out",
    "Lunch In",
    "Out Time",
    "OT In",
    "OT Out",
    "Total Hrs",
    "Total OT Hrs",
    "Present",
    "OT",
]
KEYS = [
    "employeeCode",
    "employeeName",
    "department",
    "designation",
    "sno",
    "date",
    "inTime",
    "lunchOut",
    "lunchIn",
    "outTime",
    "otIn",
    "otOut",
    "totalHrs",
    "totalOtHrs",
    "present",
    "ot",
]


def d(day: int) -> date:
    return date(2026, 8, day)


def tm(text: str) -> time:
    return time(*[int(x) for x in text.split(":")])


def punch(emp, day: date, *times: str, source: str = SRC) -> None:
    for text in times:
        AttendanceLog.objects.create(employee=emp, date=day, punch_time=tm(text), punch_type="IN", source=source)


def hdr(user) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def seed(emp, first: date, last: date) -> None:
    """Persist the engine's verdicts the way opening the Attendance screens does (test set-up only)."""
    with mock.patch("api.attendance_final.ist_today", return_value=TODAY):
        compute_range_records(emp, first, last, PayrollSettings.get())


class _World(TestCase):
    @classmethod
    def setUpTestData(cls):
        _holiday_dates_for_month.cache_clear()
        ps = PayrollSettings.get()
        ps.attendance_mode = "simple"
        ps.morning_late_in_enabled = True
        ps.evening_early_out_enabled = True
        ps.slip_company_name = "UK TEXTILES - H.O"
        ps.slip_company_address = "TIRUPUR"
        ps.save()

        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.cutting = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.packing = Department.objects.create(name="PACKING", branch=cls.b2)
        cls.cutter = Designation.objects.create(title="Cutter")
        cls.shift = ShiftTemplate.objects.create(
            name="General",
            shift_type="staff",
            start_time=time(9, 0),
            end_time=time(18, 0),
            grace_period_minutes=15,
            first_half_end=time(13, 30),
            lunch_duration_minutes=60,
        )

        def mk(code, first, branch=None, dept=None, **kw):
            kw.setdefault("employment_type", "staff")
            return Employee.objects.create(
                employee_code=code,
                first_name=first,
                last_name="T",
                department=dept or cls.cutting,
                branch=branch or cls.b1,
                designation=cls.cutter,
                **kw,
            )

        cls.alice = mk("TC001", "Alice", father_name="Raman", join_date="2024-02-01")
        cls.bob = mk("TC002", "Bob", join_date="2024-03-01", salary_type="weekly")
        cls.carol = mk("TC003", "Carol", join_date="2024-03-01")
        cls.dan = mk("TC004", "Dan", join_date="2024-03-01")  # nobody has opened any of his days
        cls.eve = mk("TC005", "Eve", join_date="2026-08-05")  # joins on the 5th
        cls.pat = mk("TC006", "Pat", employment_type="production", join_date="2024-03-01")
        cls.zed = mk("TC007", "Zed", branch=cls.b2, dept=cls.packing, join_date="2024-03-01")
        for emp in (cls.alice, cls.bob, cls.carol, cls.dan, cls.eve, cls.zed):
            EmployeeShiftAssignment.objects.create(employee=emp, shift=cls.shift, effective_from=date(2020, 1, 1))
        Holiday.objects.create(name="Test Holiday", date=d(6), holiday_type="national")

        # Alice: four punches Monday, two Tuesday, nothing Wednesday (absent), a single punch Friday, Sunday off
        punch(cls.alice, d(3), "08:55", "13:00", "13:45", "18:05")
        punch(cls.alice, d(4), "09:00", "18:00")
        punch(cls.alice, d(7), "09:00")
        # Bob works late: Tuesday 90 min over (HR announced 0.5 shift of OT pay), Wednesday 3 hours over (no record)
        punch(cls.bob, d(4), "09:00", "19:30")
        punch(cls.bob, d(5), "09:00", "21:00")
        # Carol is on approved sick leave Wednesday
        punch(cls.carol, d(3), "09:00", "18:00")
        sl, _ = LeaveType.objects.get_or_create(code="SL", defaults={"name": "Sick Leave"})
        LeaveRequest.objects.create(
            employee=cls.carol,
            leave_type_ref=sl,
            type="sick",
            start_date="2026-08-05",
            end_date="2026-08-05",
            total_days=1,
            status="approved",
        )
        # Dan has punches but no stored verdict at all
        punch(cls.dan, d(3), "09:00", "18:00")
        punch(cls.eve, d(7), "09:00", "18:00")
        # Production: works Sundays like any day
        punch(cls.pat, d(9), "08:00", "17:00")
        punch(cls.zed, d(3), "09:00", "18:00")
        # HR typed one day by hand: a frozen manual verdict with its times
        AttendanceDayRecord.objects.create(
            employee=cls.alice,
            date=d(10),
            status="present",
            shifts_earned="1.00",
            source="manual",
            override_by="HR Anita",
            override_note="Missed punch",
            first_punch=time(9, 0),
            last_punch=time(18, 0),
        )

        for emp in (cls.alice, cls.bob, cls.carol, cls.eve, cls.pat, cls.zed):
            seed(emp, d(3), d(10))
        OvertimeRecord.objects.create(
            employee=cls.bob,
            date=d(4),
            shift_end_time=time(18, 0),
            last_punch_out=time(19, 30),
            ot_minutes=90,
            status="announced",
            compensation_type="pay",
            ot_type="shift",
            shifts=Decimal("0.50"),
        )
        ps.ot_threshold_minutes = 60
        ps.ot_detection_enabled = True
        ps.compensation_feature_enabled = True
        ps.save()
        AttendanceDayRecord.objects.filter(employee=cls.bob, date=d(5)).update(ot_extra_minutes=180)
        AttendanceDayRecord.objects.filter(employee=cls.bob, date=d(4)).update(ot_extra_minutes=90)

        cls.admin = HRUser.objects.create(username="tc_admin", password_hash="x", is_super_admin=True)
        cls.role_ok = Role.objects.create(name="tc_ok", permissions={"reports": "view", "attendance": "view"})
        cls.branch_user = HRUser.objects.create(username="tc_b1", password_hash="x", role=cls.role_ok, branch=cls.b1)
        cls.no_module = HRUser.objects.create(
            username="tc_none",
            password_hash="x",
            role=Role.objects.create(name="tc_none", permissions={"reports": "view"}),
        )
        cls.payroll_only = HRUser.objects.create(
            username="tc_pay",
            password_hash="x",
            role=Role.objects.create(name="tc_pay", permissions={"reports": "view", "payroll": "view"}),
        )

    @classmethod
    def tearDownClass(cls):
        _holiday_dates_for_month.cache_clear()
        super().tearDownClass()

    def setUp(self):
        p = mock.patch("api.reporting.filters.ist_today", return_value=TODAY)
        p.start()
        self.addCleanup(p.stop)
        _holiday_dates_for_month.cache_clear()

    # -- helpers ---------------------------------------------------------------
    def get(self, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/run/{RID}", {**PERIOD, **params}, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, fmt, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/export/{RID}", {"fmt": fmt, **PERIOD, **params}, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r

    def pdf_pages(self, **params) -> list[str]:
        with pdfplumber.open(io.BytesIO(self.export("pdf", **params).content)) as pdf:
            return [page.extract_text() or "" for page in pdf.pages]

    @staticmethod
    def data(body) -> list[dict]:
        return [r for r in body["rows"] if not r.get("_kind")]

    def row(self, body, emp, day: int) -> dict:
        for r in self.data(body):
            if r["employeeCode"] == emp.employee_code and r["date"].startswith(f"{day:02d}-08-2026"):
                return r
        raise AssertionError(f"no row for {emp.employee_code} {day}")

    def subtotal(self, body, emp) -> dict:
        return next(r for r in body["rows"] if r.get("_kind") and f"({emp.employee_code})" in r["employeeName"])


# ═════════════════════════════════════════════════════════════════════════════
# Registration, catalog and access
# ═════════════════════════════════════════════════════════════════════════════


class CatalogAndAccessTests(_World):
    def test_registered_in_the_hr_group_with_the_printed_columns(self):
        spec = registry.get_spec(RID)
        self.assertEqual(registry.LOAD_ERRORS, {})
        self.assertEqual((spec.title, spec.category, spec.modules), ("Time Card (Form 25 B)", "hr", ("attendance",)))
        self.assertEqual([c.label for c in spec.columns], LABELS)
        self.assertEqual([c.key for c in spec.columns], KEYS)
        self.assertTrue(set(spec.modules) <= set(all_module_keys()))
        self.assertIsNotNone(spec.pdf_builder)
        self.assertIsNone(spec.xlsx_builder)
        self.assertFalse(spec.super_admin_only)

    def test_the_catalog_offers_a_month_wide_period_the_scope_filters_and_the_lunch_switch(self):
        body = self.client.get("/api/reports/catalog", **hdr(self.admin)).json()
        report = next(r for r in body["reports"] if r["id"] == RID)
        self.assertEqual(report["category"], "hr")
        self.assertEqual(
            [(f["key"], f["kind"]) for f in report["filters"]],
            [
                ("dateRange", "dateRange"),
                ("branch", "branch"),
                ("department", "department"),
                ("designation", "designation"),
                ("employmentType", "employmentType"),
                ("employee", "employee"),
                ("employeeStatus", "employeeStatus"),
                ("deductLunch", "boolean"),
            ],
        )
        period = next(f for f in report["filters"] if f["key"] == "dateRange")
        self.assertEqual((period["label"], period["maxDays"]), ("Period", 31))
        self.assertFalse(next(f for f in report["filters"] if f["key"] == "deductLunch")["default"])

    def test_attendance_access_opens_it_and_nothing_else_does(self):
        for who in (self.branch_user, self.admin):
            self.assertTrue(self.data(self.get(user=who)))
            self.assertEqual(self.export("xlsx", user=who).status_code, 200)
        for who in (self.no_module, self.payroll_only):
            catalog = self.client.get("/api/reports/catalog", **hdr(who)).json()
            self.assertNotIn(RID, [r["id"] for r in catalog["reports"]], who.username)
            self.get(user=who, expect=403)
            self.export("pdf", user=who, expect=403)

    def test_it_needs_an_hr_login(self):
        self.assertEqual(self.client.get(f"/api/reports/run/{RID}").status_code, 401)


# ═════════════════════════════════════════════════════════════════════════════
# What the rows say
# ═════════════════════════════════════════════════════════════════════════════


class RowTests(_World):
    def test_one_row_per_employee_per_day_in_department_then_code_order_with_a_total_line_each(self):
        body = self.get()
        data = self.data(body)
        self.assertEqual(len(data), 7 * 8)  # an unscoped admin sees both units: seven employees x eight days
        self.assertEqual(
            list(dict.fromkeys(r["employeeCode"] for r in data)),
            ["TC001", "TC002", "TC003", "TC004", "TC005", "TC006", "TC007"],
        )  # CUTTING before PACKING, then by code
        totals = [r for r in body["rows"] if r.get("_kind") == "subtotal"]
        self.assertEqual(len(totals), 7)
        self.assertEqual(body["rows"][8]["_kind"], "subtotal")  # each employee's total line follows his eighth day

    def test_a_four_punch_day_shows_the_lunch_punches_and_the_hours_between_the_pairs(self):
        r = self.row(self.get(), self.alice, 3)
        self.assertEqual(
            (r["employeeCode"], r["employeeName"], r["department"], r["designation"], r["sno"]),
            ("TC001", "Alice T", "CUTTING", "Cutter", 1),
        )
        self.assertEqual(r["date"], "03-08-2026 Mon")
        self.assertEqual(
            (r["inTime"], r["lunchOut"], r["lunchIn"], r["outTime"]), ("08:55 AM", "01:00 PM", "01:45 PM", "06:05 PM")
        )
        self.assertEqual(r["totalHrs"], "08:25 Hrs")  # 08:55-13:00 + 13:45-18:05
        self.assertEqual((r["present"], r["ot"], r["totalOtHrs"]), (1.0, 0.0, "00:00 Hrs"))
        self.assertEqual((r["otIn"], r["otOut"]), (None, None))

    def test_a_two_punch_day_has_dashes_for_lunch_and_the_whole_span_as_hours(self):
        r = self.row(self.get(), self.alice, 4)
        self.assertEqual((r["inTime"], r["lunchOut"], r["lunchIn"], r["outTime"]), ("09:00 AM", None, None, "06:00 PM"))
        self.assertEqual(r["totalHrs"], "09:00 Hrs")
        self.assertEqual(r["present"], 1.0)

    def test_the_lunch_switch_takes_the_shifts_lunch_off_a_two_punch_day(self):
        self.assertEqual(self.row(self.get(deductLunch="true"), self.alice, 4)["totalHrs"], "08:00 Hrs")

    def test_a_single_punch_day_has_an_in_time_and_no_hours(self):
        r = self.row(self.get(), self.alice, 7)
        self.assertEqual((r["inTime"], r["outTime"], r["totalHrs"]), ("09:00 AM", None, None))

    def test_a_weekly_off_day_prints_wo_and_dashes(self):
        r = self.row(self.get(), self.alice, 9)  # Sunday
        self.assertEqual(r["date"], "09-08-2026 Sun")
        self.assertEqual((r["inTime"], r["outTime"], r["totalHrs"], r["totalOtHrs"]), ("WO", None, None, None))
        self.assertEqual((r["present"], r["ot"]), (0.0, 0.0))

    def test_an_absent_day_is_zero_present_with_no_marker(self):
        r = self.row(self.get(), self.alice, 5)
        self.assertEqual((r["inTime"], r["outTime"], r["present"]), (None, None, 0.0))

    def test_a_holiday_prints_its_code_by_the_holidays_type(self):
        self.assertEqual(self.row(self.get(), self.alice, 6)["inTime"], "NH")
        for htype, code in (("regional", "FH"), ("company", "OH")):
            Holiday.objects.filter(date=d(6)).update(holiday_type=htype)
            self.assertEqual(self.row(self.get(), self.alice, 6)["inTime"], code)

    def test_a_leave_day_prints_the_leave_types_code(self):
        r = self.row(self.get(), self.carol, 5)
        self.assertEqual((r["inTime"], r["present"]), ("SL", 0.0))

    def test_days_before_joining_are_not_applicable(self):
        body = self.get()
        self.assertEqual([self.row(body, self.eve, n)["inTime"] for n in (3, 4)], ["NA", "NA"])
        self.assertIsNone(self.row(body, self.eve, 3)["present"])
        self.assertEqual(self.row(body, self.eve, 7)["inTime"], "09:00 AM")

    def test_a_manual_day_shows_the_times_hr_typed(self):
        r = self.row(self.get(), self.alice, 10)
        self.assertEqual((r["inTime"], r["outTime"], r["present"]), ("09:00 AM", "06:00 PM", 1.0))

    def test_a_day_nobody_has_opened_is_not_processed_and_never_computed(self):
        body = self.get()
        self.assertEqual(self.row(body, self.dan, 4)["inTime"], "Not processed")
        self.assertIsNone(self.row(body, self.dan, 4)["present"])
        self.assertEqual(AttendanceDayRecord.objects.filter(employee=self.dan).count(), 0)  # a GET wrote nothing
        r = self.row(body, self.dan, 3)  # punches are facts: they are shown, the verdict is not invented
        self.assertEqual((r["inTime"], r["outTime"], r["present"]), ("09:00 AM", "06:00 PM", None))
        self.assertTrue(any("have no attendance record yet" in n and "have punches" in n for n in body["notes"]))

    def test_production_works_sundays_so_there_is_no_weekly_off(self):
        r = self.row(self.get(), self.pat, 9)
        self.assertEqual((r["inTime"], r["outTime"]), ("08:00 AM", "05:00 PM"))
        self.assertEqual(r["totalOtHrs"], "00:00 Hrs")  # overtime is a staff concept: nothing is added for production


class OvertimeTests(_World):
    def test_an_announced_pay_record_gives_its_window_minutes_and_shifts(self):
        r = self.row(self.get(), self.bob, 4)
        self.assertEqual((r["otIn"], r["otOut"]), ("06:00 PM", "07:30 PM"))
        self.assertEqual((r["totalOtHrs"], r["ot"], r["present"]), ("01:30 Hrs", 0.5, 1.0))

    def test_stored_minutes_past_shift_end_show_without_a_record_but_pay_no_shifts(self):
        r = self.row(self.get(), self.bob, 5)
        self.assertEqual((r["otIn"], r["otOut"], r["totalOtHrs"], r["ot"]), ("06:00 PM", "09:00 PM", "03:00 Hrs", 0.0))

    def test_a_rejected_record_counts_nothing(self):
        OvertimeRecord.objects.filter(employee=self.bob, date=d(4)).update(status="rejected")
        r = self.row(self.get(), self.bob, 4)
        self.assertEqual((r["otIn"], r["otOut"], r["totalOtHrs"], r["ot"]), (None, None, "00:00 Hrs", 0.0))

    def test_relaxation_and_detected_records_show_hours_but_pay_no_shifts(self):
        OvertimeRecord.objects.filter(employee=self.bob, date=d(4)).update(compensation_type="relaxation")
        self.assertEqual(self.row(self.get(), self.bob, 4)["ot"], 0.0)
        OvertimeRecord.objects.filter(employee=self.bob, date=d(4)).update(status="detected", compensation_type=None)
        r = self.row(self.get(), self.bob, 4)
        self.assertEqual((r["totalOtHrs"], r["ot"]), ("01:30 Hrs", 0.0))

    def test_a_legacy_pay_record_without_shifts_is_one_whole_day(self):
        OvertimeRecord.objects.filter(employee=self.bob, date=d(4)).update(ot_type="legacy", shifts=None)
        self.assertEqual(self.row(self.get(), self.bob, 4)["ot"], 1.0)

    def test_below_the_settings_threshold_is_not_overtime(self):
        r = self.row(self.get(), self.alice, 3)  # left 5 minutes after the shift end
        self.assertEqual(r["totalOtHrs"], "00:00 Hrs")

    def test_the_total_line_adds_hours_present_and_ot_shifts(self):
        sub = self.subtotal(self.get(), self.bob)
        self.assertEqual(sub["totalOtHrs"], "04:30 Hrs")
        self.assertEqual(sub["ot"], 0.5)
        self.assertEqual(sub["totalHrs"], "22:30 Hrs")  # 10:30 on Tuesday + 12:00 on Wednesday

    def test_the_notes_say_overtime_is_shift_based_not_double_wage(self):
        text = " ".join(self.get()["notes"])
        self.assertIn("not 'Double The Wage'", text)
        self.assertIn("the engine keeps no separate overtime punch times", text)


class TotalsTests(_World):
    def test_the_total_line_counts_present_shifts_and_the_other_day_kinds(self):
        body = self.get()
        sub = self.subtotal(body, self.alice)
        # Mon 1.00 + Tue 1.00 + Fri (single punch, stored by the engine) + Mon 10 manual 1.00; Wed absent, Sat absent
        self.assertEqual(sub["employeeName"], "Alice T (TC001) total")
        self.assertEqual(sub["date"], "Holidays 1.00  Weekly off 1.00  Absent 2.00  Leave 0.00")
        rows = [r for r in self.data(body) if r["employeeCode"] == "TC001"]
        self.assertEqual(sub["present"], round(sum(r["present"] or 0 for r in rows), 2))
        self.assertEqual(len(rows), 8)
        self.assertEqual([r["sno"] for r in rows], list(range(1, 9)))

    def test_leave_and_not_processed_days_are_counted_apart(self):
        self.assertIn("Leave 1.00", self.subtotal(self.get(), self.carol)["date"])
        sub = self.subtotal(self.get(), self.dan)
        self.assertEqual(sub["present"], 0.0)
        self.assertEqual(sub["date"], "Holidays 0.00  Weekly off 0.00  Absent 0.00  Leave 0.00")  # nothing is judged

    def test_the_cards_cover_the_summary(self):
        body = self.get()
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(cards["Employees"], 7)
        self.assertEqual(cards["Leave days"], 1)
        self.assertEqual(cards["Days not processed"], 8)  # Dan's eight days
        self.assertEqual(body["totals"]["present"], round(sum(r["present"] or 0 for r in self.data(body)), 2))


class FilterTests(_World):
    def test_a_range_wider_than_a_month_is_a_400_naming_the_field(self):
        r = self.client.get(
            f"/api/reports/run/{RID}", {"dateFrom": "2026-07-01", "dateTo": "2026-08-10"}, **hdr(self.admin)
        )
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json().get("field"), "dateTo")
        self.assertIn("max 31", r.json()["message"])

    def test_a_bad_or_reversed_date_is_a_400_never_a_500(self):
        for params in (
            {"dateFrom": "nope", "dateTo": "2026-08-10"},
            {"dateFrom": "2026-08-10", "dateTo": "2026-08-01"},
        ):
            self.assertEqual(self.client.get(f"/api/reports/run/{RID}", params, **hdr(self.admin)).status_code, 400)

    def test_scope_filters_narrow_the_cards(self):
        codes = lambda **p: list(dict.fromkeys(r["employeeCode"] for r in self.data(self.get(**p))))  # noqa: E731
        self.assertEqual(codes(employeeIds=f"{self.alice.id},{self.bob.id}"), ["TC001", "TC002"])
        self.assertEqual(codes(employmentType="production"), ["TC006"])
        self.assertEqual(codes(departmentIds=str(self.packing.id)), ["TC007"])

    def test_more_employees_than_a_run_can_print_is_a_400_asking_to_narrow(self):
        with mock.patch.object(TC, "MAX_EMPLOYEES", 3):
            r = self.client.get(f"/api/reports/run/{RID}", PERIOD, **hdr(self.admin))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["field"], "employee")
        self.assertIn("at most 3 employees", r.json()["message"])

    def test_an_employee_who_left_with_no_activity_is_not_listed_but_a_chosen_one_is(self):
        leo = Employee.objects.create(
            employee_code="TC099", first_name="Leo", last_name="T", status="inactive", branch=self.b1,
            department=self.cutting, employment_type="staff", join_date="2024-01-01",
        )  # fmt: skip
        self.assertNotIn("TC099", {r["employeeCode"] for r in self.data(self.get())})
        self.assertIn("TC099", {r["employeeCode"] for r in self.data(self.get(employeeIds=str(leo.id)))})

    def test_days_after_the_last_working_day_are_not_applicable(self):
        ResignationRequest.objects.create(
            employee=self.carol, status="approved", last_working_date=d(4), reason="moving"
        )
        body = self.get()
        self.assertEqual(self.row(body, self.carol, 5)["inTime"], "NA")
        self.assertEqual(self.row(body, self.carol, 3)["inTime"], "09:00 AM")


class BranchIsolationTests(_World):
    def test_a_branch_user_sees_only_their_own_branch_and_cannot_widen_it(self):
        mine = {r["employeeCode"] for r in self.data(self.get(user=self.branch_user))}
        self.assertNotIn("TC007", mine)
        self.assertIn("TC001", mine)
        for params in ({"branchIds": str(self.b2.id)}, {"employeeIds": str(self.zed.id)}):
            body = self.get(user=self.branch_user, **params)
            self.assertEqual(self.data(body), [], params)
            self.assertNotIn("Zed", str(body))

    def test_the_pdf_and_excel_of_a_branch_user_are_scoped_too(self):
        pages = self.pdf_pages(user=self.branch_user)
        self.assertNotIn("Zed", " ".join(pages))
        ws = load_workbook(io.BytesIO(self.export("xlsx", user=self.branch_user).content)).active
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertNotIn("Zed", cells)
        self.assertIn("Alice", cells)


# ═════════════════════════════════════════════════════════════════════════════
# Exports
# ═════════════════════════════════════════════════════════════════════════════


class ExportTests(_World):
    def test_excel_carries_the_same_columns_and_rows_as_the_screen(self):
        screen = self.get(employeeIds=f"{self.alice.id},{self.bob.id}")
        r = self.export("xlsx", employeeIds=f"{self.alice.id},{self.bob.id}")
        self.assertRegex(r["Content-Disposition"], r'filename="time_card_form_25_b_\d{8}_\d{4}\.xlsx"')
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[HEADER_ROW]][:16], LABELS)
        rows = [
            [c.value for c in row]
            for row in ws.iter_rows(min_row=HEADER_ROW + 1)
            if row[0].value in ("TC001", "TC002") and row[4].value
        ]
        self.assertEqual(len(rows), len(self.data(screen)))
        wo = next(row for row in rows if row[5] == "09-08-2026 Sun" and row[0] == "TC001")
        self.assertEqual(wo[6], "WO")
        four = next(row for row in rows if row[5] == "03-08-2026 Mon" and row[0] == "TC001")
        self.assertEqual(four[6:10], ["08:55 AM", "01:00 PM", "01:45 PM", "06:05 PM"])

    def test_the_pdf_is_one_page_per_employee_with_page_x_of_y(self):
        ids = ",".join(str(e.id) for e in (self.alice, self.bob, self.carol))
        r = self.export("pdf", employeeIds=ids)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertRegex(r["Content-Disposition"], r'\.pdf"$')
        pages = self.pdf_pages(employeeIds=ids)
        self.assertEqual(len(pages), 3)
        for i, text in enumerate(pages, 1):
            self.assertIn(f"Page {i} of 3", text)
        self.assertIn("Alice T", pages[0])
        self.assertIn("Bob T", pages[1])
        self.assertIn("Carol T", pages[2])

    def test_the_pdf_prints_the_header_the_table_the_footer_and_the_signature_lines(self):
        text = self.pdf_pages(employeeIds=str(self.alice.id))[0]
        for part in (
            "UK TEXTILES - H.O",
            "TIRUPUR",
            "Time Card",
            "Monthly Consolidated Report",
            "Form 25 B, see Rule 103",
            "License No :",
            "DOJ : 01-02-2024",
            "Normal Wages : Month",
            "OT : OT Shifts x Daily Wage",
            "Working Hours : 09:00 AM To 06:00 PM",
            "Sunday Holiday",
            "From 03-08-2026 To 10-08-2026",
            "Lunch Hours : 01:30 PM To 02:30 PM",
            "Employee Name : Alice T",
            "Employee Code : TC001",
            "Department : CUTTING",
            "Father Name : Raman",
            "Designation : Cutter",
            "SNo",
            "AttDate",
            "Shift Time",
            "OT Time",
            "InTime",
            "LunchOut",
            "LunchIn",
            "OutTime",
            "OT In",
            "OTOut",
            "03-08-2026 Mon 08:55 AM 01:00 PM 01:45 PM 06:05 PM",
            "08:25 Hrs",
            "09-08-2026 Sun WO",
            "WO - Weekly Off",
            "EL - Earned Leave",
            "NH - National Holiday",
            "FH - Festival Holiday",
            "OH - Other Holiday",
            "NA - Not Applicable",
            "Total Present : 3.50",
            "Total Holidays : 1.00",
            "Total Absent : 2.00",
            "Total OT Hours : 0.00",
            "Total weeklyOff : 1.00",
            "Employee signature",
            "Authorised Signatory",
        ):
            self.assertIn(part, text)

    def test_the_card_header_follows_the_employee_wages_ot_and_weekly_off_rules(self):
        bob = self.pdf_pages(employeeIds=str(self.bob.id))[0]
        self.assertIn("Normal Wages : Week", bob)  # a weekly-paid employee
        self.assertIn("Total OT Hours : 4.50", bob)
        pat = self.pdf_pages(employeeIds=str(self.pat.id))[0]
        self.assertIn("Normal Wages : Shift", pat)
        self.assertIn("OT : Extra shifts (in shift credit)", pat)
        self.assertIn("Sunday Working Day", pat)
        self.assertIn("Working Hours : -", pat)  # no shift assigned to the production employee
        self.assertIn("Lunch Hours : -", pat)

    def test_a_saturday_off_schedule_is_named_in_the_header(self):
        EmployeeShiftAssignment.objects.filter(employee=self.alice).update(saturday_off=True)
        self.assertIn("Sunday Holiday and Saturday Holiday", self.pdf_pages(employeeIds=str(self.alice.id))[0])

    def test_the_pdf_prints_nothing_but_a_notice_when_no_card_matches(self):
        text = self.pdf_pages(employeeIds="999999")[0]
        self.assertIn("No time-card rows match the selected filters.", text)

    def test_a_full_31_day_month_still_fits_one_page_per_employee(self):
        ids = f"{self.alice.id},{self.bob.id}"
        pages = self.pdf_pages(dateFrom="2026-08-01", dateTo="2026-08-31", employeeIds=ids)
        self.assertEqual(len(pages), 2)
        for text in pages:
            self.assertIn("31-08-2026", text)
            self.assertIn("Authorised Signatory", text)

    def test_the_export_is_audited(self):
        from .models import AuditLog

        before = AuditLog.objects.count()
        self.export("xlsx")
        entry = AuditLog.objects.order_by("-id").first()
        self.assertEqual(AuditLog.objects.count(), before + 1)
        self.assertIn("Time Card (Form 25 B)", entry.record_description)


# ═════════════════════════════════════════════════════════════════════════════
# Cost and safety
# ═════════════════════════════════════════════════════════════════════════════


class CostAndSafetyTests(_World):
    def _queries(self) -> int:
        self.get()  # warm-up (permission caches)
        with CaptureQueriesContext(connection) as q:
            self.get()
        return len(q)

    def test_the_query_count_does_not_grow_with_the_number_of_employees(self):
        before = self._queries()
        for i in range(15):
            emp = Employee.objects.create(
                employee_code=f"X{i:02d}", first_name=f"Xa{i}", last_name="Extra", employment_type="staff",
                branch=self.b1, department=self.cutting, designation=self.cutter, join_date="2024-01-01",
            )  # fmt: skip
            EmployeeShiftAssignment.objects.create(employee=emp, shift=self.shift, effective_from=date(2020, 1, 1))
            punch(emp, d(3), "09:00", "13:00", "14:00", "18:00")
            punch(emp, d(4), "09:00", "18:30")
            seed(emp, d(3), d(10))
        self.assertEqual(self._queries(), before)
        self.assertEqual(len({r["employeeCode"] for r in self.data(self.get())}), 7 + 15)

    def test_photo_and_password_columns_are_never_selected(self):
        with CaptureQueriesContext(connection) as q:
            self.get()
            self.export("xlsx")
        for query in q.captured_queries:
            self.assertNotIn("photo_url", query["sql"])
            if '"hr_users"' not in query["sql"]:
                self.assertNotIn("password_hash", query["sql"])

    def test_running_and_exporting_never_writes(self):
        def counts():
            return {
                m.__name__: m.objects.count()
                for m in apps.get_app_config("api").get_models()
                if m.__name__ != "AuditLog"  # an export is supposed to leave an audit entry
            }

        before = counts()
        self.get()
        self.export("xlsx")
        self.export("pdf")
        self.assertEqual(counts(), before)

    def test_a_row_over_the_limit_is_refused_for_export_not_silently_cut(self):
        with mock.patch("api.reporting.runner.XLSX_ROW_LIMIT", 10):
            r = self.client.get(f"/api/reports/export/{RID}", {"fmt": "xlsx", **PERIOD}, **hdr(self.admin))
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.json()["error"], "too_many_rows")

    def test_an_empty_roster_is_an_empty_report_not_an_error(self):
        body = self.get(employeeIds="999999")
        self.assertEqual(body["rows"], [])
        self.assertIsNone(body["totals"])
        self.assertEqual(self.export("xlsx", employeeIds="999999").status_code, 200)


class FormatterTests(_World):
    def test_twelve_hour_clock_and_hours_text(self):
        self.assertEqual(TC._t12(time(0, 5)), "12:05 AM")
        self.assertEqual(TC._t12(time(12, 0)), "12:00 PM")
        self.assertEqual(TC._t12(time(23, 59)), "11:59 PM")
        self.assertEqual(TC._hrs(0), "00:00 Hrs")
        self.assertEqual(TC._hrs(661), "11:01 Hrs")
        self.assertIsNone(TC._hrs(None))
        self.assertIsNone(TC._t12(None))

    def test_a_switched_off_overtime_feature_is_said_on_the_card_and_in_the_notes(self):
        ps = PayrollSettings.get()
        ps.ot_detection_enabled = False
        ps.save()
        self.assertIn("OT : Not detected (switched off)", self.pdf_pages(employeeIds=str(self.alice.id))[0])
        self.assertTrue(any("Overtime detection is switched off" in n for n in self.get()["notes"]))
