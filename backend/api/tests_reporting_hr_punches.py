"""
Report Center - HR Reports (category ``hr``): the two printed punch registers.

* hr-daily-movement-register   - "Daily Movement Register": every punch of the day (Time1..Time12), by department.
* hr-punching-register         - "Punching Register (Employees Present List)": shift in-time, first punch, late hours.

"Today" is pinned to 2026-09-29 (a Tuesday) so every expected value is hand-derived and deterministic.

Run via: python manage.py test api.tests_reporting_hr_punches -v 2
"""

import io
from datetime import date, time
from unittest.mock import patch

from django.apps import apps
from django.db import connection
from django.test import SimpleTestCase, TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .models import (
    AttendanceDayRecord,
    AttendanceLog,
    Branch,
    Department,
    Designation,
    Employee,
    EmployeeShiftAssignment,
    HRUser,
    PayrollSettings,
    Role,
    ShiftTemplate,
)
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import hr_reports_punches as H
from .reporting.export_xlsx import HEADER_ROW

TODAY = date(2026, 9, 29)  # Tuesday
D29 = TODAY
D28 = date(2026, 9, 28)  # Monday
D30 = date(2026, 9, 30)  # Wednesday: in the future, nothing can be listed
MOVE = "hr-daily-movement-register"
PRESENT = "hr-punching-register"

MOVE_LABELS = ["SNo", "TicketNo", "EmpName", "Date"] + [f"Time{n}" for n in range(1, 13)]
MOVE_KEYS = ["sno", "ticketNo", "empName", "date"] + [f"time{n}" for n in range(1, 13)]
PRESENT_LABELS = [
    "SNo",
    "Ticket No",
    "Employee Name",
    "Designation",
    "Nature of Emp",
    "Date",
    "Shift InTime",
    "Act InTime",
    "Late Hours",
]
PRESENT_KEYS = ["sno", "ticketNo", "empName", "designation", "nature", "date", "shiftIn", "actIn", "lateHours"]


def hdr(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def at(text: str) -> time:
    return time(*[int(x) for x in text.split(":")])


def punch(emp, day, *times, kind="IN", source="biometric:adms:SN1"):
    for text in times:
        AttendanceLog.objects.create(employee=emp, date=day, punch_time=at(text), punch_type=kind, source=source)


def record(emp, day, status="present", first=None, **kw):
    return AttendanceDayRecord.objects.create(
        employee=emp, date=day, status=status, first_punch=at(first) if first else None, **kw
    )


def data_rows(body) -> list[dict]:
    return [r for r in body["rows"] if not r.get("_kind")]


def headings(body) -> list[str]:
    return [r["empName"] for r in body["rows"] if r.get("_kind") == "subtotal"]


def cards(body) -> dict:
    return {c["label"]: c["value"] for c in body["summary"]}


def times_of(row) -> list:
    return [row[f"time{n}"] for n in range(1, 13) if row.get(f"time{n}")]


class _World(TestCase):
    """Staff and production people in two units, with every shape of punch day the registers meet."""

    RID = MOVE

    @classmethod
    def setUpTestData(cls):
        PayrollSettings.get()  # the app's singleton settings row exists on every installed system (created at first use)
        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.d_acc = Department.objects.create(name="ACCOUNTS", branch=cls.b1)
        cls.d_adm = Department.objects.create(name="ADMIN", branch=cls.b1)
        cls.d_cad = Department.objects.create(name="CAD", branch=cls.b1)
        cls.d_pk = Department.objects.create(name="PACKING", branch=cls.b2)
        cls.des_exec = Designation.objects.create(title="ACCOUNTS EXECUTIVE", department=cls.d_acc, level="junior")
        cls.des_off = Designation.objects.create(title="OFFICE ASST", department=cls.d_adm, level="junior")

        def mk(code, first, last, dept, branch=None, **kw):
            kw.setdefault("employment_type", "staff")
            kw.setdefault("designation", cls.des_off)
            return Employee.objects.create(
                employee_code=code, first_name=first, last_name=last, department=dept,
                branch=branch if branch is not None else (dept.branch if dept else None), **kw,
            )  # fmt: skip

        cls.acc = mk("2316", "Meenatchi", "T", cls.d_acc, designation=cls.des_exec)
        cls.adm1 = mk("189", "Murugan", "K", cls.d_adm)
        cls.adm2 = mk("226", "Krishnan", "Nair", cls.d_adm)
        cls.adm3 = mk("2234", "Murugesan", "K", cls.d_adm)  # never punches
        cls.adm4 = mk("2503", "Narayanan", "Kuppayi", cls.d_adm)  # one punch only
        cls.prod = mk("300", "Pradeep", "Raj", cls.d_adm, employment_type="production")
        cls.eve = mk("EVE", "Eve", "Night", cls.d_cad)
        cls.nodept = mk("999", "Zed", "Lone", None)
        cls.pk = mk("P5", "Pooja", "Devi", cls.d_pk)
        cls.old = mk("OLD1", "Old", "Hand", cls.d_adm, status="inactive")

        # ── punches ─────────────────────────────────────────────────────────
        punch(cls.acc, D29, "13:41:00", "08:58:10", "19:01:00", "14:37:00")  # entered out of order
        punch(cls.adm1, D29, "09:00:00")
        punch(cls.adm1, D29, "19:13:00", kind="OUT")
        # thirteen punches: 08:00, 08:30 ... 14:00
        punch(cls.adm2, D29, *[f"{8 + n // 2:02d}:{30 * (n % 2):02d}:00" for n in range(13)])
        punch(cls.adm4, D29, "09:20:00")
        punch(cls.prod, D29, "08:00:00", "17:00:00")
        punch(cls.nodept, D29, "08:30:00", "17:30:00")
        punch(cls.pk, D29, "09:00:00", "18:00:00")
        punch(cls.old, D29, "09:00:00", "18:00:00")
        # yesterday: Meenatchi two punches; Eve a normal Monday whose exit (01:05) the device stamped on Tuesday
        punch(cls.acc, D28, "09:05:00", "18:10:00")
        punch(cls.eve, D28, "09:00:00", "13:00:00", "13:45:00")
        punch(cls.eve, D29, "01:05:00", kind="OUT")
        punch(cls.eve, D29, "09:00:00", "18:00:00")
        punch(cls.adm3, D28, "09:00:00")  # Murugesan punched yesterday only

        # ── shifts ──────────────────────────────────────────────────────────
        def shift(name, start, end, branch=None):
            return ShiftTemplate.objects.create(
                name=name, shift_type="staff", start_time=start, end_time=end, branch=branch
            )

        cls.general = shift("General", time(9, 0), time(18, 0), cls.b1)
        cls.early = shift("Early", time(7, 30), time(16, 30), cls.b1)

        def assign(emp, sh, start, end=None, **kw):
            return EmployeeShiftAssignment.objects.create(
                employee=emp, shift=sh, effective_from=start, effective_to=end, **kw
            )

        assign(cls.acc, cls.general, date(2026, 1, 1))
        assign(cls.adm1, cls.general, date(2026, 1, 1))
        assign(cls.adm1, cls.early, date(2026, 9, 29))  # starts today: the later start wins from today on
        assign(cls.adm2, cls.general, date(2026, 1, 1), date(2026, 9, 28))  # ended yesterday: no shift today
        assign(cls.adm4, cls.general, date(2026, 1, 1), custom_start_time=time(10, 0))
        assign(cls.eve, cls.general, date(2026, 1, 1))
        assign(cls.pk, cls.general, date(2026, 1, 1))

        # ── users ───────────────────────────────────────────────────────────
        cls.admin = HRUser.objects.create(username="hr_admin", password_hash="x", is_super_admin=True)

        def user(name, perms, **kw):
            role = Role.objects.create(name=name, permissions={"reports": "view", **perms})
            return HRUser.objects.create(username=name, password_hash="x", role=role, **kw)

        cls.att_user = user("hr_att", {"attendance": "view"})
        cls.salary_user = user("hr_sal", {"salary": "view", "payroll": "view"})
        cls.branch_user = user("hr_b1", {"attendance": "view"}, branch=cls.b1)

    def setUp(self):
        p = patch("api.reporting.filters.ist_today", return_value=TODAY)
        p.start()
        self.addCleanup(p.stop)

    def get(self, user=None, expect=200, rid=None, **params):
        r = self.client.get(f"/api/reports/run/{rid or self.RID}", params, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, fmt, user=None, expect=200, rid=None, **params):
        r = self.client.get(f"/api/reports/export/{rid or self.RID}", {"fmt": fmt, **params}, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r

    def codes(self, body) -> list[str]:
        return [r["ticketNo"] for r in data_rows(body)]


# ═════════════════════════════════════════════════════════════════════════════
# Pure helpers
# ═════════════════════════════════════════════════════════════════════════════


class ClockTextTests(SimpleTestCase):
    def test_twelve_hour_clock_without_a_leading_zero(self):
        self.assertEqual(H._clock12(time(8, 58), pad=False), "8:58 AM")
        self.assertEqual(H._clock12(time(13, 41, 30), pad=False), "1:41 PM")
        self.assertEqual(H._clock12(time(0, 5), pad=False), "12:05 AM")
        self.assertEqual(H._clock12(time(12, 0), pad=False), "12:00 PM")
        self.assertEqual(H._clock12(time(23, 59), pad=False), "11:59 PM")

    def test_padded_clock_for_the_punching_register(self):
        self.assertEqual(H._clock12(time(9, 5), pad=True), "09:05 AM")
        self.assertEqual(H._clock12(time(18, 0), pad=True), "06:00 PM")

    def test_direction_letter_is_the_stored_one(self):
        self.assertEqual(H._direction_letter("IN"), "I")
        self.assertEqual(H._direction_letter("OUT"), "O")
        self.assertEqual(H._direction_letter(" out "), "O")
        self.assertEqual(H._direction_letter(None), "I")

    def test_late_hours_as_hh_mm(self):
        self.assertEqual(H._hhmm(40), "00:40")
        self.assertEqual(H._hhmm(95), "01:35")
        self.assertEqual(H._hhmm(600), "10:00")


# ═════════════════════════════════════════════════════════════════════════════
# Registration, catalog and access (both reports)
# ═════════════════════════════════════════════════════════════════════════════


class CatalogAndAccessTests(_World):
    def test_both_are_registered_in_the_hr_group_with_the_printouts_column_labels(self):
        self.assertEqual(registry.LOAD_ERRORS, {})
        specs = {s.id: s for s in registry.all_specs()}
        move, pres = specs[MOVE], specs[PRESENT]
        self.assertEqual((move.title, move.category, move.modules), ("Daily Movement Register", "hr", ("attendance",)))
        self.assertEqual(
            (pres.title, pres.category, pres.modules),
            ("Punching Register (Employees Present List)", "hr", ("attendance",)),
        )
        self.assertEqual([c.label for c in move.columns], MOVE_LABELS)
        self.assertEqual([c.key for c in move.columns], MOVE_KEYS)
        self.assertEqual([c.label for c in pres.columns], PRESENT_LABELS)
        self.assertEqual([c.key for c in pres.columns], PRESENT_KEYS)
        for spec in (move, pres):
            self.assertFalse(spec.super_admin_only)
            self.assertTrue(set(spec.modules) <= set(all_module_keys()), spec.modules)

    def test_the_catalog_offers_a_date_range_defaulting_to_today_and_the_scope_filters(self):
        body = self.client.get("/api/reports/catalog", **hdr(self.admin)).json()
        for rid in (MOVE, PRESENT):
            report = next(r for r in body["reports"] if r["id"] == rid)
            self.assertEqual(report["category"], "hr")
            self.assertEqual(
                [(f["key"], f["kind"]) for f in report["filters"]],
                [
                    ("dateRange", "dateRange"),
                    ("branch", "branch"),
                    ("department", "department"),
                    ("designation", "designation"),
                    ("employmentType", "employmentType"),
                    ("employee", "employee"),
                    ("employeeStatus", "employeeStatus"),
                ],
            )
            date_filter = next(f for f in report["filters"] if f["key"] == "dateRange")
            self.assertEqual(date_filter["maxDays"], 31)
            self.assertEqual(date_filter["default"]["dateFrom"], date_filter["default"]["dateTo"])  # opens on one day
            status = next(f for f in report["filters"] if f["key"] == "employeeStatus")
            self.assertEqual(status["default"], "active")

    def test_attendance_access_opens_both_and_salary_access_alone_does_not(self):
        for rid in (MOVE, PRESENT):
            for who in (self.att_user, self.admin):
                self.assertTrue(data_rows(self.get(user=who, rid=rid)), (rid, who.username))
                self.assertEqual(self.export("xlsx", user=who, rid=rid).status_code, 200)
            catalog = self.client.get("/api/reports/catalog", **hdr(self.salary_user)).json()
            self.assertNotIn(rid, [r["id"] for r in catalog["reports"]])
            self.get(user=self.salary_user, rid=rid, expect=403)
            self.export("pdf", user=self.salary_user, rid=rid, expect=403)

    def test_both_need_an_hr_login(self):
        for rid in (MOVE, PRESENT):
            self.assertEqual(self.client.get(f"/api/reports/run/{rid}").status_code, 401)

    def test_a_range_wider_than_a_month_or_a_bad_date_is_a_400_never_a_500(self):
        for rid in (MOVE, PRESENT):
            wide = self.client.get(
                f"/api/reports/run/{rid}", {"dateFrom": "2026-08-01", "dateTo": "2026-09-29"}, **hdr(self.admin)
            )
            self.assertEqual(wide.status_code, 400)
            self.assertEqual(wide.json().get("field"), "dateTo")
            self.assertIn("max 31", wide.json()["message"])
            for params in (
                {"dateFrom": "nope", "dateTo": "2026-09-29"},
                {"dateFrom": "2026-09-29", "dateTo": "2026-09-01"},
            ):
                r = self.client.get(f"/api/reports/run/{rid}", params, **hdr(self.admin))
                self.assertEqual(r.status_code, 400, params)


# ═════════════════════════════════════════════════════════════════════════════
# Daily Movement Register
# ═════════════════════════════════════════════════════════════════════════════


class MovementRowTests(_World):
    RID = MOVE

    def test_one_line_per_employee_listed_by_department_then_ticket_number_with_continuous_sno(self):
        body = self.get()
        self.assertEqual(
            self.codes(body), ["2316", "189", "226", "300", "2503", "EVE", "P5", "999"]
        )  # natural ticket order (189 < 226 < 2503), departments A-Z, no-department last
        self.assertEqual([r["sno"] for r in data_rows(body)], list(range(1, 9)))
        self.assertEqual(headings(body), ["ACCOUNTS", "ADMIN", "CAD", "PACKING", "Unassigned"])
        shape = [("H" if r.get("_kind") else r["ticketNo"]) for r in body["rows"]]
        self.assertEqual(
            shape, ["H", "2316", "H", "189", "226", "300", "2503", "H", "EVE", "H", "P5", "H", "999"]
        )  # a heading line before each department's people

    def test_golden_four_punch_day_in_time_order_with_the_printouts_text(self):
        row = data_rows(self.get())[0]
        self.assertEqual(
            row,
            {
                "sno": 1,
                "ticketNo": "2316",
                "empName": "Meenatchi T",
                "date": "29/09/26 Tue",
                "time1": "8:58 AM-I",  # entered third, listed first
                "time2": "1:41 PM-I",
                "time3": "2:37 PM-I",
                "time4": "7:01 PM-I",
                **{f"time{n}": None for n in range(5, 13)},
            },
        )

    def test_a_two_punch_day_leaves_the_rest_blank_and_the_suffix_is_the_stored_direction(self):
        rows = {r["ticketNo"]: r for r in data_rows(self.get())}
        self.assertEqual(times_of(rows["189"]), ["9:00 AM-I", "7:13 PM-O"])  # stored IN then OUT
        self.assertEqual(times_of(rows["2503"]), ["9:20 AM-I"])  # a single punch is a line too
        self.assertIsNone(rows["189"]["time3"])
        self.assertEqual(times_of(rows["300"]), ["8:00 AM-I", "5:00 PM-I"])  # production

    def test_the_suffix_is_kept_as_stored_even_when_the_device_flags_everything_in(self):
        # the first punch stored OUT and the second IN: not 'fixed' by position
        emp = Employee.objects.create(employee_code="7", first_name="Flip", last_name="Flop", branch=self.b1,
                                      department=self.d_acc, employment_type="staff")  # fmt: skip
        punch(emp, D29, "09:00:00", kind="OUT")
        punch(emp, D29, "17:00:00", kind="IN")
        row = next(r for r in data_rows(self.get()) if r["ticketNo"] == "7")
        self.assertEqual(times_of(row), ["9:00 AM-O", "5:00 PM-I"])

    def test_more_than_twelve_punches_show_twelve_and_the_count_of_the_rest_in_the_last_cell(self):
        body = self.get()
        row = next(r for r in data_rows(body) if r["ticketNo"] == "226")
        self.assertEqual(row["time1"], "8:00 AM-I")
        self.assertEqual(row["time11"], "1:00 PM-I")
        self.assertEqual(row["time12"], "1:30 PM-I (+1 more)")  # the 13th (2:00 PM) is not dropped silently
        self.assertEqual(cards(body)["Days with more than 12 punches"], 1)
        self.assertTrue(any("more than 12 punches" in n for n in body["notes"]), body["notes"])

    def test_exactly_twelve_punches_have_no_more_marker(self):
        AttendanceLog.objects.filter(employee=self.adm2, punch_time=at("14:00:00")).delete()
        row = next(r for r in data_rows(self.get()) if r["ticketNo"] == "226")
        self.assertEqual(row["time12"], "1:30 PM-I")
        self.assertEqual(cards(self.get())["Days with more than 12 punches"], 0)

    def test_an_exit_after_midnight_is_on_the_day_it_closes_and_not_on_the_next(self):
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-29", departmentIds=str(self.d_cad.id))
        rows = {r["date"]: r for r in data_rows(body)}
        self.assertEqual(
            times_of(rows["28/09/26 Mon"]), ["9:00 AM-I", "1:00 PM-I", "1:45 PM-I", "1:05 AM-O (+1)"]
        )  # Monday's forgotten exit, stamped Tuesday 01:05; its stored OUT is kept
        self.assertEqual(times_of(rows["29/09/26 Tue"]), ["9:00 AM-I", "6:00 PM-I"])  # not repeated as Tuesday's first
        self.assertTrue(any("after midnight" in n for n in body["notes"]), body["notes"])
        self.assertTrue(any("day the attendance engine places it" in n for n in body["notes"]), body["notes"])

    def test_an_employee_with_no_punch_on_the_date_is_not_listed(self):
        body = self.get()
        self.assertNotIn("2234", self.codes(body))  # active, never punched today
        self.assertNotIn("2234", self.codes(self.get(dateFrom="2026-09-27", dateTo="2026-09-27")))
        self.assertIn("2234", self.codes(self.get(dateFrom="2026-09-28", dateTo="2026-09-28")))  # punched on the 28th

    def test_a_date_range_lists_each_employees_dates_together_in_date_order(self):
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-29", departmentIds=f"{self.d_acc.id},{self.d_adm.id}")
        shown = [(r["ticketNo"], r["date"], r["sno"]) for r in data_rows(body)]
        self.assertEqual(
            shown[:5],
            [
                ("2316", "28/09/26 Mon", 1),
                ("2316", "29/09/26 Tue", 2),
                ("189", "29/09/26 Tue", 3),
                ("226", "29/09/26 Tue", 4),
                ("300", "29/09/26 Tue", 5),
            ],
        )
        self.assertEqual(headings(body), ["ACCOUNTS", "ADMIN"])
        self.assertIn("one line per employee per date", " ".join(body["notes"]))

    def test_a_future_or_empty_date_is_an_empty_report_not_an_error(self):
        for day in ("2026-09-30", "2026-09-27"):
            body = self.get(dateFrom=day, dateTo=day)
            self.assertEqual(body["rows"], [])
            self.assertEqual(cards(body)["Punches"], 0)

    def test_cards_and_notes(self):
        body = self.get()
        self.assertEqual(
            cards(body),
            {
                "Employees": 8,
                "Employee-days listed": 8,
                "Punches": 4 + 2 + 13 + 2 + 1 + 2 + 2 + 2,
                "Departments": 5,
                "Days with more than 12 punches": 1,
            },
        )
        text = " ".join(body["notes"])
        self.assertIn("One line per employee for 29-Sep-2026", text)
        self.assertIn("direction stored on the punch record", text)  # I / O is as stored, and says so
        self.assertIn("not a reliable In / Out", text)
        self.assertIn("not listed", text)


class MovementFilterTests(_World):
    RID = MOVE

    def test_status_defaults_to_active_and_the_leaver_comes_back_with_all(self):
        self.assertNotIn("OLD1", self.codes(self.get()))
        self.assertIn("OLD1", self.codes(self.get(employeeStatus="all")))
        self.assertEqual(self.codes(self.get(employeeStatus="inactive")), ["OLD1"])

    def test_employee_type_department_designation_branch_and_employee_filters(self):
        self.assertEqual(self.codes(self.get(employmentType="production")), ["300"])
        self.assertNotIn("300", self.codes(self.get(employmentType="staff")))
        self.assertEqual(self.codes(self.get(departmentIds=str(self.d_pk.id))), ["P5"])
        self.assertEqual(self.codes(self.get(designationIds=str(self.des_exec.id))), ["2316"])
        self.assertEqual(self.codes(self.get(branchIds=str(self.b2.id))), ["P5"])
        self.assertEqual(self.codes(self.get(employeeIds=f"{self.acc.id},{self.adm1.id}")), ["2316", "189"])

    def test_the_applied_filters_are_echoed(self):
        applied = {f["label"]: f["value"] for f in self.get(departmentIds=str(self.d_pk.id))["filters"]}
        self.assertEqual(applied["Department"], "PACKING")
        self.assertIn("29-Sep-2026", applied["Date"])


class MovementIsolationAndSafetyTests(_World):
    RID = MOVE

    def test_a_branch_user_sees_only_their_branch_and_cannot_widen_it(self):
        mine = self.get(user=self.branch_user)
        self.assertEqual(self.codes(mine), ["2316", "189", "226", "300", "2503", "EVE"])
        self.assertNotIn("PACKING", headings(mine))
        self.assertNotIn("Pooja", str(mine))
        for params in ({"branchIds": str(self.b2.id)}, {"employeeIds": str(self.pk.id)}):
            self.assertEqual(data_rows(self.get(user=self.branch_user, **params)), [], params)

    def test_a_branch_users_exports_hold_nothing_of_another_branch(self):
        ws = load_workbook(io.BytesIO(self.export("xlsx", user=self.branch_user).content)).active
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertNotIn("Pooja", cells)
        self.assertIn("Meenatchi T", cells)

    def test_running_and_exporting_never_writes(self):
        def counts():
            return {
                m.__name__: m.objects.count()
                for m in apps.get_app_config("api").get_models()
                if m.__name__ != "AuditLog"  # an export is supposed to leave an audit entry
            }

        before = counts()
        for rid in (MOVE, PRESENT):
            self.get(rid=rid, dateFrom="2026-09-28", dateTo="2026-09-29")
            self.export("xlsx", rid=rid)
            self.export("pdf", rid=rid)
        self.assertEqual(counts(), before)

    def test_the_query_count_does_not_grow_with_the_number_of_employees_or_punches(self):
        def queries(rid) -> int:
            self.get(rid=rid)  # warm-up (permission caches)
            with CaptureQueriesContext(connection) as q:
                self.get(rid=rid, dateFrom="2026-09-28", dateTo="2026-09-29")
            return len(q)

        before = {rid: queries(rid) for rid in (MOVE, PRESENT)}
        for i in range(20):
            emp = Employee.objects.create(
                employee_code=f"X{i}", first_name=f"Xa{i}", last_name="Extra", employment_type="staff",
                branch=self.b1, department=self.d_adm, designation=self.des_off,
            )  # fmt: skip
            punch(emp, D29, "08:00:00", "12:00:00", "13:00:00", "17:00:00")
            punch(emp, D28, "08:00:00", "17:00:00")
            EmployeeShiftAssignment.objects.create(employee=emp, shift=self.general, effective_from=date(2026, 1, 1))
            record(emp, D29, first="08:00", is_late=bool(i % 2))
        self.assertEqual({rid: queries(rid) for rid in (MOVE, PRESENT)}, before)
        self.assertEqual(len(data_rows(self.get(dateFrom="2026-09-28", dateTo="2026-09-29"))), 11 + 40)

    def test_photo_and_password_columns_are_never_selected(self):
        with CaptureQueriesContext(connection) as q:
            for rid in (MOVE, PRESENT):
                self.get(rid=rid)
                self.export("xlsx", rid=rid)
        for query in q.captured_queries:
            self.assertNotIn("photo_url", query["sql"])
            if '"hr_users"' not in query["sql"]:
                self.assertNotIn("password_hash", query["sql"])

    def test_a_row_over_the_limit_is_refused_for_export_not_silently_cut(self):
        for rid in (MOVE, PRESENT):
            with patch("api.reporting.runner.XLSX_ROW_LIMIT", 2):
                r = self.client.get(f"/api/reports/export/{rid}", {"fmt": "xlsx"}, **hdr(self.admin))
            self.assertEqual(r.status_code, 413, rid)
            self.assertEqual(r.json()["error"], "too_many_rows")

    def test_the_screen_limit_flags_a_cut_list_and_heading_lines_do_not_count(self):
        params = {"fmt": "xlsx", "departmentIds": str(self.d_adm.id)}  # four people and one heading line
        for rid in (MOVE, PRESENT):
            with patch("api.reporting.runner.XLSX_ROW_LIMIT", 4):
                ok = self.client.get(f"/api/reports/export/{rid}", params, **hdr(self.admin))
            self.assertEqual(ok.status_code, 200, rid)  # the heading line is furniture: it does not count
            with patch("api.reporting.runner.XLSX_ROW_LIMIT", 3):
                over = self.client.get(f"/api/reports/export/{rid}", params, **hdr(self.admin))
            self.assertEqual(over.status_code, 413, rid)


class MovementExportTests(_World):
    RID = MOVE

    def test_excel_carries_the_same_columns_and_rows_as_the_screen(self):
        params = {"dateFrom": "2026-09-28", "dateTo": "2026-09-29"}
        screen = self.get(**params)
        r = self.export("xlsx", **params)
        self.assertRegex(r["Content-Disposition"], r'filename="daily_movement_register_\d{8}_\d{4}\.xlsx"')
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[HEADER_ROW]][: len(MOVE_LABELS)], MOVE_LABELS)
        self.assertGreater(len(screen["rows"]), 10)
        for i, row in enumerate(screen["rows"]):
            got = [ws.cell(row=HEADER_ROW + 1 + i, column=n + 1).value for n in range(len(MOVE_KEYS))]
            want = [row.get(k) for k in MOVE_KEYS]
            self.assertEqual(
                [str(g if g is not None else "") for g in got], [str(w if w is not None else "") for w in want]
            )
        cells = " ".join(str(c.value) for rw in ws.iter_rows() for c in rw if c.value is not None)
        self.assertIn("ACCOUNTS", cells)
        self.assertIn("28-Sep-2026 to 29-Sep-2026", cells)  # the filter echo in the sheet header

    def test_pdf_is_a_real_pdf(self):
        r = self.export("pdf", dateFrom="2026-09-28", dateTo="2026-09-29")
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertRegex(r["Content-Disposition"], r'\.pdf"$')


# ═════════════════════════════════════════════════════════════════════════════
# Punching Register (Employees Present List)
# ═════════════════════════════════════════════════════════════════════════════


class _PresentWorld(_World):
    RID = PRESENT

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # stored verdicts of the 29th (what opening the Attendance screens persists)
        record(cls.acc, D29, first="08:58", total_punches=4)  # on time
        record(cls.adm1, D29, first="09:40", is_late=True, total_punches=2)  # late: general 09:00 / early 07:30 today
        record(cls.adm2, D29, first="08:00", total_punches=13)
        record(cls.adm4, D29, first="09:20", total_punches=1)  # its shift starts 10:00 by a custom time: on time
        record(cls.eve, D29, first="09:00", total_punches=2)
        record(cls.eve, D28, first="09:00", total_punches=4)
        record(cls.pk, D29, first="09:00", total_punches=2)
        # present with no punch (an HR entry) and a paid casual-leave day: only the first belongs on the list
        cls.hr_present = Employee.objects.create(
            employee_code="500", first_name="Hira", last_name="Entry", employment_type="staff",
            branch=cls.b1, department=cls.d_acc,
        )  # fmt: skip
        record(cls.hr_present, D29, first="09:00", source="manual", total_punches=0)
        cls.casual = Employee.objects.create(
            employee_code="501", first_name="Cas", last_name="Leave", employment_type="staff",
            branch=cls.b1, department=cls.d_acc,
        )  # fmt: skip
        record(cls.casual, D29, source="manual", override_note="Casual Leave (paid) approved by HR", total_punches=0)
        # absent with no punch, and someone on leave: never listed
        record(cls.adm3, D29, status="absent")
        record(cls.nodept, D29, status="absent", total_punches=2)  # absent verdict but it did punch: listed


class PresentRowTests(_PresentWorld):
    def test_only_present_employees_are_listed_by_department_with_continuous_sno(self):
        body = self.get()
        self.assertEqual(self.codes(body), ["500", "2316", "189", "226", "300", "2503", "EVE", "P5", "999"])
        self.assertEqual([r["sno"] for r in data_rows(body)], list(range(1, 10)))
        self.assertEqual(headings(body), ["ACCOUNTS", "ADMIN", "CAD", "PACKING", "Unassigned"])
        self.assertNotIn("2234", self.codes(body))  # absent, no punch
        self.assertNotIn("501", self.codes(body))  # a paid casual-leave day is not 'present'

    def test_golden_row(self):
        row = {r["ticketNo"]: r for r in data_rows(self.get())}["2316"]
        self.assertEqual(
            row,
            {
                "sno": 2,  # the HR-entered 500 sorts before it
                "ticketNo": "2316",
                "empName": "Meenatchi T",
                "designation": "ACCOUNTS EXECUTIVE",
                "nature": None,
                "date": "29/09/2026 Tue",
                "shiftIn": "09:00 AM",
                "actIn": "08:58 AM",  # the first of four punches
                "lateHours": None,
            },
        )

    def test_act_intime_is_the_first_punch_of_the_day_padded_to_two_digits(self):
        rows = {r["ticketNo"]: r for r in data_rows(self.get())}
        self.assertEqual(rows["189"]["actIn"], "09:00 AM")  # punch at 09:00:00, the stored verdict says 09:40
        self.assertEqual(rows["2503"]["actIn"], "09:20 AM")
        self.assertEqual(rows["226"]["actIn"], "08:00 AM")
        self.assertEqual(rows["999"]["actIn"], "08:30 AM")

    def test_a_present_day_with_no_punch_shows_the_stored_first_punch_or_a_dash(self):
        rows = {r["ticketNo"]: r for r in data_rows(self.get())}
        self.assertEqual(rows["500"]["actIn"], "09:00 AM")  # entered by HR: the stored first punch
        self.assertIsNone(rows["500"]["shiftIn"])  # and it has no shift
        record_ = AttendanceDayRecord.objects.get(employee=self.hr_present, date=D29)
        record_.first_punch = None
        record_.save()
        body = self.get()
        self.assertIsNone({r["ticketNo"]: r for r in data_rows(body)}["500"]["actIn"])
        self.assertTrue(any("marked present with no punch" in n for n in body["notes"]), body["notes"])

    def test_a_punch_without_a_present_verdict_is_still_listed_and_a_verdictless_day_is_counted(self):
        rows = {r["ticketNo"]: r for r in data_rows(self.get())}
        self.assertIn("999", rows)  # stored absent, yet it punched twice
        AttendanceDayRecord.objects.filter(employee=self.pk, date=D29).delete()  # nobody has opened the day
        body = self.get()
        self.assertIn("P5", self.codes(body))
        self.assertTrue(any("no stored attendance verdict" in n for n in body["notes"]), body["notes"])
        self.assertIsNone({r["ticketNo"]: r for r in data_rows(body)}["P5"]["lateHours"])


class PresentShiftTests(_PresentWorld):
    def test_shift_intime_is_the_start_of_the_assignment_in_force_on_that_date(self):
        rows = {r["ticketNo"]: r for r in data_rows(self.get())}
        self.assertEqual(rows["2316"]["shiftIn"], "09:00 AM")
        self.assertEqual(rows["189"]["shiftIn"], "07:30 AM")  # two cover today: the later start (Early) wins
        self.assertEqual(rows["2503"]["shiftIn"], "10:00 AM")  # the employee's custom start time
        self.assertEqual(rows["P5"]["shiftIn"], "09:00 AM")

    def test_the_same_employee_on_an_earlier_date_uses_the_earlier_assignment(self):
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-29", employeeIds=str(self.adm1.id))
        self.assertEqual([r["shiftIn"] for r in data_rows(body)], ["07:30 AM"])  # punched on the 29th only
        punch(self.adm1, D28, "09:30:00")
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-29", employeeIds=str(self.adm1.id))
        self.assertEqual(
            [(r["date"], r["shiftIn"]) for r in data_rows(body)],
            [("28/09/2026 Mon", "09:00 AM"), ("29/09/2026 Tue", "07:30 AM")],
        )

    def test_no_shift_in_force_is_a_dash_and_is_counted(self):
        body = self.get()
        rows = {r["ticketNo"]: r for r in data_rows(body)}
        self.assertIsNone(rows["226"]["shiftIn"])  # its assignment ended yesterday
        self.assertIsNone(rows["999"]["shiftIn"])  # never assigned
        self.assertIsNone(rows["300"]["shiftIn"])
        self.assertEqual(cards(body)["No shift assigned"], 4)  # 226, 300, 999 and the HR-entered 500


class PresentLateTests(_PresentWorld):
    def test_late_hours_come_from_the_stored_verdict_and_are_hh_mm_past_the_shift_start(self):
        rows = {r["ticketNo"]: r for r in data_rows(self.get())}
        self.assertEqual(rows["189"]["lateHours"], "02:10")  # first punch 09:40 stored, shift started 07:30 today
        self.assertIsNone(rows["2503"]["lateHours"])  # 09:20 is before its 10:00 start
        self.assertIsNone(rows["2316"]["lateHours"])  # on time: a dash

    def test_a_late_looking_punch_the_engine_did_not_flag_is_never_called_late_here(self):
        # Pooja's shift starts at 09:00; make her punch 11:00 but keep the stored 'on time' verdict
        AttendanceLog.objects.filter(employee=self.pk, date=D29, punch_time=at("09:00:00")).update(
            punch_time=at("11:00:00")
        )
        row = {r["ticketNo"]: r for r in data_rows(self.get())}["P5"]
        self.assertEqual(row["actIn"], "11:00 AM")
        self.assertIsNone(row["lateHours"])  # the stored verdict says on time: nothing is recomputed

    def test_a_late_day_whose_minutes_cannot_be_measured_says_late(self):
        rec = AttendanceDayRecord.objects.get(employee=self.acc, date=D29)
        rec.is_late = True
        rec.compensation_kind = "window_late_in"  # a company-announced late-in time decided it
        rec.save()
        body = self.get()
        self.assertEqual({r["ticketNo"]: r for r in data_rows(body)}["2316"]["lateHours"], "Late")
        self.assertTrue(any("show the word Late" in n for n in body["notes"]), body["notes"])

    def test_the_cards_count_late_and_on_time_days(self):
        body = self.get()
        self.assertEqual(cards(body)["Late arrivals"], 1)  # only 189 carries the stored late flag
        self.assertEqual(cards(body)["Employee-days listed"], 9)
        self.assertEqual(cards(body)["On time"], 8)

    def test_a_morning_permission_moves_the_start_the_way_the_engine_did(self):
        rec = AttendanceDayRecord.objects.get(employee=self.adm1, date=D29)
        rec.morning_permission_applied = True  # start moved one hour later: 08:30
        rec.save()
        self.assertEqual({r["ticketNo"]: r for r in data_rows(self.get())}["189"]["lateHours"], "01:10")


class PresentNatureTests(_PresentWorld):
    def test_nature_of_emp_is_a_dash_for_everyone_and_the_note_says_why(self):
        body = self.get()
        self.assertTrue(data_rows(body))
        self.assertEqual({r["nature"] for r in data_rows(body)}, {None})
        text = " ".join(body["notes"])
        self.assertIn("Nature of Emp (Day scholar / Hostel) is not recorded in the HRMS yet", text)
        self.assertIn("no such field", text)
        spec = registry.get_spec(PRESENT)
        self.assertIn("Nature of Emp", [c.label for c in spec.columns])

    def test_the_notes_state_the_rules_used(self):
        text = " ".join(self.get()["notes"])
        for part in ("Shift InTime", "latest start wins", "Act InTime", "Late Hours", "Nothing is judged again here"):
            self.assertIn(part, text)


class PresentFilterTests(_PresentWorld):
    def test_filters_narrow_the_list(self):
        self.assertEqual(self.codes(self.get(departmentIds=str(self.d_pk.id))), ["P5"])
        self.assertEqual(self.codes(self.get(employmentType="production")), ["300"])
        self.assertEqual(self.codes(self.get(designationIds=str(self.des_exec.id))), ["2316"])
        self.assertIn("OLD1", self.codes(self.get(employeeStatus="all")))
        self.assertNotIn("OLD1", self.codes(self.get()))

    def test_a_date_range_lists_each_present_day(self):
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-29", departmentIds=str(self.d_cad.id))
        self.assertEqual(
            [(r["ticketNo"], r["date"], r["actIn"]) for r in data_rows(body)],
            [("EVE", "28/09/2026 Mon", "09:00 AM"), ("EVE", "29/09/2026 Tue", "09:00 AM")],
        )  # Tuesday's 01:05 exit belongs to Monday: it is not Tuesday's first punch

    def test_a_future_date_is_an_empty_report(self):
        self.assertEqual(self.get(dateFrom="2026-09-30", dateTo="2026-09-30")["rows"], [])

    def test_a_branch_user_sees_only_their_branch(self):
        mine = self.get(user=self.branch_user)
        self.assertNotIn("P5", self.codes(mine))
        self.assertNotIn("PACKING", headings(mine))
        self.assertEqual(data_rows(self.get(user=self.branch_user, branchIds=str(self.b2.id))), [])


class PresentExportTests(_PresentWorld):
    def test_excel_carries_the_same_columns_and_rows_as_the_screen(self):
        params = {"dateFrom": "2026-09-28", "dateTo": "2026-09-29"}
        screen = self.get(**params)
        r = self.export("xlsx", **params)
        self.assertRegex(
            r["Content-Disposition"], r'filename="punching_register_employees_present_list_\d{8}_\d{4}\.xlsx"'
        )
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[HEADER_ROW]][: len(PRESENT_LABELS)], PRESENT_LABELS)
        self.assertGreater(len(screen["rows"]), 10)
        for i, row in enumerate(screen["rows"]):
            got = [ws.cell(row=HEADER_ROW + 1 + i, column=n + 1).value for n in range(len(PRESENT_KEYS))]
            want = [row.get(k) for k in PRESENT_KEYS]
            self.assertEqual(
                [str(g if g is not None else "") for g in got], [str(w if w is not None else "") for w in want]
            )

    def test_pdf_is_a_real_pdf(self):
        r = self.export("pdf", dateFrom="2026-09-28", dateTo="2026-09-29")
        self.assertTrue(r.content.startswith(b"%PDF"))
