"""Printable and coloured exports of the Attendance Sheet (muster roll). Registers no report.

Why the sheet has its own builders: it is a register with one column per date, up to 31 days -- 43 columns with
the identity and total columns. The generic PDF table gives every column a minimum width that cannot fit 43
columns on one page (it fails outright for a full month), and neither generic export can colour a cell by its
value. The Report Log's Attendance Sheet is a colour-coded register, so both exports here reproduce it: the same
palette (AttendanceSheet.tsx STATUS_META), the Strength row, the legend and the totals. The Excel file is the
framework's own workbook with the day cells coloured afterwards.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table, TableStyle

from ..branding import company
from ..export_pdf import BRAND, GRID, INK, MUTED, ensure_fonts, pdf_text
from ..export_xlsx import HEADER_ROW, build_xlsx
from ..formatting import display_date
from ..runner import KIND_SUBTOTAL, KIND_TOTAL
from .attendance_core_pdf import filter_box, kpi_strip, letterhead, numbered_canvas, style_factory, title_bar

# The colour of each code in the Report Log's Attendance Sheet (AttendanceSheet.tsx STATUS_META).
PALETTE = {
    "P": "C6EFCE",
    "CL": "C6EFCE",
    "CO": "C6EFCE",
    "½M": "FFEB9C",
    "½": "FFEB9C",
    "½E": "FED7AA",
    "A": "FFC7CE",
    "L": "BDD7EE",
    "H": "E2E8F0",
    "WO": "E2E8F0",
}
SUNDAY_HEAD = "FFE4E6"  # the Sheet's rose Sunday header
STRENGTH_FILL = "F1F5F9"
DAY_KEY = re.compile(r"^d(\d{4})(\d{2})(\d{2})$")
LEGEND_CHIPS = [
    ("P", "Present"),
    ("½M", "Half day - morning worked"),
    ("½E", "Half day - evening worked"),
    ("A", "Absent"),
    ("L", "On leave"),
    ("H", "Holiday / weekly off"),
    ("WO", "Weekly off"),
    ("CL", "Casual leave"),
    ("CO", "Comp off"),
]


def code_fill(code) -> str | None:
    """Hex fill of a cell code; a leave-type code such as SL or EL is a leave (blue). None for an empty cell."""
    if not isinstance(code, str) or not code.strip() or code == "-":
        return None
    return PALETTE.get(code, PALETTE["L"])


def _weekday_is_sunday(key: str) -> bool:
    m = DAY_KEY.match(key)
    if not m:
        return False
    return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).weekday() == 6


# ── Excel: the framework workbook + the Report Log colours ─────────────────────────────────────


def build_muster_xlsx(ctx, out) -> bytes:
    raw = build_xlsx(out)
    if not out.rows:
        return raw
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    day_cols = [(i, c) for i, c in enumerate(out.columns, start=1) if DAY_KEY.match(c.key)]
    code_font = Font(name="Calibri", size=9, bold=True)
    centre = Alignment(horizontal="center", vertical="center")
    for i, c in day_cols:
        ws.column_dimensions[get_column_letter(i)].width = 5.5
        if _weekday_is_sunday(c.key):
            head = ws.cell(row=HEADER_ROW, column=i)
            head.fill = PatternFill("solid", fgColor=SUNDAY_HEAD)
            head.font = Font(name="Calibri", size=10, bold=True, color="9F1239")
    for r, row in enumerate(out.rows, start=HEADER_ROW + 1):
        if row.get("_kind") == KIND_TOTAL:  # the Strength row
            for i, c in day_cols:
                cell = ws.cell(row=r, column=i)
                count = row.get(c.key)
                if isinstance(count, (int, float)):  # the generic sheet writes every date column as text
                    cell.value = count
                cell.fill = PatternFill("solid", fgColor=STRENGTH_FILL)
                cell.alignment = centre
            continue
        for i, c in day_cols:
            fill = code_fill(row.get(c.key))
            if fill is None:
                continue
            cell = ws.cell(row=r, column=i)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = code_font
            cell.alignment = centre
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF: A3 / A4 landscape register ───────────────────────────────────────────────────────────


def _fit(text: str, width: float, font: str, size: float, pad: float = 2.0) -> str:
    """Text cut with an ellipsis so it fits ``width`` points."""
    if pdfmetrics.stringWidth(text, font, size) <= width - 2 * pad:
        return text
    while text and pdfmetrics.stringWidth(text + "…", font, size) > width - 2 * pad:
        text = text[:-1]
    return text.rstrip() + "…"


def build_muster_pdf(ctx, out) -> bytes:
    regular, bold, rupee_ok = ensure_fonts()
    co = company()
    cols = out.columns
    day_idx = [i for i, c in enumerate(cols) if DAY_KEY.match(c.key)]
    n_days = len(day_idx)
    page = landscape(A3) if n_days > 16 else landscape(A4)
    margin = 1.0 * cm
    avail = page[0] - 2 * margin
    size = 7.0
    style = style_factory(regular, size)

    # -- column widths (points): identity + total columns are fixed, the date columns share what is left ------
    fixed = {"sno": 0.9, "employeeCode": 1.9, "employeeName": 4.2, "department": 3.0, "effective": 1.1, "shifts": 1.1}
    widths = []
    for i, c in enumerate(cols):
        if i in day_idx:
            widths.append(0.0)
        else:
            widths.append(fixed.get(c.key, 0.85) * cm)
    base = sum(widths)
    day_w = min(1.0 * cm, max(0.5 * cm, (avail - base) / n_days)) if n_days else 0.0
    for i in day_idx:
        widths[i] = day_w
    used = sum(widths)
    if used < avail:  # short ranges: give the spare width to the name and department columns
        extra = avail - used
        for i, c in enumerate(cols):
            if c.key in ("employeeName", "department"):
                widths[i] += extra / 2
    elif used > avail:
        widths = [w * avail / used for w in widths]

    s_head = style("h", fontName=bold, alignment=TA_CENTER, textColor=colors.white, fontSize=6.6, leading=7.8)
    s_head_sun = style(
        "hs", fontName=bold, alignment=TA_CENTER, textColor=colors.HexColor("#9F1239"), fontSize=6.6, leading=7.8
    )
    s_note = style("n", fontSize=7, leading=9, textColor=MUTED)

    story: list = []
    story += letterhead(co, avail, regular, bold, style)
    right = ""
    if ctx.date_from and ctx.date_to:
        right = f"{display_date(ctx.date_from)} to {display_date(ctx.date_to)}"
    story.append(title_bar(out.spec.title.upper(), right, avail, bold, style))
    story.append(filter_box(out.filters, avail, style))
    story.append(Spacer(1, 4))
    strip = kpi_strip(out.summary, avail, bold, style, rupee_ok)
    if strip is not None:
        story.append(strip)
        story.append(Spacer(1, 4))

    # -- the register ------------------------------------------------------------------------------------
    def head_cell(i: int, c):
        if i in day_idx:
            num, _, wd = c.label.partition(" ")
            return Paragraph(f"{escape(num)}<br/>{escape(wd)}", s_head_sun if _weekday_is_sunday(c.key) else s_head)
        return Paragraph(escape(c.label), s_head)

    data: list[list] = [[head_cell(i, c) for i, c in enumerate(cols)]]
    cmds: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("FONTNAME", (0, 1), (-1, -1), regular),
        ("FONTSIZE", (0, 1), (-1, -1), size),
        ("LEADING", (0, 1), (-1, -1), size * 1.2),
        ("TEXTCOLOR", (0, 1), (-1, -1), INK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, GRID),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 1.8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8),
    ]
    for i in day_idx:
        if _weekday_is_sunday(cols[i].key):
            cmds.append(("BACKGROUND", (i, 0), (i, 0), colors.HexColor("#" + SUNDAY_HEAD)))
    text_cols = [i for i, c in enumerate(cols) if c.key in ("employeeCode", "employeeName", "department")]
    for i in text_cols:
        cmds.append(("ALIGN", (i, 1), (i, -1), "LEFT"))
    for i, c in enumerate(cols):
        if i not in text_cols:
            cmds.append(("ALIGN", (i, 1), (i, -1), "CENTER"))

    for row in out.rows:
        idx = len(data)
        kind = row.get("_kind")
        cells = []
        for i, c in enumerate(cols):
            v = row.get(c.key)
            if i in day_idx:
                text = "-" if v is None or v == "" else str(v)
            elif i in text_cols:
                text = _fit("" if v is None else str(v), widths[i], bold if kind else regular, size)
            elif kind and v is None:
                text = ""
            else:
                text = pdf_text(c, v, rupee_ok)
            cells.append(text)
        data.append(cells)
        if kind in (KIND_TOTAL, KIND_SUBTOTAL):
            cmds.append(("FONTNAME", (0, idx), (-1, idx), bold))
            cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#" + STRENGTH_FILL)))
            continue
        for i in day_idx:
            fill = code_fill(row.get(cols[i].key))
            if fill:
                cmds.append(("BACKGROUND", (i, idx), (i, idx), colors.HexColor("#" + fill)))
                cmds.append(("FONTNAME", (i, idx), (i, idx), bold))
    if out.totals:
        idx = len(data)
        cells = []
        placed = False
        for i, c in enumerate(cols):
            v = out.totals.get(c.key)
            if v is None:
                cells.append("" if placed or c.key != "employeeName" else "TOTAL")
                placed = placed or c.key == "employeeName"
            else:
                cells.append(pdf_text(c, v, rupee_ok))
        data.append(cells)
        cmds.append(("FONTNAME", (0, idx), (-1, idx), bold))
        cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#DCE6F2")))
        cmds.append(("LINEABOVE", (0, idx), (-1, idx), 0.9, BRAND))

    if len(data) == 1:
        story.append(
            Paragraph("No records match the selected filters.", style("e", fontSize=10, leading=14, textColor=MUTED))
        )
    else:
        table = Table(data, colWidths=widths, repeatRows=1, splitByRow=1)
        table.setStyle(TableStyle(cmds))
        story.append(table)

    # -- legend + notes -------------------------------------------------------------------------------------
    story.append(Spacer(1, 6))
    legend_cells: list = []
    legend_cmds: list[tuple] = [
        ("FONTNAME", (0, 0), (-1, -1), regular),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
    ]
    legend_w: list[float] = []
    for n, (code, label) in enumerate(LEGEND_CHIPS):
        legend_cells += [code, label]
        legend_w += [0.9 * cm, 3.6 * cm]
        legend_cmds.append(("BACKGROUND", (2 * n, 0), (2 * n, 0), colors.HexColor("#" + PALETTE[code])))
        legend_cmds.append(("FONTNAME", (2 * n, 0), (2 * n, 0), bold))
        legend_cmds.append(("ALIGN", (2 * n, 0), (2 * n, 0), "CENTER"))
    legend = Table([legend_cells], colWidths=legend_w, hAlign="LEFT")
    legend.setStyle(TableStyle(legend_cmds))
    story.append(legend)
    if out.truncated:
        story.append(Spacer(1, 4))
        story.append(
            Paragraph(
                f"<b>Note:</b> only the first {out.row_count:,} rows are printed here. Narrow the filters to see the rest.",
                s_note,
            )
        )
    for note in out.notes:
        story.append(Spacer(1, 2))
        story.append(Paragraph(escape(note), s_note))

    buf = io.BytesIO()
    doc = BaseDocTemplate(
        buf,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=1.5 * cm,
        title=out.spec.title,
        author=co.name,
    )
    frame = Frame(
        margin,
        1.5 * cm,
        avail,
        page[1] - margin - 1.5 * cm,
        id="body",
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )
    doc.addPageTemplates([PageTemplate(id="sheet", frames=[frame])])
    stamp = f"Generated {display_date(out.generated_at[:10])} {out.generated_at[11:]} by {out.generated_by}"
    doc.build(story, canvasmaker=numbered_canvas(f"{co.name} - {out.spec.title}  |  {stamp}", regular, margin))
    return buf.getvalue()
