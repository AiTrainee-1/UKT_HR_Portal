"""
Payroll core - money movement and deductions:

    bank-advice              net pay to transfer by bank (with an HDFC bulk-upload sheet in Excel) + a separate cash list
    payroll-payment-status   which slips are marked paid / pending, and where slip and payroll disagree
    salary-deduction-summary employee-wise PF, ESI, advance recovery and late deductions

Net pay is always the SLIP's net_salary (what the employee is handed). Payroll.final_salary can differ after a
'mark paid' edit, so that difference is surfaced, never silently picked.
"""

from __future__ import annotations

import dataclasses
import io
import re
from datetime import date

from django.db.models import Q, TextField, Value
from django.db.models.functions import Replace, Upper

from api.branch_scope import get_branch_scope
from api.models import Employee, Payroll

from ..common import EMP_COLS, EMP_COLS_SHORT, emp_cells
from ..filters import boolean, period, scope, select, text
from ..formatting import fmt_dt, r2
from ..registry import register
from ..types import BADGE, CURRENCY, DATETIME, PERCENT, TEXT, ColumnSpec, ReportResult, ReportSpec
from .payroll_core_common import (
    PAYMENT_KEY,
    STATUS_PAID,
    STATUS_PENDING,
    SlipRow,
    amount_in_words,
    is_paid,
    legacy_filter,
    load_slips,
    month_q,
    nat_key,
    payroll_key,
    period_label,
    provisional_note,
    slip_filter,
    slip_key,
    status_text,
    subtotals,
    type_label,
)

CATEGORY = "payroll"
BOTH = ("payroll", "production_payroll")
MONEY = 1.3

_IFSC = re.compile(r"^[A-Z]{4}0[A-Z0-9]{6}$")
# The bank template's BENE A/C NO column: letters and digits only, at most 25 characters. A longer or otherwise odd
# number cannot be uploaded as typed and must never be cut short (a truncated number can credit a stranger).
_ACCOUNT = re.compile(r"^[A-Za-z0-9]{1,25}$")


def _cur(key: str, label: str, width: float = MONEY, total: str | None = "sum") -> ColumnSpec:
    return ColumnSpec(key, label, CURRENCY, width, total=total)


# ═════════════════════════════════════════════════════════════════════════════
#  bank-advice
# ═════════════════════════════════════════════════════════════════════════════

MODE_BANK, MODE_CASH = "Bank", "Cash"
OK = "OK"
ISSUE_NO_BANK = "No bank details"
ISSUE_IFSC = "Invalid IFSC"
ISSUE_ACCT = "Invalid account no."
ISSUE_DUP = "Duplicate account"
ISSUE_NET = "Zero/negative net"
ISSUE_DIFF = "Net differs from payroll"
BLOCKING = (ISSUE_NO_BANK, ISSUE_IFSC, ISSUE_ACCT, ISSUE_NET)

# The bank-portal template columns (mirrors frontend/src/lib/payrollExcelExport.ts) and their length limits.
HDFC_HEADERS = [
    "TXN TYPE\nRTGS-R\nNEFT-N\nHDFC TRF-I\nIMPS - M(1)",
    'BENEFICIARY CODE\n(MANDATORY ONLY FOR\nTXN TYPE "I")\n(13)',
    "BENE A/C NO\n(25)",
    "AMOUNT\n(20)",
    "BENEFICIARY NAME\n(35)",
    "IFSC CODE\n(11)",
    "BENE BANK NAME\n(40)",
    "BENE BANK BRANCH NAME\n(40)",
    "BENE EMAIL ID\n(100)",
]
HDFC_INDICATORS = ["M", "M / O", "M", "M", "M", "M", "M", "O", "M"]
HDFC_WIDTHS = [18, 22, 24, 14, 32, 16, 30, 30, 32]
HDFC_SHEET = "HDFC Bulk Upload"

BANK_FILTERS = (
    period(default="lastMonth"),
    *scope(designation=False, status=None),
    select(
        PAYMENT_KEY,
        "Payment status",
        [("pending", "Pending"), ("paid", "Paid"), ("all", "All")],
        default="pending",
        help="Defaults to pending so the list is what still has to be paid.",
    ),
    select("paymentMode", "Payment mode", [("bank", "Bank transfer"), ("cash", "Cash salary")]),
    text("bankName", "Bank name contains", "e.g. HDFC"),
    boolean(
        "combinePeriods", "One line per employee", help="Add an employee's period slips (production) into one payment."
    ),
)


def _acct(v) -> str:
    return "".join((v or "").split())


def _ifsc(v) -> str:
    return "".join((v or "").split()).upper()


def _narration(s) -> str:
    if s.period_start is not None and s.period_end is not None:
        return f"WAGES {s.period_start.strftime('%d%b').upper()}-{s.period_end.strftime('%d%b%Y').upper()}"
    return f"SALARY {date(s.year, s.month, 1).strftime('%b %Y').upper()}"


@dataclasses.dataclass
class _Line:
    slip: object
    employee: object
    net: float
    period: str
    narration: str
    paid: bool
    diff: bool


def _lines(loaded: list[SlipRow], combine: bool) -> list[_Line]:
    lines: list[_Line] = []
    for r in loaded:
        s = r.slip
        payroll_net = r.payroll["final_salary"] if r.payroll else None
        diff = payroll_net is not None and abs(float(payroll_net) - float(s.net_salary)) >= 0.01
        lines.append(_Line(s, s.employee, r.d["net"], r.d["period"], _narration(s), r.paid, diff))
    if not combine:
        return lines
    # One line per employee AND payment state: a paid period is never folded into an unpaid line (that line would ask
    # the bank to pay the paid period a second time), so a partly-paid employee gets a Paid line and a Pending line.
    merged: dict[tuple[int, bool], _Line] = {}
    for ln in lines:
        key = (ln.employee.id, ln.paid)
        prev = merged.get(key)
        if prev is None:
            merged[key] = ln
            continue
        merged[key] = _Line(
            prev.slip,
            prev.employee,
            round(prev.net + ln.net, 2),
            f"{prev.period} (+{ln.period})" if prev.period != ln.period else prev.period,
            "SALARY/WAGES " + date(prev.slip.year, prev.slip.month, 1).strftime("%b %Y").upper(),
            ln.paid,
            prev.diff or ln.diff,
        )
    return list(merged.values())


def _shared_accounts(ctx, lines: list[_Line]) -> set[str]:
    """Account numbers held by more than one person: the listed employees plus every ACTIVE employee the user may
    see (a paid colleague, or a duplicate profile, sharing an account is worth a second look). One query."""
    listed: dict[str, set[int]] = {}
    for ln in lines:
        a = _acct(ln.employee.bank_account).upper()
        if a:
            listed.setdefault(a, set()).add(ln.employee.id)
    if not listed:
        return set()
    qs = Employee.objects.filter(status="active")
    branch = get_branch_scope(ctx.request)
    if branch is not None:
        qs = qs.filter(branch_id=branch)
    norm = Upper(Replace("bank_account", Value(" "), Value(""), output_field=TextField()), output_field=TextField())
    for eid, n in qs.annotate(norm=norm).filter(norm__in=list(listed)).values_list("id", "norm"):
        listed[n].add(eid)
    return {a for a, ids in listed.items() if len(ids) > 1}


def _bank_run(ctx) -> ReportResult:
    combine = bool(ctx.params.get("combinePeriods"))
    needle = (ctx.params.get("bankName") or "").strip()
    extra = Q(employee__bank_name__icontains=needle) if needle else None
    params = dict(ctx.params)
    if params.get(PAYMENT_KEY) == "all":
        params[PAYMENT_KEY] = None
    loaded, notes = load_slips(
        dataclasses.replace(ctx, params=params),
        legacy_key=None,
        extra_q=extra,
        order="code",
    )
    lines = _lines(loaded, combine)

    shared = _shared_accounts(ctx, lines)

    bank_rows, cash_rows = [], []
    for ln in lines:
        e = ln.employee
        acct, ifsc = _acct(e.bank_account), _ifsc(e.bank_ifsc)
        issues: list[str] = []
        if not acct:
            issues.append(ISSUE_NO_BANK)
        else:
            if not _ACCOUNT.match(acct):
                issues.append(ISSUE_ACCT)
            if not _IFSC.match(ifsc):
                issues.append(ISSUE_IFSC)
            if acct.upper() in shared:
                issues.append(ISSUE_DUP)
        if ln.net <= 0:
            issues.append(ISSUE_NET)
        if ln.diff:
            issues.append(ISSUE_DIFF)
        row = {
            **{k: v for k, v in emp_cells(e).items() if k != "designation"},
            "bankName": (e.bank_name or "").strip() or None,
            "bankAccount": acct or None,
            "bankIfsc": ifsc or None,
            "netPay": ln.net,
            "paymentMode": MODE_BANK if acct else MODE_CASH,
            "period": ln.period,
            "narration": ln.narration,
            "paymentStatus": STATUS_PAID if ln.paid else STATUS_PENDING,
            "check": "; ".join(issues) if issues else OK,
        }
        (bank_rows if acct else cash_rows).append(row)

    mode = ctx.params.get("paymentMode")
    if mode == "bank":
        cash_rows = []
    elif mode == "cash":
        bank_rows = []
    bank_rows.sort(key=lambda r: ((r["bankName"] or "~").upper(), nat_key(r["employeeCode"]), r["period"]))
    cash_rows.sort(key=lambda r: (r["department"].lower(), nat_key(r["employeeCode"]), r["period"]))
    rows = subtotals(
        bank_rows + cash_rows,
        lambda r: r["paymentMode"],
        ["netPay"],
        label=lambda g: "Bank transfer total" if g == MODE_BANK else "Cash salary total",
    )
    bank_total = round(sum(r["netPay"] for r in bank_rows), 2)
    cash_total = round(sum(r["netPay"] for r in cash_rows), 2)
    flagged = sum(1 for r in bank_rows + cash_rows if r["check"] not in (OK, ISSUE_NO_BANK))
    no_bank = sum(
        1
        for r in bank_rows + cash_rows
        if ISSUE_NO_BANK in r["check"] or ISSUE_IFSC in r["check"] or ISSUE_ACCT in r["check"]
    )
    summary = [
        {"label": "Total net pay", "value": round(bank_total + cash_total, 2), "format": "currency"},
        {"label": "Payments", "value": len(bank_rows) + len(cash_rows), "format": "integer"},
        {"label": "Bank transfers", "value": bank_total, "format": "currency"},
        {"label": "Bank payments", "value": len(bank_rows), "format": "integer"},
        {"label": "Cash salaries", "value": cash_total, "format": "currency"},
        {"label": "Cash payments", "value": len(cash_rows), "format": "integer"},
        {"label": "Missing / invalid bank details", "value": no_bank, "format": "integer"},
        {"label": "Rows needing attention", "value": flagged, "format": "integer"},
    ]
    all_notes = [
        "Amount = the salary slip's net pay (what the employee is handed). Payroll's own final salary can differ after a "
        "'mark paid' edit - such rows are flagged 'Net differs from payroll' instead of being silently changed.",
        "Employees with no bank account are listed separately as cash salaries. Account number, IFSC and bank name come "
        "from the employee profile; spaces are stripped, IFSC must look like AAAA0XXXXXX and the account number must be "
        "letters and digits only, at most 25 characters (an account number is never cut short). Rows flagged No bank "
        "details, Invalid IFSC, Invalid account no. or Zero/negative net - and rows already marked Paid - are left out "
        "of the HDFC upload sheet, so uploading it can never pay a salary twice.",
        "With 'One line per employee', an employee whose periods are only partly paid gets two lines - one Paid, one "
        "Pending - so a paid period is never requested again.",
        "Payment status reflects the payroll row's flag only - there is no payment date or company debit account in the "
        "system, so this is a payment list rather than a bank-portal file; the Excel download adds an 'HDFC Bulk Upload' "
        "sheet in the bank's template for the bank-ready rows.",
        f"Amount in words (bank transfers): {amount_in_words(bank_total)}. Total including cash: "
        f"{amount_in_words(round(bank_total + cash_total, 2))}.",
        *notes,
    ]
    prov = provisional_note(loaded)
    if prov:
        all_notes.append(prov)
    if not rows:
        all_notes.append("Nothing to pay for these filters (try Payment status = All, or generate payroll first).")
    return ReportResult(rows=rows, summary=summary, notes=all_notes)


def _bank_xlsx(ctx, out) -> bytes:
    """The generic Excel report plus, on a second sheet, the bank's bulk-upload template for the bank-ready rows."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..export_xlsx import build_xlsx, safe_cell

    generic = build_xlsx(out)
    # Bank-ready = a bank row with no blocking issue that is still to be paid. Paid rows are never repeated in an
    # upload file, whatever Payment status the report was run with.
    ready = [
        r
        for r in out.rows
        if r.get("_kind") is None
        and r.get("paymentMode") == MODE_BANK
        and r.get("paymentStatus") != STATUS_PAID
        and not set(str(r.get("check") or "").split("; ")) & set(BLOCKING)
        and _IFSC.match(_ifsc(r.get("bankIfsc")))
        and _ACCOUNT.match(_acct(r.get("bankAccount")))
        and (r.get("netPay") or 0) > 0
    ]
    if not ready:
        return generic
    wb = load_workbook(io.BytesIO(generic))
    ws = wb.create_sheet(HDFC_SHEET)
    fill = PatternFill("solid", fgColor="5B8C00")
    for i, w in enumerate(HDFC_WIDTHS, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.append(HDFC_HEADERS)
    ws.append(HDFC_INDICATORS)
    for row_idx, height in ((1, 72), (2, 18)):
        ws.row_dimensions[row_idx].height = height
        for cell in ws[row_idx]:
            cell.fill = fill
            cell.font = Font(bold=True, color="FF0000" if (row_idx == 1 and cell.column == 4) else "FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in ready:
        ws.append(
            [
                None,
                None,
                safe_cell(_acct(r["bankAccount"]).upper()),
                round(float(r["netPay"]), 2),
                safe_cell((r["employeeName"] or "").upper()[:35]),
                _ifsc(r["bankIfsc"])[:11],
                safe_cell((r["bankName"] or "").upper()[:40]),
                None,
                None,
            ]
        )
        row_idx = ws.max_row
        ws.row_dimensions[row_idx].height = 16
        for cell in ws[row_idx]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in (3, 5, 6, 7):
            c = ws.cell(row=row_idx, column=col)
            c.number_format = "@"
            if isinstance(c.value, str):
                c.data_type = "s"  # an account number is never a formula or a float
        ws.cell(row=row_idx, column=4).number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


register(
    ReportSpec(
        id="bank-advice",
        title="Bank Transfer Advice / Cash Salary List",
        description="Net pay to transfer by bank with account and IFSC, plus a separate cash-salary list and validity checks.",
        category=CATEGORY,
        modules=BOTH,
        icon="Landmark",
        tags=("bank", "neft", "rtgs", "hdfc", "transfer", "cash salary", "net pay", "disbursement", "ifsc"),
        filters=BANK_FILTERS,
        columns=(
            EMP_COLS[0],
            EMP_COLS[1],
            EMP_COLS[2],
            ColumnSpec("bankName", "Bank", TEXT, 1.4),
            ColumnSpec("bankAccount", "Account no.", TEXT, 2.4),
            ColumnSpec("bankIfsc", "IFSC", TEXT, 1.6),
            _cur("netPay", "Net pay", 1.6),
            ColumnSpec("paymentMode", "Mode", BADGE, 0.9),
            ColumnSpec("period", "Period", TEXT, 1.9),
            ColumnSpec("narration", "Narration", TEXT, 2.0),
            ColumnSpec("paymentStatus", "Payment", BADGE, 1.2),
            ColumnSpec("check", "Check", BADGE, 1.8),
        ),
        run=_bank_run,
        xlsx_builder=_bank_xlsx,
    )
)


# ═════════════════════════════════════════════════════════════════════════════
#  payroll-payment-status
# ═════════════════════════════════════════════════════════════════════════════

MATCH_OK, MATCH_SLIP_ONLY, MATCH_PAYROLL_ONLY, MATCH_DIFF = "OK", "Slip only", "Payroll only", "Net differs"


def _payment_status_run(ctx) -> ReportResult:
    only_issues = bool(ctx.params.get("onlyIssues"))
    wanted = ctx.params.get(PAYMENT_KEY)
    base_ctx = dataclasses.replace(ctx, params={**ctx.params, PAYMENT_KEY: None})
    loaded, notes = load_slips(base_ctx, legacy_key=None, order="dept")

    rows: list[dict] = []
    keys = set()
    for r in loaded:
        s, d = r.slip, r.d
        keys.add(slip_key(s))
        p = r.payroll
        pay_net = r2(p["final_salary"]) if p else None
        variance = round(pay_net - d["net"], 2) if pay_net is not None else None
        if p is None:
            match = MATCH_SLIP_ONLY
        elif variance is not None and abs(variance) >= 0.01:
            match = MATCH_DIFF
        else:
            match = MATCH_OK
        rows.append(
            {
                **emp_cells(s.employee),
                "employmentType": d["type"],
                "period": d["period"],
                "slipNet": d["net"],
                "payrollNet": pay_net,
                "variance": variance,
                "paymentStatus": status_text(p),
                "match": match,
                "lastUpdated": fmt_dt(p["updated_at"]) if p else None,
            }
        )

    # payroll rows that have no slip (a slip deleted, or payroll edited without one)
    q, _n = slip_filter(base_ctx, legacy_key=None)
    orphan_qs = (
        Payroll.objects.select_related("employee", "employee__department", "employee__designation")
        .filter(q)
        .filter(month_q(*ctx.period))
        .order_by("employee__employee_code", "id")[: ctx.row_limit]
    )
    for p in orphan_qs:
        if payroll_key(p.employee_id, p.period_start, p.period_end, p.week_number, p.year, p.month) in keys:
            continue
        rows.append(
            {
                **emp_cells(p.employee),
                "employmentType": type_label(p),
                "period": period_label(p),
                "slipNet": None,
                "payrollNet": r2(p.final_salary),
                "variance": None,
                "paymentStatus": STATUS_PAID if is_paid({"status": p.status}) else STATUS_PENDING,
                "match": MATCH_PAYROLL_ONLY,
                "lastUpdated": fmt_dt(p.updated_at),
            }
        )
    rows.sort(key=lambda r: (r["department"].lower(), nat_key(r["employeeCode"]), r["period"]))

    if wanted in ("paid", "pending"):
        rows = [r for r in rows if r["paymentStatus"] == (STATUS_PAID if wanted == "paid" else STATUS_PENDING)]
    if only_issues:
        rows = [r for r in rows if r["match"] != MATCH_OK]

    def amount(r):
        return r["slipNet"] if r["slipNet"] is not None else (r["payrollNet"] or 0)

    paid = [r for r in rows if r["paymentStatus"] == STATUS_PAID]
    pending = [r for r in rows if r["paymentStatus"] != STATUS_PAID]
    summary = [
        {"label": "Paid amount", "value": round(sum(amount(r) for r in paid), 2), "format": "currency"},
        {"label": "Paid slips", "value": len(paid), "format": "integer"},
        {"label": "Pending amount", "value": round(sum(amount(r) for r in pending), 2), "format": "currency"},
        {"label": "Pending slips", "value": len(pending), "format": "integer"},
        {
            "label": "Net variance (payroll - slip)",
            "value": round(sum(r["variance"] or 0 for r in rows), 2),
            "format": "currency",
        },
        {
            "label": "Rows needing attention",
            "value": sum(1 for r in rows if r["match"] != MATCH_OK),
            "format": "integer",
        },
    ]
    all_notes = [
        "Payment status comes from the payroll row's status: only 'paid' counts as paid, anything else (or no payroll row) is "
        "pending. There is no payment date in the system - 'Last updated' is when the payroll row was last changed, which "
        "regenerating payroll also does (regeneration resets a paid row to pending).",
        "Slip net is the printed salary slip; payroll net is the payroll row's final salary. They differ after a 'mark "
        "paid' edit (which drops overtime) or a manual bonus/deduction edit. 'Slip only' = a slip with no payroll row; "
        "'Payroll only' = a payroll row with no slip.",
        *notes,
    ]
    if not rows:
        all_notes.append("No payroll or slips exist for these filters.")
    return ReportResult(rows=rows, summary=summary, notes=all_notes)


register(
    ReportSpec(
        id="payroll-payment-status",
        title="Payroll Payment Status",
        description="Which payrolls are marked paid or pending, with slip vs payroll net and mismatches to review.",
        category=CATEGORY,
        modules=BOTH,
        icon="BadgeCheck",
        tags=("paid", "pending", "payment", "status", "variance", "reconciliation"),
        filters=(
            period(default="lastMonth"),
            *scope(designation=False, status="all"),
            select(PAYMENT_KEY, "Payment status", [("paid", "Paid"), ("pending", "Pending")]),
            boolean("onlyIssues", "Only rows needing attention", help="Slip/payroll mismatches and net differences."),
        ),
        columns=(
            *EMP_COLS_SHORT,
            ColumnSpec("employmentType", "Type", BADGE, 1.0),
            ColumnSpec("period", "Period", TEXT, 1.9),
            _cur("slipNet", "Slip net pay"),
            _cur("payrollNet", "Payroll net pay"),
            _cur("variance", "Variance"),
            ColumnSpec("paymentStatus", "Payment", BADGE, 1.0),
            ColumnSpec("match", "Match", BADGE, 1.1),
            ColumnSpec("lastUpdated", "Last updated", DATETIME, 1.5),
        ),
        run=_payment_status_run,
    )
)


# ═════════════════════════════════════════════════════════════════════════════
#  salary-deduction-summary
# ═════════════════════════════════════════════════════════════════════════════

DEDUCTION_SUM = (
    "totalEarnings",
    "pfDeduction",
    "esiDeduction",
    "advanceDeduction",
    "lateDeduction",
    "otherDeductions",
    "totalDeductions",
    "netPay",
)
HEAD_KEYS = {
    "pf": "pfDeduction",
    "esi": "esiDeduction",
    "advance": "advanceDeduction",
    "late": "lateDeduction",
    "other": "otherDeductions",
}


def _pct(part, whole) -> float | None:
    return round(part / whole * 100, 2) if whole else None


def _deduction_run(ctx) -> ReportResult:
    group = ctx.param("groupBy", "department")
    head = ctx.param("deductionHead", "any")
    loaded, notes = load_slips(ctx, order="dept" if group == "department" else "code")
    rows = []
    for r in loaded:
        s, d = r.slip, r.d
        rows.append(
            {
                **emp_cells(s.employee),
                "employmentType": d["type"],
                "period": d["period"],
                "totalEarnings": d["total_earnings"],
                "pfDeduction": d["pf"],
                "esiDeduction": d["esi"],
                "advanceDeduction": d["advance"],
                "lateDeduction": d["late"],
                "otherDeductions": d["other"],
                "totalDeductions": d["total_deductions"],
                "deductionPct": _pct(d["total_deductions"], d["total_earnings"]),
                "netPay": d["net"],
            }
        )
    if head == "any":
        rows = [r for r in rows if r["totalDeductions"] > 0]
    elif head in HEAD_KEYS:
        rows = [r for r in rows if (r[HEAD_KEYS[head]] or 0) > 0]
    if group == "employmentType":
        rows.sort(key=lambda r: (r["employmentType"], r["department"].lower(), nat_key(r["employeeCode"]), r["period"]))
        rows = subtotals(rows, lambda r: r["employmentType"], DEDUCTION_SUM)
    elif group == "department":
        rows = subtotals(rows, lambda r: r["department"], DEDUCTION_SUM)
    for r in rows:
        if r.get("_kind") == "subtotal":
            r["deductionPct"] = _pct(r["totalDeductions"], r["totalEarnings"])
    data = [r for r in rows if r.get("_kind") != "subtotal"]

    def col(key):
        return round(sum(x[key] or 0 for x in data), 2)

    total_ded = col("totalDeductions")
    # The percentage is over EVERY slip of the selection, not just the listed rows: the default "Any deduction" filter
    # hides slips without deductions, and dividing by only the remaining earnings would inflate the figure.
    all_ded = round(sum(r.d["total_deductions"] for r in loaded), 2)
    all_earn = round(sum(r.d["total_earnings"] for r in loaded), 2)
    with_ded = len({x["employeeCode"] for x in data if x["totalDeductions"] > 0})
    summary = [
        {"label": "Total deductions", "value": total_ded, "format": "currency"},
        {"label": "PF", "value": col("pfDeduction"), "format": "currency"},
        {"label": "ESI", "value": col("esiDeduction"), "format": "currency"},
        {"label": "Advance recovery", "value": col("advanceDeduction"), "format": "currency"},
        {"label": "Late deduction", "value": col("lateDeduction"), "format": "currency"},
        {"label": "Other / manual", "value": col("otherDeductions"), "format": "currency"},
        {"label": "Employees with deductions", "value": with_ded, "format": "integer"},
        {"label": "Deductions as % of earnings", "value": _pct(all_ded, all_earn), "format": "percent"},
    ]
    all_notes = [
        "Deduction % = total deductions / total earnings (gross pay + overtime); the summary card uses every slip of the "
        "selection, including the ones the Deduction head filter hides. Late deduction comes from the slip's "
        "late-detection snapshot (production slips keep it only inside total deductions); Other / manual = total less "
        "PF, ESI, advance and late - normally zero, so a value there points to a legacy or manually edited slip.",
        "Absence and unpaid leave are NOT deduction lines: they are already reduced from the pro-rated gross pay.",
        "Advance recovery is the figure stored on the slip. Regenerating a month after the instalment was recovered "
        "stores 0 there even though the instalment stays recovered, so it can be understated for regenerated months "
        "(the Advance ledger is the authority).",
    ]
    if total_ded:
        shares = ", ".join(
            f"{label} {_pct(col(key), total_ded)}%"
            for label, key in (
                ("PF", "pfDeduction"),
                ("ESI", "esiDeduction"),
                ("Advance", "advanceDeduction"),
                ("Late", "lateDeduction"),
                ("Other", "otherDeductions"),
            )
        )
        all_notes.append(f"Share of total deductions: {shares}.")
    all_notes.extend(notes)
    prov = provisional_note(loaded)
    if prov:
        all_notes.append(prov)
    if not rows:
        all_notes.append("No slips with deductions for these filters (choose 'All slips' to list everyone).")
    return ReportResult(rows=rows, summary=summary, notes=all_notes)


register(
    ReportSpec(
        id="salary-deduction-summary",
        title="Salary Deduction Summary",
        description="PF, ESI, advance recovery and late deductions per employee with total and percent of earnings.",
        category=CATEGORY,
        modules=BOTH,
        icon="Percent",
        tags=("deductions", "pf", "esi", "advance", "late", "salary detection", "recoveries"),
        filters=(
            period(default="lastMonth"),
            *scope(designation=False, status="all"),
            legacy_filter(),
            select(
                "deductionHead",
                "Deduction head",
                [
                    ("any", "Any deduction"),
                    ("pf", "PF"),
                    ("esi", "ESI"),
                    ("advance", "Advance recovery"),
                    ("late", "Late deduction"),
                    ("other", "Other / manual"),
                    ("all", "All slips"),
                ],
                default="any",
                placeholder="Any deduction",
            ),
            select(
                "groupBy",
                "Group by",
                [("department", "Department"), ("employmentType", "Staff / production"), ("none", "No grouping")],
                default="department",
                placeholder="Department",
            ),
        ),
        columns=(
            *EMP_COLS_SHORT,
            ColumnSpec("employmentType", "Type", BADGE, 1.3),
            ColumnSpec("period", "Period", TEXT, 1.9),
            _cur("totalEarnings", "Total earnings", 1.6),
            _cur("pfDeduction", "PF"),
            _cur("esiDeduction", "ESI"),
            _cur("advanceDeduction", "Advance"),
            _cur("lateDeduction", "Late deduction"),
            _cur("otherDeductions", "Other / manual"),
            _cur("totalDeductions", "Total deductions", 1.5),
            ColumnSpec("deductionPct", "Deduction %", PERCENT, 1.3),
            _cur("netPay", "Net pay", 1.5),
        ),
        run=_deduction_run,
    )
)
