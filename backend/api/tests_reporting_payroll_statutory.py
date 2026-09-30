"""
Report Center - group G2 (payroll: statutory statements, salary deductions, overtime, controls).

Every figure below is derived by hand from a small February-2026 payroll that is generated through the REAL
payroll engine (payroll_views._generate_staff_payroll / _generate_production_payroll) from manual attendance
rows, so the reports are checked against exactly what payroll stored.

February 2026 starts on a Sunday: 24 Mon-Sat working days, 28 calendar days. Salary 24,000 -> daily rate 1,000.

The cast (all Feb 2026, PF 12% of basic, ESI 0.75% while salary <= 21,000, late free allowance 3,
slabs 1 billable -> 0.5 shift and 4 -> 1 shift):

  3001 A  salary 20,400  full month, 2 paid OT days (slip OT 1,700), 2 relaxation credits, 1 detected, 1 rejected
  3002 B  salary 24,000  18 full (5 late) + 2 half + 2 absent + 2 unpaid leave -> gross 19,000, late penalty 500
  3003 C  salary 15,000  full month, advance instalment 1,000, no bank / PF / ESI details, branch 2
  3004 D  salary 60,000  2 late + 1 early-out + 5 permissions (2 excess) -> late penalty 1,250
  3006 L  salary 18,000  full month, then left the company (inactive)
  3007 E  salary 12,000  never got a slip (announced OT pay day)        3008 F  no salary configured
  P001    production 500/shift, two periods (12.5 + 4 shifts) in the month

Run through the scratchpad run_tests.py wrapper (private throw-away database).
"""

import io
from datetime import date, datetime, time, timezone
from decimal import Decimal
from unittest import mock

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .models import (
    Advance,
    AdvanceRepayment,
    AttendanceDayRecord,
    Branch,
    CompensationDayAnnouncement,
    CompensationLeaveCredit,
    Department,
    Designation,
    Employee,
    EmployeePermission,
    HRUser,
    OvertimeRecord,
    Payroll,
    PayrollSettings,
    Role,
    SalarySlip,
)
from .payroll_views import _build_working_days, _generate_production_payroll, _generate_staff_payroll
from .permission_registry import all_module_keys
from .reporting import registry

MONTH, YEAR = 2, 2026
WORKING_DAYS = _build_working_days(MONTH, YEAR, False, set())  # 24 Mon-Sat days
P = {"period": "2026-02"}
OT_FEB = {"dateFrom": "2026-02-01", "dateTo": "2026-02-28"}
PROD1 = (date(2026, 2, 2), date(2026, 2, 11))
PROD2 = (date(2026, 2, 16), date(2026, 2, 25))
ANNOUNCED = datetime(2026, 2, 27, 20, 30, tzinfo=timezone.utc)  # 28-Feb 02:00 IST

G2_IDS = (
    "pf-statement", "esi-statement", "statutory-coverage", "late-salary-impact", "attendance-salary-loss",
    "overtime-register", "overtime-payment-summary", "compensation-credits", "compensation-day-announcements",
    "payroll-exceptions", "min-wage-compliance", "payroll-slip-reconciliation",
)  # fmt: skip
# Parameters that make each report show data for the fixture payroll.
DATA_PARAMS = {
    "pf-statement": P, "esi-statement": P, "statutory-coverage": {}, "late-salary-impact": P,
    "attendance-salary-loss": P, "overtime-register": OT_FEB, "overtime-payment-summary": P,
    "compensation-credits": {}, "compensation-day-announcements": {"when": "all"}, "payroll-exceptions": P,
    "min-wage-compliance": P, "payroll-slip-reconciliation": P,
}  # fmt: skip
# Parameters under which every report is empty.
EMPTY_PARAMS = {
    "pf-statement": {"period": "2019-01"}, "esi-statement": {"period": "2019-01"},
    "statutory-coverage": {"coverage": "no-slip", "departmentIds": "999999"},
    "late-salary-impact": {"period": "2019-01"}, "attendance-salary-loss": {"period": "2019-01"},
    "overtime-register": {"dateFrom": "2019-01-01", "dateTo": "2019-01-31"},
    "overtime-payment-summary": {"period": "2019-01"},
    "compensation-credits": {"status": "used", "departmentIds": "999999"},
    "compensation-day-announcements": {"when": "upcoming"},
    "payroll-exceptions": {"period": "2019-01", "departmentIds": "999999"}, "min-wage-compliance": {"period": "2019-01"},
    "payroll-slip-reconciliation": {"period": "2019-01"},
}  # fmt: skip


def _headers(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def _staff(code, salary, dept, branch, **kw) -> Employee:
    return Employee.objects.create(
        employee_code=code, first_name=f"E{code}", last_name="Staff", employment_type="staff", status="active",
        salary_type="monthly", salary_amount=Decimal(salary) if salary is not None else None,
        department=dept, branch=branch, **kw,
    )  # fmt: skip


def _day(emp, d, status="present", shifts="1.00", **kw):
    return AttendanceDayRecord.objects.create(
        employee=emp, date=d, status=status, shifts_earned=Decimal(shifts),
        is_half_shift=(status == "half_shift"), source="manual", **kw,
    )  # fmt: skip


def _present_all(emp, **kw):
    for d in WORKING_DAYS:
        _day(emp, d, **kw)


def _ot(emp, day, minutes, status="announced", ctype="pay", announced_at=None, by="HR One"):
    return OvertimeRecord.objects.create(
        employee=emp, date=day, shift_end_time=time(17, 30), last_punch_out=time(19, 30), ot_minutes=minutes,
        status=status, compensation_type=ctype if status == "announced" else None,
        announced_by=by if status == "announced" else None,
        announced_at=announced_at if status == "announced" else None,
    )  # fmt: skip


_LATE_SUMMARY = {
    "lateInCount": 4, "earlyOutCount": 0, "excessPermissionCount": 0, "totalLateCount": 4, "freeAllowanceUsed": 3,
    "billableLateCount": 1, "shiftDeductions": 0.5, "freeAllowance": 3, "permissionMonthlyCap": 3,
}  # fmt: skip


def _hand_slip(emp, *, gross="20000", pf="1200", esi="150", adv="0", late="0", ot="0", year=YEAR, month=MONTH,
               breakdown=True, late_summary=None, period=None, payroll=True, payroll_final=None, ot_days=0):  # fmt: skip
    """A slip typed by hand (older / seed-shaped data the engine would not produce today)."""
    gross, pf, esi, adv, late, ot = (Decimal(x) for x in (gross, pf, esi, adv, late, ot))
    total = pf + esi + adv + late
    net = gross + ot - total
    prod = period is not None
    bd = None
    if breakdown:
        bd = {
            "type": "production" if prod else "staff",
            "salaryPerShift": float(emp.salary_per_shift or 0),
            "summary": (
                {"totalDays": 10, "daysWorked": 8, "daysAbsent": 2, "totalShifts": 8.0} if prod else
                {"totalWorkingDays": 24, "effectivePaidDays": 24.0, "halfShiftDays": 0, "absentDays": 0,
                 "unpaidLeaveDays": 0}
            ),
            "earnings": {"monthlySalary": float(emp.salary_amount or 0), "dailyRate": round(float(gross) / 24, 2),
                         "otDays": ot_days, "otAmount": float(ot), "grossSalary": float(gross),
                         "totalShifts": 8.0},
            "deductions": {"pf": float(pf), "esi": float(esi), "esiApplicableBelow": 21000.0,
                           "lateShiftPenalty": float(late), "lateSummary": late_summary},
        }  # fmt: skip
    basic = (gross / 2).quantize(Decimal("0.01"))
    hra = (gross / 5).quantize(Decimal("0.01"))
    start, end = period if prod else (None, None)
    number = f"SS/{emp.employee_code}/{start}_{end}" if prod else f"SS/{emp.employee_code}/{year}/{month:02d}"
    slip = SalarySlip.objects.create(
        employee=emp, month=month, year=year, period_start=start, period_end=end, slip_number=number,
        basic=basic, hra=hra, allowances=gross - basic - hra, ot_amount=ot, gross_salary=gross, pf_deduction=pf,
        esi_deduction=esi, advance_deduction=adv, other_deductions=(late if not prod else Decimal("0")),
        total_deductions=total, net_salary=net, working_days=10 if prod else 24, present_days=8 if prod else 24,
        breakdown_details=bd,
    )  # fmt: skip
    if payroll:
        Payroll.objects.create(
            employee=emp, salary_mode="shift" if prod else "monthly", month=month, year=year, period_start=start,
            period_end=end, total_working_days=10 if prod else 24, present_days=24, absent_days=0, base_salary=gross,
            gross_salary=gross, deductions=total, bonus=0,
            final_salary=Decimal(payroll_final) if payroll_final is not None else net,
        )  # fmt: skip
    return slip


class _Base(TestCase):
    """The shared payroll, generated once per test class through the real engine."""

    @classmethod
    def setUpTestData(cls):
        ps = PayrollSettings.get()
        ps.staff_payroll_rules_enabled = True
        ps.pf_rate, ps.esi_rate, ps.esi_applicable_below = Decimal("12"), Decimal("0.75"), Decimal("21000")
        ps.prod_payroll_rules_enabled = True
        ps.prod_pf_rate, ps.prod_esi_rate = Decimal("12"), Decimal("0.75")
        ps.prod_esi_applicable_below = Decimal("21000")
        ps.prod_pf_ef_enabled = False
        ps.prod_late_detection_enabled = False
        ps.late_free_allowance = 3
        ps.late_deduction_slabs = [{"fromLates": 1, "deductionShifts": 0.5}, {"fromLates": 4, "deductionShifts": 1}]
        ps.permission_monthly_cap = 3
        ps.attendance_mode = "simple"
        ps.compensation_feature_enabled = True
        ps.ot_detection_enabled = True
        ps.min_wage_rate = Decimal("0")
        ps.save()

        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.d_cut = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.d_adm = Department.objects.create(name="ADMIN", branch=cls.b1)
        cls.d_sew = Department.objects.create(name="SEWING", branch=cls.b2)
        cls.d_stitch = Department.objects.create(name="STITCHING", branch=cls.b1)
        cls.desig = Designation.objects.create(title="Operator")
        bank = dict(bank_name="SBI", bank_account="123456789012", bank_ifsc="SBIN0001234")

        cls.A = _staff("3001", "20400", cls.d_cut, cls.b1, designation=cls.desig, uan_number="100000000001",
                       pf_number="TN/A/1", esi_number="1111111111", **bank)  # fmt: skip
        cls.B = _staff("3002", "24000", cls.d_cut, cls.b1, uan_number="100000000002", pf_number="TN/B/2", **bank)
        cls.C = _staff("3003", "15000", cls.d_sew, cls.b2)  # no bank, no statutory numbers
        cls.D = _staff("3004", "60000", cls.d_adm, cls.b1, uan_number="100000000004", pf_number="TN/D/4", **bank)
        cls.L = _staff("3006", "18000", cls.d_cut, cls.b1, uan_number="100000000006", pf_number="TN/L/6",
                       esi_number="2222222222", **bank)  # fmt: skip
        cls.E = _staff("3007", "12000", cls.d_sew, cls.b2, **bank)  # never gets a slip
        cls.F = _staff("3008", None, cls.d_cut, cls.b1)  # no salary configured
        cls.P1 = Employee.objects.create(
            employee_code="P001", first_name="Prod", last_name="One", employment_type="production", status="active",
            salary_per_shift=Decimal("500"), department=cls.d_stitch, branch=cls.b1, uan_number="100000000009",
            pf_number="TN/P/9", **bank,
        )  # fmt: skip

        # A: full month + overtime (2 pay days paid in the slip, 2 relaxation credits, one detected, one rejected)
        _present_all(cls.A)
        _ot(cls.A, date(2026, 2, 2), 120, announced_at=ANNOUNCED)
        _ot(cls.A, date(2026, 2, 3), 90, announced_at=ANNOUNCED)
        r1 = _ot(cls.A, date(2026, 2, 4), 100, ctype="relaxation", announced_at=ANNOUNCED)
        _ot(cls.A, date(2026, 2, 5), 75, status="detected")
        _ot(cls.A, date(2026, 2, 6), 80, status="rejected")
        r2 = _ot(cls.A, date(2026, 2, 7), 95, ctype="relaxation", announced_at=ANNOUNCED)
        cls.credit_avail = CompensationLeaveCredit.objects.create(employee=cls.A, source_overtime_record=r1)
        cls.credit_used = CompensationLeaveCredit.objects.create(
            employee=cls.A, source_overtime_record=r2, status="used", used_date=date(2026, 2, 20)
        )
        _generate_staff_payroll(cls.A, MONTH, YEAR)

        # B: 18 full (5 of them late), 2 half, 2 absent, 2 unpaid leave
        for i, d in enumerate(WORKING_DAYS):
            if i < 18:
                _day(cls.B, d, is_late=i < 5)
            elif i < 20:
                _day(cls.B, d, "half_shift", "0.50")
            elif i < 22:
                _day(cls.B, d, "absent", "0")
            else:
                _day(cls.B, d, "on_leave", "0")
        _generate_staff_payroll(cls.B, MONTH, YEAR)
        _ot(cls.B, date(2026, 2, 10), 60, announced_at=ANNOUNCED)  # announced AFTER the slip: pending

        # C: full month + one advance instalment, and a relaxation credit
        _present_all(cls.C)
        adv = Advance.objects.create(
            employee=cls.C, advance_type="general", amount=Decimal("1000"), status="approved",
            total_repaid=Decimal("0"), outstanding=Decimal("1000"),
        )  # fmt: skip
        AdvanceRepayment.objects.create(advance=adv, month=MONTH, year=YEAR, amount=Decimal("1000"))
        _generate_staff_payroll(cls.C, MONTH, YEAR)
        r3 = _ot(cls.C, date(2026, 2, 11), 70, ctype="relaxation", announced_at=ANNOUNCED)
        cls.credit_c = CompensationLeaveCredit.objects.create(employee=cls.C, source_overtime_record=r3)

        # D: 2 lates + 1 early-out + 5 approved permissions (3 in cap, 2 excess)
        for i, d in enumerate(WORKING_DAYS):
            _day(cls.D, d, is_late=i < 2, early_leave=(i == 2))
        for day in range(2, 7):
            EmployeePermission.objects.create(employee=cls.D, date=date(2026, 2, day), status="approved")
        _generate_staff_payroll(cls.D, MONTH, YEAR)

        # L: full month, then leaves the company
        _present_all(cls.L)
        _generate_staff_payroll(cls.L, MONTH, YEAR)
        Employee.objects.filter(pk=cls.L.pk).update(status="inactive")

        # E: no slip, but an announced pay OT day
        _ot(cls.E, date(2026, 2, 12), 130, announced_at=ANNOUNCED)

        # P1: two production periods in the month
        for i in range(10):
            _day(cls.P1, date(2026, 2, 2 + i), shifts="1.50" if i < 5 else "1.00")
        _generate_production_payroll(cls.P1, *PROD1)
        for d in range(16, 20):
            _day(cls.P1, date(2026, 2, d))
        _generate_production_payroll(cls.P1, *PROD2)

        # credits: created_at is auto_now_add, so pin it to fixed instants
        CompensationLeaveCredit.objects.filter(pk__in=[cls.credit_avail.pk, cls.credit_used.pk]).update(
            created_at=ANNOUNCED
        )
        CompensationLeaveCredit.objects.filter(pk=cls.credit_c.pk).update(
            created_at=datetime(2026, 3, 10, 5, 0, tzinfo=timezone.utc)
        )

        # compensation day announcements
        made = datetime(2026, 2, 1, 4, 30, tzinfo=timezone.utc)
        cls.ann_all = CompensationDayAnnouncement.objects.create(
            date=date(2026, 2, 14), reason="Festival", announced_by="HR One"
        )
        cls.ann_b1 = CompensationDayAnnouncement.objects.create(
            date=date(2026, 2, 20),
            leave_until_time=time(15, 0),
            branch=cls.b1,
            reason="Unit 1 early release",
            announced_by="HR One",
        )
        cls.ann_b2 = CompensationDayAnnouncement.objects.create(
            date=date(2026, 2, 25),
            leave_until_time=time(14, 0),
            branch=cls.b2,
            department=cls.d_sew,
            reason="Sewing special",
            announced_by="HR Two",
        )
        cls.ann_b2.employees.set([cls.C, cls.E])
        cls.ann_dept = CompensationDayAnnouncement.objects.create(
            date=date(2026, 3, 5), department=cls.d_cut, reason="Cutting day off", announced_by="HR One"
        )
        cls.ann_named = CompensationDayAnnouncement.objects.create(
            date=date(2026, 2, 27), reason="Named employee", announced_by="HR Two"
        )
        cls.ann_named.employees.set([cls.E])
        CompensationDayAnnouncement.objects.update(created_at=made)

        # HR users
        cls.admin = HRUser.objects.create(username="g2_admin", password_hash="x", is_super_admin=True)

        def mk(name, perms, branch=None):
            role = Role.objects.create(name=name, permissions=perms)
            return HRUser.objects.create(username=name, password_hash="x", role=role, branch=branch)

        full = {"reports": "view", "payroll": "view", "production_payroll": "view", "compensation": "view"}
        cls.u_b1 = mk("g2_b1", full, cls.b1)
        cls.u_b2 = mk("g2_b2", full, cls.b2)
        cls.u_reports_only = mk("g2_plain", {"reports": "view"})
        cls.u_prod_only = mk("g2_prod", {"reports": "view", "production_payroll": "view"})
        cls.u_staff_only = mk("g2_staff", {"reports": "view", "payroll": "view"})
        cls.u_comp_only = mk("g2_comp", {"reports": "view", "compensation": "view"})

    # -- helpers -----------------------------------------------------------------------------------
    def run_report(self, report_id, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/run/{report_id}", params, **_headers(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, report_id, fmt, user=None, **params):
        return self.client.get(
            f"/api/reports/export/{report_id}", {"fmt": fmt, **params}, **_headers(user or self.admin)
        )

    @staticmethod
    def data_rows(body):
        return [r for r in body["rows"] if r.get("_kind") is None]

    def by_code(self, body):
        return {r["employeeCode"]: r for r in self.data_rows(body)}

    def codes(self, body):
        return [r["employeeCode"] for r in self.data_rows(body)]

    @staticmethod
    def summary(body):
        return {s["label"]: s["value"] for s in body["summary"]}

    def today(self, d):
        """Pin the report clock (ReportContext.today and the filter defaults) to a fixed date."""
        return mock.patch("api.reporting.filters.ist_today", return_value=d)

    def assert_totals(self, body, keys):
        rows = self.data_rows(body)
        for k in keys:
            expected = round(sum((r[k] or 0) for r in rows), 2)
            self.assertAlmostEqual(body["totals"][k], expected, places=2, msg=k)

    def assert_subtotals(self, body, keys):
        """Each subtotal row equals the sum of the data rows above it (up to the previous subtotal)."""
        group = []
        seen = 0
        for r in body["rows"]:
            if r.get("_kind") == "subtotal":
                seen += 1
                for k in keys:
                    vals = [g[k] for g in group if g[k] is not None]
                    expect = round(sum(vals), 2) if vals else None
                    self.assertEqual(r[k], expect, f"{k} in {r['employeeName']}")
                group = []
            else:
                group.append(r)
        return seen

    def assert_flat_queries(self, report_id, grow, params=None, user=None, slack=3):
        """No N+1: query count with the fixture vs with many more rows may differ by a small constant only."""
        params = params or DATA_PARAMS[report_id]
        self.run_report(report_id, user=user, **params)  # warm caches (content types, settings)
        with CaptureQueriesContext(connection) as small:
            few = self.run_report(report_id, user=user, **params)
        grow()
        with CaptureQueriesContext(connection) as big:
            many = self.run_report(report_id, user=user, **params)
        self.assertGreaterEqual(many["rowCount"], few["rowCount"] + 12, f"{report_id}: the extra data must add rows")
        self.assertLessEqual(
            abs(len(big) - len(small)), slack, f"{report_id}: {len(small)} queries -> {len(big)} queries"
        )

    def extras(self, n, start=5000, **slip_kw):
        """n more staff employees with hand-typed slips and (mismatching) payroll rows."""
        out = []
        for i in range(n):
            emp = _staff(str(start + i), "20000", self.d_cut, self.b1)
            slip_kw.setdefault("late", "416.67")
            slip_kw.setdefault("late_summary", _LATE_SUMMARY)
            _hand_slip(emp, payroll_final="19000", **slip_kw)
            out.append(emp)
        return out


# ================================================================================================
#  registration
# ================================================================================================
class RegistrationTests(TestCase):
    def test_all_twelve_reports_registered_with_real_modules_and_categories(self):
        specs = {s.id: s for s in registry.all_specs()}
        self.assertEqual([], [k for k in registry.LOAD_ERRORS if "payroll_statutory" in k])
        for rid in G2_IDS:
            self.assertIn(rid, specs)
            s = specs[rid]
            self.assertEqual(s.category, "payroll", rid)
            self.assertTrue(s.modules, f"{rid} must name its owning modules")
            for m in s.modules:
                self.assertIn(m, all_module_keys(), rid)
            self.assertFalse(s.super_admin_only)
            self.assertTrue(s.description and s.title)
            self.assertTrue(s.columns, rid)

    def test_overtime_family_has_two_variants(self):
        specs = {s.id: s for s in registry.all_specs()}
        self.assertEqual(
            (specs["overtime-register"].family, specs["overtime-register"].variant), ("overtime", "Register")
        )
        self.assertEqual(
            (specs["overtime-payment-summary"].family, specs["overtime-payment-summary"].variant),
            ("overtime", "Summary"),
        )

    def test_money_and_id_columns_use_the_right_types(self):
        spec = registry.get_spec("pf-statement")
        types = {c.key: c.type for c in spec.columns}
        self.assertEqual(types["uanNumber"], "text")  # identifiers stay text: no float rounding of a 12-digit UAN
        self.assertEqual(types["pfNumber"], "text")
        self.assertEqual(types["eeShare"], "currency")
        self.assertEqual(registry.get_spec("esi-statement").columns[2].type, "text")


# ================================================================================================
#  PF statement
# ================================================================================================
class PfStatementTests(_Base):
    def test_golden_rows_ceiling_15000(self):
        body = self.run_report("pf-statement", **P)
        rows = self.by_code(body)
        self.assertEqual(list(rows), ["3001", "3002", "3003", "3004", "3006", "P001"])  # incl. the employee who left
        expect = {  # pfWages, epfWages, eeShare, epsShare, erEpfShare, erTotal, eeVariance, remittance, days, ncp, gross
            "3001": (10200, 10200, 1224, 850, 374, 1224, 0, 2448, 28, 0, 22100),
            "3002": (9500, 9500, 1140, 791, 349, 1140, 0, 2280, 23, 5, 19000),
            "3003": (7500, 7500, 900, 625, 275, 900, 0, 1800, 28, 0, 15000),
            "3004": (30000, 15000, 3600, 1250, 550, 1800, 1800, 3600, 28, 0, 60000),
            "3006": (9000, 9000, 1080, 750, 330, 1080, 0, 2160, 28, 0, 18000),
            "P001": (8250, 8250, 990, 687, 303, 990, 0, 1980, 14, 6, 8250),
        }
        keys = ("pfWages", "epfWages", "eeShare", "epsShare", "erEpfShare", "erTotal", "eeVariance",
                "totalRemittance", "daysPaid", "ncpDays", "grossWages")  # fmt: skip
        for code, vals in expect.items():
            self.assertEqual(tuple(rows[code][k] for k in keys), tuple(float(v) for v in vals), code)
        self.assertEqual(rows["3001"]["uanNumber"], "100000000001")
        self.assertEqual(rows["3001"]["employmentType"], "Staff")
        self.assertEqual(rows["P001"]["employmentType"], "Production")
        self.assertEqual(rows["3003"]["idCheck"], "Missing UAN, PF no.")
        self.assertEqual(rows["3001"]["idCheck"], "OK")

    def test_totals_and_summary_equal_the_rows(self):
        body = self.run_report("pf-statement", **P)
        self.assertEqual(
            body["totals"],
            {"grossWages": 142350.0, "pfWages": 74450.0, "epfWages": 59450.0, "epsWages": 59450.0, "eeShare": 8934.0,
             "epsShare": 4953.0, "erEpfShare": 2181.0, "erTotal": 7134.0, "eeVariance": 1800.0,
             "totalRemittance": 14268.0},
        )  # fmt: skip
        self.assert_totals(body, ["grossWages", "pfWages", "epfWages", "eeShare", "epsShare", "erEpfShare", "erTotal"])
        s = self.summary(body)
        self.assertEqual(s["PF covered employees"], 6)
        self.assertEqual(s["EE PF deducted"], body["totals"]["eeShare"])
        self.assertEqual(s["ER EPS (est.)"], body["totals"]["epsShare"])
        self.assertEqual(s["ER EPF (est.)"], body["totals"]["erEpfShare"])
        self.assertEqual(s["Statutory remittance (est.)"], body["totals"]["totalRemittance"])
        self.assertEqual(s["Covered, UAN missing"], 1)
        # EE = statutory 12% (ER total); the ER total splits exactly into EPS + EPF diff
        self.assertEqual(body["totals"]["erTotal"], body["totals"]["epsShare"] + body["totals"]["erEpfShare"])
        self.assertEqual(round(body["totals"]["eeShare"] - body["totals"]["erTotal"], 2), body["totals"]["eeVariance"])

    def test_pf_deducted_equals_the_slips(self):
        body = self.run_report("pf-statement", **P)
        slips = sum(float(s.pf_deduction) for s in SalarySlip.objects.filter(month=MONTH, year=YEAR))
        self.assertEqual(body["totals"]["eeShare"], round(slips, 2))

    def test_uncapped_wage_basis(self):
        body = self.run_report("pf-statement", wageCeiling="none", **P)
        d = self.by_code(body)["3004"]
        self.assertEqual((d["epfWages"], d["epsWages"]), (30000.0, 15000.0))
        self.assertEqual((d["erTotal"], d["epsShare"], d["erEpfShare"], d["eeVariance"]), (3600.0, 1250.0, 2350.0, 0.0))
        self.assertTrue(any("no cap" in n for n in body["notes"]))

    def test_coverage_all_lists_uncovered_without_employer_share(self):
        ps = PayrollSettings.get()
        ps.staff_payroll_rules_enabled = False
        ps.save()
        emp = _staff("3010", "20000", self.d_cut, self.b1)
        _present_all(emp)
        _generate_staff_payroll(emp, MONTH, YEAR)  # rules off: PF 0
        self.assertEqual(SalarySlip.objects.get(employee=emp).pf_deduction, 0)
        self.assertNotIn("3010", self.codes(self.run_report("pf-statement", **P)))
        row = self.by_code(self.run_report("pf-statement", coverage="all", **P))["3010"]
        self.assertEqual(row["eeShare"], 0.0)
        for k in (
            "pfWages",
            "epfWages",
            "epsWages",
            "epsShare",
            "erEpfShare",
            "erTotal",
            "totalRemittance",
            "eeVariance",
        ):
            self.assertIsNone(row[k], k)  # no PF deducted: no PF wages and no employer share are shown
        self.assertEqual(row["grossWages"], 20000.0)  # the gross wages stay visible for the audit view
        self.assertIsNone(row["idCheck"])
        # the wage totals and the summary cards still describe the covered people only
        body = self.run_report("pf-statement", coverage="all", **P)
        self.assertEqual(self.summary(body)["PF wages"], body["totals"]["pfWages"])
        self.assertEqual(self.summary(body)["EPF wages"], body["totals"]["epfWages"])
        self.assertEqual(body["totals"]["pfWages"], 74450.0)  # unchanged by the uncovered employee

    def test_filters_narrow(self):
        def codes(**q):
            return self.codes(self.run_report("pf-statement", **{**P, **q}))

        self.assertEqual(codes(departmentIds=str(self.d_sew.id)), ["3003"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), ["3003"])
        self.assertEqual(codes(employmentType="production"), ["P001"])
        self.assertEqual(codes(employmentType="staff"), ["3001", "3002", "3003", "3004", "3006"])
        self.assertEqual(codes(employeeIds=f"{self.A.id},{self.D.id}"), ["3001", "3004"])
        self.assertEqual(codes(designationIds=str(self.desig.id)), ["3001"])
        self.assertEqual(codes(period="2026-01"), [])

    def test_production_slips_of_a_month_are_summed_per_employee(self):
        rows = self.data_rows(self.run_report("pf-statement", **P))
        self.assertEqual(len([r for r in rows if r["employeeCode"] == "P001"]), 1)
        p = self.by_code(self.run_report("pf-statement", **P))["P001"]
        self.assertEqual((p["pfWages"], p["eeShare"], p["daysPaid"], p["ncpDays"]), (8250.0, 990.0, 14.0, 6.0))

    def test_employee_without_department_or_designation(self):
        emp = _staff("3011", "20000", None, None)
        _hand_slip(emp)
        row = self.by_code(self.run_report("pf-statement", **P))["3011"]
        self.assertEqual(row["department"], "Unassigned")
        self.assertIsNone(row["uanNumber"])
        self.assertEqual(row["idCheck"], "Missing UAN, PF no.")
        # a branch-less employee is invisible to a branch-limited login
        self.assertNotIn("3011", self.codes(self.run_report("pf-statement", user=self.u_b1, **P)))

    def test_slip_without_breakdown_still_reports(self):
        emp = _staff("3012", "20000", self.d_cut, self.b1, uan_number="1", pf_number="P/1")
        _hand_slip(emp, breakdown=False)
        row = self.by_code(self.run_report("pf-statement", **P))["3012"]
        self.assertEqual((row["pfWages"], row["eeShare"]), (10000.0, 1200.0))
        self.assertEqual((row["daysPaid"], row["ncpDays"]), (28.0, 0.0))  # falls back to the slip's absent days

    def test_duplicate_staff_slips_are_not_double_counted(self):
        emp = _staff("3013", "20000", self.d_cut, self.b1, uan_number="1", pf_number="P/1")
        _hand_slip(emp, payroll=False)
        SalarySlip.objects.create(  # a second staff slip for the same month (NULL week_number defeats the DB unique key)
            employee=emp, month=MONTH, year=YEAR, slip_number="SS/3013/2026/02-dup", basic=Decimal("5000"),
            pf_deduction=Decimal("600"), gross_salary=Decimal("10000"), net_salary=Decimal("9400"),
            total_deductions=Decimal("600"), working_days=24,
        )  # fmt: skip
        body = self.run_report("pf-statement", **P)
        row = self.by_code(body)["3013"]
        self.assertEqual((row["pfWages"], row["eeShare"]), (5000.0, 600.0))  # the newest slip only
        self.assertTrue(any("duplicate staff slip" in n for n in body["notes"]))

    def test_branch_isolation(self):
        body = self.run_report("pf-statement", user=self.u_b1, **P)
        self.assertEqual(self.codes(body), ["3001", "3002", "3004", "3006", "P001"])
        widened = self.run_report("pf-statement", user=self.u_b1, branchIds=str(self.b2.id), **P)
        self.assertEqual(widened["rows"], [])
        other = self.run_report("pf-statement", user=self.u_b1, employeeIds=str(self.C.id), **P)
        self.assertEqual(other["rows"], [])
        self.assertEqual(self.codes(self.run_report("pf-statement", user=self.u_b2, **P)), ["3003"])

    def test_role_without_the_owning_module_is_forbidden(self):
        r = self.client.get("/api/reports/run/pf-statement", P, **_headers(self.u_reports_only))
        self.assertEqual(r.status_code, 403)
        self.assertEqual(r.json()["error"], "report_forbidden")

    def test_production_only_role_sees_only_production_rows(self):
        self.assertEqual(self.codes(self.run_report("pf-statement", user=self.u_prod_only, **P)), ["P001"])
        self.assertEqual(
            self.codes(self.run_report("pf-statement", user=self.u_staff_only, **P)),
            ["3001", "3002", "3003", "3004", "3006"],
        )

    def test_xlsx_round_trip_keeps_identifiers_as_text(self):
        r = self.export("pf-statement", "xlsx", **P)
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:5], ["Emp Code", "Employee", "UAN", "PF No.", "Department"])
        first = [c.value for c in ws[8]]
        self.assertEqual(first[0], "3001")
        self.assertEqual(first[2], "100000000001")
        self.assertEqual(ws["C8"].data_type, "s")
        self.assertEqual(ws["M8"].value, 1224.0)  # EE PF deducted
        total = [c.value for c in ws[14]]
        self.assertEqual(total[0], "TOTAL")
        self.assertEqual(total[12], 8934.0)

    def test_pdf_export(self):
        r = self.export("pf-statement", "pdf", **P)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.assert_flat_queries("pf-statement", lambda: self.extras(12))


# ================================================================================================
#  ESI statement
# ================================================================================================
class EsiStatementTests(_Base):
    def test_golden_rows(self):
        body = self.run_report("esi-statement", **P)
        rows = self.by_code(body)
        self.assertEqual(list(rows), ["3001", "3003", "3006", "P001"])
        expect = {  # esiWages, eeShare, erShare, total, days
            "3001": (20400, 153.00, 663.00, 816.00, 28),
            "3003": (15000, 112.50, 487.50, 600.00, 28),
            "3006": (18000, 135.00, 585.00, 720.00, 28),
            "P001": (8250, 61.88, 268.13, 330.01, 14),  # 46.88 + 15.00 ; 3.25% of 8250 = 268.125 -> half up
        }
        for code, (wages, ee, er, total, days) in expect.items():
            r = rows[code]
            self.assertEqual((r["esiWages"], r["eeShare"], r["erShare"], r["totalContribution"], r["daysPaid"]),
                             (float(wages), ee, er, total, float(days)), code)  # fmt: skip
            self.assertEqual(r["coverage"], "Covered")
        self.assertEqual(rows["3001"]["idCheck"], "OK")
        self.assertEqual(rows["3003"]["idCheck"], "Missing IP no.")
        self.assertIsNone(rows["3003"]["esiNumber"])

    def test_totals_and_summary(self):
        body = self.run_report("esi-statement", **P)
        self.assertEqual(
            body["totals"], {"esiWages": 61650.0, "eeShare": 462.38, "erShare": 2003.63, "totalContribution": 2466.01}
        )
        self.assert_totals(body, ["esiWages", "eeShare", "erShare", "totalContribution"])
        s = self.summary(body)
        self.assertEqual(s["ESI covered employees"], 4)
        self.assertEqual(s["Employee contribution"], 462.38)
        self.assertEqual(s["Total contribution"], 2466.01)
        self.assertEqual(s["Above ESI ceiling (exempt)"], 2)  # B (24,000) and D (60,000)
        self.assertEqual(s["Covered, IP number missing"], 2)
        slips = sum(float(x.esi_deduction) for x in SalarySlip.objects.filter(month=MONTH, year=YEAR))
        self.assertEqual(body["totals"]["eeShare"], round(slips, 2))

    def test_coverage_filter(self):
        exempt = self.run_report("esi-statement", coverage="exempt", **P)
        rows = self.by_code(exempt)
        self.assertEqual(list(rows), ["3002", "3004"])
        for r in rows.values():
            self.assertEqual((r["coverage"], r["eeShare"]), ("Above ceiling", 0.0))
            self.assertIsNone(r["erShare"])  # nobody contributes: no employer share
        self.assertEqual(rows["3002"]["grossSalary"], 19000.0)  # the salary is still shown ...
        self.assertIsNone(rows["3002"]["esiWages"])  # ... but no ESI wages: nothing is contributed on it
        allrows = self.run_report("esi-statement", coverage="all", **P)
        self.assertEqual(self.codes(allrows), ["3001", "3002", "3003", "3004", "3006", "P001"])
        self.assertEqual(allrows["totals"]["erShare"], 2003.63)  # only covered rows carry an employer share
        # totals row and summary cards describe the same (covered) people whatever the coverage filter shows
        for cov, wages in (("covered", 61650.0), ("exempt", None), ("all", 61650.0)):
            body = self.run_report("esi-statement", coverage=cov, **P)
            self.assertEqual(body["totals"]["esiWages"], wages, cov)  # a dash, not a zero, when nobody contributes
            self.assertEqual(self.summary(body)["ESI wages"], wages or 0.0, cov)

    def test_not_covered_when_rules_were_off(self):
        ps = PayrollSettings.get()
        ps.staff_payroll_rules_enabled = False
        ps.save()
        emp = _staff("3010", "20000", self.d_cut, self.b1)
        _present_all(emp)
        _generate_staff_payroll(emp, MONTH, YEAR)
        row = self.by_code(self.run_report("esi-statement", coverage="exempt", **P))["3010"]
        self.assertEqual(row["coverage"], "Not covered")  # under the ceiling, but nothing was deducted

    def test_filters_narrow(self):
        def codes(**q):
            return self.codes(self.run_report("esi-statement", **{**P, **q}))

        self.assertEqual(codes(departmentIds=str(self.d_sew.id)), ["3003"])
        self.assertEqual(codes(employmentType="production"), ["P001"])
        self.assertEqual(codes(branchIds=str(self.b1.id)), ["3001", "3006", "P001"])
        self.assertEqual(codes(employeeIds=str(self.L.id)), ["3006"])
        self.assertEqual(codes(period="2026-03"), [])

    def test_branch_isolation(self):
        self.assertEqual(self.codes(self.run_report("esi-statement", user=self.u_b1, **P)), ["3001", "3006", "P001"])
        self.assertEqual(self.run_report("esi-statement", user=self.u_b1, branchIds=str(self.b2.id), **P)["rows"], [])
        self.assertEqual(self.codes(self.run_report("esi-statement", user=self.u_b2, **P)), ["3003"])

    def test_production_only_role_sees_only_production(self):
        self.assertEqual(self.codes(self.run_report("esi-statement", user=self.u_prod_only, **P)), ["P001"])

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("esi-statement", "xlsx", **P).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Emp Code", "Employee", "ESI IP No."])
        self.assertEqual(ws["A8"].value, "3001")
        self.assertEqual(ws["C8"].value, "1111111111")
        self.assertEqual(ws["G8"].value, 20400.0)  # gross salary
        self.assertEqual(ws["H8"].value, 20400.0)  # ESI wages
        self.assertEqual(ws["I8"].value, 153.0)  # employee share
        self.assertTrue(self.export("esi-statement", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.assert_flat_queries("esi-statement", lambda: self.extras(12))


# ================================================================================================
#  PF / ESI coverage register
# ================================================================================================
class StatutoryCoverageTests(_Base):
    def test_default_lists_active_employees_with_coverage(self):
        body = self.run_report("statutory-coverage")
        rows = self.by_code(body)
        self.assertEqual(list(rows), ["3001", "3002", "3003", "3004", "3007", "3008", "P001"])  # L left: not active
        a = rows["3001"]
        self.assertEqual(
            (a["pfCovered"], a["esiCovered"], a["salaryBasis"], a["esiCeiling"]), ("Yes", "Yes", 20400.0, 21000.0)
        )
        self.assertEqual(a["lastSlipPeriod"], "February 2026")
        self.assertEqual(rows["3002"]["esiCovered"], "Above ceiling")
        self.assertEqual(rows["3003"]["idCheck"], "Missing UAN, PF no., ESI no.")
        self.assertEqual((rows["3007"]["pfCovered"], rows["3007"]["esiCovered"]), ("No slip", "No slip"))
        self.assertEqual(rows["3007"]["salaryBasis"], 12000.0)  # the profile salary, even without a slip
        p = rows["P001"]
        self.assertEqual(
            (p["salaryBasis"], p["lastSlipPeriod"]), (4000.0, "16-Feb-2026 to 25-Feb-2026")
        )  # latest slip: gross 2000 x 2
        self.assertEqual(p["idCheck"], "Missing ESI no.")
        s = self.summary(body)
        self.assertEqual(
            (s["Employees"], s["PF covered"], s["ESI covered"], s["Covered by neither"], s["No slip yet"],
             s["Covered, IDs missing"], s["Within 10% of ESI ceiling"]),
            (7, 5, 3, 0, 2, 2, 1),
        )  # fmt: skip

    def test_coverage_filter(self):
        def codes(cov, **q):
            return self.codes(self.run_report("statutory-coverage", coverage=cov, **q))

        self.assertEqual(codes("pf"), ["3001", "3002", "3003", "3004", "P001"])
        self.assertEqual(codes("esi"), ["3001", "3003", "P001"])
        self.assertEqual(codes("missing-ids"), ["3003", "P001"])
        self.assertEqual(codes("no-slip"), ["3007", "3008"])
        self.assertEqual(codes("neither"), [])
        emp = _staff("3010", "30000", self.d_cut, self.b1)
        _hand_slip(emp, pf="0", esi="0")
        self.assertEqual(codes("neither"), ["3010"])

    def test_status_and_scope_filters(self):
        self.assertEqual(
            self.codes(self.run_report("statutory-coverage", employeeStatus="all")),
            ["3001", "3002", "3003", "3004", "3006", "3007", "3008", "P001"],
        )
        self.assertEqual(self.codes(self.run_report("statutory-coverage", employeeStatus="inactive")), ["3006"])
        self.assertEqual(
            self.codes(self.run_report("statutory-coverage", departmentIds=str(self.d_sew.id))), ["3003", "3007"]
        )
        self.assertEqual(self.codes(self.run_report("statutory-coverage", employmentType="production")), ["P001"])

    def test_notes_state_the_payroll_rule_switches(self):
        body = self.run_report("statutory-coverage")
        self.assertTrue(any("staff PF/ESI rules ON" in n for n in body["notes"]))
        ps = PayrollSettings.get()
        ps.staff_payroll_rules_enabled = False
        ps.save()
        self.assertTrue(any("staff PF/ESI rules OFF" in n for n in self.run_report("statutory-coverage")["notes"]))

    def test_branch_isolation(self):
        body = self.run_report("statutory-coverage", user=self.u_b1)
        self.assertEqual(self.codes(body), ["3001", "3002", "3004", "3008", "P001"])
        self.assertEqual(self.run_report("statutory-coverage", user=self.u_b1, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(self.codes(self.run_report("statutory-coverage", user=self.u_b2)), ["3003", "3007"])

    def test_production_only_role(self):
        self.assertEqual(self.codes(self.run_report("statutory-coverage", user=self.u_prod_only)), ["P001"])

    def test_latest_slip_is_chosen_across_months(self):
        _hand_slip(self.A, month=3, year=2026, pf="0", esi="0", payroll=False)
        row = self.by_code(self.run_report("statutory-coverage"))["3001"]
        self.assertEqual(row["lastSlipPeriod"], "March 2026")
        self.assertEqual((row["pfCovered"], row["esiCovered"]), ("No", "No"))  # coverage follows the newest slip

    def test_no_n_plus_one(self):
        self.assert_flat_queries("statutory-coverage", lambda: self.extras(12))


# ================================================================================================
#  Late detection - salary impact
# ================================================================================================
class LateSalaryImpactTests(_Base):
    def test_golden_rows_and_subtotals(self):
        body = self.run_report("late-salary-impact", **P)
        rows = self.by_code(body)
        self.assertEqual(list(rows), ["3004", "3002"])  # department order: ADMIN, CUTTING
        d, b = rows["3004"], rows["3002"]
        self.assertEqual(
            (d["lateInCount"], d["earlyOutCount"], d["excessPermissionCount"], d["totalLateCount"], d["freeAllowance"],
             d["freeUsed"], d["billableLateCount"], d["shiftDeductions"], d["unitRate"], d["lateDeductionAmount"],
             d["pctOfEarnings"], d["netPay"]),
            (2, 1, 2, 5, 3, 3, 2, 0.5, 2500.0, 1250.0, 2.08, 55150.0),
        )  # fmt: skip
        self.assertEqual(
            (b["lateInCount"], b["earlyOutCount"], b["excessPermissionCount"], b["totalLateCount"], b["billableLateCount"],
             b["shiftDeductions"], b["unitRate"], b["lateDeductionAmount"], b["pctOfEarnings"], b["netPay"]),
            (5, 0, 0, 5, 2, 0.5, 1000.0, 500.0, 2.63, 17360.0),
        )  # fmt: skip
        self.assertEqual((d["employmentType"], d["period"], b["department"]), ("Staff", "February 2026", "CUTTING"))
        self.assertEqual(
            self.assert_subtotals(
                body, ["totalLateCount", "billableLateCount", "shiftDeductions", "lateDeductionAmount"]
            ),
            2,
        )

    def test_totals_summary_and_agreement_with_the_slips(self):
        body = self.run_report("late-salary-impact", **P)
        self.assertEqual(
            body["totals"],
            {"lateInCount": 7, "earlyOutCount": 1, "excessPermissionCount": 2, "totalLateCount": 10,
             "billableLateCount": 4, "shiftDeductions": 1.0, "lateDeductionAmount": 1750.0, "netPay": 72510.0},
        )  # fmt: skip
        self.assert_totals(body, ["totalLateCount", "billableLateCount", "shiftDeductions", "lateDeductionAmount"])
        s = self.summary(body)
        self.assertEqual(
            (s["Employees with lates"], s["Late occurrences"], s["Billable occurrences"], s["Shifts deducted"],
             s["Late deduction"], s["Average per deducted employee"]),
            (2, 10, 4, 1.0, 1750.0, 875.0),
        )  # fmt: skip
        # what the slips actually took for lateness (other_deductions is the staff late penalty)
        taken = sum(float(x.other_deductions) for x in SalarySlip.objects.filter(month=MONTH, year=YEAR))
        self.assertEqual(body["totals"]["lateDeductionAmount"], round(taken, 2))

    def test_min_lates_filter(self):
        def codes(n):
            return self.codes(self.run_report("late-salary-impact", minLates=n, **P))

        self.assertEqual(codes("1"), ["3004", "3002"])
        self.assertEqual(codes("5"), ["3004", "3002"])
        self.assertEqual(codes("10"), [])
        self.assertEqual(
            codes("all"), ["3004", "3001", "3002", "3006", "3003", "P001", "P001"]
        )  # every slip (P001 has two)

    def test_all_slips_show_a_dash_for_production_without_late_data(self):
        body = self.run_report("late-salary-impact", minLates="all", **P)
        p = self.by_code(body)["P001"]
        for k in ("lateInCount", "totalLateCount", "billableLateCount", "shiftDeductions", "lateDeductionAmount",
                  "pctOfEarnings", "freeAllowance"):  # fmt: skip
            self.assertIsNone(p[k], k)  # unknown, not zero
        self.assertEqual(p["period"], "16-Feb-2026 to 25-Feb-2026")
        a = self.by_code(body)["3001"]
        self.assertEqual((a["totalLateCount"], a["lateDeductionAmount"]), (0, 0.0))  # known zero for staff
        stitching = [
            r for r in body["rows"] if r.get("_kind") == "subtotal" and r["employeeName"] == "STITCHING total"
        ][0]
        self.assertIsNone(stitching["totalLateCount"])  # a group with nothing known stays a dash
        self.assertIsNone(stitching["lateDeductionAmount"])

    def test_only_deducted(self):
        self.assertEqual(
            self.codes(self.run_report("late-salary-impact", minLates="all", onlyDeducted="true", **P)),
            ["3004", "3002"],
        )
        self.assertEqual(
            self.codes(self.run_report("late-salary-impact", minLates="all", onlyDeducted="true", period="2026-01")), []
        )

    def test_production_late_deduction_is_read_from_the_breakdown(self):
        emp = Employee.objects.create(
            employee_code="P002",
            first_name="Prod",
            last_name="Two",
            employment_type="production",
            status="active",
            salary_per_shift=Decimal("500"),
            department=self.d_stitch,
            branch=self.b1,
        )
        slip = _hand_slip(
            emp,
            gross="4000",
            pf="0",
            esi="0",
            late="125",
            period=(date(2026, 2, 2), date(2026, 2, 15)),
            late_summary={"totalLateCount": 4, "billableLateCount": 1, "shiftDeductions": 0.25},
        )
        self.assertEqual(slip.other_deductions, 0)  # production keeps the penalty inside total_deductions only
        r = self.by_code(self.run_report("late-salary-impact", **P))["P002"]
        self.assertEqual(
            (r["employmentType"], r["period"], r["totalLateCount"], r["billableLateCount"], r["shiftDeductions"],
             r["unitRate"], r["lateDeductionAmount"], r["lateInCount"], r["freeAllowance"]),
            ("Production", "02-Feb-2026 to 15-Feb-2026", 4, 1, 0.25, 500.0, 125.0, None, None),
        )  # fmt: skip

    def test_older_slips_without_late_data_fall_back_to_the_stored_penalty(self):
        emp = _staff("3020", "20000", self.d_cut, self.b1)
        _hand_slip(emp, late="300", breakdown=False)
        emp2 = _staff("3021", "20000", self.d_cut, self.b1)
        _hand_slip(emp2, late="0", late_summary=None)  # breakdown without a late summary
        rows = self.by_code(self.run_report("late-salary-impact", **P))
        r = rows["3020"]
        self.assertEqual((r["lateDeductionAmount"], r["totalLateCount"], r["unitRate"]), (300.0, None, None))
        self.assertNotIn("3021", rows)  # nothing known and nothing deducted
        r2 = self.by_code(self.run_report("late-salary-impact", minLates="all", **P))["3021"]
        self.assertEqual((r2["totalLateCount"], r2["lateDeductionAmount"]), (None, 0.0))

    def test_scope_filters(self):
        def codes(**q):
            return self.codes(self.run_report("late-salary-impact", **{**P, **q}))

        self.assertEqual(codes(departmentIds=str(self.d_adm.id)), ["3004"])
        self.assertEqual(codes(employeeIds=str(self.B.id)), ["3002"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), [])
        self.assertEqual(codes(employmentType="production"), [])
        self.assertEqual(codes(period="2026-01"), [])

    def test_branch_isolation(self):
        self.assertEqual(self.codes(self.run_report("late-salary-impact", user=self.u_b1, minLates="all", **P)),
                         ["3004", "3001", "3002", "3006", "P001", "P001"])  # fmt: skip
        self.assertEqual(
            self.run_report("late-salary-impact", user=self.u_b1, branchIds=str(self.b2.id), minLates="all", **P)[
                "rows"
            ],
            [],
        )
        self.assertEqual(
            self.codes(self.run_report("late-salary-impact", user=self.u_b2, minLates="all", **P)), ["3003"]
        )

    def test_module_gate_and_kind_restriction(self):
        r = self.client.get("/api/reports/run/late-salary-impact", P, **_headers(self.u_reports_only))
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))
        self.assertEqual(
            self.codes(self.run_report("late-salary-impact", user=self.u_prod_only, minLates="all", **P)),
            ["P001", "P001"],
        )
        self.assertEqual(
            self.codes(self.run_report("late-salary-impact", user=self.u_staff_only, **P)), ["3004", "3002"]
        )

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("late-salary-impact", "xlsx", **P).content)).active
        self.assertEqual([c.value for c in ws[7]][:5], ["Emp Code", "Employee", "Department", "Type", "Period"])
        self.assertEqual(ws["A8"].value, "3004")
        self.assertEqual(ws["O8"].value, 1250.0)
        self.assertTrue(self.export("late-salary-impact", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.assert_flat_queries("late-salary-impact", lambda: self.extras(12))


# ================================================================================================
#  Salary detection (loss of pay)
# ================================================================================================
class AttendanceSalaryLossTests(_Base):
    def test_golden_rows(self):
        body = self.run_report("attendance-salary-loss", **P)
        rows = self.by_code(body)
        self.assertEqual(list(rows), ["3004", "3002"])  # anyone with LOP or a late deduction
        b = rows["3002"]
        self.assertEqual(
            (b["monthlySalary"], b["workingDays"], b["paidDays"], b["absentDays"], b["unpaidLeaveDays"], b["halfDayCount"],
             b["halfDayLossDays"], b["lopDays"], b["dailyRate"], b["lopAmount"], b["lateDeduction"],
             b["totalSalaryImpact"], b["impactPct"], b["monthStatus"]),
            (24000.0, 24, 19.0, 2.0, 2.0, 2, 1.0, 5.0, 1000.0, 5000.0, 500.0, 5500.0, 22.92, "Complete"),
        )  # fmt: skip
        d = rows["3004"]
        self.assertEqual(
            (d["lopDays"], d["lopAmount"], d["lateDeduction"], d["totalSalaryImpact"]), (0.0, 0.0, 1250.0, 1250.0)
        )

    def test_loss_of_pay_equals_salary_minus_slip_gross(self):
        body = self.run_report("attendance-salary-loss", impact="all", **P)
        self.assertEqual(self.codes(body), ["3004", "3001", "3002", "3006", "3003"])  # every staff slip, no production
        for row in self.data_rows(body):
            slip = SalarySlip.objects.get(employee__employee_code=row["employeeCode"], month=MONTH, year=YEAR)
            self.assertEqual(
                row["lopAmount"], float(Decimal(str(row["monthlySalary"])) - slip.gross_salary), row["employeeCode"]
            )
            self.assertEqual(row["lateDeduction"], float(slip.other_deductions))

    def test_totals_summary_and_subtotals(self):
        body = self.run_report("attendance-salary-loss", **P)
        self.assertEqual(
            body["totals"],
            {"absentDays": 2.0, "unpaidLeaveDays": 2.0, "halfDayCount": 2, "halfDayLossDays": 1.0, "lopDays": 5.0,
             "lopAmount": 5000.0, "lateDeduction": 1750.0, "totalSalaryImpact": 6750.0},
        )  # fmt: skip
        self.assert_totals(body, ["absentDays", "lopDays", "lopAmount", "lateDeduction", "totalSalaryImpact"])
        keys = [
            "absentDays",
            "unpaidLeaveDays",
            "halfDayCount",
            "lopDays",
            "lopAmount",
            "lateDeduction",
            "totalSalaryImpact",
        ]
        self.assertEqual(self.assert_subtotals(body, keys), 2)
        s = self.summary(body)
        self.assertEqual(
            (s["Loss of pay"], s["Late deduction"], s["Total salary impact"], s["Employees affected"], s["Absent days"],
             s["Unpaid leave days"], s["Half-days"]),
            (5000.0, 1750.0, 6750.0, 2, 2.0, 2.0, 2),
        )  # fmt: skip

    def test_impact_filter(self):
        def codes(kind):
            return self.codes(self.run_report("attendance-salary-loss", impact=kind, **P))

        self.assertEqual(codes("absent"), ["3002"])
        self.assertEqual(codes("half"), ["3002"])
        self.assertEqual(codes("unpaid"), ["3002"])
        self.assertEqual(codes("late"), ["3004", "3002"])
        self.assertEqual(codes("all"), ["3004", "3001", "3002", "3006", "3003"])

    def test_production_is_excluded_and_explained(self):
        body = self.run_report("attendance-salary-loss", impact="all", **P)
        self.assertNotIn("P001", self.codes(body))
        self.assertTrue(any("2 production slip(s)" in n for n in body["notes"]))

    def test_provisional_month_is_flagged(self):
        with self.today(date(2026, 2, 15)):
            body = self.run_report("attendance-salary-loss", **P)
        self.assertTrue(all(r["monthStatus"] == "Provisional" for r in self.data_rows(body)))
        self.assertTrue(any("has not ended" in n for n in body["notes"]))
        with self.today(date(2026, 3, 1)):
            self.assertTrue(
                all(
                    r["monthStatus"] == "Complete"
                    for r in self.data_rows(self.run_report("attendance-salary-loss", **P))
                )
            )

    def test_slip_without_breakdown_uses_the_profile_salary(self):
        emp = _staff("3020", "20000", self.d_cut, self.b1)
        slip = _hand_slip(emp, gross="15000", breakdown=False, payroll=False)
        SalarySlip.objects.filter(pk=slip.pk).update(absent_days=Decimal("6"), unpaid_leave_days=Decimal("1"))
        body = self.run_report("attendance-salary-loss", **P)
        r = self.by_code(body)["3020"]
        self.assertEqual((r["monthlySalary"], r["lopAmount"]), (20000.0, 5000.0))
        self.assertEqual((r["absentDays"], r["unpaidLeaveDays"], r["lopDays"]), (5.0, 1.0, 6.0))
        self.assertIsNone(r["halfDayLossDays"])  # not derivable without the breakdown
        self.assertIsNone(r["paidDays"])
        self.assertTrue(any("profile salary" in n for n in body["notes"]))

    def test_scope_filters_and_no_employment_filter(self):
        self.assertEqual(
            self.codes(self.run_report("attendance-salary-loss", departmentIds=str(self.d_cut.id), **P)), ["3002"]
        )
        self.assertEqual(
            self.codes(self.run_report("attendance-salary-loss", employeeIds=str(self.D.id), **P)), ["3004"]
        )
        self.assertEqual(self.run_report("attendance-salary-loss", period="2026-01")["rows"], [])
        spec = registry.get_spec("attendance-salary-loss")
        self.assertNotIn("employmentType", [f.key for f in spec.filters])

    def test_branch_isolation(self):
        self.assertEqual(
            self.codes(self.run_report("attendance-salary-loss", user=self.u_b1, impact="all", **P)),
            ["3004", "3001", "3002", "3006"],
        )
        self.assertEqual(
            self.run_report("attendance-salary-loss", user=self.u_b1, branchIds=str(self.b2.id), impact="all", **P)[
                "rows"
            ],
            [],
        )

    def test_module_gate(self):
        for user in (self.u_reports_only, self.u_prod_only):
            r = self.client.get("/api/reports/run/attendance-salary-loss", P, **_headers(user))
            self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))
        self.assertEqual(self.run_report("attendance-salary-loss", user=self.u_staff_only, **P)["rowCount"], 4)

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("attendance-salary-loss", "xlsx", **P).content)).active
        self.assertEqual([c.value for c in ws[7]][:4], ["Emp Code", "Employee", "Department", "Monthly salary"])
        self.assertEqual(ws["A8"].value, "3004")
        self.assertTrue(self.export("attendance-salary-loss", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.assert_flat_queries("attendance-salary-loss", lambda: self.extras(12))


# ================================================================================================
#  Overtime register
# ================================================================================================
class OvertimeRegisterTests(_Base):
    def test_golden_rows(self):
        body = self.run_report("overtime-register", **OT_FEB)
        rows = self.data_rows(body)
        self.assertEqual(len(rows), 9)
        self.assertEqual([(r["employeeCode"], r["date"]) for r in rows], [
            ("3001", "2026-02-02"), ("3001", "2026-02-03"), ("3001", "2026-02-04"), ("3001", "2026-02-05"),
            ("3001", "2026-02-06"), ("3001", "2026-02-07"), ("3002", "2026-02-10"), ("3003", "2026-02-11"),
            ("3007", "2026-02-12"),
        ])  # fmt: skip
        r = rows[0]
        self.assertEqual(
            (r["shiftEndTime"], r["lastPunchOut"], r["otMinutes"], r["otHours"], r["status"], r["compensationType"],
             r["outcome"], r["announcedBy"], r["announcedAt"], r["payableAmount"]),
            ("17:30", "19:30", 120, 2.0, "Announced", "Pay", "Paid in slip", "HR One", "2026-02-28 02:00", 850.0),
        )  # fmt: skip
        outcomes = [r["outcome"] for r in rows]
        self.assertEqual(outcomes, ["Paid in slip", "Paid in slip", "Credit available", None, None, "Credit used",
                                    "Awaiting payroll", "Credit available", "Awaiting payroll"])  # fmt: skip
        self.assertEqual([r["payableAmount"] for r in rows],
                         [850.0, 850.0, None, None, None, None, 1000.0, None, None])  # fmt: skip
        self.assertEqual([r["status"] for r in rows][3:5], ["Detected", "Rejected"])
        self.assertIsNone(rows[3]["compensationType"])

    def test_announced_at_is_shown_in_ist(self):
        # 2026-02-27 20:30 UTC is 02:00 on the 28th at the factory
        rows = self.data_rows(self.run_report("overtime-register", **OT_FEB))
        self.assertEqual(rows[0]["announcedAt"], "2026-02-28 02:00")

    def test_totals_and_summary_equal_the_rows(self):
        body = self.run_report("overtime-register", **OT_FEB)
        self.assertEqual(body["totals"], {"otMinutes": 820, "otHours": 13.67, "payableAmount": 2700.0})
        self.assert_totals(body, ["otMinutes", "otHours", "payableAmount"])
        s = self.summary(body)
        self.assertEqual(
            (s["Records"], s["Total OT hours (listed)"], s["Detected"], s["Announced"], s["Rejected"],
             s["Pay days awaiting payroll"], s["Payable (where a slip exists)"], s["Relaxation credits granted"]),
            (9, 13.67, 1, 7, 1, 2, 2700.0, 3),
        )  # fmt: skip

    def test_filters_narrow(self):
        def rows(**q):
            return [
                (r["employeeCode"], r["date"][-2:])
                for r in self.data_rows(self.run_report("overtime-register", **{**OT_FEB, **q}))
            ]

        self.assertEqual(rows(status="detected"), [("3001", "05")])
        self.assertEqual(rows(status="rejected"), [("3001", "06")])
        self.assertEqual(len(rows(status="announced")), 7)
        self.assertEqual(rows(compensationType="pay"), [("3001", "02"), ("3001", "03"), ("3002", "10"), ("3007", "12")])
        self.assertEqual(len(rows(compensationType="relaxation")), 3)
        self.assertEqual(
            rows(dateFrom="2026-02-03", dateTo="2026-02-05"), [("3001", "03"), ("3001", "04"), ("3001", "05")]
        )
        self.assertEqual(rows(departmentIds=str(self.d_sew.id)), [("3003", "11"), ("3007", "12")])
        self.assertEqual(rows(branchIds=str(self.b2.id)), [("3003", "11"), ("3007", "12")])
        self.assertEqual(rows(employeeIds=str(self.B.id)), [("3002", "10")])
        self.assertEqual(rows(dateFrom="2026-03-01", dateTo="2026-03-31"), [])

    def test_regenerating_payroll_turns_awaiting_into_paid(self):
        _generate_staff_payroll(self.B, MONTH, YEAR)  # picks up the announced day (real engine path)
        self.assertEqual(SalarySlip.objects.get(employee=self.B, month=MONTH).ot_amount, Decimal("1000.00"))
        rows = self.by_code(self.run_report("overtime-register", **OT_FEB))
        self.assertEqual((rows["3002"]["outcome"], rows["3002"]["payableAmount"]), ("Paid in slip", 1000.0))

    def test_partly_paid_when_only_some_days_are_in_the_slip(self):
        _generate_staff_payroll(self.B, MONTH, YEAR)
        _ot(self.B, date(2026, 2, 11), 65, announced_at=ANNOUNCED)  # announced after the regenerated slip
        rows = [
            r for r in self.data_rows(self.run_report("overtime-register", **OT_FEB)) if r["employeeCode"] == "3002"
        ]
        self.assertEqual([r["outcome"] for r in rows], ["Partly paid", "Partly paid"])

    def test_switched_off_features(self):
        ps = PayrollSettings.get()
        ps.compensation_feature_enabled = False
        ps.save()
        body = self.run_report("overtime-register", **OT_FEB)
        self.assertEqual(body["rows"], [])  # an explanation, not zeros
        self.assertTrue(any("Compensation feature is switched off" in n for n in body["notes"]))
        ps.compensation_feature_enabled = True
        ps.ot_detection_enabled = False
        ps.save()
        body = self.run_report("overtime-register", **OT_FEB)
        self.assertEqual(len(self.data_rows(body)), 9)  # records already on file (and paid) stay visible
        self.assertTrue(any("detection is switched off" in n for n in body["notes"]))

    def test_read_only_even_without_a_settings_row(self):
        PayrollSettings.objects.all().delete()
        self.run_report("overtime-register", **OT_FEB)
        self.assertEqual(PayrollSettings.objects.count(), 0)  # PayrollSettings.get() would have created it

    def test_branch_isolation(self):
        body = self.run_report("overtime-register", user=self.u_b1, **OT_FEB)
        self.assertEqual({r["employeeCode"] for r in self.data_rows(body)}, {"3001", "3002"})
        self.assertEqual(
            self.run_report("overtime-register", user=self.u_b1, branchIds=str(self.b2.id), **OT_FEB)["rows"], []
        )
        self.assertEqual(
            {r["employeeCode"] for r in self.data_rows(self.run_report("overtime-register", user=self.u_b2, **OT_FEB))},
            {"3003", "3007"},
        )

    def test_module_gate(self):
        self.assertEqual(
            self.client.get("/api/reports/run/overtime-register", OT_FEB, **_headers(self.u_reports_only)).status_code,
            403,
        )
        self.assertEqual(self.run_report("overtime-register", user=self.u_comp_only, **OT_FEB)["rowCount"], 9)
        # the day-wise HR decisions are Compensation-page data: Payroll alone does not open them
        r = self.client.get("/api/reports/run/overtime-register", OT_FEB, **_headers(self.u_staff_only))
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("overtime-register", "xlsx", **OT_FEB).content)).active
        self.assertEqual([c.value for c in ws[7]][:4], ["Emp Code", "Employee", "Department", "Date"])
        self.assertEqual(ws["A8"].value, "3001")
        self.assertEqual(ws["H8"].value, 2.0)
        self.assertEqual(ws["N8"].value, 850.0)
        self.assertEqual(ws["M8"].value, datetime(2026, 2, 28, 2, 0))
        self.assertTrue(self.export("overtime-register", "pdf", **OT_FEB).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        def grow():
            for emp in self.extras(12, late="0", late_summary=None):
                _ot(emp, date(2026, 2, 9), 100, announced_at=ANNOUNCED)
                _ot(emp, date(2026, 2, 10), 100, ctype="relaxation", announced_at=ANNOUNCED)

        self.assert_flat_queries("overtime-register", grow)


# ================================================================================================
#  Overtime payment summary
# ================================================================================================
class OvertimePaymentSummaryTests(_Base):
    def test_golden_rows(self):
        body = self.run_report("overtime-payment-summary", **P)
        rows = self.by_code(body)
        self.assertEqual(list(rows), ["3001", "3002", "3003", "3007"])
        keys = ("detectedDays", "announcedPayDays", "relaxationDays", "rejectedDays", "totalOtHours", "dailyRate",
                "expectedOtPay", "slipOtAmount", "pendingPayDays", "pendingPayAmount")  # fmt: skip
        expect = {
            "3001": (1, 2, 2, 1, 8.0, 850.0, 1700.0, 1700.0, 0, 0.0),
            "3002": (0, 1, 0, 0, 1.0, 1000.0, 1000.0, 0.0, 1, 1000.0),
            "3003": (0, 0, 1, 0, 1.17, 625.0, 0.0, 0.0, 0, 0.0),
            "3007": (0, 1, 0, 0, 2.17, None, None, None, 1, None),  # no slip yet: rate and amounts unknown
        }
        for code, vals in expect.items():
            self.assertEqual(tuple(rows[code][k] for k in keys), vals, code)
        self.assertEqual(rows["3001"]["period"], "February 2026")

    def test_totals_and_summary(self):
        body = self.run_report("overtime-payment-summary", **P)
        self.assertEqual(
            body["totals"],
            {"detectedDays": 1, "announcedPayDays": 4, "relaxationDays": 3, "rejectedDays": 1, "totalOtHours": 12.34,
             "expectedOtPay": 2700.0, "slipOtAmount": 1700.0, "pendingPayDays": 2, "pendingPayAmount": 1000.0},
        )  # fmt: skip
        self.assert_totals(
            body, ["announcedPayDays", "totalOtHours", "expectedOtPay", "slipOtAmount", "pendingPayAmount"]
        )
        s = self.summary(body)
        self.assertEqual(
            (s["Employees with overtime"], s["OT hours (excl. rejected)"], s["Announced pay days"], s["OT paid in slips"],
             s["Pay days pending payroll"], s["Pending amount (where rate known)"], s["Relaxation days announced"],
             s["Credits redeemed"]),
            (4, 12.34, 4, 1700.0, 2, 1000.0, 3, 1),
        )  # fmt: skip
        slips = sum(float(x.ot_amount) for x in SalarySlip.objects.filter(month=MONTH, year=YEAR))
        self.assertEqual(s["OT paid in slips"], round(slips, 2))

    def test_compensation_type_filter(self):
        self.assertEqual(
            self.codes(self.run_report("overtime-payment-summary", compensationType="pay", **P)),
            ["3001", "3002", "3007"],
        )
        self.assertEqual(
            self.codes(self.run_report("overtime-payment-summary", compensationType="relaxation", **P)),
            ["3001", "3003"],
        )
        self.assertEqual(len(self.codes(self.run_report("overtime-payment-summary", compensationType="all", **P))), 4)

    def test_scope_filters(self):
        self.assertEqual(
            self.codes(self.run_report("overtime-payment-summary", departmentIds=str(self.d_sew.id), **P)),
            ["3003", "3007"],
        )
        self.assertEqual(
            self.codes(self.run_report("overtime-payment-summary", employeeIds=str(self.A.id), **P)), ["3001"]
        )
        self.assertEqual(self.codes(self.run_report("overtime-payment-summary", period="2026-03")), [])

    def test_slip_overtime_without_a_current_record_is_still_listed(self):
        emp = _staff("3020", "20000", self.d_cut, self.b1)
        _hand_slip(emp, ot="833.33", ot_days=1)  # a slip that paid OT whose OT record no longer exists
        r = self.by_code(self.run_report("overtime-payment-summary", **P))["3020"]
        self.assertEqual(
            (r["announcedPayDays"], r["slipOtAmount"], r["pendingPayDays"], r["expectedOtPay"]), (0, 833.33, 0, 0.0)
        )

    def test_switched_off_feature(self):
        ps = PayrollSettings.get()
        ps.compensation_feature_enabled = False
        ps.save()
        body = self.run_report("overtime-payment-summary", **P)
        self.assertEqual(body["rows"], [])
        self.assertTrue(any("switched off" in n for n in body["notes"]))

    def test_branch_isolation(self):
        self.assertEqual(self.codes(self.run_report("overtime-payment-summary", user=self.u_b1, **P)), ["3001", "3002"])
        self.assertEqual(
            self.run_report("overtime-payment-summary", user=self.u_b1, branchIds=str(self.b2.id), **P)["rows"], []
        )
        self.assertEqual(self.codes(self.run_report("overtime-payment-summary", user=self.u_b2, **P)), ["3003", "3007"])

    def test_module_gate(self):
        self.assertEqual(
            self.client.get(
                "/api/reports/run/overtime-payment-summary", P, **_headers(self.u_reports_only)
            ).status_code,
            403,
        )
        self.assertEqual(self.run_report("overtime-payment-summary", user=self.u_comp_only, **P)["rowCount"], 4)
        # the pay side (what slips paid, what is pending) is payroll data too
        self.assertEqual(self.run_report("overtime-payment-summary", user=self.u_staff_only, **P)["rowCount"], 4)

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("overtime-payment-summary", "xlsx", **P).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Emp Code", "Employee", "Department"])
        self.assertEqual(ws["A8"].value, "3001")
        self.assertEqual(ws["L8"].value, 1700.0)
        self.assertTrue(self.export("overtime-payment-summary", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        def grow():
            for emp in self.extras(12, late="0", late_summary=None):
                _ot(emp, date(2026, 2, 9), 100, announced_at=ANNOUNCED)

        self.assert_flat_queries("overtime-payment-summary", grow)


# ================================================================================================
#  Compensation credits
# ================================================================================================
class CompensationCreditsTests(_Base):
    def test_golden_rows_with_a_pinned_clock(self):
        with self.today(date(2026, 3, 15)):
            body = self.run_report("compensation-credits")
        rows = self.data_rows(body)
        self.assertEqual([(r["employeeCode"], r["sourceDate"], r["status"], r["usedDate"], r["ageDays"], r["createdAt"]) for r in rows], [
            ("3001", "2026-02-04", "Available", None, 15, "2026-02-28 02:00"),
            ("3001", "2026-02-07", "Used", "2026-02-20", None, "2026-02-28 02:00"),
            ("3003", "2026-02-11", "Available", None, 5, "2026-03-10 10:30"),
        ])  # fmt: skip
        s = self.summary(body)
        self.assertEqual(
            (
                s["Credits earned"],
                s["Available"],
                s["Redeemed"],
                s["Employees holding credits"],
                s["Oldest unused (days)"],
            ),
            (3, 2, 1, 2, 15),
        )

    def test_status_filter(self):
        with self.today(date(2026, 3, 15)):
            self.assertEqual(len(self.run_report("compensation-credits", status="available")["rows"]), 2)
            used = self.run_report("compensation-credits", status="used")
        self.assertEqual([r["usedDate"] for r in used["rows"]], ["2026-02-20"])

    def test_earned_in_window_uses_ist_days(self):
        with self.today(date(2026, 3, 15)):

            def codes(w):
                return [
                    (r["employeeCode"], r["sourceDate"])
                    for r in self.run_report("compensation-credits", earnedIn=w)["rows"]
                ]

            self.assertEqual(codes("lastMonth"), [("3001", "2026-02-04"), ("3001", "2026-02-07")])
            self.assertEqual(codes("thisMonth"), [("3003", "2026-02-11")])
            self.assertEqual(len(codes("all")), 3)
            self.assertEqual(len(codes("last90")), 3)
            self.assertEqual(len(codes("thisYear")), 3)
            # 28-Feb 20:00 UTC is 01-Mar 01:30 IST: the credit belongs to March, not February
            CompensationLeaveCredit.objects.filter(pk=self.credit_used.pk).update(
                created_at=datetime(2026, 2, 28, 20, 0, tzinfo=timezone.utc)
            )
            self.assertEqual(codes("lastMonth"), [("3001", "2026-02-04")])
            self.assertEqual(len(codes("thisMonth")), 2)

    def test_credit_whose_overtime_record_was_removed(self):
        OvertimeRecord.objects.filter(pk=self.credit_avail.source_overtime_record_id).delete()
        with self.today(date(2026, 3, 15)):
            rows = self.run_report("compensation-credits", employeeIds=str(self.A.id))["rows"]
        self.assertIsNone(rows[0]["sourceDate"])  # SET_NULL: blank, not an error

    def test_scope_filters(self):
        self.assertEqual(
            self.codes(self.run_report("compensation-credits", departmentIds=str(self.d_sew.id))), ["3003"]
        )
        self.assertEqual(
            self.codes(self.run_report("compensation-credits", employeeIds=str(self.A.id))), ["3001", "3001"]
        )
        self.assertEqual(self.run_report("compensation-credits", branchIds=str(self.b2.id))["rowCount"], 1)

    def test_switched_off_feature(self):
        ps = PayrollSettings.get()
        ps.compensation_feature_enabled = False
        ps.save()
        body = self.run_report("compensation-credits")
        self.assertEqual(body["rows"], [])
        self.assertTrue(any("switched off" in n for n in body["notes"]))

    def test_branch_isolation(self):
        self.assertEqual(self.codes(self.run_report("compensation-credits", user=self.u_b1)), ["3001", "3001"])
        self.assertEqual(self.run_report("compensation-credits", user=self.u_b1, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(self.codes(self.run_report("compensation-credits", user=self.u_b2)), ["3003"])

    def test_module_gate_needs_compensation(self):
        for user in (self.u_reports_only, self.u_staff_only):
            r = self.client.get("/api/reports/run/compensation-credits", **_headers(user))
            self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))
        self.assertEqual(self.run_report("compensation-credits", user=self.u_comp_only)["rowCount"], 3)

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("compensation-credits", "xlsx").content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Emp Code", "Employee", "Department"])
        self.assertEqual(ws["A8"].value, "3001")
        self.assertEqual(ws["G9"].value.date(), date(2026, 2, 20))  # redeemed-on is a real date cell
        self.assertTrue(self.export("compensation-credits", "pdf").content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        def grow():
            for emp in self.extras(12, late="0", late_summary=None, payroll=False):
                rec = _ot(emp, date(2026, 2, 9), 100, ctype="relaxation", announced_at=ANNOUNCED)
                CompensationLeaveCredit.objects.create(employee=emp, source_overtime_record=rec)

        self.assert_flat_queries("compensation-credits", grow)


# ================================================================================================
#  Compensation day announcements
# ================================================================================================
class CompensationDayAnnouncementsTests(_Base):
    def rows(self, user=None, **q):
        with self.today(date(2026, 6, 15)):
            return self.data_rows(self.run_report("compensation-day-announcements", user=user, **q))

    def test_golden_rows_newest_first(self):
        rows = self.rows(when="all")
        self.assertEqual(
            [r["date"] for r in rows], ["2026-03-05", "2026-02-27", "2026-02-25", "2026-02-20", "2026-02-14"]
        )
        by = {r["date"]: r for r in rows}
        allday = by["2026-02-14"]
        self.assertEqual(
            (allday["releaseType"], allday["leaveUntilTime"], allday["branch"], allday["department"], allday["employeeCount"],
             allday["reason"], allday["announcedBy"], allday["createdAt"]),
            ("Whole day", None, "All branches", "All departments", None, "Festival", "HR One", "2026-02-01 10:00"),
        )  # fmt: skip
        early = by["2026-02-20"]
        self.assertEqual(
            (early["releaseType"], early["leaveUntilTime"], early["branch"]), ("Early release", "15:00", "Unit 1")
        )
        sew = by["2026-02-25"]
        self.assertEqual((sew["branch"], sew["department"], sew["employeeCount"]), ("Unit 2", "SEWING", 2))
        self.assertEqual(by["2026-03-05"]["department"], "CUTTING")

    def test_summary(self):
        with self.today(date(2026, 6, 15)):
            body = self.run_report("compensation-day-announcements", when="all")
        s = self.summary(body)
        self.assertEqual((s["Announcements"], s["Whole-day"], s["Early release"], s["Distinct dates"]), (5, 3, 2, 5))

    def test_when_filter(self):
        self.assertEqual(len(self.rows()), 5)  # default: this year (2026)
        self.assertEqual([r["date"] for r in self.rows(when="lastMonth")], [])  # May 2026
        with self.today(date(2026, 3, 1)):
            body = self.run_report("compensation-day-announcements", when="thisMonth")
        self.assertEqual([r["date"] for r in self.data_rows(body)], ["2026-03-05"])
        with self.today(date(2026, 2, 21)):
            up = self.run_report("compensation-day-announcements", when="upcoming")
        self.assertEqual([r["date"] for r in self.data_rows(up)], ["2026-03-05", "2026-02-27", "2026-02-25"])
        with self.today(date(2027, 1, 5)):
            self.assertEqual(self.run_report("compensation-day-announcements")["rows"], [])  # default window: 2027

    def test_department_and_branch_filters_include_company_wide_days(self):
        dates = lambda **q: [r["date"] for r in self.rows(when="all", **q)]  # noqa: E731
        # the engine (attendance_final._compensation_day_for) applies an announcement that NAMES employees to those
        # employees only (02-27 names E of SEWING: not a CUTTING day), and a department's day to its own branch
        # (03-05 is CUTTING, a Unit 1 department: not a Unit 2 day)
        self.assertEqual(dates(departmentIds=str(self.d_cut.id)), ["2026-03-05", "2026-02-20", "2026-02-14"])
        self.assertEqual(dates(branchIds=str(self.b2.id)), ["2026-02-27", "2026-02-25", "2026-02-14"])
        self.assertEqual(dates(departmentIds=str(self.d_sew.id)), ["2026-02-27", "2026-02-25", "2026-02-14"])

    def test_branch_scoped_users_only_see_their_branch_and_company_wide_days(self):
        self.assertEqual(
            [r["date"] for r in self.rows(user=self.u_b1, when="all")], ["2026-03-05", "2026-02-20", "2026-02-14"]
        )
        self.assertEqual(
            [r["date"] for r in self.rows(user=self.u_b2, when="all")], ["2026-02-27", "2026-02-25", "2026-02-14"]
        )
        widened = [r["date"] for r in self.rows(user=self.u_b1, when="all", branchIds=str(self.b2.id))]
        self.assertEqual(widened, ["2026-02-14"])  # only what branch 1 may see; branch 2's days stay hidden
        sew = [r for r in self.rows(user=self.u_b2, when="all") if r["date"] == "2026-02-25"][0]
        self.assertEqual(sew["employeeCount"], 2)  # both named employees are in branch 2
        for r in self.rows(user=self.u_b1, when="all"):
            self.assertNotIn(r["date"], ("2026-02-25", "2026-02-27"))

    def test_named_employees_from_another_branch_are_not_counted_for_a_scoped_user(self):
        self.ann_b2.employees.add(self.A)  # a branch-1 employee joins the branch-2 announcement
        scoped = [r for r in self.rows(user=self.u_b1, when="all") if r["date"] == "2026-02-25"]
        self.assertEqual([r["employeeCount"] for r in scoped], [1])  # visible via A, counts only branch-1 people
        full = [r for r in self.rows(when="all") if r["date"] == "2026-02-25"]
        self.assertEqual(full[0]["employeeCount"], 3)

    def test_switched_off_feature(self):
        ps = PayrollSettings.get()
        ps.compensation_feature_enabled = False
        ps.save()
        with self.today(date(2026, 6, 15)):
            body = self.run_report("compensation-day-announcements", when="all")
        self.assertEqual(body["rows"], [])
        self.assertTrue(any("switched off" in n for n in body["notes"]))

    def test_module_gate(self):
        r = self.client.get("/api/reports/run/compensation-day-announcements", **_headers(self.u_reports_only))
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))
        self.assertEqual(
            self.run_report("compensation-day-announcements", user=self.u_comp_only, when="all")["rowCount"], 5
        )

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("compensation-day-announcements", "xlsx", when="all").content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Date", "Leave until", "Type"])
        self.assertEqual(ws["A8"].value.date(), date(2026, 3, 5))
        self.assertTrue(self.export("compensation-day-announcements", "pdf", when="all").content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        def grow():
            for i in range(12):
                a = CompensationDayAnnouncement.objects.create(date=date(2026, 4, 1 + i), reason="x", branch=self.b1)
                a.employees.set([self.A, self.B])

        self.assert_flat_queries("compensation-day-announcements", grow, params={"when": "all"})


# ================================================================================================
#  Payroll exceptions
# ================================================================================================
class PayrollExceptionsTests(_Base):
    def test_golden_findings(self):
        body = self.run_report("payroll-exceptions", **P)
        rows = self.data_rows(body)
        self.assertEqual(
            [(r["employeeCode"], r["exceptionType"]) for r in rows],
            [("3003", "Bank details"), ("3003", "PF ID missing"), ("3003", "ESI ID missing"), ("3007", "No slip"),
             ("3008", "No salary set"), ("P001", "ESI ID missing")],
        )  # fmt: skip
        by = {(r["employeeCode"], r["exceptionType"]): r for r in rows}
        self.assertEqual(by[("3003", "Bank details")]["detail"], "Missing bank account and IFSC on the profile.")
        self.assertEqual(
            by[("3003", "PF ID missing")]["detail"], "PF of Rs. 900.00 deducted but UAN and PF no. not on the profile."
        )
        self.assertEqual(by[("3003", "ESI ID missing")]["netPay"], 12987.5)
        self.assertEqual(by[("P001", "ESI ID missing")]["period"], "02-Feb-2026 to 11-Feb-2026")
        self.assertEqual(
            by[("P001", "ESI ID missing")]["detail"], "ESI of Rs. 61.88 deducted but no ESI IP number on the profile."
        )
        self.assertIsNone(by[("3007", "No slip")]["netPay"])
        self.assertEqual(
            by[("3008", "No salary set")]["detail"], "No monthly salary on the profile - payroll skips this employee."
        )
        s = self.summary(body)
        self.assertEqual(
            (s["Employees checked"], s["No slip"], s["No salary set"], s["Bank details missing"], s["PF / ESI ID missing"],
             s["Zero / negative net"], s["Provisional slips"], s["Findings listed"]),
            (7, 1, 1, 1, 3, 0, 0, 6),
        )  # fmt: skip

    def test_a_clean_check_says_so(self):
        body = self.run_report("payroll-exceptions", exceptionType="zero-net", **P)
        self.assertEqual(body["rows"], [])
        self.assertIn("No exceptions found for the selected check in February 2026", body["notes"][0])
        self.assertNotIn("No exceptions found", self.run_report("payroll-exceptions", **P)["notes"][0])

    def test_check_filter_keeps_all_counts(self):
        body = self.run_report("payroll-exceptions", exceptionType="no-slip", **P)
        self.assertEqual([(r["employeeCode"], r["exceptionType"]) for r in body["rows"]], [("3007", "No slip")])
        s = self.summary(body)
        self.assertEqual(
            (s["No slip"], s["PF / ESI ID missing"], s["Findings listed"]), (1, 3, 1)
        )  # counts still describe all checks
        self.assertEqual(self.run_report("payroll-exceptions", exceptionType="pf-id-missing", **P)["rowCount"], 1)
        self.assertEqual(self.run_report("payroll-exceptions", exceptionType="esi-id-missing", **P)["rowCount"], 2)

    def test_zero_and_negative_net_pay(self):
        for code, advance in (("3030", "10000"), ("3031", "12000")):
            emp = _staff(code, "10000", self.d_cut, self.b1, bank_name="SBI", bank_account="1", bank_ifsc="SBIN0000001",
                         uan_number="1", pf_number="1", esi_number="1")  # fmt: skip
            _present_all(emp)
            adv = Advance.objects.create(employee=emp, advance_type="general", amount=Decimal(advance), status="approved",
                                         total_repaid=Decimal("0"), outstanding=Decimal(advance))  # fmt: skip
            AdvanceRepayment.objects.create(advance=adv, month=MONTH, year=YEAR, amount=Decimal(advance))
            PayrollSettings.objects.filter(pk=1).update(staff_payroll_rules_enabled=False)
            _generate_staff_payroll(emp, MONTH, YEAR)
        self.assertEqual(SalarySlip.objects.get(employee__employee_code="3030").net_salary, Decimal("0.00"))
        self.assertEqual(SalarySlip.objects.get(employee__employee_code="3031").net_salary, Decimal("-2000.00"))
        rows = {
            (r["employeeCode"], r["exceptionType"]): r
            for r in self.data_rows(self.run_report("payroll-exceptions", **P))
        }
        self.assertEqual(rows[("3030", "Zero net pay")]["netPay"], 0.0)
        self.assertEqual(rows[("3031", "Negative net pay")]["netPay"], -2000.0)
        self.assertIn("exceed earnings", rows[("3031", "Negative net pay")]["detail"])
        self.assertEqual(self.summary(self.run_report("payroll-exceptions", **P))["Zero / negative net"], 2)

    def test_provisional_staff_slips_while_the_month_is_running(self):
        with self.today(date(2026, 2, 15)):
            body = self.run_report("payroll-exceptions", exceptionType="provisional", **P)
        self.assertEqual(
            self.codes(body), ["3001", "3002", "3003", "3004"]
        )  # staff slips of active employees; production is never provisional
        self.assertEqual(self.summary(body)["Provisional slips"], 4)
        with self.today(date(2026, 3, 2)):
            self.assertEqual(self.run_report("payroll-exceptions", exceptionType="provisional", **P)["rows"], [])

    def test_no_slip_detail_says_when_the_month_has_not_ended(self):
        with self.today(date(2026, 2, 15)):
            rows = self.data_rows(self.run_report("payroll-exceptions", exceptionType="no-slip", **P))
        self.assertIn("has not ended yet", rows[0]["detail"])

    def test_employees_who_joined_after_the_month_are_not_expected_to_have_a_slip(self):
        _staff("3050", "20000", self.d_cut, self.b1, join_date="2026-03-05")  # joined after February
        _staff("3051", "20000", self.d_cut, self.b1, join_date="15/02/2026")  # joined during February (DD/MM/YYYY)
        _staff("3052", "20000", self.d_cut, self.b1, join_date="soon")  # unreadable text: assume employed
        body = self.run_report("payroll-exceptions", exceptionType="no-slip", **P)
        self.assertEqual(self.codes(body), ["3007", "3051", "3052"])
        self.assertTrue(any("joining date is after the month" in n for n in body["notes"]))
        self.assertEqual(self.summary(body)["No slip"], 3)

    def test_inactive_employees_are_not_expected_to_have_slips(self):
        body = self.run_report("payroll-exceptions", period="2026-03")
        self.assertNotIn("3006", self.codes(body))  # L left: no 'No slip' for an inactive employee
        no_slip = {r["employeeCode"] for r in self.data_rows(body) if r["exceptionType"] == "No slip"}
        self.assertEqual(no_slip, {"3001", "3002", "3003", "3004", "3007", "P001"})

    def test_scope_filters(self):
        self.assertEqual(
            self.codes(self.run_report("payroll-exceptions", departmentIds=str(self.d_sew.id), **P)),
            ["3003", "3003", "3003", "3007"],
        )
        self.assertEqual(self.codes(self.run_report("payroll-exceptions", employmentType="production", **P)), ["P001"])
        self.assertEqual(
            self.codes(self.run_report("payroll-exceptions", branchIds=str(self.b2.id), **P)),
            ["3003", "3003", "3003", "3007"],
        )

    def test_branch_isolation(self):
        body = self.run_report("payroll-exceptions", user=self.u_b1, **P)
        self.assertEqual(self.codes(body), ["3008", "P001"])
        self.assertEqual(
            self.run_report("payroll-exceptions", user=self.u_b1, branchIds=str(self.b2.id), **P)["rows"], []
        )
        self.assertEqual(self.summary(body)["Employees checked"], 5)  # A, B, D, F, P1 (L is inactive)
        self.assertEqual(
            self.codes(self.run_report("payroll-exceptions", user=self.u_b2, **P)), ["3003", "3003", "3003", "3007"]
        )

    def test_production_only_role(self):
        body = self.run_report("payroll-exceptions", user=self.u_prod_only, **P)
        self.assertEqual(self.codes(body), ["P001"])
        self.assertEqual(self.summary(body)["Employees checked"], 1)

    def test_module_gate(self):
        r = self.client.get("/api/reports/run/payroll-exceptions", P, **_headers(self.u_reports_only))
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))

    def test_export_round_trip(self):
        ws = load_workbook(io.BytesIO(self.export("payroll-exceptions", "xlsx", **P).content)).active
        self.assertEqual(
            [c.value for c in ws[7]][:6], ["Emp Code", "Employee", "Department", "Type", "Period", "Check"]
        )
        self.assertEqual((ws["A8"].value, ws["F8"].value), ("3003", "Bank details"))
        self.assertTrue(self.export("payroll-exceptions", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.assert_flat_queries("payroll-exceptions", lambda: self.extras(12))


# ================================================================================================
#  Minimum wage compliance
# ================================================================================================
class MinWageComplianceTests(_Base):
    def set_rate(self, rate):
        PayrollSettings.objects.filter(pk=1).update(min_wage_rate=Decimal(rate))

    def test_not_configured_explains_instead_of_guessing(self):
        body = self.run_report("min-wage-compliance", **P)
        self.assertEqual(body["rows"], [])
        self.assertTrue(any("not set" in n for n in body["notes"]))

    def test_per_day_basis(self):
        self.set_rate("900")
        body = self.run_report("min-wage-compliance", **P)
        rows = body["rows"]
        by = {(r["employeeCode"], r["period"]): r for r in rows}
        expect = {  # paid days/shifts, earned, wage rate, shortfall, result
            ("3001", "February 2026"): (24.0, 20400.0, 850.0, 1200.0, "Below minimum"),
            ("3002", "February 2026"): (19.0, 19000.0, 1000.0, None, "OK"),
            ("3003", "February 2026"): (24.0, 15000.0, 625.0, 6600.0, "Below minimum"),
            ("3004", "February 2026"): (24.0, 60000.0, 2500.0, None, "OK"),
            ("3006", "February 2026"): (24.0, 18000.0, 750.0, 3600.0, "Below minimum"),
            ("P001", "02-Feb-2026 to 11-Feb-2026"): (12.5, 6250.0, 500.0, 5000.0, "Below minimum"),
            ("P001", "16-Feb-2026 to 25-Feb-2026"): (4.0, 2000.0, 500.0, 1600.0, "Below minimum"),
        }
        self.assertEqual(set(by), set(expect))
        for k, (paid, earned, rate, short, res) in expect.items():
            r = by[k]
            self.assertEqual(
                (r["paidDays"], r["earnedWages"], r["wageRate"], r["shortfallAmount"], r["result"]),
                (paid, earned, rate, short, res),
                k,
            )
            self.assertEqual(r["minWageRate"], 900.0)
        s = self.summary(body)
        self.assertEqual(
            (s["Slips checked"], s["Below minimum"], s["Total shortfall"], s["Minimum rate used"]),
            (7, 5, 18000.0, 900.0),
        )
        self.assertEqual(body["totals"]["shortfallAmount"], 18000.0)
        self.assert_totals(body, ["earnedWages", "shortfallAmount"])

    def test_per_month_basis_is_staff_only(self):
        self.set_rate("20000")
        body = self.run_report("min-wage-compliance", basis="month", **P)
        by = self.by_code(body)
        self.assertEqual(set(by), {"3001", "3002", "3003", "3004", "3006"})  # production has no monthly wage
        self.assertEqual(
            (by["3003"]["wageRate"], by["3003"]["shortfallAmount"], by["3003"]["result"]),
            (15000.0, 5000.0, "Below minimum"),
        )
        self.assertEqual(
            (by["3006"]["shortfallAmount"], by["3001"]["result"], by["3002"]["result"]), (2000.0, "OK", "OK")
        )
        self.assertTrue(any("2 production slip(s) skipped" in n for n in body["notes"]))
        self.assertEqual(self.summary(body)["Total shortfall"], 7000.0)

    def test_partial_attendance_prorates_the_monthly_shortfall(self):
        self.set_rate("30000")
        b = self.by_code(self.run_report("min-wage-compliance", basis="month", **P))["3002"]
        # salary 24,000 -> gap 6,000 x (19 paid / 24 working days)
        self.assertEqual(b["shortfallAmount"], 4750.0)

    def test_only_below_filter(self):
        self.set_rate("900")
        body = self.run_report("min-wage-compliance", onlyBelow="true", **P)
        self.assertEqual({r["result"] for r in body["rows"]}, {"Below minimum"})
        self.assertEqual(body["rowCount"], 5)
        self.assertEqual(self.summary(body)["Slips checked"], 7)  # the summary still covers everyone

    def test_scope_filters(self):
        self.set_rate("900")
        self.assertEqual(
            self.codes(self.run_report("min-wage-compliance", departmentIds=str(self.d_sew.id), **P)), ["3003"]
        )
        self.assertEqual(
            self.codes(self.run_report("min-wage-compliance", employmentType="production", **P)), ["P001", "P001"]
        )
        self.assertEqual(self.run_report("min-wage-compliance", period="2026-01")["rows"], [])

    def test_slip_without_a_paid_day_has_no_data(self):
        self.set_rate("900")
        emp = _staff("3040", "20000", self.d_cut, self.b1)
        slip = _hand_slip(emp, gross="0", pf="0", esi="0", breakdown=False)
        SalarySlip.objects.filter(pk=slip.pk).update(present_days=Decimal("0"), paid_leave_days=Decimal("0"))
        r = self.by_code(self.run_report("min-wage-compliance", **P))["3040"]
        self.assertEqual((r["result"], r["wageRate"], r["shortfallAmount"]), ("No data", None, None))

    def test_branch_isolation(self):
        self.set_rate("900")
        self.assertEqual(
            self.codes(self.run_report("min-wage-compliance", user=self.u_b1, **P)),
            ["3001", "3002", "3004", "3006", "P001", "P001"],
        )
        self.assertEqual(
            self.run_report("min-wage-compliance", user=self.u_b1, branchIds=str(self.b2.id), **P)["rows"], []
        )

    def test_module_gate(self):
        r = self.client.get("/api/reports/run/min-wage-compliance", P, **_headers(self.u_reports_only))
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))

    def test_export_round_trip(self):
        self.set_rate("900")
        ws = load_workbook(io.BytesIO(self.export("min-wage-compliance", "xlsx", **P).content)).active
        self.assertEqual([c.value for c in ws[7]][:5], ["Emp Code", "Employee", "Department", "Type", "Period"])
        self.assertEqual(ws["A8"].value, "3001")
        self.assertEqual(ws["K8"].value, "Below minimum")
        self.assertTrue(self.export("min-wage-compliance", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.set_rate("900")
        self.assert_flat_queries("min-wage-compliance", lambda: self.extras(12))


# ================================================================================================
#  Payroll vs slip reconciliation
# ================================================================================================
class PayrollSlipReconciliationTests(_Base):
    def patch_payroll(self, emp, **body):
        p = Payroll.objects.get(employee=emp, month=MONTH, year=YEAR)
        r = self.client.patch(f"/api/payroll/{p.id}", body, content_type="application/json", **_headers(self.admin))
        self.assertEqual(r.status_code, 200, r.content)
        return p

    def issues(self, body):
        return [(r["employeeCode"], r["issue"]) for r in self.data_rows(body)]

    def test_clean_payroll_has_no_findings(self):
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(body["rows"], [])
        self.assertTrue(all(v in (0, 0.0) for v in self.summary(body).values()))
        self.assertIn("No differences found", body["notes"][0])  # "all is well" is stated, not left as a blank table
        # a month nobody generated is a different message: nothing to reconcile yet
        empty = self.run_report("payroll-slip-reconciliation", period="2019-01")
        self.assertIn("generate payroll first", empty["notes"][0])

    def test_marking_paid_drops_overtime_and_is_detected(self):
        self.patch_payroll(self.A, status="paid")  # the real mark-paid path recomputes final pay without OT
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(self.issues(body), [("3001", "Net differs")])
        r = body["rows"][0]
        self.assertEqual((r["slipNet"], r["payrollNet"], r["variance"]), (20723.0, 19023.0, -1700.0))
        self.assertEqual((r["slipDeductions"], r["payrollDeductions"], r["payrollBonus"]), (1377.0, 1377.0, 0.0))
        self.assertIn("overtime", r["detail"])
        s = self.summary(body)
        self.assertEqual(
            (s["Net pay differs"], s["Net variance (payroll - slip)"], s["Overtime dropped by status change"]),
            (1, -1700.0, 1),
        )
        self.assertEqual(body["totals"]["variance"], -1700.0)

    def test_manual_bonus_edit_is_detected(self):
        self.patch_payroll(self.L, bonus=500)
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(self.issues(body), [("3006", "Net differs"), ("3006", "Manual edit")])
        net = body["rows"][0]
        self.assertEqual(
            (net["slipNet"], net["payrollNet"], net["variance"], net["payrollBonus"]), (16785.0, 17285.0, 500.0, 500.0)
        )

    def test_regenerating_after_instalments_were_processed_loses_the_advance(self):
        _generate_staff_payroll(self.C, MONTH, YEAR)  # the known defect: instalment already processed
        self.assertEqual(SalarySlip.objects.get(employee=self.C, month=MONTH).advance_deduction, Decimal("0.00"))
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(self.issues(body), [("3003", "Advance mismatch")])
        r = body["rows"][0]
        self.assertEqual((r["advanceProcessed"], r["slipAdvance"]), (1000.0, 0.0))
        self.assertEqual(self.summary(body)["Advance mismatches"], 1)

    def test_processed_instalment_for_an_employee_without_a_slip(self):
        adv = Advance.objects.create(employee=self.E, advance_type="general", amount=Decimal("500"), status="approved",
                                     total_repaid=Decimal("500"), outstanding=Decimal("0"))  # fmt: skip
        AdvanceRepayment.objects.create(advance=adv, month=MONTH, year=YEAR, amount=Decimal("500"), is_processed=True)
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(self.issues(body), [("3007", "Advance mismatch")])
        r = body["rows"][0]
        self.assertEqual(
            (r["advanceProcessed"], r["slipAdvance"], r["slipNet"], r["period"]), (500.0, None, None, "February 2026")
        )

    def test_slip_without_payroll_and_payroll_without_slip(self):
        Payroll.objects.filter(employee=self.D).delete()
        SalarySlip.objects.filter(employee=self.B).delete()
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(self.issues(body), [("3002", "Payroll without slip"), ("3004", "Slip without payroll")])
        orphan_payroll, orphan_slip = body["rows"]
        self.assertEqual((orphan_payroll["slipNet"], orphan_payroll["payrollNet"]), (None, 17360.0))
        self.assertEqual(
            (orphan_slip["slipNet"], orphan_slip["payrollNet"], orphan_slip["variance"]), (55150.0, None, None)
        )

    def test_arithmetic_that_does_not_close(self):
        SalarySlip.objects.filter(employee=self.D).update(net_salary=Decimal("55000.00"))
        body = self.run_report("payroll-slip-reconciliation", issue="arithmetic", **P)
        self.assertEqual(self.issues(body), [("3004", "Arithmetic")])
        self.assertEqual(body["rows"][0]["detail"], "Net Rs. 55,000.00 != gross + OT - deductions (Rs. 55,150.00).")
        SalarySlip.objects.filter(employee=self.D).update(
            net_salary=Decimal("55150.00"), total_deductions=Decimal("4000.00")
        )
        self.assertEqual(
            self.issues(self.run_report("payroll-slip-reconciliation", issue="arithmetic", **P)),
            [("3004", "Arithmetic")],
        )

    def test_production_late_penalty_inside_total_deductions_is_not_an_error(self):
        emp = Employee.objects.create(employee_code="P002", first_name="Prod", last_name="Two", employment_type="production",
                                      status="active", salary_per_shift=Decimal("500"), department=self.d_stitch, branch=self.b1)  # fmt: skip
        _hand_slip(emp, gross="4000", pf="0", esi="0", late="125", period=(date(2026, 2, 2), date(2026, 2, 15)),
                   late_summary={"totalLateCount": 4, "billableLateCount": 1, "shiftDeductions": 0.25})  # fmt: skip
        self.assertEqual(self.run_report("payroll-slip-reconciliation", **P)["rows"], [])

    def test_issue_filter_and_scope_filters(self):
        self.patch_payroll(self.A, status="paid")
        self.patch_payroll(self.L, bonus=500)
        self.assertEqual(
            self.issues(self.run_report("payroll-slip-reconciliation", issue="manual-edit", **P)),
            [("3006", "Manual edit")],
        )
        self.assertEqual(self.issues(self.run_report("payroll-slip-reconciliation", issue="net-mismatch", **P)),
                         [("3001", "Net differs"), ("3006", "Net differs")])  # fmt: skip
        self.assertEqual(self.issues(self.run_report("payroll-slip-reconciliation", departmentIds=str(self.d_cut.id), **P)),
                         [("3001", "Net differs"), ("3006", "Net differs"), ("3006", "Manual edit")])  # fmt: skip
        self.assertEqual(self.run_report("payroll-slip-reconciliation", employeeIds=str(self.B.id), **P)["rows"], [])
        self.assertEqual(self.run_report("payroll-slip-reconciliation", period="2026-01")["rows"], [])

    def test_branch_isolation(self):
        self.patch_payroll(self.A, status="paid")
        _generate_staff_payroll(self.C, MONTH, YEAR)
        self.assertEqual(
            self.issues(self.run_report("payroll-slip-reconciliation", **P)),
            [("3001", "Net differs"), ("3003", "Advance mismatch")],
        )
        self.assertEqual(
            self.issues(self.run_report("payroll-slip-reconciliation", user=self.u_b1, **P)), [("3001", "Net differs")]
        )
        self.assertEqual(
            self.issues(self.run_report("payroll-slip-reconciliation", user=self.u_b2, **P)),
            [("3003", "Advance mismatch")],
        )
        self.assertEqual(
            self.run_report("payroll-slip-reconciliation", user=self.u_b1, branchIds=str(self.b2.id), **P)["rows"], []
        )

    def test_production_only_role_cannot_see_staff_findings(self):
        self.patch_payroll(self.A, status="paid")
        self.assertEqual(self.run_report("payroll-slip-reconciliation", user=self.u_prod_only, **P)["rows"], [])
        self.assertEqual(
            self.issues(self.run_report("payroll-slip-reconciliation", user=self.u_staff_only, **P)),
            [("3001", "Net differs")],
        )

    def test_module_gate(self):
        r = self.client.get("/api/reports/run/payroll-slip-reconciliation", P, **_headers(self.u_reports_only))
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))

    def test_export_round_trip(self):
        self.patch_payroll(self.A, status="paid")
        ws = load_workbook(io.BytesIO(self.export("payroll-slip-reconciliation", "xlsx", **P).content)).active
        self.assertEqual([c.value for c in ws[7]][:5], ["Emp Code", "Employee", "Department", "Type", "Period"])
        self.assertEqual(ws["A8"].value, "3001")
        self.assertEqual(ws["H8"].value, -1700.0)
        self.assertTrue(self.export("payroll-slip-reconciliation", "pdf", **P).content.startswith(b"%PDF"))

    def test_no_n_plus_one(self):
        self.patch_payroll(self.A, status="paid")
        self.assert_flat_queries("payroll-slip-reconciliation", lambda: self.extras(12))


# ================================================================================================
#  Cross-cutting: exports, empty results, permissions, read-only
# ================================================================================================
class ExportAndAccessTests(_Base):
    def setUp(self):
        PayrollSettings.objects.filter(pk=1).update(min_wage_rate=Decimal("900"))
        p = Payroll.objects.get(employee=self.A, month=MONTH, year=YEAR)
        self.client.patch(
            f"/api/payroll/{p.id}", {"status": "paid"}, content_type="application/json", **_headers(self.admin)
        )

    def test_every_report_exports_valid_excel_and_pdf_with_data(self):
        for rid in G2_IDS:
            spec = registry.get_spec(rid)
            params = DATA_PARAMS[rid]
            with self.subTest(report=rid):
                x = self.export(rid, "xlsx", **params)
                self.assertEqual(x.status_code, 200, rid)
                self.assertTrue(x.content.startswith(b"PK"))
                ws = load_workbook(io.BytesIO(x.content)).active
                self.assertEqual([c.value for c in ws[7]], [c.label for c in spec.columns])
                self.assertIsNotNone(ws["A8"].value, f"{rid} exported no data row")
                d = self.export(rid, "pdf", **params)
                self.assertEqual(d.status_code, 200, rid)
                self.assertTrue(d.content.startswith(b"%PDF"))
                self.assertIn(b"%%EOF", d.content[-64:])

    def test_every_report_handles_an_empty_result_on_screen_and_in_both_exports(self):
        for rid in G2_IDS:
            params = EMPTY_PARAMS[rid]
            with self.subTest(report=rid), self.today(date(2030, 1, 1)):  # nothing is "upcoming" in 2030
                body = self.run_report(rid, **params)
                self.assertEqual(body["rows"], [], rid)
                self.assertIsInstance(body["summary"], list)
                self.assertEqual(self.export(rid, "xlsx", **params).status_code, 200)
                self.assertTrue(self.export(rid, "pdf", **params).content.startswith(b"%PDF"))

    def test_slip_based_reports_say_when_payroll_was_not_generated(self):
        for rid in (
            "pf-statement",
            "esi-statement",
            "late-salary-impact",
            "attendance-salary-loss",
            "min-wage-compliance",
        ):
            with self.subTest(report=rid):
                body = self.run_report(rid, **EMPTY_PARAMS[rid])
                self.assertEqual(body["rows"], [])
                self.assertIn("Generate Payroll", body["notes"][0])
        # slips exist but the late-count filter hides every row (employee 3001 has no lates): no "generate" hint
        hidden = self.run_report("late-salary-impact", employeeIds=str(self.A.id), **P)
        self.assertEqual(hidden["rows"], [])
        self.assertFalse(any("Generate Payroll" in n for n in hidden["notes"]))

    def test_every_report_runs_with_default_filters(self):
        for rid in G2_IDS:
            for user in (self.admin, self.u_b1):
                with self.subTest(report=rid, user=user.username):
                    self.assertEqual(self.client.get(f"/api/reports/run/{rid}", **_headers(user)).status_code, 200)

    def test_role_with_reports_but_no_owning_module_gets_report_forbidden(self):
        for rid in G2_IDS:
            for verb, extra in (("run", {}), ("export", {"fmt": "xlsx"})):
                with self.subTest(report=rid, verb=verb):
                    r = self.client.get(f"/api/reports/{verb}/{rid}", extra, **_headers(self.u_reports_only))
                    self.assertEqual(r.status_code, 403)
                    self.assertEqual(r.json()["error"], "report_forbidden")

    def test_catalog_lists_them_only_for_roles_that_can_open_them(self):
        def ids(user):
            body = self.client.get("/api/reports/catalog", **_headers(user)).json()
            return {r["id"] for r in body["reports"]}

        self.assertTrue(set(G2_IDS) <= ids(self.admin))
        self.assertFalse(set(G2_IDS) & ids(self.u_reports_only))
        comp = ids(self.u_comp_only) & set(G2_IDS)
        self.assertEqual(
            comp,
            {"overtime-register", "overtime-payment-summary", "compensation-credits", "compensation-day-announcements"},
        )
        self.assertNotIn("compensation-credits", ids(self.u_staff_only))
        self.assertNotIn("overtime-register", ids(self.u_staff_only))
        self.assertIn("overtime-payment-summary", ids(self.u_staff_only))
        self.assertIn("pf-statement", ids(self.u_prod_only))
        self.assertNotIn("attendance-salary-loss", ids(self.u_prod_only))

    def test_running_reports_never_writes(self):
        writes = ("INSERT", "UPDATE", "DELETE")
        for rid in G2_IDS:
            with self.subTest(report=rid), CaptureQueriesContext(connection) as ctx:
                self.run_report(rid, **DATA_PARAMS[rid])
            bad = [q["sql"][:120] for q in ctx.captured_queries if q["sql"].lstrip().upper().startswith(writes)]
            self.assertEqual(bad, [], rid)

    def test_exports_are_audited_and_carry_the_filters(self):
        from .models import AuditLog

        before = AuditLog.objects.filter(action="export", module="reports").count()
        self.export("pf-statement", "xlsx", **P)
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), before + 1)
        log = AuditLog.objects.filter(action="export", module="reports").latest("id")
        self.assertIn("PF Statement", log.record_description)
        self.assertIn("XLSX", log.record_description)
        self.assertIn("February 2026", log.record_description)

    def test_invalid_filter_values_are_400(self):
        for rid, q in (("pf-statement", {"wageCeiling": "bogus"}), ("esi-statement", {"coverage": "x"}),
                       ("late-salary-impact", {"minLates": "2"}), ("attendance-salary-loss", {"impact": "??"}),
                       ("payroll-exceptions", {"exceptionType": "nope"}), ("min-wage-compliance", {"basis": "week"}),
                       ("overtime-register", {"status": "maybe"}), ("compensation-credits", {"earnedIn": "never"}),
                       ("payroll-slip-reconciliation", {"issue": "other"})):  # fmt: skip
            r = self.client.get(f"/api/reports/run/{rid}", q, **_headers(self.admin))
            self.assertEqual(r.status_code, 400, rid)
            self.assertEqual(r.json()["error"], "invalid_filter")


# ================================================================================================
#  Adversarial review (second engineer): each test states what a payroll manager would expect; a failing test is
#  a proven defect in the report layer, not in the payroll engine.
# ================================================================================================
class AdversarialReviewTests(_Base):
    def patch_payroll(self, emp, **body):
        p = Payroll.objects.get(employee=emp, month=MONTH, year=YEAR)
        r = self.client.patch(f"/api/payroll/{p.id}", body, content_type="application/json", **_headers(self.admin))
        self.assertEqual(r.status_code, 200, r.content)
        return p

    # -- payroll-slip-reconciliation ---------------------------------------------------------------
    def test_recon_variance_total_row_equals_the_net_variance_card(self):
        """One bonus edit = ONE 500 rupee net difference, however many finding rows describe it."""
        self.patch_payroll(self.L, bonus=500)
        body = self.run_report("payroll-slip-reconciliation", **P)
        self.assertEqual(self.summary(body)["Net variance (payroll - slip)"], 500.0)
        self.assertEqual(body["totals"]["variance"], 500.0, "the Variance total row double counts the same 500")

    def test_recon_variance_is_not_repeated_on_the_advance_finding_of_the_same_employee(self):
        # C: regenerated after the instalment was processed (advance finding) AND marked paid with a bonus
        _generate_staff_payroll(self.C, MONTH, YEAR)
        self.patch_payroll(self.C, bonus=300)
        body = self.run_report("payroll-slip-reconciliation", **P)
        rows = [r for r in self.data_rows(body) if r["employeeCode"] == "3003"]
        self.assertGreaterEqual(len(rows), 2)
        self.assertEqual(body["totals"]["variance"], self.summary(body)["Net variance (payroll - slip)"])

    # -- overtime register vs payment summary ----------------------------------------------------------
    def test_ot_register_and_summary_agree_on_pay_days_still_awaiting_payroll(self):
        _generate_staff_payroll(self.B, MONTH, YEAR)  # picks up the 10-Feb day
        _ot(self.B, date(2026, 2, 11), 65, announced_at=ANNOUNCED)  # announced after that slip: exactly ONE day pending
        reg = self.run_report("overtime-register", employeeIds=str(self.B.id), **OT_FEB)
        summ = self.run_report("overtime-payment-summary", employeeIds=str(self.B.id), **P)
        self.assertEqual(self.by_code(summ)["3002"]["pendingPayDays"], 1)
        self.assertEqual(self.summary(summ)["Pay days pending payroll"], 1)
        self.assertEqual(
            self.summary(reg)["Pay days awaiting payroll"], 1, "the register counts the paid day as awaiting too"
        )

    # -- provisional flag ---------------------------------------------------------------------------------
    def test_a_slip_computed_mid_month_stays_provisional_after_the_month_ends(self):
        mid = datetime(2026, 2, 15, 6, 0, tzinfo=timezone.utc)
        SalarySlip.objects.filter(employee=self.B, month=MONTH, year=YEAR).update(generated_at=mid)
        Payroll.objects.filter(employee=self.B, month=MONTH, year=YEAR).update(updated_at=mid)
        with self.today(date(2026, 3, 15)):
            loss = self.by_code(self.run_report("attendance-salary-loss", **P))
            exc = self.run_report("payroll-exceptions", exceptionType="provisional", **P)
        self.assertEqual(loss["3002"]["monthStatus"], "Provisional", "computed on 15-Feb and never regenerated")
        self.assertIn("3002", self.codes(exc))

    # -- late report: employees vs slips ---------------------------------------------------------------------
    def test_late_impact_counts_employees_not_production_slips(self):
        emp = Employee.objects.create(employee_code="P002", first_name="Prod", last_name="Two", employment_type="production",
                                      status="active", salary_per_shift=Decimal("500"), department=self.d_stitch, branch=self.b1)  # fmt: skip
        summary = {"totalLateCount": 4, "billableLateCount": 1, "shiftDeductions": 0.25}
        for period in (PROD1, PROD2):
            _hand_slip(emp, gross="4000", pf="0", esi="0", late="125", period=period, late_summary=summary)
        body = self.run_report("late-salary-impact", **P)
        s = self.summary(body)
        self.assertEqual(s["Late deduction"], 2000.0)
        self.assertEqual(s["Employees with lates"], 3, "3004, 3002 and P002 (two slips)")
        self.assertEqual(s["Average per deducted employee"], round(2000 / 3, 2))

    # -- announcements: filters --------------------------------------------------------------------------------
    def test_announcement_filters_follow_the_engine_scoping_rules(self):
        def dates(**q):
            with self.today(date(2026, 6, 15)):
                return [
                    r["date"]
                    for r in self.data_rows(self.run_report("compensation-day-announcements", when="all", **q))
                ]

        # an announcement naming employees applies to those employees only: it is not a "CUTTING" day when the
        # only person named (E) works in SEWING
        self.assertNotIn("2026-02-27", dates(departmentIds=str(self.d_cut.id)))
        # a CUTTING-department day (a Unit 1 department) does not apply to any Unit 2 employee
        self.assertNotIn("2026-03-05", dates(branchIds=str(self.b2.id)))

    # -- filter echo: a branch-limited login must not learn who else exists -----------------------------------
    def test_out_of_branch_employee_ids_are_not_echoed_in_the_filter_summary(self):
        for rid, params in (("pf-statement", P), ("payroll-exceptions", P), ("statutory-coverage", {})):
            body = self.run_report(rid, user=self.u_b1, employeeIds=str(self.C.id), **params)
            self.assertEqual(self.data_rows(body), [], rid)
            shown = " ".join(f"{f['label']}={f['value']}" for f in body["filters"])
            self.assertNotIn("3003", shown, f"{rid}: Unit 2 employee {self.C.employee_code} leaked via 'filters'")

    # -- generic guards -------------------------------------------------------------------------------------------
    def test_no_report_builds_a_cell_that_the_exporter_silently_drops(self):
        """runner.normalise_rows keeps declared columns only: a mistyped row key would vanish without any error."""
        from .reporting import runner

        seen: list[tuple[str, set]] = []
        original = runner.normalise_rows

        def spy(columns, rows):
            keys = {c.key for c in columns} | {"_kind"}
            for r in rows:
                seen.append((columns[0].key, set(r) - keys))
            return original(columns, rows)

        with mock.patch.object(runner, "normalise_rows", spy):
            for rid in G2_IDS:
                with self.subTest(report=rid):
                    seen.clear()
                    self.run_report(rid, **DATA_PARAMS[rid])
                    self.assertEqual([extra for _first, extra in seen if extra], [], rid)

    def test_statements_for_a_month_in_progress_say_so(self):
        """Slips generated mid-month count the remaining days as absent: PF / ESI / late figures are provisional."""
        with self.today(date(2026, 2, 15)):
            for rid in ("pf-statement", "esi-statement", "late-salary-impact", "min-wage-compliance"):
                PayrollSettings.objects.filter(pk=1).update(min_wage_rate=Decimal("900"))
                notes = " ".join(self.run_report(rid, **P)["notes"]).lower()
                self.assertTrue(
                    "not ended" in notes or "provisional" in notes, f"{rid}: no warning for a running month"
                )


# ================================================================================================
#  Provisional slips after the month ended, variance counted once, OT days awaiting payroll
# ================================================================================================
class StaleSlipAndTotalsTests(_Base):
    MID = datetime(2026, 2, 15, 6, 0, tzinfo=timezone.utc)

    def _stamp(self, emp, slip_at, payroll_at):
        SalarySlip.objects.filter(employee=emp, month=MONTH, year=YEAR).update(generated_at=slip_at)
        Payroll.objects.filter(employee=emp, month=MONTH, year=YEAR).update(updated_at=payroll_at)

    def test_a_slip_regenerated_after_the_month_ended_is_complete(self):
        # slip first created mid-month, payroll regenerated on 2-Mar (the upsert refreshes Payroll.updated_at only)
        self._stamp(self.B, self.MID, datetime(2026, 3, 2, 6, 0, tzinfo=timezone.utc))
        with self.today(date(2026, 3, 15)):
            loss = self.by_code(self.run_report("attendance-salary-loss", **P))
            exc = self.run_report("payroll-exceptions", exceptionType="provisional", **P)
        self.assertEqual(loss["3002"]["monthStatus"], "Complete")
        self.assertNotIn("3002", self.codes(exc))

    def test_statement_notes_warn_about_stale_slips_only(self):
        self._stamp(self.B, self.MID, self.MID)
        with self.today(date(2026, 3, 15)):
            for rid in ("pf-statement", "esi-statement", "late-salary-impact"):
                notes = [n for n in self.run_report(rid, **P)["notes"] if n.startswith("PROVISIONAL")]
                self.assertEqual(len(notes), 1, rid)
                self.assertIn("1 staff slip(s)", notes[0], rid)
        with self.today(date(2026, 3, 15)):
            PayrollSettings.objects.filter(pk=1).update(min_wage_rate=Decimal("900"))
            self.assertTrue(
                any(n.startswith("PROVISIONAL") for n in self.run_report("min-wage-compliance", **P)["notes"])
            )

    def test_no_provisional_warning_once_every_slip_was_generated_after_the_month(self):
        with self.today(date(2026, 3, 15)):
            for rid in ("pf-statement", "esi-statement", "late-salary-impact", "attendance-salary-loss"):
                notes = " ".join(self.run_report(rid, **P)["notes"])
                self.assertNotIn("PROVISIONAL", notes, rid)

    def test_recon_variance_column_is_filled_on_the_net_differs_finding_only(self):
        p = Payroll.objects.get(employee=self.L, month=MONTH, year=YEAR)
        self.client.patch(
            f"/api/payroll/{p.id}", {"bonus": 500}, content_type="application/json", **_headers(self.admin)
        )
        rows = self.data_rows(self.run_report("payroll-slip-reconciliation", **P))
        self.assertEqual([(r["issue"], r["variance"]) for r in rows], [("Net differs", 500.0), ("Manual edit", None)])
        # the other finding still carries both nets so it can be read on its own
        self.assertEqual((rows[1]["slipNet"], rows[1]["payrollNet"]), (16785.0, 17285.0))
