"""
Report Center - HR Reports (category ``hr``): the two printed registers.

hr-strength-register           "Attendance Abstract For <date>": ON roll / present strength, Male / Female, by department
                               and designation, with SUB TOTAL lines and a GRAND TOTAL.
hr-attendance-register-summary "Attendance Register From <date> To <date>": Present, OT, Total and Leave Days per employee,
                               grouped by department, with SUB TOTAL lines and a GRAND TOTAL.

"Today" is pinned to 2026-09-29 (a Tuesday) so every expected value is hand-derived and deterministic.

Run via: python manage.py test api.tests_reporting_hr_strength -v 2
"""

import io
from datetime import date, time
from decimal import Decimal
from unittest.mock import patch

import pdfplumber
from django.apps import apps
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .models import (
    AttendanceDayRecord,
    AttendanceLog,
    Branch,
    Department,
    Designation,
    Employee,
    Holiday,
    HRUser,
    OvertimeRecord,
    PayrollSettings,
    ResignationRequest,
    Role,
)
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.export_xlsx import HEADER_ROW

TODAY = date(2026, 9, 29)  # Tuesday
SID = "hr-strength-register"
RID = "hr-attendance-register-summary"

STRENGTH_LABELS = [
    "S.No.",
    "Department",
    "Designation",
    "ON Roll Strength - Male",
    "ON Roll Strength - Female",
    "Present Strength - Male",
    "Present Strength - Female",
]
STRENGTH_KEYS = ["sno", "department", "designation", "onRollMale", "onRollFemale", "presentMale", "presentFemale"]
REGISTER_LABELS = [
    "SNo",
    "Department",
    "TicketNo",
    "EmpName",
    "Designation",
    "Present",
    "OT",
    "Total",
    "Leave Days",
]
REGISTER_KEYS = [
    "sno",
    "department",
    "employeeCode",
    "employeeName",
    "designation",
    "present",
    "ot",
    "total",
    "leaveDays",
]


def hdr(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def rec(emp, day, status, shifts="0"):
    return AttendanceDayRecord.objects.create(
        employee=emp, date=day, status=status, shifts_earned=Decimal(shifts), total_punches=0
    )


def punch(emp, day, hh=9, mm=0):
    AttendanceLog.objects.create(
        employee=emp, date=day, punch_time=time(hh, mm), punch_type="IN", source="biometric:dev1"
    )


def data_rows(body) -> list[dict]:
    return [r for r in body["rows"] if not r.get("_kind")]


def subtotals(body) -> list[dict]:
    return [r for r in body["rows"] if r.get("_kind") == "subtotal"]


def cards(body) -> dict:
    return {c["label"]: c["value"] for c in body["summary"]}


class _Base(TestCase):
    report_id = ""

    def setUp(self):
        p = patch("api.reporting.filters.ist_today", return_value=TODAY)
        p.start()
        self.addCleanup(p.stop)

    @classmethod
    def make_users(cls):
        cls.admin = HRUser.objects.create(username="hr_admin", password_hash="x", is_super_admin=True)

        def user(name, perms, **kw):
            role = Role.objects.create(name=name, permissions={"reports": "view", **perms})
            return HRUser.objects.create(username=name, password_hash="x", role=role, **kw)

        cls.attendance_user = user("hr_att", {"attendance": "view"})
        cls.employees_user = user("hr_emp", {"employees": "view"})
        cls.salary_user = user("hr_sal", {"salary": "view", "payroll": "view"})
        cls.plain_user = user("hr_plain", {})
        cls.branch_user = user("hr_branch", {"attendance": "view"}, branch=cls.b1)

    def get(self, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/run/{self.report_id}", params, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, fmt, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/export/{self.report_id}", {"fmt": fmt, **params}, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r

    def excel_rows(self, **params) -> list[list]:
        ws = load_workbook(io.BytesIO(self.export("xlsx", **params).content)).active
        out = []
        for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            if all(v is None for v in row):
                break  # the notes follow after a blank line
            out.append(list(row))
        return out

    def assert_excel_equals_screen(self, **params):
        screen = self.get(**params)
        excel = self.excel_rows(**params)
        keys = [c["key"] for c in screen["columns"]]
        self.assertEqual(len(excel), len(screen["rows"]))  # sub totals and the grand total are in the file too
        for shown, cells in zip(screen["rows"], excel, strict=True):
            for key, cell in zip(keys, cells, strict=True):
                want = shown.get(key)
                self.assertEqual(cell if cell != "" else None, want, (key, shown))

    def assert_never_writes(self, **params):
        def counts():
            return {
                m.__name__: m.objects.count()
                for m in apps.get_app_config("api").get_models()
                if m.__name__ != "AuditLog"  # an export is supposed to leave an audit entry
            }

        before = counts()
        self.get(**params)
        self.export("xlsx", **params)
        self.export("pdf", **params)
        self.assertEqual(counts(), before)

    def pdf_text(self, **params) -> str:
        r = self.export("pdf", **params)
        self.assertTrue(r.content.startswith(b"%PDF"))
        with pdfplumber.open(io.BytesIO(r.content)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)


# ═════════════════════════════════════════════════════════════════════════════
# Strength register
# ═════════════════════════════════════════════════════════════════════════════


class _StrengthWorld(_Base):
    """Thirteen people on 29-09-2026: every on-roll / present / gender situation the register can meet."""

    report_id = SID

    @classmethod
    def setUpTestData(cls):
        PayrollSettings.get()  # a GET must not be the thing that creates it
        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.acc = Department.objects.create(name="ACCOUNTS", branch=cls.b1)
        cls.adm = Department.objects.create(name="ADMIN", branch=cls.b1)
        cls.cut = Department.objects.create(name="CUTTING", branch=cls.b2)
        # created in this order: the register lists a department's designations in that order
        cls.manager = Designation.objects.create(title="MANAGER", department=cls.acc)
        cls.executive = Designation.objects.create(title="EXECUTIVE", department=cls.acc)
        cls.officer = Designation.objects.create(title="OFFICER", department=cls.adm)
        cls.designer = Designation.objects.create(title="DESIGNER", department=cls.adm)

        def mk(code, gender, dept, desig, branch=None, **kw):
            return Employee.objects.create(
                employee_code=code, first_name=code, last_name="X", gender=gender, department=dept, designation=desig,
                branch=branch or cls.b1, join_date=kw.pop("join_date", "2024-01-01"), **kw,
            )  # fmt: skip

        cls.a1 = mk("A1", "male", cls.acc, cls.manager)
        cls.a2 = mk("A2", "male", cls.acc, cls.manager)
        cls.a3 = mk("A3", "female", cls.acc, cls.executive)
        cls.a4 = mk("A4", "female", cls.acc, cls.executive)
        cls.a5 = mk("A5", "female", cls.acc, cls.manager)
        cls.d1 = mk("D1", "male", cls.adm, cls.officer)
        cls.d2 = mk("D2", "female", cls.adm, cls.designer)
        cls.d3 = mk("D3", "male", cls.adm, cls.designer, join_date="2026-10-01")  # joins after the date
        cls.d4 = mk("D4", "", cls.adm, cls.officer)  # no gender on record
        cls.d5 = mk("D5", "male", cls.adm, cls.officer, status="inactive")  # left, nothing stored
        cls.d6 = mk("D6", "Female", cls.adm, cls.officer, status="inactive")  # left, but the day has a record
        cls.d7 = mk("D7", "Male", cls.adm, cls.designer, join_date="29-09-2026")  # typed dd-mm-yyyy, joins that day
        cls.c1 = mk("C1", "male", cls.cut, None, branch=cls.b2, employment_type="production")

        rec(cls.a1, TODAY, "present", "1")
        rec(cls.a2, TODAY, "absent")
        punch(cls.a3, TODAY)  # no verdict yet: the punch makes her present
        rec(cls.a4, TODAY, "half_shift", "0.5")
        rec(cls.d1, TODAY, "on_leave")
        punch(cls.d3, TODAY)  # not on roll yet, so never present
        rec(cls.d4, TODAY, "present", "1")
        rec(cls.d6, TODAY, "absent")
        rec(cls.d7, TODAY, "present", "1")
        punch(cls.c1, TODAY, 8, 30)
        punch(cls.a5, date(2026, 9, 28))  # yesterday's punch is not today's

        cls.make_users()

    def lines(self, body) -> list[tuple]:
        """(department, designation, on roll M, F[, other], present M, F[, other]) of every line, sub totals included."""
        keys = [c["key"] for c in body["columns"] if c["key"] not in ("sno", "department", "designation")]
        return [
            (r["department"], r["designation"], *(r[k] for k in keys))
            for r in body["rows"]
        ]  # fmt: skip


class StrengthCatalogTests(_StrengthWorld):
    def test_it_is_registered_in_the_hr_group_with_the_printed_columns(self):
        spec = next(s for s in registry.all_specs() if s.id == SID)
        self.assertEqual(registry.LOAD_ERRORS, {})
        self.assertEqual((spec.title, spec.category), ("Strength Register (Attendance Abstract)", "hr"))
        self.assertEqual(spec.modules, ("attendance", "employees"))
        self.assertTrue(set(spec.modules) <= set(all_module_keys()))
        self.assertEqual([c.label for c in spec.columns], STRENGTH_LABELS)
        self.assertEqual([c.key for c in spec.columns], STRENGTH_KEYS)

    def test_the_catalog_offers_one_date_the_scope_filters_and_the_split_switch(self):
        body = self.client.get("/api/reports/catalog", **hdr(self.admin)).json()
        report = next(r for r in body["reports"] if r["id"] == SID)
        self.assertEqual(report["category"], "hr")
        self.assertEqual(
            [(f["key"], f["kind"]) for f in report["filters"]],
            [
                ("dateRange", "dateRange"),
                ("branch", "branch"),
                ("department", "department"),
                ("designation", "designation"),
                ("employmentType", "employmentType"),
                ("splitByGender", "boolean"),
            ],
        )
        date_filter = next(f for f in report["filters"] if f["key"] == "dateRange")
        self.assertEqual(date_filter["label"], "Date")
        self.assertEqual(
            date_filter["default"]["dateFrom"], date_filter["default"]["dateTo"]
        )  # opens on one day (today)

    def test_attendance_or_employee_access_opens_it_and_a_reports_only_role_does_not(self):
        for who in (self.attendance_user, self.employees_user, self.admin):
            self.assertTrue(self.get(user=who)["rows"], who.username)
            self.assertEqual(self.export("xlsx", user=who).status_code, 200, who.username)
        for who in (self.salary_user, self.plain_user):
            catalog = self.client.get("/api/reports/catalog", **hdr(who)).json()
            self.assertNotIn(SID, [r["id"] for r in catalog["reports"]], who.username)
            self.get(user=who, expect=403)
            self.export("pdf", user=who, expect=403)

    def test_it_needs_an_hr_login(self):
        self.assertEqual(self.client.get(f"/api/reports/run/{SID}").status_code, 401)


class StrengthRowTests(_StrengthWorld):
    def test_golden_register_one_line_per_designation_and_gender(self):
        body = self.get()
        self.assertEqual(
            [c["label"] for c in body["columns"]],
            [
                *STRENGTH_LABELS[:5],
                "ON Roll Strength - Other / Not set",
                *STRENGTH_LABELS[5:],
                "Present Strength - Other / Not set",
            ],  # D4 has no gender: the third column appears
        )
        self.assertEqual(
            self.lines(body),
            [
                # (department, designation, roll M, F, other, present M, F, other)
                ("ACCOUNTS", "MANAGER", 2, 0, 0, 1, 0, 0),
                ("ACCOUNTS", "MANAGER", 0, 1, 0, 0, 0, 0),
                ("ACCOUNTS", "EXECUTIVE", 0, 2, 0, 0, 2, 0),
                ("ACCOUNTS", "SUB TOTAL", 2, 3, 0, 1, 2, 0),
                ("ADMIN", "OFFICER", 1, 0, 0, 0, 0, 0),  # D1 is on leave: on roll, not present
                ("ADMIN", "OFFICER", 0, 1, 0, 0, 0, 0),  # D6 left, but the day has a record: on roll
                ("ADMIN", "OFFICER", 0, 0, 1, 0, 0, 1),  # D4: no gender
                ("ADMIN", "DESIGNER", 1, 0, 0, 1, 0, 0),  # D7 joins that very day; D3 joins later: not on roll
                ("ADMIN", "DESIGNER", 0, 1, 0, 0, 0, 0),
                ("ADMIN", "SUB TOTAL", 2, 2, 1, 1, 0, 1),
                ("CUTTING", "No designation", 1, 0, 0, 1, 0, 0),
                ("CUTTING", "SUB TOTAL", 1, 0, 0, 1, 0, 0),
                (None, "GRAND TOTAL", 5, 5, 1, 3, 2, 1),
            ],
        )

    def test_lines_are_numbered_continuously_across_departments_and_total_lines_are_marked(self):
        body = self.get()
        self.assertEqual([r["sno"] for r in data_rows(body)], list(range(1, 10)))
        self.assertEqual([r["designation"] for r in subtotals(body)], ["SUB TOTAL"] * 3)
        self.assertEqual([r["department"] for r in subtotals(body)], ["ACCOUNTS", "ADMIN", "CUTTING"])
        last = body["rows"][-1]
        self.assertEqual((last["_kind"], last["designation"]), ("total", "GRAND TOTAL"))
        self.assertIsNone(body["totals"])  # the GRAND TOTAL line is the total: no second TOTAL row under it

    def test_without_the_split_a_designation_is_one_line_men_and_women_side_by_side(self):
        body = self.get(splitByGender="false")
        self.assertEqual(
            self.lines(body),
            [
                ("ACCOUNTS", "MANAGER", 2, 1, 0, 1, 0, 0),
                ("ACCOUNTS", "EXECUTIVE", 0, 2, 0, 0, 2, 0),
                ("ACCOUNTS", "SUB TOTAL", 2, 3, 0, 1, 2, 0),
                ("ADMIN", "OFFICER", 1, 1, 1, 0, 0, 1),
                ("ADMIN", "DESIGNER", 1, 1, 0, 1, 0, 0),
                ("ADMIN", "SUB TOTAL", 2, 2, 1, 1, 0, 1),
                ("CUTTING", "No designation", 1, 0, 0, 1, 0, 0),
                ("CUTTING", "SUB TOTAL", 1, 0, 0, 1, 0, 0),
                (None, "GRAND TOTAL", 5, 5, 1, 3, 2, 1),
            ],
        )

    def test_the_third_gender_column_only_exists_when_someone_on_roll_needs_it(self):
        self.d4.delete()
        body = self.get()
        self.assertEqual([c["label"] for c in body["columns"]], STRENGTH_LABELS)
        self.assertEqual(self.lines(body)[-1], (None, "GRAND TOTAL", 5, 5, 3, 2))
        self.assertIn("Every on-roll employee has Male or Female on record", " ".join(body["notes"]))
        # an unknown gender text counts as "other" too
        Employee.objects.filter(employee_code="A5").update(gender="Transgender")
        body = self.get()
        self.assertEqual(len(body["columns"]), 9)
        self.assertEqual(self.lines(body)[-1][:2], (None, "GRAND TOTAL"))
        self.assertEqual(cards(body)["Other / not set on roll"], 1)

    def test_on_roll_and_present_are_not_the_same_thing(self):
        body = self.get()
        self.assertEqual(
            cards(body),
            {
                "On roll": 11,
                "Present": 6,
                "Not present": 5,
                "Male on roll": 5,
                "Female on roll": 5,
                "Other / not set on roll": 1,
            },
        )

    def test_someone_who_joined_after_the_date_is_not_on_roll_and_on_the_day_is(self):
        body = self.get()
        designer_male = next(r for r in data_rows(body) if r["designation"] == "DESIGNER" and r["onRollMale"])
        self.assertEqual(designer_male["onRollMale"], 1)  # D7 only, D3 (joins 1-Oct) is out although he punched
        Employee.objects.filter(employee_code="D3").update(join_date="2026-09-29")
        body = self.get()
        designer_male = next(r for r in data_rows(body) if r["designation"] == "DESIGNER" and r["onRollMale"])
        self.assertEqual((designer_male["onRollMale"], designer_male["presentMale"]), (2, 2))  # now on roll, punched
        Employee.objects.filter(employee_code="D3").update(join_date="soon")  # unreadable: counted, and said so
        body = self.get()
        self.assertTrue(any("no readable joining date" in n for n in body["notes"]), body["notes"])

    def test_inactive_employees_count_only_on_the_days_they_still_belonged(self):
        officers = lambda b: [  # noqa: E731
            (r["onRollMale"], r["onRollFemale"]) for r in data_rows(b) if r["designation"] == "OFFICER"
        ]
        self.assertEqual(officers(self.get()), [(1, 0), (0, 1), (0, 0)])  # D5 (nothing stored) is out, D6 is in
        # an approved last working day after the date keeps a leaver on the roll
        ResignationRequest.objects.create(employee=self.d5, status="approved", last_working_date=date(2026, 9, 30))
        self.assertEqual(officers(self.get()), [(2, 0), (0, 1), (0, 0)])
        # ... and one before the date takes an active employee off it
        ResignationRequest.objects.create(employee=self.a2, status="approved", last_working_date=date(2026, 9, 28))
        managers = [(r["onRollMale"], r["presentMale"]) for r in data_rows(self.get()) if r["designation"] == "MANAGER"]
        self.assertEqual(managers[0], (1, 1))

    def test_a_day_with_a_verdict_and_a_punch_is_counted_once(self):
        punch(self.a1, TODAY)
        punch(self.a1, TODAY, 17, 0)
        body = self.get()
        self.assertEqual(self.lines(body)[-1], (None, "GRAND TOTAL", 5, 5, 1, 3, 2, 1))

    def test_another_date_reads_that_dates_own_verdicts(self):
        yesterday = date(2026, 9, 28)
        rec(self.a2, yesterday, "present", "1")
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-28")
        manager_male = next(r for r in data_rows(body) if r["designation"] == "MANAGER" and r["onRollMale"])
        self.assertEqual((manager_male["onRollMale"], manager_male["presentMale"]), (2, 1))  # A2 only
        woman = next(r for r in data_rows(body) if r["designation"] == "MANAGER" and r["onRollFemale"])
        self.assertEqual(woman["presentFemale"], 1)  # A5's punch of yesterday

    def test_a_future_date_has_nobody_present(self):
        body = self.get(dateFrom="2026-10-05", dateTo="2026-10-05")
        self.assertEqual(cards(body)["Present"], 0)
        self.assertEqual(cards(body)["On roll"], 11)  # D3 has joined by then
        self.assertTrue(any("in the future" in n for n in body["notes"]))

    def test_the_notes_state_the_rules(self):
        text = " ".join(self.get()["notes"])
        for part in (
            "29-Sep-2026",
            "ON ROLL = employees employed that day",
            "joins after the date is not on roll",
            "PRESENT = on-roll employees with a stored Present or Half-day verdict",
            "counts as not present",
            "2 employee(s) are counted present from a punch alone",
        ):
            self.assertIn(part, text)

    def test_nobody_gives_an_empty_register_not_an_error(self):
        Employee.objects.all().delete()
        body = self.get()
        self.assertEqual(body["rows"], [])
        self.assertEqual(cards(body)["On roll"], 0)
        self.assertEqual([c["label"] for c in body["columns"]], STRENGTH_LABELS)

    def test_an_employee_with_no_department_is_listed_last(self):
        Employee.objects.create(
            employee_code="Z9", first_name="Z", last_name="Z", gender="male", join_date="2024-01-01",
            branch=self.b1, designation=self.officer,
        )  # fmt: skip
        body = self.get()
        self.assertEqual(self.lines(body)[-2][:2], ("Unassigned", "SUB TOTAL"))
        self.assertEqual(self.lines(body)[-3][:2], ("Unassigned", "OFFICER"))

    def test_same_named_departments_of_two_units_stay_apart(self):
        other = Department.objects.create(name="ADMIN", branch=self.b2)
        Employee.objects.create(
            employee_code="Y9", first_name="Y", last_name="Y", gender="male", join_date="2024-01-01",
            branch=self.b2, department=other, designation=self.officer,
        )  # fmt: skip
        depts = [r["department"] for r in subtotals(self.get())]
        self.assertIn("ADMIN (Unit 1)", depts)
        self.assertIn("ADMIN (Unit 2)", depts)


class StrengthFilterTests(_StrengthWorld):
    def test_an_abstract_is_for_the_to_date_and_says_so_when_a_range_is_sent(self):
        body = self.get(dateFrom="2026-09-28", dateTo="2026-09-29")
        self.assertEqual(cards(body)["On roll"], 11)  # the figures of the 29th
        self.assertTrue(any("the To date, 29-Sep-2026, is used" in n for n in body["notes"]), body["notes"])
        self.assertFalse(any("the To date" in n for n in self.get()["notes"]))  # one day: nothing to explain

    def test_bad_or_too_wide_dates_are_a_400_naming_the_field(self):
        r = self.client.get(
            f"/api/reports/run/{SID}", {"dateFrom": "2026-08-01", "dateTo": "2026-09-29"}, **hdr(self.admin)
        )
        self.assertEqual((r.status_code, r.json()["field"]), (400, "dateTo"))
        for params in (
            {"dateFrom": "nope", "dateTo": "2026-09-29"},
            {"dateFrom": "2026-09-29", "dateTo": "2026-09-01"},
        ):
            self.assertEqual(self.client.get(f"/api/reports/run/{SID}", params, **hdr(self.admin)).status_code, 400)

    def test_department_designation_branch_and_type_filters_narrow_the_register(self):
        body = self.get(departmentIds=str(self.cut.id))
        self.assertEqual([r["department"] for r in subtotals(body)], ["CUTTING"])
        self.assertEqual(self.lines(body)[-1], (None, "GRAND TOTAL", 1, 0, 1, 0))
        body = self.get(designationIds=str(self.executive.id))
        self.assertEqual(self.lines(body)[-1], (None, "GRAND TOTAL", 0, 2, 0, 2))
        body = self.get(branchIds=str(self.b2.id))
        self.assertEqual([r["department"] for r in subtotals(body)], ["CUTTING"])
        body = self.get(employmentType="production")
        self.assertEqual(self.lines(body)[-1], (None, "GRAND TOTAL", 1, 0, 1, 0))
        body = self.get(employmentType="staff")
        self.assertEqual(self.lines(body)[-1][:2], (None, "GRAND TOTAL"))
        self.assertEqual(cards(body)["On roll"], 10)

    def test_the_chosen_filters_are_echoed(self):
        body = self.get(departmentIds=str(self.acc.id))
        applied = {f["label"]: f["value"] for f in body["filters"]}
        self.assertEqual(applied["Department"], "ACCOUNTS")
        self.assertEqual(applied["Date"], "29-Sep-2026 to 29-Sep-2026")

    def test_bad_ids_are_a_400(self):
        for params in ({"departmentIds": "x"}, {"employmentType": "contract"}):
            self.assertEqual(self.client.get(f"/api/reports/run/{SID}", params, **hdr(self.admin)).status_code, 400)


class StrengthBranchIsolationTests(_StrengthWorld):
    def test_a_branch_user_sees_only_their_own_branch(self):
        body = self.get(user=self.branch_user)
        self.assertNotIn("CUTTING", str(body["rows"]))
        self.assertEqual(self.lines(body)[-1], (None, "GRAND TOTAL", 4, 5, 1, 2, 2, 1))
        self.assertEqual(cards(body)["On roll"], 10)

    def test_a_branch_user_cannot_widen_their_scope_with_a_parameter(self):
        body = self.get(user=self.branch_user, branchIds=str(self.b2.id))
        self.assertEqual(body["rows"], [])
        body = self.get(user=self.branch_user, departmentIds=str(self.cut.id))
        self.assertEqual(body["rows"], [])

    def test_a_branch_user_gets_no_branch_filter_and_a_scoped_excel(self):
        catalog = self.client.get("/api/reports/catalog", **hdr(self.branch_user)).json()
        report = next(r for r in catalog["reports"] if r["id"] == SID)
        self.assertNotIn("branch", [f["key"] for f in report["filters"]])
        ws = load_workbook(io.BytesIO(self.export("xlsx", user=self.branch_user).content)).active
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertNotIn("CUTTING", cells)
        self.assertIn("ACCOUNTS", cells)

    def test_another_branchs_punches_and_verdicts_do_not_leak_into_a_branch_users_counts(self):
        rec(self.c1, TODAY, "present", "1")
        body = self.get(user=self.branch_user)
        self.assertEqual(cards(body)["Present"], 5)


class StrengthExportAndSafetyTests(_StrengthWorld):
    def test_excel_carries_the_same_columns_and_rows_as_the_screen(self):
        self.assert_excel_equals_screen()
        self.assert_excel_equals_screen(splitByGender="false")
        self.assert_excel_equals_screen(departmentIds=str(self.adm.id))
        ws = load_workbook(io.BytesIO(self.export("xlsx").content)).active
        self.assertEqual([c.value for c in ws[HEADER_ROW]][:5], [*STRENGTH_LABELS[:5]])
        labels = [r[2] for r in self.excel_rows()]
        self.assertEqual(labels.count("SUB TOTAL"), 3)
        self.assertEqual(labels[-1], "GRAND TOTAL")

    def test_the_pdf_is_a_real_pdf_with_the_printed_headings_and_totals(self):
        text = self.pdf_text()
        for part in ("Strength Register (Attendance Abstract)", "SUB TOTAL", "GRAND TOTAL", "ACCOUNTS", "OFFICER"):
            self.assertIn(part, text)

    def test_running_and_exporting_never_writes(self):
        self.assert_never_writes()

    def test_the_query_count_does_not_grow_with_the_number_of_employees(self):
        def queries() -> int:
            self.get()  # warm-up (permission caches)
            with CaptureQueriesContext(connection) as q:
                self.get()
            return len(q)

        before = queries()
        for i in range(30):
            emp = Employee.objects.create(
                employee_code=f"X{i}", first_name=f"X{i}", last_name="Extra", gender=("male", "female", "")[i % 3],
                branch=self.b1, department=(self.acc, self.adm)[i % 2], designation=(self.manager, self.officer)[i % 2],
                join_date="2024-01-01",
            )  # fmt: skip
            if i % 2:
                rec(emp, TODAY, "present", "1")
            else:
                punch(emp, TODAY)
        self.assertEqual(queries(), before)
        self.assertEqual(cards(self.get())["On roll"], 11 + 30)

    def test_photo_and_password_columns_are_never_selected(self):
        with CaptureQueriesContext(connection) as q:
            self.get()
            self.export("xlsx")
        for query in q.captured_queries:
            self.assertNotIn("photo_url", query["sql"])
            if '"hr_users"' not in query["sql"]:
                self.assertNotIn("password_hash", query["sql"])


# ═════════════════════════════════════════════════════════════════════════════
# Attendance Register (Summary)
# ═════════════════════════════════════════════════════════════════════════════

FROM, TO = "2026-09-21", "2026-09-27"  # Monday .. Sunday
RANGE = {"dateFrom": FROM, "dateTo": TO}


def ot(emp, day, shifts, *, status="announced", kind="pay", ot_type="shift"):
    return OvertimeRecord.objects.create(
        employee=emp, date=day, last_punch_out=time(20, 0), ot_minutes=120, status=status, compensation_type=kind,
        ot_type=ot_type, shifts=None if shifts is None else Decimal(shifts),
    )  # fmt: skip


class _RegisterWorld(_Base):
    report_id = RID

    @classmethod
    def setUpTestData(cls):
        PayrollSettings.get()
        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.acc = Department.objects.create(name="ACCOUNTS", branch=cls.b1)
        cls.adm = Department.objects.create(name="ADMIN", branch=cls.b1)
        cls.cut = Department.objects.create(name="CUTTING", branch=cls.b2)
        cls.des = Designation.objects.create(title="AI ENGINEER", department=cls.acc)
        Holiday.objects.create(name="Festival", date=date(2026, 9, 24))

        def mk(code, name, dept, branch=None, **kw):
            return Employee.objects.create(
                employee_code=code, first_name=name, last_name=".M", department=dept, branch=branch or cls.b1,
                designation=cls.des, join_date=kw.pop("join_date", "2024-01-01"), **kw,
            )  # fmt: skip

        cls.r1 = mk("30020", "SURYA", cls.acc)
        cls.r2 = mk("30021", "RAVI", cls.acc)
        cls.r3 = mk("30100", "PRIYA", cls.adm, employment_type="production")
        cls.r4 = mk("30200", "NOBODY", None)
        cls.r5 = mk("30300", "LEFT", cls.adm, status="inactive")
        cls.r6 = mk("30400", "NEWBIE", cls.adm, join_date="2026-09-23")
        cls.r7 = mk("30500", "OTHERUNIT", cls.cut, branch=cls.b2)

        # R1: 21 P, 22 P, 23 half, 24 company holiday, 25 absent, 26 on leave, 27 Sunday (holiday verdict, no Holiday row)
        for d, s, v in (
            (21, "present", "1"), (22, "present", "1"), (23, "half_shift", "0.5"), (24, "holiday", "0"),
            (25, "absent", "0"), (26, "on_leave", "0"), (27, "holiday", "0"),
        ):  # fmt: skip
            rec(cls.r1, date(2026, 9, d), s, v)
        rec(cls.r2, date(2026, 9, 21), "present", "1")
        rec(cls.r2, date(2026, 9, 22), "absent")  # 23..27 never opened: no record
        rec(cls.r3, date(2026, 9, 21), "present", "1.5")  # production: shifts, not days
        rec(cls.r3, date(2026, 9, 22), "half_shift", "0.5")
        rec(cls.r3, date(2026, 9, 27), "absent")  # a production Sunday nobody worked is an absence
        rec(cls.r4, date(2026, 9, 21), "present", "1")
        rec(cls.r6, date(2026, 9, 22), "absent")  # before the joining date: ignored
        rec(cls.r6, date(2026, 9, 23), "present", "1")
        rec(cls.r7, date(2026, 9, 21), "present", "1")

        ot(cls.r1, date(2026, 9, 22), "0.50")  # Type 2 shift record, announced as pay: counts 0.50
        ot(cls.r1, date(2026, 9, 23), None, ot_type="legacy")  # a record from before shift counting: one shift
        ot(cls.r1, date(2026, 9, 21), "1", kind="relaxation")  # relaxation: never pays
        ot(cls.r1, date(2026, 9, 25), "1", status="detected", kind=None)  # only detected: not decided
        ot(cls.r1, date(2026, 9, 26), "0.25", ot_type="allowance")  # the automatic next-day allowance: never
        ot(cls.r1, date(2026, 9, 28), "1")  # outside the dates

        cls.make_users()


class RegisterCatalogTests(_RegisterWorld):
    def test_it_is_registered_in_the_hr_group_with_the_printed_columns(self):
        spec = next(s for s in registry.all_specs() if s.id == RID)
        self.assertEqual(
            (spec.title, spec.category, spec.modules), ("Attendance Register (Summary)", "hr", ("attendance",))
        )
        self.assertEqual([c.label for c in spec.columns], REGISTER_LABELS)
        self.assertEqual([c.key for c in spec.columns], REGISTER_KEYS)
        self.assertTrue(set(spec.modules) <= set(all_module_keys()))

    def test_the_catalog_offers_a_range_no_wider_than_a_month_and_the_scope_filters(self):
        body = self.client.get("/api/reports/catalog", **hdr(self.admin)).json()
        report = next(r for r in body["reports"] if r["id"] == RID)
        self.assertEqual(
            [(f["key"], f["kind"]) for f in report["filters"]],
            [
                ("dateRange", "dateRange"),
                ("branch", "branch"),
                ("department", "department"),
                ("designation", "designation"),
                ("employmentType", "employmentType"),
                ("employee", "employee"),
                ("employeeStatus", "employeeStatus"),
            ],
        )
        date_filter = next(f for f in report["filters"] if f["key"] == "dateRange")
        self.assertEqual(date_filter["maxDays"], 31)
        self.assertTrue(date_filter["default"]["dateFrom"].endswith("-01"))  # opens on the 1st of this month

    def test_attendance_access_opens_it_and_other_roles_do_not(self):
        for who in (self.attendance_user, self.admin):
            self.assertTrue(self.get(user=who, **RANGE)["rows"], who.username)
        for who in (self.employees_user, self.salary_user, self.plain_user):
            catalog = self.client.get("/api/reports/catalog", **hdr(who)).json()
            self.assertNotIn(RID, [r["id"] for r in catalog["reports"]], who.username)
            self.get(user=who, expect=403, **RANGE)

    def test_it_needs_an_hr_login(self):
        self.assertEqual(self.client.get(f"/api/reports/run/{RID}").status_code, 401)

    def test_a_range_wider_than_a_month_is_a_400_naming_the_field(self):
        r = self.client.get(
            f"/api/reports/run/{RID}", {"dateFrom": "2026-08-01", "dateTo": "2026-09-29"}, **hdr(self.admin)
        )
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["field"], "dateTo")
        self.assertIn("max 31", r.json()["message"])


class RegisterRowTests(_RegisterWorld):
    def lines(self, body):
        return [
            (r.get("sno"), r["department"], r.get("employeeCode"), r.get("employeeName"), r["present"], r["ot"], r["total"], r["leaveDays"])
            for r in body["rows"]
        ]  # fmt: skip

    def test_golden_register_grouped_by_department_with_totals(self):
        body = self.get(**RANGE)
        self.assertEqual([c["label"] for c in body["columns"]], REGISTER_LABELS)
        self.assertEqual(
            self.lines(body),
            [
                # sno, department, ticket, name, present, ot, total, leave
                (1, "ACCOUNTS", "30020", "SURYA .M", 2.5, 1.5, 4.0, 2.5),
                (2, "ACCOUNTS", "30021", "RAVI .M", 1.0, 0.0, 1.0, 1.0),
                (None, "ACCOUNTS", None, "SUB TOTAL", 3.5, 1.5, 5.0, 3.5),
                (3, "ADMIN", "30100", "PRIYA .M", 2.0, 0.0, 2.0, 1.5),
                (4, "ADMIN", "30400", "NEWBIE .M", 1.0, 0.0, 1.0, 0.0),
                (None, "ADMIN", None, "SUB TOTAL", 3.0, 0.0, 3.0, 1.5),
                (5, "CUTTING", "30500", "OTHERUNIT .M", 1.0, 0.0, 1.0, 0.0),
                (None, "CUTTING", None, "SUB TOTAL", 1.0, 0.0, 1.0, 0.0),
                (6, "Unassigned", "30200", "NOBODY .M", 1.0, 0.0, 1.0, 0.0),
                (None, "Unassigned", None, "SUB TOTAL", 1.0, 0.0, 1.0, 0.0),
                (None, None, None, "GRAND TOTAL", 8.5, 1.5, 10.0, 5.0),
            ],
        )
        self.assertEqual(body["rows"][-1]["_kind"], "total")
        self.assertIsNone(body["totals"])

    def test_the_leave_rule_counts_absent_and_leave_whole_and_a_half_shift_half(self):
        row = next(r for r in data_rows(self.get(**RANGE)) if r["employeeCode"] == "30020")
        # 23 half-shift 0.5 + 25 absent 1 + 26 on leave 1 = 2.5; the company holiday (24) and the Sunday (27) are not leave
        self.assertEqual(row["leaveDays"], 2.5)
        self.assertEqual(row["present"], 2.5)  # 1 + 1 + 0.5 shifts earned

    def test_a_staff_saturday_off_is_not_leave(self):
        from .models import EmployeeShiftAssignment, ShiftTemplate

        shift = ShiftTemplate.objects.create(
            name="General", shift_type="staff", start_time=time(9, 0), end_time=time(18, 0)
        )
        EmployeeShiftAssignment.objects.create(
            employee=self.r2, shift=shift, effective_from=date(2026, 1, 1), saturday_off=True
        )
        rec(self.r2, date(2026, 9, 26), "absent")  # the engine stores a Saturday-off Saturday as absent
        row = next(r for r in data_rows(self.get(**RANGE)) if r["employeeCode"] == "30021")
        self.assertEqual(row["leaveDays"], 1.0)  # only the 22nd

    def test_a_production_sunday_with_nobody_there_is_an_absence(self):
        row = next(r for r in data_rows(self.get(**RANGE)) if r["employeeCode"] == "30100")
        # the half shift on the 22nd (0.5) + the empty Sunday (1)
        self.assertEqual((row["present"], row["leaveDays"]), (2.0, 1.5))

    def test_ot_is_what_payroll_pays_and_nothing_else(self):
        row = next(r for r in data_rows(self.get(**RANGE)) if r["employeeCode"] == "30020")
        # 0.50 (shift record, pay) + 1.00 (older announced record); not the relaxation, the detected one, the
        # automatic allowance or the record outside the dates
        self.assertEqual((row["ot"], row["total"]), (1.5, 4.0))

    def test_ot_is_zero_and_said_so_when_the_compensation_feature_is_off(self):
        PayrollSettings.objects.filter(pk=1).update(compensation_feature_enabled=False)
        body = self.get(**RANGE)
        row = next(r for r in data_rows(body) if r["employeeCode"] == "30020")
        self.assertEqual((row["ot"], row["total"]), (0.0, 2.5))
        self.assertTrue(any("Compensation feature is switched off" in n for n in body["notes"]), body["notes"])

    def test_days_before_joining_and_after_an_approved_last_day_are_left_out(self):
        row = next(r for r in data_rows(self.get(**RANGE)) if r["employeeCode"] == "30400")
        self.assertEqual(
            (row["present"], row["leaveDays"]), (1.0, 0.0)
        )  # the absent record of the 22nd is before joining
        ResignationRequest.objects.create(employee=self.r1, status="approved", last_working_date=date(2026, 9, 22))
        row = next(r for r in data_rows(self.get(**RANGE)) if r["employeeCode"] == "30020")
        self.assertEqual((row["present"], row["leaveDays"]), (2.0, 0.0))

    def test_days_nobody_has_opened_are_counted_and_said_not_invented(self):
        body = self.get(**RANGE)
        # R2 5 + R3 4 + R4 6 + R6 4 (from the 23rd) + R7 6; R1 has every day
        self.assertEqual(cards(body)["Days with no record"], 25)
        self.assertTrue(any("25 past employee-day(s) have no stored attendance record" in n for n in body["notes"]))

    def test_todays_absence_is_provisional_not_leave(self):
        rec(self.r2, TODAY, "absent")
        body = self.get(dateFrom="2026-09-29", dateTo="2026-09-29")
        row = next(r for r in data_rows(body) if r["employeeCode"] == "30021")
        self.assertEqual(row["leaveDays"], 0.0)
        self.assertTrue(any("absence(s) today are not counted as leave" in n for n in body["notes"]))

    def test_a_leaver_with_nothing_in_the_dates_is_not_listed_but_one_with_activity_is(self):
        self.assertNotIn("30300", [r["employeeCode"] for r in data_rows(self.get(**RANGE))])
        rec(self.r5, date(2026, 9, 22), "present", "1")
        self.assertIn("30300", [r["employeeCode"] for r in data_rows(self.get(**RANGE))])
        self.assertNotIn("30300", [r["employeeCode"] for r in data_rows(self.get(employeeStatus="active", **RANGE))])

    def test_tickets_are_in_natural_order_inside_a_department(self):
        mk = Employee.objects.create
        for code in ("30009", "300100"):
            mk(
                employee_code=code,
                first_name="N",
                last_name="X",
                department=self.acc,
                branch=self.b1,
                join_date="2024-01-01",
            )
        codes = [r["employeeCode"] for r in data_rows(self.get(**RANGE)) if r["department"] == "ACCOUNTS"]
        self.assertEqual(codes, ["30009", "30020", "30021", "300100"])

    def test_the_summary_cards_and_notes(self):
        body = self.get(**RANGE)
        self.assertEqual(
            cards(body),
            {
                "Employees": 6,
                "Present (shifts)": 8.5,
                "OT (shifts)": 1.5,
                "Total (shifts)": 10.0,
                "Leave days": 5.0,
                "Days with no record": 25,
            },
        )
        text = " ".join(body["notes"])
        for part in (
            "21-Sep-2026 to 27-Sep-2026",
            "Total = Present + OT",
            "Weekly offs",
            "company holidays are not leave",
            "announced as Pay",
        ):
            self.assertIn(part, text)

    def test_nobody_gives_an_empty_register_not_an_error(self):
        Employee.objects.all().delete()
        body = self.get(**RANGE)
        self.assertEqual(body["rows"], [])
        self.assertEqual(cards(body)["Employees"], 0)
        self.assertIsNone(body["totals"])


class RegisterFilterAndIsolationTests(_RegisterWorld):
    def test_scope_filters_narrow_the_register(self):
        codes = lambda **p: [r["employeeCode"] for r in data_rows(self.get(**RANGE, **p))]  # noqa: E731
        self.assertEqual(codes(departmentIds=str(self.adm.id)), ["30100", "30400"])
        self.assertEqual(codes(employmentType="production"), ["30100"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), ["30500"])
        self.assertEqual(codes(employeeIds=f"{self.r1.id},{self.r2.id}"), ["30020", "30021"])
        self.assertEqual(codes(designationIds=str(self.des.id))[0], "30020")

    def test_a_branch_user_sees_only_their_own_branch_and_cannot_widen_it(self):
        body = self.get(user=self.branch_user, **RANGE)
        self.assertNotIn("30500", str(body["rows"]))
        self.assertNotIn("CUTTING", str(body["rows"]))
        self.assertEqual(body["rows"][-1]["present"], 7.5)  # the grand total leaves the other unit out
        self.assertEqual(data_rows(self.get(user=self.branch_user, branchIds=str(self.b2.id), **RANGE)), [])
        self.assertEqual(data_rows(self.get(user=self.branch_user, employeeIds=str(self.r7.id), **RANGE)), [])

    def test_the_other_branchs_overtime_does_not_leak_into_a_branch_users_totals(self):
        ot(self.r7, date(2026, 9, 22), "1")
        body = self.get(user=self.branch_user, **RANGE)
        self.assertEqual(body["rows"][-1]["ot"], 1.5)

    def test_a_branch_user_gets_a_scoped_excel(self):
        ws = load_workbook(io.BytesIO(self.export("xlsx", user=self.branch_user, **RANGE).content)).active
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertNotIn("OTHERUNIT", cells)
        self.assertIn("SURYA", cells)


class RegisterExportAndSafetyTests(_RegisterWorld):
    def test_excel_carries_the_same_columns_and_rows_as_the_screen(self):
        self.assert_excel_equals_screen(**RANGE)
        self.assert_excel_equals_screen(departmentIds=str(self.acc.id), **RANGE)
        ws = load_workbook(io.BytesIO(self.export("xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[HEADER_ROW]], REGISTER_LABELS)
        names = [r[3] for r in self.excel_rows(**RANGE)]
        self.assertEqual(names.count("SUB TOTAL"), 4)
        self.assertEqual(names[-1], "GRAND TOTAL")

    def test_the_pdf_is_a_real_pdf_with_the_printed_headings_and_totals(self):
        text = self.pdf_text(**RANGE)
        for part in (
            "Attendance Register (Summary)",
            "TicketNo",
            "EmpName",
            "Leave Days",
            "SUB TOTAL",
            "GRAND TOTAL",
            "30020",
        ):
            self.assertIn(part, text)

    def test_running_and_exporting_never_writes(self):
        self.assert_never_writes(**RANGE)

    def test_the_query_count_does_not_grow_with_the_number_of_employees(self):
        def queries() -> int:
            self.get(**RANGE)  # warm-up (permission caches)
            with CaptureQueriesContext(connection) as q:
                self.get(**RANGE)
            return len(q)

        before = queries()
        for i in range(30):
            emp = Employee.objects.create(
                employee_code=f"X{i}", first_name=f"X{i}", last_name="Extra", branch=self.b1,
                department=(self.acc, self.adm, None)[i % 3], designation=self.des, join_date="2024-01-01",
            )  # fmt: skip
            for d in (21, 22, 23):
                rec(emp, date(2026, 9, d), ("present", "half_shift", "absent")[d % 3], "1")
            ot(emp, date(2026, 9, 22), "0.25")
        self.assertEqual(queries(), before)
        self.assertEqual(cards(self.get(**RANGE))["Employees"], 6 + 30)

    def test_photo_and_password_columns_are_never_selected(self):
        with CaptureQueriesContext(connection) as q:
            self.get(**RANGE)
            self.export("xlsx", **RANGE)
        for query in q.captured_queries:
            self.assertNotIn("photo_url", query["sql"])
            if '"hr_users"' not in query["sql"]:
                self.assertNotIn("password_hash", query["sql"])

    def test_a_row_over_the_limit_is_refused_for_export_not_silently_cut(self):
        with patch("api.reporting.runner.XLSX_ROW_LIMIT", 3):
            r = self.client.get(f"/api/reports/export/{RID}", {"fmt": "xlsx", **RANGE}, **hdr(self.admin))
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.json()["error"], "too_many_rows")
