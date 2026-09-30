"""
Every registered report, end to end, on a realistic seeded month (seed_data: 35 employees, attendance, payroll,
advances, leave and requests for June 2026).

Whatever a domain author forgot, this catches: a crash on real data or on empty data, a screen/Excel mismatch,
a report that writes on GET, a branch-scoped user seeing another branch, a view-only role being refused,
a permission module that does not exist, an N+1 query pattern.

Run via: python manage.py test api.tests_reporting_all -v 2
"""

import io
from django.apps import apps
from django.core.management import call_command
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .models import Branch, Department, Employee, HRUser, PayrollSettings, Role
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.export_xlsx import HEADER_ROW

JUNE = {"period": "2026-06", "dateFrom": "2026-06-01", "dateTo": "2026-06-30", "year": "2026"}
# A screen run may not issue more queries than this on the 35-employee seed (guards N+1 patterns).
QUERY_BUDGET = 150


def _headers(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def _default_query(spec_json: dict, overrides: dict | None = None) -> dict:
    """The parameters the UI would send for a report's default filters (with optional overrides)."""
    q: dict = {}
    for f in spec_json["filters"]:
        kind, d = f["kind"], f.get("default")
        if kind == "period" and d:
            q["period"] = d
        elif kind == "year" and d:
            q["year"] = d
        elif kind == "dateRange" and d:
            q["dateFrom"], q["dateTo"] = d["dateFrom"], d["dateTo"]
        elif kind == "employeeStatus" and d:
            q["employeeStatus"] = d
        elif kind == "select" and d:
            q[f["key"]] = d
        elif kind == "boolean" and d is True:
            q[f["key"]] = "true"
        elif kind == "number" and d is not None:
            q[f["key"]] = d
    if overrides:
        for key, value in overrides.items():
            if key in ("period", "year") and not any(f["kind"] == key for f in spec_json["filters"]):
                continue
            if key in ("dateFrom", "dateTo") and not any(f["kind"] == "dateRange" for f in spec_json["filters"]):
                continue
            q[key] = value
    return q


def _row_counts() -> dict[str, int]:
    """Row count of every table of the app -- the audit trail excluded (exports are supposed to add rows there)."""
    out = {}
    for model in apps.get_app_config("api").get_models():
        if model.__name__ == "AuditLog":
            continue
        out[model.__name__] = model.objects.count()
    return out


class AllReportsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command(
            "seed_data", confirm=True, verbosity=0, stdout=io.StringIO(), stderr=io.StringIO()
        )  # its check mark cannot be printed to a cp1252 console
        PayrollSettings.get()  # the app's singleton settings row exists on every installed system (created at first use)
        cls.seed_branch = Branch.objects.order_by("id").first()
        cls.other_branch = Branch.objects.create(name="Other Unit", code="OTH")
        other_dept = Department.objects.create(name="OTHER-DEPT", branch=cls.other_branch)
        cls.stranger = Employee.objects.create(
            employee_code="OTHER001",
            first_name="Stranger",
            last_name="Elsewhere",
            employment_type="staff",
            department=other_dept,
            branch=cls.other_branch,
            status="active",
            join_date="2026-01-05",
        )
        cls.admin = HRUser.objects.create(username="all_admin", password_hash="x", is_super_admin=True)
        every_module = {key: "view" for key in all_module_keys()}
        cls.scoped = HRUser.objects.create(
            username="all_scoped",
            password_hash="x",
            branch=cls.seed_branch,
            role=Role.objects.create(name="all_scoped_role", permissions=every_module),
        )
        cls.viewer = HRUser.objects.create(
            username="all_viewer",
            password_hash="x",
            role=Role.objects.create(name="all_viewer_role", permissions=every_module),
        )
        cls.reports_only = HRUser.objects.create(
            username="all_reports_only",
            password_hash="x",
            role=Role.objects.create(name="all_reports_only_role", permissions={"reports": "view"}),
        )

    def catalog(self, user):
        r = self.client.get("/api/reports/catalog", **_headers(user))
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    def run_report(self, user, spec, params):
        return self.client.get(f"/api/reports/run/{spec['id']}", params, **_headers(user))

    # ── catalog ──────────────────────────────────────────────────────────────
    def test_the_catalog_is_large_and_every_definition_loaded(self):
        self.assertEqual(registry.LOAD_ERRORS, {})
        reports = self.catalog(self.admin)["reports"]
        self.assertGreaterEqual(len(reports), 60, "expected the full Report Center catalog")
        cats = {c["id"] for c in self.catalog(self.admin)["categories"]}
        self.assertTrue({"payroll", "attendance", "leave", "gate", "employees"} <= cats, cats)

    # ── every report runs, on real data and on empty data ──────────────────
    def test_every_report_runs_on_the_seeded_month(self):
        for spec in self.catalog(self.admin)["reports"]:
            with self.subTest(report=spec["id"]):
                with CaptureQueriesContext(connection) as ctx:
                    r = self.run_report(self.admin, spec, _default_query(spec, JUNE))
                self.assertEqual(r.status_code, 200, r.content[:400])
                body = r.json()
                self.assertEqual(body["id"], spec["id"])
                self.assertTrue(body["columns"], "no columns")
                self.assertEqual(body["rowCount"], len(body["rows"]))
                keys = {c["key"] for c in body["columns"]}
                for row in body["rows"][:50]:
                    self.assertTrue(set(row) - {"_kind"} <= keys)
                self.assertLessEqual(len(ctx), QUERY_BUDGET, f"{spec['id']} ran {len(ctx)} queries (N+1?)")

    def test_every_report_runs_with_default_filters_on_current_dates(self):
        # "today" is nowhere near the seeded month: mostly empty results - must not crash or divide by zero.
        for spec in self.catalog(self.admin)["reports"]:
            with self.subTest(report=spec["id"]):
                r = self.run_report(self.admin, spec, _default_query(spec))
                self.assertEqual(r.status_code, 200, r.content[:400])

    def test_reports_without_any_employees_do_not_crash(self):
        Employee.objects.all().delete()
        for spec in self.catalog(self.admin)["reports"]:
            with self.subTest(report=spec["id"]):
                r = self.run_report(self.admin, spec, _default_query(spec, JUNE))
                self.assertEqual(r.status_code, 200, r.content[:400])

    # ── exports agree with the screen ──────────────────────────────────────
    def test_every_report_exports_to_excel_and_pdf_and_matches_the_screen(self):
        for spec in self.catalog(self.admin)["reports"]:
            with self.subTest(report=spec["id"]):
                params = _default_query(spec, JUNE)
                screen = self.run_report(self.admin, spec, params).json()
                x = self.client.get(
                    f"/api/reports/export/{spec['id']}", {**params, "fmt": "xlsx"}, **_headers(self.admin)
                )
                self.assertEqual(x.status_code, 200, x.content[:300])
                self.assertTrue(x.content.startswith(b"PK"))
                ws = load_workbook(io.BytesIO(x.content)).active
                header = [c.value for c in ws[HEADER_ROW]]
                self.assertEqual(header[: len(screen["columns"])], [c["label"] for c in screen["columns"]])
                # the sheet holds exactly the screen's rows, in order (compared on the first column when it is text)
                first = screen["columns"][0]
                if first["type"] == "text":
                    got = [ws.cell(row=HEADER_ROW + 1 + i, column=1).value for i in range(len(screen["rows"]))]
                    want = [r.get(first["key"]) for r in screen["rows"]]
                    self.assertEqual([str(g or "") for g in got], [str(w or "") for w in want])
                p = self.client.get(
                    f"/api/reports/export/{spec['id']}", {**params, "fmt": "pdf"}, **_headers(self.admin)
                )
                self.assertEqual(p.status_code, 200, p.content[:300])
                self.assertTrue(p.content.startswith(b"%PDF"))

    # ── reading a report never changes data ────────────────────────────────
    def test_running_and_exporting_reports_never_writes_to_the_database(self):
        before = _row_counts()
        for spec in self.catalog(self.admin)["reports"]:
            params = _default_query(spec, JUNE)
            self.run_report(self.admin, spec, params)
            self.client.get(f"/api/reports/export/{spec['id']}", {**params, "fmt": "xlsx"}, **_headers(self.admin))
        after = _row_counts()
        changed = {k: (before[k], after[k]) for k in before if before[k] != after[k]}
        self.assertEqual(changed, {}, "a report wrote to the database (table: (before, after))")

    # ── hostile / malformed input never crashes a report ───────────────────
    def test_malformed_filters_are_a_400_never_a_500(self):
        nasty = [
            {"dateFrom": "9999-12-31", "dateTo": "9999-12-31"},
            {"dateFrom": "0001-01-01", "dateTo": "0001-01-02"},
            {"period": "9999-12"},
            {"period": "2026-13"},
            {"year": "99999999999999999999"},
            {"departmentIds": "99999999999999999999"},
            {"employeeIds": "1,,2,x"},
            {"employeeIds": ",".join(str(i) for i in range(1, 1500))},
            {"branchIds": "-1"},
            {"employmentType": "x" * 500},
            {"employeeStatus": "\u0000"},
        ]
        for spec in self.catalog(self.admin)["reports"]:
            for bad in nasty:
                with self.subTest(report=spec["id"], params=str(bad)[:60]):
                    params = _default_query(spec, JUNE)
                    params.update(bad)
                    r = self.run_report(self.admin, spec, params)
                    self.assertIn(r.status_code, (200, 400), r.content[:300])

    def test_free_text_filters_survive_nul_bytes_and_unicode(self):
        for spec in self.catalog(self.admin)["reports"]:
            text_filters = [f["key"] for f in spec["filters"] if f["kind"] == "text"]
            for key in text_filters:
                for value in ("a\u0000b", "\u0000", "நிர்வாகி", "x" * 5000, "%_\\"):
                    with self.subTest(report=spec["id"], filter=key):
                        params = _default_query(spec, JUNE)
                        params[key] = value
                        r = self.run_report(self.admin, spec, params)
                        self.assertIn(r.status_code, (200, 400), r.content[:300])

    # ── access ───────────────────────────────────────────────────────────────
    def test_a_view_only_role_can_run_and_export_everything_it_may_see(self):
        reports = self.catalog(self.viewer)["reports"]
        self.assertGreater(len(reports), 30)
        for spec in reports:
            with self.subTest(report=spec["id"]):
                params = _default_query(spec, JUNE)
                self.assertEqual(self.run_report(self.viewer, spec, params).status_code, 200)
                x = self.client.get(
                    f"/api/reports/export/{spec['id']}", {**params, "fmt": "xlsx"}, **_headers(self.viewer)
                )
                self.assertEqual(x.status_code, 200)

    def test_admin_only_reports_are_hidden_from_every_other_role(self):
        admin_ids = {s["id"] for s in self.catalog(self.admin)["reports"]}
        viewer_ids = {s["id"] for s in self.catalog(self.viewer)["reports"]}
        specs = {s.id: s for s in registry.all_specs()}
        self.assertTrue(any(s.super_admin_only for s in specs.values()), "no admin-only report registered")
        for rid in admin_ids - viewer_ids:
            self.assertTrue(
                specs[rid].super_admin_only, f"{rid} is hidden from an all-access role but is not super_admin_only"
            )
        for spec in specs.values():
            if spec.super_admin_only:
                self.assertIn(spec.id, admin_ids)
                self.assertNotIn(spec.id, viewer_ids)
                r = self.client.get(f"/api/reports/run/{spec.id}", **_headers(self.viewer))
                self.assertEqual(r.status_code, 404)

    def test_a_role_with_only_the_reports_module_sees_only_reports_that_need_no_other_module(self):
        ids = {s["id"] for s in self.catalog(self.reports_only)["reports"]}
        gated = [s for s in registry.all_specs() if s.modules or s.super_admin_only]
        self.assertTrue(gated, "expected gated reports")
        for spec in gated:
            self.assertNotIn(spec.id, ids, f"{spec.id} leaks to a role without its owning module")

    def test_sensitive_reports_are_gated_by_an_owning_module(self):
        specs = {s.id: s for s in registry.all_specs()}
        for rid in (
            "salary-register",
            "salary-slip",
            "pf-statement",
            "esi-statement",
            "bank-advice",
            "advance-ledger",
            "visitor-register",
        ):
            if rid in specs:
                self.assertTrue(
                    specs[rid].modules or specs[rid].super_admin_only, f"{rid} must name its owning module(s)"
                )

    # ── branch isolation ─────────────────────────────────────────────────────
    def test_a_branch_scoped_user_never_sees_another_branchs_people(self):
        # the stranger belongs to another branch: their code / name must not appear anywhere a scoped user looks
        for spec in self.catalog(self.scoped)["reports"]:
            with self.subTest(report=spec["id"]):
                params = _default_query(spec, JUNE)
                r = self.run_report(self.scoped, spec, params)
                self.assertEqual(r.status_code, 200, r.content[:300])
                self.assertNotIn("OTHER001", r.content.decode(), "another branch's employee leaked")
                x = self.client.get(
                    f"/api/reports/export/{spec['id']}", {**params, "fmt": "xlsx"}, **_headers(self.scoped)
                )
                self.assertEqual(x.status_code, 200)
                ws = load_workbook(io.BytesIO(x.content)).active
                cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
                self.assertNotIn("OTHER001", cells)

    def test_the_stranger_is_visible_to_an_unscoped_admin_in_the_employee_master(self):
        specs = {s["id"]: s for s in self.catalog(self.admin)["reports"]}
        if "employee-master" not in specs:
            self.skipTest("employee-master not registered")
        body = self.run_report(
            self.admin, specs["employee-master"], _default_query(specs["employee-master"], JUNE)
        ).json()
        self.assertIn("OTHER001", str(body["rows"]))

    # ── contract on the specs themselves ─────────────────────────────────────
    def test_spec_contract(self):
        keys = set(all_module_keys())
        seen_titles: dict[str, str] = {}
        for spec in registry.all_specs():
            self.assertTrue(set(spec.modules) <= keys, f"{spec.id}: unknown module in {spec.modules}")
            self.assertTrue(spec.description.strip(), spec.id)
            self.assertLess(len(spec.description), 220, f"{spec.id}: keep the catalog card description short")
            key = spec.title.strip().lower()
            self.assertNotIn(key, seen_titles, f"{spec.id} has the same title as {seen_titles.get(key)}")
            seen_titles[key] = spec.id
