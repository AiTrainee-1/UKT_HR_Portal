"""Designations and departments named in text are created when they do not exist (bulk Add new / Update, and the single
employee API), and the bulk upload stays light and bounded: a few queries per row, a sheet read as a stream, a cap on
rows, size and time, and nothing half-done when a file is refused (employee_bulk_views.py, org_lookup.py)."""

import io
import json
import re
import time
from unittest import mock

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import SimpleTestCase, TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl.styles import Font

from . import employee_bulk_views
from .employee_bulk_views import _collect_rows
from .jwt_utils import sign_token
from .models import AuditLog, Branch, Department, Designation, Employee, HRUser, Role
from .org_lookup import OrgLookup, clean_name, is_blank_name
from .tests_employee_bulk import (
    PRODUCTION_EXPORT,
    PRODUCTION_UPLOAD_HEADERS,
    STAFF_EXPORT,
    STAFF_UPLOAD_HEADERS,
    Base,
    _row,
    _sheet,
)


class Fixtures(Base):
    def setUp(self):
        super().setUp()
        self.unit1 = Branch.objects.create(name="Unit1", code="U1")
        self.dept_u1 = Department.objects.create(name="Stitching", branch=self.unit1)

    def edit(self, emp, **cells):
        """An update-sheet row for `emp`: its code and only the cells given (blank cells leave everything as is)."""
        return _row(STAFF_EXPORT, **{"Employee Code": emp.employee_code, **cells})

    def update_staff(self, rows, **fields):
        return self.update(STAFF_EXPORT, rows, category="staff", employeeStatus="active", **fields)

    def designation_of(self, emp):
        emp.refresh_from_db()
        return Designation.objects.get(pk=emp.designation_id) if emp.designation_id else None

    def scoped_headers(self, branch):
        role = Role.objects.create(name=f"HR {branch.name}", permissions={"employees": "edit"})
        user = HRUser.objects.create(username=f"hr_{branch.name}", password_hash="x", role=role, branch=branch)
        return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


class UpdateCreatesDesignations(Fixtures):
    def test_a_designation_that_does_not_exist_is_created_in_the_employees_department_and_assigned(self):
        emp = self.person("S1")
        r = self.update_staff([self.edit(emp, **{"Designation": "Quality Checker"})])
        self.assertEqual(r.status_code, 200, r.content)
        made = self.designation_of(emp)
        self.assertEqual(made.title, "Quality Checker")
        self.assertEqual(made.department_id, self.dept.id)  # in the department, so the Edit form and the branch see it
        row = r.json()["rows"][0]
        self.assertEqual((row["status"], row["changes"]), ("updated", ["Designation"]))
        self.assertEqual(row["notes"], ["Created designation 'Quality Checker' in Stitching"])
        self.assertEqual(
            r.json()["newDesignations"], [{"title": "Quality Checker", "department": "Stitching", "rows": 1}]
        )
        self.assertEqual(r.json()["warnings"], [])  # it used to warn "not found" and keep the old one

    def test_an_existing_designation_is_reused_ignoring_case_and_stray_spaces(self):
        existing = Designation.objects.create(title="Sr. Tailor", department=self.dept)
        emp = self.person("S1")
        r = self.update_staff([self.edit(emp, **{"Designation": "  sr.   TAILOR "})])
        self.assertEqual(self.designation_of(emp).id, existing.id)
        self.assertEqual(Designation.objects.count(), 1)
        self.assertEqual(r.json()["newDesignations"], [])
        self.assertEqual(r.json()["rows"][0]["notes"], [])

    def test_the_same_new_title_for_many_rows_makes_one_designation(self):
        a, b, c = self.person("S1"), self.person("S2"), self.person("S3")
        r = self.update_staff([self.edit(e, **{"Designation": "Line Leader"}) for e in (a, b, c)])
        self.assertEqual(Designation.objects.filter(title="Line Leader").count(), 1)
        for e in (a, b, c):
            self.assertEqual(self.designation_of(e).title, "Line Leader")
        self.assertEqual(r.json()["newDesignations"], [{"title": "Line Leader", "department": "Stitching", "rows": 3}])
        # only the row that caused it says "Created"; the others simply got it
        self.assertEqual(len([x for x in r.json()["rows"] if x["notes"]]), 1)

    def test_a_title_that_exists_only_under_another_department_is_created_for_this_one(self):
        other = Department.objects.create(name="Cutting", branch=self.branch)
        theirs = Designation.objects.create(title="Helper", department=other)
        emp = self.person("S1")
        self.update_staff([self.edit(emp, **{"Designation": "Helper"})])
        mine = self.designation_of(emp)
        self.assertNotEqual(mine.id, theirs.id)
        self.assertEqual(mine.department_id, self.dept.id)
        theirs.refresh_from_db()
        self.assertEqual(theirs.department_id, other.id)

    def test_a_row_that_changes_the_department_gets_its_designation_in_the_new_one(self):
        emp = self.person("S1")
        r = self.update_staff([self.edit(emp, **{"Department": "Packing", "Designation": "Packer"})])
        emp.refresh_from_db()
        self.assertEqual(emp.department.name, "Packing")
        self.assertEqual(self.designation_of(emp).department_id, emp.department_id)
        self.assertEqual(r.json()["rows"][0]["changes"], ["Department", "Designation"])
        self.assertEqual(
            r.json()["rows"][0]["notes"], ["Created department 'Packing'", "Created designation 'Packer' in Packing"]
        )

    def test_inactive_employees_get_new_designations_too(self):
        emp = self.person("S1", status="inactive")
        r = self.update(
            STAFF_EXPORT,
            [self.edit(emp, **{"Designation": "Ex Supervisor"})],
            category="staff",
            employeeStatus="inactive",
        )
        self.assertEqual(r.json()["counts"]["updated"], 1, r.content)
        self.assertEqual(self.designation_of(emp).title, "Ex Supervisor")
        emp.refresh_from_db()
        self.assertEqual(emp.status, "inactive")

    def test_production_employees_get_new_designations_too(self):
        emp = self.person("P1", kind="production")
        r = self.update(
            PRODUCTION_EXPORT,
            [_row(PRODUCTION_EXPORT, **{"Employee Code": "P1", "Designation": "Machine Operator"})],
            category="production",
            employeeStatus="active",
        )
        self.assertEqual(r.json()["counts"]["updated"], 1, r.content)
        self.assertEqual(self.designation_of(emp).title, "Machine Operator")

    def test_nothing_is_created_by_a_check_but_it_says_what_would_be(self):
        emp = self.person("S1")
        r = self.update_staff([self.edit(emp, **{"Department": "Packing", "Designation": "Packer"})], mode="preview")
        body = r.json()
        self.assertEqual(body["rows"][0]["status"], "updated")
        self.assertEqual(body["newDesignations"][0]["title"], "Packer")
        self.assertEqual(body["newDepartments"][0]["name"], "Packing")
        self.assertEqual(Designation.objects.count(), 0)
        self.assertFalse(Department.objects.filter(name="Packing").exists())
        emp.refresh_from_db()
        self.assertEqual((emp.department_id, emp.designation_id), (self.dept.id, None))

    def test_placeholders_in_the_cell_are_not_titles(self):
        emp = self.person("S1")
        for placeholder in (
            "-",
            "N/A",
            "nil",
            "None",
        ):  # one upload each: the same person twice in a sheet is a duplicate
            r = self.update_staff([self.edit(emp, **{"Designation": placeholder})])
            self.assertEqual(r.json()["rows"][0]["status"], "unchanged", placeholder)
        self.assertEqual(Designation.objects.count(), 0)

    def test_a_title_that_is_clearly_not_one_is_refused_for_that_cell_only(self):
        emp = self.person("S1")
        r = self.update_staff([self.edit(emp, **{"Designation": "x" * 150, "Last Name": "Renamed"})])
        row = r.json()["rows"][0]
        self.assertEqual((row["status"], row["changes"]), ("updated", ["Last Name"]))
        self.assertIn("too long", row["warnings"][0])
        self.assertEqual(Designation.objects.count(), 0)

    def test_a_blank_cell_never_clears_or_creates_anything(self):
        d = Designation.objects.create(title="Keeper", department=self.dept)
        emp = self.person("S1", designation=d)
        r = self.update_staff([self.edit(emp)])
        self.assertEqual(r.json()["rows"][0]["status"], "unchanged")
        self.assertEqual(self.designation_of(emp).id, d.id)

    def test_a_department_is_matched_in_the_employees_own_branch_first(self):
        Department.objects.create(name="Cutting", branch=self.branch)  # the older one, in another branch
        mine = Department.objects.create(name="Cutting", branch=self.unit1)
        emp = self.person("S1", branch=self.unit1, department=self.dept_u1)
        self.update_staff([self.edit(emp, **{"Department": "cutting"})])
        emp.refresh_from_db()
        self.assertEqual(emp.department_id, mine.id)
        self.assertEqual(Department.objects.filter(name__iexact="cutting").count(), 2)

    def test_a_new_department_belongs_to_the_employees_branch_never_to_nobody(self):
        emp = self.person("S1", branch=self.unit1, department=self.dept_u1)
        self.update_staff([self.edit(emp, **{"Department": "Embroidery"})])
        made = Department.objects.get(name="Embroidery")
        self.assertEqual(made.branch_id, self.unit1.id)

    def test_a_branch_change_makes_the_new_department_in_the_new_branch(self):
        emp = self.person("S1")
        self.update_staff([self.edit(emp, **{"Branch": "Unit1", "Department": "Finishing"})])
        emp.refresh_from_db()
        self.assertEqual(emp.branch_id, self.unit1.id)
        self.assertEqual(emp.department.branch_id, self.unit1.id)
        self.assertEqual(emp.unit_code, "U1-1")

    def test_a_failed_row_does_not_stop_the_others_and_the_designation_it_made_stays_usable(self):
        bad, good = self.person("S1"), self.person("S2")
        r = self.update_staff(
            [
                self.edit(bad, **{"Designation": "Planner", "Salary Amount": 10**12}),  # too big for the column
                self.edit(good, **{"Designation": "Planner"}),
            ]
        )
        by = self.by_row(r.json())
        self.assertEqual(by[2]["status"], "failed")
        self.assertEqual(by[3]["status"], "updated")
        self.assertEqual(self.designation_of(good).title, "Planner")
        self.assertIsNone(self.designation_of(bad))
        self.assertEqual(Designation.objects.filter(title="Planner").count(), 1)


class AddNewCreatesDesignations(Fixtures):
    def test_new_employees_get_new_designations_and_departments(self):
        r = self.upload(
            STAFF_UPLOAD_HEADERS,
            [self.staff_row("N1", Department="Quality", Designation="Inspector")],
            category="staff",
        )
        self.assertEqual(r.status_code, 201, r.content)
        emp = Employee.objects.get(employee_code="N1")
        self.assertEqual((emp.department.name, emp.designation.title), ("Quality", "Inspector"))
        self.assertEqual(emp.department.branch_id, emp.branch_id)  # never a department of no branch
        self.assertEqual(emp.designation.department_id, emp.department_id)
        body = r.json()
        self.assertEqual(
            body["rows"][0]["notes"], ["Created department 'Quality'", "Created designation 'Inspector' in Quality"]
        )
        self.assertEqual(body["warnings"], [])  # it used to say "Designation 'Inspector' not found - left blank"

    def test_a_second_row_with_the_same_title_reuses_it(self):
        r = self.upload(
            STAFF_UPLOAD_HEADERS,
            [self.staff_row(f"N{i}", Department="Stitching", Designation="Tailor") for i in range(3)],
            category="staff",
        )
        self.assertEqual(r.json()["counts"]["created"], 3)
        self.assertEqual(Designation.objects.filter(title="Tailor").count(), 1)
        self.assertEqual(r.json()["newDesignations"], [{"title": "Tailor", "department": "Stitching", "rows": 3}])

    def test_production_sheets_too(self):
        r = self.upload(
            PRODUCTION_UPLOAD_HEADERS,
            [self.production_row("N1", Department="Stitching", Designation="Operator")],
            category="production",
        )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Employee.objects.get(employee_code="N1").designation.title, "Operator")

    def test_a_check_creates_nothing(self):
        r = self.upload(
            STAFF_UPLOAD_HEADERS,
            [self.staff_row("N1", Department="Quality", Designation="Inspector")],
            category="staff",
            mode="preview",
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["newDesignations"][0]["title"], "Inspector")
        self.assertFalse(Employee.objects.filter(employee_code="N1").exists())
        self.assertEqual(Designation.objects.count(), 0)
        self.assertFalse(Department.objects.filter(name="Quality").exists())

    def test_a_branch_login_creates_them_in_its_own_branch_and_sees_them(self):
        hr = self.scoped_headers(self.unit1)
        sheet = [self.staff_row("N1", Branch="Head Office", Department="Stitching", Designation="Loom Fitter")]
        r = self.client.post(
            "/api/employees/bulk-upload", {"file": _sheet(STAFF_UPLOAD_HEADERS, sheet), "category": "staff"}, **hr
        )
        self.assertEqual(r.status_code, 201, r.content)
        emp = Employee.objects.get(employee_code="N1")
        self.assertEqual(emp.branch_id, self.unit1.id)  # the login's branch wins over the sheet
        self.assertEqual(emp.department_id, self.dept_u1.id)
        self.assertEqual(emp.designation.department_id, self.dept_u1.id)
        listing = self.client.get("/api/designations", **hr).json()
        self.assertEqual([d["title"] for d in listing], ["Loom Fitter"])

    def test_a_row_the_database_refuses_fails_alone_and_unit_codes_stay_unique(self):
        r = self.upload(
            STAFF_UPLOAD_HEADERS,
            [
                self.staff_row("N1", Branch="Unit1", Department="Stitching"),
                self.staff_row("N2", Branch="Unit1", Department="Stitching", **{"Salary Amount": 10**12}),
                self.staff_row("N3", Branch="Unit1", Department="Stitching"),
            ],
            category="staff",
        )
        by = self.by_row(r.json())
        self.assertEqual([by[2]["status"], by[3]["status"], by[4]["status"]], ["created", "failed", "created"])
        codes = list(Employee.objects.filter(unit_code__isnull=False).values_list("unit_code", flat=True))
        self.assertEqual(len(codes), len(set(codes)))
        self.unit1.refresh_from_db()
        self.assertGreaterEqual(self.unit1.next_employee_seq, 2)


class SingleEmployeeApi(Fixtures):
    def patch(self, emp, body):
        return self.client.patch(
            f"/api/employees/{emp.id}", data=json.dumps(body), content_type="application/json", **self.hr
        )

    def test_a_designation_sent_as_a_name_is_created_and_assigned(self):
        emp = self.person("S1")
        r = self.patch(emp, {"designation": "Floor Manager"})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(self.designation_of(emp).title, "Floor Manager")
        self.assertEqual(self.designation_of(emp).department_id, self.dept.id)

    def test_a_designation_sent_as_an_id_still_works(self):
        d = Designation.objects.create(title="Existing", department=self.dept)
        emp = self.person("S1")
        self.assertEqual(self.patch(emp, {"designationId": d.id}).status_code, 200)
        self.assertEqual(self.designation_of(emp).id, d.id)

    def test_a_department_name_that_two_branches_share_no_longer_crashes_the_request(self):
        Department.objects.create(name="Cutting", branch=self.branch)
        mine = Department.objects.create(name="Cutting", branch=self.unit1)
        emp = self.person("S1", branch=self.unit1, department=self.dept_u1)
        r = self.patch(emp, {"department": "Cutting"})  # used to raise MultipleObjectsReturned -> HTTP 500
        self.assertEqual(r.status_code, 200, r.content)
        emp.refresh_from_db()
        self.assertEqual(emp.department_id, mine.id)

    def test_a_null_department_name_is_ignored_not_a_crash(self):
        emp = self.person("S1")
        self.assertEqual(self.patch(emp, {"department": None, "designation": None}).status_code, 200)

    def test_adding_an_employee_with_a_new_designation_name(self):
        r = self.client.post(
            "/api/employees",
            data=json.dumps(
                {
                    "employeeCode": "A1",
                    "firstName": "Asha",
                    "lastName": "K",
                    "phone": "9000000001",
                    "branchId": self.branch.id,
                    "department": "Stitching",
                    "designation": "Shift Incharge",
                }
            ),
            content_type="application/json",
            **self.hr,
        )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Employee.objects.get(employee_code="A1").designation.title, "Shift Incharge")


class WorkStaysBounded(Fixtures):
    def queries(self, call):
        with CaptureQueriesContext(connection) as ctx:
            response = call()
        self.assertIn(response.status_code, (200, 201), response.content)
        return len(ctx)

    def test_unchanged_rows_cost_no_queries_at_all(self):
        d = Designation.objects.create(title="Operator", department=self.dept)

        def run(n):
            Employee.objects.all().delete()
            emps = [self.person(f"U{i}", designation=d) for i in range(n)]
            rows = [self.edit(e, **{"Designation": "operator", "Department": "stitching"}) for e in emps]
            return self.queries(lambda: self.update_staff(rows))

        small, large = run(5), run(60)
        self.assertEqual(small, large)

    def test_changed_rows_cost_a_savepoint_and_a_write_not_a_query_per_field(self):
        def run(n):
            Employee.objects.all().delete()
            emps = [self.person(f"C{i}") for i in range(n)]
            rows = [
                self.edit(e, **{"Last Name": "Changed", "Designation": "Operator", "Salary Amount": 31000})
                for e in emps
            ]
            return self.queries(lambda: self.update_staff(rows))

        small, large = run(5), run(45)
        self.assertLessEqual((large - small) / 40, 4)

    def test_new_employees_cost_a_few_queries_each(self):
        def run(n, tag):
            rows = [self.staff_row(f"{tag}{i}", Department="Stitching", Designation="Tailor") for i in range(n)]
            return self.queries(lambda: self.upload(STAFF_UPLOAD_HEADERS, rows, category="staff"))

        small, large = run(5, "A"), run(45, "B")
        self.assertLessEqual((large - small) / 40, 6)  # it was about fifteen

    def test_saving_a_row_writes_only_what_the_sheet_changed(self):
        emp = self.person("S1", password_hash="keep-me", photo_url="data:image/png;base64,AAAA")
        with CaptureQueriesContext(connection) as ctx:
            self.update_staff([self.edit(emp, **{"Last Name": "Renamed"})])
        writes = [q["sql"] for q in ctx if q["sql"].startswith('UPDATE "employees"')]
        self.assertEqual(len(writes), 1)
        assigned = set(re.findall(r'"(\w+)" =', writes[0].split(" SET ")[1].split(" WHERE ")[0]))
        # not the whole row: a login stamp or a password set while the file was being read must survive the save
        self.assertEqual(assigned, {"last_name", "updated_at"})
        reloaded = Employee.objects.get(pk=emp.pk)
        self.assertEqual((reloaded.password_hash, reloaded.photo_url), ("keep-me", "data:image/png;base64,AAAA"))

    def test_the_employee_photo_is_never_read_for_a_bulk_update(self):
        emp = self.person("S1", photo_url="data:image/png;base64," + "A" * 5000)
        with CaptureQueriesContext(connection) as ctx:
            self.update_staff([self.edit(emp, **{"Last Name": "Renamed"})])
        selects = [q["sql"] for q in ctx if q["sql"].startswith("SELECT") and '"employees"' in q["sql"]]
        self.assertTrue(selects)
        self.assertTrue(all("photo_url" not in sql for sql in selects))

    def test_formatting_left_on_hundreds_of_thousands_of_empty_rows_is_not_loaded(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(STAFF_UPLOAD_HEADERS)
        ws.append(self.staff_row("F1"))
        ws.cell(row=600000, column=1).font = Font(bold=True)  # what a formatted column leaves behind in Excel
        buf = io.BytesIO()
        wb.save(buf)
        started = time.monotonic()
        r = self.client.post(
            "/api/employees/bulk-upload",
            {"file": SimpleUploadedFile("big.xlsx", buf.getvalue()), "category": "staff"},
            **self.hr,
        )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["counts"]["created"], 1)
        self.assertLess(time.monotonic() - started, 10)

    def test_too_many_rows_are_refused_before_anything_is_done(self):
        rows = [self.staff_row(f"R{i}") for i in range(6)]
        with mock.patch.object(employee_bulk_views, "MAX_ROWS", 5):
            r = self.upload(STAFF_UPLOAD_HEADERS, rows, category="staff")
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.json()["error"], "too_large")
        self.assertIn("Split it", r.json()["message"])
        self.assertEqual(Employee.objects.count(), 0)

    def test_a_file_over_the_size_limit_is_refused(self):
        with mock.patch.object(employee_bulk_views, "MAX_UPLOAD_BYTES", 100):
            r = self.upload(STAFF_UPLOAD_HEADERS, [self.staff_row("R1")], category="staff")
        self.assertEqual(r.status_code, 400)
        self.assertIn("limit", r.json()["error"])
        self.assertEqual(Employee.objects.count(), 0)

    @staticmethod
    def clock(ticks_before_late):
        """A stand-in for `time` whose clock jumps past any deadline after `ticks_before_late` readings."""
        calls = {"n": 0}
        fake = mock.Mock()

        def monotonic():
            calls["n"] += 1
            return 0 if calls["n"] <= ticks_before_late else 10_000

        fake.monotonic = monotonic
        return fake

    def test_a_new_employee_upload_that_runs_out_of_time_changes_nothing(self):
        rows = [
            self.staff_row(f"T{i}", Branch="Unit1", Designation="Slow Role", Department="Stitching") for i in range(6)
        ]
        with mock.patch.object(employee_bulk_views, "time", self.clock(4)):  # the start, then three rows, then late
            r = self.upload(STAFF_UPLOAD_HEADERS, rows, category="staff")
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.json()["error"], "too_large")
        self.assertIn("Nothing was changed", r.json()["message"])
        self.assertEqual(Employee.objects.count(), 0)
        self.assertEqual(Designation.objects.count(), 0)
        self.unit1.refresh_from_db()
        self.assertEqual(self.unit1.next_employee_seq, 0)
        self.assertEqual(AuditLog.objects.filter(module="employees").count(), 0)

    def test_an_update_that_runs_out_of_time_changes_nothing_and_says_how_to_split(self):
        emps = [self.person(f"T{i}") for i in range(6)]
        rows = [self.edit(e, **{"Last Name": "Late", "Designation": "Slow Role"}) for e in emps]
        with mock.patch.object(employee_bulk_views, "time", self.clock(4)):
            r = self.update_staff(rows)
        self.assertEqual(r.status_code, 413)
        self.assertIn("leaving", r.json()["message"])  # "...on Keep for each part": splitting must not remove anyone
        self.assertFalse(Employee.objects.filter(last_name="Late").exists())
        self.assertEqual(Designation.objects.count(), 0)

    def test_a_check_that_runs_out_of_time_is_refused_the_same_way(self):
        emps = [self.person(f"T{i}") for i in range(6)]
        rows = [self.edit(e, **{"Last Name": "Late"}) for e in emps]
        with mock.patch.object(employee_bulk_views, "time", self.clock(3)):
            r = self.update_staff(rows, mode="preview")
        self.assertEqual(r.status_code, 413)


class AuditTrail(Fixtures):
    def test_one_entry_per_employee_in_the_same_words_as_before(self):
        a, b = self.person("S1"), self.person("S2")
        self.update_staff([self.edit(a, **{"Last Name": "Renamed"}), self.edit(b)])
        entries = list(AuditLog.objects.filter(module="employees", action="update"))
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0].record_id, a.id)
        self.assertEqual(entries[0].record_description, "Bulk-updated employee S1 -changed Last Name")

        self.upload(STAFF_UPLOAD_HEADERS, [self.staff_row("N1")], category="staff")
        made = AuditLog.objects.get(module="employees", action="create")
        self.assertEqual(made.record_description, "Bulk-imported employee N1 -Asha N1")
        self.assertEqual(made.record_id, Employee.objects.get(employee_code="N1").id)

    def test_a_check_writes_no_audit_entries(self):
        a = self.person("S1")
        self.update_staff([self.edit(a, **{"Last Name": "Renamed"})], mode="preview")
        self.upload(STAFF_UPLOAD_HEADERS, [self.staff_row("N1")], category="staff", mode="preview")
        self.assertEqual(AuditLog.objects.filter(module="employees").count(), 0)

    def test_making_someone_inactive_or_deleting_them_is_still_recorded(self):
        keep, gone, left = self.person("S1"), self.person("S2"), self.person("S3")
        self.update_staff(
            [self.edit(keep)], missingDecisions=json.dumps({"S2": "inactive", "S3": "delete"}), confirmDelete="true"
        )
        kinds = {(e.action, e.record_id) for e in AuditLog.objects.filter(module="employees")}
        self.assertEqual(kinds, {("update", gone.id), ("delete", left.id)})
        gone.refresh_from_db()
        self.assertEqual(gone.status, "inactive")
        self.assertFalse(Employee.objects.filter(pk=left.pk).exists())


class LiteDirectory(Fixtures):
    """The Bulk Upload page lists people without ever showing a photo, so it asks for the directory without them."""

    def test_the_lite_list_sends_a_link_instead_of_the_photo_and_is_otherwise_the_same(self):
        big = "data:image/jpeg;base64," + "A" * 60_000
        with_photo = self.person("P1", photo_url=big)
        self.person("P2", photo_url="https://example.com/p2.png")
        self.person("P3")
        full = self.client.get("/api/employees", **self.hr)
        lite = self.client.get("/api/employees?lite=1", **self.hr)
        self.assertEqual(lite.status_code, 200)
        by_code = {e["employeeCode"]: e for e in lite.json()}
        self.assertEqual(by_code["P1"]["photoUrl"], f"/api/employees/{with_photo.id}/photo")
        self.assertEqual(by_code["P2"]["photoUrl"], "https://example.com/p2.png")
        self.assertIsNone(by_code["P3"]["photoUrl"])
        self.assertLess(len(lite.content), 20_000)
        self.assertGreater(len(full.content), 60_000)  # the plain list is unchanged: the picture is inline
        plain = {e["employeeCode"]: e for e in full.json()}
        self.assertEqual(plain["P1"]["photoUrl"], big)
        for code in ("P1", "P2", "P3"):
            without_photo = lambda e: {k: v for k, v in e.items() if k != "photoUrl"}  # noqa: E731
            self.assertEqual(without_photo(plain[code]), without_photo(by_code[code]), code)

    def test_it_is_one_query_for_the_people_however_many_there_are(self):
        for i in range(8):
            self.person(f"Q{i}", photo_url="data:image/png;base64,AAAA")
        with CaptureQueriesContext(connection) as ctx:
            self.assertEqual(self.client.get("/api/employees?lite=1", **self.hr).status_code, 200)
        self.assertEqual(len([q for q in ctx if 'FROM "employees"' in q["sql"]]), 1)


class SheetReading(SimpleTestCase):
    def test_trailing_blank_rows_are_dropped_and_blank_rows_between_are_kept_in_place(self):
        stream = iter([("H",), ("a",), (None,), ("b",), (None, None), ()])
        rows = _collect_rows(stream)
        self.assertEqual(rows, [("H",), ("a",), (None,), ("b",)])  # row numbers in a report are Excel's

    def test_a_long_run_of_blank_rows_ends_the_data(self):
        def stream():
            yield ("H",)
            yield ("a",)
            for _ in range(5000):
                yield (None,)
            yield ("never reached",)

        with mock.patch.object(employee_bulk_views, "MAX_BLANK_RUN", 100):
            rows = _collect_rows(stream())
        self.assertEqual(rows, [("H",), ("a",)])

    def test_more_rows_than_the_limit_are_refused(self):
        with mock.patch.object(employee_bulk_views, "MAX_ROWS", 3):
            with self.assertRaises(employee_bulk_views._SheetTooBig):
                _collect_rows(iter([("H",)] + [(str(i),) for i in range(4)]))
            self.assertEqual(len(_collect_rows(iter([("H",)] + [(str(i),) for i in range(3)]))), 4)


class NameHelpers(SimpleTestCase):
    def test_clean_name_collapses_whitespace(self):
        self.assertEqual(clean_name("  Sr.   Tailor \n"), "Sr. Tailor")
        self.assertEqual(clean_name(None), "")
        self.assertEqual(clean_name(1234), "1234")

    def test_placeholders_are_blank(self):
        for v in (None, "", "  ", "-", "--", "—", "N/A", "n/a", "NA", "Nil", "none", "NULL", "0", "."):
            self.assertTrue(is_blank_name(v), repr(v))
        for v in ("Tailor", "Na Vin", "0 Defect Inspector", "A"):
            self.assertFalse(is_blank_name(v), repr(v))


class LookupIsLazy(TestCase):
    def test_an_empty_lookup_object_does_nothing_until_asked(self):
        with CaptureQueriesContext(connection) as ctx:
            OrgLookup()
        self.assertEqual(len(ctx), 0)
