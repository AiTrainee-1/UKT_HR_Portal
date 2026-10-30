"""Geo-Tracking: the HTTP surface (state, consent, journeys, points, export, rebuild, live, settings) and who may call what.

Run:  DB_NAME=uktex_fix_geo python manage.py test api.tests_geo_tracking_views --noinput
"""

import io
from datetime import datetime, timedelta
from unittest.mock import patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .geo_tracking import journey
from .models import (
    AuditLog,
    DutyTrack,
    DutyTrackEvent,
    DutyTrackPoint,
    DutyTrackSegment,
    GeoTrackingSettings,
    OnDutyPunchVerification,
    TrackingConsent,
    TrackingDevice,
)
from .tests_geo_tracking_support import LAT0, LNG0, T0, UTC, GeoBase, bearer, ms, offset_latlng, point, snake_deep

IST = ZoneInfo("Asia/Kolkata")


def drive_points(track, n=30, *, seq0=1, t0=T0, step_m=150.0, step_s=10, north0=0.0, device="d1"):
    rows = []
    for k in range(n):
        lat, lng = offset_latlng(LAT0, LNG0, north0 + k * step_m, 0.0)
        rows.append(
            DutyTrackPoint(
                track=track,
                employee=track.employee,
                device_id=device,
                seq=seq0 + k,
                recorded_at=t0 + timedelta(seconds=k * step_s),
                latitude=round(lat, 6),
                longitude=round(lng, 6),
                accuracy_m=8.0,
                activity="vehicle",
                battery_pct=70,
                network="cell",
            )
        )
    DutyTrackPoint.objects.bulk_create(rows)
    DutyTrack.objects.filter(pk=track.pk).update(points_count=DutyTrackPoint.objects.filter(track=track).count())
    return rows


def dwell_points(track, *, seq0, t0, north, seconds=600, device="d1"):
    lat, lng = offset_latlng(LAT0, LNG0, north, 0.0)
    rows = [
        DutyTrackPoint(
            track=track,
            employee=track.employee,
            device_id=device,
            seq=seq0 + k,
            recorded_at=t0 + timedelta(seconds=10 * k),
            latitude=round(lat, 6),
            longitude=round(lng, 6),
            accuracy_m=8.0,
        )
        for k in range(seconds // 10 + 1)
    ]
    DutyTrackPoint.objects.bulk_create(rows)
    DutyTrack.objects.filter(pk=track.pk).update(points_count=DutyTrackPoint.objects.filter(track=track).count())


class Journeys(GeoBase):
    """A finished journey of Asha's (a drive, a stop, a drive, a gap) with its summary built, and one of Ravi's in Unit 2."""

    def setUp(self):
        super().setUp()
        self.end = T0 + timedelta(minutes=60)
        self.track = self.make_track(self.asha, T0, ended_at=self.end, last_device_id="d1")
        drive_points(self.track, 30, seq0=1)
        dwell_points(self.track, seq0=100, t0=T0 + timedelta(seconds=300), north=4350.0)
        drive_points(self.track, 25, seq0=200, t0=T0 + timedelta(seconds=910), north0=4350.0)
        DutyTrackEvent.objects.create(
            track=self.track, employee=self.asha, kind="TRACKING_STARTED", source="server", started_at=T0
        )
        DutyTrackEvent.objects.create(
            track=self.track,
            employee=self.asha,
            kind="GPS_OFF",
            device_id="d1",
            seq=500,
            started_at=T0 + timedelta(minutes=30),
            ended_at=T0 + timedelta(minutes=40),
            meta={"endSeq": 501},
        )
        DutyTrackEvent.objects.create(
            track=self.track,
            employee=self.asha,
            kind="TRACKING_STOPPED",
            source="server",
            started_at=self.end,
            meta={"reason": "manual"},
        )
        DutyTrack.objects.filter(pk=self.track.pk).update(
            events_count=3, last_latitude=11.04, last_longitude=77.0, last_battery_pct=64, last_network="cell"
        )
        TrackingDevice.objects.create(
            employee=self.asha,
            device_id="d1",
            platform="android",
            os_version="14",
            model="Pixel 6",
            app_version="3.2.0",
            permission_level="always",
        )
        journey.rebuild(self.track)
        self.track.refresh_from_db()
        self.other = self.make_track(self.ravi, T0, ended_at=self.end)
        drive_points(self.other, 12, seq0=1, device="r1")
        journey.rebuild(self.other)


class EmployeeStateTests(GeoBase):
    def test_state_without_a_track(self):
        r = self.as_employee(self.asha, "get", "/api/tracking/state")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(
            set(body),
            {
                "enabled",
                "active",
                "consentRequired",
                "noticeVersion",
                "noticeText",
                "serverTime",
                "track",
                "config",
                "punchPhotos",
            },
        )
        self.assertEqual(
            (
                body["enabled"],
                body["active"],
                body["track"],
                body["consentRequired"],
                body["noticeVersion"],
                body["punchPhotos"],
            ),
            (True, False, None, True, 1, 1),
        )
        self.assertTrue(body["noticeText"])
        self.assertEqual(body["config"]["sampleMovingSeconds"], 3)
        self.assertEqual(body["config"]["uploadBatchPoints"], 30)
        datetime.fromisoformat(body["serverTime"])

    def test_state_with_an_active_track(self):
        track = self.make_track(self.asha, T0)
        body = self.as_employee(self.asha, "get", "/api/tracking/state").json()
        self.assertTrue(body["active"])
        self.assertEqual(set(body["track"]), {"id", "sessionId", "startedAt", "endsAt", "state"})
        self.assertEqual((body["track"]["id"], body["track"]["sessionId"]), (track.pk, track.session_id))
        self.assertEqual(datetime.fromisoformat(body["track"]["startedAt"]), T0)

    def test_the_end_of_the_day_is_the_next_day_end_time_in_indian_time(self):
        self.make_track(self.asha, T0)
        ends = datetime.fromisoformat(
            self.as_employee(self.asha, "get", "/api/tracking/state").json()["track"]["endsAt"]
        )
        local = ends.astimezone(IST)
        self.assertEqual((local.hour, local.minute), (23, 50))
        self.assertEqual(ends.utcoffset(), timedelta(hours=5, minutes=30))
        GeoTrackingSettings.objects.update_or_create(pk=1, defaults={"day_end_time": "22:15"})
        ends = datetime.fromisoformat(
            self.as_employee(self.asha, "get", "/api/tracking/state").json()["track"]["endsAt"]
        ).astimezone(IST)
        self.assertEqual((ends.hour, ends.minute), (22, 15))

    def test_consent_required_follows_the_notice_version(self):
        self.assertTrue(self.as_employee(self.asha, "get", "/api/tracking/state").json()["consentRequired"])
        self.consent(self.asha)
        self.assertFalse(self.as_employee(self.asha, "get", "/api/tracking/state").json()["consentRequired"])
        GeoTrackingSettings.objects.update_or_create(pk=1, defaults={"notice_version": 2})
        body = self.as_employee(self.asha, "get", "/api/tracking/state").json()
        self.assertEqual((body["consentRequired"], body["noticeVersion"]), (True, 2))

    def test_punch_photos_and_config_follow_the_settings(self):
        GeoTrackingSettings.objects.update_or_create(
            pk=1, defaults={"onduty_punch_photos": 2, "upload_batch_points": 50, "sample_moving_seconds": 5}
        )
        body = self.as_employee(self.asha, "get", "/api/tracking/state").json()
        self.assertEqual(
            (body["punchPhotos"], body["config"]["uploadBatchPoints"], body["config"]["sampleMovingSeconds"]),
            (2, 50, 5),
        )

    def test_with_tracking_off_nothing_is_active(self):
        self.make_track(self.asha, T0)
        GeoTrackingSettings.objects.update_or_create(pk=1, defaults={"enabled": False})
        body = self.as_employee(self.asha, "get", "/api/tracking/state").json()
        self.assertEqual((body["enabled"], body["active"], body["track"]), (False, False, None))

    def test_an_ended_track_is_not_active(self):
        self.make_track(self.asha, T0, ended_at=T0 + timedelta(hours=1))
        self.assertFalse(self.as_employee(self.asha, "get", "/api/tracking/state").json()["active"])

    def test_a_session_in_progress_without_a_track_gets_one_when_the_app_asks(self):
        session = self.make_session(self.asha)  # the request was made before tracking existed
        self.assertFalse(DutyTrack.objects.filter(session=session).exists())
        body = self.as_employee(self.asha, "get", "/api/tracking/state").json()
        self.assertTrue(body["active"])
        track = DutyTrack.objects.get(session=session)
        self.assertEqual(track.status, "active")
        self.assertTrue(DutyTrackEvent.objects.filter(track=track, kind="TRACKING_STARTED").exists())
        self.as_employee(self.asha, "get", "/api/tracking/state")
        self.assertEqual(DutyTrack.objects.filter(employee=self.asha).count(), 1)

    def test_a_finished_session_does_not_get_a_track(self):
        self.make_session(self.asha, employee_ended_at=T0)
        self.assertFalse(self.as_employee(self.asha, "get", "/api/tracking/state").json()["active"])
        self.assertEqual(DutyTrack.objects.count(), 0)

    def test_an_hr_token_is_refused(self):
        self.assertEqual(self.as_hr(self.hr_super, "get", "/api/tracking/state").status_code, 403)


class ConsentTests(GeoBase):
    def test_accepting_the_current_notice(self):
        r = self.as_employee(self.asha, "post", "/api/tracking/consent", {"deviceId": "dev-1", "noticeVersion": 1})
        self.assertEqual((r.status_code, r.json()), (200, {"ok": True}))
        c = TrackingConsent.objects.get()
        self.assertEqual((c.employee_id, c.device_id, c.notice_version), (self.asha.pk, "dev-1", 1))
        self.assertIsNotNone(c.accepted_at)

    def test_snake_case_works_and_accepting_twice_records_once(self):
        body = snake_deep({"deviceId": "dev-1", "noticeVersion": 1})
        self.assertEqual(self.as_employee(self.asha, "post", "/api/tracking/consent", body).status_code, 200)
        self.assertEqual(self.as_employee(self.asha, "post", "/api/tracking/consent", body).status_code, 200)
        self.assertEqual(TrackingConsent.objects.count(), 1)
        self.as_employee(self.asha, "post", "/api/tracking/consent", {"deviceId": "dev-2", "noticeVersion": 1})
        self.assertEqual(TrackingConsent.objects.count(), 2)

    def test_a_stale_notice_version_is_refused_with_the_current_one(self):
        GeoTrackingSettings.objects.update_or_create(pk=1, defaults={"notice_version": 3})
        r = self.as_employee(self.asha, "post", "/api/tracking/consent", {"deviceId": "d", "noticeVersion": 2})
        self.assertEqual(
            (r.status_code, r.json()["code"], r.json()["noticeVersion"]), (400, "NOTICE_VERSION_MISMATCH", 3)
        )
        self.assertEqual(TrackingConsent.objects.count(), 0)

    def test_the_notice_version_is_required(self):
        self.assertEqual(
            self.as_employee(self.asha, "post", "/api/tracking/consent", {"deviceId": "d"}).status_code, 400
        )

    def test_an_hr_token_is_refused(self):
        self.assertEqual(
            self.as_hr(self.hr_super, "post", "/api/tracking/consent", {"noticeVersion": 1}).status_code, 403
        )

    def test_after_consent_a_batch_is_taken(self):
        self.make_track(self.asha, T0)
        self.as_employee(self.asha, "post", "/api/tracking/consent", {"deviceId": "dev-1", "noticeVersion": 1})
        r = self.batch(self.asha, [point(1, T0 + timedelta(seconds=5))])
        self.assertEqual((r.status_code, r.json()["accepted"]), (200, 1))


class BatchViewTests(GeoBase):
    def setUp(self):
        super().setUp()
        self.make_track(self.asha, T0)
        self.consent(self.asha)

    def test_camel_and_snake_bodies_over_http(self):
        rows = [point(1, T0 + timedelta(seconds=5)), point(2, T0 + timedelta(seconds=15), LAT0 + 0.001, LNG0)]
        a = self.batch(self.asha, rows, device="dev-a")
        b = self.batch(self.asha, rows, device="dev-b", snake=True)
        self.assertEqual((a.status_code, b.status_code), (200, 200))
        self.assertEqual((a.json()["accepted"], b.json()["accepted"]), (2, 2))

    def test_a_body_that_is_not_json_is_a_400(self):
        r = self.client.post(
            "/api/tracking/batch",
            "{not json",
            content_type="application/json",
            **bearer({"role": "employee", "employeeId": self.asha.id}),
        )
        self.assertEqual(r.status_code, 400)

    def test_a_json_list_instead_of_an_object_is_a_400(self):
        self.assertEqual(self.as_employee(self.asha, "post", "/api/tracking/batch", [1, 2]).status_code, 400)

    def test_no_token_is_a_401_and_an_hr_token_a_403(self):
        self.assertEqual(self.client.post("/api/tracking/batch", {}, content_type="application/json").status_code, 401)
        self.assertEqual(self.as_hr(self.hr_super, "post", "/api/tracking/batch", {"deviceId": "d"}).status_code, 403)

    def test_get_is_not_allowed(self):
        self.assertEqual(self.as_employee(self.asha, "get", "/api/tracking/batch").status_code, 405)


class EmployeeJourneyTests(Journeys):
    def test_my_journeys_lists_only_my_own(self):
        body = self.as_employee(self.asha, "get", "/api/tracking/my-journeys").json()
        self.assertEqual([i["id"] for i in body["results"]], [self.track.pk])
        self.assertEqual((body["total"], body["page"], body["pageSize"]), (1, 1, 25))
        item = body["results"][0]
        self.assertEqual(item["employeeCode"], "G1")
        self.assertEqual(
            self.as_employee(self.ravi, "get", "/api/tracking/my-journeys").json()["results"][0]["id"], self.other.pk
        )

    def test_the_list_item_has_the_contracted_fields(self):
        item = self.as_employee(self.asha, "get", "/api/tracking/my-journeys").json()["results"][0]
        self.assertEqual(
            set(item),
            {
                "id",
                "sessionId",
                "employeeId",
                "employeeCode",
                "employeeName",
                "department",
                "branchName",
                "destination",
                "startedAt",
                "endedAt",
                "endReason",
                "status",
                "state",
                "durationS",
                "distanceM",
                "travelS",
                "stopS",
                "gapS",
                "stopCount",
                "longestStopS",
                "pointsCount",
                "eventsCount",
                "hasMock",
                "approvalStatus",
                "lastPointAt",
                "punchCount",
            },
        )
        self.assertEqual(
            (
                item["employeeName"],
                item["department"],
                item["branchName"],
                item["destination"],
                item["endReason"],
                item["status"],
                item["state"],
                item["approvalStatus"],
            ),
            (
                "Asha Kumar",
                "Sales",
                "Geo Head Office",
                "ABC Traders, Coimbatore",
                "manual",
                "ended",
                "ended",
                "pending_hod",
            ),
        )
        self.assertEqual(
            (item["durationS"], item["stopCount"], item["longestStopS"], item["hasMock"]), (3600, 1, 620, False)
        )
        self.assertAlmostEqual(item["distanceM"], 29 * 150.0 + 24 * 150.0, delta=10)

    def punch(self, session, number, *, status="approved", photo2=False, minutes=10):
        punched = T0.astimezone(IST) + timedelta(minutes=minutes)
        v = OnDutyPunchVerification(
            session=session,
            employee=session.employee,
            punch_date=punched.date(),
            punch_time=punched.time().replace(microsecond=0),
            punch_type="IN" if number % 2 else "OUT",
            punch_number=number,
            latitude=11.0123,
            longitude=77.0456,
            accuracy_m=9.5,
            status=status,
            photo="on_duty_punch_verifications/2026/10/p.jpg",
            photo2="on_duty_punch_verifications/2026/10/p2.jpg" if photo2 else None,
        )
        v.save()
        return v

    def test_the_list_item_counts_the_geo_punches_of_its_session(self):
        # slot 1 and 2 recorded, slot 2 was rejected once and redone (counts once), slot 3 was rejected and never redone
        self.punch(self.track.session, 1)
        self.punch(self.track.session, 2, status="rejected")
        self.punch(self.track.session, 2, status="pending")
        self.punch(self.track.session, 3, status="rejected")
        # another employee's punches never leak into this journey
        self.punch(self.other.session, 1)
        body = self.as_employee(self.asha, "get", "/api/tracking/my-journeys").json()
        self.assertEqual(body["results"][0]["punchCount"], 2)
        self.assertEqual(
            self.as_employee(self.ravi, "get", "/api/tracking/my-journeys").json()["results"][0]["punchCount"], 1
        )

    def test_a_journey_without_punches_says_zero(self):
        self.assertEqual(
            self.as_employee(self.asha, "get", "/api/tracking/my-journeys").json()["results"][0]["punchCount"], 0
        )

    def test_the_detail_lists_the_geo_punches_of_its_session_in_slot_order(self):
        self.punch(self.track.session, 2, status="rejected", photo2=True, minutes=20)
        self.punch(self.track.session, 1, photo2=True, minutes=10)
        self.punch(self.track.session, 2, status="pending", minutes=25)
        self.punch(self.other.session, 1)
        body = self.as_employee(self.asha, "get", f"/api/tracking/journeys/{self.track.pk}").json()
        punches = body["punches"]
        self.assertEqual(
            [(p["punchNumber"], p["status"]) for p in punches], [(1, "approved"), (2, "rejected"), (2, "pending")]
        )
        first = punches[0]
        self.assertEqual(
            set(first),
            {
                "id",
                "punchNumber",
                "punchType",
                "punchDate",
                "punchTime",
                "latitude",
                "longitude",
                "accuracyM",
                "isMocked",
                "hasPhoto",
                "hasPhoto2",
                "status",
                "hrReviewComment",
                "hrReviewedAt",
                "createdAt",
            },
        )
        self.assertEqual((first["punchType"], first["hasPhoto"], first["hasPhoto2"]), ("IN", True, True))
        self.assertEqual((first["latitude"], first["longitude"], first["accuracyM"]), (11.0123, 77.0456, 9.5))
        # the IST wall clock: T0 + 10 minutes
        punched = T0.astimezone(IST) + timedelta(minutes=10)
        self.assertEqual((first["punchDate"], first["punchTime"]), (str(punched.date()), punched.strftime("%H:%M:%S")))
        self.assertEqual((punches[1]["hasPhoto2"], punches[2]["hasPhoto2"]), (True, False))

    def test_the_detail_of_a_journey_with_no_punches_has_an_empty_list(self):
        body = self.as_employee(self.asha, "get", f"/api/tracking/journeys/{self.track.pk}").json()
        self.assertEqual(body["punches"], [])

    def test_my_own_journey_detail_is_readable_and_someone_elses_is_not_found(self):
        r = self.as_employee(self.asha, "get", f"/api/tracking/journeys/{self.track.pk}")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["track"]["id"], self.track.pk)
        self.assertEqual(self.as_employee(self.asha, "get", f"/api/tracking/journeys/{self.other.pk}").status_code, 404)
        self.assertEqual(self.as_employee(self.ravi, "get", f"/api/tracking/journeys/{self.track.pk}").status_code, 404)

    def test_an_employee_reading_their_journey_writes_no_audit_entry(self):
        self.as_employee(self.asha, "get", f"/api/tracking/journeys/{self.track.pk}")
        self.assertFalse(AuditLog.objects.filter(module="geo_tracking").exists())

    def test_filters_and_paging(self):
        later = self.make_track(self.asha, T0 + timedelta(days=1, hours=2), ended_at=T0 + timedelta(days=1, hours=3))
        r = self.as_employee(self.asha, "get", "/api/tracking/my-journeys", pageSize=1).json()
        self.assertEqual((len(r["results"]), r["total"], r["results"][0]["id"]), (1, 2, later.pk))
        r = self.as_employee(self.asha, "get", "/api/tracking/my-journeys", pageSize=1, page=2).json()
        self.assertEqual(r["results"][0]["id"], self.track.pk)
        r = self.as_employee(self.asha, "get", "/api/tracking/my-journeys", status="active").json()
        self.assertEqual(r["total"], 0)

    def test_an_hr_token_on_my_journeys_is_refused(self):
        self.assertEqual(self.as_hr(self.hr_super, "get", "/api/tracking/my-journeys").status_code, 403)


class HrListTests(Journeys):
    def test_the_list_shape_and_the_unscoped_admin_sees_everyone(self):
        body = self.as_hr(self.hr_super, "get", "/api/tracking/journeys").json()
        self.assertEqual(set(body), {"results", "page", "pageSize", "total"})
        self.assertEqual(body["total"], 2)
        self.assertEqual({i["id"] for i in body["results"]}, {self.track.pk, self.other.pk})

    def test_a_branch_hr_sees_only_their_branch(self):
        ids = lambda user: {i["id"] for i in self.as_hr(user, "get", "/api/tracking/journeys").json()["results"]}
        self.assertEqual(ids(self.hr_ho), {self.track.pk})
        self.assertEqual(ids(self.hr_u2), {self.other.pk})

    def test_filter_by_employee(self):
        body = self.as_hr(self.hr_super, "get", "/api/tracking/journeys", employeeId=self.ravi.pk).json()
        self.assertEqual([i["id"] for i in body["results"]], [self.other.pk])
        self.assertEqual(self.as_hr(self.hr_super, "get", "/api/tracking/journeys", employeeId="abc").status_code, 400)

    def test_the_date_filter_uses_indian_calendar_days(self):
        # 18:40 UTC on the 27th is 00:10 IST on the 28th
        night = self.make_track(
            self.meena, datetime(2026, 10, 27, 18, 40, tzinfo=UTC), ended_at=datetime(2026, 10, 27, 20, 0, tzinfo=UTC)
        )
        ids = lambda **q: {
            i["id"] for i in self.as_hr(self.hr_super, "get", "/api/tracking/journeys", **q).json()["results"]
        }
        self.assertIn(night.pk, ids(**{"from": "2026-10-28", "to": "2026-10-28"}))
        self.assertNotIn(night.pk, ids(**{"from": "2026-10-27", "to": "2026-10-27"}))
        self.assertIn(self.track.pk, ids(**{"from": "2026-10-28", "to": "2026-10-28"}))
        self.assertEqual(ids(**{"from": "2026-10-29"}), set())
        self.assertEqual(ids(**{"to": "2026-10-27"}), set())

    def test_the_date_filter_takes_iso_times_too(self):
        ids = {
            i["id"]
            for i in self.as_hr(
                self.hr_super,
                "get",
                "/api/tracking/journeys",
                **{"from": (T0 - timedelta(minutes=1)).isoformat(), "to": (T0 + timedelta(minutes=1)).isoformat()},
            ).json()["results"]
        }
        self.assertEqual(ids, {self.track.pk, self.other.pk})

    def test_a_bad_date_is_a_400(self):
        self.assertEqual(
            self.as_hr(self.hr_super, "get", "/api/tracking/journeys", **{"from": "last tuesday"}).status_code, 400
        )

    def test_status_filter(self):
        active = self.make_track(self.meena, T0 + timedelta(hours=2))
        ids = lambda status: {
            i["id"] for i in self.as_hr(self.hr_super, "get", "/api/tracking/journeys", status=status).json()["results"]
        }
        self.assertEqual(ids("active"), {active.pk})
        self.assertEqual(ids("ended"), {self.track.pk, self.other.pk})

    def test_the_end_reason_filter(self):
        rejected = self.make_track(
            self.meena, T0 + timedelta(hours=3), ended_at=T0 + timedelta(hours=4), end_reason="rejected"
        )
        ids = lambda **q: {
            i["id"] for i in self.as_hr(self.hr_super, "get", "/api/tracking/journeys", **q).json()["results"]
        }
        self.assertEqual(ids(endReason="rejected"), {rejected.pk})
        self.assertEqual(ids(endReason="manual"), {self.track.pk, self.other.pk})
        self.assertEqual(ids(end_reason="auto_day_end"), set())
        self.assertEqual(len(ids(endReason="not-a-reason")), 3)  # an unknown value filters nothing

    def test_paging_and_the_page_size_cap(self):
        for k in range(5):
            self.make_track(self.meena, T0 + timedelta(days=2 + k), ended_at=T0 + timedelta(days=2 + k, hours=1))
        body = self.as_hr(self.hr_super, "get", "/api/tracking/journeys", pageSize=3, page=2).json()
        self.assertEqual((len(body["results"]), body["total"], body["page"], body["pageSize"]), (3, 7, 2, 3))
        self.assertEqual(
            self.as_hr(self.hr_super, "get", "/api/tracking/journeys", pageSize=9999).json()["pageSize"], 200
        )

    def test_a_viewer_may_read_the_list_and_a_role_without_geo_attendance_may_not(self):
        self.assertEqual(self.as_hr(self.hr_viewer, "get", "/api/tracking/journeys").status_code, 200)
        self.assertEqual(self.as_hr(self.hr_nogeo, "get", "/api/tracking/journeys").status_code, 403)

    def test_a_live_track_reads_as_offline_once_it_has_been_quiet_too_long(self):
        quiet = self.make_track(
            self.meena,
            datetime.now(UTC) - timedelta(hours=1),
            last_point_at=datetime.now(UTC) - timedelta(minutes=30),
            state="moving",
        )
        item = [
            i
            for i in self.as_hr(self.hr_super, "get", "/api/tracking/journeys").json()["results"]
            if i["id"] == quiet.pk
        ][0]
        self.assertEqual((item["status"], item["state"]), ("active", "offline"))


class HrDetailTests(Journeys):
    def detail(self, user=None, pk=None):
        return self.as_hr(user or self.hr_super, "get", f"/api/tracking/journeys/{pk or self.track.pk}")

    def test_the_detail_has_every_part_of_the_contract(self):
        body = self.detail().json()
        self.assertEqual(
            set(body),
            {"track", "summary", "segments", "events", "path", "pathSource", "device", "timeline", "punches"},
        )
        track = body["track"]
        self.assertEqual(
            (track["id"], track["lastLatitude"], track["lastLongitude"], track["lastBatteryPct"], track["lastNetwork"]),
            (self.track.pk, 11.04, 77.0, 64, "cell"),
        )
        self.assertEqual(
            set(body["summary"]),
            {
                "distanceM",
                "travelS",
                "stopS",
                "gapS",
                "stopCount",
                "longestStopS",
                "maxSpeedKmh",
                "durationS",
                "firstPointAt",
                "summaryThrough",
                "rebuiltAt",
                "pointsCount",
                "eventsCount",
                "mockPoints",
                "flags",
            },
        )
        self.assertEqual(
            (body["summary"]["stopCount"], body["summary"]["stopS"], body["summary"]["durationS"]), (1, 620, 3600)
        )

    def test_the_segments_carry_the_stationary_period_with_its_place_and_times(self):
        segs = self.detail().json()["segments"]
        self.assertEqual([s["kind"] for s in segs], ["travel", "stop", "travel", "gap"])
        stop = segs[1]
        self.assertEqual(
            set(stop),
            {
                "id",
                "seqNo",
                "kind",
                "startedAt",
                "endedAt",
                "durationS",
                "distanceM",
                "startLat",
                "startLng",
                "endLat",
                "endLng",
                "centerLat",
                "centerLng",
                "radiusM",
                "pointCount",
                "avgSpeedKmh",
                "maxSpeedKmh",
                "polyline",
                "address",
                "cause",
                "estimated",
                "eventId",
                "meta",
            },
        )
        self.assertEqual((stop["durationS"], stop["seqNo"]), (620, 2))
        self.assertEqual(datetime.fromisoformat(stop["startedAt"]), T0 + timedelta(seconds=290))
        self.assertAlmostEqual(stop["centerLat"], offset_latlng(LAT0, LNG0, 4350.0, 0.0)[0], places=4)
        travel = segs[0]
        self.assertGreater(len(travel["polyline"]), 1)
        self.assertEqual(travel["polyline"][0], [round(LAT0, 6), round(LNG0, 6)])

    def test_the_events_are_listed_without_internal_bookkeeping(self):
        events = self.detail().json()["events"]
        self.assertEqual([e["kind"] for e in events], ["TRACKING_STARTED", "GPS_OFF", "TRACKING_STOPPED"])
        gps = events[1]
        self.assertEqual((gps["durationS"], gps["open"], gps["source"]), (600, False, "device"))
        self.assertNotIn("endSeq", gps["meta"])

    def test_the_path_is_lat_lng_and_epoch_milliseconds_capped_at_1500(self):
        body = self.detail().json()
        self.assertEqual(body["pathSource"], "points")
        self.assertLessEqual(len(body["path"]), 1500)
        lat, lng, t = body["path"][0]
        self.assertAlmostEqual(lat, LAT0, places=5)
        self.assertEqual(t, ms(T0))
        self.assertEqual(body["path"][-1][2], ms(T0 + timedelta(seconds=910 + 24 * 10)))

    def test_the_path_falls_back_to_the_segments_once_the_raw_points_are_gone(self):
        DutyTrackPoint.objects.filter(track=self.track).delete()
        body = self.detail().json()
        self.assertEqual(body["pathSource"], "segments")
        self.assertGreater(len(body["path"]), 2)

    def test_the_device_card(self):
        device = self.detail().json()["device"]
        self.assertEqual(
            (device["model"], device["osVersion"], device["appVersion"], device["permissionLevel"]),
            ("Pixel 6", "14", "3.2.0", "always"),
        )
        self.assertIsNone(self.detail(pk=self.other.pk).json()["device"])

    def test_the_timeline_merges_segments_and_events_in_time_order(self):
        timeline = self.detail().json()["timeline"]
        self.assertEqual(timeline[0]["kind"], "TRACKING_STARTED")
        self.assertEqual(timeline[-1]["kind"], "TRACKING_STOPPED")
        starts = [i["startedAt"] for i in timeline]
        self.assertEqual(starts, sorted(starts))
        self.assertEqual({i["type"] for i in timeline}, {"segment", "event"})
        self.assertEqual(len(timeline), 4 + 3)
        stop = [i for i in timeline if i["kind"] == "stop"][0]
        self.assertEqual((stop["title"], stop["durationS"]), ("Stopped", 620))

    def test_another_branchs_journey_is_not_found(self):
        self.assertEqual(self.detail(self.hr_ho, self.other.pk).status_code, 404)
        self.assertEqual(self.detail(self.hr_u2, self.track.pk).status_code, 404)
        self.assertEqual(self.detail(self.hr_ho, self.track.pk).status_code, 200)
        self.assertEqual(self.detail(pk=999999).status_code, 404)

    def test_every_hr_view_of_a_journey_is_audit_logged(self):
        self.detail(self.hr_ho)
        entry = AuditLog.objects.get(module="geo_tracking", action="view")
        self.assertEqual(entry.record_id, self.track.pk)
        self.assertEqual(entry.user_type, "hr")
        self.assertEqual(entry.branch_id, self.ho.pk)
        self.assertIn("Asha Kumar", entry.record_description)
        self.detail(self.hr_ho)
        self.assertEqual(AuditLog.objects.filter(module="geo_tracking", action="view").count(), 2)

    def test_a_refused_view_writes_no_audit_entry(self):
        self.detail(self.hr_ho, self.other.pk)
        self.assertFalse(AuditLog.objects.filter(module="geo_tracking").exists())

    def test_the_raw_points_are_paged_by_cursor(self):
        url = f"/api/tracking/journeys/{self.track.pk}/points"
        first = self.as_hr(self.hr_super, "get", url, limit=40).json()
        self.assertEqual(len(first["results"]), 40)
        self.assertIsNotNone(first["nextAfter"])
        self.assertEqual(
            set(first["results"][0]),
            {
                "id",
                "seq",
                "recordedAt",
                "receivedAt",
                "latitude",
                "longitude",
                "accuracyM",
                "speedMps",
                "bearingDeg",
                "altitudeM",
                "activity",
                "batteryPct",
                "charging",
                "network",
                "provider",
                "isMock",
                "timeSuspect",
            },
        )
        seen = [p["id"] for p in first["results"]]
        after = first["nextAfter"]
        while after:
            page = self.as_hr(self.hr_super, "get", url, limit=40, after=after).json()
            seen += [p["id"] for p in page["results"]]
            after = page["nextAfter"]
        total = DutyTrackPoint.objects.filter(track=self.track).count()
        self.assertEqual((len(seen), len(set(seen)), seen == sorted(seen)), (total, total, True))
        self.assertEqual(
            self.as_hr(self.hr_ho, "get", f"/api/tracking/journeys/{self.other.pk}/points").status_code, 404
        )

    def test_viewing_the_raw_points_is_audit_logged(self):
        self.as_hr(self.hr_super, "get", f"/api/tracking/journeys/{self.track.pk}/points")
        self.assertEqual(
            AuditLog.objects.filter(module="geo_tracking", action="view", record_id=self.track.pk).count(), 1
        )


class ExportTests(Journeys):
    def export(self, user=None, pk=None, **q):
        return self.as_hr(user or self.hr_super, "get", f"/api/tracking/journeys/{pk or self.track.pk}/export", **q)

    def test_the_default_export_is_an_excel_file_with_summary_segments_and_events(self):
        r = self.export()
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn(f"journey-{self.track.pk}-G1-20261028.xlsx", r["Content-Disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Summary", "Segments", "Events"])
        summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(
            (summary["Employee"], summary["Employee code"], summary["Stops"], summary["End reason"]),
            ("Asha Kumar", "G1", 1, "manual"),
        )
        self.assertEqual(summary["Longest stop"], "0:10:20")
        self.assertEqual(summary["Tracking started (IST)"], "28-Oct-2026 10:00:00")
        segments = list(wb["Segments"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([row[1] for row in segments], ["Travel", "Stop", "Travel", "No data"])
        self.assertEqual(segments[1][4], "0:10:20")
        events = list(wb["Events"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([row[1] for row in events], ["TRACKING_STARTED", "GPS_OFF", "TRACKING_STOPPED"])

    def test_csv_export(self):
        r = self.export(format="csv")
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/csv", r["Content-Type"])
        text = r.content.decode("utf-8")
        for needle in ("Summary", "Segments", "Events", "Asha Kumar", "GPS_OFF", "Stop"):
            self.assertIn(needle, text)

    def test_an_unknown_format_is_a_400(self):
        self.assertEqual(self.export(format="docx").status_code, 400)

    def test_the_pdf_is_the_journey_report_with_a_map_and_every_section(self):
        from pypdf import PdfReader

        with patch(
            "api.geo_tracking.static_map._fetch", return_value=None
        ):  # no tile server in a test: the grid fallback
            r = self.export(format="pdf")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertIn(f"journey-{self.track.pk}-G1-20261028.pdf", r["Content-Disposition"])
        self.assertTrue(r.content.startswith(b"%PDF"))
        reader = PdfReader(io.BytesIO(r.content))
        text = " ".join(page.extract_text() for page in reader.pages)
        for needle in (
            "ON-DUTY JOURNEY REPORT",
            "Asha Kumar",
            "1. Journey map",
            "2. Journey details",
            "3. Route breakdown",
            "4. Stop details",
            "5. Attendance details",
            "7. Location history",
            "Stop 1",
            "Page 1 of",
        ):
            self.assertIn(needle, text)
        # the picture of the map is in it
        self.assertTrue(any(page.images for page in reader.pages))

    def test_the_pdf_is_scoped_and_audit_logged_like_the_other_exports(self):
        self.assertEqual(self.export(self.hr_ho, self.other.pk, format="pdf").status_code, 404)
        with patch("api.geo_tracking.static_map._fetch", return_value=None):
            self.export(self.hr_ho, format="pdf")
        entry = AuditLog.objects.get(module="geo_tracking", action="export")
        self.assertEqual(entry.record_id, self.track.pk)
        self.assertIn("PDF", entry.record_description)

    def test_a_spreadsheet_cannot_be_made_to_run_a_formula_from_a_destination(self):
        self.track.session.destination = '=HYPERLINK("http://evil")'
        self.track.session.save()
        wb = load_workbook(io.BytesIO(self.export().content))
        summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
        self.assertTrue(summary["Destination"].startswith("'="))

    def test_scope_and_audit(self):
        self.assertEqual(self.export(self.hr_ho, self.other.pk).status_code, 404)
        self.export(self.hr_ho)
        entry = AuditLog.objects.get(module="geo_tracking", action="export")
        self.assertEqual(entry.record_id, self.track.pk)

    def test_an_employee_token_is_refused(self):
        self.assertEqual(
            self.as_employee(self.asha, "get", f"/api/tracking/journeys/{self.track.pk}/export").status_code, 403
        )


class RebuildViewTests(Journeys):
    def rebuild(self, user=None, pk=None):
        return self.as_hr(user or self.hr_super, "post", f"/api/tracking/journeys/{pk or self.track.pk}/rebuild")

    def test_rebuild_recomputes_from_the_raw_rows_and_answers_the_detail(self):
        DutyTrackSegment.objects.filter(track=self.track).delete()
        DutyTrack.objects.filter(pk=self.track.pk).update(stop_count=0, distance_m=0)
        r = self.rebuild()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            set(r.json()),
            {"track", "summary", "segments", "events", "path", "pathSource", "device", "timeline", "punches"},
        )
        self.assertEqual(r.json()["summary"]["stopCount"], 1)
        self.assertEqual(DutyTrackSegment.objects.filter(track=self.track).count(), 4)

    def test_rebuild_needs_edit_access_and_branch_scope(self):
        self.assertEqual(self.rebuild(self.hr_viewer).status_code, 403)
        self.assertEqual(self.rebuild(self.hr_ho, self.other.pk).status_code, 404)
        self.assertEqual(self.rebuild(self.hr_ho).status_code, 200)

    def test_rebuild_is_audit_logged(self):
        self.rebuild(self.hr_ho)
        self.assertEqual(
            AuditLog.objects.filter(module="geo_tracking", action="update", record_id=self.track.pk).count(), 1
        )

    def test_a_purged_journey_cannot_be_rebuilt(self):
        DutyTrack.objects.filter(pk=self.track.pk).update(flags={"purgedAt": "2027-10-28T00:00:00+00:00"})
        r = self.rebuild()
        self.assertEqual(r.status_code, 409)
        self.assertEqual(DutyTrackSegment.objects.filter(track=self.track).count(), 4)

    def test_get_is_not_allowed(self):
        self.assertEqual(
            self.as_hr(self.hr_super, "get", f"/api/tracking/journeys/{self.track.pk}/rebuild").status_code, 405
        )


class LiveTests(GeoBase):
    def setUp(self):
        super().setUp()
        self.now = datetime.now(UTC)
        self.asha_track = self.make_track(
            self.asha,
            self.now - timedelta(hours=1),
            last_point_at=self.now - timedelta(seconds=30),
            last_latitude=11.01,
            last_longitude=77.02,
            last_accuracy_m=9.5,
            last_battery_pct=55,
            last_network="cell",
            state="moving",
            state_since=self.now - timedelta(minutes=5),
            points_count=40,
            mock_points=0,
            stop_count=2,
            distance_m=3200.0,
        )
        self.ravi_track = self.make_track(
            self.ravi,
            self.now - timedelta(hours=2),
            last_point_at=self.now - timedelta(seconds=40),
            last_latitude=11.5,
            last_longitude=77.5,
            state="stationary",
            mock_points=3,
        )
        self.make_track(self.meena, self.now - timedelta(hours=5), ended_at=self.now - timedelta(hours=3))  # not live

    def live(self, user):
        return self.as_hr(user, "get", "/api/tracking/live")

    def test_live_lists_only_the_active_tracks_with_the_contracted_fields(self):
        rows = self.live(self.hr_super).json()
        self.assertEqual({r["trackId"] for r in rows}, {self.asha_track.pk, self.ravi_track.pk})
        row = [r for r in rows if r["trackId"] == self.asha_track.pk][0]
        self.assertEqual(
            set(row),
            {
                "trackId",
                "sessionId",
                "employeeId",
                "employeeCode",
                "employeeName",
                "department",
                "branchName",
                "destination",
                "state",
                "stateSince",
                "latitude",
                "longitude",
                "accuracyM",
                "lastPointAt",
                "batteryPct",
                "network",
                "isMocked",
                "startedAt",
                "approvalStatus",
                "stopCount",
                "distanceM",
            },
        )
        self.assertEqual(
            (
                row["employeeName"],
                row["state"],
                row["latitude"],
                row["longitude"],
                row["accuracyM"],
                row["batteryPct"],
                row["network"],
            ),
            ("Asha Kumar", "moving", 11.01, 77.02, 9.5, 55, "cell"),
        )
        self.assertEqual(
            (row["approvalStatus"], row["stopCount"], row["distanceM"], row["isMocked"]),
            ("pending_hod", 2, 3200.0, False),
        )
        self.assertTrue([r for r in rows if r["trackId"] == self.ravi_track.pk][0]["isMocked"])

    def test_live_is_branch_scoped(self):
        self.assertEqual({r["trackId"] for r in self.live(self.hr_ho).json()}, {self.asha_track.pk})
        self.assertEqual({r["trackId"] for r in self.live(self.hr_u2).json()}, {self.ravi_track.pk})

    def test_a_quiet_track_reads_offline_without_waiting_for_the_watchdog(self):
        DutyTrack.objects.filter(pk=self.asha_track.pk).update(last_point_at=self.now - timedelta(minutes=20))
        row = [r for r in self.live(self.hr_super).json() if r["trackId"] == self.asha_track.pk][0]
        self.assertEqual(row["state"], "offline")

    def test_location_off_is_kept_even_when_quiet(self):
        DutyTrack.objects.filter(pk=self.asha_track.pk).update(
            last_point_at=self.now - timedelta(minutes=20), state="location_off"
        )
        row = [r for r in self.live(self.hr_super).json() if r["trackId"] == self.asha_track.pk][0]
        self.assertEqual(row["state"], "location_off")

    def test_an_employee_token_is_refused_and_live_needs_the_geo_module(self):
        self.assertEqual(self.as_employee(self.asha, "get", "/api/tracking/live").status_code, 403)
        self.assertEqual(self.live(self.hr_nogeo).status_code, 403)
        self.assertEqual(self.live(self.hr_viewer).status_code, 200)


class SettingsApiTests(GeoBase):
    URL = "/api/tracking/settings"

    def test_get_returns_the_whole_object(self):
        body = self.as_hr(self.hr_super, "get", self.URL).json()
        self.assertEqual(
            set(body),
            {
                "enabled",
                "stopSeconds",
                "stopRadiusM",
                "stopAccuracyFactor",
                "mergeStopGapSeconds",
                "gapSeconds",
                "signalLostSeconds",
                "maxAccuracyM",
                "maxSpeedKmh",
                "sampleMovingSeconds",
                "sampleStillSeconds",
                "distanceFilterM",
                "uploadBatchPoints",
                "uploadIntervalSeconds",
                "retentionDays",
                "dayEndTime",
                "ondutyPunchPhotos",
                "punchPhotos",
                "noticeVersion",
                "noticeText",
                "reverseGeocodeProvider",
                "updatedBy",
                "updatedAt",
            },
        )
        self.assertEqual(
            (
                body["stopSeconds"],
                body["dayEndTime"],
                body["ondutyPunchPhotos"],
                body["reverseGeocodeProvider"],
                body["enabled"],
            ),
            (180, "23:50", 1, "off", True),
        )
        self.assertEqual((body["updatedBy"], body["updatedAt"]), (None, None))

    def test_put_is_a_partial_update_and_returns_the_full_object(self):
        r = self.as_hr(self.hr_super, "put", self.URL, {"stopSeconds": 240, "retentionDays": 90})
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(
            (body["stopSeconds"], body["retentionDays"], body["stopRadiusM"], body["gapSeconds"]), (240, 90, 50, 300)
        )
        self.assertEqual((body["updatedBy"], bool(body["updatedAt"])), ("Admin", True))
        s = GeoTrackingSettings.get()
        self.assertEqual((s.stop_seconds, s.retention_days, s.updated_by_id), (240, 90, self.hr_super.pk))

    def test_snake_case_keys_and_the_photo_policy_aliases(self):
        r = self.as_hr(self.hr_super, "put", self.URL, {"stop_radius_m": 80, "punchPhotos": 2})
        self.assertEqual(
            (r.status_code, r.json()["stopRadiusM"], r.json()["ondutyPunchPhotos"], r.json()["punchPhotos"]),
            (200, 80, 2, 2),
        )
        self.assertEqual(
            self.as_hr(self.hr_super, "put", self.URL, {"onduty_punch_photos": 1}).json()["ondutyPunchPhotos"], 1
        )

    def test_unknown_keys_are_ignored_and_an_empty_body_changes_nothing(self):
        r = self.as_hr(self.hr_super, "put", self.URL, {"colour": "red"})
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.json()["updatedAt"])

    def test_every_validation_rule_answers_400_with_one_sentence_and_saves_nothing(self):
        bad = [
            {"stopSeconds": 10},
            {"stopSeconds": 99999},
            {"stopSeconds": "soon"},
            {"stopSeconds": 90.5},
            {"stopRadiusM": 5},
            {"stopRadiusM": 900},
            {"stopAccuracyFactor": 0.1},
            {"stopAccuracyFactor": 9},
            {"stopAccuracyFactor": "x"},
            {"mergeStopGapSeconds": -1},
            {"mergeStopGapSeconds": 99999},
            {"gapSeconds": 5},
            {"signalLostSeconds": 30},
            {"maxAccuracyM": 3},
            {"maxSpeedKmh": 5},
            {"maxSpeedKmh": 9000},
            {"sampleMovingSeconds": 0},
            {"sampleStillSeconds": 5},
            {"distanceFilterM": 500},
            {"uploadBatchPoints": 1},
            {"uploadBatchPoints": 5000},
            {"uploadIntervalSeconds": 1},
            {"retentionDays": 1},
            {"retentionDays": 100000},
            {"dayEndTime": "25:00"},
            {"dayEndTime": "tea time"},
            {"dayEndTime": "09:00"},
            {"dayEndTime": 2350},
            {"ondutyPunchPhotos": 3},
            {"ondutyPunchPhotos": "two"},
            {"ondutyPunchPhotos": 0},
            {"reverseGeocodeProvider": "google"},
            {"reverseGeocodeProvider": 5},
            {"noticeText": "   "},
            {"noticeText": "x" * 5001},
            {"noticeText": 12},
            {"enabled": "maybe"},
            {"enabled": 7},
            {"noticeVersion": 0},
        ]
        for payload in bad:
            r = self.as_hr(self.hr_super, "put", self.URL, payload)
            self.assertEqual(r.status_code, 400, payload)
            self.assertIsInstance(r.json()["error"], str, payload)
            self.assertTrue(r.json()["error"], payload)
        self.assertIsNone(GeoTrackingSettings.get().updated_at)
        self.assertEqual(GeoTrackingSettings.get().stop_seconds, 180)

    def test_one_bad_value_rejects_the_whole_update(self):
        r = self.as_hr(self.hr_super, "put", self.URL, {"stopSeconds": 300, "maxSpeedKmh": 5})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(GeoTrackingSettings.get().stop_seconds, 180)

    def test_accepted_edges(self):
        ok = {
            "stopSeconds": 60,
            "stopRadiusM": 500,
            "stopAccuracyFactor": 5,
            "mergeStopGapSeconds": 0,
            "distanceFilterM": 0,
            "dayEndTime": "12:00",
            "enabled": False,
            "reverseGeocodeProvider": "nominatim",
        }
        r = self.as_hr(self.hr_super, "put", self.URL, ok)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(
            (r.json()["dayEndTime"], r.json()["enabled"], r.json()["reverseGeocodeProvider"]),
            ("12:00", False, "nominatim"),
        )
        self.assertEqual(
            self.as_hr(self.hr_super, "put", self.URL, {"dayEndTime": "23:59:00", "enabled": "true"}).json()[
                "dayEndTime"
            ],
            "23:59",
        )

    def test_changing_the_notice_text_raises_the_version_unless_one_is_given(self):
        r = self.as_hr(self.hr_super, "put", self.URL, {"noticeText": "New wording about location."})
        self.assertEqual((r.json()["noticeText"], r.json()["noticeVersion"]), ("New wording about location.", 2))
        r = self.as_hr(self.hr_super, "put", self.URL, {"noticeText": "New wording about location."})
        self.assertEqual(r.json()["noticeVersion"], 2)  # unchanged text: no new version
        r = self.as_hr(self.hr_super, "put", self.URL, {"noticeText": "Another one.", "noticeVersion": 7})
        self.assertEqual(r.json()["noticeVersion"], 7)
        r = self.as_hr(self.hr_super, "put", self.URL, {"noticeVersion": 3})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(GeoTrackingSettings.get().notice_version, 7)

    def test_a_branch_login_can_read_the_rules_but_not_change_them(self):
        self.assertEqual(self.as_hr(self.hr_ho, "get", self.URL).status_code, 200)
        r = self.as_hr(self.hr_ho, "put", self.URL, {"stopSeconds": 200})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(GeoTrackingSettings.get().stop_seconds, 180)

    def test_view_only_and_no_access_follow_the_settings_attendance_permission(self):
        self.assertEqual(self.as_hr(self.hr_viewer, "get", self.URL).status_code, 200)
        self.assertEqual(self.as_hr(self.hr_viewer, "put", self.URL, {"stopSeconds": 200}).status_code, 403)
        self.assertEqual(self.as_hr(self.hr_nogeo, "get", self.URL).status_code, 403)
        self.assertEqual(self.as_hr(self.hr_editor, "put", self.URL, {"stopSeconds": 200}).status_code, 200)
        # a sub-key can be given on its own: settings.attendance alone is enough, geo_attendance is not
        only_geo = self.hr_user("onlygeo", {"geo_attendance": "edit"})
        self.assertEqual(self.as_hr(only_geo, "get", self.URL).status_code, 403)
        only_attendance = self.hr_user("onlyatt", {"settings.attendance": "edit"})
        self.assertEqual(self.as_hr(only_attendance, "put", self.URL, {"stopSeconds": 210}).status_code, 200)

    def test_changes_are_audit_logged_with_old_and_new_values(self):
        self.as_hr(self.hr_editor, "put", self.URL, {"stopSeconds": 240, "dayEndTime": "23:30"})
        entry = AuditLog.objects.get(module="geo_tracking", action="update")
        self.assertEqual(entry.old_values, {"stop_seconds": 180, "day_end_time": "23:50"})
        self.assertEqual(entry.new_values, {"stop_seconds": 240, "day_end_time": "23:30"})
        self.as_hr(self.hr_editor, "put", self.URL, {"stopSeconds": 240})  # nothing changed: nothing logged
        self.assertEqual(AuditLog.objects.filter(module="geo_tracking").count(), 1)

    def test_an_employee_token_is_refused(self):
        self.assertEqual(self.as_employee(self.asha, "get", self.URL).status_code, 403)
        self.assertEqual(self.as_employee(self.asha, "put", self.URL, {"stopSeconds": 200}).status_code, 403)


class PermissionMatrixTests(Journeys):
    """Employee tokens on HR routes are 403, HR tokens on employee routes are 403, no token is 401."""

    def hr_routes(self):
        pk = self.track.pk
        return [
            ("get", "/api/tracking/journeys"),
            ("get", f"/api/tracking/journeys/{pk}/points"),
            ("get", f"/api/tracking/journeys/{pk}/trail"),
            ("get", f"/api/tracking/journeys/{pk}/export"),
            ("post", f"/api/tracking/journeys/{pk}/rebuild"),
            ("get", "/api/tracking/live"),
            ("get", "/api/tracking/settings"),
            ("put", "/api/tracking/settings"),
        ]

    def employee_routes(self):
        return [
            ("get", "/api/tracking/state"),
            ("post", "/api/tracking/consent"),
            ("post", "/api/tracking/batch"),
            ("get", "/api/tracking/my-journeys"),
        ]

    def test_an_employee_token_on_every_hr_route_is_forbidden(self):
        for method, path in self.hr_routes():
            self.assertEqual(self.as_employee(self.asha, method, path).status_code, 403, f"{method} {path}")

    def test_an_hr_token_on_every_employee_route_is_forbidden(self):
        for method, path in self.employee_routes():
            self.assertEqual(self.as_hr(self.hr_super, method, path).status_code, 403, f"{method} {path}")

    def test_no_token_is_unauthorised_everywhere(self):
        for method, path in (
            self.hr_routes() + self.employee_routes() + [("get", f"/api/tracking/journeys/{self.track.pk}")]
        ):
            r = getattr(self.client, method)(path, **({"content_type": "application/json"} if method != "get" else {}))
            self.assertEqual(r.status_code, 401, f"{method} {path}")

    def test_the_shared_detail_serves_both_roles_but_scopes_each(self):
        url = f"/api/tracking/journeys/{self.track.pk}"
        self.assertEqual(self.as_employee(self.asha, "get", url).status_code, 200)
        self.assertEqual(self.as_hr(self.hr_super, "get", url).status_code, 200)
        self.assertEqual(self.as_employee(self.meena, "get", url).status_code, 404)
        self.assertEqual(self.as_hr(self.hr_nogeo, "get", url).status_code, 403)
