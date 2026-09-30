"""
Report Center - Leave & Requests group (requests_leave): leave register, leave balance statement, monthly leave
summary, daily leave load by department, holiday calendar, casual leave records, casual leave eligibility,
employee requests register, pending approvals ageing and approver workload.

Every number asserted here is derived by hand from the fixtures in each class's setUpTestData (dates are fixed;
"today" is pinned to 2026-09-29 by patching the report context's clock).

Run via: python manage.py test api.tests_reporting_requests_leave -v 2
"""

import io
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import patch

from django.apps import apps
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .clock import FACTORY_TZ
from .jwt_utils import sign_token
from .models import (
    Advance,
    AttendanceDayRecord,
    AttendanceOverrideRequest,
    Branch,
    CasualLeaveRequest,
    Department,
    DepartmentManager,
    Designation,
    Employee,
    EmployeePermission,
    EmployeeRequest,
    HRUser,
    Holiday,
    LeaveBalance,
    LeaveRequest,
    LeaveType,
    ManagerDepartmentAssignment,
    ManagerEmployeeAssignment,
    MissingPunchRequest,
    OnDutyPunchVerification,
    OnDutySession,
    OutpassRequest,
    PayrollSettings,
    ResignationRequest,
    Role,
)
from .permission_registry import all_module_keys
from .reporting import registry

TODAY = date(2026, 9, 29)
MODULES_ALL = {
    "reports": "view",
    "leave": "view",
    "casual_leave": "view",
    "requests": "view",
    "missing_punch": "view",
    "geo_attendance": "view",
    "attendance": "view",
    "recruitment.resignations": "view",
    "settlement": "view",
}
MY_IDS = (
    "leave-register",
    "leave-balance-statement",
    "leave-summary-monthly",
    "leave-department-calendar",
    "casual-leave-register",
    "casual-leave-eligibility-usage",
    "holiday-list",
    "employee-requests-register",
    "pending-approvals-ageing",
    "approver-workload-turnaround",
)


def ist(y, m, d, hh=0, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=FACTORY_TZ)


def hdr(user):
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def mk_hr(name, perms, branch=None):
    role = Role.objects.create(name=name, permissions=perms)
    return HRUser.objects.create(username=name, password_hash="x", role=role, branch=branch)


def mk_emp(code, dept, branch, **kw):
    kw.setdefault("employment_type", "staff")
    return Employee.objects.create(
        employee_code=code, first_name=code, last_name="T", department=dept, branch=branch, **kw
    )


def leave_type(code, name, is_paid=True):
    """The default leave types (AL SL CL EL ML PL) are seeded by migration 0088, so reuse them."""
    lt, _ = LeaveType.objects.get_or_create(code=code, defaults={"name": name, "is_paid": is_paid})
    LeaveType.objects.filter(pk=lt.pk).update(name=name, is_paid=is_paid, is_active=True)
    lt.refresh_from_db()
    return lt


def stamp(obj, **fields):
    """Set columns Django would otherwise auto-fill (created_at is auto_now_add)."""
    type(obj).objects.filter(pk=obj.pk).update(**fields)
    return obj


def data_rows(body):
    return [r for r in body["rows"] if "_kind" not in r]


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.b1 = Branch.objects.create(name="Unit 1", code="RLB1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="RLB2")
        cls.cutting = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.sewing = Department.objects.create(name="SEWING", branch=cls.b2)
        cls.packing = Department.objects.create(name="PACKING", branch=cls.b1)
        cls.tailor = Designation.objects.create(title="Tailor")
        cls.a = mk_emp("RL_A", cls.cutting, cls.b1, designation=cls.tailor, join_date="2025-01-10")
        cls.b = mk_emp("RL_B", cls.cutting, cls.b1, employment_type="production", join_date="2025-02-01")
        cls.c = mk_emp("RL_C", cls.sewing, cls.b2, join_date="15/06/2024")
        cls.d = mk_emp("RL_D", cls.cutting, cls.b1, status="inactive")

        cls.admin = HRUser.objects.create(username="rl_admin", password_hash="x", is_super_admin=True)

        def hr(name, perms, branch=None):
            role = Role.objects.create(name=name, permissions=perms)
            return HRUser.objects.create(username=name, password_hash="x", role=role, branch=branch)

        cls.branch_user = hr("rl_branch1", MODULES_ALL, cls.b1)
        cls.branch_user2 = hr("rl_branch2", MODULES_ALL, cls.b2)
        cls.reports_only = hr("rl_reports_only", {"reports": "view"})
        cls.leave_only = hr("rl_leave_only", {"reports": "view", "leave": "view"})
        cls.requests_only = hr("rl_requests_only", {"reports": "view", "requests": "view"})
        cls.attendance_only = hr("rl_attendance_only", {"reports": "view", "attendance": "view"})

    def setUp(self):
        patcher = patch("api.reporting.filters.ist_today", return_value=TODAY)
        patcher.start()
        self.addCleanup(patcher.stop)

    # ── request helpers ────────────────────────────────────────────────────
    def get(self, path, user=None, **params):
        return self.client.get(path, params, **hdr(user or self.admin))

    def run_report(self, rid, user=None, **params):
        r = self.get(f"/api/reports/run/{rid}", user, **params)
        self.assertEqual(r.status_code, 200, r.content[:600])
        return r.json()

    def sheet(self, rid, user=None, **params):
        r = self.get(f"/api/reports/export/{rid}", user, fmt="xlsx", **params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertEqual(r.content[:2], b"PK")
        return load_workbook(io.BytesIO(r.content)).active

    def pdf(self, rid, user=None, **params):
        r = self.get(f"/api/reports/export/{rid}", user, fmt="pdf", **params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertTrue(r.content.startswith(b"%PDF"))
        return r.content

    def query_count(self, rid, user=None, **params):
        with CaptureQueriesContext(connection) as cap:
            self.run_report(rid, user, **params)
        return len(cap)


class ContractMixin:
    """The checks every report must pass: 403 without the owning module, empty result on screen and in both
    exports, an Excel round trip, a PDF, and no N+1 as the row count grows."""

    report_id = None
    params: dict = {}
    empty_params: dict = {}
    first_key = "employeeCode"
    first_value = None
    grow_params: dict | None = None  # defaults to ``params``

    def grow(self, start, count):  # create ``count`` more employees' worth of data
        raise NotImplementedError

    def test_role_without_the_owning_module_gets_report_forbidden(self):
        for kind, extra in (("run", {}), ("export", {"fmt": "xlsx"})):
            r = self.get(f"/api/reports/{kind}/{self.report_id}", self.reports_only, **self.params, **extra)
            self.assertEqual(r.status_code, 403, kind)
            self.assertEqual(r.json()["error"], "report_forbidden", kind)
        ids = {x["id"] for x in self.get("/api/reports/catalog", self.reports_only).json()["reports"]}
        self.assertNotIn(self.report_id, ids)

    def test_empty_result_on_screen_and_in_both_exports(self):
        body = self.run_report(self.report_id, **self.empty_params)
        self.assertEqual(data_rows(body), [])
        self.assertFalse(body["truncated"])
        ws = self.sheet(self.report_id, **self.empty_params)
        self.assertEqual(ws["A7"].value, body["columns"][0]["label"])
        self.assertIn(ws["A8"].value, (None, "TOTAL"))  # nothing but an (empty) totals row under the header
        self.pdf(self.report_id, **self.empty_params)

    def test_excel_round_trip_and_pdf(self):
        body = self.run_report(self.report_id, **self.params)
        self.assertTrue(data_rows(body))
        ws = self.sheet(self.report_id, **self.params)
        header = [c.value for c in ws[7]]
        self.assertEqual(header, [c["label"] for c in body["columns"]])
        first = ws["A8"].value
        if isinstance(first, datetime):
            first = first.date().isoformat()
        expected = self.first_value or data_rows(body)[0][self.first_key]
        self.assertEqual(first, expected)
        self.pdf(self.report_id, **self.params)

    def test_runs_with_no_query_string_at_all_on_its_default_filters(self):
        r = self.get(f"/api/reports/run/{self.report_id}")
        self.assertEqual(r.status_code, 200, r.content[:400])
        body = r.json()
        self.assertEqual((body["id"], body["category"], body["truncated"]), (self.report_id, "leave", False))
        keys = {c["key"] for c in body["columns"]}
        for row in body["rows"]:
            self.assertTrue(set(row) - {"_kind"} <= keys)
        for kind in ("xlsx", "pdf"):
            self.assertEqual(self.get(f"/api/reports/export/{self.report_id}", fmt=kind).status_code, 200, kind)

    def test_excel_repeats_the_applied_filters_the_notes_and_the_totals_row(self):
        body = self.run_report(self.report_id, **self.params)
        ws = self.sheet(self.report_id, **self.params)
        column_a = [c.value for c in ws["A"]]
        self.assertTrue(body["notes"], "every report states its assumptions")
        for note in body["notes"]:
            self.assertIn(note, column_a)
        applied = "   |   ".join(f"{f['label']}: {f['value']}" for f in body["filters"]) or "All records"
        self.assertEqual(ws["A4"].value, applied)
        if body["totals"]:
            self.assertIn("TOTAL", column_a)

    def test_query_count_does_not_grow_with_rows(self):
        params = self.grow_params or self.params
        self.grow(0, 3)
        small = self.query_count(self.report_id, **params)
        self.grow(3, 15)
        big = self.query_count(self.report_id, **params)
        self.assertLessEqual(big - small, 3, f"{self.report_id}: {small} queries for 3 extra rows, {big} for 15")


# ══════════════════════════════════════════════════════════════════════════════
#  Registry / catalog
# ══════════════════════════════════════════════════════════════════════════════


class RegistryTests(Base):
    def test_all_ten_reports_are_registered_with_valid_metadata(self):
        registry.all_specs()
        self.assertEqual([k for k in registry.LOAD_ERRORS if "requests_leave" in k], [])
        for rid in MY_IDS:
            spec = registry.get_spec(rid)
            self.assertIsNotNone(spec, rid)
            self.assertEqual(spec.category, "leave", rid)
            self.assertTrue(spec.modules, rid)
            for m in spec.modules:
                self.assertIn(m, all_module_keys(), f"{rid}: unknown module {m}")
            self.assertRegex(spec.id, r"^[a-z0-9-]+$")

    def test_families_have_both_variants(self):
        fam = {}
        for rid in MY_IDS:
            spec = registry.get_spec(rid)
            if spec.family:
                fam.setdefault(spec.family, {})[spec.variant] = rid
        self.assertEqual(set(fam), {"leave", "casual-leave"})
        self.assertEqual(set(fam["leave"].values()), {"leave-register", "leave-summary-monthly"})
        self.assertEqual(set(fam["casual-leave"].values()), {"casual-leave-register", "casual-leave-eligibility-usage"})

    def test_catalog_lists_them_for_admin_and_hides_them_from_a_reports_only_role(self):
        ids = {x["id"] for x in self.get("/api/reports/catalog").json()["reports"]}
        self.assertTrue(set(MY_IDS) <= ids)
        hidden = {x["id"] for x in self.get("/api/reports/catalog", self.reports_only).json()["reports"]}
        self.assertFalse(set(MY_IDS) & hidden)

    def test_leave_only_role_is_offered_leave_reports_but_not_the_others(self):
        ids = {x["id"] for x in self.get("/api/reports/catalog", self.leave_only).json()["reports"]}
        for rid in ("leave-register", "leave-balance-statement", "leave-summary-monthly", "holiday-list"):
            self.assertIn(rid, ids)
        self.assertNotIn("casual-leave-register", ids)
        self.assertNotIn("employee-requests-register", ids)
        # the two blended reports open for any owning module
        self.assertIn("pending-approvals-ageing", ids)
        self.assertIn("approver-workload-turnaround", ids)


# ══════════════════════════════════════════════════════════════════════════════
#  Leave Register
# ══════════════════════════════════════════════════════════════════════════════


class LeaveRegisterTests(ContractMixin, Base):
    report_id = "leave-register"
    params = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
    empty_params = {"dateFrom": "2030-01-01", "dateTo": "2030-01-31"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cl = leave_type("CL", "Casual Leave", True)
        cls.sl = leave_type("SL", "Sick Leave", False)

        def leave(
            emp,
            start,
            end,
            total,
            status="pending",
            ltype=None,
            text="casual",
            half=False,
            slot=None,
            by=None,
            role=None,
            comment=None,
            created=None,
            reason=None,
        ):
            lr = LeaveRequest.objects.create(
                employee=emp,
                leave_type_ref=ltype,
                type=text,
                start_date=start,
                end_date=end,
                total_days=total,
                is_half_day=half,
                half_day_slot=slot,
                status=status,
                approved_by=by,
                approver_role=role,
                hr_comment=comment,
                reason=reason,
            )
            if created:
                stamp(lr, created_at=created)
            return lr

        # RL_A (staff, cutting)
        leave(
            cls.a,
            "2026-08-31",
            "2026-09-01",
            "2.0",
            "approved",
            cls.cl,
            by="Priya HR",
            role="hr",
            created=ist(2026, 8, 25, 9, 30),
        )
        cls.a_sick = leave(
            cls.a,
            "2026-09-07",
            "2026-09-09",
            "3.0",
            "approved",
            cls.sl,
            by="Priya HR",
            role="hr",
            comment="Get well soon",
            created=ist(2026, 9, 1, 10, 0),
            reason="Fever",
        )
        leave(cls.a, "2026-09-08", "2026-09-08", "1.0", "approved", cls.sl, by="Priya HR", role="hr")  # duplicate day
        leave(cls.a, "2026-09-13", "2026-09-13", "1.0", "approved", cls.cl, by="Priya HR", role="hr")  # a Sunday
        leave(
            cls.a,
            "2026-09-15",
            "2026-09-15",
            "0.5",
            "approved",
            cls.cl,
            half=True,
            slot="morning",
            by="Asha Test",
            role="dept_head",
        )
        leave(cls.a, "2026-09-21", "2026-09-22", "2.0", "rejected", cls.cl, by="Priya HR", role="hr")
        # applied at 01:00 IST on the 21st = 19:30 UTC on the 20th: notice must use the IST day
        leave(
            cls.a,
            "2026-09-28",
            "2026-10-03",
            "6.0",
            "approved",
            cls.cl,
            by="Priya HR",
            role="hr",
            created=ist(2026, 9, 21, 1, 0),
        )
        leave(cls.a, "2026-08-10", "2026-08-12", "3.0", "approved", cls.cl, by="Priya HR", role="hr")  # outside
        # RL_B (production)
        leave(cls.b, "2026-09-10", "2026-09-11", "2.0", "approved", cls.cl, by="Priya HR", role="hr")
        leave(cls.b, "15-09-2026", "15-09-2026", "1.0", "pending")  # dd-mm-yyyy text
        # RL_C (staff, other branch)
        leave(cls.c, "2026-09-14", "2026-09-14", "1.0", "pending", text="Sick")
        leave(cls.c, "2026-09-23", "2026-09-23", "1.0", "on_hold")
        cls.c_bad = leave(cls.c, "soon", "soon", "1.0", "approved", cls.cl)
        # RL_D (inactive)
        leave(cls.d, "2026-09-03", "2026-09-03", "1.0", "approved", cls.sl, by="Priya HR", role="hr")

    def test_golden_rows_and_order(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        self.assertEqual(
            [(r["employeeCode"], r["startDate"]) for r in rows],
            [
                ("RL_A", "2026-08-31"),
                ("RL_A", "2026-09-07"),
                ("RL_A", "2026-09-08"),
                ("RL_A", "2026-09-13"),
                ("RL_A", "2026-09-15"),
                ("RL_A", "2026-09-21"),
                ("RL_A", "2026-09-28"),
                ("RL_B", "2026-09-10"),
                ("RL_B", "2026-09-15"),
                ("RL_C", "2026-09-14"),
                ("RL_C", "2026-09-23"),
                ("RL_D", "2026-09-03"),
            ],
        )
        by_start = {(r["employeeCode"], r["startDate"]): r for r in rows}
        cross_in = by_start[("RL_A", "2026-08-31")]
        self.assertEqual((cross_in["totalDays"], cross_in["daysInPeriod"]), (2.0, 1.0))  # only 1 Sep is in September
        sick = by_start[("RL_A", "2026-09-07")]
        self.assertEqual(sick["leaveType"], "Sick Leave")
        self.assertEqual((sick["totalDays"], sick["daysInPeriod"], sick["dayType"]), (3.0, 3.0, "Full day"))
        self.assertEqual((sick["appliedOn"], sick["advanceDays"]), ("2026-09-01 10:00", 6))
        self.assertEqual(
            (sick["approvedBy"], sick["approverRole"], sick["hrComment"]), ("Priya HR", "HR", "Get well soon")
        )
        self.assertEqual((sick["status"], sick["payImpact"], sick["reason"]), ("Approved", "Unpaid (LOP)", "Fever"))
        sunday = by_start[("RL_A", "2026-09-13")]
        self.assertEqual((sunday["totalDays"], sunday["daysInPeriod"]), (1.0, 0.0))  # Sundays are never counted
        half = by_start[("RL_A", "2026-09-15")]
        self.assertEqual((half["dayType"], half["totalDays"], half["daysInPeriod"]), ("Half day (morning)", 0.5, 0.5))
        self.assertEqual(
            (half["approverRole"], half["approvedBy"], half["payImpact"]), ("HOD", "Asha Test", "Half-day leave")
        )
        rejected = by_start[("RL_A", "2026-09-21")]
        self.assertEqual(
            (rejected["status"], rejected["payImpact"], rejected["daysInPeriod"]), ("Rejected", "None", 2.0)
        )
        crossing = by_start[("RL_A", "2026-09-28")]
        self.assertEqual((crossing["totalDays"], crossing["daysInPeriod"]), (6.0, 3.0))  # 28-30 Sep of 28 Sep-3 Oct
        self.assertEqual((crossing["appliedOn"], crossing["advanceDays"]), ("2026-09-21 01:00", 7))
        prod = by_start[("RL_B", "2026-09-10")]
        self.assertEqual(prod["payImpact"], "No effect (production)")
        odd = by_start[("RL_B", "2026-09-15")]
        self.assertEqual(
            (odd["endDate"], odd["status"], odd["payImpact"]), ("2026-09-15", "Pending", "Awaiting decision")
        )
        self.assertEqual(odd["leaveType"], "Casual (untyped)")
        self.assertEqual(by_start[("RL_C", "2026-09-14")]["leaveType"], "Sick (untyped)")
        other = by_start[("RL_C", "2026-09-23")]
        self.assertEqual((other["status"], other["payImpact"]), ("Other", None))

    def test_summary_notes_and_no_totals_row(self):
        body = self.run_report(self.report_id, **self.params)
        summary = {s["label"]: s["value"] for s in body["summary"]}
        self.assertEqual(summary["Leave requests"], 12)
        # RL_A 1,7,8,9,28,29,30 Sep + a half day = 7.5; RL_B 10,11 = 2; RL_D 3 = 1 (the duplicate 8 Sep counts once)
        self.assertEqual(summary["Approved days (in period)"], 10.5)
        self.assertEqual(summary["Pending days (in period)"], 2.0)
        self.assertEqual(summary["Rejected days (in period)"], 2.0)
        self.assertEqual(summary["Half-day requests"], 1)
        self.assertEqual(summary["Unpaid (LOP) days - staff"], 8.5)  # production RL_B excluded
        self.assertEqual(summary["Employees on approved leave"], 3)
        rows = data_rows(body)
        approved_sum = sum(r["daysInPeriod"] for r in rows if r["status"] == "Approved")
        self.assertEqual(approved_sum, 11.5)
        self.assertEqual(approved_sum - 1, summary["Approved days (in period)"])  # the duplicate 8 Sep
        self.assertIsNone(body["totals"])
        notes = " ".join(body["notes"])
        self.assertIn("1 leave request(s) with unreadable start/end dates", notes)
        self.assertIn(f"ids: {self.c_bad.id}", notes)

    def test_filters_narrow_the_result(self):
        def codes(**p):
            return [
                (r["employeeCode"], r["startDate"])
                for r in data_rows(self.run_report(self.report_id, **{**self.params, **p}))
            ]

        self.assertEqual(len(codes(status="approved")), 8)
        self.assertEqual(codes(status="other"), [("RL_C", "2026-09-23")])
        self.assertEqual(len(codes(status="pending")), 2)
        self.assertEqual(codes(dayType="half"), [("RL_A", "2026-09-15")])
        self.assertEqual(len(codes(dayType="full")), 11)
        self.assertEqual(codes(approverRole="dept_head"), [("RL_A", "2026-09-15")])
        self.assertEqual({c for c, _ in codes(employeeIds=str(self.a.id))}, {"RL_A"})
        self.assertEqual(len(codes(employeeIds=str(self.a.id))), 7)
        self.assertEqual(len(codes(departmentIds=str(self.cutting.id))), 10)
        self.assertEqual({c for c, _ in codes(departmentIds=str(self.sewing.id))}, {"RL_C"})
        self.assertEqual({c for c, _ in codes(employmentType="production")}, {"RL_B"})
        self.assertNotIn("RL_D", {c for c, _ in codes(employeeStatus="active")})
        self.assertEqual({c for c, _ in codes(employeeStatus="inactive")}, {"RL_D"})
        self.assertEqual(len(codes(designationIds=str(self.tailor.id))), 7)
        self.assertEqual(sorted(s for _c, s in codes(leaveType="SL")), ["2026-09-03", "2026-09-07", "2026-09-08"])
        self.assertEqual(len(codes(leaveType="sick")), 4)  # SL by name + the untyped 'Sick' row
        self.assertEqual(codes(branchIds=str(self.b2.id)), [("RL_C", "2026-09-14"), ("RL_C", "2026-09-23")])

    def test_date_filter_is_an_overlap_not_a_start_month_match(self):
        october = data_rows(self.run_report(self.report_id, dateFrom="2026-10-01", dateTo="2026-10-31"))
        self.assertEqual([(r["employeeCode"], r["startDate"]) for r in october], [("RL_A", "2026-09-28")])
        self.assertEqual((october[0]["totalDays"], october[0]["daysInPeriod"]), (6.0, 3.0))  # Thu Fri Sat
        one_day = data_rows(self.run_report(self.report_id, dateFrom="2026-09-08", dateTo="2026-09-08"))
        self.assertEqual(
            [(r["startDate"], r["daysInPeriod"]) for r in one_day], [("2026-09-07", 1.0), ("2026-09-08", 1.0)]
        )
        sunday = data_rows(self.run_report(self.report_id, dateFrom="2026-09-13", dateTo="2026-09-13"))
        self.assertEqual([(r["startDate"], r["daysInPeriod"]) for r in sunday], [("2026-09-13", 0.0)])
        august = data_rows(self.run_report(self.report_id, dateFrom="2026-08-01", dateTo="2026-08-31"))
        self.assertEqual([r["startDate"] for r in august], ["2026-08-10", "2026-08-31"])
        self.assertEqual([r["daysInPeriod"] for r in august], [3.0, 1.0])  # 31 Aug of 31 Aug-1 Sep

    def test_branch_isolation(self):
        body = self.run_report(self.report_id, self.branch_user, **self.params)
        codes = {r["employeeCode"] for r in data_rows(body)}
        self.assertEqual(codes, {"RL_A", "RL_B", "RL_D"})
        self.assertNotIn("unreadable", " ".join(body["notes"]))  # RL_C's bad row belongs to the other branch
        widened = self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params)
        self.assertEqual(data_rows(widened), [])
        other = self.run_report(self.report_id, self.branch_user, employeeIds=str(self.c.id), **self.params)
        self.assertEqual(data_rows(other), [])
        two = self.run_report(self.report_id, self.branch_user2, **self.params)
        self.assertEqual({r["employeeCode"] for r in data_rows(two)}, {"RL_C"})

    def test_summary_of_a_filtered_view_follows_the_rows(self):
        body = self.run_report(self.report_id, **{**self.params, "employeeIds": str(self.a.id), "status": "approved"})
        summary = {s["label"]: s["value"] for s in body["summary"]}
        self.assertEqual(summary["Leave requests"], 6)
        self.assertEqual(summary["Approved days (in period)"], 7.5)
        self.assertEqual(summary["Pending days (in period)"], 0.0)
        self.assertEqual(summary["Employees on approved leave"], 1)

    def test_bad_filters_are_400(self):
        r = self.get(f"/api/reports/run/{self.report_id}", dateFrom="2026-09-30", dateTo="2026-09-01")
        self.assertEqual(r.status_code, 400)
        r = self.get(f"/api/reports/run/{self.report_id}", **self.params, status="nonsense")
        self.assertEqual(r.status_code, 400)

    def test_report_is_read_only(self):
        before = LeaveRequest.objects.count(), LeaveBalance.objects.count(), AttendanceDayRecord.objects.count()
        self.run_report(self.report_id, **self.params)
        self.pdf(self.report_id, **self.params)
        self.assertEqual(
            before, (LeaveRequest.objects.count(), LeaveBalance.objects.count(), AttendanceDayRecord.objects.count())
        )

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1)
            LeaveRequest.objects.create(
                employee=emp,
                leave_type_ref=self.cl,
                start_date="2026-09-16",
                end_date="2026-09-16",
                total_days=1,
                status="approved",
                approved_by="Priya HR",
                approver_role="hr",
            )


# ══════════════════════════════════════════════════════════════════════════════
#  Leave Balance Statement
# ══════════════════════════════════════════════════════════════════════════════


class LeaveBalanceTests(ContractMixin, Base):
    report_id = "leave-balance-statement"
    params = {"year": "2026"}
    empty_params = {"year": "2027"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cl = leave_type("CL", "Casual Leave", True)
        cls.sl = leave_type("SL", "Sick Leave", False)
        cls.e = mk_emp("RL_E", cls.cutting, cls.b1)  # no balance and no leave at all

        def bal(emp, lt, year, alloc, used, remaining, carried="0"):
            return LeaveBalance.objects.create(
                employee=emp,
                leave_type=lt,
                year=year,
                allocated=alloc,
                used=used,
                remaining=remaining,
                carried_forward=carried,
            )

        bal(cls.a, cls.cl, 2026, "12", "3", "9", "2")
        bal(cls.a, cls.sl, 2026, "8", "0", "8")
        bal(cls.c, cls.cl, 2026, "12", "1.5", "10.5")
        bal(cls.b, cls.cl, 2026, "12", "0", "12")
        bal(cls.d, cls.cl, 2026, "12", "0", "12")
        bal(cls.a, cls.cl, 2025, "12", "12", "0")

        def leave(emp, start, end, total, status, ltype=None, text="casual", half=False, slot=None):
            return LeaveRequest.objects.create(
                employee=emp,
                leave_type_ref=ltype,
                type=text,
                start_date=start,
                end_date=end,
                total_days=total,
                status=status,
                is_half_day=half,
                half_day_slot=slot,
            )

        leave(cls.a, "2026-03-02", "2026-03-04", "3.0", "approved", cls.cl)
        leave(cls.a, "2026-04-06", "2026-04-06", "0.5", "approved", cls.cl, half=True, slot="morning")
        leave(cls.a, "2026-05-04", "2026-05-04", "1.0", "pending", cls.cl)
        leave(cls.a, "2026-06-01", "2026-06-01", "1.0", "approved", cls.sl)
        leave(cls.a, "2026-07-01", "2026-07-02", "2.0", "approved")  # untyped, text 'casual'
        leave(cls.a, "2025-12-29", "2026-01-02", "5.0", "approved", cls.cl)  # starts in 2025
        leave(cls.a, "bad", "bad", "1.0", "approved", cls.cl)
        leave(cls.a, "2026-08-03", "2026-08-03", "1.0", "rejected", cls.cl)  # rejected: never counted
        leave(cls.c, "2026-02-02", "2026-02-02", "1.0", "approved", cls.cl)
        leave(cls.c, "2026-07-06", "2026-07-25", "18.0", "approved", cls.cl)
        leave(cls.b, "10-08-2026", "10-08-2026", "1.0", "approved", cls.cl)  # dd-mm-yyyy text
        leave(cls.d, "2026-01-05", "2026-01-05", "1.0", "approved", cls.cl)
        for emp, dt, st in (
            (cls.a, date(2026, 3, 10), "approved"),
            (cls.a, date(2026, 5, 12), "approved"),
            (cls.a, date(2026, 6, 9), "pending"),
            (cls.c, date(2026, 4, 14), "approved"),
        ):
            CasualLeaveRequest.objects.create(employee=emp, date=dt, status=st)

    def rows_by_key(self, body):
        return {(r["employeeCode"], r["leaveType"]): r for r in data_rows(body)}

    def test_golden_rows(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        self.assertEqual(
            [(r["employeeCode"], r["leaveType"]) for r in rows],
            [
                ("RL_A", "Casual (untyped)"),
                ("RL_A", "Casual Leave"),
                ("RL_A", "Casual Leave (CL system)"),
                ("RL_A", "Sick Leave"),
                ("RL_B", "Casual Leave"),
                ("RL_C", "Casual Leave"),
                ("RL_C", "Casual Leave (CL system)"),
            ],
        )
        r = self.rows_by_key(body)
        a_cl = r[("RL_A", "Casual Leave")]
        self.assertEqual(
            (a_cl["leaveCode"], a_cl["isPaid"], a_cl["allocated"], a_cl["carriedForward"], a_cl["ledgerUsed"]),
            ("CL", "Paid", 12.0, 2.0, 3.0),
        )
        self.assertEqual(
            (a_cl["approvedDays"], a_cl["pendingDays"]), (3.5, 1.0)
        )  # 3.0 + a half day; 2025-start excluded
        self.assertEqual((a_cl["ledgerRemaining"], a_cl["computedRemaining"], a_cl["variance"]), (9.0, 8.5, -0.5))
        a_sl = r[("RL_A", "Sick Leave")]
        self.assertEqual(
            (a_sl["isPaid"], a_sl["allocated"], a_sl["ledgerUsed"], a_sl["approvedDays"]), ("Unpaid", 8.0, 0.0, 1.0)
        )
        self.assertEqual((a_sl["computedRemaining"], a_sl["variance"]), (7.0, -1.0))
        untyped = r[("RL_A", "Casual (untyped)")]
        self.assertEqual(untyped["approvedDays"], 2.0)
        for key in (
            "allocated",
            "ledgerUsed",
            "ledgerRemaining",
            "computedRemaining",
            "variance",
            "leaveCode",
            "isPaid",
        ):
            self.assertIsNone(untyped[key], key)
        a_sys = r[("RL_A", "Casual Leave (CL system)")]
        self.assertEqual(
            (
                a_sys["allocated"],
                a_sys["approvedDays"],
                a_sys["pendingDays"],
                a_sys["computedRemaining"],
                a_sys["ledgerUsed"],
            ),
            (12.0, 2.0, 1.0, 10.0, None),
        )
        b_cl = r[("RL_B", "Casual Leave")]  # dd-mm-yyyy text date still counts, in 2026
        self.assertEqual((b_cl["approvedDays"], b_cl["computedRemaining"], b_cl["variance"]), (1.0, 11.0, -1.0))
        c_cl = r[("RL_C", "Casual Leave")]
        self.assertEqual(
            (c_cl["ledgerUsed"], c_cl["approvedDays"], c_cl["computedRemaining"], c_cl["variance"]),
            (1.5, 19.0, -7.0, -17.5),
        )
        self.assertEqual((r[("RL_C", "Casual Leave (CL system)")]["approvedDays"]), 1.0)

    def test_totals_summary_and_notes_agree_with_the_rows(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        t = body["totals"]
        self.assertEqual(t["allocated"], round(sum(r["allocated"] or 0 for r in rows), 2))
        self.assertEqual(t["allocated"], 68.0)
        self.assertEqual(t["approvedDays"], 29.5)
        self.assertEqual(t["pendingDays"], 2.0)
        self.assertEqual(t["computedRemaining"], 40.5)
        summary = {s["label"]: s["value"] for s in body["summary"]}
        self.assertEqual(summary["Employees with an allocation"], 3)
        self.assertEqual(summary["Days allocated"], t["allocated"])
        self.assertEqual(summary["Approved days taken"], t["approvedDays"])
        self.assertEqual(summary["Remaining (computed)"], t["computedRemaining"])
        self.assertEqual(summary["Employees over-drawn"], 1)
        self.assertEqual(summary["Ledger mismatches"], 4)
        notes = " ".join(body["notes"])
        self.assertIn("1 employee(s) in this selection have no allocation and no leave in 2026", notes)  # RL_E
        self.assertIn("unreadable dates are not counted", notes)
        self.assertIn("fixed yearly entitlement of 12", notes)

    def test_filters(self):
        def keys(**p):
            return [
                (r["employeeCode"], r["leaveType"])
                for r in data_rows(self.run_report(self.report_id, **{**self.params, **p}))
            ]

        self.assertEqual(len(keys(includeCasual="false")), 5)
        self.assertNotIn("Casual Leave (CL system)", {k[1] for k in keys(includeCasual="false")})
        self.assertEqual(keys(leaveType="SL"), [("RL_A", "Sick Leave")])
        self.assertEqual(len(keys(leaveType="CL")), 5)
        self.assertEqual({k[0] for k in keys(departmentIds=str(self.sewing.id))}, {"RL_C"})
        self.assertEqual({k[0] for k in keys(employmentType="production")}, {"RL_B"})
        self.assertEqual({k[0] for k in keys(employeeIds=str(self.a.id))}, {"RL_A"})
        with_inactive = keys(employeeStatus="all")
        self.assertIn(("RL_D", "Casual Leave"), with_inactive)
        self.assertEqual(len(with_inactive), 8)
        self.assertEqual(keys(employeeStatus="inactive"), [("RL_D", "Casual Leave")])
        # only rows needing attention: ledger differs from approved leave, or overdrawn
        self.assertEqual(
            keys(onlyMismatch="true"),
            [("RL_A", "Casual Leave"), ("RL_A", "Sick Leave"), ("RL_B", "Casual Leave"), ("RL_C", "Casual Leave")],
        )

    def test_designation_and_branch_filters(self):
        def keys(**p):
            return [
                (r["employeeCode"], r["leaveType"])
                for r in data_rows(self.run_report(self.report_id, **{**self.params, **p}))
            ]

        tailors = keys(designationIds=str(self.tailor.id))
        self.assertEqual({k[0] for k in tailors}, {"RL_A"})
        self.assertEqual(len(tailors), 4)
        self.assertEqual(
            keys(branchIds=str(self.b2.id)), [("RL_C", "Casual Leave"), ("RL_C", "Casual Leave (CL system)")]
        )
        self.assertEqual({k[0] for k in keys(branchIds=str(self.b1.id))}, {"RL_A", "RL_B"})

    def test_previous_year_uses_only_requests_that_start_in_that_year(self):
        body = self.run_report(self.report_id, year="2025")
        rows = data_rows(body)
        self.assertEqual([(r["employeeCode"], r["leaveType"]) for r in rows], [("RL_A", "Casual Leave")])
        r = rows[0]
        self.assertEqual(
            (r["ledgerUsed"], r["approvedDays"], r["computedRemaining"], r["variance"]), (12.0, 5.0, 7.0, 7.0)
        )

    def test_no_allocations_at_all_says_so(self):
        body = self.run_report(self.report_id, year="2027")
        self.assertEqual(data_rows(body), [])
        self.assertTrue(body["notes"][0].startswith("No leave balances have been allocated for 2027"))

    def test_branch_isolation_and_casual_module_gate(self):
        body = self.run_report(self.report_id, self.branch_user, **self.params)
        self.assertEqual({r["employeeCode"] for r in data_rows(body)}, {"RL_A", "RL_B"})
        self.assertEqual(
            data_rows(self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params)), []
        )
        no_cl = self.run_report(self.report_id, self.leave_only, **self.params)
        self.assertNotIn("Casual Leave (CL system)", {r["leaveType"] for r in data_rows(no_cl)})
        self.assertTrue(any("does not have access to the Casual Leave module" in n for n in no_cl["notes"]))

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1)
            LeaveBalance.objects.create(employee=emp, leave_type=self.cl, year=2026, allocated=12, used=1, remaining=11)
            LeaveRequest.objects.create(
                employee=emp,
                leave_type_ref=self.cl,
                start_date="2026-05-05",
                end_date="2026-05-05",
                total_days=1,
                status="approved",
            )
            CasualLeaveRequest.objects.create(employee=emp, date=date(2026, 6, 10), status="approved")


# ══════════════════════════════════════════════════════════════════════════════
#  Monthly Leave Summary
# ══════════════════════════════════════════════════════════════════════════════


class LeaveMonthlyTests(ContractMixin, Base):
    report_id = "leave-summary-monthly"
    params = {"year": "2026"}
    empty_params = {"year": "2031"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cl = leave_type("CL", "Casual Leave", True)
        cls.sl = leave_type("SL", "Sick Leave", False)

        def leave(emp, start, end, ltype=None, status="approved", text="casual", half=False, slot=None):
            return LeaveRequest.objects.create(
                employee=emp,
                leave_type_ref=ltype,
                type=text,
                start_date=start,
                end_date=end,
                total_days=1,
                status=status,
                is_half_day=half,
                half_day_slot=slot,
            )

        leave(cls.a, "2026-01-29", "2026-02-03", cls.cl)  # Thu Fri Sat + Mon Tue (Sunday 1 Feb skipped)
        leave(cls.a, "2026-01-30", "2026-01-30", cls.cl)  # duplicate day
        leave(cls.a, "2025-12-30", "2026-01-02", cls.cl)  # only 1-2 Jan fall in 2026
        leave(cls.a, "2026-03-04", "2026-03-05", cls.sl)
        leave(cls.a, "2026-03-06", "2026-03-06", cls.cl, half=True, slot="morning")
        leave(cls.a, "2025-11-03", "2025-11-03", cls.cl)
        leave(cls.b, "2026-05-11", "2026-05-12", cls.sl)
        leave(cls.c, "2026-12-24", "2026-12-26", cls.cl)
        leave(cls.c, "2026-06-01", "2026-06-01", cls.cl, status="pending")
        leave(cls.c, "2026-06-02", "2026-06-02", cls.cl, status="rejected")
        cls.c_bad = leave(cls.c, "xx", "xx", cls.cl)
        leave(cls.d, "2026-07-06", "2026-07-06")  # untyped 'casual'
        for emp, dt, st in (
            (cls.a, date(2026, 2, 10), "approved"),
            (cls.a, date(2026, 8, 11), "approved"),
            (cls.c, date(2026, 8, 12), "approved"),
            (cls.a, date(2026, 9, 9), "pending"),
        ):
            CasualLeaveRequest.objects.create(employee=emp, date=dt, status=st)

    def keyed(self, body):
        return {(r["employeeCode"], r["leaveType"]): r for r in data_rows(body)}

    def test_golden_grid_subtotals_and_totals(self):
        body = self.run_report(self.report_id, **self.params)
        rows = body["rows"]
        self.assertEqual(
            [(r.get("employeeCode"), r["employeeName"] if r.get("_kind") else r["leaveType"]) for r in rows],
            [
                ("RL_A", "Casual Leave"),
                ("RL_A", "Casual Leave (CL system)"),
                ("RL_A", "Sick Leave"),
                ("RL_B", "Sick Leave"),
                ("RL_D", "Casual (untyped)"),
                (None, "CUTTING total"),
                ("RL_C", "Casual Leave"),
                ("RL_C", "Casual Leave (CL system)"),
                (None, "SEWING total"),
            ],
        )
        k = self.keyed(body)

        def months(row):
            return {
                m: row[m]
                for m in ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
                if row[m]
            }

        a_cl = k[("RL_A", "Casual Leave")]
        self.assertEqual(
            months(a_cl), {"jan": 5.0, "feb": 2.0, "mar": 0.5}
        )  # duplicate day and Sunday not double counted
        self.assertEqual(a_cl["total"], 7.5)
        self.assertEqual(months(k[("RL_A", "Casual Leave (CL system)")]), {"feb": 1.0, "aug": 1.0})
        self.assertEqual(months(k[("RL_A", "Sick Leave")]), {"mar": 2.0})
        self.assertEqual(months(k[("RL_B", "Sick Leave")]), {"may": 2.0})
        self.assertEqual(months(k[("RL_D", "Casual (untyped)")]), {"jul": 1.0})
        self.assertEqual(months(k[("RL_C", "Casual Leave")]), {"dec": 3.0})
        self.assertEqual(months(k[("RL_C", "Casual Leave (CL system)")]), {"aug": 1.0})
        self.assertIsNone(a_cl["apr"])  # a dash, not a zero

        cutting = rows[5]
        self.assertEqual(cutting["_kind"], "subtotal")
        self.assertEqual(months(cutting), {"jan": 5.0, "feb": 3.0, "mar": 2.5, "may": 2.0, "jul": 1.0, "aug": 1.0})
        self.assertEqual(cutting["total"], 14.5)
        sewing = rows[8]
        self.assertEqual((sewing["_kind"], sewing["aug"], sewing["dec"], sewing["total"]), ("subtotal", 1.0, 3.0, 4.0))

        totals = body["totals"]  # subtotal rows are excluded from the totals
        self.assertEqual(totals["total"], 18.5)
        self.assertEqual(
            (totals["jan"], totals["feb"], totals["mar"], totals["aug"], totals["dec"]), (5.0, 3.0, 2.5, 2.0, 3.0)
        )
        self.assertEqual(totals["total"], round(sum(r["total"] for r in data_rows(body)), 2))
        for r in data_rows(body):
            self.assertEqual(
                r["total"],
                round(
                    sum(
                        r[m] or 0
                        for m in ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
                    ),
                    2,
                ),
            )

    def test_summary_and_notes(self):
        body = self.run_report(self.report_id, **self.params)
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Leave days in the year"], 18.5)
        self.assertEqual(s["Employees with leave"], 4)
        self.assertEqual(s["Average days per employee"], round(18.5 / 4, 2))
        self.assertEqual(s["Department with most leave"], "CUTTING (14.5 days)")
        self.assertTrue(any("1 approved leave request(s) with unreadable dates" in n for n in body["notes"]))

    def test_filters(self):
        def total(**p):
            body = self.run_report(self.report_id, **{**self.params, **p})
            return body["totals"]["total"] if body["totals"] else 0

        self.assertEqual(total(includeCasual="false"), 15.5)
        self.assertEqual(total(leaveType="SL"), 4.0)
        self.assertEqual(total(departmentIds=str(self.sewing.id)), 4.0)
        self.assertEqual(total(employeeStatus="active"), 17.5)  # without inactive RL_D
        self.assertEqual(total(employmentType="production"), 2.0)
        self.assertEqual(total(employeeIds=str(self.a.id)), 11.5)
        self.assertEqual(total(designationIds=str(self.tailor.id)), 11.5)  # only RL_A is a Tailor
        self.assertEqual(total(branchIds=str(self.b2.id)), 4.0)  # RL_C, Unit 2
        self.assertEqual(total(branchIds=str(self.b1.id)), 14.5)  # RL_A, RL_B, RL_D, Unit 1

    def test_previous_year_split_of_a_cross_year_request(self):
        body = self.run_report(self.report_id, year="2025")
        row = data_rows(body)[0]
        self.assertEqual((row["employeeCode"], row["leaveType"]), ("RL_A", "Casual Leave"))
        self.assertEqual((row["nov"], row["dec"], row["total"]), (1.0, 2.0, 3.0))  # 3 Nov, plus 30-31 Dec

    def test_branch_isolation_and_casual_module_gate(self):
        body = self.run_report(self.report_id, self.branch_user, **self.params)
        self.assertEqual({r["employeeCode"] for r in data_rows(body)}, {"RL_A", "RL_B", "RL_D"})
        self.assertNotIn("unreadable", " ".join(body["notes"]))
        self.assertEqual(
            data_rows(self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params)), []
        )
        no_cl = self.run_report(self.report_id, self.leave_only, **self.params)
        self.assertNotIn("Casual Leave (CL system)", {r["leaveType"] for r in data_rows(no_cl)})
        self.assertEqual(no_cl["totals"]["total"], 15.5)

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1)
            LeaveRequest.objects.create(
                employee=emp,
                leave_type_ref=self.sl,
                start_date="2026-04-13",
                end_date="2026-04-14",
                total_days=2,
                status="approved",
            )
            CasualLeaveRequest.objects.create(employee=emp, date=date(2026, 6, 10), status="approved")


# ══════════════════════════════════════════════════════════════════════════════
#  Daily Leave Load by Department
# ══════════════════════════════════════════════════════════════════════════════


class LeaveCalendarTests(ContractMixin, Base):
    report_id = "leave-department-calendar"
    params = {"dateFrom": "2026-09-07", "dateTo": "2026-09-13"}
    empty_params = {"dateFrom": "2026-09-07", "dateTo": "2026-09-13", "departmentIds": "999999"}
    first_key = "date"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a2 = mk_emp("RL_A2", cls.cutting, cls.b1, join_date="2025-03-03")
        cls.a3 = mk_emp("RL_A3", cls.cutting, cls.b1, join_date="2026-09-11")  # joins mid-range
        Holiday.objects.create(name="Ganesh Chaturthi", date=date(2026, 9, 10))

        def leave(emp, start, end, half=False, slot=None):
            LeaveRequest.objects.create(
                employee=emp,
                type="casual",
                start_date=start,
                end_date=end,
                total_days=1,
                status="approved",
                is_half_day=half,
                half_day_slot=slot,
            )

        leave(cls.a, "2026-09-07", "2026-09-08")
        leave(cls.a2, "2026-09-09", "2026-09-09", half=True, slot="morning")
        leave(cls.a2, "2026-09-08", "2026-09-08")
        leave(cls.b, "2026-09-11", "2026-09-11")
        LeaveRequest.objects.create(  # pending: never counted
            employee=cls.a3,
            type="casual",
            start_date="2026-09-07",
            end_date="2026-09-07",
            total_days=1,
            status="pending",
        )
        CasualLeaveRequest.objects.create(employee=cls.c, date=date(2026, 9, 7), status="approved")
        CasualLeaveRequest.objects.create(employee=cls.a2, date=date(2026, 9, 8), status="approved")  # full leave wins
        # on duty: 00:30 IST on the 9th is still the 8th in UTC
        s = OnDutySession.objects.create(employee=cls.a, destination="Yarn mill", branch=cls.b1, status="active")
        stamp(s, created_at=ist(2026, 9, 9, 0, 30))
        s2 = OnDutySession.objects.create(employee=cls.b, destination="Dye house", branch=cls.b1, status="rejected")
        stamp(s2, created_at=ist(2026, 9, 9, 10, 0))
        s3 = OnDutySession.objects.create(employee=cls.a2, destination="Bank", branch=cls.b1, status="pending_hod")
        stamp(s3, created_at=ist(2026, 9, 9, 11, 0))
        EmployeePermission.objects.create(
            employee=cls.a, date=date(2026, 9, 7), status="approved", type="morning_late_in"
        )
        EmployeePermission.objects.create(employee=cls.a2, date=date(2026, 9, 7), status="pending")
        for emp, dt in (
            (cls.b, date(2026, 9, 8)),
            (cls.a2, date(2026, 9, 8)),
            (cls.c, date(2026, 9, 9)),
            (cls.a, date(2026, 9, 10)),
        ):
            AttendanceDayRecord.objects.create(employee=emp, date=dt, status="absent")

    def by_day(self, body, dept):
        return {r["date"]: r for r in data_rows(body) if r["department"] == dept}

    def test_golden_numbers_for_cutting(self):
        body = self.run_report(self.report_id, **self.params)
        cut = self.by_day(body, "CUTTING")
        expected = {
            # date: (dayType, strength, full, half, cl, duty, perm, absent, pct)
            "2026-09-07": ("Working day", 3, 1, 0, 0, 0, 1, 0, 33.3),
            "2026-09-08": ("Working day", 3, 2, 0, 0, 0, 0, 1, 66.7),  # RL_A2's absence is explained by leave
            "2026-09-09": ("Working day", 3, 0, 1, 0, 1, 0, 0, 16.7),
            "2026-09-10": ("Holiday", 3, None, None, None, None, None, None, None),
            "2026-09-11": ("Working day", 4, 1, 0, 0, 0, 0, 0, 25.0),  # RL_A3 has joined
            "2026-09-12": ("Working day", 4, 0, 0, 0, 0, 0, 0, 0.0),
            "2026-09-13": ("Sunday", 4, 0, 0, 0, 0, 0, 0, 0.0),  # only production RL_B works on a Sunday
        }
        self.assertEqual(set(cut), set(expected))
        for day, (dtype, strength, full, half, cl, duty, perm, absent, pct) in expected.items():
            r = cut[day]
            got = (
                r["dayType"],
                r["strength"],
                r["onLeaveFull"],
                r["onLeaveHalf"],
                r["casualLeave"],
                r["onDuty"],
                r["permission"],
                r["absentUnexplained"],
                r["leavePct"],
            )
            want = (
                dtype,
                strength,
                *(None if v is None else float(v) for v in (full, half, cl, duty, perm, absent)),
                pct,
            )
            self.assertEqual(got, want, day)
        self.assertEqual(cut["2026-09-07"]["weekday"], "Mon")

    def test_golden_numbers_for_sewing_and_row_order(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        self.assertEqual(len(rows), 14)
        self.assertEqual(
            [(r["date"], r["department"]) for r in rows][:4],
            [("2026-09-07", "CUTTING"), ("2026-09-07", "SEWING"), ("2026-09-08", "CUTTING"), ("2026-09-08", "SEWING")],
        )
        sew = self.by_day(body, "SEWING")
        mon = sew["2026-09-07"]
        self.assertEqual((mon["strength"], mon["casualLeave"], mon["leavePct"]), (1, 1.0, 100.0))
        self.assertEqual(sew["2026-09-09"]["absentUnexplained"], 1.0)
        self.assertIsNone(sew["2026-09-10"]["leavePct"])  # holiday
        self.assertEqual(sew["2026-09-13"]["dayType"], "Sunday")
        self.assertIsNone(sew["2026-09-13"]["casualLeave"])  # every staff member is on the weekly off

    def test_summary(self):
        body = self.run_report(self.report_id, **self.params)
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Days in range"], 7)
        self.assertEqual(s["Departments"], 2)
        self.assertEqual(s["Peak leave day"], "07-Sep-2026 (2 on leave)")  # 7 and 8 Sep tie; earliest wins
        self.assertEqual(s["Average daily leave %"], 22.1)  # (50 + 50 + 12.5 + 20 + 0 + 0) / 6 days that had staff on
        self.assertIsNone(body["totals"])

    def test_filters_and_branch_isolation(self):
        sewing = data_rows(self.run_report(self.report_id, departmentIds=str(self.sewing.id), **self.params))
        self.assertEqual({r["department"] for r in sewing}, {"SEWING"})
        prod = data_rows(self.run_report(self.report_id, employmentType="production", **self.params))
        self.assertEqual({r["department"] for r in prod}, {"CUTTING"})
        self.assertEqual({r["strength"] for r in prod if r["dayType"] != "Holiday"}, {1})
        b1 = data_rows(self.run_report(self.report_id, self.branch_user, **self.params))
        self.assertEqual({r["department"] for r in b1}, {"CUTTING"})
        widened = data_rows(self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params))
        self.assertEqual(widened, [])
        unit2 = data_rows(self.run_report(self.report_id, branchIds=str(self.b2.id), **self.params))
        self.assertEqual({r["department"] for r in unit2}, {"SEWING"})
        r = self.get(f"/api/reports/run/{self.report_id}", dateFrom="2026-09-01", dateTo="2026-10-15")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["field"], "dateTo")

    def test_modules_the_role_cannot_open_leave_their_columns_blank(self):
        body = self.run_report(self.report_id, self.leave_only, **self.params)
        rows = data_rows(body)
        for r in rows:
            for key in ("casualLeave", "onDuty", "permission", "absentUnexplained"):
                self.assertIsNone(r[key], key)
        mon = self.by_day(body, "CUTTING")["2026-09-07"]
        self.assertEqual((mon["onLeaveFull"], mon["leavePct"]), (1.0, 33.3))
        sew_mon = self.by_day(body, "SEWING")["2026-09-07"]
        self.assertEqual(sew_mon["leavePct"], 0.0)  # casual leave is not visible to this role
        self.assertTrue(any("Columns left blank" in n for n in body["notes"]))

    def test_same_department_name_in_two_branches_is_told_apart(self):
        d1 = Department.objects.create(name="ADMIN", branch=self.b1)
        d2 = Department.objects.create(name="ADMIN", branch=self.b2)
        mk_emp("RL_X1", d1, self.b1)
        mk_emp("RL_X2", d2, self.b2)
        body = self.run_report(self.report_id, **self.params)
        labels = {r["department"] for r in data_rows(body)}
        self.assertTrue({"ADMIN (Unit 1)", "ADMIN (Unit 2)"} <= labels)

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1, join_date="2024-01-01")
            LeaveRequest.objects.create(
                employee=emp,
                type="casual",
                start_date="2026-09-08",
                end_date="2026-09-09",
                total_days=2,
                status="approved",
            )
            CasualLeaveRequest.objects.create(employee=emp, date=date(2026, 9, 11), status="approved")
            EmployeePermission.objects.create(employee=emp, date=date(2026, 9, 12), status="approved")
            AttendanceDayRecord.objects.create(employee=emp, date=date(2026, 9, 7), status="absent")


# ══════════════════════════════════════════════════════════════════════════════
#  Holiday Calendar
# ══════════════════════════════════════════════════════════════════════════════


class HolidayListTests(ContractMixin, Base):
    report_id = "holiday-list"
    params = {"year": "2026"}
    empty_params = {"year": "2030"}
    first_key = "date"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        mk = lambda name, d, htype="national", branch=None, dept=None, rec=False, desc=None: Holiday.objects.create(  # noqa: E731
            name=name, date=d, holiday_type=htype, branch=branch, department=dept, is_recurring=rec, description=desc
        )
        mk("Republic Day", date(2026, 1, 26), rec=True, desc="Flag hoisting")
        mk("Republic Day (Unit 2)", date(2026, 1, 26), branch=cls.b2)
        mk("Unit 1 Foundation Day", date(2026, 3, 5), "company", branch=cls.b1)
        mk("Unit 2 Anniversary", date(2026, 4, 10), "regional", branch=cls.b2)
        mk("Sewing Shutdown", date(2026, 5, 15), "company", dept=cls.sewing)  # company-wide but tied to Unit 2's dept
        mk("Independence Special", date(2026, 8, 16), "company")  # a Sunday
        mk("Diwali", date(2026, 11, 8), "company")  # a Sunday
        mk("Christmas", date(2025, 12, 25))

    def names(self, user=None, **p):
        return [r["name"] for r in data_rows(self.run_report(self.report_id, user, **{**self.params, **p}))]

    def test_golden_rows(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        self.assertEqual(
            [r["name"] for r in rows],
            [
                "Republic Day",
                "Republic Day (Unit 2)",
                "Unit 1 Foundation Day",
                "Unit 2 Anniversary",
                "Sewing Shutdown",
                "Independence Special",
                "Diwali",
            ],
        )
        first = rows[0]
        self.assertEqual((first["date"], first["weekday"], first["holidayType"]), ("2026-01-26", "Monday", "National"))
        self.assertEqual(
            (first["scope"], first["department"], first["isRecurring"], first["description"]),
            ("All branches", "All departments", "Yes", "Flag hoisting"),
        )
        self.assertEqual((rows[1]["scope"], rows[1]["isRecurring"]), ("Unit 2", "No"))
        self.assertEqual((rows[4]["scope"], rows[4]["department"]), ("All branches", "SEWING"))
        self.assertEqual([r["onSunday"] for r in rows], ["No", "No", "No", "No", "No", "Yes", "Yes"])
        self.assertEqual(rows[5]["weekday"], "Sunday")

    def test_summary_and_notes(self):
        body = self.run_report(self.report_id, **self.params)
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Holidays listed"], 7)
        self.assertEqual(s["Distinct dates"], 6)  # 26 Jan appears twice
        self.assertEqual((s["National"], s["Regional"], s["Company"]), (2, 1, 4))
        self.assertEqual(s["Fall on a Sunday"], 2)
        self.assertEqual(s["Working-day holidays"], 4)
        self.assertEqual(s["National"] + s["Regional"] + s["Company"], s["Holidays listed"])
        self.assertEqual(s["Fall on a Sunday"] + s["Working-day holidays"], s["Distinct dates"])
        self.assertTrue(any("1 date(s) have more than one holiday row" in n for n in body["notes"]))

    def test_filters(self):
        self.assertEqual(self.names(year="2025"), ["Christmas"])
        self.assertEqual(self.names(holidayType="national"), ["Republic Day", "Republic Day (Unit 2)"])
        self.assertEqual(self.names(holidayType="regional"), ["Unit 2 Anniversary"])
        # a branch filter keeps that branch's rows plus the company-wide ones
        self.assertEqual(
            self.names(branchIds=str(self.b1.id)),
            ["Republic Day", "Unit 1 Foundation Day", "Sewing Shutdown", "Independence Special", "Diwali"],
        )
        self.assertEqual(len(self.names(departmentIds=str(self.sewing.id))), 7)
        self.assertEqual(len(self.names(departmentIds=str(self.cutting.id))), 6)  # not the SEWING-only shutdown

    def test_branch_isolation_keeps_company_wide_rows_but_not_other_branches(self):
        self.assertEqual(
            self.names(self.branch_user),
            ["Republic Day", "Unit 1 Foundation Day", "Independence Special", "Diwali"],
        )
        self.assertEqual(
            self.names(self.branch_user2),
            [
                "Republic Day",
                "Republic Day (Unit 2)",
                "Unit 2 Anniversary",
                "Sewing Shutdown",
                "Independence Special",
                "Diwali",
            ],
        )
        self.assertEqual(
            self.names(self.branch_user, branchIds=str(self.b2.id)), ["Republic Day", "Independence Special", "Diwali"]
        )

    def grow(self, start, count):
        for i in range(start, start + count):
            Holiday.objects.create(name=f"Extra {i}", date=date(2026, 6, 1 + i % 25), holiday_type="company")


# ══════════════════════════════════════════════════════════════════════════════
#  Casual Leave Records
# ══════════════════════════════════════════════════════════════════════════════


class CasualLeaveRegisterTests(ContractMixin, Base):
    report_id = "casual-leave-register"
    params = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
    empty_params = {"dateFrom": "2030-09-01", "dateTo": "2030-09-30"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.e = mk_emp("RL_E", cls.cutting, cls.b1, join_date="2026-08-01")
        cls.f = mk_emp("RL_F", cls.cutting, cls.b1, join_date="2025-05-05")

        def cl(emp, dt, status, by=None, role=None, created=None, reviewed=None, reason=None, comment=None):
            r = CasualLeaveRequest.objects.create(
                employee=emp,
                date=dt,
                status=status,
                reviewed_by=by,
                reviewer_role=role,
                reviewed_at=reviewed,
                reason=reason,
                review_comment=comment,
            )
            if created:
                stamp(r, created_at=created)
            return r

        def rec(emp, dt, status, note):
            AttendanceDayRecord.objects.create(
                employee=emp, date=dt, status=status, override_note=note, source="manual"
            )

        cl(
            cls.a,
            date(2026, 9, 7),
            "approved",
            "Priya HR",
            "hr",
            ist(2026, 9, 3, 10),
            ist(2026, 9, 4, 16, 30),
            reason="Family function",
            comment="Enjoy",
        )
        rec(cls.a, date(2026, 9, 7), "present", "Casual Leave (paid) -approved")
        cl(cls.a, date(2026, 8, 10), "approved", "Priya HR", "hr", ist(2026, 8, 5, 9), ist(2026, 8, 6, 9))
        cl(cls.c, date(2026, 9, 14), "rejected", "Asha Test", "dept_head", ist(2026, 9, 10, 9), ist(2026, 9, 10, 11))
        rec(cls.c, date(2026, 9, 14), "on_leave", "Casual Leave rejected -marked as leave")
        cl(cls.b, date(2026, 9, 16), "pending", created=ist(2026, 9, 12, 9))
        cl(cls.d, date(2026, 9, 21), "approved", "Priya HR", "hr", ist(2026, 9, 18, 9), ist(2026, 9, 19, 9))
        cl(cls.f, date(2026, 9, 28), "approved", "Priya HR", "hr", ist(2026, 9, 25, 9), ist(2026, 9, 26, 10, 30))
        rec(cls.f, date(2026, 9, 28), "present", "HR override")  # later overwritten by an HR entry
        cl(cls.e, date(2026, 9, 30), "pending", created=ist(2026, 9, 28, 9))

    def test_golden_rows(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        self.assertEqual(
            [(r["clDate"], r["employeeCode"]) for r in rows],
            [
                ("2026-09-07", "RL_A"),
                ("2026-09-14", "RL_C"),
                ("2026-09-16", "RL_B"),
                ("2026-09-21", "RL_D"),
                ("2026-09-28", "RL_F"),
                ("2026-09-30", "RL_E"),
            ],
        )
        a, c, b, d, f, e = rows
        self.assertEqual((a["weekday"], a["joinDate"], a["serviceMonths"]), ("Mon", "2025-01-10", 19))
        self.assertEqual(
            (a["appliedOn"], a["reviewedAt"], a["turnaroundHours"]), ("2026-09-03 10:00", "2026-09-04 16:30", 30.5)
        )
        self.assertEqual(
            (a["status"], a["reviewedBy"], a["reviewerRole"], a["reason"], a["reviewComment"]),
            ("Approved", "Priya HR", "HR", "Family function", "Enjoy"),
        )
        self.assertEqual(a["attendanceOutcome"], "Paid present")
        self.assertEqual(
            (c["joinDate"], c["serviceMonths"], c["reviewerRole"], c["turnaroundHours"]), ("2024-06-15", 26, "HOD", 2.0)
        )  # 15/06/2024 read as a date
        self.assertEqual(c["attendanceOutcome"], "Marked unpaid leave")
        self.assertEqual(
            (b["status"], b["turnaroundHours"], b["reviewedBy"], b["attendanceOutcome"]),
            ("Pending", None, None, "Awaiting decision"),
        )
        self.assertEqual(b["serviceMonths"], 19)
        self.assertEqual((d["joinDate"], d["serviceMonths"], d["turnaroundHours"]), (None, None, 24.0))
        self.assertEqual(d["attendanceOutcome"], "Not reflected")  # no day record at all
        self.assertEqual((f["turnaroundHours"], f["attendanceOutcome"]), (25.5, "Not reflected"))  # overridden
        self.assertEqual((e["serviceMonths"], e["status"]), (1, "Pending"))

    def test_summary_and_no_totals(self):
        body = self.run_report(self.report_id, **self.params)
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Casual leave requests"], 6)
        self.assertEqual((s["Approved (paid days)"], s["Rejected"], s["Pending"]), (3, 1, 2))
        self.assertEqual(s["Employees"], 6)
        self.assertEqual(s["Average turnaround (hrs)"], 20.5)  # (30.5 + 2 + 24 + 25.5) / 4
        self.assertEqual(s["Approved (paid days)"] + s["Rejected"] + s["Pending"], s["Casual leave requests"])
        self.assertIsNone(body["totals"])

    def test_filters(self):
        def codes(**p):
            return [r["employeeCode"] for r in data_rows(self.run_report(self.report_id, **{**self.params, **p}))]

        self.assertEqual(codes(status="approved"), ["RL_A", "RL_D", "RL_F"])
        self.assertEqual(codes(status="pending"), ["RL_B", "RL_E"])
        self.assertEqual(codes(reviewerRole="dept_head"), ["RL_C"])
        self.assertEqual(codes(employeeIds=str(self.a.id)), ["RL_A"])
        self.assertEqual(codes(employmentType="production"), ["RL_B"])
        self.assertEqual(codes(employeeStatus="inactive"), ["RL_D"])
        self.assertEqual(set(codes(departmentIds=str(self.sewing.id))), {"RL_C"})
        self.assertEqual(codes(designationIds=str(self.tailor.id)), ["RL_A"])
        self.assertEqual(codes(dateFrom="2026-08-01", dateTo="2026-08-31"), ["RL_A"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), ["RL_C"])

    def test_branch_isolation(self):
        codes = {r["employeeCode"] for r in data_rows(self.run_report(self.report_id, self.branch_user, **self.params))}
        self.assertEqual(codes, {"RL_A", "RL_B", "RL_D", "RL_E", "RL_F"})
        self.assertEqual(
            data_rows(self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params)), []
        )
        self.assertEqual(
            data_rows(self.run_report(self.report_id, self.branch_user, employeeIds=str(self.c.id), **self.params)), []
        )

    def test_report_does_not_touch_attendance(self):
        before = list(AttendanceDayRecord.objects.order_by("id").values_list("id", "status", "override_note"))
        self.run_report(self.report_id, **self.params)
        self.sheet(self.report_id, **self.params)
        self.assertEqual(
            before, list(AttendanceDayRecord.objects.order_by("id").values_list("id", "status", "override_note"))
        )

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1, join_date="2024-01-01")
            CasualLeaveRequest.objects.create(
                employee=emp,
                date=date(2026, 9, 2 + i % 20),
                status="approved",
                reviewed_by="Priya HR",
                reviewer_role="hr",
                reviewed_at=ist(2026, 9, 25, 9),
            )
            AttendanceDayRecord.objects.create(
                employee=emp,
                date=date(2026, 9, 2 + i % 20),
                status="present",
                override_note="Casual Leave (paid) -approved",
            )


# ══════════════════════════════════════════════════════════════════════════════
#  Casual Leave Eligibility & Yearly Usage
# ══════════════════════════════════════════════════════════════════════════════


class CasualLeaveEligibilityTests(ContractMixin, Base):
    report_id = "casual-leave-eligibility-usage"
    params = {"period": "2026-09"}
    empty_params = {"period": "2026-09", "departmentIds": "999999"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        mk = lambda code, join, dept=None, branch=None, **kw: mk_emp(  # noqa: E731
            code, dept or cls.cutting, branch or cls.b1, join_date=join, **kw
        )
        cls.e1 = mk("RLE1", "2025-01-10")
        cls.e2 = mk("RLE2", "2026-06-20")
        cls.e3 = mk("RLE3", "15/03/2026")
        cls.e4 = mk("RLE4", None)
        cls.e5 = mk("RLE5", "2024-01-01")
        cls.e6 = mk("RLE6", "2024-01-01")
        cls.e7 = mk("RLE7", "2024-01-01")
        cls.e8 = mk("RLE8", "2024-01-01", employment_type="production")
        cls.e9 = mk("RLE9", "2024-01-01", status="inactive")
        cls.e10 = mk("RLE10", "2024-01-01", cls.sewing, cls.b2)
        cl = lambda emp, dt, status: CasualLeaveRequest.objects.create(employee=emp, date=dt, status=status)  # noqa: E731
        cl(cls.e5, date(2026, 9, 3), "approved")
        cl(cls.e6, date(2026, 9, 20), "pending")
        cl(cls.e7, date(2026, 9, 5), "rejected")  # a rejected request does not use the monthly slot
        cl(cls.e7, date(2026, 2, 10), "approved")
        cl(cls.e7, date(2026, 3, 10), "approved")
        cl(cls.e8, date(2026, 9, 9), "approved")  # production is out of scope
        cl(cls.e5, date(2025, 9, 3), "approved")  # last year: not counted

    def keyed(self, body):
        return {r["employeeCode"]: r for r in data_rows(body) if r["employeeCode"].startswith("RLE")}

    def test_golden_rows(self):
        body = self.run_report(self.report_id, **self.params)
        r = self.keyed(body)
        self.assertEqual(
            sorted(r), ["RLE1", "RLE10", "RLE2", "RLE3", "RLE4", "RLE5", "RLE6", "RLE7"]
        )  # no production RLE8, no inactive RLE9
        e1 = r["RLE1"]
        self.assertEqual(
            (e1["serviceMonths"], e1["eligible"], e1["eligibleFrom"], e1["reason"]),
            (20, "Eligible", "2025-07-10", None),
        )
        self.assertEqual((e1["approvedThisYear"], e1["remainingThisYear"], e1["monthsUsed"]), (0, 12, None))
        e2 = r["RLE2"]
        self.assertEqual(
            (e2["serviceMonths"], e2["eligible"], e2["reason"], e2["eligibleFrom"]),
            (2, "Not yet eligible", "2/6 months of service", "2026-12-20"),
        )
        e3 = r["RLE3"]  # exactly six months on the 15th, joined 15/03/2026
        self.assertEqual(
            (e3["joinDate"], e3["serviceMonths"], e3["eligible"], e3["eligibleFrom"]),
            ("2026-03-15", 6, "Eligible", "2026-09-15"),
        )
        e4 = r["RLE4"]
        self.assertEqual(
            (e4["serviceMonths"], e4["eligible"], e4["joinDate"], e4["eligibleFrom"]),
            (None, "No join date", None, None),
        )
        e5 = r["RLE5"]
        self.assertEqual(
            (e5["eligible"], e5["approvedThisYear"], e5["remainingThisYear"], e5["monthsUsed"]),
            ("Used this month", 1, 11, "Sep"),
        )
        e6 = r["RLE6"]
        self.assertEqual(
            (e6["eligible"], e6["pendingThisYear"], e6["approvedThisYear"], e6["monthsUsed"]),
            ("Used this month", 1, 0, "Sep (pending)"),
        )
        e7 = r["RLE7"]
        self.assertEqual(
            (e7["eligible"], e7["approvedThisYear"], e7["remainingThisYear"], e7["monthsUsed"]),
            ("Eligible", 2, 10, "Feb, Mar"),
        )
        self.assertEqual(r["RLE10"]["eligible"], "Eligible")

    def test_totals_and_summary_agree_with_the_rows(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        t = body["totals"]
        self.assertEqual(t["approvedThisYear"], sum(x["approvedThisYear"] for x in rows))
        # the ten staff rows: RLE1-7 and RLE10, plus the shared fixture's RL_A and RL_C (both long-serving staff)
        self.assertEqual(len(rows), 10)
        self.assertEqual((t["approvedThisYear"], t["pendingThisYear"], t["remainingThisYear"]), (3, 1, 117))
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Eligible"], 6)  # RLE1 RLE3 RLE7 RLE10 RL_A RL_C
        self.assertEqual((s["Not yet eligible"], s["Used this month"], s["No join date"]), (1, 2, 1))
        self.assertEqual(s["Eligible"] + s["Not yet eligible"] + s["Used this month"] + s["No join date"], len(rows))
        self.assertEqual(s["Approved casual leave in 2026"], 3)
        self.assertEqual(s["No approved casual leave in 2026"], 8)

    def test_filters(self):
        def codes(**p):
            return sorted(
                x["employeeCode"]
                for x in data_rows(self.run_report(self.report_id, **{**self.params, **p}))
                if x["employeeCode"].startswith("RLE")
            )

        self.assertEqual(codes(eligibility="used_this_month"), ["RLE5", "RLE6"])
        self.assertEqual(codes(eligibility="no_join_date"), ["RLE4"])
        self.assertEqual(codes(eligibility="not_eligible_service"), ["RLE2"])
        self.assertEqual(codes(eligibility="eligible"), ["RLE1", "RLE10", "RLE3", "RLE7"])
        self.assertEqual(codes(departmentIds=str(self.sewing.id)), ["RLE10"])
        self.assertEqual(
            codes(employeeStatus="all"), ["RLE1", "RLE10", "RLE2", "RLE3", "RLE4", "RLE5", "RLE6", "RLE7", "RLE9"]
        )
        self.assertEqual(codes(employeeStatus="inactive"), ["RLE9"])
        self.assertEqual(codes(employeeIds=f"{self.e1.id},{self.e2.id}"), ["RLE1", "RLE2"])

        def everyone(**p):
            return sorted(x["employeeCode"] for x in data_rows(self.run_report(self.report_id, **{**self.params, **p})))

        self.assertEqual(everyone(designationIds=str(self.tailor.id)), ["RL_A"])
        self.assertEqual(everyone(branchIds=str(self.b2.id)), ["RLE10", "RL_C"])
        # an October check date: RLE2 still has under six months, RLE3 has seven
        october = self.keyed(self.run_report(self.report_id, period="2026-10"))
        self.assertEqual((october["RLE2"]["serviceMonths"], october["RLE3"]["serviceMonths"]), (3, 7))
        self.assertEqual(october["RLE5"]["eligible"], "Eligible")  # September's use no longer blocks October

    def test_branch_isolation(self):
        codes = {x["employeeCode"] for x in data_rows(self.run_report(self.report_id, self.branch_user, **self.params))}
        self.assertNotIn("RLE10", codes)
        self.assertIn("RLE1", codes)
        widened = data_rows(self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params))
        self.assertEqual(widened, [])

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1, join_date="2024-01-01")
            CasualLeaveRequest.objects.create(employee=emp, date=date(2026, 5, 4), status="approved")


# ══════════════════════════════════════════════════════════════════════════════
#  Employee Requests Register
# ══════════════════════════════════════════════════════════════════════════════


class EmployeeRequestsTests(ContractMixin, Base):
    report_id = "employee-requests-register"
    params = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
    empty_params = {"dateFrom": "2030-09-01", "dateTo": "2030-09-30"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()

        def req(emp, rtype, subject, status, created, desc="", handled_by=None, handled_at=None, notes=None):
            r = EmployeeRequest.objects.create(
                employee=emp,
                request_type=rtype,
                subject=subject,
                description=desc,
                status=status,
                handled_by=handled_by,
                handled_at=handled_at,
                hr_notes=notes,
            )
            return stamp(r, created_at=created)

        cls.r1 = req(cls.a, "salary_enquiry", "Salary query", "pending", ist(2026, 9, 2, 9), "Why is my PF lower?")
        cls.r2 = req(
            cls.a,
            "shift_correction",
            "Shift fix",
            "approved",
            ist(2026, 9, 4, 9),
            "Moved to B shift",
            "Priya HR",
            ist(2026, 9, 5, 13, 30),
            "Done",
        )
        cls.r3 = req(cls.b, "general", "Canteen", "in_review", ist(2026, 9, 10, 9))
        cls.r4 = req(cls.c, "advance", "Advance", "rejected", ist(2026, 9, 12, 9), "", "Asha", ist(2026, 9, 12, 10))
        cls.r5 = req(cls.d, "permission", "Permission", "more_info", ist(2026, 9, 15, 0, 30))  # 14 Sep in UTC
        cls.r6 = req(
            cls.a, "leave", "Leave ticket", "approved", ist(2026, 8, 30, 9), "", "Priya HR", ist(2026, 8, 31, 9)
        )
        cls.r7 = req(cls.a, "general", "Long one", "pending", ist(2026, 9, 20, 9), "x" * 500)

    def test_golden_rows_and_order(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        self.assertEqual(
            [r["subject"] for r in rows], ["Salary query", "Shift fix", "Canteen", "Advance", "Permission", "Long one"]
        )
        r1, r2, r3, r4, r5, r7 = rows
        self.assertEqual(
            (r1["requestType"], r1["status"], r1["createdAt"], r1["openAgeDays"], r1["turnaroundHours"]),
            ("Salary Enquiry", "Pending", "2026-09-02 09:00", 27, None),
        )
        self.assertEqual(
            (r2["status"], r2["handledBy"], r2["handledAt"], r2["turnaroundHours"], r2["openAgeDays"], r2["hrNotes"]),
            ("Approved", "Priya HR", "2026-09-05 13:30", 28.5, None, "Done"),
        )
        self.assertEqual((r3["status"], r3["openAgeDays"]), ("In Review", 19))
        self.assertEqual((r4["status"], r4["turnaroundHours"], r4["openAgeDays"]), ("Rejected", 1.0, None))
        self.assertEqual(
            (r5["status"], r5["openAgeDays"], r5["createdAt"]), ("More Info Needed", 14, "2026-09-15 00:30")
        )
        self.assertEqual(len(r7["description"]), 300)  # screen text is clipped
        self.assertTrue(r7["description"].endswith("…"))

    def test_summary_notes_and_no_totals(self):
        body = self.run_report(self.report_id, **self.params)
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Requests"], 6)
        self.assertEqual((s["Open"], s["Approved"], s["Rejected"]), (4, 1, 1))
        self.assertEqual(s["Open"] + s["Approved"] + s["Rejected"], s["Requests"])
        self.assertEqual(s["Average handling time (hrs)"], 14.75)  # (28.5 + 1) / 2
        self.assertEqual(s["Oldest open request (days)"], 27)
        self.assertIsNone(body["totals"])
        self.assertTrue(any(n.startswith("By type:") and "General Query 2" in n for n in body["notes"]))

    def test_filters(self):
        def subjects(**p):
            return [r["subject"] for r in data_rows(self.run_report(self.report_id, **{**self.params, **p}))]

        self.assertEqual(subjects(requestType="salary_enquiry"), ["Salary query"])
        self.assertEqual(subjects(status="pending"), ["Salary query", "Long one"])
        self.assertEqual(subjects(handledBy="priya"), ["Shift fix"])
        self.assertEqual(subjects(employeeIds=str(self.b.id)), ["Canteen"])
        self.assertEqual(subjects(employmentType="production"), ["Canteen"])
        self.assertEqual(subjects(departmentIds=str(self.sewing.id)), ["Advance"])
        self.assertEqual(subjects(employeeStatus="inactive"), ["Permission"])
        self.assertEqual(subjects(branchIds=str(self.b2.id)), ["Advance"])
        self.assertEqual(subjects(designationIds=str(self.tailor.id)), ["Salary query", "Shift fix", "Long one"])
        self.assertEqual(subjects(dateFrom="2026-08-30", dateTo="2026-08-30"), ["Leave ticket"])
        # bucketed by the IST day: 00:30 on the 15th is the 14th in UTC
        self.assertEqual(subjects(dateFrom="2026-09-15", dateTo="2026-09-15"), ["Permission"])
        self.assertEqual(subjects(dateFrom="2026-09-14", dateTo="2026-09-14"), [])

    def test_branch_isolation_the_legacy_list_did_not_scope(self):
        rows = data_rows(self.run_report(self.report_id, self.branch_user, **self.params))
        self.assertNotIn("Advance", [r["subject"] for r in rows])  # RL_C belongs to the other branch
        self.assertEqual(len(rows), 5)
        self.assertEqual(
            data_rows(self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params)), []
        )
        self.assertEqual(
            [r["subject"] for r in data_rows(self.run_report(self.report_id, self.branch_user2, **self.params))],
            ["Advance"],
        )

    def test_excel_carries_the_full_description(self):
        ws = self.sheet(self.report_id, **self.params)
        found = [
            c.value
            for row in ws.iter_rows(min_row=8)
            for c in row
            if isinstance(c.value, str) and c.value.startswith("xxxx")
        ]
        self.assertEqual([len(v) for v in found], [500])

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1)
            r = EmployeeRequest.objects.create(employee=emp, request_type="general", subject=f"S{i}", description="d")
            stamp(r, created_at=ist(2026, 9, 20, 9))


# ══════════════════════════════════════════════════════════════════════════════
#  Approval workflows: shared fixture
# ══════════════════════════════════════════════════════════════════════════════


class ApprovalBase(Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # HODs: RL_H1 covers CUTTING (all rights); RL_H2 covers SEWING but may not act on missing punches.
        cls.h1e = mk_emp("RL_H1", cls.cutting, cls.b1)
        cls.h2e = mk_emp("RL_H2", cls.sewing, cls.b2)
        cls.h1 = DepartmentManager.objects.create(employee=cls.h1e)
        cls.h2 = DepartmentManager.objects.create(employee=cls.h2e, can_approve_missing_punch=False)
        ManagerDepartmentAssignment.objects.create(manager=cls.h1, department=cls.cutting)
        ManagerDepartmentAssignment.objects.create(manager=cls.h2, department=cls.sewing)
        cls.f = mk_emp("RL_F", cls.packing, cls.b1)  # PACKING has no HOD


def _leave(emp, status="pending", start="2026-10-05", end="2026-10-06", created=None, by=None, role=None):
    lr = LeaveRequest.objects.create(
        employee=emp,
        type="casual",
        start_date=start,
        end_date=end,
        total_days=2,
        status=status,
        approved_by=by,
        approver_role=role,
    )
    return stamp(lr, created_at=created) if created else lr


# ══════════════════════════════════════════════════════════════════════════════
#  Pending Approvals Ageing
# ══════════════════════════════════════════════════════════════════════════════


class PendingApprovalsTests(ContractMixin, ApprovalBase):
    report_id = "pending-approvals-ageing"
    params: dict = {}
    empty_params = {"departmentIds": "999999"}
    first_key = "module"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # 1 leave (HOD or HR), 28 Sep
        cls.leave_a = _leave(cls.a, created=ist(2026, 9, 28, 9))
        _leave(cls.a, "approved", created=ist(2026, 9, 1, 9), by="Priya HR", role="hr")  # decided: ignored
        # 2 permission, no HOD, 20 Sep
        cls.perm_f = stamp(
            EmployeePermission.objects.create(
                employee=cls.f,
                date=date(2026, 10, 1),
                permission_time="09:30",
                type="morning_late_in",
                status="pending",
            ),
            created_at=ist(2026, 9, 20, 10),
        )
        EmployeePermission.objects.create(employee=cls.f, date=date(2026, 10, 2), status="rejected")
        # 3 casual leave, C (RL_H2 holds the right), 25 Sep
        cls.cl_c = stamp(
            CasualLeaveRequest.objects.create(employee=cls.c, date=date(2026, 10, 14)), created_at=ist(2026, 9, 25, 9)
        )
        # missing punches: C waits at a HOD lacking the right (stuck), A at HR, F has no HOD (stuck)
        mp = lambda emp, status, created: stamp(
            MissingPunchRequest.objects.create(  # noqa: E731
                employee=emp,
                date=date(2026, 9, 8),
                punch_time="09:05",
                punch_type="IN",
                punch_slot="morning_in",
                reason="forgot",
                status=status,
            ),
            created_at=created,
        )
        cls.mp_c = mp(cls.c, "pending_hod", ist(2026, 9, 10, 9))
        cls.mp_a = mp(cls.a, "pending_hr", ist(2026, 9, 22, 9))
        cls.mp_f = mp(cls.f, "pending_hod", ist(2026, 9, 27, 9))
        mp(cls.a, "approved", ist(2026, 9, 3, 9))  # decided: ignored
        # on duty + its punch; 00:30 IST on the 29th is still the 28th in UTC
        cls.od_a = stamp(
            OnDutySession.objects.create(employee=cls.a, destination="Tirupur Yarn Mills", branch=cls.b1),
            created_at=ist(2026, 9, 29, 0, 30),
        )
        cls.pv_a = stamp(
            OnDutyPunchVerification.objects.create(
                session=cls.od_a,
                employee=cls.a,
                punch_date=date(2026, 9, 29),
                punch_time="09:00",
                punch_type="IN",
                punch_number=1,
                latitude=Decimal("11.1"),
                longitude=Decimal("77.3"),
                photo="p.jpg",
            ),
            created_at=ist(2026, 9, 29, 9),
        )
        # attendance correction proposed for the production employee
        cls.ao_b = stamp(
            AttendanceOverrideRequest.objects.create(employee=cls.b, date=date(2026, 9, 9), requested_by="Priya HR"),
            created_at=ist(2026, 9, 14, 9),
        )
        cls.op_a = stamp(
            OutpassRequest.objects.create(employee=cls.a, destination="Bank", reason="cash"),
            created_at=ist(2026, 9, 29, 8),
        )
        OutpassRequest.objects.create(  # created by an on-duty approval: never pending, and not manual
            employee=cls.a, destination="x", reason="y", status="approved", source="on_duty"
        )
        cls.er_a = stamp(
            EmployeeRequest.objects.create(
                employee=cls.a, request_type="salary_enquiry", subject="PF", description="d"
            ),
            created_at=ist(2026, 9, 1, 9),
        )
        cls.res_a = stamp(
            ResignationRequest.objects.create(employee=cls.a, reason="move"), created_at=ist(2026, 9, 16, 9)
        )
        cls.res_f = stamp(
            ResignationRequest.objects.create(employee=cls.f, reason="move", status="dept_approved"),
            created_at=ist(2026, 9, 18, 9),
        )
        cls.adv_c = stamp(
            Advance.objects.create(employee=cls.c, advance_type="general", amount=Decimal("5000"), status="pending"),
            created_at=ist(2026, 9, 26, 9),
        )
        cls.leave_h1 = _leave(cls.h1e, created=ist(2026, 9, 29, 9))  # the HOD's own request

    def rows(self, user=None, **p):
        return data_rows(self.run_report(self.report_id, user, **p))

    def test_golden_rows_ordered_oldest_first(self):
        rows = self.rows()
        self.assertEqual(len(rows), 15)
        self.assertEqual(
            [(r["module"], r["employeeCode"], r["ageDays"]) for r in rows],
            [
                ("Employee Request", "RL_A", 28),
                ("Missing Punch", "RL_C", 19),
                ("Attendance Correction", "RL_B", 15),
                ("Resignation", "RL_A", 13),
                ("Resignation", "RL_F", 11),
                ("Permission", "RL_F", 9),
                ("Missing Punch", "RL_A", 7),
                ("Casual Leave", "RL_C", 4),
                ("Advance / Loan", "RL_C", 3),
                ("Missing Punch", "RL_F", 2),
                ("Leave", "RL_A", 1),
                ("Leave", "RL_H1", 0),
                ("On-Duty", "RL_A", 0),
                ("On-Duty Punch", "RL_A", 0),
                ("Outpass", "RL_A", 0),
            ],
        )

    def test_who_each_request_is_waiting_on(self):
        rows = self.rows()
        by = {(r["module"], r["employeeCode"]): r for r in rows}
        leave = by[("Leave", "RL_A")]
        self.assertEqual(
            (leave["pendingWith"], leave["hodAssigned"], leave["stage"], leave["ageBucket"], leave["requestId"]),
            ("HOD RL_H1 T or HR", "Yes", "Pending", "0-1 days", self.leave_a.id),
        )
        self.assertIn("05-Oct-2026 to 06-Oct-2026 (2 day(s))", leave["subject"])
        self.assertEqual(leave["requestDate"], "2026-10-05")
        self.assertEqual(leave["submittedAt"], "2026-09-28 09:00")
        perm = by[("Permission", "RL_F")]
        self.assertEqual(
            (perm["pendingWith"], perm["hodAssigned"], perm["ageBucket"]), ("HR (no HOD assigned)", "No", "8-15 days")
        )
        self.assertEqual(perm["subject"], "Morning Late-In on 01-Oct-2026 at 09:30")
        self.assertEqual(by[("Casual Leave", "RL_C")]["pendingWith"], "HOD RL_H2 T or HR")
        mp_c = by[("Missing Punch", "RL_C")]
        self.assertEqual(mp_c["stage"], "HOD stage")
        self.assertEqual(mp_c["ageBucket"], "Over 15 days")
        self.assertEqual(
            mp_c["pendingWith"],
            "No approver can act (HOD RL_H2 T lacks the approval right); HR cannot decide this step",
        )
        self.assertEqual(mp_c["subject"], "Morning Check-In 09:05 on 08-Sep-2026")
        mp_a = by[("Missing Punch", "RL_A")]
        self.assertEqual((mp_a["stage"], mp_a["pendingWith"]), ("HR stage", "HR"))
        self.assertEqual(
            by[("Missing Punch", "RL_F")]["pendingWith"],
            "No approver can act (no HOD assigned); HR cannot decide this step",
        )
        od = by[("On-Duty", "RL_A")]
        self.assertEqual(
            (od["pendingWith"], od["ageDays"], od["subject"]), ("HOD RL_H1 T or HR", 0, "Tirupur Yarn Mills")
        )
        self.assertEqual(od["requestDate"], "2026-09-29")  # the IST day, not the UTC one
        punch = by[("On-Duty Punch", "RL_A")]
        self.assertEqual(punch["pendingWith"], "HR - after the On-Duty request is approved")
        self.assertIsNone(punch["hodAssigned"])
        self.assertEqual(by[("Attendance Correction", "RL_B")]["pendingWith"], "HOD RL_H1 T")
        self.assertEqual(by[("Outpass", "RL_A")]["pendingWith"], "HOD RL_H1 T or HR")
        er = by[("Employee Request", "RL_A")]
        self.assertEqual((er["pendingWith"], er["hodAssigned"], er["ageBucket"]), ("HR", None, "Over 15 days"))
        res = {r["stage"]: r for r in rows if r["module"] == "Resignation"}
        self.assertEqual(res["HOD stage"]["pendingWith"], "HOD RL_H1 T (HR can only reject)")
        self.assertEqual((res["HR stage"]["pendingWith"], res["HR stage"]["hodAssigned"]), ("HR", "No"))
        self.assertEqual(by[("Advance / Loan", "RL_C")]["subject"], "General Advance of Rs. 5,000.00")
        own = by[("Leave", "RL_H1")]
        self.assertEqual(own["pendingWith"], "HR (the employee is their own HOD)")

    def test_summary_and_notes(self):
        body = self.run_report(self.report_id)
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Pending requests"], 15)
        self.assertEqual(s["Oldest (days)"], 28)
        self.assertEqual(s["Waiting over 7 days"], 6)
        self.assertEqual(s["Stuck - no approver can act"], 2)
        self.assertEqual((s["With HR only"], s["With the HOD only"], s["HOD or HR"]), (7, 2, 4))
        self.assertEqual(
            s["Stuck - no approver can act"] + s["With HR only"] + s["With the HOD only"] + s["HOD or HR"],
            s["Pending requests"],
        )
        self.assertIsNone(body["totals"])
        notes = " ".join(body["notes"])
        self.assertIn("By age: 0-1 days 5, 2-3 days 2, 4-7 days 2, 8-15 days 4, Over 15 days 2.", notes)
        self.assertIn("Missing Punch 3", notes)

    def test_filters(self):
        def keys(user=None, **p):
            return [(r["module"], r["employeeCode"]) for r in self.rows(user, **p)]

        self.assertEqual(len(keys(pendingWith="stuck")), 2)
        self.assertEqual(len(keys(pendingWith="hod")), 6)  # HOD-only plus HOD-or-HR
        self.assertEqual(len(keys(pendingWith="hr")), 11)  # HR-only plus HOD-or-HR
        self.assertEqual(
            keys(minAgeDays="15"),
            [("Employee Request", "RL_A"), ("Missing Punch", "RL_C"), ("Attendance Correction", "RL_B")],
        )
        self.assertEqual(len(keys(module="missing_punch")), 3)
        self.assertEqual({m for m, _ in keys(module="leave,permission")}, {"Leave", "Permission"})
        self.assertEqual(len(keys(module="leave,permission")), 3)
        self.assertEqual(len(keys(employeeIds=str(self.a.id))), 7)
        self.assertEqual(keys(employmentType="production"), [("Attendance Correction", "RL_B")])
        self.assertEqual({c for _, c in keys(departmentIds=str(self.sewing.id))}, {"RL_C"})
        self.assertEqual(len(keys(departmentIds=str(self.sewing.id))), 3)
        self.assertEqual(len(keys(designationIds=str(self.tailor.id))), 7)  # RL_A is the only Tailor
        self.assertEqual({c for _, c in keys(branchIds=str(self.b2.id))}, {"RL_C"})
        self.assertEqual(len(keys(employeeStatus="active")), 15)
        self.assertEqual(keys(employeeStatus="inactive"), [])  # nobody who has left is waiting on a decision here
        r = self.get(f"/api/reports/run/{self.report_id}", module="nonsense")
        self.assertEqual(r.status_code, 400)

    def test_branch_isolation(self):
        keys = {(r["module"], r["employeeCode"]) for r in self.rows(self.branch_user)}
        self.assertEqual(len(keys), 12)
        self.assertFalse({c for _, c in keys} & {"RL_C"})
        self.assertEqual(self.rows(self.branch_user, employeeIds=str(self.c.id)), [])
        self.assertEqual(self.rows(self.branch_user, branchIds=str(self.b2.id)), [])
        self.assertEqual(
            {c for _, c in ((r["module"], r["employeeCode"]) for r in self.rows(self.branch_user2))}, {"RL_C"}
        )

    def test_a_role_only_sees_the_workflows_it_can_open(self):
        body = self.run_report(self.report_id, self.leave_only)
        rows = data_rows(body)
        self.assertEqual({r["module"] for r in rows}, {"Leave"})
        self.assertEqual(len(rows), 2)
        self.assertTrue(
            any("Not shown because your role cannot open them" in n and "Missing Punch" in n for n in body["notes"])
        )
        asked = self.run_report(self.report_id, self.leave_only, module="missing_punch")
        self.assertEqual(data_rows(asked), [])  # asking for a hidden workflow yields nothing, never the others

    def test_a_request_waiting_for_more_information_is_the_employees_move(self):
        stamp(
            EmployeeRequest.objects.create(
                employee=self.a, request_type="general", subject="Docs", description="d", status="more_info"
            ),
            created_at=ist(2026, 9, 24, 9),
        )
        body = self.run_report(self.report_id)
        row = next(r for r in data_rows(body) if r["subject"] == "General Query: Docs")
        self.assertEqual(
            (row["pendingWith"], row["stage"], row["ageDays"]),
            ("Employee (HR asked for more information)", "More info needed", 5),
        )
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual((s["Pending requests"], s["With HR only"], s["Waiting on the employee"]), (16, 7, 1))
        self.assertNotIn(row["subject"], [r["subject"] for r in self.rows(pendingWith="hr")])

    def test_no_writes(self):
        before = (MissingPunchRequest.objects.count(), OnDutySession.objects.count(), EmployeeRequest.objects.count())
        self.run_report(self.report_id)
        self.pdf(self.report_id)
        self.assertEqual(
            before,
            (MissingPunchRequest.objects.count(), OnDutySession.objects.count(), EmployeeRequest.objects.count()),
        )

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting if i % 2 else self.sewing, self.b1 if i % 2 else self.b2)
            stamp(
                LeaveRequest.objects.create(
                    employee=emp, type="casual", start_date="2026-10-01", end_date="2026-10-01", total_days=1
                ),
                created_at=ist(2026, 9, 20, 9),
            )
            stamp(
                MissingPunchRequest.objects.create(
                    employee=emp, date=date(2026, 9, 8), punch_time="09:00", punch_type="IN", reason="r"
                ),
                created_at=ist(2026, 9, 21, 9),
            )
            stamp(
                OnDutySession.objects.create(employee=emp, destination="d", branch=emp.branch),
                created_at=ist(2026, 9, 22, 9),
            )


# ══════════════════════════════════════════════════════════════════════════════
#  Approver Workload & Turnaround
# ══════════════════════════════════════════════════════════════════════════════


class ApproverWorkloadTests(ContractMixin, ApprovalBase):
    report_id = "approver-workload-turnaround"
    params = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
    empty_params = {"dateFrom": "2030-09-01", "dateTo": "2030-09-30", "departmentIds": "999999"}
    first_key = "approverName"
    first_value = "(not recorded)"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        P, H1, H2 = "Priya HR", "RL_H1 T", "RL_H2 T"

        # Leave and permission store no decision time: they are placed by the request date (September).
        _leave(cls.a, "approved", created=ist(2026, 9, 3, 9), by=P, role="hr")
        _leave(cls.a, "rejected", created=ist(2026, 9, 4, 9), by=P, role="hr")
        _leave(cls.b, "approved", created=ist(2026, 9, 5, 9), by=H1, role="dept_head")
        _leave(cls.a, "approved", created=ist(2026, 8, 3, 9), by=P, role="hr")  # August: out of range
        _leave(cls.a, "approved", created=ist(2026, 9, 6, 9))  # approver never recorded
        perm = lambda emp, st, by, role: stamp(
            EmployeePermission.objects.create(  # noqa: E731
                employee=emp, date=date(2026, 9, 10), status=st, approved_by=by, approver_role=role
            ),
            created_at=ist(2026, 9, 7, 9),
        )
        perm(cls.a, "approved", P, "hr")
        perm(cls.c, "rejected", H2, "dept_head")
        # casual leave: decision time is known
        cl = lambda emp, st, by, role, created, reviewed: stamp(
            CasualLeaveRequest.objects.create(  # noqa: E731
                employee=emp,
                date=date(2026, 9, 20),
                status=st,
                reviewed_by=by,
                reviewer_role=role,
                reviewed_at=reviewed,
            ),
            created_at=created,
        )
        cl(cls.a, "approved", P, "hr", ist(2026, 9, 4, 10), ist(2026, 9, 7, 16, 30))  # 78.5 h
        cl(cls.c, "rejected", H2, "dept_head", ist(2026, 9, 10, 9), ist(2026, 9, 10, 11))  # 2 h
        cl(cls.a, "approved", P, "hr", ist(2026, 9, 28, 9), ist(2026, 10, 1, 9))  # decided in October
        # missing punch: two stages, each with its own timestamps
        mp = lambda emp, st, created, hod_by, hod_at, hr_by, hr_at: stamp(
            MissingPunchRequest.objects.create(  # noqa: E731
                employee=emp,
                date=date(2026, 9, 8),
                punch_time="09:00",
                punch_type="IN",
                reason="r",
                status=st,
                hod_reviewed_by=hod_by,
                hod_reviewed_at=hod_at,
                hr_reviewed_by=hr_by,
                hr_reviewed_at=hr_at,
            ),
            created_at=created,
        )
        mp(cls.a, "approved", ist(2026, 9, 1, 9), H1, ist(2026, 9, 1, 12), P, ist(2026, 9, 2, 9))  # 3 h then 21 h
        mp(cls.c, "rejected", ist(2026, 9, 5, 9), H2, ist(2026, 9, 5, 10), None, None)  # HOD rejection is terminal
        mp(cls.a, "pending_hr", ist(2026, 9, 20, 9), H1, ist(2026, 9, 20, 13), None, None)  # HOD passed it on
        # on duty
        od = lambda emp, st, created, hod_by, hod_at, hr_by, hr_at: stamp(
            OnDutySession.objects.create(  # noqa: E731
                employee=emp,
                destination="d",
                branch=emp.branch,
                status=st,
                hod_reviewed_by=hod_by,
                hod_reviewed_at=hod_at,
                hr_reviewed_by=hr_by,
                hr_reviewed_at=hr_at,
            ),
            created_at=created,
        )
        cls.od_direct = od(
            cls.a, "active", ist(2026, 9, 8, 9), None, None, P, ist(2026, 9, 8, 11)
        )  # HR decided directly
        od(cls.b, "rejected", ist(2026, 9, 9, 9), H1, ist(2026, 9, 9, 9, 30), None, None)
        pv = lambda st, created, at, comment=None: stamp(
            OnDutyPunchVerification.objects.create(  # noqa: E731
                session=cls.od_direct,
                employee=cls.a,
                punch_date=date(2026, 9, 8),
                punch_time="09:00",
                punch_type="IN",
                punch_number=1,
                latitude=Decimal("11.1"),
                longitude=Decimal("77.3"),
                photo="p.jpg",
                status=st,
                hr_reviewed_by=P,
                hr_reviewed_at=at,
                hr_review_comment=comment,
            ),
            created_at=created,
        )
        pv("approved", ist(2026, 9, 8, 9, 30), ist(2026, 9, 8, 12))  # 2.5 h
        pv("rejected", ist(2026, 9, 8, 9, 30), ist(2026, 9, 8, 13), "blurred")  # 3.5 h
        pv(
            "rejected",
            ist(2026, 9, 8, 9, 30),
            ist(2026, 9, 8, 13),
            "Voided automatically -the On-Duty request was rejected by HR.",
        )  # system, excluded
        # attendance corrections
        ao = lambda st, by, created, reviewed: stamp(
            AttendanceOverrideRequest.objects.create(  # noqa: E731
                employee=cls.a if st == "approved" else cls.b,
                date=date(2026, 9, 9),
                status=st,
                reviewed_by=by,
                reviewed_at=reviewed,
                requested_by=P,
            ),
            created_at=created,
        )
        ao("approved", H1, ist(2026, 9, 11, 10), ist(2026, 9, 12, 10))
        ao("rejected", None, ist(2026, 9, 11, 10), None)  # superseded: excluded
        ao("rejected", H1, ist(2026, 9, 12, 9), ist(2026, 9, 13, 9))
        # outpass: approved ones have approved_at, a rejected one only its request date
        stamp(
            OutpassRequest.objects.create(
                employee=cls.a,
                destination="x",
                reason="r",
                status="approved",
                approved_by=H1,
                approver_role="dept_head",
                approved_at=ist(2026, 9, 15, 10),
            ),
            created_at=ist(2026, 9, 15, 9),
        )
        stamp(
            OutpassRequest.objects.create(
                employee=cls.a, destination="x", reason="r", status="rejected", approved_by=P, approver_role="hr"
            ),
            created_at=ist(2026, 9, 16, 9),
        )
        stamp(
            OutpassRequest.objects.create(
                employee=cls.a,
                destination="x",
                reason="r",
                status="approved",
                approved_by=P,
                approver_role="system",
                source="on_duty",
                approved_at=ist(2026, 9, 17, 10),
            ),
            created_at=ist(2026, 9, 17, 10),
        )  # auto-created by an on-duty approval: excluded
        # employee requests
        stamp(
            EmployeeRequest.objects.create(
                employee=cls.a,
                request_type="general",
                subject="s",
                description="d",
                status="approved",
                handled_by=P,
                handled_at=ist(2026, 9, 3, 15),
            ),
            created_at=ist(2026, 9, 2, 15),
        )  # 24 h
        # work still waiting: a leave (HOD or HR), a missing punch at the HOD stage, an employee request (HR)
        stamp(
            MissingPunchRequest.objects.create(
                employee=cls.a,
                date=date(2026, 9, 8),
                punch_time="09:00",
                punch_type="IN",
                reason="r",
                status="pending_hod",
            ),
            created_at=ist(2026, 9, 26, 9),
        )
        _leave(cls.a, "pending", created=ist(2026, 9, 27, 9))
        stamp(
            EmployeeRequest.objects.create(employee=cls.a, request_type="general", subject="s", description="d"),
            created_at=ist(2026, 9, 28, 9),
        )

    def keyed(self, body):
        return {(r["approverName"], r["approverRole"], r["module"]): r for r in data_rows(body)}

    def test_golden_rows(self):
        body = self.run_report(self.report_id, **self.params)
        k = self.keyed(body)

        def row(name, role, module):
            r = k[(name, role, module)]
            return (
                r["approved"],
                r["rejected"],
                r["total"],
                r["rejectionRate"],
                r["avgTurnaroundHours"],
                r["maxTurnaroundHours"],
                r["timingBasis"],
            )

        P, H1, H2 = "Priya HR", "RL_H1 T", "RL_H2 T"
        self.assertEqual(row(P, "HR", "Leave"), (1, 1, 2, 50.0, None, None, "N/A"))
        self.assertEqual(row(P, "HR", "Permission"), (1, 0, 1, 0.0, None, None, "N/A"))
        self.assertEqual(row(P, "HR", "Casual Leave"), (1, 0, 1, 0.0, 78.5, 78.5, "Decision time"))
        self.assertEqual(row(P, "HR", "Missing Punch - HR stage"), (1, 0, 1, 0.0, 21.0, 21.0, "Stage time"))
        self.assertEqual(row(P, "HR", "On-Duty - HR stage"), (1, 0, 1, 0.0, 2.0, 2.0, "Stage time"))
        self.assertEqual(row(P, "HR", "On-Duty Punch"), (1, 1, 2, 50.0, 3.0, 3.5, "Decision time"))
        self.assertEqual(row(P, "HR", "Outpass"), (0, 1, 1, 100.0, None, None, "Decision time"))
        self.assertEqual(row(P, "HR", "Employee Request"), (1, 0, 1, 0.0, 24.0, 24.0, "Last handled (approx.)"))
        self.assertEqual(row(H1, "HOD", "Leave"), (1, 0, 1, 0.0, None, None, "N/A"))
        self.assertEqual(row(H1, "HOD", "Missing Punch - HOD stage"), (2, 0, 2, 0.0, 3.5, 4.0, "Stage time"))
        self.assertEqual(row(H1, "HOD", "On-Duty - HOD stage"), (0, 1, 1, 100.0, 0.5, 0.5, "Stage time"))
        self.assertEqual(row(H1, "HOD", "Attendance Correction"), (1, 1, 2, 50.0, 24.0, 24.0, "Decision time"))
        self.assertEqual(row(H1, "HOD", "Outpass"), (1, 0, 1, 0.0, 1.0, 1.0, "Decision time"))
        self.assertEqual(row(H2, "HOD", "Permission"), (0, 1, 1, 100.0, None, None, "N/A"))
        self.assertEqual(row(H2, "HOD", "Casual Leave"), (0, 1, 1, 100.0, 2.0, 2.0, "Decision time"))
        self.assertEqual(row(H2, "HOD", "Missing Punch - HOD stage"), (0, 1, 1, 100.0, 1.0, 1.0, "Stage time"))
        self.assertEqual(row("(not recorded)", "Unknown", "Leave"), (1, 0, 1, 0.0, None, None, "N/A"))
        names = [r["approverName"] for r in data_rows(body)]
        self.assertEqual(names, sorted(names, key=str.lower))
        # nothing that must be left out sneaked in
        self.assertNotIn((P, "System", "Outpass"), k)
        self.assertEqual(sum(1 for r in data_rows(body) if r["module"] == "On-Duty Punch"), 1)

    def test_pending_now_is_shown_for_hods_and_the_hr_queue_only(self):
        k = self.keyed(self.run_report(self.report_id, **self.params))
        self.assertEqual(k[("RL_H1 T", "HOD", "Leave")]["pendingNow"], 1)
        self.assertEqual(k[("RL_H1 T", "HOD", "Missing Punch - HOD stage")]["pendingNow"], 1)
        self.assertEqual(k[("RL_H1 T", "HOD", "Attendance Correction")]["pendingNow"], 0)
        self.assertEqual(k[("RL_H2 T", "HOD", "Casual Leave")]["pendingNow"], 0)
        self.assertEqual(k[("HR queue (shared)", "HR", "Leave")]["pendingNow"], 1)
        self.assertEqual(k[("HR queue (shared)", "HR", "Employee Request")]["pendingNow"], 1)
        self.assertEqual(k[("HR queue (shared)", "HR", "Missing Punch - HR stage")]["pendingNow"], 1)
        self.assertNotIn(("HR queue (shared)", "HR", "Missing Punch - HOD stage"), k)  # HR cannot act at the HOD stage
        queue = k[("HR queue (shared)", "HR", "Leave")]
        self.assertEqual(
            (queue["approved"], queue["rejected"], queue["total"], queue["avgTurnaroundHours"]), (0, 0, 0, None)
        )
        self.assertIsNone(k[("Priya HR", "HR", "Leave")]["pendingNow"])  # pending work is not assigned to one HR person

    def test_totals_and_summary_agree_with_the_rows(self):
        body = self.run_report(self.report_id, **self.params)
        rows = data_rows(body)
        t = body["totals"]
        self.assertEqual((t["approved"], t["rejected"], t["total"]), (13, 8, 21))
        self.assertEqual(t["total"], sum(r["total"] for r in rows))
        self.assertEqual(t["approved"] + t["rejected"], t["total"])
        self.assertNotIn("pendingNow", {k for k, v in t.items() if v is not None})
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(s["Decisions in the period"], 21)
        self.assertEqual(s["Average turnaround (hrs)"], 13.64)  # 191 hours over 14 timed decisions
        self.assertEqual((s["Decided by HR"], s["Decided by HODs"]), (10, 10))
        self.assertEqual(s["Fastest approver"], "RL_H2 T (1.5 h)")
        self.assertEqual(s["Slowest approver"], "Priya HR (21.9 h)")
        self.assertEqual(
            s["Requests waiting now"], 4
        )  # leave, missing punch at HOD stage, missing punch at HR stage, request
        self.assertTrue(any("not additive" in n for n in body["notes"]))

    def test_filters(self):
        def keys(user=None, **p):
            return {
                (r["approverName"], r["module"])
                for r in data_rows(self.run_report(self.report_id, user, **{**self.params, **p}))
            }

        hods = keys(approverRole="dept_head")
        self.assertEqual({n for n, _ in hods}, {"RL_H1 T", "RL_H2 T"})
        self.assertEqual({n for n, _ in keys(approverRole="hr")}, {"Priya HR", "HR queue (shared)"})
        self.assertEqual({n for n, _ in keys(approverName="priya")}, {"Priya HR"})
        self.assertEqual({m for _, m in keys(module="leave")}, {"Leave"})
        self.assertEqual({m for _, m in keys(module="casual_leave,permission")}, {"Casual Leave", "Permission"})
        sewing = keys(departmentIds=str(self.sewing.id))
        self.assertEqual(
            sewing, {("RL_H2 T", "Permission"), ("RL_H2 T", "Casual Leave"), ("RL_H2 T", "Missing Punch - HOD stage")}
        )
        only_a = keys(employeeIds=str(self.a.id))
        self.assertNotIn(("RL_H2 T", "Permission"), only_a)
        self.assertEqual({n for n, _ in keys(branchIds=str(self.b2.id))}, {"RL_H2 T"})  # RL_C's decisions only
        october = self.run_report(self.report_id, dateFrom="2026-10-01", dateTo="2026-10-31")
        decided = [r for r in data_rows(october) if r["total"]]
        self.assertEqual(
            [(r["approverName"], r["module"], r["total"]) for r in decided], [("Priya HR", "Casual Leave", 1)]
        )
        self.assertEqual(decided[0]["avgTurnaroundHours"], 72.0)  # 28 Sep 09:00 to 1 Oct 09:00

    def test_branch_isolation(self):
        body = self.run_report(self.report_id, self.branch_user, **self.params)
        names = {r["approverName"] for r in data_rows(body)}
        self.assertNotIn("RL_H2 T", names)  # every decision RL_H2 took was on the other branch's employee RL_C
        self.assertIn("RL_H1 T", names)
        widened = self.run_report(self.report_id, self.branch_user, branchIds=str(self.b2.id), **self.params)
        self.assertEqual(data_rows(widened), [])
        other = self.run_report(self.report_id, self.branch_user2, **self.params)
        self.assertEqual({r["approverName"] for r in data_rows(other)}, {"RL_H2 T"})

    def test_a_role_only_sees_the_workflows_it_can_open(self):
        body = self.run_report(self.report_id, self.leave_only, **self.params)
        self.assertEqual({r["module"] for r in data_rows(body)}, {"Leave"})
        self.assertTrue(any("Not shown because your role cannot open them" in n for n in body["notes"]))

    def grow(self, start, count):
        for i in range(start, start + count):
            emp = mk_emp(f"RL_G{i:02d}", self.cutting, self.b1)
            stamp(
                CasualLeaveRequest.objects.create(
                    employee=emp,
                    date=date(2026, 9, 21),
                    status="approved",
                    reviewed_by="Priya HR",
                    reviewer_role="hr",
                    reviewed_at=ist(2026, 9, 22, 9),
                ),
                created_at=ist(2026, 9, 21, 9),
            )
            stamp(
                MissingPunchRequest.objects.create(
                    employee=emp,
                    date=date(2026, 9, 8),
                    punch_time="09:00",
                    punch_type="IN",
                    reason="r",
                    status="approved",
                    hod_reviewed_by="RL_H1 T",
                    hod_reviewed_at=ist(2026, 9, 9, 9),
                    hr_reviewed_by="Priya HR",
                    hr_reviewed_at=ist(2026, 9, 10, 9),
                ),
                created_at=ist(2026, 9, 8, 9),
            )
            stamp(
                EmployeeRequest.objects.create(employee=emp, request_type="general", subject="s", description="d"),
                created_at=ist(2026, 9, 25, 9),
            )
            _leave(emp, "pending", created=ist(2026, 9, 26, 9))


# ══════════════════════════════════════════════════════════════════════════════
#  Employees with no department / designation / branch
# ══════════════════════════════════════════════════════════════════════════════


class OrphanEmployeeTests(ApprovalBase):
    """A record with no department, designation or branch (legacy imports) still appears for an unscoped viewer,
    under 'Unassigned', and never leaks to a branch-scoped one."""

    RANGE = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
    RANGE_WEEK = {"dateFrom": "2026-09-14", "dateTo": "2026-09-20"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cl = leave_type("CL", "Casual Leave", True)
        cls.o = mk_emp("RL_O", None, None, join_date="2024-01-01")
        _leave(cls.o, "approved", "2026-09-16", "2026-09-17", created=ist(2026, 9, 10, 9), by="Priya HR", role="hr")
        LeaveBalance.objects.create(employee=cls.o, leave_type=cls.cl, year=2026, allocated=12, used=2, remaining=10)
        stamp(
            CasualLeaveRequest.objects.create(
                employee=cls.o,
                date=date(2026, 9, 21),
                status="approved",
                reviewed_by="Priya HR",
                reviewer_role="hr",
                reviewed_at=ist(2026, 9, 22, 9),
            ),
            created_at=ist(2026, 9, 20, 9),
        )
        stamp(
            EmployeeRequest.objects.create(
                employee=cls.o, request_type="general", subject="Orphan ticket", description="d"
            ),
            created_at=ist(2026, 9, 12, 9),
        )
        stamp(
            EmployeePermission.objects.create(employee=cls.o, date=date(2026, 9, 25), status="pending"),
            created_at=ist(2026, 9, 20, 10),
        )

    def rows(self, rid, user=None, **p):
        return data_rows(self.run_report(rid, user, **p))

    def test_registers_show_the_employee_as_unassigned(self):
        reg = self.rows("leave-register", **self.RANGE)
        self.assertEqual(
            [(r["employeeCode"], r["department"], r["daysInPeriod"], r["advanceDays"]) for r in reg],
            [("RL_O", "Unassigned", 2.0, 6)],
        )
        cl = self.rows("casual-leave-register", **self.RANGE)
        self.assertEqual(
            [
                (r["employeeCode"], r["department"], r["designation"], r["serviceMonths"], r["turnaroundHours"])
                for r in cl
            ],
            [("RL_O", "Unassigned", None, 32, 48.0)],
        )
        er = self.rows("employee-requests-register", **self.RANGE)
        self.assertEqual([(r["employeeCode"], r["department"]) for r in er], [("RL_O", "Unassigned")])

    def test_monthly_grid_groups_them_under_an_unassigned_subtotal(self):
        body = self.run_report("leave-summary-monthly", year="2026")
        self.assertEqual(
            [(r.get("employeeCode"), r["employeeName"] if r.get("_kind") else r["leaveType"]) for r in body["rows"]],
            [("RL_O", "Casual (untyped)"), ("RL_O", "Casual Leave (CL system)"), (None, "Unassigned total")],
        )
        sub = body["rows"][-1]
        self.assertEqual((sub["_kind"], sub["sep"], sub["total"]), ("subtotal", 3.0, 3.0))  # 16-17 Sep + one casual day
        self.assertEqual(body["totals"]["total"], 3.0)

    def test_balance_statement_keeps_the_ledger_row_next_to_the_untyped_and_casual_system_rows(self):
        rows = self.rows("leave-balance-statement", year="2026")
        got = {(r["employeeCode"], r["leaveType"]): r for r in rows}
        self.assertEqual(
            set(got), {("RL_O", "Casual (untyped)"), ("RL_O", "Casual Leave"), ("RL_O", "Casual Leave (CL system)")}
        )
        self.assertEqual({r["department"] for r in rows}, {"Unassigned"})
        ledger = got[("RL_O", "Casual Leave")]
        self.assertEqual(
            (
                ledger["allocated"],
                ledger["ledgerUsed"],
                ledger["approvedDays"],
                ledger["computedRemaining"],
                ledger["variance"],
            ),
            (12.0, 2.0, 0.0, 12.0, 2.0),
        )
        self.assertEqual(got[("RL_O", "Casual (untyped)")]["approvedDays"], 2.0)

    def test_calendar_and_eligibility_list_an_unassigned_department(self):
        rows = self.rows("leave-department-calendar", **self.RANGE_WEEK)
        un = {r["date"]: r for r in rows if r["department"] == "Unassigned"}
        self.assertEqual(set(un), {f"2026-09-{d}" for d in range(14, 21)})
        got = [
            (un[f"2026-09-{d}"]["strength"], un[f"2026-09-{d}"]["onLeaveFull"], un[f"2026-09-{d}"]["leavePct"])
            for d in (14, 16, 17, 20)
        ]
        self.assertEqual(got, [(1, 0.0, 0.0), (1, 1.0, 100.0), (1, 1.0, 100.0), (1, None, None)])  # 20 Sep: Sunday off
        elig = self.rows("casual-leave-eligibility-usage", period="2026-09", employeeIds=str(self.o.id))
        self.assertEqual(
            [(r["department"], r["designation"], r["serviceMonths"], r["eligible"], r["monthsUsed"]) for r in elig],
            [("Unassigned", None, 32, "Used this month", "Sep")],
        )

    def test_pending_approvals_have_no_head_to_wait_on(self):
        rows = self.rows("pending-approvals-ageing")
        got = {r["module"]: r for r in rows}
        self.assertEqual(set(got), {"Permission", "Employee Request"})
        perm, req = got["Permission"], got["Employee Request"]
        self.assertEqual(
            (perm["department"], perm["pendingWith"], perm["hodAssigned"], perm["ageDays"]),
            ("Unassigned", "HR (no HOD assigned)", "No", 9),
        )
        self.assertEqual((req["pendingWith"], req["ageBucket"]), ("HR", "Over 15 days"))

    def test_workload_counts_their_decisions_and_pending_work(self):
        body = self.run_report("approver-workload-turnaround", **self.RANGE)
        got = {(r["approverName"], r["module"]): r for r in data_rows(body)}
        self.assertEqual(
            set(got),
            {
                ("Priya HR", "Leave"),
                ("Priya HR", "Casual Leave"),
                ("HR queue (shared)", "Permission"),
                ("HR queue (shared)", "Employee Request"),
            },
        )
        self.assertEqual(got[("Priya HR", "Casual Leave")]["avgTurnaroundHours"], 48.0)
        self.assertEqual(got[("HR queue (shared)", "Permission")]["pendingNow"], 1)

    def test_a_branch_scoped_user_never_sees_an_employee_with_no_branch(self):
        for rid, params in (
            ("leave-register", self.RANGE),
            ("casual-leave-register", self.RANGE),
            ("employee-requests-register", self.RANGE),
            ("leave-summary-monthly", {"year": "2026"}),
            ("leave-balance-statement", {"year": "2026"}),
            ("casual-leave-eligibility-usage", {"period": "2026-09"}),
            ("pending-approvals-ageing", {}),
        ):
            codes = {r["employeeCode"] for r in self.rows(rid, self.branch_user, **params)}
            self.assertNotIn("RL_O", codes, rid)
        self.assertEqual(self.rows("approver-workload-turnaround", self.branch_user, **self.RANGE), [])
        departments = {
            r["department"] for r in self.rows("leave-department-calendar", self.branch_user, **self.RANGE_WEEK)
        }
        self.assertNotIn("Unassigned", departments)


# ══════════════════════════════════════════════════════════════════════════════
#  Text date columns in every shape they have been stored
# ══════════════════════════════════════════════════════════════════════════════


class TextDateEdgeTests(Base):
    """LeaveRequest.start_date/end_date are TEXT: readable variants are placed by date, unreadable ones are counted
    in a note (never a crash, never silently dropped)."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()

        def leave(start, end, total="1.0"):
            return LeaveRequest.objects.create(
                employee=cls.a,
                type="casual",
                start_date=start,
                end_date=end,
                total_days=total,
                status="approved",
                approved_by="Priya HR",
                approver_role="hr",
            )

        cls.stamped = leave("2026-09-30T09:00:00", "2026-09-30T18:00:00")  # Wed, time suffix kept by an old client
        cls.padded = leave(" 2026-09-10", " 2026-09-11", "2.0")  # Thu-Fri, stray spaces
        cls.slashed = leave("14/09/2026", "14/09/2026")  # Mon, dd/mm/yyyy
        cls.ymd_slashed = leave("2026/09/15", "2026/09/15")  # Tue, yyyy/mm/dd
        cls.october_stamped = leave("2026-10-01T00:00:00", "2026-10-01T00:00:00")  # Thu
        cls.october = leave("2026-10-05", "2026-10-05")  # Mon
        cls.reversed = leave("2026-09-09", "2026-09-05")  # ends before it starts
        cls.blank_end = leave("2026-09-08", "")
        cls.gibberish = leave("N/A", "N/A")

    def ids_note(self, *requests):
        return "(ids: " + ", ".join(str(i) for i in sorted(r.id for r in requests)) + ")"

    def test_register_reads_every_readable_format_and_counts_the_rest(self):
        body = self.run_report("leave-register", dateFrom="2026-09-01", dateTo="2026-09-30")
        self.assertEqual(
            [(r["startDate"], r["endDate"], r["daysInPeriod"]) for r in data_rows(body)],
            [
                ("2026-09-10", "2026-09-11", 2.0),
                ("2026-09-14", "2026-09-14", 1.0),
                ("2026-09-15", "2026-09-15", 1.0),
                ("2026-09-30", "2026-09-30", 1.0),
            ],
        )
        note = next(n for n in body["notes"] if "unreadable" in n)
        self.assertIn("3 leave request(s)", note)
        self.assertIn(self.ids_note(self.reversed, self.blank_end, self.gibberish), note)

    def test_the_same_readable_rows_are_placed_in_october_and_unreadable_text_is_always_surfaced(self):
        body = self.run_report("leave-register", dateFrom="2026-10-01", dateTo="2026-10-31")
        self.assertEqual(
            [(r["startDate"], r["daysInPeriod"]) for r in data_rows(body)], [("2026-10-01", 1.0), ("2026-10-05", 1.0)]
        )
        # text that cannot be compared in SQL is fetched whatever the dates; the reversed pair is ISO, so SQL drops it
        note = next(n for n in body["notes"] if "unreadable" in n)
        self.assertIn("2 leave request(s)", note)
        self.assertIn(self.ids_note(self.blank_end, self.gibberish), note)

    def test_monthly_grid_splits_them_by_month(self):
        body = self.run_report("leave-summary-monthly", year="2026")
        (row,) = data_rows(body)
        self.assertEqual((row["sep"], row["oct"], row["total"]), (5.0, 2.0, 7.0))
        note = next(n for n in body["notes"] if "unreadable" in n)
        self.assertIn("3 approved leave request(s)", note)
        self.assertIn(self.ids_note(self.reversed, self.blank_end, self.gibberish), note)

    def test_balance_statement_charges_the_start_year_and_counts_an_unreadable_start(self):
        body = self.run_report("leave-balance-statement", year="2026")
        (row,) = data_rows(body)  # no ledger at all: one computed row
        self.assertEqual((row["leaveType"], row["allocated"], row["ledgerUsed"]), ("Casual (untyped)", None, None))
        # stored totals of every request whose START date is readable: 1 + 2 + 1 + 1 + 1 + 1 + 1 + 1
        self.assertEqual(row["approvedDays"], 9.0)
        note = next(n for n in body["notes"] if "unreadable" in n)
        self.assertIn("1 leave request(s)", note)
        self.assertIn(self.ids_note(self.gibberish), note)

    def test_calendar_marks_the_same_days(self):
        body = self.run_report("leave-department-calendar", dateFrom="2026-09-01", dateTo="2026-09-30")
        cut = {r["date"]: r for r in data_rows(body) if r["department"] == "CUTTING"}
        for d in ("2026-09-10", "2026-09-11", "2026-09-14", "2026-09-15", "2026-09-30"):
            self.assertEqual(cut[d]["onLeaveFull"], 1.0, d)
        self.assertEqual(cut["2026-09-16"]["onLeaveFull"], 0.0)
        note = next(n for n in body["notes"] if "unreadable" in n)
        self.assertIn("3 approved leave request(s)", note)


# ══════════════════════════════════════════════════════════════════════════════
#  Who a request waits on: one HOD per employee, and each workflow's own right
# ══════════════════════════════════════════════════════════════════════════════


class HodResolutionTests(ApprovalBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.leave = _leave(cls.a, created=ist(2026, 9, 28, 9))  # RL_A is in CUTTING, covered by RL_H1
        stamp(
            CasualLeaveRequest.objects.create(employee=cls.a, date=date(2026, 10, 14)), created_at=ist(2026, 9, 28, 9)
        )
        stamp(
            MissingPunchRequest.objects.create(
                employee=cls.a,
                date=date(2026, 9, 8),
                punch_time="09:05",
                punch_type="IN",
                reason="r",
                status="pending_hod",
            ),
            created_at=ist(2026, 9, 27, 9),
        )
        stamp(
            OnDutySession.objects.create(employee=cls.a, destination="Mill", branch=cls.b1, status="pending_hod"),
            created_at=ist(2026, 9, 27, 9),
        )
        stamp(
            AttendanceOverrideRequest.objects.create(employee=cls.a, date=date(2026, 9, 9), requested_by="Priya HR"),
            created_at=ist(2026, 9, 26, 9),
        )

    def waiting(self):
        rows = data_rows(self.run_report("pending-approvals-ageing", employeeIds=str(self.a.id)))
        return {r["module"]: (r["pendingWith"], r["hodAssigned"]) for r in rows}

    def pending_now(self):
        body = self.run_report("approver-workload-turnaround", dateFrom="2026-09-01", dateTo="2026-09-30")
        return {(r["approverName"], r["approverRole"], r["module"]): r["pendingNow"] for r in data_rows(body)}

    def test_the_department_head_is_the_default_owner(self):
        w = self.waiting()
        self.assertEqual(w["Leave"], ("HOD RL_H1 T or HR", "Yes"))
        self.assertEqual(w["Casual Leave"], ("HOD RL_H1 T or HR", "Yes"))
        self.assertEqual(w["On-Duty"], ("HOD RL_H1 T or HR", "Yes"))
        self.assertEqual(w["Missing Punch"], ("HOD RL_H1 T", "Yes"))
        self.assertEqual(w["Attendance Correction"], ("HOD RL_H1 T", "Yes"))
        self.assertEqual(self.pending_now()[("RL_H1 T", "HOD", "Leave")], 1)

    def test_an_individual_assignment_beats_the_department_and_the_workload_follows_it(self):
        ManagerEmployeeAssignment.objects.create(manager=self.h2, employee=self.a)
        w = self.waiting()
        self.assertEqual(w["Leave"], ("HOD RL_H2 T or HR", "Yes"))
        # RL_H2 owns RL_A now but may not act on missing punches, so nobody can decide that request
        self.assertEqual(
            w["Missing Punch"],
            ("No approver can act (HOD RL_H2 T lacks the approval right); HR cannot decide this step", "Yes"),
        )
        now = self.pending_now()
        self.assertEqual(now[("RL_H2 T", "HOD", "Leave")], 1)
        self.assertNotIn(("RL_H1 T", "HOD", "Leave"), now)

    def test_an_inactive_head_is_ignored_and_without_any_head_hr_takes_what_it_may(self):
        ManagerEmployeeAssignment.objects.create(manager=self.h2, employee=self.a)
        DepartmentManager.objects.filter(pk=self.h2.pk).update(is_active=False)
        self.assertEqual(self.waiting()["Leave"], ("HOD RL_H1 T or HR", "Yes"))  # falls back to the department's head
        DepartmentManager.objects.filter(pk=self.h1.pk).update(is_active=False)
        w = self.waiting()
        stuck = "No approver can act (no HOD assigned); HR cannot decide this step"
        self.assertEqual(w["Leave"], ("HR (no HOD assigned)", "No"))
        self.assertEqual(w["Casual Leave"], ("HR (no HOD assigned)", "No"))
        self.assertEqual(w["On-Duty"], ("HR (no HOD assigned)", "No"))  # HR may decide an on-duty request directly
        self.assertEqual(w["Missing Punch"], (stuck, "No"))
        self.assertEqual(w["Attendance Correction"], (stuck, "No"))

    def test_each_workflow_has_its_own_approval_right(self):
        DepartmentManager.objects.filter(pk=self.h1.pk).update(
            can_approve_leaves=False, can_approve_on_duty=False, can_approve_attendance=False
        )
        w = self.waiting()
        self.assertEqual(w["Leave"], ("HR (HOD RL_H1 T lacks the approval right)", "Yes"))
        self.assertEqual(w["On-Duty"], ("HR (HOD RL_H1 T lacks the approval right)", "Yes"))
        self.assertEqual(w["Casual Leave"], ("HOD RL_H1 T or HR", "Yes"))  # a right that was left switched on
        self.assertEqual(w["Missing Punch"], ("HOD RL_H1 T", "Yes"))
        self.assertEqual(
            w["Attendance Correction"],
            ("No approver can act (HOD RL_H1 T lacks the approval right); HR cannot decide this step", "Yes"),
        )
        body = self.run_report("pending-approvals-ageing", employeeIds=str(self.a.id))
        s = {x["label"]: x["value"] for x in body["summary"]}
        self.assertEqual(
            (s["Stuck - no approver can act"], s["With HR only"], s["With the HOD only"], s["HOD or HR"]), (1, 2, 1, 1)
        )
        # the head with no right to a workflow is not shown as owing a decision on it
        now = self.pending_now()
        self.assertNotIn(("RL_H1 T", "HOD", "Leave"), now)
        self.assertEqual(now[("RL_H1 T", "HOD", "Casual Leave")], 1)


# ══════════════════════════════════════════════════════════════════════════════
#  Per-module access to the blended approval reports
# ══════════════════════════════════════════════════════════════════════════════


def seed_every_request(cls):
    """One or two rows of every request workflow for RL_A, pending and decided, in September 2026."""
    p, h1 = "Priya HR", "RL_H1 T"
    _leave(cls.a, created=ist(2026, 9, 20, 9))
    _leave(cls.a, "approved", "2026-09-15", "2026-09-16", created=ist(2026, 9, 3, 9), by=p, role="hr")
    stamp(
        EmployeePermission.objects.create(employee=cls.a, date=date(2026, 9, 25), status="pending"),
        created_at=ist(2026, 9, 21, 9),
    )
    stamp(
        EmployeePermission.objects.create(
            employee=cls.a, date=date(2026, 9, 10), status="approved", approved_by=p, approver_role="hr"
        ),
        created_at=ist(2026, 9, 7, 9),
    )
    stamp(CasualLeaveRequest.objects.create(employee=cls.a, date=date(2026, 10, 14)), created_at=ist(2026, 9, 22, 9))
    stamp(
        CasualLeaveRequest.objects.create(
            employee=cls.a,
            date=date(2026, 9, 21),
            status="approved",
            reviewed_by=p,
            reviewer_role="hr",
            reviewed_at=ist(2026, 9, 22, 9),
        ),
        created_at=ist(2026, 9, 20, 9),
    )
    stamp(
        MissingPunchRequest.objects.create(
            employee=cls.a,
            date=date(2026, 9, 8),
            punch_time="09:05",
            punch_type="IN",
            reason="r",
            status="pending_hr",
            hod_reviewed_by=h1,
            hod_reviewed_at=ist(2026, 9, 9, 9),
        ),
        created_at=ist(2026, 9, 8, 9),
    )
    od = stamp(
        OnDutySession.objects.create(employee=cls.a, destination="Mill", branch=cls.b1), created_at=ist(2026, 9, 23, 9)
    )
    stamp(
        OnDutyPunchVerification.objects.create(
            session=od,
            employee=cls.a,
            punch_date=date(2026, 9, 23),
            punch_time="09:00",
            punch_type="IN",
            punch_number=1,
            latitude=Decimal("11.1"),
            longitude=Decimal("77.3"),
            photo="p.jpg",
        ),
        created_at=ist(2026, 9, 23, 9),
    )
    stamp(
        AttendanceOverrideRequest.objects.create(employee=cls.a, date=date(2026, 9, 9), requested_by=p),
        created_at=ist(2026, 9, 14, 9),
    )
    stamp(
        AttendanceOverrideRequest.objects.create(
            employee=cls.a,
            date=date(2026, 9, 10),
            status="approved",
            requested_by=p,
            reviewed_by=h1,
            reviewed_at=ist(2026, 9, 12, 10),
        ),
        created_at=ist(2026, 9, 11, 10),
    )
    stamp(
        OutpassRequest.objects.create(employee=cls.a, destination="Bank", reason="cash"), created_at=ist(2026, 9, 24, 8)
    )
    stamp(
        OutpassRequest.objects.create(
            employee=cls.a,
            destination="Post",
            reason="r",
            status="approved",
            approved_by=h1,
            approver_role="dept_head",
            approved_at=ist(2026, 9, 15, 10),
        ),
        created_at=ist(2026, 9, 15, 9),
    )
    stamp(
        EmployeeRequest.objects.create(employee=cls.a, request_type="general", subject="PF", description="d"),
        created_at=ist(2026, 9, 1, 9),
    )
    stamp(
        EmployeeRequest.objects.create(
            employee=cls.a,
            request_type="general",
            subject="Done",
            description="d",
            status="approved",
            handled_by=p,
            handled_at=ist(2026, 9, 3, 15),
        ),
        created_at=ist(2026, 9, 2, 15),
    )
    stamp(ResignationRequest.objects.create(employee=cls.a, reason="move"), created_at=ist(2026, 9, 16, 9))
    stamp(
        Advance.objects.create(employee=cls.a, advance_type="general", amount=Decimal("5000"), status="pending"),
        created_at=ist(2026, 9, 26, 9),
    )


ALL_WORKFLOWS = {
    "Leave",
    "Permission",
    "Casual Leave",
    "Missing Punch",
    "On-Duty",
    "On-Duty Punch",
    "Attendance Correction",
    "Outpass",
    "Employee Request",
    "Resignation",
    "Advance / Loan",
}


class ModuleAccessTests(ApprovalBase):
    RANGE = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        seed_every_request(cls)
        cls.only = {
            "leave": mk_hr("rl_m_leave", {"reports": "view", "leave": "view"}),
            "casual": mk_hr("rl_m_casual", {"reports": "view", "casual_leave": "view"}),
            "missing": mk_hr("rl_m_missing", {"reports": "view", "missing_punch": "view"}),
            "geo": mk_hr("rl_m_geo", {"reports": "view", "geo_attendance": "view"}),
            "settlement": mk_hr("rl_m_settlement", {"reports": "view", "settlement": "edit"}),
            "resignations": mk_hr(
                "rl_m_resign", {"reports": "view", "recruitment": "view"}
            ),  # a parent grant is inherited
        }

    def modules(self, rid, user, **p):
        return {r["module"] for r in data_rows(self.run_report(rid, user, **p))}

    def status(self, rid, user, **p):
        return self.get(f"/api/reports/run/{rid}", user, **p).status_code

    def test_the_pending_report_shows_each_role_only_the_workflows_it_owns(self):
        rid = "pending-approvals-ageing"
        self.assertEqual(self.modules(rid, self.admin), ALL_WORKFLOWS)
        self.assertEqual(self.modules(rid, self.requests_only), {"Permission", "Outpass", "Employee Request"})
        self.assertEqual(self.modules(rid, self.attendance_only), {"Attendance Correction"})
        self.assertEqual(self.modules(rid, self.only["leave"]), {"Leave"})
        self.assertEqual(self.modules(rid, self.only["casual"]), {"Casual Leave"})
        self.assertEqual(self.modules(rid, self.only["missing"]), {"Missing Punch"})
        self.assertEqual(self.modules(rid, self.only["geo"]), {"On-Duty", "On-Duty Punch"})
        self.assertEqual(self.modules(rid, self.only["settlement"]), {"Advance / Loan"})
        self.assertEqual(self.modules(rid, self.only["resignations"]), {"Resignation"})

    def test_the_workload_report_shows_each_role_only_the_workflows_it_owns(self):
        rid = "approver-workload-turnaround"
        self.assertEqual(
            self.modules(rid, self.requests_only, **self.RANGE), {"Permission", "Outpass", "Employee Request"}
        )
        self.assertEqual(self.modules(rid, self.attendance_only, **self.RANGE), {"Attendance Correction"})
        self.assertEqual(self.modules(rid, self.only["casual"], **self.RANGE), {"Casual Leave"})
        # resignations and advances are not approver-workload workflows, so a role that owns only those gets nothing
        self.assertEqual(self.modules(rid, self.only["settlement"], **self.RANGE), set())
        self.assertEqual(self.modules(rid, self.only["resignations"], **self.RANGE), set())
        admin = self.modules(rid, self.admin, **self.RANGE)
        self.assertTrue(
            {"Leave", "Permission", "Casual Leave", "Attendance Correction", "Outpass", "Employee Request"} <= admin
        )
        self.assertNotIn("Resignation", admin)

    def test_direct_report_access_follows_the_owning_module(self):
        for rid in (
            "leave-register",
            "leave-balance-statement",
            "leave-summary-monthly",
            "holiday-list",
            "casual-leave-register",
            "casual-leave-eligibility-usage",
        ):
            self.assertEqual(self.status(rid, self.requests_only), 403, rid)
        self.assertEqual(self.status("employee-requests-register", self.requests_only), 200)
        self.assertEqual(self.status("holiday-list", self.leave_only), 200)
        self.assertEqual(self.status("employee-requests-register", self.leave_only), 403)
        self.assertEqual(self.status("casual-leave-register", self.leave_only), 403)
        self.assertEqual(self.status("casual-leave-register", self.only["casual"]), 200)
        self.assertEqual(self.status("casual-leave-eligibility-usage", self.only["casual"]), 200)
        self.assertEqual(self.status("leave-register", self.only["casual"]), 403)
        hidden = mk_hr("rl_m_hidden", {"reports": "view", "leave": "hidden"})
        self.assertEqual(self.status("leave-register", hidden), 403)
        editor = mk_hr("rl_m_editor", {"reports": "view", "leave": "edit"})
        self.assertEqual(self.status("leave-register", editor), 200)

    def test_the_calendar_opens_for_any_one_of_its_modules_and_fills_only_the_columns_that_role_may_see(self):
        params = {"dateFrom": "2026-09-14", "dateTo": "2026-09-16"}
        for user in (self.attendance_only, self.requests_only, self.only["casual"], self.only["geo"]):
            self.assertEqual(self.status("leave-department-calendar", user, **params), 200)
        rows = data_rows(self.run_report("leave-department-calendar", self.attendance_only, **params))
        cut = {r["date"]: r for r in rows if r["department"] == "CUTTING"}["2026-09-15"]
        self.assertEqual(cut["absentUnexplained"], 0.0)
        for key in ("onLeaveFull", "onLeaveHalf", "casualLeave", "onDuty", "permission", "leavePct"):
            self.assertIsNone(cut[key], key)


# ══════════════════════════════════════════════════════════════════════════════
#  Reading never changes anything
# ══════════════════════════════════════════════════════════════════════════════


class NoWritesTests(ApprovalBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cl = leave_type("CL", "Casual Leave", True)
        LeaveBalance.objects.create(employee=cls.a, leave_type=cls.cl, year=2026, allocated=12, used=1, remaining=11)
        Holiday.objects.create(name="Ganesh", date=date(2026, 9, 10))
        seed_every_request(cls)
        # The company letterhead of every export reads the PayrollSettings singleton, which PayrollSettings.get()
        # creates on first use. It always exists in a live system; create it up front so the test isolates reports.
        PayrollSettings.get()

    def snapshot(self):
        models = [
            Employee,
            LeaveType,
            LeaveBalance,
            Holiday,
            LeaveRequest,
            CasualLeaveRequest,
            EmployeePermission,
            EmployeeRequest,
            MissingPunchRequest,
            OnDutySession,
            OnDutyPunchVerification,
            AttendanceOverrideRequest,
            AttendanceDayRecord,
            OutpassRequest,
            ResignationRequest,
            Advance,
            DepartmentManager,
        ]
        return {m.__name__: list(m.objects.order_by("pk").values()) for m in models}

    def table_counts(self):
        return {
            m.__name__: m.objects.count() for m in apps.get_app_config("api").get_models() if m.__name__ != "AuditLog"
        }

    def test_screen_excel_and_pdf_of_all_ten_reports_change_nothing(self):
        params = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30", "year": "2026", "period": "2026-09"}
        before, counts = self.snapshot(), self.table_counts()
        for rid in MY_IDS:
            self.assertEqual(self.get(f"/api/reports/run/{rid}", **params).status_code, 200, rid)
            for fmt in ("xlsx", "pdf"):
                self.assertEqual(self.get(f"/api/reports/export/{rid}", fmt=fmt, **params).status_code, 200, (rid, fmt))
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(self.table_counts(), counts)
