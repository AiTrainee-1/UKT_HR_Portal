"""
Report Center - Employees & Manpower reports (group ``employees_master``).

Covers: employee-master, employee-contact-list, strength-statement, new-joinings, exits-register,
resignation-register, manpower-movement, statutory-compliance, data-quality-audit, birthdays,
work-anniversaries, service-milestones, age-compliance, workforce-profile, family-dependents, org-structure.

"Today" is pinned to 2026-09-29 (a Tuesday) so every expected number is hand-derived and deterministic.

Run via: python manage.py test api.tests_reporting_employees_master -v 2
"""

import io
from datetime import date, datetime, timedelta
from datetime import timezone as dt_tz
from decimal import Decimal
from unittest.mock import patch

from django.db import connection
from django.test import SimpleTestCase, TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .casual_leave_views import ELIGIBILITY_MONTHS, _service_months
from .jwt_utils import sign_token
from .models import (
    Branch,
    Department,
    DepartmentHeadcount,
    DepartmentManager,
    Designation,
    Employee,
    EmployeeDocument,
    FamilyDependent,
    HRUser,
    ManagerDepartmentAssignment,
    ResignationRequest,
    Role,
    SalarySlip,
)
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import employees_master_base as B
from .reporting.definitions import employees_master_compliance as C

TODAY = date(2026, 9, 29)  # Tuesday
MY_REPORTS = (
    "employee-master", "employee-contact-list", "strength-statement", "new-joinings", "exits-register",
    "resignation-register", "manpower-movement", "statutory-compliance", "data-quality-audit", "birthdays",
    "work-anniversaries", "service-milestones", "age-compliance", "workforce-profile", "family-dependents",
    "org-structure",
)  # fmt: skip


def utc(y, m, d, hh=0, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=dt_tz.utc)


def hdr(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def codes(body, key="employeeCode"):
    return [r[key] for r in body["rows"] if r.get("_kind") not in ("subtotal", "total")]


def cards(body):
    return {c["label"]: c["value"] for c in body["summary"]}


def data_rows(body):
    return [r for r in body["rows"] if r.get("_kind") not in ("subtotal", "total")]


class _World(TestCase):
    """Ten employees across two units (plus one legacy record with no branch / department)."""

    @classmethod
    def setUpTestData(cls):
        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.b3 = Branch.objects.create(name="Old Unit", code="OLD", is_active=False)
        cls.d_cut1 = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.d_sew1 = Department.objects.create(name="SEWING", branch=cls.b1)
        cls.d_cut2 = Department.objects.create(name="CUTTING", branch=cls.b2)  # same name, other unit
        cls.d_legacy = Department.objects.create(name="LEGACY", branch=None)
        cls.d_scrap = Department.objects.create(name="SCRAP", branch=cls.b3)
        cls.des_mgr = Designation.objects.create(title="Manager", department=cls.d_cut1, level="manager")
        cls.des_op = Designation.objects.create(title="Operator", department=cls.d_sew1, level="junior")
        cls.des_sup = Designation.objects.create(title="Supervisor", department=cls.d_sew1, level="senior")

        def mk(code, first, last, **kw):
            return Employee.objects.create(employee_code=code, first_name=first, last_name=last, **kw)

        cls.e1 = mk(
            "E1", "Arun", "Kumar", employment_type="staff", gender="male", branch=cls.b1, department=cls.d_cut1,
            designation=cls.des_mgr, join_date="2024-01-15", date_of_birth=date(1990, 3, 10), phone="9876543210",
            email="arun@example.com", emergency_contact="Ravi 9000000001", blood_group="O+", father_name="Kumar",
            address="12 Main St", salary_amount=Decimal("30000"), unit_code="U1-1", pf_number="TN/TPR/001",
            esi_number="1234567890", uan_number="100000000001", bank_name="HDFC", bank_account="5566778899001",
            bank_ifsc="HDFC0001234", id_proof="123412341234",
        )  # fmt: skip
        cls.e2 = mk(
            "E2", "Bala", "Devi", employment_type="production", gender="female", branch=cls.b1, department=cls.d_sew1,
            designation=cls.des_op, join_date="2026-09-05", date_of_birth=date(2001, 9, 29), phone="9876543211",
            blood_group="b+", address="45 Lake Rd", salary_per_shift=Decimal("500"), salary_type="weekly",
        )  # fmt: skip
        cls.e3 = mk(
            "E3", "Chitra", "Raj", employment_type="staff", gender="female", branch=cls.b2, department=cls.d_cut2,
            join_date="05-09-2026", date_of_birth=date(2000, 2, 29), phone="+91 98765 43212", father_name="Raj",
            pf_number="TN/TPR/002", esi_number="12345", uan_number="12345678901", bank_account="ABC123",
            bank_ifsc="BADIFSC",
        )  # fmt: skip
        cls.e4 = mk(
            "E4", "Dinesh", "Mani", employment_type="staff", gender="male", status="inactive", branch=cls.b1,
            department=cls.d_sew1, designation=cls.des_sup, join_date="2020-03-01", date_of_birth=date(1985, 7, 20),
            phone="9876543214", salary_amount=Decimal("25000"),
        )  # fmt: skip
        cls.e5 = mk(
            "E5", "Elango", "Raja", employment_type="production", gender="male", status="inactive", branch=cls.b1,
            department=cls.d_sew1, designation=cls.des_op, join_date="2025-06-10", date_of_birth=date(1992, 1, 1),
            phone="9876543215", salary_per_shift=Decimal("450"), salary_type="weekly",
        )  # fmt: skip
        cls.e6 = mk(
            "E6", "Farida", "Banu", employment_type="staff", gender="female", branch=None, department=None,
            join_date="2023-11-30", date_of_birth=date(1995, 12, 31), phone=None, address="9 Park St",
            father_name="Banu", blood_group="A+", emergency_contact="Salim 9000000006", salary_amount=Decimal("0"),
        )  # fmt: skip
        cls.e7 = mk(
            "E7", "Gopal", "Nair", employment_type="staff", gender=None, branch=cls.b1, department=cls.d_sew1,
            designation=cls.des_op, join_date="soon", date_of_birth=None, phone=None, salary_amount=Decimal("20000"),
            pf_number="PF-DUP-9",
        )  # fmt: skip
        cls.e8 = mk(
            "E8", "Hari", "Om", employment_type="production", gender="Male", branch=cls.b2, department=cls.d_cut2,
            join_date="2026-08-31", date_of_birth=date(2010, 5, 5), phone="9876543218", blood_group="AB-",
            esi_number="12345678901234567", uan_number="100000000008", bank_account="123456789",
            bank_ifsc="sbin0001234", salary_type="weekly",
        )  # fmt: skip
        cls.e9 = mk(
            "E9", "Indu", "Shree", employment_type="staff", gender="female", status="resigned", branch=cls.b2,
            department=cls.d_cut2, join_date="2019-01-01", date_of_birth=date(1980, 11, 11), phone="9876543219",
            salary_amount=Decimal("22000"),
        )  # fmt: skip
        cls.e10 = mk(
            "E10", "Jai", "Singh", employment_type="staff", gender="male", branch=cls.b1, department=cls.d_cut1,
            join_date="", date_of_birth=date(1960, 9, 1), phone="", address="7 Hill Rd", father_name="Singh",
            blood_group=" o+ ", salary_amount=Decimal("18000"), pf_number="pf dup 9",
        )  # fmt: skip
        # updated_at is auto_now: only .update() can pin it. It stands in for the (undated) manual-deactivation date.
        Employee.objects.filter(pk=cls.e5.pk).update(updated_at=utc(2026, 9, 10, 8, 0))
        Employee.objects.filter(pk=cls.e9.pk).update(updated_at=utc(2026, 1, 15, 20, 0))  # = 16-Jan 01:30 IST

        # HODs: E1 heads SEWING; E10 heads CUTTING (Unit 1); an INACTIVE HOD row covers CUTTING (Unit 2).
        m1 = DepartmentManager.objects.create(employee=cls.e1)
        ManagerDepartmentAssignment.objects.create(manager=m1, department=cls.d_sew1)
        m10 = DepartmentManager.objects.create(employee=cls.e10)
        ManagerDepartmentAssignment.objects.create(manager=m10, department=cls.d_cut1)
        m3 = DepartmentManager.objects.create(employee=cls.e3, is_active=False)
        ManagerDepartmentAssignment.objects.create(manager=m3, department=cls.d_cut2)
        DepartmentHeadcount.objects.create(department=cls.d_cut1, required_count=5)
        DepartmentHeadcount.objects.create(department=cls.d_sew1, required_count=1)
        DepartmentHeadcount.objects.create(department=cls.d_cut2, required_count=0)

        # Documents on file
        for cat in ("aadhaar_card", "pan_card", "bank_passbook"):
            cls.doc(cls.e1, cat)
        cls.doc(cls.e3, "aadhaar_card")

        # Payroll evidence for August 2026 (E8 is production: two weekly slips)
        SalarySlip.objects.create(
            employee=cls.e1, month=8, year=2026, slip_number="S-E1-08", pf_deduction=1800, esi_deduction=225
        )
        SalarySlip.objects.create(
            employee=cls.e8, month=8, year=2026, week_number=1, slip_number="S-E8-08-1", pf_deduction=60
        )
        SalarySlip.objects.create(
            employee=cls.e8, month=8, year=2026, week_number=2, slip_number="S-E8-08-2", pf_deduction=40
        )

        # Resignations
        def resign(emp, created, **kw):
            r = ResignationRequest.objects.create(employee=emp, **kw)
            ResignationRequest.objects.filter(pk=r.pk).update(created_at=created)
            return r

        cls.r1 = resign(
            cls.e4, utc(2026, 8, 10, 4, 0), reason="Better opportunity", last_working_date=date(2026, 8, 31),
            survey_q1_answer="Salary", survey_q2_answer="Yes", survey_q3_answer="Higher pay", status="approved",
            dept_head=cls.e1, dept_head_status="approved", dept_head_comment="Fine",
            dept_head_approved_at=utc(2026, 8, 12, 5, 0), hr_comment="OK", approved_by="HR Admin",
            approved_at=utc(2026, 8, 20, 10, 0),
        )  # fmt: skip
        cls.r2 = resign(
            cls.e1,
            utc(2026, 9, 20, 20, 0),
            reason="Relocation",
            last_working_date=date(2026, 10, 31),
            status="pending",
        )  # 21-Sep 01:30 IST
        cls.r3 = resign(
            cls.e3, utc(2026, 9, 1, 5, 0), reason="Health", status="rejected", rejected_by="dept_head",
            dept_head=cls.e10, dept_head_status="rejected", dept_head_approved_at=utc(2026, 9, 3, 5, 0),
        )  # fmt: skip
        cls.r4 = resign(
            cls.e7, utc(2026, 9, 15, 6, 0), reason="Studies", status="dept_approved", dept_head=cls.e1,
            dept_head_status="approved",
        )  # fmt: skip
        cls.r5 = resign(
            cls.e6, utc(2026, 9, 25, 6, 0), reason="Personal", status="rejected", rejected_by="hr",
            hr_comment="Not now",
        )  # fmt: skip

        # Users
        full = {
            "reports": "view",
            "employees": "view",
            "recruitment": "view",
            "casual_leave": "view",
            "payroll": "view",
        }
        cls.admin = HRUser.objects.create(username="em_admin", password_hash="x", is_super_admin=True)
        cls.role_full = Role.objects.create(name="em_full", permissions=full)
        cls.branch_user = HRUser.objects.create(username="em_b1", password_hash="x", role=cls.role_full, branch=cls.b1)
        cls.plain_user = HRUser.objects.create(
            username="em_plain",
            password_hash="x",
            role=Role.objects.create(name="em_plain", permissions={"reports": "view"}),
        )
        cls.no_salary_user = HRUser.objects.create(
            username="em_nosal",
            password_hash="x",
            role=Role.objects.create(
                name="em_nosal", permissions={"reports": "view", "employees": "view", "recruitment": "view"}
            ),
        )

    @staticmethod
    def doc(emp, category):
        return EmployeeDocument.objects.create(
            employee=emp, category=category, file="employee_documents/x.pdf", original_filename="x.pdf"
        )

    def setUp(self):
        p = patch("api.reporting.filters.ist_today", return_value=TODAY)
        self.mock_today = p.start()
        self.addCleanup(p.stop)

    # -- helpers ---------------------------------------------------------------------------
    def get(self, rid, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/run/{rid}", params, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, rid, fmt, user=None, expect=200, **params):
        r = self.client.get(f"/api/reports/export/{rid}", {"fmt": fmt, **params}, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r

    def make(self, code, **kw):
        base = dict(first_name=code, last_name="Extra", employment_type="staff", status="active")
        base.update(kw)
        return Employee.objects.create(employee_code=code, **base)


# ═════════════════════════════════════════════════════════════════════════════
# Pure helpers
# ═════════════════════════════════════════════════════════════════════════════


class HelperTests(SimpleTestCase):
    def test_natural_code_order(self):
        codes_ = ["E10", "E2", "E1", "A9", "E02", "e3"]
        self.assertEqual(sorted(codes_, key=B.code_key), ["A9", "E1", "E02", "E2", "e3", "E10"])

    def test_age_and_tenure(self):
        self.assertEqual(B.age_on(date(2001, 9, 29), TODAY), 25)  # birthday today counts
        self.assertEqual(B.age_on(date(2001, 9, 30), TODAY), 24)
        self.assertEqual(B.age_on(date(2000, 2, 29), date(2026, 2, 28)), 25)  # leap-day baby, day before
        self.assertIsNone(B.age_on(date(2030, 1, 1), TODAY))  # future DOB
        self.assertIsNone(B.age_on(None, TODAY))
        self.assertEqual(B.months_between(date(2024, 1, 15), TODAY), 32)
        self.assertEqual(B.months_between(date(2023, 11, 30), TODAY), 33)
        self.assertEqual(B.tenure_text(32), "2y 8m")
        self.assertEqual(B.tenure_text(24), "2y")
        self.assertEqual(B.tenure_text(5), "5m")
        self.assertEqual(B.tenure_text(0), "<1m")
        self.assertIsNone(B.tenure_text(None))

    def test_add_months_clamps_to_month_end(self):
        self.assertEqual(B.add_months(date(2026, 8, 31), 6), date(2027, 2, 28))
        self.assertEqual(B.add_months(date(2023, 11, 30), 3), date(2024, 2, 29))
        self.assertEqual(B.add_months(date(2026, 11, 15), 3), date(2027, 2, 15))

    def test_casual_leave_date_agrees_with_the_casual_leave_engine(self):
        """clEligibleFrom must be the first day the app's own service-month rule reaches 6 months."""
        for joined in (
            date(2026, 3, 31), date(2026, 8, 31), date(2025, 8, 30), date(2024, 2, 29), date(2026, 1, 15),
            date(2026, 5, 31), date(2025, 12, 31),
        ):  # fmt: skip
            when = B.months_reached_on(joined, ELIGIBILITY_MONTHS)
            emp = Employee(join_date=joined.isoformat())
            self.assertGreaterEqual(_service_months(emp, when), ELIGIBILITY_MONTHS, joined)
            self.assertLess(_service_months(emp, when - timedelta(days=1)), ELIGIBILITY_MONTHS, joined)

    def test_month_arithmetic_saturates_instead_of_overflowing(self):
        self.assertEqual(B.add_months(date(9999, 12, 31), 6), date.max)
        self.assertEqual(B.add_months(date(9999, 11, 30), 24), date.max)
        self.assertEqual(B.months_reached_on(date(9999, 12, 31), 6), date.max)
        self.assertEqual(B.months_reached_on(date(9999, 8, 31), 6), date.max)

    def test_a_placeholder_join_date_is_unreadable_not_a_date(self):
        for raw in ("9999-12-31", "1900-01-01", "0001-01-01", "2101-01-01", "soon"):
            self.assertEqual(B.join_date_of(Employee(join_date=raw)), (None, "unreadable"), raw)
        self.assertEqual(B.join_date_of(Employee(join_date=" ")), (None, "missing"))
        self.assertEqual(B.join_date_of(Employee(join_date="05-09-2026")), (date(2026, 9, 5), "ok"))
        self.assertEqual(B.join_date_of(Employee(join_date="1950-01-01")), (date(1950, 1, 1), "ok"))

    def test_identifier_placeholders(self):
        for raw in ("-", "--", " ", "NA", "n/a", "N.A.", "NIL", "None", "null", "0", "0000000000", "00-00", "not applicable",
                    "Not Available", "TBD", "  na "):  # fmt: skip
            self.assertTrue(C._is_placeholder(raw), raw)
            self.assertIsNone(C._real_id(raw), raw)
            self.assertEqual(C._norm_id(raw), "", raw)
        for raw in ("TN/TPR/001", "1234567890", "PF-DUP-9", "100000000001", "NAIDU1", "0123456789", "N1"):
            self.assertFalse(C._is_placeholder(raw), raw)
            self.assertEqual(C._real_id(raw), raw)
        self.assertEqual(C._norm_id("pf dup-9"), "PFDUP9")
        self.assertIsNone(C._real_id(None))

    def test_feb_29_anniversaries_and_year_wrap(self):
        self.assertEqual(B.occurrence_in_year(2, 29, 2027), date(2027, 2, 28))
        self.assertEqual(B.occurrence_in_year(2, 29, 2028), date(2028, 2, 29))
        start, end = date(2026, 12, 20), date(2027, 1, 18)
        self.assertEqual(B.occurrence_in_window(12, 31, start, end), date(2026, 12, 31))
        self.assertEqual(B.occurrence_in_window(1, 5, start, end), date(2027, 1, 5))
        self.assertIsNone(B.occurrence_in_window(1, 19, start, end))
        self.assertEqual(B.occurrence_in_window(2, 29, date(2026, 2, 1), date(2026, 2, 28)), date(2026, 2, 28))

    def test_window_choices(self):
        self.assertEqual(B.window_bounds("thisMonth", TODAY), (date(2026, 9, 1), date(2026, 9, 30)))
        self.assertEqual(B.window_bounds("nextMonth", date(2026, 12, 5)), (date(2027, 1, 1), date(2027, 1, 31)))
        self.assertEqual(B.window_bounds("next7", TODAY), (TODAY, date(2026, 10, 5)))
        self.assertEqual(B.window_bounds("next30", TODAY), (TODAY, date(2026, 10, 28)))
        self.assertEqual(B.window_bounds("m02", TODAY), (date(2026, 2, 1), date(2026, 2, 28)))

    def test_masking_and_blood_group(self):
        self.assertEqual(B.mask_tail("1234567890123"), "*********0123")
        self.assertEqual(B.mask_tail("12 34"), "****")
        self.assertEqual(B.mask_tail("abc"), "***")
        self.assertIsNone(B.mask_tail("  "))
        self.assertIsNone(B.mask_tail(None))
        self.assertEqual(B.blood_group_of(" o+ "), "O+")
        self.assertEqual(B.blood_group_of("ab -"), "AB-")
        self.assertIsNone(B.blood_group_of("Z"))
        self.assertIsNone(B.blood_group_of(None))

    def test_gender_and_status_buckets(self):
        self.assertEqual([B.gender_bucket(g) for g in ("Male", " FEMALE ", "other", "Non-binary", "", None)],
                         ["male", "female", "other", "other", "unspecified", "unspecified"])  # fmt: skip
        self.assertEqual(B.gender_label("male"), "Male")
        self.assertEqual(B.gender_label("Non-binary"), "Non-binary")
        self.assertIsNone(B.gender_label(" "))
        self.assertEqual(B.status_label("resigned"), "Resigned")
        self.assertEqual(B.status_label(""), "Unknown")

    def test_ist_bounds_use_the_factory_day(self):
        start, end = B.ist_bounds(date(2026, 9, 21), date(2026, 9, 21))
        self.assertEqual(
            start, datetime(2026, 9, 20, 18, 30, tzinfo=dt_tz.utc)
        )  # IST midnight = 18:30 UTC previous day
        self.assertEqual(end - start, timedelta(days=1))
        self.assertEqual(B.ist_date(utc(2026, 9, 20, 20, 0)), date(2026, 9, 21))


# ═════════════════════════════════════════════════════════════════════════════
# Catalog / access
# ═════════════════════════════════════════════════════════════════════════════


class CatalogAndAccessTests(_World):
    def test_all_sixteen_reports_are_registered_in_the_employees_category(self):
        specs = {s.id: s for s in registry.all_specs()}
        self.assertEqual(registry.LOAD_ERRORS, {})
        for rid in MY_REPORTS:
            self.assertIn(rid, specs)
            s = specs[rid]
            self.assertEqual(s.category, "employees", rid)
            self.assertTrue(s.modules, rid)
            for m in s.modules:
                self.assertIn(m, all_module_keys(), rid)
            self.assertFalse(s.super_admin_only, rid)

    def test_celebrations_family_has_both_variants(self):
        specs = {s.id: s for s in registry.all_specs()}
        self.assertEqual((specs["birthdays"].family, specs["birthdays"].variant), ("celebrations", "Birthdays"))
        self.assertEqual(specs["work-anniversaries"].family, "celebrations")
        self.assertNotEqual(specs["birthdays"].variant, specs["work-anniversaries"].variant)

    def test_catalog_lists_them_for_an_authorised_role_and_hides_them_otherwise(self):
        ids = {r["id"] for r in self.client.get("/api/reports/catalog", **hdr(self.branch_user)).json()["reports"]}
        self.assertTrue(set(MY_REPORTS) <= ids)
        ids = {r["id"] for r in self.client.get("/api/reports/catalog", **hdr(self.plain_user)).json()["reports"]}
        self.assertFalse(set(MY_REPORTS) & ids)

    def test_a_reports_only_role_gets_403_report_forbidden_on_run_and_export(self):
        for rid in MY_REPORTS:
            r = self.client.get(f"/api/reports/run/{rid}", **hdr(self.plain_user))
            self.assertEqual(r.status_code, 403, rid)
            self.assertEqual(r.json()["error"], "report_forbidden", rid)
            for fmt in ("xlsx", "pdf"):
                r = self.client.get(f"/api/reports/export/{rid}", {"fmt": fmt}, **hdr(self.plain_user))
                self.assertEqual(r.status_code, 403, (rid, fmt))
                self.assertEqual(r.json()["error"], "report_forbidden", (rid, fmt))

    def test_owning_module_grants(self):
        def user(name, perms):
            return HRUser.objects.create(
                username=name,
                password_hash="x",
                role=Role.objects.create(name=name, permissions={"reports": "view", **perms}),
            )

        payroll_only = user("em_pay_only", {"payroll": "view"})
        resign_only = user("em_res_only", {"recruitment.resignations": "view"})
        employees_only = user("em_emp_only", {"employees": "view"})
        # Statutory numbers: payroll, salary or employees access is enough.
        self.get("statutory-compliance", user=payroll_only)
        self.get("statutory-compliance", user=employees_only)
        self.get("employee-master", user=payroll_only, expect=403)
        # Resignation requests belong to the recruitment.resignations module only ...
        self.get("resignation-register", user=resign_only)
        self.get("resignation-register", user=employees_only, expect=403)
        # ... while exits are visible with either the employees or the resignations module.
        self.get("exits-register", user=resign_only)
        self.get("exits-register", user=employees_only)
        # New joinings: employees or recruitment.new_joinees
        self.get("new-joinings", user=user("em_nj", {"recruitment.new_joinees": "view"}))


# ═════════════════════════════════════════════════════════════════════════════
# employee-master
# ═════════════════════════════════════════════════════════════════════════════


class EmployeeMasterTests(_World):
    RID = "employee-master"

    def test_default_lists_active_employees_in_natural_code_order(self):
        body = self.get(self.RID)
        # E10 comes after E8 (a plain string sort would put it right after E1)
        self.assertEqual(codes(body), ["E1", "E2", "E3", "E6", "E7", "E8", "E10"])
        self.assertIsNone(body["totals"])
        self.assertEqual(
            cards(body),
            {
                "Total employees": 7, "Active": 7, "Inactive / left": 0, "Staff": 5, "Production": 2, "Male": 3,
                "Female": 3, "Other / not set": 1,
            },
        )  # fmt: skip

    def test_golden_row_values_and_standard_columns(self):
        body = self.get(self.RID)
        keys = [c["key"] for c in body["columns"]]
        self.assertEqual(
            keys,
            ["employeeCode", "employeeName", "gender", "age", "department", "designation", "branch", "employmentType",
             "joinDate", "tenure", "status", "phone"],
        )  # fmt: skip
        rows = {r["employeeCode"]: r for r in body["rows"]}
        e1 = rows["E1"]
        self.assertEqual(e1["employeeName"], "Arun Kumar")
        self.assertEqual(
            (e1["gender"], e1["age"], e1["department"], e1["designation"]), ("Male", 36, "CUTTING", "Manager")
        )
        self.assertEqual(
            (e1["branch"], e1["employmentType"], e1["joinDate"], e1["tenure"]),
            ("Unit 1", "Staff", "2024-01-15", "2y 8m"),
        )
        self.assertEqual(rows["E2"]["age"], 25)  # birthday is today
        self.assertEqual(rows["E3"]["joinDate"], "2026-09-05")  # dd-mm-yyyy text parsed
        self.assertEqual(rows["E3"]["tenure"], "<1m")
        self.assertEqual(rows["E3"]["age"], 26)  # born 29-Feb-2000
        self.assertEqual(rows["E3"]["designation"], None)
        self.assertEqual(rows["E6"]["department"], "Unassigned")
        self.assertEqual(rows["E6"]["branch"], "No branch")
        self.assertIsNone(rows["E7"]["joinDate"])  # 'soon' is unreadable
        self.assertIsNone(rows["E7"]["age"])
        self.assertEqual(rows["E7"]["gender"], None)
        self.assertEqual(rows["E8"]["gender"], "Male")  # 'Male' typed with a capital

    def test_full_layout_adds_contact_fields_and_the_hod(self):
        body = self.get(self.RID, layout="full")
        self.assertEqual(len(body["columns"]), 19)
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["E1"]["unitCode"], "U1-1")
        self.assertEqual(rows["E1"]["fatherName"], "Kumar")
        self.assertEqual(rows["E1"]["email"], "arun@example.com")
        self.assertEqual(rows["E1"]["emergencyContact"], "Ravi 9000000001")
        self.assertEqual(rows["E1"]["bloodGroup"], "O+")
        self.assertEqual(rows["E1"]["dateOfBirth"], "1990-03-10")
        # one active HOD per employee: E1 heads SEWING, E10 heads CUTTING (Unit 1)
        self.assertEqual(rows["E2"]["hod"], "Arun Kumar")
        self.assertEqual(rows["E7"]["hod"], "Arun Kumar")
        self.assertEqual(rows["E1"]["hod"], "Jai Singh")
        self.assertIsNone(rows["E10"]["hod"])  # the HOD of the department has nobody above them
        self.assertIsNone(rows["E3"]["hod"])  # the HOD row covering CUTTING (Unit 2) is inactive
        self.assertIsNone(rows["E6"]["hod"])
        self.assertTrue(any("HOD" in n for n in body["notes"]))

    def test_an_individual_hod_assignment_beats_the_department_one(self):
        from .models import ManagerEmployeeAssignment

        ManagerEmployeeAssignment.objects.create(
            manager=DepartmentManager.objects.get(employee=self.e10), employee=self.e2
        )
        rows = {r["employeeCode"]: r["hod"] for r in self.get(self.RID, layout="full")["rows"]}
        self.assertEqual(rows["E2"], "Jai Singh")  # E1 still holds SEWING as a whole
        self.assertEqual(rows["E7"], "Arun Kumar")

    def test_sensitive_columns_never_appear(self):
        for layout in ("standard", "full"):
            body = self.get(self.RID, layout=layout, employeeStatus="all")
            keys = {c["key"] for c in body["columns"]}
            forbidden = {
                "salary", "salaryAmount", "salaryPerShift", "salaryType", "bankName", "bankAccount", "bankIfsc", "idProof",
                "address", "pfNumber", "esiNumber", "uanNumber", "passwordHash", "photoUrl",
            }  # fmt: skip
            self.assertFalse(keys & forbidden, keys & forbidden)
            for row in body["rows"]:
                self.assertFalse(set(row) & forbidden)
            values = str(body["rows"]).lower()
            for secret in (
                "1234567890123",
                "5566778899001",
                "123412341234",
                "hdfc",
                "12 main st",
                "30000",
                "tn/tpr/001",
                "100000000001",
            ):
                self.assertNotIn(secret, values, secret)

    def test_status_filter_treats_any_non_active_text_as_inactive(self):
        self.assertEqual(codes(self.get(self.RID, employeeStatus="inactive")), ["E4", "E5", "E9"])  # E9 is 'resigned'
        self.assertEqual(len(codes(self.get(self.RID, employeeStatus="all"))), 10)
        body = self.get(self.RID, employeeStatus="all")
        self.assertEqual(cards(body)["Inactive / left"], 3)
        statuses = {r["employeeCode"]: r["status"] for r in body["rows"]}
        self.assertEqual((statuses["E4"], statuses["E9"], statuses["E1"]), ("Inactive", "Resigned", "Active"))

    def test_each_filter_narrows(self):
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_sew1.id))), ["E2", "E7"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2", "E7"])
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E2", "E8"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E3", "E8"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=f"{self.e1.id},{self.e6.id}")), ["E1", "E6"])
        self.assertEqual(codes(self.get(self.RID, gender="male")), ["E1", "E8", "E10"])
        self.assertEqual(codes(self.get(self.RID, gender="female")), ["E2", "E3", "E6"])
        self.assertEqual(codes(self.get(self.RID, gender="unspecified")), ["E7"])
        self.assertEqual(codes(self.get(self.RID, gender="other")), [])

    def test_unrecognised_gender_text_is_bucketed_as_other(self):
        self.make("E11", gender="Non-binary", branch=self.b1)
        self.assertEqual(codes(self.get(self.RID, gender="other")), ["E11"])
        row = self.get(self.RID, gender="other")["rows"][0]
        self.assertEqual(row["gender"], "Non-binary")  # the raw value stays visible so it can be corrected

    def test_notes_report_unreadable_join_dates_and_missing_dob(self):
        notes = " ".join(self.get(self.RID)["notes"])
        self.assertIn("1 employee(s) have a join date that could not be read", notes)
        self.assertIn("1 employee(s) have no date of birth", notes)

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user, employeeStatus="all")
        self.assertEqual(codes(body), ["E1", "E2", "E4", "E5", "E7", "E10"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, employeeIds=str(self.e3.id))), [])
        self.assertEqual(
            codes(self.get(self.RID, user=self.branch_user, employeeIds=str(self.e6.id))), []
        )  # no-branch record

    def test_hod_scope_shows_no_hod_across_branches_for_a_scoped_user(self):
        # Scoped user still gets HODs for their own employees
        body = self.get(self.RID, user=self.branch_user, layout="full")
        self.assertEqual({r["employeeCode"]: r["hod"] for r in body["rows"]}["E2"], "Arun Kumar")


# ═════════════════════════════════════════════════════════════════════════════
# employee-contact-list
# ═════════════════════════════════════════════════════════════════════════════


class ContactListTests(_World):
    RID = "employee-contact-list"

    def test_default_rows_summary_and_notes(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E1", "E2", "E3", "E6", "E7", "E8", "E10"])
        self.assertEqual(
            cards(body),
            {"Employees listed": 7, "Missing phone": 3, "Missing emergency contact": 5, "Blood group not recorded": 2},
        )
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["E2"]["bloodGroup"], "B+")  # 'b+' normalised
        self.assertEqual(rows["E10"]["bloodGroup"], "O+")  # ' o+ ' normalised
        self.assertEqual(rows["E3"]["phone"], "+91 98765 43212")
        self.assertIsNone(rows["E10"]["phone"])  # '' -> no value
        self.assertIsNone(rows["E3"]["bloodGroup"])
        self.assertTrue(any("Blood groups:" in n and "O+ 2" in n and "not recorded 2" in n for n in body["notes"]))
        self.assertNotIn("address", " ".join(c["key"].lower() for c in body["columns"]))
        self.assertNotIn("12 Main St", str(body))

    def test_blood_group_filter_including_not_recorded(self):
        self.assertEqual(codes(self.get(self.RID, bloodGroup="O+")), ["E1", "E10"])
        self.assertEqual(codes(self.get(self.RID, bloodGroup="unknown")), ["E3", "E7"])
        self.assertEqual(codes(self.get(self.RID, bloodGroup="A+")), ["E6"])

    def test_sort_orders(self):
        self.assertEqual(codes(self.get(self.RID, sortBy="bloodGroup")), ["E6", "E2", "E8", "E1", "E10", "E3", "E7"])
        self.assertEqual(codes(self.get(self.RID, sortBy="department")), ["E1", "E3", "E8", "E10", "E2", "E7", "E6"])
        self.assertEqual(codes(self.get(self.RID, sortBy="name")), ["E1", "E2", "E3", "E6", "E7", "E8", "E10"])
        self.assertEqual(self.get(self.RID, sortBy="bogus", expect=400)["error"], "invalid_filter")

    def test_scope_filters_and_isolation(self):
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_cut1.id))), ["E1", "E10"])
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E2", "E8"])
        self.assertEqual(codes(self.get(self.RID, employeeStatus="inactive")), ["E4", "E5", "E9"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E3", "E8"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2", "E7"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=f"{self.e2.id},{self.e10.id}")), ["E2", "E10"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user)), ["E1", "E2", "E7", "E10"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# strength-statement
# ═════════════════════════════════════════════════════════════════════════════


class StrengthTests(_World):
    RID = "strength-statement"

    def test_department_view_golden(self):
        body = self.get(self.RID)
        self.assertEqual([c["label"] for c in body["columns"]][:2], ["Department", "Branch"])
        rows = body["rows"]
        self.assertEqual([r["group"] for r in rows], ["CUTTING", "CUTTING", "SEWING", "Unassigned"])
        # same-named departments of different units are separate rows, distinguished by branch
        self.assertEqual([r["branch"] for r in rows], ["Unit 1", "Unit 2", "Unit 1", None])
        keys = ("staff", "production", "male", "female", "otherOrUnset", "total")
        self.assertEqual(
            [tuple(r[k] for k in keys) for r in rows],
            [(2, 0, 2, 0, 0, 2), (1, 1, 1, 1, 0, 2), (1, 1, 0, 1, 1, 2), (1, 0, 0, 1, 0, 1)],
        )
        self.assertEqual([r["sharePct"] for r in rows], [28.57, 28.57, 28.57, 14.29])
        self.assertEqual(
            body["totals"], {"staff": 5, "production": 2, "male": 3, "female": 3, "otherOrUnset": 1, "total": 7}
        )
        # 'Unassigned' is a catch-all bucket, not a department: 3 real departments
        self.assertEqual(
            cards(body), {"Total strength": 7, "Staff": 5, "Production": 2, "Male": 3, "Female": 3, "Departments": 3}
        )

    def test_totals_equal_the_sum_of_the_rows(self):
        for group_by in ("department", "designation", "branch"):
            body = self.get(self.RID, groupBy=group_by, employeeStatus="all")
            for k in ("staff", "production", "male", "female", "otherOrUnset", "total"):
                self.assertEqual(body["totals"][k], sum(r[k] for r in body["rows"]), (group_by, k))
            self.assertEqual(body["totals"]["total"], 10)
            self.assertEqual(body["totals"]["staff"] + body["totals"]["production"], 10)

    def test_designation_and_branch_views(self):
        body = self.get(self.RID, groupBy="designation")
        self.assertEqual([c["label"] for c in body["columns"]][:3], ["Designation", "Department", "Level"])
        rows = {r["group"]: r for r in body["rows"]}
        self.assertEqual(rows["Manager"]["total"], 1)
        self.assertEqual((rows["Manager"]["department"], rows["Manager"]["level"]), ("CUTTING", "manager"))
        self.assertEqual((rows["Operator"]["total"], rows["Operator"]["department"]), (2, "SEWING"))
        self.assertEqual(rows["No designation"]["total"], 4)
        self.assertEqual(body["rows"][-1]["group"], "No designation")  # unassigned always last
        body = self.get(self.RID, groupBy="branch")
        self.assertEqual(
            [(r["group"], r["code"], r["total"]) for r in body["rows"]],
            [("Unit 1", "U1", 4), ("Unit 2", "U2", 2), ("No branch", None, 1)],
        )

    def test_filters(self):
        self.assertEqual([r["total"] for r in self.get(self.RID, employmentType="production")["rows"]], [1, 1])
        body = self.get(self.RID, employeeStatus="inactive")
        self.assertEqual([(r["group"], r["branch"], r["staff"], r["production"], r["male"], r["female"]) for r in body["rows"]],
                         [("CUTTING", "Unit 2", 1, 0, 0, 1), ("SEWING", "Unit 1", 1, 1, 2, 0)])  # fmt: skip
        self.assertEqual([r["total"] for r in self.get(self.RID, branchIds=str(self.b2.id))["rows"]], [2])
        self.assertEqual(
            [r["group"] for r in self.get(self.RID, departmentIds=str(self.d_sew1.id))["rows"]], ["SEWING"]
        )
        self.assertEqual(
            self.get(self.RID, groupBy="designation", designationIds=str(self.des_mgr.id))["totals"]["total"], 1
        )

    def test_include_empty_groups(self):
        body = self.get(self.RID, includeEmpty="true")
        groups = [(r["group"], r["total"]) for r in body["rows"]]
        self.assertIn(("LEGACY", 0), groups)
        self.assertIn(("SCRAP", 0), groups)
        self.assertEqual(len(groups), 6)
        self.assertEqual(body["totals"]["total"], 7)

    def test_unknown_employment_type_and_gender_are_reported_not_hidden(self):
        self.make("E11", employment_type="contract", gender="Non-binary", branch=self.b1, department=self.d_cut1)
        body = self.get(self.RID)
        self.assertEqual(body["totals"]["total"], 8)
        self.assertEqual(body["totals"]["staff"] + body["totals"]["production"], 7)
        self.assertTrue(any("employment type other than staff or production" in n for n in body["notes"]))
        cutting = next(r for r in body["rows"] if r["group"] == "CUTTING" and r["branch"] == "Unit 1")
        self.assertEqual((cutting["total"], cutting["otherOrUnset"]), (3, 1))

    def test_designation_without_a_department_is_admin_only(self):
        Designation.objects.create(title="Floater", department=None)
        body = self.get(self.RID, groupBy="designation", includeEmpty="true")
        self.assertIn(("Floater", 0), [(r["group"], r["total"]) for r in body["rows"]])
        body = self.get(self.RID, user=self.branch_user, groupBy="designation", includeEmpty="true")
        self.assertNotIn("Floater", [r["group"] for r in body["rows"]])

    def test_branch_isolation_and_include_empty_do_not_leak(self):
        body = self.get(self.RID, user=self.branch_user)
        self.assertEqual([(r["group"], r["total"]) for r in body["rows"]], [("CUTTING", 2), ("SEWING", 2)])
        body = self.get(self.RID, user=self.branch_user, includeEmpty="true")
        self.assertEqual(
            [r["group"] for r in body["rows"]], ["CUTTING", "SEWING"]
        )  # LEGACY / SCRAP / Unit 2 stay hidden
        body = self.get(self.RID, user=self.branch_user, groupBy="branch", includeEmpty="true")
        self.assertEqual([r["group"] for r in body["rows"]], ["Unit 1"])
        self.assertEqual(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))["rows"], [])


# ═════════════════════════════════════════════════════════════════════════════
# new-joinings
# ═════════════════════════════════════════════════════════════════════════════


class NewJoiningsTests(_World):
    RID = "new-joinings"

    def test_default_period_is_this_month_and_parses_text_dates(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E2", "E3"])  # E3's join date is the text '05-09-2026'
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["E3"]["joinDate"], "2026-09-05")
        self.assertEqual((rows["E2"]["ageAtJoining"], rows["E3"]["ageAtJoining"]), (24, 26))
        self.assertEqual((rows["E2"]["docsMissing"], rows["E3"]["docsMissing"]), (6, 5))
        self.assertEqual(
            cards(body),
            {
                "Joined in period": 2,
                "Staff": 1,
                "Production": 1,
                "Still active": 2,
                "Left since joining": 0,
                "With documents pending": 2,
            },
        )
        notes = " ".join(body["notes"])
        self.assertIn("1 employee(s) in scope have a join date that could not be read", notes)  # E7 'soon'
        self.assertIn("1 employee(s) in scope have no join date on file", notes)  # E10 ''

    def test_document_gap_counts_required_documents_only(self):
        for cat in (
            "pan_card",
            "aadhaar_card",
            "educational_certificate",
            "voter_id_or_birth_certificate",
            "bank_passbook",
            "production_employee_documents",
        ):
            self.doc(self.e2, cat)
        self.doc(self.e3, "offer_letter")  # not a required category
        body = self.get(self.RID)
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual((rows["E2"]["docsMissing"], rows["E3"]["docsMissing"]), (0, 5))
        self.assertEqual(cards(body)["With documents pending"], 1)

    def test_ranges_are_inclusive_and_bounded(self):
        self.assertEqual(codes(self.get(self.RID, dateFrom="2026-08-01", dateTo="2026-08-31")), ["E8"])
        self.assertEqual(codes(self.get(self.RID, dateFrom="2023-11-30", dateTo="2023-11-30")), ["E6"])
        self.assertEqual(codes(self.get(self.RID, dateFrom="2026-01-01", dateTo="2026-12-31")), ["E8", "E2", "E3"])
        self.assertEqual(self.get(self.RID, dateFrom="2000-01-01", dateTo="2026-12-31", expect=400)["field"], "dateTo")

    def test_status_after_joining_and_filters(self):
        body = self.get(self.RID, dateFrom="2020-01-01", dateTo="2026-12-31")
        self.assertEqual(codes(body), ["E4", "E6", "E1", "E5", "E8", "E2", "E3"])  # by join date
        self.assertEqual(cards(body)["Still active"], 5)
        self.assertEqual(cards(body)["Left since joining"], 2)
        self.assertEqual(
            codes(self.get(self.RID, dateFrom="2020-01-01", dateTo="2026-12-31", employeeStatus="inactive")),
            ["E4", "E5"],
        )
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E2"])
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_sew1.id))), ["E2"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E3"])

    def test_future_join_dates_are_listed_when_the_range_covers_them(self):
        self.make("E11", join_date="2026-10-15", branch=self.b1)
        self.assertEqual(codes(self.get(self.RID, dateFrom="2026-10-01", dateTo="2026-10-31")), ["E11"])

    def test_salary_columns_are_permission_gated_and_null_safe(self):
        body = self.get(self.RID)
        self.assertIn("salaryAmount", [c["key"] for c in body["columns"]])
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["E2"]["salaryPerShift"], 500.0)
        self.assertIsNone(rows["E2"]["salaryAmount"])  # unknown is not zero
        self.assertIsNone(rows["E3"]["salaryAmount"])
        body = self.get(self.RID, user=self.no_salary_user)
        keys = [c["key"] for c in body["columns"]]
        self.assertNotIn("salaryAmount", keys)
        self.assertNotIn("salaryPerShift", keys)
        self.assertNotIn("salaryPerShift", body["rows"][0])
        self.assertTrue(any("Salary columns are hidden" in n for n in body["notes"]))
        self.assertIn(
            "salaryAmount", [c["key"] for c in self.get(self.RID, user=self.branch_user)["columns"]]
        )  # payroll: view

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user, dateFrom="2019-01-01", dateTo="2026-12-31")
        self.assertEqual(codes(body), ["E4", "E1", "E5", "E2"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# exits-register
# ═════════════════════════════════════════════════════════════════════════════


class ExitsRegisterTests(_World):
    RID = "exits-register"
    WIDE = dict(dateFrom="2025-01-01", dateTo="2026-12-31")

    def test_golden_exit_dates_basis_and_service(self):
        body = self.get(self.RID, **self.WIDE)
        self.assertEqual(codes(body), ["E9", "E4", "E5"])  # by exit date
        rows = {r["employeeCode"]: r for r in body["rows"]}
        # E4: approved resignation, dated by the last working date
        self.assertEqual((rows["E4"]["exitDate"], rows["E4"]["exitBasis"]), ("2026-08-31", "Resignation"))
        self.assertEqual(
            (rows["E4"]["approvedBy"], rows["E4"]["approvedOn"], rows["E4"]["reason"]),
            ("HR Admin", "2026-08-20", "Better opportunity"),
        )
        self.assertEqual(rows["E4"]["tenureMonths"], 77)
        # E5: manually deactivated, last modified 10-Sep 08:00 UTC = 13:30 IST
        self.assertEqual((rows["E5"]["exitDate"], rows["E5"]["exitBasis"]), ("2026-09-10", "Deactivated (approx.)"))
        self.assertEqual(rows["E5"]["tenureMonths"], 15)
        self.assertIsNone(rows["E5"]["approvedBy"])
        # E9: status 'resigned', last modified 15-Jan 20:00 UTC = 16-Jan 01:30 IST (the IST date, not the UTC one)
        self.assertEqual((rows["E9"]["exitDate"], rows["E9"]["exitBasis"]), ("2026-01-16", "Deactivated (approx.)"))
        self.assertEqual(rows["E9"]["tenureMonths"], 84)
        self.assertEqual(rows["E9"]["status"], "Resigned")
        self.assertEqual(
            cards(body),
            {
                "Total exits": 3,
                "Resignations": 1,
                "Manual deactivations": 2,
                "Staff": 2,
                "Production": 1,
                "Average service (months)": 58.7,
            },
        )
        notes = " ".join(body["notes"])
        self.assertIn("approx", notes)
        self.assertIn("deleted from the system do not appear", notes)

    def test_default_range_is_this_month(self):
        self.assertEqual(codes(self.get(self.RID)), ["E5"])

    def test_exit_type_and_scope_filters(self):
        self.assertEqual(codes(self.get(self.RID, exitType="resignation", **self.WIDE)), ["E4"])
        self.assertEqual(codes(self.get(self.RID, exitType="manual", **self.WIDE)), ["E9", "E5"])
        self.assertEqual(codes(self.get(self.RID, employmentType="staff", **self.WIDE)), ["E9", "E4"])
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_cut2.id), **self.WIDE)), ["E9"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_sup.id), **self.WIDE)), ["E4"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id), **self.WIDE)), ["E9"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=str(self.e5.id), **self.WIDE)), ["E5"])
        self.assertEqual(codes(self.get(self.RID, dateFrom="2026-01-16", dateTo="2026-01-16")), ["E9"])
        self.assertEqual(
            codes(self.get(self.RID, dateFrom="2026-01-15", dateTo="2026-01-15")), []
        )  # the UTC date must not match

    def test_an_approved_resignation_of_an_active_employee_is_not_an_exit(self):
        ResignationRequest.objects.create(
            employee=self.e1, status="approved", last_working_date=date(2026, 9, 15), approved_at=utc(2026, 9, 1, 5, 0)
        )
        self.assertNotIn("E1", codes(self.get(self.RID, **self.WIDE)))  # E1 is still active (re-activated)

    def test_latest_approved_resignation_wins_and_approval_date_is_the_fallback(self):
        e11 = self.make("E11", status="inactive", branch=self.b1, join_date="2024-06-15")
        ResignationRequest.objects.create(
            employee=e11, status="approved", last_working_date=date(2026, 3, 31), approved_at=utc(2026, 3, 1, 5, 0)
        )
        ResignationRequest.objects.create(
            employee=e11, status="approved", last_working_date=date(2026, 6, 30), approved_at=utc(2026, 6, 1, 5, 0)
        )
        e12 = self.make("E12", status="inactive", branch=self.b1)
        ResignationRequest.objects.create(
            employee=e12, status="approved", approved_at=utc(2026, 5, 31, 20, 0)
        )  # no LWD; 1-Jun IST
        ResignationRequest.objects.create(
            employee=e12, status="rejected", last_working_date=date(2026, 2, 1)
        )  # ignored
        rows = {r["employeeCode"]: r for r in self.get(self.RID, **self.WIDE)["rows"]}
        self.assertEqual((rows["E11"]["exitDate"], rows["E11"]["exitBasis"]), ("2026-06-30", "Resignation"))
        self.assertEqual(
            (rows["E12"]["exitDate"], rows["E12"]["exitBasis"]), ("2026-06-01", "Resignation (approval date)")
        )
        self.assertIsNone(rows["E12"]["tenureMonths"])  # no join date -> no service figure
        self.assertEqual(self.get(self.RID, exitType="resignation", **self.WIDE)["rowCount"], 3)

    def test_exit_before_join_date_gives_no_service_figure(self):
        # deactivated (approx.) on 5-Oct although the join date says 20-Oct: the join date is wrong, so no service figure
        e11 = self.make("E11", status="inactive", branch=self.b1, join_date="2026-10-20")
        Employee.objects.filter(pk=e11.pk).update(updated_at=utc(2026, 10, 5, 5, 0))
        body = self.get(self.RID, dateFrom="2026-10-01", dateTo="2026-10-31")
        self.assertEqual(codes(body), ["E11"])
        self.assertIsNone(body["rows"][0]["tenureMonths"])

    def test_future_last_working_date_is_noted(self):
        e11 = self.make("E11", status="inactive", branch=self.b1, join_date="2026-01-05")
        ResignationRequest.objects.create(employee=e11, status="approved", last_working_date=date(2026, 10, 15))
        body = self.get(self.RID, dateFrom="2026-10-01", dateTo="2026-10-31")
        self.assertEqual(codes(body), ["E11"])
        self.assertTrue(any("last working date after today" in n for n in body["notes"]))

    def test_a_resignation_before_the_join_date_is_an_earlier_stint_and_falls_back(self):
        # E11 resigned (LWD 31-Mar-2024), was re-hired 15-Jun-2024 on the same record, deactivated again 10-Sep-2026
        e11 = self.make("E11", status="inactive", branch=self.b1, join_date="2024-06-15")
        ResignationRequest.objects.create(employee=e11, status="approved", last_working_date=date(2024, 3, 31))
        Employee.objects.filter(pk=e11.pk).update(updated_at=utc(2026, 9, 10, 8, 0))
        rows = {r["employeeCode"]: r for r in self.get(self.RID, **self.WIDE)["rows"]}
        self.assertEqual((rows["E11"]["exitDate"], rows["E11"]["exitBasis"]), ("2026-09-10", "Deactivated (approx.)"))
        self.assertEqual(rows["E11"]["tenureMonths"], 26)  # 15-Jun-2024 -> 10-Sep-2026, not negative and not blank
        # a resignation of the CURRENT stint still wins
        ResignationRequest.objects.create(employee=e11, status="approved", last_working_date=date(2026, 8, 31))
        rows = {r["employeeCode"]: r for r in self.get(self.RID, **self.WIDE)["rows"]}
        self.assertEqual((rows["E11"]["exitDate"], rows["E11"]["exitBasis"]), ("2026-08-31", "Resignation"))
        self.assertEqual(rows["E11"]["tenureMonths"], 26)

    def test_reason_and_approver_columns_are_shown_to_a_role_with_the_resignations_module(self):
        body = self.get(self.RID, user=self.branch_user, **self.WIDE)
        e4 = next(r for r in body["rows"] if r["employeeCode"] == "E4")
        self.assertEqual(
            (e4["reason"], e4["approvedBy"], e4["approvedOn"]), ("Better opportunity", "HR Admin", "2026-08-20")
        )
        self.assertFalse(any("hidden: your role has no access to Resignations" in n for n in body["notes"]))

    def test_branch_isolation(self):
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, **self.WIDE)), ["E4", "E5"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id), **self.WIDE)), [])


# ═════════════════════════════════════════════════════════════════════════════
# resignation-register
# ═════════════════════════════════════════════════════════════════════════════


class ResignationRegisterTests(_World):
    RID = "resignation-register"

    def test_default_month_golden(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E3", "E7", "E1", "E6"])  # by request time
        rows = {r["employeeCode"]: r for r in body["rows"]}
        # request raised 20-Sep 20:00 UTC = 21-Sep 01:30 IST
        r2 = rows["E1"]
        self.assertEqual(r2["requestedOn"], "2026-09-21 01:30")
        self.assertEqual(
            (r2["stage"], r2["status"], r2["pendingDays"], r2["noticeDays"]), ("Awaiting HOD", "Pending", 8, 40)
        )
        self.assertEqual(
            (rows["E7"]["stage"], rows["E7"]["pendingDays"], rows["E7"]["deptHead"]), ("Awaiting HR", 14, "Arun Kumar")
        )
        self.assertEqual(
            (rows["E3"]["stage"], rows["E3"]["deptHead"], rows["E3"]["pendingDays"]),
            ("Rejected by HOD", "Jai Singh", None),
        )
        self.assertEqual(rows["E6"]["stage"], "Rejected by HR")
        self.assertEqual(
            cards(body),
            {
                "Total requests": 4, "Awaiting HOD": 1, "Awaiting HR": 1, "Approved": 0, "Rejected by HOD": 1,
                "Rejected by HR": 1, "Avg days to decision": 2.0,
            },
        )  # fmt: skip

    def test_approved_request_and_average_decision_time(self):
        body = self.get(self.RID, dateFrom="2026-08-01", dateTo="2026-09-29", layout="full")
        self.assertEqual(codes(body), ["E4", "E3", "E7", "E1", "E6"])
        r1 = body["rows"][0]
        self.assertEqual(r1["requestedOn"], "2026-08-10 09:30")
        self.assertEqual((r1["lastWorkingDate"], r1["noticeDays"], r1["stage"]), ("2026-08-31", 21, "Approved"))
        self.assertEqual((r1["approvedBy"], r1["approvedAt"], r1["hrComment"]), ("HR Admin", "2026-08-20 15:30", "OK"))
        self.assertEqual(
            (r1["deptHeadStatus"], r1["deptHeadAt"], r1["deptHeadComment"]), ("Approved", "2026-08-12 10:30", "Fine")
        )
        self.assertEqual(
            (r1["surveyReason"], r1["surveyRecommend"], r1["surveyRetain"]), ("Salary", "Yes", "Higher pay")
        )
        self.assertIsNone(r1["pendingDays"])
        # (approved: 20-Aug - 10-Aug = 10) and (HOD rejection: 3-Sep - 1-Sep = 2) -> 6.0; the HR rejection has no time stamp
        self.assertEqual(cards(body)["Avg days to decision"], 6.0)

    def test_summary_layout_omits_comments_and_survey_answers(self):
        summary = self.get(self.RID)
        full = self.get(self.RID, layout="full")
        s_keys = {c["key"] for c in summary["columns"]}
        f_keys = {c["key"] for c in full["columns"]}
        self.assertTrue(s_keys < f_keys)
        for k in ("deptHeadComment", "hrComment", "surveyReason", "surveyRecommend", "surveyRetain"):
            self.assertNotIn(k, s_keys)
            self.assertIn(k, f_keys)
        self.assertEqual(len(f_keys), 21)

    def test_date_range_uses_the_ist_day(self):
        self.assertEqual(codes(self.get(self.RID, dateFrom="2026-09-21", dateTo="2026-09-21")), ["E1"])
        self.assertEqual(
            codes(self.get(self.RID, dateFrom="2026-09-20", dateTo="2026-09-20")), []
        )  # UTC date must not match

    def test_filters(self):
        wide = dict(dateFrom="2026-08-01", dateTo="2026-09-29")
        self.assertEqual(codes(self.get(self.RID, status="pending,dept_approved", **wide)), ["E7", "E1"])
        self.assertEqual(codes(self.get(self.RID, status="approved", **wide)), ["E4"])
        self.assertEqual(codes(self.get(self.RID, status="rejected", **wide)), ["E3", "E6"])
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_cut2.id), **wide)), ["E3"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id), **wide)), ["E3"])
        self.assertEqual(codes(self.get(self.RID, employmentType="production", **wide)), [])
        self.assertEqual(codes(self.get(self.RID, employeeIds=str(self.e4.id), **wide)), ["E4"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_sup.id), **wide)), ["E4"])
        self.assertEqual(self.get(self.RID, status="bogus", expect=400)["field"], "status")

    def test_requests_of_inactive_employees_still_appear(self):
        # The approved resignation's employee (E4) is inactive: a transactional register must not drop them.
        self.assertIn("E4", codes(self.get(self.RID, dateFrom="2026-08-01", dateTo="2026-08-31")))

    def test_production_requests_from_the_mobile_flow_are_listed(self):
        r = ResignationRequest.objects.create(employee=self.e2, reason="Moving", status="pending")
        ResignationRequest.objects.filter(pk=r.pk).update(created_at=utc(2026, 9, 28, 5, 0))
        self.assertIn("E2", codes(self.get(self.RID, employmentType="production")))

    def test_branch_isolation(self):
        wide = dict(dateFrom="2026-08-01", dateTo="2026-09-29")
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, **wide)), ["E4", "E7", "E1"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id), **wide)), [])
        body = self.get(self.RID, user=self.branch_user, **wide)
        self.assertNotIn("E3", str(body))
        self.assertNotIn("E6", str(body))
        # the counters are computed over the scoped rows only
        self.assertEqual(cards(body)["Total requests"], 3)


# ═════════════════════════════════════════════════════════════════════════════
# manpower-movement
# ═════════════════════════════════════════════════════════════════════════════


class ManpowerMovementTests(_World):
    RID = "manpower-movement"

    def test_monthly_view_golden(self):
        body = self.get(self.RID, year="2026")
        rows = {r["group"]: r for r in body["rows"]}
        self.assertEqual(len(body["rows"]), 12)
        self.assertEqual([r["group"] for r in body["rows"]][:2], ["Jan 2026", "Feb 2026"])
        keys = ("joinedStaff", "joinedProduction", "joinedTotal", "leftStaff", "leftProduction", "leftTotal", "net")
        self.assertEqual(tuple(rows["Jan 2026"][k] for k in keys), (0, 0, 0, 1, 0, 1, -1))  # E9 (approx 16-Jan)
        self.assertEqual(tuple(rows["Aug 2026"][k] for k in keys), (0, 1, 1, 1, 0, 1, 0))  # E8 joined; E4 left 31-Aug
        self.assertEqual(tuple(rows["Sep 2026"][k] for k in keys), (1, 1, 2, 0, 1, 1, 1))  # E3+E2 joined; E5 left
        self.assertEqual(tuple(rows["Feb 2026"][k] for k in keys), (0, 0, 0, 0, 0, 0, 0))
        self.assertEqual(
            body["totals"],
            {
                "joinedStaff": 1,
                "joinedProduction": 2,
                "joinedTotal": 3,
                "leftStaff": 2,
                "leftProduction": 1,
                "leftTotal": 3,
                "net": 0,
                "exitRatePct": 30.0,
            },
        )
        # exit rate = leavers / (active today 7 + leavers in the year 3): 3/10 overall, 1/10 in each leaving month
        self.assertEqual(
            (rows["Jan 2026"]["exitRatePct"], rows["Aug 2026"]["exitRatePct"], rows["Sep 2026"]["exitRatePct"]),
            (10.0, 10.0, 10.0),
        )
        self.assertEqual(rows["Feb 2026"]["exitRatePct"], 0.0)
        self.assertEqual(cards(body), {"Joined": 3, "Left": 3, "Net movement": 0, "Exit rate": 30.0, "Active today": 7})
        self.assertTrue(any("approximation" in n for n in body["notes"]))
        self.assertTrue(
            any("1 leaver(s)" not in n and "2 leaver(s) have only an approximate" in n for n in body["notes"])
        )  # E5, E9

    def test_department_view_golden(self):
        body = self.get(self.RID, year="2026", groupBy="department")
        # the two CUTTING departments are told apart by unit; SEWING has no namesake, so it stays plain
        self.assertEqual(
            [r["group"] for r in body["rows"]], ["CUTTING (Unit 1)", "CUTTING (Unit 2)", "SEWING", "Unassigned"]
        )
        keys = ("joinedTotal", "leftTotal", "net", "exitRatePct")
        got = [tuple(r[k] for k in keys) for r in body["rows"]]
        # Unit 1 CUTTING: nobody moved (2 active). Unit 2 CUTTING: E3+E8 joined, E9 left 1/(2+1). SEWING: E2 joined,
        # E4+E5 left 2/(2+2). Unassigned: only E6, active since 2023.
        self.assertEqual(got, [(0, 0, 0, 0.0), (2, 1, 1, 33.33), (1, 2, -1, 50.0), (0, 0, 0, 0.0)])
        self.assertEqual(
            (body["totals"]["joinedTotal"], body["totals"]["leftTotal"], body["totals"]["exitRatePct"]), (3, 3, 30.0)
        )
        self.assertTrue(any("own active strength" in n for n in body["notes"]))

    def test_another_year_and_filters(self):
        body = self.get(self.RID, year="2025")
        self.assertEqual(body["totals"]["joinedTotal"], 1)  # E5 joined 10-Jun-2025 (a leaver later, in 2026)
        self.assertEqual(next(r for r in body["rows"] if r["group"] == "Jun 2025")["joinedProduction"], 1)
        self.assertEqual(body["totals"]["leftTotal"], 0)
        # the rate divides by TODAY's headcount: meaningless (and blanked, with a note) for any year but this one
        self.assertIsNone(body["totals"]["exitRatePct"])
        self.assertTrue(all(r["exitRatePct"] is None for r in body["rows"]))
        self.assertIsNone(cards(body)["Exit rate"])
        self.assertTrue(any("current year only" in n for n in body["notes"]))
        body = self.get(self.RID, year="2026", employmentType="production")
        self.assertEqual((body["totals"]["joinedTotal"], body["totals"]["leftTotal"]), (2, 1))
        body = self.get(self.RID, year="2026", branchIds=str(self.b2.id))
        self.assertEqual((body["totals"]["joinedTotal"], body["totals"]["leftTotal"]), (2, 1))
        body = self.get(self.RID, year="2026", departmentIds=str(self.d_sew1.id))
        self.assertEqual((body["totals"]["joinedTotal"], body["totals"]["leftTotal"]), (1, 2))

    def test_default_year_is_this_year(self):
        self.assertEqual(self.get(self.RID)["rows"][0]["group"], "Jan 2026")

    def test_a_person_who_joined_and_left_in_the_year_counts_in_both(self):
        e11 = self.make("E11", status="inactive", branch=self.b1, department=self.d_cut1, join_date="2026-02-10")
        ResignationRequest.objects.create(employee=e11, status="approved", last_working_date=date(2026, 4, 30))
        body = self.get(self.RID, year="2026")
        rows = {r["group"]: r for r in body["rows"]}
        self.assertEqual((rows["Feb 2026"]["joinedStaff"], rows["Apr 2026"]["leftStaff"]), (1, 1))
        self.assertEqual(body["totals"]["joinedTotal"], 4)
        self.assertEqual(body["totals"]["leftTotal"], 4)

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user, year="2026")
        # Unit 1 only: E2 joined; E4 + E5 left; active today E1, E2, E7, E10 -> 2 / (4 + 2)
        self.assertEqual((body["totals"]["joinedTotal"], body["totals"]["leftTotal"]), (1, 2))
        self.assertEqual(body["totals"]["exitRatePct"], 33.33)
        self.assertEqual(cards(body)["Active today"], 4)
        self.assertEqual(
            self.get(self.RID, user=self.branch_user, year="2026", branchIds=str(self.b2.id))["totals"]["joinedTotal"],
            0,
        )

    def test_empty_department_view_has_no_totals_row(self):
        body = self.get(self.RID, groupBy="department", departmentIds="999999")
        self.assertEqual(body["rows"], [])
        self.assertTrue(all(v is None for v in (body["totals"] or {}).values()))


# ═════════════════════════════════════════════════════════════════════════════
# statutory-compliance
# ═════════════════════════════════════════════════════════════════════════════


class StatutoryComplianceTests(_World):
    RID = "statutory-compliance"

    def rows(self, **params):
        body = self.get(self.RID, **params)
        return body, {r["employeeCode"]: r for r in body["rows"]}

    def test_default_lists_only_employees_with_gaps_and_summarises_everyone(self):
        body, rows = self.rows()
        self.assertEqual(codes(body), ["E2", "E3", "E6", "E7", "E8", "E10"])  # E1 is fully compliant
        self.assertEqual(
            cards(body),
            {
                "Employees checked": 7, "Fully compliant": 1, "PF number missing": 3, "ESI number missing": 4,
                "UAN missing": 4, "Bank details missing": 4, "Deducted, number missing": 1, "Completeness": 14.3,
            },
        )  # fmt: skip
        self.assertEqual(
            {k: v["issueCount"] for k, v in rows.items()}, {"E2": 7, "E3": 6, "E6": 7, "E7": 7, "E8": 5, "E10": 7}
        )

    def test_each_issue_is_named(self):
        _b, rows = self.rows()
        self.assertEqual(
            rows["E2"]["issues"],
            "PF number missing; ESI number missing; UAN missing; Bank a/c and IFSC missing; "
            "Aadhaar not uploaded; PAN not uploaded; Bank passbook not uploaded",
        )
        self.assertEqual(
            rows["E3"]["issues"],
            "IFSC format is invalid; Bank a/c should be 9-18 digits; UAN should be 12 digits; "
            "ESI number should be 10 or 17 digits; PAN not uploaded; Bank passbook not uploaded",
        )
        self.assertIn("PF deducted in payroll but no PF number", rows["E8"]["issues"])

    def test_identifiers_are_text_and_bank_details_are_masked(self):
        body, rows = self.rows(includeCompliant="true")
        e1 = rows["E1"]
        self.assertEqual(e1["pfNumber"], "TN/TPR/001")
        self.assertEqual(e1["esiNumber"], "1234567890")
        self.assertEqual(e1["uanNumber"], "100000000001")
        self.assertEqual(e1["bankAccount"], "*********9001")
        self.assertEqual(e1["idProof"], "********1234")
        self.assertEqual(e1["bankIfsc"], "HDFC0001234")
        self.assertEqual(rows["E8"]["bankIfsc"], "SBIN0001234")  # upper-cased
        self.assertEqual(rows["E8"]["bankAccount"], "*****6789")
        self.assertEqual(rows["E8"]["esiNumber"], "12345678901234567")  # 17-digit legacy ESI kept in full, as text
        self.assertIsInstance(e1["uanNumber"], str)
        self.assertNotIn("5566778899001", str(body))
        self.assertNotIn("123412341234", str(body))

    def test_documents_and_payroll_evidence(self):
        _b, rows = self.rows(includeCompliant="true")
        self.assertEqual(
            (rows["E1"]["aadhaarDoc"], rows["E1"]["panDoc"], rows["E1"]["passbookDoc"]), ("Yes", "Yes", "Yes")
        )
        self.assertEqual((rows["E3"]["aadhaarDoc"], rows["E3"]["panDoc"]), ("Yes", "No"))
        # August 2026 is "last month": E1 has a PF+ESI slip, E8's two weekly slips (60 + 40) add up to a PF deduction
        self.assertEqual((rows["E1"]["pfDeducted"], rows["E1"]["esiDeducted"]), ("Yes", "Yes"))
        self.assertEqual((rows["E8"]["pfDeducted"], rows["E8"]["esiDeducted"]), ("Yes", "No"))
        self.assertIsNone(rows["E2"]["pfDeducted"])  # no slip that month: unknown, not "No"
        body, rows = self.rows(includeCompliant="true", period="2026-09")
        self.assertIsNone(rows["E8"]["pfDeducted"])
        self.assertEqual(cards(body)["Deducted, number missing"], 0)
        self.assertEqual(rows["E8"]["issueCount"], 4)

    def test_issue_filter(self):
        expect = {
            "missing_pf": ["E2", "E6", "E8"],
            "missing_esi": ["E2", "E6", "E7", "E10"],
            "missing_uan": ["E2", "E6", "E7", "E10"],
            "missing_bank": ["E2", "E6", "E7", "E10"],
            "invalid_format": ["E3"],
            "no_aadhaar_doc": ["E2", "E6", "E7", "E8", "E10"],
            "no_pan_doc": ["E2", "E3", "E6", "E7", "E8", "E10"],
            "no_passbook_doc": ["E2", "E3", "E6", "E7", "E8", "E10"],
            "duplicate_id": ["E7", "E10"],
            "deducted_no_number": ["E8"],
        }
        for issue, want in expect.items():
            self.assertEqual(codes(self.get(self.RID, issue=issue)), want, issue)

    def test_duplicates_normalise_and_name_the_other_employee(self):
        _b, rows = self.rows()
        self.assertIn("PF number shared with E10", rows["E7"]["issues"])  # 'PF-DUP-9' == 'pf dup 9'
        self.assertIn("PF number shared with E7", rows["E10"]["issues"])

    def test_duplicates_are_found_across_departments_even_when_filtered(self):
        _b, rows = self.rows(departmentIds=str(self.d_sew1.id))
        self.assertEqual(sorted(rows), ["E2", "E7"])
        self.assertIn("shared with E10", rows["E7"]["issues"])  # E10 is in CUTTING, outside the filter

    def test_inactive_employees_are_never_flagged_as_duplicates(self):
        self.make("E11", status="inactive", branch=self.b1, pf_number="TN/TPR/001")
        body, rows = self.rows(employeeStatus="all", includeCompliant="true")
        self.assertNotIn("shared with", rows["E11"]["issues"] or "")
        self.assertEqual(rows["E1"]["issueCount"], 0)  # a rejoiner's old record does not taint the active one

    def test_duplicate_detection_respects_branch_isolation(self):
        self.make("E11", branch=self.b2, pf_number="TN/TPR/001")
        _b, rows = self.rows(includeCompliant="true")
        self.assertIn("shared with E11", rows["E1"]["issues"])  # the admin sees the clash
        body = self.get(self.RID, user=self.branch_user, includeCompliant="true")
        e1 = next(r for r in body["rows"] if r["employeeCode"] == "E1")
        self.assertEqual(e1["issueCount"], 0)  # a Unit 1 user cannot see (or be told about) Unit 2's people
        self.assertNotIn("E11", str(body))

    def test_scope_filters_and_status(self):
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E2", "E8"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E3", "E8"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2", "E7"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=str(self.e6.id))), ["E6"])
        self.assertEqual(codes(self.get(self.RID, employeeStatus="inactive")), ["E4", "E5", "E9"])
        self.assertEqual(cards(self.get(self.RID, employeeStatus="all"))["Employees checked"], 10)
        self.assertEqual(len(codes(self.get(self.RID, employeeStatus="all", includeCompliant="true"))), 10)

    def test_missing_uan_note_appears_when_almost_nobody_has_one(self):
        # 4 of 7 active employees lack a UAN: no note. Wipe the rest and the data-capture gap is stated.
        self.assertFalse(any("UAN cannot currently be entered" in n for n in self.get(self.RID)["notes"]))
        Employee.objects.filter(status="active").update(uan_number=None)
        self.assertTrue(any("UAN cannot currently be entered" in n for n in self.get(self.RID)["notes"]))

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user, includeCompliant="true")
        self.assertEqual(codes(body), ["E1", "E2", "E7", "E10"])
        self.assertEqual(cards(body)["Employees checked"], 4)
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# data-quality-audit
# ═════════════════════════════════════════════════════════════════════════════


class DataQualityTests(_World):
    RID = "data-quality-audit"

    def test_golden_default(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E2", "E3", "E6", "E7", "E8", "E10"])
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(
            {k: v["issueCount"] for k, v in rows.items()}, {"E2": 2, "E3": 5, "E6": 5, "E7": 8, "E8": 5, "E10": 4}
        )
        self.assertEqual(rows["E2"]["issues"], "Father's name missing; Emergency contact missing")
        self.assertEqual(
            rows["E7"]["issues"],
            "Date of birth missing; Gender not set; Phone missing; Address missing; Join date 'soon' cannot be read; "
            "Father's name missing; Blood group missing; Emergency contact missing",
        )
        self.assertEqual(
            rows["E6"]["issues"], "Phone missing; No department; No designation; No branch; Monthly salary not set"
        )
        self.assertEqual(
            rows["E8"]["issues"],
            "Address missing; No designation; Per-shift rate not set; Father's name missing; Emergency contact missing",
        )
        self.assertEqual(
            rows["E10"]["issues"], "Phone missing; Join date missing; No designation; Emergency contact missing"
        )
        self.assertEqual(rows["E7"]["joinDate"], "soon")  # shown exactly as entered, so the fix is visible
        self.assertEqual(
            cards(body),
            {
                "Employees checked": 7, "Clean records": 1, "Records with issues": 6, "Profile completeness": 14.3,
                "Date of birth: issues": 1, "Gender: issues": 1, "Phone: issues": 3, "Address: issues": 3,
                "Join date: issues": 2, "Department / designation / branch: issues": 4, "Salary set-up: issues": 3,
                "Father's name, blood group, emergency contact: issues": 5,
            },
        )  # fmt: skip

    def test_include_clean_lists_everyone(self):
        body = self.get(self.RID, includeClean="true")
        self.assertEqual(len(body["rows"]), 7)
        e1 = next(r for r in body["rows"] if r["employeeCode"] == "E1")
        self.assertEqual((e1["issueCount"], e1["issues"]), (0, None))

    def test_check_selection(self):
        self.assertEqual(codes(self.get(self.RID, check="dob")), ["E7"])
        body = self.get(self.RID, check="phone,address")
        self.assertEqual(codes(body), ["E3", "E6", "E7", "E8", "E10"])
        self.assertEqual(
            {r["employeeCode"]: r["issueCount"] for r in body["rows"]}, {"E3": 1, "E6": 1, "E7": 2, "E8": 1, "E10": 1}
        )
        self.assertEqual(codes(self.get(self.RID, check="salary")), ["E3", "E6", "E8"])
        self.assertEqual(codes(self.get(self.RID, check="join_date")), ["E7", "E10"])
        self.assertEqual(codes(self.get(self.RID, check="org_unit")), ["E3", "E6", "E8", "E10"])
        self.assertEqual(self.get(self.RID, check="bogus", expect=400)["field"], "check")

    def test_status_and_scope_filters(self):
        self.assertEqual(codes(self.get(self.RID, employeeStatus="inactive")), ["E4", "E5", "E9"])
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_sew1.id))), ["E2", "E7"])
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E2", "E8"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E3", "E8"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2", "E7"])

    def test_implausible_dates_of_birth(self):
        self.make("E11", date_of_birth=date(2030, 1, 1), branch=self.b1, join_date="2026-01-01")
        self.make("E12", date_of_birth=date(2015, 1, 1), branch=self.b1, join_date="2026-01-01")  # age 11
        self.make("E13", date_of_birth=date(1930, 1, 1), branch=self.b1, join_date="2026-01-01")  # age 96
        self.make("E14", date_of_birth=date(1990, 1, 1), branch=self.b1, join_date="1989-06-01")  # born after joining
        rows = {r["employeeCode"]: r["issues"] for r in self.get(self.RID, check="dob")["rows"]}
        self.assertIn("Date of birth is in the future", rows["E11"])
        self.assertIn("Age 11 looks implausible", rows["E12"])
        self.assertIn("Age 96 looks implausible", rows["E13"])
        self.assertIn("Date of birth is not before the join date", rows["E14"])

    def test_join_date_in_the_future_and_unrecognised_values(self):
        self.make("E11", join_date="2027-01-01", gender="M", blood_group="Z", branch=self.b1)
        row = next(
            r for r in self.get(self.RID, check="join_date,gender,profile")["rows"] if r["employeeCode"] == "E11"
        )
        self.assertIn("Join date is in the future", row["issues"])
        self.assertIn("Gender 'M' is not recognised", row["issues"])
        self.assertIn("Blood group 'Z' is not recognised", row["issues"])

    def test_phone_format_and_duplicates(self):
        self.make("E11", phone="12345", branch=self.b1)
        self.make("E12", phone="+91-98765-43210", branch=self.b1)  # same number as E1
        self.make("E13", phone="9000000013", branch=self.b1, status="inactive")
        self.make("E14", phone="9000000013", branch=self.b1)  # shares only with an inactive employee: not flagged
        rows = {r["employeeCode"]: r["issues"] for r in self.get(self.RID, check="phone", employeeStatus="all")["rows"]}
        self.assertIn("Phone is not a 10-digit number", rows["E11"])
        self.assertIn("Phone also used by E12", rows["E1"])
        self.assertIn("Phone also used by E1", rows["E12"])
        self.assertNotIn("E14", rows)
        self.assertNotIn("E13", rows)

    def test_duplicate_people_and_stray_spaces_in_codes(self):
        self.make("E11", first_name="Arun", last_name="Kumar", date_of_birth=date(1990, 3, 10), branch=self.b1)
        self.make("E12 ", branch=self.b1)  # trailing space
        body = self.get(self.RID, check="duplicates,code")
        rows = {r["employeeCode"]: r["issues"] for r in body["rows"]}
        self.assertIn("Same name and date of birth as E11", rows["E1"])
        self.assertIn("Same name and date of birth as E1", rows["E11"])
        self.assertIn("Employee code has stray spaces", rows["E12"])  # shown trimmed
        self.assertNotIn("E12 ", rows)

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user, includeClean="true")
        self.assertEqual(codes(body), ["E1", "E2", "E7", "E10"])
        self.assertEqual(cards(body)["Employees checked"], 4)
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# age-compliance
# ═════════════════════════════════════════════════════════════════════════════


class AgeComplianceTests(_World):
    RID = "age-compliance"

    def test_default_shows_exceptions(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E7", "E8", "E10"])
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual((rows["E7"]["flag"], rows["E7"]["age"]), ("DOB missing", None))
        self.assertEqual((rows["E8"]["flag"], rows["E8"]["age"]), ("Under age", 16))
        self.assertEqual((rows["E10"]["flag"], rows["E10"]["age"]), ("Over age", 66))
        self.assertEqual(
            cards(body), {"Under minimum age": 1, "Over maximum age": 1, "DOB missing / invalid": 1, "Total checked": 7}
        )

    def test_show_choices(self):
        self.assertEqual(codes(self.get(self.RID, show="under_min")), ["E8"])
        self.assertEqual(codes(self.get(self.RID, show="over_max")), ["E10"])
        self.assertEqual(codes(self.get(self.RID, show="dob_missing")), ["E7"])
        body = self.get(self.RID, show="all")
        self.assertEqual(len(body["rows"]), 7)
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["E1"]["flag"], "OK")
        self.assertEqual(rows["E1"]["ageAtJoining"], 33)  # born 10-Mar-1990, joined 15-Jan-2024
        self.assertEqual(rows["E8"]["ageAtJoining"], 16)
        self.assertIsNone(rows["E7"]["ageAtJoining"])  # join date unreadable

    def test_thresholds_are_parameters(self):
        self.assertEqual(codes(self.get(self.RID, show="under_min", minAge="26")), ["E2", "E8"])  # E2 is 25
        self.assertEqual(codes(self.get(self.RID, show="under_min", minAge="25")), ["E8"])  # 25 is not under 25
        self.assertEqual(codes(self.get(self.RID, show="over_max", maxAge="30")), ["E1", "E10"])  # E6 is exactly 30
        self.assertEqual(self.get(self.RID, minAge="5", expect=400)["field"], "minAge")

    def test_future_date_of_birth_is_invalid_not_ok(self):
        self.make("E11", date_of_birth=date(2030, 1, 1), branch=self.b1)
        body = self.get(self.RID, show="dob_missing")
        self.assertEqual(codes(body), ["E7", "E11"])
        self.assertEqual(next(r for r in body["rows"] if r["employeeCode"] == "E11")["flag"], "DOB invalid")
        self.assertEqual(cards(body)["DOB missing / invalid"], 2)

    def test_age_proof_is_document_presence(self):
        rows = {r["employeeCode"]: r["birthProof"] for r in self.get(self.RID, show="all")["rows"]}
        self.assertEqual((rows["E1"], rows["E3"], rows["E8"]), ("Yes", "Yes", "No"))  # E1 / E3 uploaded an Aadhaar card
        self.doc(self.e8, "voter_id_or_birth_certificate")
        self.doc(self.e10, "pan_card")  # not an age proof
        rows = {r["employeeCode"]: r["birthProof"] for r in self.get(self.RID, show="all")["rows"]}
        self.assertEqual((rows["E8"], rows["E10"]), ("Yes", "No"))

    def test_scope_filters_and_isolation(self):
        self.assertEqual(len(codes(self.get(self.RID, employeeStatus="all", show="all"))), 10)
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E8"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E8"])
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_cut1.id))), ["E10"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E7"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=str(self.e8.id))), ["E8"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user)), ["E7", "E10"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# family-dependents
# ═════════════════════════════════════════════════════════════════════════════


class FamilyDependentsTests(_World):
    RID = "family-dependents"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        FamilyDependent.objects.create(
            employee=cls.e1, name="Priya", relation="spouse", date_of_birth=date(1992, 5, 1),
            is_insurance_nominee=True, covered_under_health_scheme=True,
        )  # fmt: skip
        FamilyDependent.objects.create(
            employee=cls.e1, name="Kiran", relation="child", date_of_birth=date(2018, 12, 1),
            covered_under_health_scheme=True,
        )  # fmt: skip
        FamilyDependent.objects.create(employee=cls.e2, name="Raman", relation="father")
        FamilyDependent.objects.create(employee=cls.e3, name="Lakshmi", relation="mother", is_insurance_nominee=True)
        FamilyDependent.objects.create(employee=cls.e4, name="Meena", relation="spouse")

    def test_golden_rows_and_summary(self):
        body = self.get(self.RID)
        self.assertEqual([(r["employeeCode"], r["dependentName"]) for r in body["rows"]],
                         [("E1", "Kiran"), ("E1", "Priya"), ("E2", "Raman"), ("E3", "Lakshmi")])  # fmt: skip
        rows = {r["dependentName"]: r for r in body["rows"]}
        self.assertEqual((rows["Priya"]["age"], rows["Kiran"]["age"]), (34, 7))
        self.assertIsNone(rows["Raman"]["age"])  # no date of birth -> blank, not 0
        self.assertEqual(
            (rows["Priya"]["relation"], rows["Priya"]["isInsuranceNominee"], rows["Priya"]["coveredUnderHealthScheme"]),
            ("Spouse", "Yes", "Yes"),
        )
        self.assertEqual((rows["Kiran"]["isInsuranceNominee"], rows["Raman"]["coveredUnderHealthScheme"]), ("No", "No"))
        self.assertEqual(
            cards(body),
            {
                "Dependents": 4,
                "Employees with dependents": 3,
                "Insurance nominees": 2,
                "Covered under health scheme": 2,
            },
        )
        self.assertTrue(any("cannot be added or edited from the HR portal" in n for n in body["notes"]))

    def test_filters(self):
        self.assertEqual([r["dependentName"] for r in self.get(self.RID, relation="child")["rows"]], ["Kiran"])
        self.assertEqual([r["dependentName"] for r in self.get(self.RID, nominee="true")["rows"]], ["Priya", "Lakshmi"])
        self.assertEqual(
            [r["dependentName"] for r in self.get(self.RID, healthScheme="true")["rows"]], ["Kiran", "Priya"]
        )
        self.assertEqual([r["dependentName"] for r in self.get(self.RID, employeeStatus="inactive")["rows"]], ["Meena"])
        self.assertEqual(len(self.get(self.RID, employeeStatus="all")["rows"]), 5)
        self.assertEqual(
            [r["dependentName"] for r in self.get(self.RID, branchIds=str(self.b2.id))["rows"]], ["Lakshmi"]
        )
        self.assertEqual(
            [r["dependentName"] for r in self.get(self.RID, departmentIds=str(self.d_sew1.id))["rows"]], ["Raman"]
        )
        self.assertEqual(
            [r["dependentName"] for r in self.get(self.RID, employmentType="production")["rows"]], ["Raman"]
        )
        self.assertEqual(
            [r["dependentName"] for r in self.get(self.RID, designationIds=str(self.des_mgr.id))["rows"]],
            ["Kiran", "Priya"],
        )
        self.assertEqual(
            [r["dependentName"] for r in self.get(self.RID, employeeIds=str(self.e3.id))["rows"]], ["Lakshmi"]
        )

    def test_empty_table_is_a_valid_report(self):
        FamilyDependent.objects.all().delete()
        body = self.get(self.RID)
        self.assertEqual(body["rows"], [])
        self.assertEqual(cards(body)["Dependents"], 0)

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user)
        self.assertEqual([r["dependentName"] for r in body["rows"]], ["Kiran", "Priya", "Raman"])
        self.assertNotIn("Lakshmi", str(body))
        self.assertEqual(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))["rows"], [])


# ═════════════════════════════════════════════════════════════════════════════
# birthdays / work-anniversaries
# ═════════════════════════════════════════════════════════════════════════════


class BirthdayTests(_World):
    RID = "birthdays"

    def test_this_month_golden(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E10", "E2"])  # 1-Sep then 29-Sep
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(
            (rows["E10"]["birthday"], rows["E10"]["weekday"], rows["E10"]["turningAge"]), ("2026-09-01", "Tuesday", 66)
        )
        self.assertEqual(
            (rows["E2"]["birthday"], rows["E2"]["weekday"], rows["E2"]["turningAge"]), ("2026-09-29", "Tuesday", 25)
        )
        self.assertEqual(rows["E2"]["dateOfBirth"], "2001-09-29")
        self.assertEqual(cards(body), {"Birthdays in window": 2, "No date of birth (not listed)": 1})
        self.assertTrue(any("Window: 01-Sep-2026 to 30-Sep-2026" in n for n in body["notes"]))
        self.assertTrue(any("1 employee(s) in scope have no date of birth" in n for n in body["notes"]))

    def test_named_months_and_windows(self):
        body = self.get(self.RID, window="m12")
        self.assertEqual(
            [(r["employeeCode"], r["birthday"], r["turningAge"]) for r in body["rows"]], [("E6", "2026-12-31", 31)]
        )
        self.assertEqual(codes(self.get(self.RID, window="m05")), ["E8"])
        self.assertEqual(codes(self.get(self.RID, window="nextMonth")), [])  # nobody in October
        self.assertEqual(codes(self.get(self.RID, window="next7")), ["E2"])  # today counts
        self.assertEqual(codes(self.get(self.RID, window="next30")), ["E2"])
        self.assertEqual(self.get(self.RID, window="m13", expect=400)["field"], "window")

    def test_feb_29_birthday_is_celebrated_on_28_feb_in_a_non_leap_year(self):
        body = self.get(self.RID, window="m02")
        self.assertEqual(
            [(r["employeeCode"], r["birthday"], r["turningAge"], r["weekday"]) for r in body["rows"]],
            [("E3", "2026-02-28", 26, "Saturday")],
        )
        self.mock_today.return_value = date(2028, 2, 10)  # 2028 is a leap year
        body = self.get(self.RID)
        self.assertEqual(
            [(r["employeeCode"], r["birthday"], r["turningAge"]) for r in body["rows"]], [("E3", "2028-02-29", 28)]
        )

    def test_window_wraps_from_december_into_january(self):
        self.make("E11", date_of_birth=date(1990, 1, 5), branch=self.b1)
        self.make("E12", date_of_birth=date(1990, 1, 19), branch=self.b1)  # one day past the window
        self.mock_today.return_value = date(2026, 12, 20)
        body = self.get(self.RID, window="next30")
        self.assertEqual([(r["employeeCode"], r["birthday"], r["turningAge"]) for r in body["rows"]],
                         [("E6", "2026-12-31", 31), ("E11", "2027-01-05", 37)])  # fmt: skip
        self.mock_today.return_value = date(2026, 12, 5)
        self.assertEqual(codes(self.get(self.RID, window="nextMonth")), ["E11", "E12"])  # January 2027

    def test_status_filter_and_dates_of_birth_from_the_future_or_this_year_are_skipped(self):
        self.make("E11", date_of_birth=date(2026, 9, 10), branch=self.b1)  # born this year: not "turning" anything yet
        self.make("E12", date_of_birth=date(1970, 9, 12), branch=self.b1, status="inactive")
        self.assertEqual(codes(self.get(self.RID)), ["E10", "E2"])
        self.assertEqual(codes(self.get(self.RID, employeeStatus="inactive")), ["E12"])
        self.assertEqual(codes(self.get(self.RID, employeeStatus="all")), ["E10", "E12", "E2"])

    def test_scope_filters_and_isolation(self):
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_cut1.id))), ["E10"])
        self.assertEqual(codes(self.get(self.RID, employmentType="production")), ["E2"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=str(self.e2.id))), ["E2"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), [])
        self.make("E11", date_of_birth=date(1985, 9, 15), branch=self.b2, department=self.d_cut2)
        self.assertEqual(codes(self.get(self.RID)), ["E10", "E11", "E2"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user)), ["E10", "E2"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


class WorkAnniversaryTests(_World):
    RID = "work-anniversaries"

    def test_default_month_lists_only_completed_years(self):
        body = self.get(self.RID)  # September: E2 and E3 joined this September -> 0 years, not an anniversary
        self.assertEqual(body["rows"], [])
        self.assertEqual(
            cards(body), {"Anniversaries in window": 0, "Join date missing / unreadable (not listed)": 2}
        )  # E7, E10

    def test_named_months(self):
        body = self.get(self.RID, window="m01")
        self.assertEqual(
            [(r["employeeCode"], r["anniversary"], r["yearsCompleted"], r["joinDate"]) for r in body["rows"]],
            [("E1", "2026-01-15", 2, "2024-01-15")],
        )
        body = self.get(self.RID, window="m11")
        self.assertEqual([(r["employeeCode"], r["yearsCompleted"]) for r in body["rows"]], [("E6", 3)])
        self.assertEqual(codes(self.get(self.RID, window="m08")), [])  # E8 joined 31-Aug-2026: zero years

    def test_inactive_employees_have_no_anniversary(self):
        self.assertEqual(codes(self.get(self.RID, window="m03")), [])  # E4 joined 1-Mar-2020 but is inactive
        self.assertEqual(codes(self.get(self.RID, window="m01")), ["E1"])  # E9 (inactive, 1-Jan-2019) excluded

    def test_minimum_years(self):
        self.assertEqual(codes(self.get(self.RID, window="m11", minYears="3")), ["E6"])
        self.assertEqual(codes(self.get(self.RID, window="m11", minYears="4")), [])
        self.assertEqual(self.get(self.RID, minYears="0", expect=400)["field"], "minYears")

    def test_feb_29_join_date_and_year_wrap(self):
        self.make("E11", join_date="2020-02-29", branch=self.b1)
        body = self.get(self.RID, window="m02")
        self.assertEqual(
            [(r["employeeCode"], r["anniversary"], r["yearsCompleted"]) for r in body["rows"]],
            [("E11", "2026-02-28", 6)],
        )
        self.mock_today.return_value = date(2028, 2, 1)
        self.assertEqual(
            [(r["employeeCode"], r["anniversary"]) for r in self.get(self.RID)["rows"]], [("E11", "2028-02-29")]
        )
        self.make("E12", join_date="03-01-2019", branch=self.b1)  # dd-mm-yyyy text: 3-Jan-2019
        self.mock_today.return_value = date(2026, 12, 20)
        body = self.get(self.RID, window="next30")
        self.assertEqual(
            [(r["employeeCode"], r["anniversary"], r["yearsCompleted"]) for r in body["rows"]],
            [("E12", "2027-01-03", 8), ("E1", "2027-01-15", 3)],
        )

    def test_ordering_and_filters(self):
        self.make("E11", join_date="2020-01-15", branch=self.b2, department=self.d_cut2, employment_type="production")
        body = self.get(self.RID, window="m01")
        self.assertEqual(
            [(r["employeeCode"], r["yearsCompleted"]) for r in body["rows"]], [("E11", 6), ("E1", 2)]
        )  # same day: more years first
        self.assertEqual(codes(self.get(self.RID, window="m01", employmentType="production")), ["E11"])
        self.assertEqual(codes(self.get(self.RID, window="m01", branchIds=str(self.b2.id))), ["E11"])
        self.assertEqual(codes(self.get(self.RID, window="m01", departmentIds=str(self.d_cut1.id))), ["E1"])
        self.assertEqual(codes(self.get(self.RID, window="m01", designationIds=str(self.des_mgr.id))), ["E1"])

    def test_branch_isolation(self):
        self.make("E11", join_date="2020-01-15", branch=self.b2, department=self.d_cut2)
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, window="m01")), ["E1"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, window="m01", branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# service-milestones
# ═════════════════════════════════════════════════════════════════════════════


class ServiceMilestoneTests(_World):
    RID = "service-milestones"

    def test_golden_default(self):
        body = self.get(self.RID)
        self.assertEqual(codes(body), ["E1", "E2", "E3", "E6", "E8"])  # E7 / E10 have no readable join date
        rows = {r["employeeCode"]: r for r in body["rows"]}
        e1 = rows["E1"]
        self.assertEqual(
            (e1["serviceMonths"], e1["probationEnd"], e1["probationSource"]), (32, "2024-04-15", "Derived")
        )
        self.assertEqual((e1["clEligibleFrom"], e1["clEligible"]), ("2024-07-15", "Eligible"))
        e6 = rows["E6"]
        self.assertEqual(
            (e6["serviceMonths"], e6["probationEnd"], e6["clEligibleFrom"], e6["clEligible"]),
            (33, "2024-02-29", "2024-05-30", "Eligible"),
        )
        e3 = rows["E3"]
        self.assertEqual(
            (e3["serviceMonths"], e3["probationEnd"], e3["clEligibleFrom"], e3["clEligible"]),
            (0, "2026-12-05", "2027-03-05", "Not yet"),
        )
        self.assertEqual(
            (rows["E2"]["clEligible"], rows["E2"]["clEligibleFrom"]), ("Not applicable", None)
        )  # production
        self.assertEqual(rows["E8"]["probationEnd"], "2026-11-30")  # 31-Aug + 3 months, clamped to the month end
        self.assertEqual(
            cards(body),
            {
                "On probation": 3, "Probation ending in 30 days": 0, "Casual leave eligible": 2,
                "Casual leave not yet eligible": 1, "Join date missing (not listed)": 2,
            },
        )  # fmt: skip

    def test_month_end_joiner_matches_the_casual_leave_rule(self):
        self.make("E11", join_date="2026-03-31", branch=self.b1)
        row = next(r for r in self.get(self.RID)["rows"] if r["employeeCode"] == "E11")
        self.assertEqual(row["clEligibleFrom"], "2026-10-01")  # no 31-Sep: the sixth month completes on 1-Oct
        self.assertEqual((row["serviceMonths"], row["clEligible"]), (5, "Not yet"))  # today is 29-Sep

    def test_milestone_filter(self):
        self.assertEqual(codes(self.get(self.RID, milestone="cl_eligible")), ["E1", "E6"])
        self.assertEqual(codes(self.get(self.RID, milestone="cl_not_yet")), ["E3"])
        self.assertEqual(codes(self.get(self.RID, milestone="probation_over")), ["E1", "E6"])
        self.assertEqual(codes(self.get(self.RID, milestone="probation_ending")), [])

    def test_probation_policy_is_a_parameter(self):
        body = self.get(self.RID, probationMonths="1", milestone="probation_ending")  # 1 month: ends 5-Oct / 30-Sep
        self.assertEqual(codes(body), ["E2", "E3", "E8"])
        self.assertEqual(cards(body)["Probation ending in 30 days"], 3)
        body = self.get(self.RID, probationMonths="24", milestone="probation_over")  # 24 months: E1 ended 15-Jan-2026
        self.assertEqual(codes(body), ["E1", "E6"])
        self.assertEqual(body["rows"][0]["probationEnd"], "2026-01-15")

    def test_recorded_dates_beat_derived_ones(self):
        Employee.objects.filter(pk=self.e3.pk).update(probation_end_date=date(2026, 10, 15))
        Employee.objects.filter(pk=self.e2.pk).update(
            probation_end_date=date(2026, 10, 15), confirmation_date=date(2026, 9, 1)
        )
        body = self.get(self.RID)
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual((rows["E3"]["probationEnd"], rows["E3"]["probationSource"]), ("2026-10-15", "Recorded"))
        self.assertEqual(rows["E2"]["confirmationDate"], "2026-09-01")
        self.assertEqual(cards(body)["On probation"], 2)  # E2 is confirmed -> off probation
        self.assertEqual(codes(self.get(self.RID, milestone="probation_ending")), ["E3"])  # 15-Oct is within 30 days

    def test_people_who_have_left_are_neither_on_probation_nor_casual_leave_eligible(self):
        body = self.get(self.RID, employeeStatus="inactive")
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(sorted(rows), ["E4", "E5", "E9"])
        self.assertEqual(rows["E4"]["clEligible"], "Not applicable")  # staff for 6+ years, but no longer active
        self.assertIsNone(rows["E4"]["clEligibleFrom"])
        self.assertEqual(cards(body)["On probation"], 0)
        self.assertEqual(cards(body)["Casual leave eligible"], 0)

    def test_scope_filters_and_isolation(self):
        self.assertEqual(codes(self.get(self.RID, employmentType="staff")), ["E1", "E3", "E6"])
        self.assertEqual(codes(self.get(self.RID, employeeStatus="inactive")), ["E4", "E5", "E9"])
        self.assertEqual(codes(self.get(self.RID, branchIds=str(self.b2.id))), ["E3", "E8"])
        self.assertEqual(codes(self.get(self.RID, departmentIds=str(self.d_cut1.id))), ["E1"])
        self.assertEqual(codes(self.get(self.RID, designationIds=str(self.des_op.id))), ["E2"])
        self.assertEqual(codes(self.get(self.RID, employeeIds=str(self.e6.id))), ["E6"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user)), ["E1", "E2"])
        self.assertEqual(codes(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])


# ═════════════════════════════════════════════════════════════════════════════
# workforce-profile
# ═════════════════════════════════════════════════════════════════════════════


class WorkforceProfileTests(_World):
    RID = "workforce-profile"

    @staticmethod
    def table(body):
        keys = ("staff", "production", "male", "female", "otherOrUnset", "total")
        return {r["bucket"]: tuple(r[k] for k in keys) for r in body["rows"]}

    def test_age_bands_golden(self):
        body = self.get(self.RID)
        self.assertEqual(
            [r["bucket"] for r in body["rows"]], ["Under 18", "18-25", "26-35", "36-45", "46-55", "56+", "Unknown"]
        )
        self.assertEqual(
            self.table(body),
            {
                "Under 18": (0, 1, 1, 0, 0, 1), "18-25": (0, 1, 0, 1, 0, 1), "26-35": (2, 0, 0, 2, 0, 2),
                "36-45": (1, 0, 1, 0, 0, 1), "46-55": (0, 0, 0, 0, 0, 0), "56+": (1, 0, 1, 0, 0, 1),
                "Unknown": (1, 0, 0, 0, 1, 1),
            },
        )  # fmt: skip
        self.assertEqual(body["totals"]["total"], 7)
        self.assertEqual(sum(r["total"] for r in body["rows"]), 7)
        self.assertEqual(
            cards(body),
            {
                "Employees": 7,
                "Average age (years)": 33.2,
                "Average service (years)": 1.1,
                "Date of birth unknown": 1,
                "Join date unknown": 2,
            },
        )

    def test_tenure_bands(self):
        body = self.get(self.RID, dimension="tenureBand")
        self.assertEqual(
            self.table(body),
            {
                "Under 6 months": (1, 2, 1, 2, 0, 3), "6-12 months": (0, 0, 0, 0, 0, 0), "1-3 years": (2, 0, 1, 1, 0, 2),
                "3-5 years": (0, 0, 0, 0, 0, 0), "5+ years": (0, 0, 0, 0, 0, 0), "Unknown": (2, 0, 1, 0, 1, 2),
            },
        )  # fmt: skip

    def test_gender_blood_group_and_salary_type(self):
        body = self.get(self.RID, dimension="gender")
        self.assertEqual(
            {k: v[-1] for k, v in self.table(body).items()}, {"Male": 3, "Female": 3, "Other": 0, "Not set": 1}
        )
        body = self.get(self.RID, dimension="bloodGroup")
        t = self.table(body)
        self.assertEqual({k: v[-1] for k, v in t.items() if v[-1]}, {"A+": 1, "B+": 1, "AB-": 1, "O+": 2, "Unknown": 2})
        self.assertEqual([r["bucket"] for r in body["rows"]][-1], "Unknown")
        body = self.get(self.RID, dimension="salaryType")
        self.assertEqual(
            self.table(body),
            {"Monthly": (5, 0, 2, 2, 1, 5), "Weekly": (0, 2, 1, 1, 0, 2), "Unknown": (0, 0, 0, 0, 0, 0)},
        )

    def test_unrecognised_values_get_their_own_row_instead_of_vanishing(self):
        self.make("E11", salary_type="fortnightly", branch=self.b1)
        body = self.get(self.RID, dimension="salaryType")
        self.assertEqual(self.table(body)["Fortnightly"][-1], 1)
        self.assertEqual(sum(r["total"] for r in body["rows"]), 8)

    def test_filters_and_status(self):
        self.assertEqual(self.get(self.RID, employeeStatus="all")["totals"]["total"], 10)
        self.assertEqual(self.get(self.RID, employeeStatus="inactive")["totals"]["total"], 3)
        self.assertEqual(self.get(self.RID, departmentIds=str(self.d_cut1.id))["totals"]["total"], 2)
        self.assertEqual(self.get(self.RID, branchIds=str(self.b2.id))["totals"]["total"], 2)
        self.assertEqual(self.get(self.RID, dimension="bogus", expect=400)["field"], "dimension")

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user)
        self.assertEqual(body["totals"]["total"], 4)  # E1, E2, E7, E10
        self.assertEqual(self.table(body)["Unknown"][-1], 1)
        self.assertEqual(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))["totals"]["total"], 0)


# ═════════════════════════════════════════════════════════════════════════════
# org-structure
# ═════════════════════════════════════════════════════════════════════════════


class OrgStructureTests(_World):
    RID = "org-structure"

    def test_golden(self):
        body = self.get(self.RID)
        labels = [(r["branch"] if r.get("_kind") != "subtotal" else "-", r["department"]) for r in body["rows"]]
        self.assertEqual(
            labels,
            [
                ("Head Office", "No departments set up"), ("-", "Head Office total"),  # seeded by migration 0036
                ("Old Unit (inactive)", "SCRAP"), ("-", "Old Unit (inactive) total"),
                ("Unit 1", "CUTTING"), ("Unit 1", "SEWING"), ("-", "Unit 1 total"),
                ("Unit 2", "CUTTING"), ("-", "Unit 2 total"),
                ("No branch", "LEGACY"), ("No branch", "Unassigned (no department)"), ("-", "No branch total"),
            ],
        )  # fmt: skip
        by = {(r["branch"], r["department"]): r for r in body["rows"] if r.get("_kind") != "subtotal"}
        keys = ("designations", "staff", "production", "hods", "required", "vacancy")
        self.assertEqual(tuple(by[("Unit 1", "CUTTING")][k] for k in keys), (1, 2, 0, "Jai Singh", 5, 3))
        self.assertEqual(
            tuple(by[("Unit 1", "SEWING")][k] for k in keys), (2, 1, 1, "Arun Kumar", 1, 0)
        )  # plan met: vacancy 0, not blank
        self.assertEqual(
            tuple(by[("Unit 2", "CUTTING")][k] for k in keys), (0, 1, 1, None, None, None)
        )  # inactive HOD; plan 0 = not planned
        self.assertEqual(by[("Old Unit (inactive)", "SCRAP")]["branchCode"], "OLD")
        self.assertEqual(by[("No branch", "Unassigned (no department)")]["staff"], 1)  # E6 has no department
        sub = {r["department"]: r for r in body["rows"] if r.get("_kind") == "subtotal"}
        self.assertEqual(
            (
                sub["Unit 1 total"]["designations"],
                sub["Unit 1 total"]["staff"],
                sub["Unit 1 total"]["production"],
                sub["Unit 1 total"]["required"],
                sub["Unit 1 total"]["vacancy"],
            ),
            (3, 3, 1, 6, 3),
        )
        self.assertEqual(body["totals"], {"designations": 3, "staff": 5, "production": 2, "required": 6, "vacancy": 3})
        self.assertEqual(
            cards(body),
            {"Branches": 4, "Departments": 5, "Active employees": 7, "Planned positions": 6, "Open vacancies": 3},
        )  # Head Office + Unit 1 + Unit 2 + Old Unit: 'No branch' is a bucket for legacy rows, not a fifth branch

    def test_totals_equal_the_rows(self):
        body = self.get(self.RID)
        rows = data_rows(body)
        for k in ("designations", "staff", "production", "required", "vacancy"):
            self.assertEqual(body["totals"][k], sum(r[k] or 0 for r in rows), k)
        self.assertEqual(
            body["totals"]["staff"] + body["totals"]["production"], 7
        )  # every active employee is counted once

    def test_subtotals_add_up_to_the_grand_total(self):
        body = self.get(self.RID)
        subs = [r for r in body["rows"] if r.get("_kind") == "subtotal"]
        self.assertEqual(len(subs), 5)  # Head Office (no departments yet), Old Unit, Unit 1, Unit 2, No branch
        for k in ("designations", "staff", "production", "required", "vacancy"):
            self.assertEqual(sum(r[k] for r in subs), body["totals"][k], k)

    def test_inactive_employees_are_not_counted_and_hod_row_must_be_active(self):
        rows = {(r["branch"], r["department"]): r for r in data_rows(self.get(self.RID))}
        self.assertEqual(rows[("Unit 1", "SEWING")]["production"], 1)  # E5 (inactive production) not counted
        self.assertEqual(rows[("Unit 1", "SEWING")]["staff"], 1)  # E4 (inactive) not counted

    def test_branch_and_department_filters(self):
        body = self.get(self.RID, branchIds=str(self.b2.id))
        self.assertEqual([(r["branch"], r["department"]) for r in data_rows(body)], [("Unit 2", "CUTTING")])
        self.assertEqual((data_rows(body)[0]["staff"], data_rows(body)[0]["production"]), (1, 1))
        body = self.get(self.RID, departmentIds=str(self.d_sew1.id))
        self.assertEqual([(r["branch"], r["department"]) for r in data_rows(body)], [("Unit 1", "SEWING")])

    def test_branch_isolation(self):
        body = self.get(self.RID, user=self.branch_user)
        self.assertEqual(
            [(r["branch"], r["department"]) for r in data_rows(body)], [("Unit 1", "CUTTING"), ("Unit 1", "SEWING")]
        )
        self.assertEqual((body["totals"]["staff"], body["totals"]["production"]), (3, 1))
        text = str(body)
        for hidden in ("Unit 2", "Old Unit", "LEGACY", "SCRAP", "No branch"):
            self.assertNotIn(hidden, text)
        self.assertEqual(data_rows(self.get(self.RID, user=self.branch_user, branchIds=str(self.b2.id))), [])

    def test_employee_in_a_department_of_another_branch_still_counts(self):
        # legacy data: an employee of Unit 1 sitting in Unit 2's department must not vanish from the totals ...
        self.make("E11", branch=self.b1, department=self.d_cut2)
        body = self.get(self.RID, user=self.branch_user)
        outside = [r for r in data_rows(body) if r["department"].startswith("Departments outside this listing")]
        # ... but they are counted under the user's OWN branch, on a line that names nothing of the other unit
        self.assertEqual([(r["branch"], r["staff"], r["production"]) for r in outside], [("Unit 1", 1, 0)])
        self.assertEqual(body["totals"]["staff"] + body["totals"]["production"], 5)
        self.assertNotIn("Unit 2", str(body))
        self.assertTrue(any("not part of this listing" in n for n in body["notes"]))


# ═════════════════════════════════════════════════════════════════════════════
# Cross-cutting contract: exports, empty results, query counts, isolation
# ═════════════════════════════════════════════════════════════════════════════

WIDE = {
    "employee-master": {"employeeStatus": "all", "layout": "full"},
    "employee-contact-list": {"employeeStatus": "all"},
    "strength-statement": {"employeeStatus": "all"},
    "new-joinings": {"dateFrom": "2019-01-01", "dateTo": "2026-12-31"},
    "exits-register": {"dateFrom": "2025-01-01", "dateTo": "2026-12-31"},
    "resignation-register": {"dateFrom": "2026-08-01", "dateTo": "2026-09-29", "layout": "full"},
    "manpower-movement": {"year": "2026"},
    "statutory-compliance": {"employeeStatus": "all", "includeCompliant": "true"},
    "data-quality-audit": {"employeeStatus": "all", "includeClean": "true"},
    "birthdays": {"window": "next30"},
    "work-anniversaries": {"window": "thisMonth"},
    "service-milestones": {"employeeStatus": "all"},
    "age-compliance": {"employeeStatus": "all", "show": "all"},
    "workforce-profile": {"employeeStatus": "all"},
    "family-dependents": {"employeeStatus": "all"},
    "org-structure": {},
}
# Reports that always return a fixed grid of rows (all 12 months / all age bands) even when nothing matches.
FIXED_GRID = {"manpower-movement", "workforce-profile"}


class ContractTests(_World):
    def test_xlsx_round_trip_for_every_report(self):
        for rid in MY_REPORTS:
            body = self.get(rid, **WIDE[rid])
            r = self.export(rid, "xlsx", **WIDE[rid])
            self.assertTrue(r.content.startswith(b"PK"), rid)
            ws = load_workbook(io.BytesIO(r.content)).active
            header = [c.value for c in ws[7]]
            self.assertEqual(header, [c["label"] for c in body["columns"]], rid)
            self.assertIsNotNone(ws["A3"].value)
            if not body["rows"]:
                continue
            # the first TEXT column of the first data row round-trips exactly
            idx = next(
                (i for i, c in enumerate(body["columns"]) if c["type"] == "text" and body["rows"][0].get(c["key"])),
                None,
            )
            if idx is not None:
                self.assertEqual(
                    ws.cell(row=8, column=idx + 1).value, body["rows"][0][body["columns"][idx]["key"]], rid
                )
            self.assertGreaterEqual(ws.max_row, 8 + len(body["rows"]) - 1, rid)

    def test_pdf_for_every_report(self):
        for rid in MY_REPORTS:
            r = self.export(rid, "pdf", **WIDE[rid])
            self.assertTrue(r.content.startswith(b"%PDF"), rid)
            self.assertIn("application/pdf", r["Content-Type"], rid)
            self.assertGreater(len(r.content), 1500, rid)

    def test_identifiers_stay_text_cells_in_excel(self):
        ws = load_workbook(
            io.BytesIO(self.export("statutory-compliance", "xlsx", **WIDE["statutory-compliance"]).content)
        ).active
        header = [c.value for c in ws[7]]
        col = header.index("UAN") + 1
        cells = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=col) for r in range(8, ws.max_row + 1)}
        self.assertEqual(cells["E1"].value, "100000000001")
        self.assertEqual(cells["E1"].data_type, "s")
        self.assertEqual(ws.cell(row=8, column=1).data_type, "s")

    def test_empty_results_work_on_screen_and_in_both_exports(self):
        for rid in MY_REPORTS:
            params = {**WIDE[rid], "departmentIds": "999999"}
            body = self.get(rid, **params)
            if rid not in FIXED_GRID:
                self.assertEqual(body["rows"], [], rid)
                self.assertEqual(body["rowCount"], 0, rid)
            else:
                self.assertTrue(
                    all(
                        all(not v for k, v in r.items() if k not in ("group", "bucket", "exitRatePct", "sharePct"))
                        for r in body["rows"]
                    ),
                    rid,
                )
            for fmt in ("xlsx", "pdf"):
                self.assertEqual(self.export(rid, fmt, **params).status_code, 200, (rid, fmt))

    def test_no_query_count_growth_with_more_employees(self):
        # Warm caches first (roles, settings), then compare 10 employees against 22.
        variants = {rid: [WIDE[rid]] for rid in MY_REPORTS}
        variants["strength-statement"] = [
            {"groupBy": g, "employeeStatus": "all", "includeEmpty": "true"}
            for g in ("department", "designation", "branch")
        ]
        variants["manpower-movement"] = [{"year": "2026", "groupBy": g} for g in ("month", "department")]
        variants["workforce-profile"] = [
            {"dimension": d, "employeeStatus": "all"}
            for d in ("ageBand", "tenureBand", "gender", "bloodGroup", "salaryType")
        ]
        variants["work-anniversaries"] = [{"window": "thisMonth"}, {"window": "next30"}]

        def measure():
            out = {}
            for rid, plist in variants.items():
                for i, params in enumerate(plist):
                    self.client.get(f"/api/reports/run/{rid}", params, **hdr(self.admin))  # warm-up
                    with CaptureQueriesContext(connection) as q:
                        r = self.client.get(f"/api/reports/run/{rid}", params, **hdr(self.admin))
                    self.assertEqual(r.status_code, 200, (rid, params))
                    out[(rid, i)] = len(q)
            return out

        before = measure()
        for i in range(12):
            e = self.make(
                f"X{i}", first_name=f"Xa{i}", employment_type="staff" if i % 2 else "production", gender=("male", "female", None)[i % 3],
                status="active" if i % 3 else "inactive", branch=(self.b1, self.b2)[i % 2],
                department=(self.d_cut1, self.d_sew1, self.d_cut2)[i % 3], designation=(None, self.des_op, self.des_mgr)[i % 3],
                join_date=f"2026-09-{i + 1:02d}", date_of_birth=date(1985 + i, (i % 12) + 1, i + 1), phone=f"90000000{i:02d}",
                pf_number=f"PF{i}", bank_account=f"1234567{i:03d}", bank_ifsc="HDFC0001234", salary_amount=Decimal("1000"),
            )  # fmt: skip
            self.doc(e, "pan_card")
            FamilyDependent.objects.create(employee=e, name=f"Dep{i}", relation="child", date_of_birth=date(2015, 1, 1))
            if e.status != "active":
                ResignationRequest.objects.create(
                    employee=e, status="approved", last_working_date=date(2026, 9, 1 + i), reason="x"
                )
            else:
                ResignationRequest.objects.create(employee=e, status="pending", reason="y")
            SalarySlip.objects.create(employee=e, month=8, year=2026, slip_number=f"S-X{i}", pf_deduction=10)
        after = measure()
        for key, n in before.items():
            self.assertLessEqual(after[key] - n, 2, f"{key}: {n} -> {after[key]} queries")

    def test_branch_isolation_across_every_employee_level_report(self):
        # Give every date-driven report something to show in BOTH units.
        self.make(
            "E11", date_of_birth=date(1985, 10, 3), join_date="2020-09-12", branch=self.b2, department=self.d_cut2
        )
        self.make(
            "E12", date_of_birth=date(1985, 10, 4), join_date="2020-09-13", branch=self.b1, department=self.d_cut1
        )
        FamilyDependent.objects.create(employee=self.e1, name="Priya", relation="spouse")
        FamilyDependent.objects.create(employee=self.e3, name="Lakshmi", relation="mother")
        wide = {**WIDE, "birthdays": {"window": "next30"}, "work-anniversaries": {"window": "thisMonth"}}
        b2_only = {"E3", "E6", "E8", "E9", "E11", "Lakshmi"}
        for rid in MY_REPORTS:
            if rid in ("strength-statement", "manpower-movement", "workforce-profile", "org-structure"):
                continue
            admin_body = self.get(rid, **wide[rid])
            scoped = self.get(rid, user=self.branch_user, **wide[rid])
            self.assertTrue(
                b2_only & {str(v) for r in admin_body["rows"] for v in r.values()},
                f"{rid}: fixture shows nothing from other units",
            )
            leaked = b2_only & {str(v) for r in scoped["rows"] for v in r.values()}
            self.assertFalse(leaked, f"{rid} leaks {leaked}")
            # a scoped user cannot widen the scope with params
            for extra in (
                {"branchIds": str(self.b2.id)},
                {"employeeIds": ",".join(str(e.id) for e in (self.e3, self.e6, self.e8, self.e9))},
            ):
                widened = self.get(rid, user=self.branch_user, **wide[rid], **extra)
                self.assertFalse(b2_only & {str(v) for r in widened["rows"] for v in r.values()}, f"{rid} {extra}")

    def test_employee_and_designation_filters_narrow_every_register(self):
        registers = (
            "employee-master", "employee-contact-list", "statutory-compliance", "data-quality-audit", "age-compliance",
            "service-milestones", "resignation-register",
        )  # fmt: skip
        for rid in registers:
            body = self.get(rid, employeeIds=str(self.e1.id), **WIDE[rid])
            self.assertEqual(codes(body), ["E1"], rid)
        for rid in (*registers, "new-joinings"):
            body = self.get(rid, designationIds=str(self.des_mgr.id), **WIDE[rid])
            self.assertEqual(codes(body), ["E1"], rid)  # Manager is held by E1 alone

    def test_totals_row_equals_the_sum_of_the_rows_for_every_report(self):
        for rid in MY_REPORTS:
            body = self.get(rid, **WIDE[rid])
            summed = [c for c in body["columns"] if c["total"] == "sum"]
            if not summed:
                self.assertIsNone(body["totals"], rid)
                continue
            for c in summed:
                want = sum(r[c["key"]] or 0 for r in data_rows(body))
                self.assertEqual(body["totals"][c["key"]], want, (rid, c["key"]))

    def test_a_large_register_paginates_and_stays_complete(self):
        Employee.objects.bulk_create(
            Employee(
                employee_code=f"B{i:04d}", first_name=f"Bulk{i}", last_name="Person", employment_type="staff",
                status="active", branch=self.b1, department=self.d_cut1, join_date="2025-01-01",
                date_of_birth=date(1990, 1, 1),
            )
            for i in range(300)
        )  # fmt: skip
        body = self.get("employee-master", employeeStatus="all")
        self.assertEqual(body["rowCount"], 310)
        self.assertEqual(
            [r["employeeCode"] for r in body["rows"]][:2], ["B0000", "B0001"]
        )  # 'B...' sorts before 'E...'
        self.assertEqual(body["rows"][-1]["employeeCode"], "E10")
        r = self.export("employee-master", "pdf", employeeStatus="all")
        self.assertGreater(r.content.count(b"/Type /Page\n") + r.content.count(b"/Type /Page "), 3)
        ws = load_workbook(io.BytesIO(self.export("employee-master", "xlsx", employeeStatus="all").content)).active
        listed = [ws.cell(row=n, column=1).value for n in range(8, ws.max_row + 1)]
        self.assertEqual(len([c for c in listed if c and c[0] in "BE" and c[1:].isdigit()]), 310)

    def test_photo_and_password_columns_are_never_loaded(self):
        """Employee.photo_url is a ~40 KB base64 string per row and password_hash is a secret: no report may even SELECT them."""
        for rid in MY_REPORTS:
            with CaptureQueriesContext(connection) as q:
                self.get(rid, **WIDE[rid])
                self.export(rid, "xlsx", **WIDE[rid])
            for query in q.captured_queries:
                sql = query["sql"]
                self.assertNotIn("photo_url", sql, rid)
                if '"hr_users"' not in sql:  # the login lookup reads the HR user's own hash; nothing else may read one
                    self.assertNotIn("password_hash", sql, rid)

    def test_reports_never_write(self):
        from django.apps import apps

        from .models import PayrollSettings

        PayrollSettings.get()  # the export letterhead reads this singleton (creating it once): make it exist up front

        def snapshot():
            counts = {
                m.__name__: m.objects.count()
                for m in apps.get_app_config("api").get_models()
                if m.__name__ != "AuditLog"
            }  # exports are audited on purpose, so the audit trail is the one table allowed to grow
            return counts, list(Employee.objects.values_list("pk", "updated_at").order_by("pk"))

        before = snapshot()
        for rid in MY_REPORTS:
            self.get(rid, **WIDE[rid])
            for fmt in ("xlsx", "pdf"):
                self.export(rid, fmt, **WIDE[rid])
        self.assertEqual(before, snapshot())

    def test_exports_are_audited_and_row_counts_match(self):
        from .models import AuditLog

        n = AuditLog.objects.filter(action="export", module="reports").count()
        self.export("employee-master", "xlsx", employeeStatus="all")
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), n + 1)
        log = AuditLog.objects.filter(action="export", module="reports").latest("id")
        self.assertIn("Employee Master Register", log.record_description)
        self.assertIn("10 rows", log.record_description)


# ═════════════════════════════════════════════════════════════════════════════
# Adversarial review. Every test states the CORRECT behaviour: a failing test is a defect in the report.
# ═════════════════════════════════════════════════════════════════════════════


class AdversarialReviewTests(_World):
    def _employees_only_user(self, name="em_emp_only_adv"):
        return HRUser.objects.create(
            username=name,
            password_hash="x",
            role=Role.objects.create(name=name, permissions={"reports": "view", "employees": "view"}),
        )

    # -- org-structure -----------------------------------------------------------------------
    def test_org_structure_hides_another_units_plan_and_hods_from_a_scoped_user(self):
        DepartmentHeadcount.objects.filter(department=self.d_cut2).update(required_count=9)
        boss = self.make("E20", branch=self.b2, department=self.d_cut2)
        ManagerDepartmentAssignment.objects.create(
            manager=DepartmentManager.objects.create(employee=boss), department=self.d_cut2
        )
        self.make("E11", branch=self.b1, department=self.d_cut2)  # legacy: a Unit 1 employee in Unit 2's department
        scoped = self.get("org-structure", user=self.branch_user)
        for r in data_rows(scoped):
            if r["branch"] == "Unit 2":
                self.assertIsNone(r["hods"], "Unit 2's HOD is shown to a Unit 1 user")
                self.assertIsNone(r["required"], "Unit 2's manpower plan is shown to a Unit 1 user")
                self.assertIsNone(r["designations"], "Unit 2's designation count is shown to a Unit 1 user")
        self.assertNotIn("E20 Extra", str(scoped))

    def test_org_structure_vacancy_does_not_depend_on_who_is_asking(self):
        DepartmentHeadcount.objects.filter(department=self.d_cut2).update(required_count=5)
        self.make("E11", branch=self.b1, department=self.d_cut2)
        admin = next(r for r in data_rows(self.get("org-structure")) if r["branch"] == "Unit 2")
        scoped = [r for r in data_rows(self.get("org-structure", user=self.branch_user)) if r["branch"] == "Unit 2"]
        for r in scoped:
            self.assertEqual(r["vacancy"], admin["vacancy"])

    def test_org_structure_branch_card_counts_real_branches_only(self):
        # Unit 1, Unit 2 and the closed Old Unit. 'No branch' is a bucket for legacy rows, not a branch. (Migration 0036
        # seeds a real 'Head Office' branch into every database: take it out so the count is the three built here.)
        Branch.objects.filter(is_head_office=True).delete()
        self.assertEqual(Branch.objects.count(), 3)
        self.assertEqual(cards(self.get("org-structure"))["Branches"], 3)

    def test_org_structure_lists_a_branch_that_has_no_departments_yet(self):
        Branch.objects.create(name="Empty Unit", code="EMP")
        self.assertIn("Empty Unit", {r["branch"] for r in data_rows(self.get("org-structure"))})

    # -- labels that cannot be told apart -------------------------------------------------------
    def test_manpower_movement_department_rows_can_be_told_apart(self):
        labels = [r["group"] for r in self.get("manpower-movement", year="2026", groupBy="department")["rows"]]
        self.assertEqual(len(labels), len(set(labels)), f"two rows carry the same label: {labels}")

    def test_strength_designation_rows_can_be_told_apart(self):
        sew2 = Department.objects.create(name="SEWING", branch=self.b2)
        op2 = Designation.objects.create(title="Operator", department=sew2)
        self.make("E21", branch=self.b2, department=sew2, designation=op2)
        rows = self.get("strength-statement", groupBy="designation")["rows"]
        keys = [(r["group"], r["department"]) for r in rows]
        self.assertEqual(len(keys), len(set(keys)), f"identical-looking rows: {keys}")

    # -- service length of people who have left ---------------------------------------------------
    def test_a_leavers_tenure_stops_at_the_exit_date(self):
        rows = {r["employeeCode"]: r for r in self.get("employee-master", employeeStatus="inactive")["rows"]}
        # E9 joined 1-Jan-2019 and left about 16-Jan-2026: 7 years. As of today it would read 7y 8m.
        self.assertIn(rows["E9"]["tenure"], (None, "7y"))
        # E4 joined 1-Mar-2020, last working day 31-Aug-2026: 77 months = 6y 5m (what the exits register says).
        self.assertIn(rows["E4"]["tenure"], (None, "6y 5m"))

    def test_workforce_profile_bands_a_leaver_by_service_at_exit(self):
        e = self.make("E30", status="inactive", branch=self.b1, department=self.d_cut1, join_date="2026-01-01")
        Employee.objects.filter(pk=e.pk).update(updated_at=utc(2026, 2, 1, 5, 0))  # left about 1-Feb-2026: one month
        by = {
            r["bucket"]: r
            for r in self.get("workforce-profile", dimension="tenureBand", employeeStatus="inactive")["rows"]
        }
        self.assertEqual(by["Under 6 months"]["total"], 1)

    def test_service_milestones_service_months_of_a_leaver_stop_at_the_exit(self):
        e = self.make("E30", status="inactive", branch=self.b1, department=self.d_cut1, join_date="2026-01-01")
        Employee.objects.filter(pk=e.pk).update(updated_at=utc(2026, 2, 1, 5, 0))
        rows = {r["employeeCode"]: r for r in self.get("service-milestones", employeeStatus="inactive")["rows"]}
        self.assertLessEqual(rows["E30"]["serviceMonths"], 1)

    # -- statutory-compliance ----------------------------------------------------------------
    def test_placeholder_identifiers_are_missing_numbers_not_duplicates(self):
        for code in ("P1", "P2"):
            self.make(code, branch=self.b1, pf_number="NA", esi_number="N/A", uan_number="-")
        rows = {r["employeeCode"]: r for r in self.get("statutory-compliance")["rows"]}
        for code in ("P1", "P2"):
            issues = rows[code]["issues"]
            self.assertIn("PF number missing", issues)
            self.assertIn("ESI number missing", issues)
            self.assertIn("UAN missing", issues)
            self.assertNotIn("shared with", issues)

    def test_esi_number_is_not_demanded_of_staff_above_the_esi_wage_ceiling(self):
        e = self.make(
            "HIGH1", branch=self.b1, department=self.d_cut1, salary_amount=Decimal("50000"), pf_number="TN/TPR/900",
            uan_number="100000000099", bank_account="1234567890123", bank_ifsc="HDFC0001234",
        )  # fmt: skip
        for cat in ("aadhaar_card", "pan_card", "bank_passbook"):
            self.doc(e, cat)
        rows = {r["employeeCode"]: r for r in self.get("statutory-compliance", includeCompliant="true")["rows"]}
        self.assertNotIn("ESI number missing", rows["HIGH1"]["issues"] or "")

    # -- permissions / leaks ----------------------------------------------------------------------
    def test_exits_register_hides_resignation_reasons_from_a_role_without_the_resignations_module(self):
        user = self._employees_only_user()
        self.get("resignation-register", user=user, expect=403)  # the same reasons are gated there
        body = self.get("exits-register", user=user, dateFrom="2025-01-01", dateTo="2026-12-31")
        e4 = next(r for r in body["rows"] if r["employeeCode"] == "E4")
        # the columns are dropped altogether (not left as a column of dashes) and the values never leave the server
        shown = {c["key"] for c in body["columns"]}
        for key in ("reason", "approvedBy", "approvedOn"):
            self.assertNotIn(key, shown)
            self.assertNotIn(key, e4)
        self.assertNotIn("Better opportunity", str(body))
        self.assertNotIn("HR Admin", str(body))
        self.assertEqual(e4["exitBasis"], "Resignation")  # the dated exit itself is not confidential
        self.assertTrue(any("no access to Resignations" in n for n in body["notes"]))
        for fmt in ("xlsx", "pdf"):
            self.export("exits-register", fmt, user=user, dateFrom="2025-01-01", dateTo="2026-12-31")
        ws = load_workbook(
            io.BytesIO(
                self.export("exits-register", "xlsx", user=user, dateFrom="2025-01-01", dateTo="2026-12-31").content
            )
        ).active
        cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertFalse(any("Better opportunity" in c or "Reason" == c for c in cells))

    # The filter summary is written by the framework (reporting/filters.py describe_params); it scopes the
    # employee / department / branch name lookups to the viewer's branch, so an id from another unit names nothing.
    def test_filter_summary_does_not_name_employees_or_units_outside_the_users_branch(self):
        body = self.get(
            "employee-master", user=self.branch_user, employeeIds=str(self.e3.id), branchIds=str(self.b2.id),
            departmentIds=str(self.d_scrap.id),
        )  # fmt: skip
        text = str(body["filters"])
        for other_unit in ("Chitra", "Unit 2", "SCRAP"):
            self.assertNotIn(other_unit, text)

    # -- disclosure -------------------------------------------------------------------------------
    def test_manpower_movement_says_how_many_employees_have_no_usable_join_date(self):
        notes = [n.lower() for n in self.get("manpower-movement", year="2026")["notes"]]
        # E7 ('soon') and E10 (blank) can never be counted as joiners of any year
        self.assertTrue(
            any("join date" in n and ("unreadable" in n or "missing" in n or "could not" in n) for n in notes), notes
        )

    # -- hostile data ---------------------------------------------------------------------------
    def _every_report_variant(self):
        for rid in MY_REPORTS:
            yield rid, dict(WIDE[rid])
        for g in ("designation", "branch"):
            yield "strength-statement", {"groupBy": g, "employeeStatus": "all", "includeEmpty": "true"}
        yield "manpower-movement", {"year": "2026", "groupBy": "department"}
        for d in ("tenureBand", "gender", "bloodGroup", "salaryType"):
            yield "workforce-profile", {"dimension": d, "employeeStatus": "all"}
        for w in ("thisMonth", "nextMonth", "next7", "next30", "m02", "m12"):
            yield "birthdays", {"window": w, "employeeStatus": "all"}
            yield "work-anniversaries", {"window": w}
        for m in ("probation_ending", "probation_over", "cl_eligible", "cl_not_yet"):
            yield "service-milestones", {"employeeStatus": "all", "milestone": m}
        yield "resignation-register", {"dateFrom": "2020-01-01", "dateTo": "2026-12-31", "layout": "full"}
        yield "data-quality-audit", {"employeeStatus": "all", "includeClean": "true", "check": "dob,phone,duplicates"}

    def test_hostile_but_plausible_data_never_crashes_a_report_or_an_export(self):
        junk = [
            dict(code="W1", first_name="<b>&Bold", last_name="O'Neil", join_date="2026-13-45", gender="M", phone="abc",
                 blood_group="??", pf_number="=1+1", esi_number="@SUM(A1)", uan_number="+123", bank_account="0012345678",
                 bank_ifsc="hdfc0abc123", id_proof="12", date_of_birth=date(1900, 1, 1)),
            dict(code=" W3 ", first_name="", last_name="", join_date="   ", employment_type="", status="", gender="  ",
                 date_of_birth=date(2026, 9, 29)),
            dict(code="W4", join_date="2026-09-05T10:00:00+05:30", salary_amount=Decimal("99999999.99"),
                 salary_per_shift=Decimal("0"), date_of_birth=date(2027, 1, 1)),
            dict(code="W5", employment_type="Contract", status="Active", join_date="31/12/2026",
                 date_of_birth=date(2000, 2, 29), blood_group="ab +", gender="Female "),
            dict(code="W6", join_date="1-1-2026", status="inactive", date_of_birth=date(1970, 1, 1)),
        ]  # fmt: skip
        for kw in junk:
            e = self.make(kw.pop("code"), branch=self.b1, department=self.d_cut1, **kw)
            FamilyDependent.objects.create(employee=e, name="<i>Kid</i>", relation="child")
            ResignationRequest.objects.create(employee=e, status="approved", last_working_date=date(2020, 1, 1))
            ResignationRequest.objects.create(employee=e, status="rejected", rejected_by=None, reason="<x>&")
        for user in (self.admin, self.branch_user):
            for rid, params in self._every_report_variant():
                self.get(rid, user=user, **params)
                for fmt in ("xlsx", "pdf"):
                    self.export(rid, fmt, user=user, **params)

    def test_an_absurd_join_year_does_not_take_service_milestones_down(self):
        self.make("W9", branch=self.b1, department=self.d_cut1, join_date="9999-12-31")
        self.get("service-milestones", employeeStatus="all")
        self.get("service-milestones", employeeStatus="all", milestone="probation_ending")
        self.get("employee-master", employeeStatus="all")
        self.get("work-anniversaries", window="m12")

    def test_a_rejoiner_who_left_again_is_not_dated_by_the_first_resignation(self):
        # Rejoining is common: E31 resigned (LWD 31-Mar-2024), came back on 15-Jun-2024 with the same record, and
        # was switched off again this month. The 2024 resignation belongs to the first stint.
        e = self.make("E31", status="inactive", branch=self.b1, department=self.d_cut1, join_date="2024-06-15")
        ResignationRequest.objects.create(
            employee=e, status="approved", last_working_date=date(2024, 3, 31), approved_at=utc(2024, 3, 1, 5, 0)
        )
        Employee.objects.filter(pk=e.pk).update(updated_at=utc(2026, 9, 10, 8, 0))
        body = self.get("exits-register", dateFrom="2026-09-01", dateTo="2026-09-29")
        self.assertIn("E31", codes(body))


# ═════════════════════════════════════════════════════════════════════════════
# Fixes for the adversarial review: each test pins a behaviour that was wrong (or that the review's own test does
# not reach: branch overlays, filters, exports, other report variants).
# ═════════════════════════════════════════════════════════════════════════════


class ReviewFixTests(_World):
    # -- labels ------------------------------------------------------------------------------------
    def test_department_labels_only_add_the_unit_where_two_names_clash(self):
        class D:  # the three attributes department_labels reads
            def __init__(self, id, name, branch_id):
                self.id, self.name, self.branch_id = id, name, branch_id

        b1 = Branch.objects.create(name="Label Unit 1", code="LU1")
        b2 = Branch.objects.create(name="Label Unit 2", code="LU2")
        got = B.department_labels(
            [D(1, "CUTTING", b1.id), D(2, "cutting", b2.id), D(3, "SEWING", b1.id), D(4, "PACKING", b1.id),
             D(5, "PACKING", b1.id), D(6, "STORES", None), D(7, "STORES", b1.id)]
        )  # fmt: skip
        self.assertEqual(
            got,
            {1: "CUTTING (Label Unit 1)", 2: "cutting (Label Unit 2)", 3: "SEWING", 4: "PACKING (Label Unit 1) #4",
             5: "PACKING (Label Unit 1) #5", 6: "STORES (No branch)", 7: "STORES (Label Unit 1)"},
        )  # fmt: skip

    def test_strength_designation_view_says_which_unit_each_department_is(self):
        sew2 = Department.objects.create(name="SEWING", branch=self.b2)
        op2 = Designation.objects.create(title="Operator", department=sew2)
        self.make("E21", branch=self.b2, department=sew2, designation=op2)
        rows = self.get("strength-statement", groupBy="designation")["rows"]
        self.assertEqual(
            [(r["group"], r["department"], r["total"]) for r in rows],
            [("Manager", "CUTTING", 1), ("Operator", "SEWING (Unit 1)", 2), ("Operator", "SEWING (Unit 2)", 1),
             ("No designation", "Unassigned", 4)],
        )  # fmt: skip

    def test_strength_cards_count_real_groups_not_the_catch_all_bucket(self):
        cards_by = {
            g: cards(self.get("strength-statement", groupBy=g)) for g in ("department", "designation", "branch")
        }
        self.assertEqual(cards_by["department"]["Departments"], 3)  # CUTTING x2 + SEWING; 'Unassigned' is not one
        self.assertEqual(cards_by["designation"]["Designations"], 2)  # Manager + Operator; 'No designation' is not one
        self.assertEqual(cards_by["branch"]["Branches"], 2)  # Unit 1 + Unit 2; 'No branch' is not one
        # an empty group that really exists still counts
        self.assertEqual(
            cards(self.get("strength-statement", groupBy="branch", includeEmpty="true"))["Branches"],
            Branch.objects.count(),  # Head Office (seeded, nobody in it yet) + Unit 1 + Unit 2 + Old Unit
        )

    def test_manpower_movement_labels_only_the_clashing_departments_and_counts_unusable_join_dates(self):
        body = self.get("manpower-movement", year="2026", groupBy="department")
        self.assertEqual(
            [r["group"] for r in body["rows"]], ["CUTTING (Unit 1)", "CUTTING (Unit 2)", "SEWING", "Unassigned"]
        )
        notes = " ".join(body["notes"])
        self.assertIn("2 employee(s) have a missing or unreadable join date", notes)  # E7 'soon', E10 blank
        # once the namesake is gone, the plain name comes back
        Department.objects.filter(pk=self.d_cut2.pk).update(name="CUTTING 2")
        self.assertEqual(
            [r["group"] for r in self.get("manpower-movement", year="2026", groupBy="department")["rows"]],
            ["CUTTING", "CUTTING 2", "SEWING", "Unassigned"],
        )

    def test_manpower_movement_dates_a_rejoiner_by_the_real_exit(self):
        e = self.make("E31", status="inactive", branch=self.b1, department=self.d_cut1, join_date="2024-06-15")
        ResignationRequest.objects.create(
            employee=e, status="approved", last_working_date=date(2024, 3, 31), approved_at=utc(2024, 3, 1, 5, 0)
        )
        Employee.objects.filter(pk=e.pk).update(updated_at=utc(2026, 9, 10, 8, 0))
        rows = {r["group"]: r for r in self.get("manpower-movement", year="2026")["rows"]}
        self.assertEqual(rows["Sep 2026"]["leftTotal"], 2)  # E5 and E31
        self.assertEqual(
            {r["group"]: r for r in self.get("manpower-movement", year="2024")["rows"]}["Jun 2024"]["joinedTotal"], 1
        )
        self.assertEqual(
            {r["group"]: r for r in self.get("manpower-movement", year="2024")["rows"]}["Mar 2024"]["leftTotal"], 0
        )

    # -- service length of leavers ---------------------------------------------------------------------
    def test_workforce_profile_average_service_of_leavers_stops_at_the_exit(self):
        # E4 77 months (to 31-Aug-2026), E5 15 (to 10-Sep-2026), E9 84 (to 16-Jan-2026): (77 + 15 + 84) / 3 / 12
        body = self.get("workforce-profile", dimension="tenureBand", employeeStatus="inactive")
        self.assertEqual(cards(body)["Average service (years)"], 4.9)
        by = {r["bucket"]: r["total"] for r in body["rows"]}
        self.assertEqual((by["1-3 years"], by["5+ years"]), (1, 2))
        self.assertTrue(any("counted up to the exit date" in n for n in body["notes"]))

    def test_the_active_are_still_measured_to_today_and_no_leaver_note_is_shown_for_them(self):
        body = self.get("employee-master")
        self.assertEqual({r["employeeCode"]: r["tenure"] for r in body["rows"]}["E1"], "2y 8m")
        self.assertFalse(any("has left" in n for n in body["notes"]))
        body = self.get("employee-master", employeeStatus="all")
        self.assertTrue(any("has left" in n for n in body["notes"]))

    def test_leaver_tenure_agrees_between_the_master_register_and_the_exits_register(self):
        master = {
            r["employeeCode"]: r["tenure"] for r in self.get("employee-master", employeeStatus="inactive")["rows"]
        }
        exits = {
            r["employeeCode"]: r["tenureMonths"]
            for r in self.get("exits-register", dateFrom="2025-01-01", dateTo="2026-12-31")["rows"]
        }
        for code, months in exits.items():
            self.assertEqual(master[code], B.tenure_text(months), code)

    # -- absurd dates --------------------------------------------------------------------------------
    def test_an_absurd_join_date_is_reported_as_unreadable_everywhere(self):
        self.make("W9", branch=self.b1, department=self.d_cut1, join_date="9999-12-31")
        body = self.get("service-milestones", employeeStatus="all")
        self.assertNotIn("W9", codes(body))
        self.assertEqual(cards(body)["Join date missing (not listed)"], 3)  # E7 'soon', E10 blank, W9
        rows = {r["employeeCode"]: r for r in self.get("data-quality-audit", employeeStatus="all")["rows"]}
        self.assertIn("Join date '9999-12-31' cannot be read", rows["W9"]["issues"])
        rows = {r["employeeCode"]: r for r in self.get("employee-master", employeeStatus="all")["rows"]}
        self.assertEqual((rows["W9"]["joinDate"], rows["W9"]["tenure"]), (None, None))
        for fmt in ("xlsx", "pdf"):
            self.export("service-milestones", fmt, employeeStatus="all")

    # -- statutory: placeholders and the ESI wage ceiling ----------------------------------------------
    def test_every_placeholder_spelling_is_a_missing_number_and_never_a_duplicate(self):
        for code, ph in (("P1", "NIL"), ("P2", "not applicable"), ("P3", "0000000000"), ("P4", "n.a."), ("P5", "None")):
            self.make(code, branch=self.b1, pf_number=ph, esi_number=ph, uan_number=ph, bank_account=ph, bank_ifsc=ph)
        body = self.get("statutory-compliance")
        rows = {r["employeeCode"]: r for r in body["rows"]}
        for code in ("P1", "P2", "P3", "P4", "P5"):
            issues = rows[code]["issues"]
            for missing in ("PF number missing", "ESI number missing", "UAN missing", "Bank a/c and IFSC missing"):
                self.assertIn(missing, issues, (code, missing))
            self.assertNotIn("shared with", issues, code)
            self.assertNotIn("format", issues, code)
            self.assertIsNone(rows[code]["pfNumber"])  # a dash on screen, not the typed placeholder
            self.assertIsNone(rows[code]["bankIfsc"])
        self.assertTrue(any("placeholder" in n for n in body["notes"]))
        self.assertEqual(cards(body)["PF number missing"], 3 + 5)  # E2, E6, E8 + the five

    def test_placeholder_bank_and_id_proof_values_are_not_masked_as_if_real(self):
        self.make("P1", branch=self.b1, id_proof="NA", bank_account="NA")
        rows = {r["employeeCode"]: r for r in self.get("statutory-compliance")["rows"]}
        self.assertIsNone(rows["P1"]["idProof"])
        self.assertIsNone(rows["P1"]["bankAccount"])

    def _high_earner(self, code="HIGH1", branch=None, **kw):
        self._seq = getattr(self, "_seq", 0) + 1  # every identifier unique: a shared number would be a finding too
        base = dict(
            branch=branch or self.b1, department=self.d_cut1, salary_amount=Decimal("50000"),
            pf_number=f"TN/TPR/9{self._seq:03d}", uan_number=f"1000009{self._seq:05d}",
            bank_account=f"98765{self._seq:08d}", bank_ifsc="HDFC0001234",
        )  # fmt: skip
        e = self.make(code, **{**base, **kw})
        for cat in ("aadhaar_card", "pan_card", "bank_passbook"):
            self.doc(e, cat)
        return e

    def _issues(self, user=None, **params):
        body = self.get("statutory-compliance", user=user, includeCompliant="true", **params)
        return body, {r["employeeCode"]: r["issues"] or "" for r in body["rows"]}

    def test_esi_number_is_owed_up_to_the_ceiling_and_not_above_it(self):
        self._high_earner("AT", salary_amount=Decimal("21000"))  # payroll deducts ESI at or below the ceiling
        self._high_earner("JUST", salary_amount=Decimal("21000.01"))
        self._high_earner("HIGH1")
        self._high_earner("NOSAL", salary_amount=None)  # unknown salary is never assumed to be above the ceiling
        self._high_earner("PROD", employment_type="production", salary_amount=None, salary_per_shift=Decimal("5000"))
        body, issues = self._issues()
        self.assertIn("ESI number missing", issues["AT"])
        self.assertNotIn("ESI number missing", issues["JUST"])
        self.assertNotIn("ESI number missing", issues["HIGH1"])
        self.assertIn("ESI number missing", issues["NOSAL"])
        self.assertIn("ESI number missing", issues["PROD"])  # production pay comes from shifts: not exempted here
        # the counts follow: the exempt two are neither 'missing ESI' nor spoil 'fully compliant'
        self.assertEqual(cards(body)["ESI number missing"], 4 + 3)  # E2, E6, E7, E10 + AT, NOSAL, PROD
        self.assertEqual(issues["HIGH1"], "")
        note = next(n for n in body["notes"] if "ESI wage ceiling" in n)
        self.assertIn("2 staff member(s)", note)
        self.assertIn("Rs 21,000", note)

    def test_an_esi_number_that_is_present_is_still_format_checked_above_the_ceiling(self):
        self._high_earner("HIGH1", esi_number="9999")
        _body, issues = self._issues()
        self.assertEqual(issues["HIGH1"], "ESI number should be 10 or 17 digits")

    def test_the_esi_ceiling_follows_the_settings_of_the_employees_own_branch(self):
        from .models import BranchSettingsOverride, PayrollSettings

        self._high_earner("U1_HIGH", branch=self.b1)
        self._high_earner("U2_HIGH", branch=self.b2)
        BranchSettingsOverride.objects.create(branch=self.b1, overrides={"esi_applicable_below": "60000"})
        _body, issues = self._issues(employeeStatus="active")
        self.assertIn("ESI number missing", issues["U1_HIGH"])  # Unit 1's own ceiling is Rs 60,000
        self.assertNotIn("ESI number missing", issues["U2_HIGH"])  # Unit 2 keeps the company one (Rs 21,000)
        # ... whoever is asking: the Unit 1 user gets the same answer for their people
        _body, issues = self._issues(user=self.branch_user)
        self.assertIn("ESI number missing", issues["U1_HIGH"])
        # and the company-wide ceiling moves both
        PayrollSettings.objects.update(esi_applicable_below=Decimal("80000"))
        BranchSettingsOverride.objects.all().delete()
        _body, issues = self._issues()
        self.assertIn("ESI number missing", issues["U1_HIGH"])
        self.assertIn("ESI number missing", issues["U2_HIGH"])

    def test_the_esi_ceiling_lookup_does_not_grow_with_the_number_of_staff(self):
        def queries():
            self.client.get("/api/reports/run/statutory-compliance", {"includeCompliant": "true"}, **hdr(self.admin))
            with CaptureQueriesContext(connection) as q:
                self.get("statutory-compliance", includeCompliant="true")
            return len(q)

        before = queries()
        for i in range(15):
            self._high_earner(f"H{i}", branch=(self.b1, self.b2)[i % 2])
        self.assertLessEqual(queries() - before, 1)

    # -- exits: confidential reasons -------------------------------------------------------------------
    def test_manpower_and_exit_reports_never_expose_a_reason_to_a_role_without_resignations(self):
        user = HRUser.objects.create(
            username="em_emp_only_2",
            password_hash="x",
            role=Role.objects.create(name="em_emp_only_2", permissions={"reports": "view", "employees": "view"}),
        )
        for rid, params in (
            ("exits-register", {"dateFrom": "2020-01-01", "dateTo": "2026-12-31"}),
            ("manpower-movement", {"year": "2026"}),
            ("employee-master", {"employeeStatus": "all", "layout": "full"}),
        ):
            text = str(self.get(rid, user=user, **params))
            for secret in ("Better opportunity", "Relocation", "Health", "Studies", "HR Admin"):
                self.assertNotIn(secret, text, (rid, secret))

    # -- org structure ---------------------------------------------------------------------------------
    def test_org_structure_admin_filtered_by_branch_gets_the_same_safe_line_for_stray_departments(self):
        DepartmentHeadcount.objects.filter(department=self.d_cut2).update(required_count=5)
        self.make("E11", branch=self.b1, department=self.d_cut2)  # a Unit 1 person in Unit 2's department
        body = self.get("org-structure", branchIds=str(self.b1.id))
        rows = data_rows(body)
        self.assertEqual(
            [(r["branch"], r["department"]) for r in rows],
            [("Unit 1", "CUTTING"), ("Unit 1", "SEWING"), ("Unit 1", "Departments outside this listing (legacy data)")],
        )
        stray = rows[-1]
        self.assertEqual((stray["staff"], stray["production"]), (1, 0))
        self.assertEqual((stray["hods"], stray["required"], stray["vacancy"], stray["designations"]), (None,) * 4)
        # unfiltered, the same department is listed properly with everybody in it, and its vacancy is not partial
        full = next(
            r for r in data_rows(self.get("org-structure")) if (r["branch"], r["department"]) == ("Unit 2", "CUTTING")
        )
        self.assertEqual((full["staff"], full["required"], full["vacancy"]), (2, 5, 3))

    def test_org_structure_lists_a_branch_without_departments_for_its_own_user_and_for_admin(self):
        empty = Branch.objects.create(name="Empty Unit", code="EMP")
        user = HRUser.objects.create(username="em_empty", password_hash="x", role=self.role_full, branch=empty)
        body = self.get("org-structure", user=user)
        self.assertEqual([(r["branch"], r["branchCode"], r["department"]) for r in data_rows(body)],
                         [("Empty Unit", "EMP", "No departments set up")])  # fmt: skip
        self.assertEqual(cards(body)["Branches"], 1)
        self.assertEqual((body["totals"]["staff"], body["totals"]["production"]), (0, 0))
        admin = self.get("org-structure")
        self.assertEqual(cards(admin)["Branches"], Branch.objects.count())  # every real branch, none twice, no bucket
        self.assertEqual(
            [r["department"] for r in data_rows(admin) if r["branch"] == "Empty Unit"], ["No departments set up"]
        )
        # a department filter is a request for those departments, not for the branches that have none
        only = self.get("org-structure", departmentIds=str(self.d_sew1.id))
        self.assertNotIn("Empty Unit", {r["branch"] for r in data_rows(only)})
        self.assertEqual(cards(only)["Branches"], 1)
        # picking one branch lists that branch only
        picked = self.get("org-structure", branchIds=str(empty.id))
        self.assertEqual([r["branch"] for r in data_rows(picked)], ["Empty Unit"])

    def test_a_branch_user_never_sees_another_units_hods_plan_or_designations_anywhere_in_the_payload(self):
        boss = self.make("E20", first_name="Zed", last_name="Boss", branch=self.b2, department=self.d_cut2)
        ManagerDepartmentAssignment.objects.create(
            manager=DepartmentManager.objects.create(employee=boss), department=self.d_cut2
        )
        DepartmentHeadcount.objects.filter(department=self.d_cut2).update(required_count=9)
        for i in range(3):
            Designation.objects.create(title=f"Secret{i}", department=self.d_cut2)
        self.make("E11", branch=self.b1, department=self.d_cut2)
        body = self.get("org-structure", user=self.branch_user)
        self.assertNotIn("Zed", str(body))
        for r in data_rows(body):
            self.assertNotEqual(r["required"], 9)
            self.assertNotEqual(r["designations"], 3)
        self.assertEqual([r["hods"] for r in data_rows(body)], ["Jai Singh", "Arun Kumar", None])
