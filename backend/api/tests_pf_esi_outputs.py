"""PF / ESI / TDS in what payroll hands out: the salary slip (JSON and PDF), the legacy payroll report and the coverage
Excel / PDF export (plus its Report Center twin).

The rules, the arithmetic and the Co Emp / TDS categories are tested in tests_pf_esi.py (whose helpers and base class are
reused here: February 2026, 24 Mon-Sat working days, a 24,000 salary has the automatic split with a first portion of 12,000
of which Basic is 4,000; PF below 25,000 Basic at 12%, ESI below 30,000 Basic at 0.75%). The insights pages (details lines,
Salary Statement, rules) are tested in tests_payroll_insights.py, the Report Center statements in tests_reporting_*.py.

Run via: python manage.py test api.tests_pf_esi_outputs
"""

import io
import re
from datetime import date
from decimal import Decimal

import pdfplumber
from openpyxl import load_workbook

from . import tests_pf_esi as pe
from .clock import ist_today
from .company_documents_views import _tds_deduction_row, build_salary_slip_pdf
from .jwt_utils import sign_token
from .models import AuditLog, Employee, SalarySlip
from .payroll_views import _generate_production_payroll
from .pf_esi_coverage_export import SUMMARY_LINES
from .salary_slip_bulk_pdf import build_bulk_salary_slip_pdf
from .tests_payroll_engine import MONTH, YEAR, _configure, _day, _present_all, _rule

EXPORT = "/api/pf-esi/coverage/export"
SLIPS = "/api/salary-slips"


def _pdf_text(content: bytes) -> str:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return "\n".join((page.extract_text() or "") for page in pdf.pages)


# ─────────────────────────────────────────────────────────────────────────────
#  Salary slips
# ─────────────────────────────────────────────────────────────────────────────


class SalarySlipTests(pe.Base):
    """T = TDS at 10% (salary 30,000: TDS 3,000), C = Co Emp (PF 1,440 + ESI 90 of the 12,000 first portion), P = neither."""

    def setUp(self):
        super().setUp()
        self.t = pe._staff("T1", salary="30000.00", tds_enabled=True, tds_percent=Decimal("10"))
        self.c = pe._staff("C1", co_emp_enabled=True)
        self.p = pe._staff("P1")
        self.slips = {}
        for emp in (self.t, self.c, self.p):
            _present_all(emp)
            self.slips[emp.employee_code] = pe._generate(emp)

    # ── JSON ──

    def test_the_slip_list_carries_the_tds_fields_of_every_slip(self):
        rows = {r["employeeCode"]: r for r in self.get(SLIPS, month=MONTH, year=YEAR, employmentType="staff").json()}
        self.assertEqual((rows["T1"]["tdsDeduction"], rows["T1"]["tdsPercent"]), (3000.0, 10.0))
        self.assertEqual(
            (rows["T1"]["pfDeduction"], rows["T1"]["esiDeduction"], rows["T1"]["totalDeductions"]), (0.0, 0.0, 3000.0)
        )
        self.assertEqual((rows["C1"]["tdsDeduction"], rows["C1"]["tdsPercent"]), (0.0, None))
        self.assertEqual(
            (rows["C1"]["pfDeduction"], rows["C1"]["esiDeduction"], rows["C1"]["totalDeductions"]),
            (1440.0, 90.0, 1530.0),
        )
        self.assertEqual(
            (rows["P1"]["tdsDeduction"], rows["P1"]["tdsPercent"], rows["P1"]["totalDeductions"]), (0.0, None, 0.0)
        )

    def test_every_deduction_of_a_slip_adds_up_to_its_total(self):
        for row in self.get(SLIPS, month=MONTH, year=YEAR).json():
            parts = (
                row["pfDeduction"]
                + row["esiDeduction"]
                + row["tdsDeduction"]
                + row["advanceDeduction"]
                + row["otherDeductions"]
            )
            self.assertAlmostEqual(parts, row["totalDeductions"], places=2, msg=row["employeeCode"])

    def test_the_single_slip_and_the_employees_own_list_carry_them_too(self):
        one = self.get(f"{SLIPS}/{self.slips['T1'].id}").json()
        self.assertEqual((one["tdsDeduction"], one["tdsPercent"]), (3000.0, 10.0))
        token = sign_token({"role": "employee", "employeeId": self.t.id})
        mine = self.client.get("/api/my/salary-slips", HTTP_AUTHORIZATION=f"Bearer {token}")  # mobile app / web app
        self.assertEqual(mine.status_code, 200, mine.content[:300])
        body = mine.json()
        rows = body["items"] if isinstance(body, dict) else body
        self.assertEqual({(r["tdsDeduction"], r["tdsPercent"]) for r in rows}, {(3000.0, 10.0)})

    def test_the_legacy_payroll_report_lists_tds_and_totals_it(self):
        body = self.get("/api/reports/payroll", month=MONTH, year=YEAR).json()
        rows = {r["employeeCode"]: r for r in body["results"]}
        self.assertEqual((rows["T1"]["tdsDeduction"], rows["T1"]["tdsPercent"]), (3000.0, 10.0))
        self.assertEqual((rows["P1"]["tdsDeduction"], rows["P1"]["tdsPercent"]), (0.0, None))
        self.assertEqual(body["totals"]["tdsDeduction"], 3000.0)
        self.assertAlmostEqual(
            body["totals"]["pfDeduction"] + body["totals"]["esiDeduction"] + body["totals"]["tdsDeduction"],
            body["totals"]["totalDeductions"],
            places=2,
        )

    # ── PDF ──

    def test_the_slip_pdf_of_a_tds_employee_shows_the_tds_line(self):
        text = _pdf_text(build_salary_slip_pdf(self.slips["T1"]))
        self.assertRegex(text, r"TDS \(10%\)\s+(₹|Rs\. )3,000\.00")
        self.assertRegex(text, r"Total Deductions\s+(₹|Rs\. )3,000\.00")
        self.assertEqual(text.count("TDS"), 1)

    def test_the_slip_pdf_of_everybody_else_has_no_tds_line_at_all(self):
        co = _pdf_text(build_salary_slip_pdf(self.slips["C1"]))
        self.assertNotIn("TDS", co)
        self.assertRegex(co, r"P\.F\s+(₹|Rs\. )1,440\.00")
        self.assertRegex(co, r"E\.S\.I\s+(₹|Rs\. )90\.00")
        self.assertNotIn("TDS", _pdf_text(build_salary_slip_pdf(self.slips["P1"])))

    def test_a_percentage_with_decimals_reads_as_people_write_it(self):
        slip = SalarySlip(tds_percent=Decimal("12.50"), tds_deduction=Decimal("1250.00"))
        self.assertEqual(_tds_deduction_row(slip)[0], "TDS (12.5%)")
        self.assertEqual(_tds_deduction_row(SalarySlip(tds_percent=Decimal("10.00"), tds_deduction=0))[0], "TDS (10%)")
        self.assertIsNone(_tds_deduction_row(SalarySlip()))  # nobody else gets a line
        self.assertEqual(
            _tds_deduction_row(SalarySlip(tds_deduction=Decimal("5.00")))[0], "TDS"
        )  # amount, no percentage

    def test_the_combined_two_per_page_pdf_shows_tds_only_on_the_tds_slip(self):
        text = _pdf_text(build_bulk_salary_slip_pdf([self.slips["T1"], self.slips["C1"], self.slips["P1"]]))
        self.assertEqual(text.count("TDS"), 1)
        self.assertIn("TDS (10%)", text)

    def test_the_bulk_pdf_endpoint_is_a_pdf(self):
        r = self.client.get(f"{SLIPS}/bulk-pdf", {"month": MONTH, "year": YEAR}, **pe._auth(self.admin))
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_the_sample_preview_slip_still_renders(self):
        r = self.client.get("/api/document-settings/salary_slip/preview", **pe._auth(self.admin))
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertNotIn("TDS", _pdf_text(r.content))


# ─────────────────────────────────────────────────────────────────────────────
#  Coverage export
# ─────────────────────────────────────────────────────────────────────────────


class CoverageExportTests(pe.Base):
    """The same world as pe.CoverageTests: staff V1 (PF + ESI), V2 (ESI only), V3 (Co Emp, neither), V4 TDS, V5 plain, V6
    inactive; production V7 (Co Emp, payroll generated, no rule fits), V8 (Co Emp, no payroll yet), V9 TDS, V10 plain."""

    def setUp(self):
        super().setUp()
        _configure(
            prod_pf_rules=[_rule("pp", 0, 20000, 12, "total")],
            prod_esi_rules=[_rule("pe", 0, 12000, 1, "total")],
        )
        self.both = pe._staff("V1", co_emp_enabled=True)
        self.esi_only = pe._staff("V2", salary="162000.00", co_emp_enabled=True)
        self.neither = pe._staff("V3", salary="300000.00", co_emp_enabled=True)
        self.tds = pe._staff("V4", tds_enabled=True, tds_percent=Decimal("5"))
        self.plain = pe._staff("V5")
        self.inactive = pe._staff("V6", co_emp_enabled=True)
        Employee.objects.filter(pk=self.inactive.pk).update(status="inactive")
        self.p1 = pe._worker("V7", per_shift="1000.00", co_emp_enabled=True)
        for offset in range(7):
            _day(self.p1, date(2026, 2, 2 + offset), shifts="1.00")
        _generate_production_payroll(self.p1, date(2026, 2, 2), date(2026, 2, 8), pay_cycle="weekly")
        self.p2 = pe._worker("V8", co_emp_enabled=True)
        self.p3 = pe._worker("V9", tds_enabled=True, tds_percent=Decimal("2"))
        self.p4 = pe._worker("V10")

    def export(self, fmt="xlsx", user=None, **params):
        return self.client.get(EXPORT, {"format": fmt, **params}, **pe._auth(user or self.admin))

    @staticmethod
    def workbook(response):
        return load_workbook(io.BytesIO(response.content))

    @staticmethod
    def summary_of(ws) -> dict:
        """{summary key: number} read back from the SUMMARY block of one worksheet."""
        labels = {label.strip(): key for key, label, _indent in SUMMARY_LINES}
        out = {}
        for row in range(1, ws.max_row + 1):
            label = ws.cell(row=row, column=1).value
            if isinstance(label, str) and label.strip() in labels:
                out[labels[label.strip()]] = ws.cell(row=row, column=4).value
        return out

    @staticmethod
    def details_of(ws) -> list[dict]:
        """The employee rows of one worksheet as {header: value}."""
        header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Emp Code")
        headers = [ws.cell(row=header_row, column=c).value for c in range(1, 13)]
        rows = []
        row = header_row + 1
        while ws.cell(row=row, column=1).value and ws.cell(row=row, column=2).value:
            rows.append({h: ws.cell(row=row, column=c).value for c, h in enumerate(headers, start=1)})
            row += 1
        return rows

    # ── Excel ──

    def test_the_excel_numbers_equal_the_coverage_endpoint(self):
        body = self.get(pe.COVERAGE, detail=1).json()
        r = self.export("xlsx", type="all")
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = self.workbook(r)
        self.assertEqual(wb.sheetnames, ["Staff PF-ESI coverage", "Production PF-ESI coverage"])
        for sheet_name, key in (("Staff PF-ESI coverage", "staff"), ("Production PF-ESI coverage", "production")):
            ws = wb[sheet_name]
            self.assertEqual(self.summary_of(ws), body[key], key)  # every number of the summary block
            rows = self.details_of(ws)
            mine = [e for e in body["employees"] if e["employmentType"] == key]
            self.assertEqual(
                [r["Emp Code"] for r in rows], sorted((e["employeeCode"] for e in mine), key=lambda c: (len(c), c)), key
            )
            self.assertEqual(len(rows), body[key]["total"])

    def test_the_summary_block_obeys_the_arithmetic(self):
        ws = self.workbook(self.export("xlsx"))["Staff PF-ESI coverage"]
        s = self.summary_of(ws)
        self.assertEqual(s["total"], s["coEmp"] + s["tds"] + s["notCoEmp"])
        self.assertEqual(s["coEmp"], s["withAny"] + s["withNeither"] + s["notEvaluated"])
        self.assertEqual(
            (s["total"], s["coEmp"], s["withPf"], s["withEsi"], s["withBoth"], s["tds"], s["notCoEmp"]),
            (5, 3, 1, 2, 1, 1, 1),
        )

    def test_the_detail_rows_say_who_has_what(self):
        wb = self.workbook(self.export("xlsx"))
        staff = {r["Emp Code"]: r for r in self.details_of(wb["Staff PF-ESI coverage"])}
        production = {r["Emp Code"]: r for r in self.details_of(wb["Production PF-ESI coverage"])}
        self.assertEqual(set(staff), {"V1", "V2", "V3", "V4", "V5"})  # V6 is inactive
        self.assertEqual(
            (
                staff["V1"]["Category"],
                staff["V1"]["PF"],
                staff["V1"]["PF %"],
                staff["V1"]["ESI"],
                staff["V1"]["ESI %"],
                staff["V1"]["Status"],
            ),
            ("Co Emp", "Yes", 12, "Yes", 0.75, "PF and ESI"),
        )
        self.assertEqual((staff["V2"]["PF"], staff["V2"]["ESI"], staff["V2"]["Status"]), ("No", "Yes", "ESI only"))
        self.assertEqual(staff["V3"]["Status"], "Co Emp, no PF or ESI")
        self.assertEqual(
            (staff["V4"]["Category"], staff["V4"]["TDS %"], staff["V4"]["PF"], staff["V4"]["Status"]),
            ("TDS", 5, None, "TDS only"),
        )
        self.assertEqual((staff["V5"]["Category"], staff["V5"]["Status"]), ("Not Co Emp", "Not Co Emp"))
        self.assertEqual(
            (production["V7"]["Status"], production["V8"]["Status"], production["V8"]["PF"]),
            ("Co Emp, no PF or ESI", "Not evaluated", None),
        )
        self.assertEqual({r["Class"] for r in staff.values()}, {"Staff"})
        self.assertEqual(staff["V1"]["Employee"], "V1 Staff")

    def test_the_file_name_and_headers(self):
        r = self.export("xlsx", type="staff")
        self.assertEqual(
            r["Content-Disposition"],
            f'attachment; filename="pf_esi_coverage_staff_{ist_today().isoformat()}.xlsx"',
        )
        self.assertEqual(r["Cache-Control"], "no-store")
        self.assertRegex(
            self.export("pdf")["Content-Disposition"],
            r'^attachment; filename="pf_esi_coverage_all_\d{4}-\d{2}-\d{2}\.pdf"$',
        )

    def test_one_class_gives_one_sheet(self):
        self.assertEqual(self.workbook(self.export("xlsx", type="staff")).sheetnames, ["Staff PF-ESI coverage"])
        self.assertEqual(
            self.workbook(self.export("xlsx", type="production")).sheetnames, ["Production PF-ESI coverage"]
        )

    def test_the_default_type_is_all(self):
        r = self.client.get(EXPORT, {"format": "xlsx"}, **pe._auth(self.admin))
        self.assertEqual(len(self.workbook(r).sheetnames), 2)
        self.assertIn("pf_esi_coverage_all_", r["Content-Disposition"])

    def test_a_name_that_looks_like_a_formula_stays_text(self):
        Employee.objects.filter(pk=self.plain.pk).update(first_name='=HYPERLINK("http://x","y")', last_name="")
        ws = self.workbook(self.export("xlsx", type="staff"))["Staff PF-ESI coverage"]
        cell = next(c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value)
        self.assertEqual(cell.data_type, "s")

    # ── PDF ──

    def test_the_pdf_is_a_printable_document_with_both_sections(self):
        r = self.export("pdf", type="all")
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        self.assertTrue(r.content.startswith(b"%PDF"))
        text = _pdf_text(r.content)
        self.assertIn("PF / ESI COVERAGE", text)
        self.assertIn("STAFF EMPLOYEES", text)
        self.assertIn("PRODUCTION EMPLOYEES", text)
        for word in (
            "V1",
            "V8",
            "Not evaluated",
            "Co Emp, no PF or ESI",
            "Active employees",
            "TDS employees",
            "Page 1 of",
        ):
            self.assertIn(word, text, word)
        self.assertNotIn("V6", text)  # inactive

    def test_the_pdf_summary_numbers_are_the_same_numbers(self):
        body = self.get(pe.COVERAGE).json()
        text = _pdf_text(self.export("pdf", type="staff").content)
        self.assertIn("STAFF EMPLOYEES", text)
        self.assertNotIn("PRODUCTION EMPLOYEES", text)
        card = re.search(r"Active employees.*?(\d+)", text, re.S)
        self.assertEqual(int(card.group(1)), body["staff"]["total"])

    def test_a_long_list_runs_over_pages_with_a_repeated_header(self):
        for n in range(120):
            pe._staff(f"W{n:03d}", co_emp_enabled=True)
        r = self.export("pdf", type="staff")
        self.assertEqual(r.status_code, 200)
        with pdfplumber.open(io.BytesIO(r.content)) as pdf:
            self.assertGreater(len(pdf.pages), 2)
            self.assertTrue(all("Emp Code" in (page.extract_text() or "") for page in pdf.pages[:3]))

    # ── errors, access ──

    def test_a_wrong_or_missing_format_is_a_400_in_the_project_error_shape(self):
        for params in ({"format": "csv"}, {"format": ""}, {}, {"format": "XLS"}):
            r = self.client.get(EXPORT, params, **pe._auth(self.admin))
            self.assertEqual(r.status_code, 400, params)
            self.assertEqual(r["Content-Type"], "application/json")
            self.assertEqual(r.json(), {"error": "format must be xlsx or pdf"})

    def test_the_format_parameter_is_not_swallowed_by_rest_framework(self):
        # DRF reads ?format= as a renderer suffix and answers 404 for unknown ones: this endpoint turns that off
        self.assertEqual(self.export("xlsx").status_code, 200)
        self.assertEqual(self.export("PDF").status_code, 200)  # the file type is not case sensitive

    def test_a_bad_type_or_branch_is_a_400(self):
        self.assertEqual(
            self.export("xlsx", type="everyone").json(), {"error": "type must be staff, production or all"}
        )
        self.assertEqual(self.export("xlsx", branchId="abc").json(), {"error": "branchId must be a number"})
        self.assertEqual(self.export("pdf", type="everyone").status_code, 400)

    def test_nobody_without_a_token_gets_a_file(self):
        self.assertEqual(self.client.get(EXPORT, {"format": "xlsx"}).status_code, 401)
        token = sign_token({"role": "employee", "employeeId": self.both.id})
        self.assertEqual(
            self.client.get(EXPORT, {"format": "xlsx"}, HTTP_AUTHORIZATION=f"Bearer {token}").status_code, 403
        )

    def test_the_permission_is_the_one_of_the_coverage_counts(self):
        nobody = pe._hr_user("exp_nobody", permissions={"employees": "hidden"})
        self.assertEqual(self.export("xlsx", nobody).status_code, 403)
        self.assertEqual(self.export("pdf", nobody).status_code, 403)
        payroll_only = pe._hr_user("exp_payroll", permissions={"payroll": "edit"})
        self.assertEqual(self.export("xlsx", payroll_only).status_code, 403)
        viewer = pe._hr_user("exp_viewer", permissions={"employees": "view"})
        self.assertEqual(self.export("xlsx", viewer).status_code, 200)
        self.assertEqual(self.export("pdf", viewer).status_code, 200)

    def test_a_branch_login_gets_only_its_own_branch(self):
        Employee.objects.filter(pk__in=[self.both.pk, self.plain.pk, self.p2.pk]).update(branch=self.north)
        Employee.objects.filter(pk__in=[self.esi_only.pk, self.tds.pk, self.p1.pk]).update(branch=self.south)
        user = pe._hr_user("exp_north", permissions={"employees": "view"}, branch=self.north)
        wb = self.workbook(self.export("xlsx", user))
        staff = self.details_of(wb["Staff PF-ESI coverage"])
        production = self.details_of(wb["Production PF-ESI coverage"])
        self.assertEqual({r["Emp Code"] for r in staff}, {"V1", "V5"})
        self.assertEqual({r["Emp Code"] for r in production}, {"V8"})
        self.assertEqual({r["Branch"] for r in staff + production}, {"North"})
        self.assertEqual(self.summary_of(wb["Staff PF-ESI coverage"])["total"], 2)
        self.assertEqual(self.summary_of(wb["Staff PF-ESI coverage"]), self.get(pe.COVERAGE, user).json()["staff"])
        # asking for another branch's id changes nothing: the scope is the login's, so nothing of South's is in the file
        other = self.workbook(self.export("xlsx", user, branchId=self.south.id))
        self.assertEqual(self.details_of(other["Staff PF-ESI coverage"]), [])
        self.assertNotIn("V2", _pdf_text(self.export("pdf", user, branchId=self.south.id).content))
        self.assertNotIn("V2", _pdf_text(self.export("pdf", user).content))

    def test_the_branch_filter_of_an_unscoped_login(self):
        Employee.objects.filter(pk__in=[self.both.pk, self.esi_only.pk]).update(branch=self.north)
        ws = self.workbook(self.export("xlsx", type="staff", branchId=self.north.id))["Staff PF-ESI coverage"]
        self.assertEqual({r["Emp Code"] for r in self.details_of(ws)}, {"V1", "V2"})
        self.assertEqual(self.summary_of(ws)["total"], 2)
        title = [ws.cell(row=r, column=1).value for r in range(1, 6)]
        self.assertTrue(any(isinstance(t, str) and "Branch: North" in t for t in title), title)

    def test_a_branch_override_shows_in_the_file(self):
        Employee.objects.filter(pk=self.both.pk).update(branch=self.north)
        self.put_rules(
            {
                "branchId": self.north.id,
                "expectedVersion": self.version(branchId=self.north.id),
                "staff": {"pfRules": [], "esiRules": []},
            }
        )
        rows = {r["Emp Code"]: r for r in self.details_of(self.workbook(self.export("xlsx"))["Staff PF-ESI coverage"])}
        self.assertEqual(rows["V1"]["Status"], "Co Emp, no PF or ESI")  # North has no rules at all
        self.assertEqual(rows["V2"]["Status"], "ESI only")  # unchanged: follows the rules for all units

    def test_an_export_is_audit_logged_and_changes_nothing(self):
        before = (Employee.objects.count(), SalarySlip.objects.count())
        self.export("xlsx", type="staff")
        self.assertEqual((Employee.objects.count(), SalarySlip.objects.count()), before)
        entry = AuditLog.objects.filter(action="export", module="employees").latest("id")
        self.assertIn("PF / ESI coverage - XLSX", entry.record_description)

    def test_an_empty_company_still_gives_a_file(self):
        Employee.objects.all().delete()
        wb = self.workbook(self.export("xlsx"))
        self.assertEqual(self.summary_of(wb["Staff PF-ESI coverage"])["total"], 0)
        self.assertEqual(self.details_of(wb["Staff PF-ESI coverage"]), [])
        self.assertIn("No active employees", _pdf_text(self.export("pdf").content))

    # ── the Report Center twin ──

    def test_the_report_center_has_the_same_report(self):
        body = self.get(pe.COVERAGE).json()
        r = self.client.get("/api/reports/run/pf-esi-coverage", **pe._auth(self.admin))
        self.assertEqual(r.status_code, 200, r.content[:300])
        out = r.json()
        cards = {c["label"]: c["value"] for c in out["summary"]}
        self.assertEqual(cards["Employees"], body["staff"]["total"] + body["production"]["total"])
        self.assertEqual(cards["Co Emp (PF and ESI apply)"], body["staff"]["coEmp"] + body["production"]["coEmp"])
        self.assertEqual(cards["With PF"], body["staff"]["withPf"] + body["production"]["withPf"])
        self.assertEqual(cards["With ESI"], body["staff"]["withEsi"] + body["production"]["withEsi"])
        self.assertEqual(cards["TDS employees"], 2)
        rows = {row["employeeCode"]: row for row in out["rows"]}
        self.assertEqual(
            (rows["V1"]["category"], rows["V1"]["pf"], rows["V1"]["pfRate"], rows["V1"]["esiRate"]),
            ("Co Emp", "Yes", 12, 0.75),
        )
        self.assertEqual(
            (rows["V4"]["category"], rows["V4"]["tdsPercent"], rows["V4"]["status"]), ("TDS", 5, "TDS only")
        )
        for fmt in ("xlsx", "pdf"):
            self.assertEqual(
                self.client.get(f"/api/reports/export/pf-esi-coverage?fmt={fmt}", **pe._auth(self.admin)).status_code,
                200,
            )

    def test_the_report_is_branch_isolated_and_needs_the_employees_module(self):
        Employee.objects.filter(pk__in=[self.both.pk]).update(branch=self.north)
        north = pe._hr_user("rep_north", permissions={"employees": "view", "reports": "view"}, branch=self.north)
        r = self.client.get("/api/reports/run/pf-esi-coverage", **pe._auth(north))
        self.assertEqual(r.status_code, 200)
        self.assertEqual([row["employeeCode"] for row in r.json()["rows"]], ["V1"])
        payroll_only = pe._hr_user("rep_payroll", permissions={"payroll": "view", "reports": "view"})
        self.assertEqual(self.client.get("/api/reports/run/pf-esi-coverage", **pe._auth(payroll_only)).status_code, 403)
