"""
Report Center framework: catalog / run / export endpoints, filter validation, branch isolation,
the Excel and PDF writers, and the audit trail.

Run via: python manage.py test api.tests_reporting_framework -v 2
"""

import io
from datetime import date, timedelta

from django.test import TestCase
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .permission_registry import all_module_keys
from .models import AuditLog, Branch, Department, Employee, HRUser, Role
from .reporting import registry
from .reporting import filters as F
from .reporting.formatting import indian_number, minutes_text, parse_date
from .reporting.types import (
    BADGE,
    CURRENCY,
    DATE,
    DURATION,
    INTEGER,
    TEXT,
    ColumnSpec,
    ReportResult,
    ReportSpec,
)


def _hr_headers(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def _run_probe(ctx):
    rows = []
    for e in ctx.employees():
        rows.append(
            {
                "code": e.employee_code,
                "name": f"{e.first_name} {e.last_name}",
                "dept": e.department.name if e.department_id else None,
                "amount": 1000.5 if e.employee_code != "RPT_C" else 250,
                "count": 2,
                "day": "2026-09-05",
                "mins": 125,
                "state": "active",
                "secret": "must never leave the server",
            }
        )
    return ReportResult(
        rows=rows, summary=[{"label": "People", "value": len(rows), "format": "integer"}], notes=["a note"]
    )


def _probe_spec(report_id="zz_probe", **kw):
    args = dict(
        id=report_id,
        title="Probe / Report",
        description="test only",
        category="employees",
        filters=(F.period(), F.date_range(max_days=31), *F.scope(status="all")),
        columns=(
            ColumnSpec("code", "Code", TEXT, 1),
            ColumnSpec("name", "Name", TEXT, 2),
            ColumnSpec("dept", "Department", TEXT, 2),
            ColumnSpec("amount", "Amount", CURRENCY, 1, total="sum"),
            ColumnSpec("count", "Count", INTEGER, 1, total="sum"),
            ColumnSpec("day", "Day", DATE, 1),
            ColumnSpec("mins", "Late", DURATION, 1, total="sum"),
            ColumnSpec("state", "State", BADGE, 1),
        ),
        run=_run_probe,
    )
    args.update(kw)
    return ReportSpec(**args)


class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.cutting = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.sewing = Department.objects.create(name="SEWING", branch=cls.b2)
        mk = lambda code, dept, branch, **kw: Employee.objects.create(  # noqa: E731
            employee_code=code, first_name=code, last_name="T", department=dept, branch=branch, **kw
        )
        cls.a = mk("RPT_A", cls.cutting, cls.b1, employment_type="staff")
        cls.b = mk("RPT_B", cls.cutting, cls.b1, employment_type="production")
        cls.c = mk("RPT_C", cls.sewing, cls.b2, employment_type="staff")
        cls.gone = mk("RPT_D", cls.cutting, cls.b1, status="inactive")
        cls.admin = HRUser.objects.create(username="rpt_admin", password_hash="x", is_super_admin=True)
        role = Role.objects.create(name="rpt_viewer", permissions={"reports": "view"})
        cls.branch_user = HRUser.objects.create(username="rpt_b1", password_hash="x", role=role, branch=cls.b1)
        cls.no_reports = HRUser.objects.create(
            username="rpt_none",
            password_hash="x",
            role=Role.objects.create(name="rpt_none", permissions={"reports": "hidden"}),
        )

    def setUp(self):
        self._registered = []
        self.reg(_probe_spec())

    def tearDown(self):
        for rid in self._registered:
            registry._REGISTRY.pop(rid, None)

    def reg(self, spec):
        registry.register(spec)
        self._registered.append(spec.id)
        return spec

    def get(self, path, user=None, **params):
        return self.client.get(path, params, **_hr_headers(user or self.admin))


class CatalogTests(_Base):
    def test_requires_hr_login(self):
        self.assertEqual(self.client.get("/api/reports/catalog").status_code, 401)
        emp = {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'employee', 'employeeId': self.a.id})}"}
        self.assertEqual(self.client.get("/api/reports/catalog", **emp).status_code, 403)

    def test_reports_module_hidden_is_denied_by_middleware(self):
        r = self.get("/api/reports/catalog", user=self.no_reports)
        self.assertEqual(r.status_code, 403)

    def test_catalog_lists_probe_with_resolved_defaults(self):
        r = self.get("/api/reports/catalog")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        probe = next(x for x in body["reports"] if x["id"] == "zz_probe")
        self.assertEqual(probe["category"], "employees")
        keys = [f["key"] for f in probe["filters"]]
        self.assertIn("period", keys)
        self.assertIn("dateRange", keys)
        period = next(f for f in probe["filters"] if f["key"] == "period")
        self.assertRegex(period["default"], r"^\d{4}-\d{2}$")
        drange = next(f for f in probe["filters"] if f["key"] == "dateRange")
        self.assertEqual(set(drange["default"]), {"dateFrom", "dateTo"})
        self.assertEqual(drange["maxDays"], 31)
        self.assertIn({"id": self.cutting.id, "name": "CUTTING"}, body["options"]["departments"])
        self.assertTrue(any(c["id"] == "employees" and c["count"] >= 1 for c in body["categories"]))

    def test_branch_scoped_user_gets_no_branch_filter_and_only_own_lists(self):
        body = self.get("/api/reports/catalog", user=self.branch_user).json()
        self.assertTrue(body["branchScoped"])
        probe = next(x for x in body["reports"] if x["id"] == "zz_probe")
        self.assertNotIn("branch", [f["key"] for f in probe["filters"]])
        self.assertEqual([b["id"] for b in body["options"]["branches"]], [self.b1.id])
        self.assertEqual([d["name"] for d in body["options"]["departments"]], ["CUTTING"])

    def test_super_admin_only_report_hidden_from_others(self):
        self.reg(_probe_spec("zz_admin", super_admin_only=True))
        ids = [x["id"] for x in self.get("/api/reports/catalog", user=self.branch_user).json()["reports"]]
        self.assertNotIn("zz_admin", ids)
        ids = [x["id"] for x in self.get("/api/reports/catalog").json()["reports"]]
        self.assertIn("zz_admin", ids)
        self.assertEqual(self.get("/api/reports/run/zz_admin", user=self.branch_user).status_code, 404)
        self.assertEqual(self.get("/api/reports/export/zz_admin", user=self.branch_user, fmt="xlsx").status_code, 404)

    def test_unknown_report_is_404(self):
        self.assertEqual(self.get("/api/reports/run/nope").status_code, 404)

    def test_family_variants_count_once_per_category(self):
        base = _probe_spec("zz_fam_a", family="zz_fam", variant="Records")
        self.reg(base)
        self.reg(_probe_spec("zz_fam_b", family="zz_fam", variant="Counts"))
        before = self.get("/api/reports/catalog").json()
        n_emp = next(c["count"] for c in before["categories"] if c["id"] == "employees")
        # probe + one family (two variants) = 2 entries, not 3
        self.assertEqual(
            n_emp, len({(s.category, s.family or s.id) for s in registry.all_specs() if s.category == "employees"})
        )


class RunTests(_Base):
    def test_rows_totals_summary_and_unknown_keys_dropped(self):
        r = self.get("/api/reports/run/zz_probe", period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-10")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body["rowCount"], 4)
        self.assertNotIn("secret", body["rows"][0])
        self.assertEqual(body["totals"]["count"], 8)
        self.assertEqual(body["totals"]["amount"], 3 * 1000.5 + 250)
        self.assertEqual(body["totals"]["mins"], 500)
        self.assertEqual(body["summary"], [{"label": "People", "value": 4, "format": "integer"}])
        self.assertEqual(body["notes"], ["a note"])
        self.assertFalse(body["truncated"])
        self.assertTrue(any(f["label"] == "Month" and f["value"] == "September 2026" for f in body["filters"]))
        self.assertEqual([c["key"] for c in body["columns"]][:2], ["code", "name"])

    def test_scope_filters(self):
        def codes(**p):
            body = self.get(
                "/api/reports/run/zz_probe", period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02", **p
            ).json()
            return [x["code"] for x in body["rows"]]

        self.assertEqual(codes(departmentIds=str(self.sewing.id)), ["RPT_C"])
        self.assertEqual(codes(employmentType="production"), ["RPT_B"])
        self.assertEqual(codes(employeeIds=f"{self.a.id},{self.c.id}"), ["RPT_A", "RPT_C"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), ["RPT_C"])
        self.assertEqual(sorted(codes()), ["RPT_A", "RPT_B", "RPT_C", "RPT_D"])  # status filter defaults to "all"

    def test_branch_isolation_cannot_be_widened_by_params(self):
        body = self.get(
            "/api/reports/run/zz_probe",
            user=self.branch_user,
            period="2026-09",
            dateFrom="2026-09-01",
            dateTo="2026-09-02",
            branchIds=str(self.b2.id),
        ).json()
        self.assertEqual(body["rows"], [])
        body = self.get(
            "/api/reports/run/zz_probe",
            user=self.branch_user,
            period="2026-09",
            dateFrom="2026-09-01",
            dateTo="2026-09-02",
            employeeIds=str(self.c.id),
        ).json()
        self.assertEqual(body["rows"], [])
        body = self.get(
            "/api/reports/run/zz_probe",
            user=self.branch_user,
            period="2026-09",
            dateFrom="2026-09-01",
            dateTo="2026-09-02",
        ).json()
        self.assertEqual(sorted(x["code"] for x in body["rows"]), ["RPT_A", "RPT_B", "RPT_D"])

    def test_employee_status_filter_active_and_inactive(self):
        spec = self.reg(_probe_spec("zz_status", filters=(F.period(), F.date_range(), F.employee_status("active"))))
        base = dict(period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02")
        act = self.get(f"/api/reports/run/{spec.id}", **base).json()
        self.assertEqual(sorted(x["code"] for x in act["rows"]), ["RPT_A", "RPT_B", "RPT_C"])
        ina = self.get(f"/api/reports/run/{spec.id}", employeeStatus="inactive", **base).json()
        self.assertEqual([x["code"] for x in ina["rows"]], ["RPT_D"])

    def test_validation_errors_are_400_with_field(self):
        cases = [
            ({"period": "2026-13"}, "period"),
            ({"period": "September"}, "period"),
            ({"dateFrom": "2026-09-10", "dateTo": "2026-09-01"}, "dateFrom"),
            ({"dateFrom": "2026-01-01", "dateTo": "2026-12-31"}, "dateTo"),  # wider than the report's 31-day cap
            ({"dateFrom": "nonsense", "dateTo": "2026-09-01"}, "dateFrom"),
            ({"departmentIds": "1,x"}, "departmentIds"),
            ({"employmentType": "contractor"}, "employmentType"),
        ]
        for extra, field in cases:
            params = {"period": "2026-09", "dateFrom": "2026-09-01", "dateTo": "2026-09-05", **extra}
            r = self.get("/api/reports/run/zz_probe", **params)
            self.assertEqual(r.status_code, 400, extra)
            self.assertEqual(r.json()["error"], "invalid_filter")
            self.assertEqual(r.json()["field"], field, extra)

    def test_defaults_apply_when_params_missing(self):
        r = self.get("/api/reports/run/zz_probe")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(any(f["label"] == "Month" for f in r.json()["filters"]))

    def test_select_boolean_number_text_filters(self):
        spec = self.reg(
            _probe_spec(
                "zz_kinds",
                filters=(
                    F.select("mode", "Mode", [("a", "Alpha"), ("b", "Beta")], default="a"),
                    F.select("multi", "Multi", [("x", "X"), ("y", "Y")], multi=True),
                    F.boolean("flag", "Flag"),
                    F.number("min", "Minimum", default=3, min=1, max=9),
                    F.text("q", "Search"),
                ),
                run=lambda ctx: ReportResult(rows=[{"code": str(sorted(ctx.params.items(), key=lambda kv: kv[0]))}]),
                columns=(ColumnSpec("code", "Params", TEXT),),
            )
        )
        r = self.get(f"/api/reports/run/{spec.id}", mode="b", multi="x,y", flag="true", min="5", q="  hello ")
        self.assertEqual(r.status_code, 200)
        text = r.json()["rows"][0]["code"]
        for frag in ("'mode', 'b'", "'multi', ['x', 'y']", "'flag', True", "'min', 5", "'q', 'hello'"):
            self.assertIn(frag, text)
        self.assertEqual(self.get(f"/api/reports/run/{spec.id}", mode="zzz").status_code, 400)
        self.assertEqual(self.get(f"/api/reports/run/{spec.id}", multi="x,q").status_code, 400)
        self.assertEqual(self.get(f"/api/reports/run/{spec.id}", min="99").status_code, 400)
        self.assertEqual(self.get(f"/api/reports/run/{spec.id}", min="abc").status_code, 400)
        d = self.get(f"/api/reports/run/{spec.id}").json()["rows"][0]["code"]
        self.assertIn("'mode', 'a'", d)
        self.assertIn("'min', 3", d)

    def test_row_limit_truncates_and_flags(self):
        spec = self.reg(
            _probe_spec(
                "zz_big",
                screen_limit=5,
                run=lambda ctx: ReportResult(
                    rows=[{"code": f"R{i}", "amount": 1} for i in range(ctx.row_limit)] + [{"code": "extra"}] * 3
                ),
            )
        )
        body = self.get(
            f"/api/reports/run/{spec.id}", period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02"
        ).json()
        self.assertEqual(body["rowCount"], 5)
        self.assertTrue(body["truncated"])
        self.assertEqual(body["limit"], 5)
        self.assertIsNone(body["totals"])  # a sum over a cut-off list would under-state the real total

    def test_export_over_the_row_limit_is_refused_not_truncated(self):
        spec = self.reg(
            _probe_spec(
                "zz_toobig",
                pdf_max_rows=3,
                run=lambda ctx: ReportResult(rows=[{"code": f"R{i}", "amount": 1} for i in range(5)]),
            )
        )
        q = dict(period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02")
        r = self.get(f"/api/reports/export/{spec.id}", fmt="pdf", **q)
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.json()["error"], "too_many_rows")
        self.assertEqual(r.json()["limit"], 3)
        self.assertIn("Narrow the filters", r.json()["message"])
        self.assertEqual(
            self.get(f"/api/reports/export/{spec.id}", fmt="xlsx", **q).status_code, 200
        )  # xlsx cap is higher

    def test_failing_report_returns_500_json_not_traceback(self):
        def boom(ctx):
            raise RuntimeError("kaput")

        spec = self.reg(_probe_spec("zz_boom", run=boom))
        with self.assertLogs("api.reporting.views", level="ERROR"):
            r = self.get(f"/api/reports/run/{spec.id}", period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02")
        self.assertEqual(r.status_code, 500)
        self.assertEqual(r.json()["error"], "report_failed")
        self.assertNotIn("kaput", r.content.decode())

    def test_subtotal_rows_do_not_count_against_the_row_limit(self):
        def rows(ctx):
            out = []
            for i in range(4):
                out.append({"code": f"R{i}", "amount": 1})
                out.append({"code": "Subtotal", "amount": 1, "_kind": "subtotal"})
            return ReportResult(rows=out)

        spec = self.reg(_probe_spec("zz_sublimit", screen_limit=4, run=rows))
        q = dict(period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02")
        body = self.get(f"/api/reports/run/{spec.id}", **q).json()
        self.assertFalse(body["truncated"])  # 4 data rows fit a limit of 4 even though 8 lines are returned
        self.assertEqual(body["rowCount"], 8)
        over = self.reg(_probe_spec("zz_sublimit2", screen_limit=3, run=rows))
        body = self.get(f"/api/reports/run/{over.id}", **q).json()
        self.assertTrue(body["truncated"])
        self.assertEqual(sum(1 for r in body["rows"] if not r.get("_kind")), 3)

    def test_structural_rows_are_excluded_from_totals(self):
        spec = self.reg(
            _probe_spec(
                "zz_sub",
                run=lambda ctx: ReportResult(
                    rows=[
                        {"code": "a", "amount": 10, "count": 1},
                        {"code": "b", "amount": 20, "count": 2},
                        {"code": "Subtotal", "amount": 30, "count": 3, "_kind": "subtotal"},
                    ]
                ),
            )
        )
        body = self.get(
            f"/api/reports/run/{spec.id}", period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02"
        ).json()
        self.assertEqual(body["totals"]["amount"], 30.0)
        self.assertEqual(body["rows"][2]["_kind"], "subtotal")

    def test_dynamic_columns_and_explicit_totals_override(self):
        spec = self.reg(
            _probe_spec(
                "zz_dyn",
                columns=(),
                run=lambda ctx: ReportResult(
                    rows=[{"a": 1, "b": "x"}],
                    columns=[ColumnSpec("a", "A", INTEGER), ColumnSpec("b", "B", TEXT)],
                    totals={"a": 99},
                ),
            )
        )
        body = self.get(
            f"/api/reports/run/{spec.id}", period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02"
        ).json()
        self.assertEqual([c["key"] for c in body["columns"]], ["a", "b"])
        self.assertEqual(body["totals"], {"a": 99})
        cat = next(x for x in self.get("/api/reports/catalog").json()["reports"] if x["id"] == spec.id)
        self.assertTrue(cat["dynamicColumns"])


class ExportTests(_Base):
    P = dict(period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-10")

    def test_xlsx_structure_types_and_totals(self):
        r = self.get("/api/reports/export/zz_probe", fmt="xlsx", **self.P)
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertRegex(r["Content-Disposition"], r'attachment; filename="probe_report_\d{8}_\d{4}\.xlsx"')
        self.assertEqual(r["Cache-Control"], "no-store")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        self.assertEqual(ws.title, "Probe   Report")  # "/" is not allowed in a sheet name
        self.assertEqual(ws["A3"].value, "PROBE / REPORT")
        self.assertIn("Month: September 2026", ws["A4"].value)
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:4], ["Code", "Name", "Department", "Amount"])
        first = [c.value for c in ws[8]]
        self.assertEqual(first[0], "RPT_A")
        self.assertEqual(first[3], 1000.5)
        self.assertEqual(ws["F8"].value.date() if hasattr(ws["F8"].value, "date") else ws["F8"].value, date(2026, 9, 5))
        self.assertEqual(
            ws["G8"].value, timedelta(minutes=125)
        )  # duration is stored as a fraction of a day, format [h]:mm
        self.assertEqual(ws["G12"].value, timedelta(minutes=500))
        self.assertIn("₹", ws["D8"].number_format)
        total_row = [c.value for c in ws[12]]
        self.assertEqual(total_row[0], "TOTAL")
        self.assertEqual(total_row[3], 3 * 1000.5 + 250)
        self.assertEqual(ws.freeze_panes, "A8")
        self.assertTrue(ws.auto_filter.ref.startswith("A7:"))

    def test_xlsx_never_turns_text_into_formulas(self):
        spec = self.reg(
            _probe_spec(
                "zz_inject",
                run=lambda ctx: ReportResult(
                    rows=[
                        {"code": '=HYPERLINK("http://evil","x")', "name": "+91 98765 43210", "dept": "@SUM(A1)"},
                        {"code": "-cmd|calc", "name": "\t=1+1"},
                    ]
                ),
            )
        )
        r = self.get(f"/api/reports/export/{spec.id}", fmt="xlsx", **self.P)
        ws = load_workbook(io.BytesIO(r.content)).active
        for cell in (ws["A8"], ws["B8"], ws["C8"], ws["A9"], ws["B9"]):
            self.assertEqual(cell.data_type, "s", cell.coordinate)
        self.assertTrue(ws["A8"].value.startswith("="))

    def test_pdf_is_a_pdf_and_paginates(self):
        r = self.get("/api/reports/export/zz_probe", fmt="pdf", **self.P)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("application/pdf", r["Content-Type"])
        many = self.reg(
            _probe_spec(
                "zz_many",
                run=lambda ctx: ReportResult(
                    rows=[{"code": f"E{i}", "name": f"Name {i}", "amount": i * 10.5, "count": i} for i in range(400)]
                ),
            )
        )
        big = self.get(f"/api/reports/export/{many.id}", fmt="pdf", **self.P)
        self.assertEqual(big.status_code, 200)
        self.assertGreater(big.content.count(b"/Type /Page\n") + big.content.count(b"/Type /Page "), 3)

    def test_pdf_and_xlsx_handle_empty_and_wide_reports(self):
        empty = self.reg(_probe_spec("zz_empty", run=lambda ctx: ReportResult(rows=[])))
        for fmt in ("pdf", "xlsx"):
            r = self.get(f"/api/reports/export/{empty.id}", fmt=fmt, **self.P)
            self.assertEqual(r.status_code, 200, fmt)
        cols = tuple(ColumnSpec(f"c{i}", f"D{i + 1}", INTEGER, 0.5) for i in range(32))
        wide = self.reg(
            _probe_spec(
                "zz_wide", columns=cols, run=lambda ctx: ReportResult(rows=[{f"c{i}": i for i in range(32)}] * 3)
            )
        )
        for fmt in ("pdf", "xlsx"):
            self.assertEqual(self.get(f"/api/reports/export/{wide.id}", fmt=fmt, **self.P).status_code, 200, fmt)

    def test_custom_builders_replace_generic_export(self):
        spec = self.reg(
            _probe_spec(
                "zz_custom", pdf_builder=lambda ctx, out: b"%PDF-custom", xlsx_builder=lambda ctx, out: b"PK-custom"
            )
        )
        self.assertEqual(self.get(f"/api/reports/export/{spec.id}", fmt="pdf", **self.P).content, b"%PDF-custom")
        self.assertEqual(self.get(f"/api/reports/export/{spec.id}", fmt="xlsx", **self.P).content, b"PK-custom")

    def test_bad_format_and_bad_filters_are_400(self):
        self.assertEqual(self.get("/api/reports/export/zz_probe", fmt="docx", **self.P).status_code, 400)
        self.assertEqual(self.get("/api/reports/export/zz_probe", fmt="pdf", period="bad").status_code, 400)

    def test_export_is_audited(self):
        before = AuditLog.objects.filter(action="export", module="reports").count()
        self.get("/api/reports/export/zz_probe", fmt="xlsx", **self.P)
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), before + 1)
        log = AuditLog.objects.filter(action="export", module="reports").latest("id")
        self.assertIn("Probe / Report", log.record_description)
        self.assertIn("XLSX", log.record_description)

    def test_export_respects_branch_isolation(self):
        r = self.get("/api/reports/export/zz_probe", user=self.branch_user, fmt="xlsx", **self.P)
        ws = load_workbook(io.BytesIO(r.content)).active
        codes = [row[0].value for row in ws.iter_rows(min_row=8, max_row=10)]
        self.assertEqual(sorted(codes), ["RPT_A", "RPT_B", "RPT_D"])
        self.assertNotIn("RPT_C", [c.value for row in ws.iter_rows() for c in row])


class ModuleGateTests(_Base):
    """A "reports" grant alone must not open data owned by a module the role cannot see."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        mk = lambda name, perms: HRUser.objects.create(  # noqa: E731
            username=name, password_hash="x", role=Role.objects.create(name=name, permissions=perms)
        )
        cls.plain = mk("gate_plain", {"reports": "view"})
        cls.payroll = mk("gate_payroll", {"reports": "view", "payroll": "view"})
        cls.slip_edit = mk("gate_slip", {"reports": "view", "salary_slip": "edit"})
        cls.payroll_hidden = mk("gate_hidden", {"reports": "view", "payroll": "hidden", "salary_slip": "hidden"})
        cls.cascade = mk("gate_cascade", {"reports": "view", "employees": "view"})

    def setUp(self):
        super().setUp()
        self.reg(_probe_spec("zz_pay", modules=("payroll", "salary_slip")))
        self.reg(_probe_spec("zz_child", modules=("employees.departments",)))

    P = dict(period="2026-09", dateFrom="2026-09-01", dateTo="2026-09-02")

    def ids(self, user):
        return {x["id"] for x in self.get("/api/reports/catalog", user=user).json()["reports"]}

    def test_catalog_only_lists_reports_the_role_can_open(self):
        self.assertNotIn("zz_pay", self.ids(self.plain))
        self.assertIn("zz_pay", self.ids(self.payroll))
        self.assertIn("zz_pay", self.ids(self.slip_edit))  # any one of the listed modules is enough
        self.assertNotIn("zz_pay", self.ids(self.payroll_hidden))
        self.assertIn("zz_pay", self.ids(self.admin))
        self.assertIn("zz_probe", self.ids(self.plain))  # reports without modules stay open to "reports" holders

    def test_run_and_export_are_403_report_forbidden(self):
        for url, extra in (("run", {}), ("export", {"fmt": "xlsx"})):
            r = self.get(f"/api/reports/{url}/zz_pay", user=self.plain, **self.P, **extra)
            self.assertEqual(r.status_code, 403, url)
            self.assertEqual(r.json()["error"], "report_forbidden", url)  # not the middleware's permission_denied
        self.assertEqual(self.get("/api/reports/run/zz_pay", user=self.payroll, **self.P).status_code, 200)
        self.assertEqual(
            self.get("/api/reports/export/zz_pay", user=self.slip_edit, fmt="pdf", **self.P).status_code, 200
        )

    def test_child_module_inherits_parent_grant(self):
        self.assertIn("zz_child", self.ids(self.cascade))
        self.assertNotIn("zz_child", self.ids(self.plain))

    def test_unknown_module_key_fails_closed_for_everyone_but_super_admin(self):
        self.reg(_probe_spec("zz_typo", modules=("payrol",)))
        self.assertNotIn("zz_typo", self.ids(self.payroll))
        self.assertIn("zz_typo", self.ids(self.admin))

    def test_denied_export_is_not_audited(self):
        before = AuditLog.objects.filter(action="export", module="reports").count()
        self.get("/api/reports/export/zz_pay", user=self.plain, fmt="xlsx", **self.P)
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), before)


class OptionsTests(_Base):
    def test_employee_search_and_resolution(self):
        r = self.get("/api/reports/options/employees", q="rpt_a")
        self.assertEqual([x["value"] for x in r.json()], [self.a.id])
        self.assertEqual(r.json()[0]["label"], "RPT_A · RPT_A T")
        ids = f"{self.a.id},{self.c.id}"
        self.assertEqual(
            sorted(x["value"] for x in self.get("/api/reports/options/employees", ids=ids).json()),
            sorted([self.a.id, self.c.id]),
        )
        only_active = self.get("/api/reports/options/employees", status="active").json()
        self.assertNotIn(self.gone.id, [x["value"] for x in only_active])
        by_dept = self.get("/api/reports/options/employees", departmentIds=str(self.sewing.id)).json()
        self.assertEqual([x["value"] for x in by_dept], [self.c.id])

    def test_employee_options_are_branch_scoped(self):
        r = self.get("/api/reports/options/employees", user=self.branch_user, ids=str(self.c.id))
        self.assertEqual(r.json(), [])
        r = self.get("/api/reports/options/employees", user=self.branch_user, q="RPT")
        self.assertNotIn(self.c.id, [x["value"] for x in r.json()])

    def test_unknown_source_404(self):
        self.assertEqual(self.get("/api/reports/options/departments").status_code, 404)

    def test_role_that_can_open_no_employee_report_cannot_browse_employees(self):
        registry.ensure_loaded()
        saved = dict(registry._REGISTRY)
        registry._REGISTRY.clear()  # only gated reports exist for this check
        try:
            self.reg(_probe_spec("zz_gated", modules=("payroll",)))
            plain = HRUser.objects.create(
                username="opt_plain",
                password_hash="x",
                role=Role.objects.create(name="opt_plain", permissions={"reports": "view"}),
            )
            r = self.get("/api/reports/options/employees", user=plain, q="RPT")
            self.assertEqual(r.status_code, 403)
            self.assertEqual(r.json()["error"], "report_forbidden")
            self.assertEqual(
                self.get("/api/reports/options/employees", q="RPT").status_code, 200
            )  # super admin still can
        finally:
            registry._REGISTRY.clear()
            registry._REGISTRY.update(saved)


class HelperTests(TestCase):
    def test_indian_number_grouping(self):
        self.assertEqual(indian_number(0), "0.00")
        self.assertEqual(indian_number(999.5), "999.50")
        self.assertEqual(indian_number(1234567.891), "12,34,567.89")
        self.assertEqual(indian_number(12345678), "1,23,45,678.00")
        self.assertEqual(indian_number(-100000), "-1,00,000.00")
        self.assertEqual(indian_number(1500, 0), "1,500")

    def test_minutes_text(self):
        self.assertEqual(minutes_text(125), "2h 05m")
        self.assertEqual(minutes_text(45), "45m")
        self.assertEqual(minutes_text(0), "0m")
        self.assertEqual(minutes_text(None), "")
        self.assertEqual(minutes_text(-70), "-1h 10m")

    def test_parse_date_is_tolerant(self):
        self.assertEqual(parse_date("2026-09-05"), date(2026, 9, 5))
        self.assertEqual(parse_date("2026-09-05T10:00:00"), date(2026, 9, 5))
        self.assertEqual(parse_date("05-09-2026"), date(2026, 9, 5))
        self.assertIsNone(parse_date("soon"))
        self.assertIsNone(parse_date(""))
        self.assertIsNone(parse_date(None))

    def test_spec_validation(self):
        with self.assertRaises(ValueError):
            ColumnSpec("x", "X", "bogus")
        with self.assertRaises(ValueError):
            _probe_spec(category="nope")
        with self.assertRaises(ValueError):
            _probe_spec(family="only")
        with self.assertRaises(ValueError):
            _probe_spec(columns=(ColumnSpec("a", "A"), ColumnSpec("a", "B")))
        with self.assertRaises(ValueError):
            registry.register(_probe_spec("zz_dup_x"))
            registry.register(_probe_spec("zz_dup_x"))

    def tearDown(self):
        registry._REGISTRY.pop("zz_dup_x", None)


class RegisteredReportsContractTests(TestCase):
    """Every report that ships must satisfy the contract the UI and exporters rely on."""

    def test_every_definition_module_imports_cleanly(self):
        registry.all_specs()
        self.assertEqual(registry.LOAD_ERRORS, {})

    def test_every_registered_report_is_well_formed(self):
        specs = registry.all_specs()
        ids = [s.id for s in specs]
        self.assertEqual(len(ids), len(set(ids)))
        families: dict[str, set[str]] = {}
        for s in specs:
            self.assertRegex(s.id, r"^[a-z0-9-]+$", s.id)
            self.assertTrue(s.title and s.description, s.id)
            if s.family:
                families.setdefault(s.family, set()).add(s.variant)
            self.assertTrue(s.columns or s.id, s.id)
            for m in s.modules:
                self.assertIn(m, all_module_keys(), f"{s.id}: unknown permission module {m!r}")
        for fam, variants in families.items():
            self.assertGreaterEqual(len(variants), 2, f"family {fam} has a single variant")
