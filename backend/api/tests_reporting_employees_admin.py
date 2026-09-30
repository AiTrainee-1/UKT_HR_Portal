"""
Report Center - group G11 (employees_admin): documents, HOD, mobile-app access, recruitment and the
admin-only audit / user-access reports.

Run via:  python manage.py test api.tests_reporting_employees_admin -v 2
(the shared runner script gives it a private throw-away database)

Everything is deterministic: "today" is pinned to 2026-09-20 and the wall clock used for session / dormancy
maths to 2026-09-20 06:30 UTC (12:00 IST); every fixture date is fixed.
"""

import io
import re
import uuid
from datetime import date, datetime, timedelta, timezone as dt_timezone
from pathlib import Path
from unittest import mock

from django.apps import apps
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import connection
from django.test import RequestFactory, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from . import hod_scope
from .clock import FACTORY_TZ
from .jwt_utils import sign_token
from .models import (
    Applicant,
    AuditLog,
    Branch,
    Department,
    DepartmentHeadcount,
    DepartmentManager,
    Designation,
    Employee,
    EmployeeDocument,
    HiringRuleSet,
    HrLoginAttempt,
    HRUser,
    Job,
    LoginSession,
    ManagerDepartmentAssignment,
    ManagerEmployeeAssignment,
    PayrollSettings,
    PushToken,
    Role,
    ScreeningCandidate,
)
from .permission_registry import MODULE_TREE, all_module_keys
from .reporting import registry
from .reporting.definitions import employees_admin_util as U
from .reporting.runner import run_report as run_spec

UTC = dt_timezone.utc
TODAY = date(2026, 9, 20)
NOW_UTC = datetime(2026, 9, 20, 6, 30, tzinfo=UTC)  # 12:00 IST
WINDOW = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
SECRET = "SECRETHASH$2b$12$abcdef"

NON_ADMIN_IDS = (
    "document-compliance",
    "document-upload-log",
    "hod-directory",
    "hod-mapping",
    "hod-conflicts",
    "manpower-requirement",
    "job-openings",
    "applicant-register",
    "screening-pipeline",
    "interview-schedule",
    "recruitment-funnel",
    "hiring-criteria",
    "mobile-app-access",
)
ADMIN_IDS = (
    "audit-log",
    "employee-change-log",
    "audit-summary",
    "hr-users",
    "role-access-matrix",
    "login-sessions",
    "login-attempts",
)
ALL_IDS = NON_ADMIN_IDS + ADMIN_IDS


def ist(y, m, d, h=0, mi=0, s=0):
    return datetime(y, m, d, h, mi, s, tzinfo=FACTORY_TZ)


def hdr(user):
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def make_user(name, permissions, branch=None, **kw):
    role = Role.objects.create(name=f"role_{name}", permissions=permissions)
    return HRUser.objects.create(username=name, password_hash=SECRET, role=role, branch=branch, **kw)


def make_emp(code, first, dept=None, branch=None, etype="staff", status="active", last="T", **kw):
    return Employee.objects.create(
        employee_code=code,
        first_name=first,
        last_name=last,
        department=dept,
        branch=branch,
        employment_type=etype,
        status=status,
        **kw,
    )


def add_doc(emp, category, when, by=None, name=None):
    doc = EmployeeDocument.objects.create(
        employee=emp,
        category=category,
        file=f"employee_documents/{uuid.uuid4().hex}.pdf",
        original_filename=name or f"{category}.pdf",
        uploaded_by=by,
    )
    EmployeeDocument.objects.filter(pk=doc.pk).update(uploaded_at=when)
    return doc


def add_log(when, user, action, module, desc=None, branch=None, ip=None, record_id=None):
    log = AuditLog.objects.create(
        user_type="hr",
        user_name=user,
        action=action,
        module=module,
        record_id=record_id,
        record_description=desc,
        ip_address=ip,
        branch=branch,
    )
    AuditLog.objects.filter(pk=log.pk).update(created_at=when)
    return log


def codes(body, key="employeeCode"):
    return [r[key] for r in body["rows"]]


def by_key(body, value, key="employeeCode"):
    return next(r for r in body["rows"] if r[key] == value)


def cards(body):
    return {s["label"]: s["value"] for s in body["summary"]}


def data_rows(body):
    return [r for r in body["rows"] if not r.get("_kind")]


class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.admin = HRUser.objects.create(username="ea_admin", password_hash=SECRET + "-admin", is_super_admin=True)
        cls.reports_only = make_user("ea_reports_only", {"reports": "view"})

    def setUp(self):
        for patcher in (
            mock.patch("api.reporting.filters.ist_today", return_value=TODAY),
            mock.patch.object(U, "utc_now", return_value=NOW_UTC),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)

    # -- HTTP helpers ------------------------------------------------------
    def raw(self, rid, user=None, fmt=None, **params):
        if fmt:
            return self.client.get(f"/api/reports/export/{rid}", {"fmt": fmt, **params}, **hdr(user or self.admin))
        return self.client.get(f"/api/reports/run/{rid}", params, **hdr(user or self.admin))

    def run_report(self, rid, user=None, **params):
        r = self.raw(rid, user, **params)
        self.assertEqual(r.status_code, 200, r.content[:400])
        return r.json()

    def sheet(self, rid, user=None, **params):
        r = self.raw(rid, user, fmt="xlsx", **params)
        self.assertEqual(r.status_code, 200, r.content[:400])
        self.assertTrue(r.content.startswith(b"PK"))
        return load_workbook(io.BytesIO(r.content)).active

    def assert_totals_match(self, body, *keys):
        rows = data_rows(body)
        for k in keys:
            self.assertEqual(body["totals"][k], sum(r[k] or 0 for r in rows), k)

    def scoped_run(self, rid, branch, **params):
        """Run a report through the framework as if the caller were branch-scoped to ``branch``.

        The admin-only reports cannot be reached by a branch-scoped user over HTTP (that is a 404, tested
        elsewhere), so their branch isolation - every query goes through ``ctx.emp_q()`` - is proven here by
        pinning the scope the middleware would have set."""
        request = RequestFactory().get("/api/reports/run/" + rid)
        request.jwt_user = {"role": "hr", "hrUserId": self.admin.id}
        with mock.patch("api.reporting.filters.get_branch_scope", return_value=branch.id):
            return run_spec(request, registry.get_spec(rid), params)


# ═══════════════════════════════════════════════════════════════════════════
#  Documents
# ═══════════════════════════════════════════════════════════════════════════


class DocumentReportTests(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cut = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.sew = Department.objects.create(name="SEWING", branch=cls.b1)
        cls.pack = Department.objects.create(name="PACKING", branch=cls.b2)
        cls.operator = Designation.objects.create(title="Operator", department=cls.cut)
        cls.e9 = make_emp("9", "Asha", cls.cut, cls.b1, designation=cls.operator)
        cls.e10 = make_emp("10", "Bala", cls.cut, cls.b1)
        cls.e20 = make_emp("20", "Chitra", cls.sew, cls.b1, etype="production")
        cls.e101 = make_emp("101", "Dev", cls.sew, cls.b1, etype="production", status="inactive")
        cls.e102 = make_emp("102", "Esha", cls.sew, cls.b1)
        cls.e201 = make_emp("201", "Farid", cls.pack, cls.b2)
        cls.no_dept = make_emp("300", "Gita", None, cls.b1)

        aug = ist(2026, 8, 15, 10)
        # 9: complete (six required) - some inside September, one on the IST 1-Sep boundary, one just after it
        add_doc(cls.e9, "pan_card", ist(2026, 9, 2, 10, 0), "HR One")
        add_doc(cls.e9, "aadhaar_card", ist(2026, 9, 2, 10, 5), "HR One")
        add_doc(cls.e9, "educational_certificate", aug, "HR One")
        add_doc(cls.e9, "voter_id_or_birth_certificate", aug, "HR One")
        add_doc(cls.e9, "bank_passbook", ist(2026, 9, 1, 0, 30), "HR One")  # 31-Aug 19:00 UTC
        add_doc(cls.e9, "staff_letter", ist(2026, 10, 1, 0, 30), "HR One")  # 30-Sep 19:00 UTC
        # 10: everything except the voter id; two education files count once
        add_doc(cls.e10, "pan_card", ist(2026, 8, 20, 9), "HR Two")
        add_doc(cls.e10, "aadhaar_card", ist(2026, 8, 20, 9, 1), "HR Two")
        add_doc(cls.e10, "educational_certificate", ist(2026, 9, 5, 9, 0), "HR Two")
        add_doc(cls.e10, "educational_certificate", ist(2026, 9, 5, 9, 1), "HR Two")
        add_doc(cls.e10, "bank_passbook", ist(2026, 9, 11, 1, 30), "HR Two")  # 10-Sep 20:00 UTC
        add_doc(cls.e10, "staff_letter", ist(2026, 8, 20, 9, 2), "HR Two")
        # 20: production employee, complete (the staff letter is irrelevant to him)
        for cat in (
            "pan_card",
            "aadhaar_card",
            "educational_certificate",
            "voter_id_or_birth_certificate",
            "bank_passbook",
            "production_employee_documents",
            "staff_letter",
        ):
            add_doc(cls.e20, cat, aug)
        add_doc(cls.e101, "pan_card", aug)
        add_doc(cls.e102, "offer_letter", ist(2026, 9, 30, 23, 30), None)  # not a required category
        add_doc(cls.e201, "pan_card", ist(2026, 9, 6, 8), "HR One")
        add_doc(cls.e201, "aadhaar_card", ist(2026, 8, 30, 8), "HR One")

        cls.doc_user = make_user("ea_docs_b1", {"reports": "view", "recruitment.documents": "view"}, cls.b1)

    # -- compliance --------------------------------------------------------
    def test_default_is_pending_active_employees_in_natural_code_order(self):
        body = self.run_report("document-compliance")
        self.assertEqual(codes(body), ["10", "102", "201", "300"])
        self.assertEqual(
            [c["key"] for c in body["columns"]][:4], ["employeeCode", "employeeName", "department", "employmentType"]
        )

    def test_golden_rows_summary_totals_and_notes(self):
        body = self.run_report("document-compliance", state="all")
        self.assertEqual(codes(body), ["9", "10", "20", "102", "201", "300"])  # natural order, 'inactive' 101 excluded
        r = by_key(body, "10")
        self.assertEqual((r["uploadedCount"], r["missingCount"], r["completionPct"]), (5, 1, 83.3))
        self.assertEqual(r["missing"], "Voter ID or Birth Certificate")
        self.assertEqual(r["voter_id_or_birth_certificate"], "Missing")
        self.assertEqual(r["pan_card"], "Uploaded")
        self.assertEqual(r["typeSpecific"], "Uploaded")
        self.assertEqual(r["lastUploadedOn"], "2026-09-11")  # 10-Sep 20:00 UTC is already the 11th in IST
        self.assertEqual(r["employmentType"], "Staff")
        r = by_key(body, "102")
        self.assertEqual((r["uploadedCount"], r["missingCount"], r["completionPct"]), (0, 6, 0.0))
        self.assertEqual(r["lastUploadedOn"], "2026-09-30")  # an offer letter counts as an upload but is not required
        self.assertEqual(r["typeSpecific"], "Missing")
        r = by_key(body, "201")
        self.assertEqual((r["uploadedCount"], r["missingCount"], r["completionPct"]), (2, 4, 33.3))
        self.assertEqual(
            r["missing"], "Educational Certificates, Voter ID or Birth Certificate, Bank Passbook, Staff Letter"
        )
        r = by_key(body, "20")
        self.assertEqual(r["missingCount"], 0)
        self.assertIsNone(r["missing"])
        self.assertEqual(r["typeSpecific"], "Uploaded")
        self.assertEqual(r["employmentType"], "Production")
        self.assertEqual(by_key(body, "300")["department"], "Unassigned")
        self.assertEqual(cards(body), {"Employees checked": 6, "Complete": 2, "Pending": 4, "Completion": 33.3})
        self.assert_totals_match(body, "uploadedCount", "missingCount")
        self.assertEqual(body["totals"]["missingCount"], 0 + 1 + 0 + 6 + 4 + 6)
        self.assertEqual(body["totals"]["completionPct"], round(sum(r["completionPct"] for r in body["rows"]) / 6, 2))
        self.assertTrue(any("Voter ID or Birth Certificate 4" in n for n in body["notes"]))

    def test_production_needs_production_documents_not_the_staff_letter(self):
        body = self.run_report("document-compliance", state="all", employeeStatus="all")
        r = by_key(body, "101")  # production, inactive, only a PAN card
        self.assertEqual(r["missingCount"], 5)
        self.assertNotIn("Staff Letter", r["missing"])
        self.assertIn("Production Employee Documents", r["missing"])

    def test_state_and_missing_category_filters(self):
        self.assertEqual(codes(self.run_report("document-compliance", state="complete")), ["9", "20"])
        body = self.run_report("document-compliance", state="all", missingCategory="voter_id_or_birth_certificate")
        self.assertEqual(codes(body), ["10", "102", "201", "300"])
        body = self.run_report("document-compliance", state="all", employeeStatus="all", missingCategory="staff_letter")
        self.assertEqual(codes(body), ["102", "201", "300"])  # production employees are never asked for it
        body = self.run_report(
            "document-compliance", state="all", employeeStatus="all", missingCategory="production_employee_documents"
        )
        self.assertEqual(codes(body), ["101"])
        # the headline cards ignore the row filters
        self.assertEqual(cards(body)["Employees checked"], 7)

    def test_scope_filters_narrow(self):
        base = dict(state="all", employeeStatus="all")
        self.assertEqual(
            codes(self.run_report("document-compliance", departmentIds=str(self.sew.id), **base)), ["20", "101", "102"]
        )
        self.assertEqual(
            codes(self.run_report("document-compliance", employmentType="production", **base)), ["20", "101"]
        )
        self.assertEqual(
            codes(self.run_report("document-compliance", employeeIds=f"{self.e9.id},{self.e201.id}", **base)),
            ["9", "201"],
        )
        self.assertEqual(codes(self.run_report("document-compliance", branchIds=str(self.b2.id), **base)), ["201"])
        self.assertEqual(
            codes(self.run_report("document-compliance", designationIds=str(self.operator.id), **base)), ["9"]
        )
        self.assertEqual(
            codes(self.run_report("document-compliance", employeeStatus="inactive", **{"state": "all"})), ["101"]
        )

    def test_branch_isolation(self):
        body = self.run_report("document-compliance", self.doc_user, state="all", employeeStatus="all")
        self.assertNotIn("201", codes(body))
        self.assertEqual(cards(body)["Employees checked"], 6)
        self.assertEqual(
            self.run_report("document-compliance", self.doc_user, branchIds=str(self.b2.id), state="all")["rows"], []
        )
        self.assertEqual(
            self.run_report("document-compliance", self.doc_user, employeeIds=str(self.e201.id), state="all")["rows"],
            [],
        )

    def test_xlsx_keeps_identifiers_as_text_and_dates_and_counts_typed(self):
        make_emp("007", "Zed", self.cut, self.b1)  # a code with leading zeros must never become the number 7
        ws = self.sheet("document-compliance", state="all")
        self.assertEqual(ws.cell(row=8, column=1).value, "007")
        rows = {r[0].value: [c.value for c in r] for r in ws.iter_rows(min_row=8)}
        nine = rows["9"]
        self.assertEqual((nine[4], nine[9], nine[10], nine[11], nine[14]), ("Uploaded", "Uploaded", 6, 0, 100.0))
        last = nine[13]
        self.assertEqual(last.date() if isinstance(last, datetime) else last, date(2026, 10, 1))
        self.assertEqual(
            rows["102"][12],
            "PAN Card, Aadhaar Card, Educational Certificates, Voter ID or Birth Certificate, Bank Passbook, Staff Letter",
        )
        total = rows["TOTAL"]
        self.assertEqual((total[10], total[11]), (19, 17 + 6))  # the new employee adds six missing documents

    # -- upload log --------------------------------------------------------
    def test_upload_log_golden_rows_ist_bounds_and_order(self):
        body = self.run_report("document-upload-log", **WINDOW)
        got = [(r["employeeCode"], r["category"]) for r in body["rows"]]
        self.assertEqual(
            got,
            [
                ("102", "Offer Letter"),
                ("10", "Bank Passbook"),
                ("201", "PAN Card"),
                ("10", "Educational Certificates"),
                ("10", "Educational Certificates"),
                ("9", "Aadhaar Card"),
                ("9", "PAN Card"),
                ("9", "Bank Passbook"),
            ],
        )
        # IST 1-Sep 00:30 (= 31-Aug UTC) is inside September; IST 1-Oct 00:30 (= 30-Sep UTC) is not
        self.assertIn("2026-09-01 00:30", codes(body, "uploadedAt"))
        self.assertNotIn("2026-10-01 00:30", codes(body, "uploadedAt"))
        self.assertEqual(body["rows"][0]["uploadedAt"], "2026-09-30 23:30")
        self.assertIsNone(body["rows"][0]["uploadedBy"])
        self.assertEqual(body["rows"][1]["uploadedBy"], "HR Two")
        self.assertEqual(cards(body), {"Files uploaded": 8, "Employees covered": 4, "Top uploader": "HR One (4)"})
        self.assertTrue(all("originalFilename" in r for r in body["rows"]))

    def test_upload_log_filters(self):
        self.assertEqual(codes(self.run_report("document-upload-log", category="pan_card", **WINDOW)), ["201", "9"])
        self.assertEqual(len(self.run_report("document-upload-log", uploadedBy="hr two", **WINDOW)["rows"]), 3)
        self.assertEqual(codes(self.run_report("document-upload-log", branchIds=str(self.b2.id), **WINDOW)), ["201"])
        self.assertEqual(
            set(codes(self.run_report("document-upload-log", departmentIds=str(self.cut.id), **WINDOW))), {"9", "10"}
        )
        self.assertEqual(
            codes(self.run_report("document-upload-log", employeeIds=str(self.e102.id), **WINDOW)), ["102"]
        )
        # August (IST): 2 (Asha) + 3 (Bala) + 7 (Chitra) + 1 (Dev) + 1 (Farid)
        self.assertEqual(
            self.run_report("document-upload-log", dateFrom="2026-08-01", dateTo="2026-08-31")["rowCount"], 14
        )

    def test_upload_log_branch_isolation_and_no_file_links(self):
        body = self.run_report("document-upload-log", self.doc_user, **WINDOW)
        self.assertNotIn("201", codes(body))
        self.assertEqual(cards(body)["Files uploaded"], 7)
        self.assertEqual(
            self.run_report("document-upload-log", self.doc_user, branchIds=str(self.b2.id), **WINDOW)["rows"], []
        )
        text = self.raw("document-upload-log", **WINDOW).content.decode()
        self.assertNotIn("employee_documents/", text)

    def test_upload_log_employee_type_and_designation_filters(self):
        aug = dict(dateFrom="2026-08-01", dateTo="2026-08-31")
        self.assertEqual(
            set(codes(self.run_report("document-upload-log", employmentType="production", **aug))), {"20", "101"}
        )
        self.assertEqual(self.run_report("document-upload-log", employmentType="production", **aug)["rowCount"], 8)
        self.assertEqual(
            set(codes(self.run_report("document-upload-log", designationIds=str(self.operator.id), **WINDOW))), {"9"}
        )

    def test_employee_without_a_branch_is_visible_to_unscoped_users_only(self):
        ghost = make_emp("400", "Nobranch", self.cut, None)
        add_doc(ghost, "pan_card", ist(2026, 9, 7, 10), "HR One")
        self.assertIn("400", codes(self.run_report("document-compliance", state="all")))
        self.assertIn("400", codes(self.run_report("document-upload-log", **WINDOW)))
        self.assertNotIn("400", codes(self.run_report("document-compliance", self.doc_user, state="all")))
        self.assertNotIn("400", codes(self.run_report("document-upload-log", self.doc_user, **WINDOW)))


# ═══════════════════════════════════════════════════════════════════════════
#  Mobile app access
# ═══════════════════════════════════════════════════════════════════════════


class MobileAccessTests(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cut = Department.objects.create(name="CUTTING", branch=cls.b1)
        pack = Department.objects.create(name="PACKING", branch=cls.b2)
        cls.cut, cls.pack = cut, pack
        cls.e9 = make_emp(
            "9",
            "Asha",
            cut,
            cls.b1,
            password_hash=SECRET + "-9",
            phone="9000000001",
            last_mobile_login_at=ist(2026, 9, 18, 9, 0),
            password_updated_at=ist(2026, 9, 1, 9, 0),
            location_tracking_enabled=True,
        )
        cls.e10 = make_emp("10", "Bala", cut, cls.b1, password_hash=SECRET + "-10")
        cls.e20 = make_emp("20", "Chitra", cut, cls.b1, etype="production")
        cls.e101 = make_emp(
            "101",
            "Dev",
            cut,
            cls.b1,
            status="inactive",
            password_hash=SECRET + "-101",
            last_mobile_login_at=ist(2026, 9, 10, 9, 0),
        )
        cls.e102 = make_emp("102", "Esha", cut, cls.b1, password_hash="")  # empty string = no password
        cls.e201 = make_emp(
            "201",
            "Farid",
            pack,
            cls.b2,
            password_hash=SECRET + "-201",
            last_mobile_login_at=ist(2026, 9, 19, 23, 45),
            co_emp_enabled=True,
        )
        PushToken.objects.create(employee=cls.e9, token="tok-a")
        PushToken.objects.create(employee=cls.e9, token="tok-b")
        PushToken.objects.create(employee=cls.e201, token="tok-c")
        cls.user = make_user("ea_mobile_b1", {"reports": "view", "mobile_app_login": "view"}, cls.b1)

    def test_golden_rows_and_summary(self):
        body = self.run_report("mobile-app-access")
        self.assertEqual(codes(body), ["9", "10", "20", "101", "102", "201"])
        r = by_key(body, "9")
        self.assertEqual((r["appPassword"], r["devices"], r["liveTracking"], r["coEmp"]), ("Set", 2, "On", "Off"))
        self.assertEqual(r["lastAppLogin"], "2026-09-18 09:00")
        self.assertEqual(r["passwordUpdatedAt"], "2026-09-01 09:00")
        self.assertEqual(r["phone"], "9000000001")
        self.assertEqual((by_key(body, "10")["appPassword"], by_key(body, "10")["lastAppLogin"]), ("Set", None))
        self.assertEqual(by_key(body, "20")["appPassword"], "Not set")
        self.assertEqual(by_key(body, "102")["appPassword"], "Not set")
        self.assertEqual(by_key(body, "101")["status"], "Inactive")
        self.assertEqual(by_key(body, "201")["lastAppLogin"], "2026-09-19 23:45")
        self.assertEqual(by_key(body, "201")["coEmp"], "On")
        self.assertEqual(
            cards(body),
            {
                "Employees": 6,
                "Has app access": 4,
                "No access": 2,
                "Signed in": 3,
                "Active without access": 2,
            },
        )
        self.assertEqual(body["totals"]["devices"], 3)
        self.assert_totals_match(body, "devices")

    def test_access_filter_and_stable_cards(self):
        body = self.run_report("mobile-app-access", access="has_access")
        self.assertEqual(codes(body), ["9", "10", "101", "201"])
        self.assertEqual(cards(body)["Employees"], 6)  # cards ignore the access tab, like the app's page
        self.assertEqual(codes(self.run_report("mobile-app-access", access="no_access")), ["20", "102"])
        self.assertEqual(codes(self.run_report("mobile-app-access", access="signed_in")), ["9", "101", "201"])
        self.assertEqual(codes(self.run_report("mobile-app-access", access="never_signed_in")), ["10", "20", "102"])

    def test_scope_filters(self):
        body = self.run_report("mobile-app-access", employmentType="staff")
        self.assertEqual(
            cards(body),
            {
                "Employees": 5,
                "Has app access": 4,
                "No access": 1,
                "Signed in": 3,
                "Active without access": 1,
            },
        )
        self.assertEqual(codes(self.run_report("mobile-app-access", employeeStatus="inactive")), ["101"])
        self.assertEqual(
            codes(self.run_report("mobile-app-access", employeeStatus="active", access="signed_in")), ["9", "201"]
        )
        self.assertEqual(codes(self.run_report("mobile-app-access", branchIds=str(self.b2.id))), ["201"])
        self.assertEqual(codes(self.run_report("mobile-app-access", employmentType="production")), ["20"])
        self.assertEqual(
            codes(self.run_report("mobile-app-access", employeeIds=f"{self.e9.id},{self.e10.id}")), ["9", "10"]
        )

    def test_branch_isolation(self):
        body = self.run_report("mobile-app-access", self.user)
        self.assertNotIn("201", codes(body))
        self.assertEqual(body["totals"]["devices"], 2)  # the other branch's push token is not counted
        self.assertEqual(self.run_report("mobile-app-access", self.user, branchIds=str(self.b2.id))["rows"], [])

    def test_no_password_hash_ever_leaves_the_server(self):
        for fmt in ("xlsx", "pdf"):
            self.assertNotIn(b"SECRETHASH", self.raw("mobile-app-access", fmt=fmt).content)
        self.assertNotIn("SECRETHASH", self.raw("mobile-app-access").content.decode())
        self.assertNotIn("passwordHash", self.raw("mobile-app-access").content.decode())

    def test_department_and_designation_filters_and_no_branch_rows(self):
        self.assertEqual(codes(self.run_report("mobile-app-access", departmentIds=str(self.pack.id))), ["201"])
        self.assertEqual(codes(self.run_report("mobile-app-access", designationIds="999999")), [])
        stray = make_emp("500", "Stray", self.cut, None, password_hash=SECRET)
        row = by_key(self.run_report("mobile-app-access"), "500")
        self.assertEqual((row["appPassword"], row["department"]), ("Set", "CUTTING"))
        self.assertNotIn("500", codes(self.run_report("mobile-app-access", self.user)))
        self.assertEqual(stray.branch_id, None)


# ═══════════════════════════════════════════════════════════════════════════
#  HOD reports
# ═══════════════════════════════════════════════════════════════════════════


class HodReportTests(_Base):
    """
    Departments CUTTING / SEWING / DYEING / STORES / FINISHING (Unit 1), PACKING (Unit 2).
    HODs: M1 Ravi (CUTTING + S3 individually), M2 Sita (SEWING; leave rights only), M3 Tara (inactive account),
    M4 Uma (PACKING; no rights), M5 Vimal (employee inactive; SEWING after Sita, and FINISHING), M6 Wasim (nothing).
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        mk = lambda n, b: Department.objects.create(name=n, branch=b)  # noqa: E731
        cls.cut, cls.sew, cls.dye, cls.stores, cls.fin = (
            mk(n, cls.b1) for n in ("CUTTING", "SEWING", "DYEING", "STORES", "FINISHING")
        )
        cls.pack = mk("PACKING", cls.b2)
        e = lambda code, first, last, dept, branch=None, **kw: make_emp(  # noqa: E731
            code, first, dept, branch or cls.b1, last=last, **kw
        )
        cls.H1, cls.H2, cls.H3 = (
            e("H1", "Ravi", "Kumar", cls.cut),
            e("H2", "Sita", "Devi", cls.sew),
            e("H3", "Tara", "Sen", cls.cut),
        )
        cls.H4 = e("H4", "Uma", "Rao", cls.pack, cls.b2)
        cls.H5 = e("H5", "Vimal", "Das", cls.sew, status="inactive")
        cls.H6 = e("H6", "Wasim", "Khan", cls.cut)
        cls.S = {
            1: e("S1", "S", "One", cls.cut),
            2: e("S2", "S", "Two", cls.cut),
            3: e("S3", "S", "Three", cls.sew),
            4: e("S4", "S", "Four", cls.sew, status="inactive"),
            5: e("S5", "S", "Five", cls.sew),
            6: e("S6", "S", "Six", None),
            7: e("S7", "S", "Seven", cls.pack, cls.b2),
            8: e("S8", "S", "Eight", cls.pack, cls.b2),
            9: e("S9", "S", "Nine", cls.stores),
            10: e("S10", "S", "Ten", cls.dye),
            11: e("S11", "S", "Eleven", cls.fin),
        }
        allrights = dict.fromkeys(
            (
                "can_approve_leaves",
                "can_approve_permissions",
                "can_approve_resignations",
                "can_approve_attendance",
                "can_approve_casual_leave",
                "can_approve_on_duty",
                "can_approve_missing_punch",
            ),
            True,
        )
        norights = {k: False for k in allrights}
        cls.M1 = DepartmentManager.objects.create(employee=cls.H1)
        cls.M2 = DepartmentManager.objects.create(employee=cls.H2, **{**norights, "can_approve_leaves": True})
        cls.M3 = DepartmentManager.objects.create(employee=cls.H3, is_active=False)
        cls.M4 = DepartmentManager.objects.create(employee=cls.H4, **norights)
        cls.M5 = DepartmentManager.objects.create(employee=cls.H5)
        cls.M6 = DepartmentManager.objects.create(employee=cls.H6)
        D, E = ManagerDepartmentAssignment.objects.create, ManagerEmployeeAssignment.objects.create
        D(manager=cls.M1, department=cls.cut)
        D(manager=cls.M2, department=cls.sew)
        D(manager=cls.M5, department=cls.sew)  # second holder of SEWING: Sita (earlier) wins
        D(manager=cls.M3, department=cls.cut)  # inactive HOD: never counts
        D(manager=cls.M3, department=cls.dye)
        E(manager=cls.M1, employee=cls.S[3])  # individual assignment beats Sita's department
        E(manager=cls.M3, employee=cls.S[1])  # under an inactive HOD: ignored
        D(manager=cls.M4, department=cls.pack)
        D(manager=cls.M5, department=cls.fin)
        cls.user = make_user("ea_hod_b1", {"reports": "view", "user_management": "view"}, cls.b1)

    # -- directory ---------------------------------------------------------
    def test_directory_golden_rows(self):
        body = self.run_report("hod-directory")
        self.assertEqual(codes(body), ["H1", "H2", "H4", "H5", "H6"])
        r = by_key(body, "H1")
        self.assertEqual((r["teamSize"], r["overlapCount"], r["directAssignments"], r["rightsCount"]), (5, 0, 1, 7))
        self.assertEqual(
            (r["departmentsCovered"], r["hodActive"], r["employeeStatus"], r["branch"]),
            ("CUTTING", "Active", "Active", "Unit 1"),
        )
        self.assertIsNone(r["issues"])
        self.assertEqual(
            r["approves"], "Leave, Permission, Resignation, Attendance, Casual Leave, On-Duty, Missing Punch"
        )
        r = by_key(body, "H2")
        self.assertEqual((r["teamSize"], r["overlapCount"], r["rightsCount"], r["approves"]), (1, 1, 1, "Leave"))
        r = by_key(body, "H4")
        self.assertEqual(
            (r["teamSize"], r["rightsCount"], r["approves"], r["issues"], r["branch"]),
            (2, 0, None, "All approval rights off", "Unit 2"),
        )
        r = by_key(body, "H5")
        self.assertEqual(
            (r["teamSize"], r["overlapCount"], r["employeeStatus"], r["hodActive"]), (1, 3, "Inactive", "Active")
        )
        self.assertEqual(r["departmentsCovered"], "FINISHING, SEWING")
        self.assertEqual(r["issues"], "HOD employee inactive")
        r = by_key(body, "H6")
        self.assertEqual((r["teamSize"], r["directAssignments"], r["departmentsCovered"]), (0, 0, None))
        self.assertEqual(r["issues"], "No departments or employees assigned")
        self.assertEqual(
            cards(body),
            {
                "HODs listed": 5,
                "Active HODs": 5,
                "Inactive HODs": 0,
                "Employees covered": 9,
                "Active employees with no HOD": 3,
                "HODs whose employee is inactive": 1,
            },
        )
        self.assert_totals_match(body, "teamSize", "overlapCount", "directAssignments")
        self.assertEqual(body["totals"]["teamSize"], 9)

    def test_directory_agrees_with_the_hod_scope_rule(self):
        body = self.run_report("hod-directory", state="all")
        for m in DepartmentManager.objects.select_related("employee"):
            managed, overridden = hod_scope.coverage(m)
            active = set(Employee.objects.filter(id__in=managed, status="active").values_list("id", flat=True))
            row = by_key(body, m.employee.employee_code)
            if m.is_active:
                self.assertEqual(row["teamSize"], len(active - {m.employee_id}), m.employee.employee_code)
            live_overlap = (
                Employee.objects.filter(id__in=list(overridden), status="active").exclude(id=m.employee_id).count()
            )
            self.assertEqual(row["overlapCount"], live_overlap, m.employee.employee_code)

    def test_directory_state_and_problem_filters(self):
        body = self.run_report("hod-directory", state="all")
        self.assertEqual(codes(body), ["H1", "H2", "H3", "H4", "H5", "H6"])
        r = by_key(body, "H3")
        self.assertEqual(
            (r["hodActive"], r["teamSize"], r["overlapCount"], r["directAssignments"]), ("Inactive", 0, 4, 1)
        )
        self.assertEqual(r["departmentsCovered"], "CUTTING, DYEING")
        self.assertEqual(r["issues"], "HOD account inactive")
        self.assertEqual(cards(body)["Inactive HODs"], 1)
        self.assertEqual(codes(self.run_report("hod-directory", state="inactive")), ["H3"])
        self.assertEqual(
            codes(self.run_report("hod-directory", state="all", problemsOnly="true")), ["H3", "H4", "H5", "H6"]
        )

    def test_directory_scope_filters_and_branch_isolation(self):
        self.assertEqual(codes(self.run_report("hod-directory", departmentIds=str(self.sew.id))), ["H2", "H5"])
        self.assertEqual(codes(self.run_report("hod-directory", branchIds=str(self.b2.id))), ["H4"])
        body = self.run_report("hod-directory", self.user)
        self.assertEqual(codes(body), ["H1", "H2", "H5", "H6"])
        self.assertEqual(cards(body)["Active employees with no HOD"], 3)
        self.assertEqual(self.run_report("hod-directory", self.user, branchIds=str(self.b2.id))["rows"], [])

    # -- mapping -----------------------------------------------------------
    def test_mapping_golden_rows(self):
        body = self.run_report("hod-mapping")
        self.assertEqual(len(body["rows"]), 15)
        pick = lambda c: {k: by_key(body, c)[k] for k in ("hodCode", "hodName", "via", "isHod", "reason")}  # noqa: E731
        self.assertEqual(
            pick("S3"), {"hodCode": "H1", "hodName": "Ravi Kumar", "via": "Direct", "isHod": None, "reason": None}
        )
        self.assertEqual(pick("S1")["via"], "Department")  # the individual assignment sits under an inactive HOD
        self.assertEqual(pick("S1")["hodCode"], "H1")
        self.assertEqual(pick("S5")["hodCode"], "H2")  # earliest holder of SEWING
        self.assertEqual(pick("H1")["isHod"], "HOD")
        self.assertEqual(pick("H3")["hodCode"], "H1")
        self.assertEqual(pick("S7")["hodCode"], "H4")
        self.assertEqual(pick("S11")["hodCode"], "H5")
        self.assertEqual(pick("S11")["reason"], "HOD's employee record is inactive")
        self.assertEqual(pick("S6")["reason"], "No department set and no direct HOD")
        self.assertEqual(pick("S9")["reason"], "Department has no active HOD")
        self.assertEqual(pick("S10")["reason"], "Assigned HOD is inactive")
        for c in ("S6", "S9", "S10"):
            self.assertIsNone(by_key(body, c)["hodCode"])
            self.assertIsNone(by_key(body, c)["via"])
        self.assertEqual(
            cards(body),
            {
                "Employees checked": 15,
                "With a HOD": 12,
                "Without a HOD": 3,
                "Covered": 80.0,
                "Without a HOD (excl. HODs)": 3,
            },
        )
        self.assertTrue(any("Uncovered (excluding HODs) by department" in n and "STORES 1" in n for n in body["notes"]))

    def test_mapping_agrees_with_effective_owner_map(self):
        body = self.run_report("hod-mapping", employeeStatus="all")
        emps = list(Employee.objects.all())
        owner = hod_scope.effective_owner_map([e.id for e in emps])
        mgr_code = {m.id: m.employee.employee_code for m in DepartmentManager.objects.select_related("employee")}
        for e in emps:
            expected = mgr_code.get(owner.get(e.id))
            self.assertEqual(by_key(body, e.employee_code)["hodCode"], expected, e.employee_code)

    def test_mapping_filters_and_isolation(self):
        self.assertEqual(codes(self.run_report("hod-mapping", coverage="uncovered")), ["S6", "S9", "S10"])
        self.assertEqual(len(self.run_report("hod-mapping", coverage="covered")["rows"]), 12)
        body = self.run_report("hod-mapping", employeeStatus="all")
        self.assertEqual(len(body["rows"]), 17)
        self.assertEqual(by_key(body, "S4")["hodCode"], "H2")
        self.assertEqual(
            codes(self.run_report("hod-mapping", employeeIds=f"{self.S[3].id},{self.S[7].id}")), ["S3", "S7"]
        )
        self.assertEqual(codes(self.run_report("hod-mapping", departmentIds=str(self.pack.id))), ["H4", "S7", "S8"])
        body = self.run_report("hod-mapping", self.user)
        self.assertFalse({"H4", "S7", "S8"} & set(codes(body)))
        self.assertEqual(self.run_report("hod-mapping", self.user, employeeIds=str(self.S[7].id))["rows"], [])
        self.assertEqual(self.run_report("hod-mapping", self.user, branchIds=str(self.b2.id))["rows"], [])

    # -- conflicts ---------------------------------------------------------
    def test_conflicts_golden_rows(self):
        body = self.run_report("hod-conflicts")
        self.assertEqual(codes(body), ["S3", "S5"])
        r = by_key(body, "S3")
        self.assertEqual(
            (r["effectiveHod"], r["effectiveVia"], r["alsoListedUnder"], r["otherListings"]),
            ("Ravi Kumar", "Direct", "Sita Devi (department), Vimal Das (department)", 2),
        )
        r = by_key(body, "S5")
        self.assertEqual(
            (r["effectiveHod"], r["effectiveVia"], r["alsoListedUnder"], r["otherListings"]),
            ("Sita Devi", "Department", "Vimal Das (department)", 1),
        )
        self.assertEqual(cards(body), {"Employees in conflict": 2, "HODs involved": 3})
        self.assertEqual(body["totals"]["otherListings"], 3)
        self.assertEqual(codes(self.run_report("hod-conflicts", employeeStatus="all")), ["S3", "S4", "S5"])

    def test_conflicts_empty_state_filters_and_isolation(self):
        body = self.run_report("hod-conflicts", departmentIds=str(self.cut.id))
        self.assertEqual(body["rows"], [])
        self.assertTrue(body["notes"][0].startswith("No conflicts found"))
        self.assertEqual(cards(body), {"Employees in conflict": 0, "HODs involved": 0})
        self.assertEqual(codes(self.run_report("hod-conflicts", self.user)), ["S3", "S5"])
        self.assertEqual(self.run_report("hod-conflicts", self.user, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(codes(self.run_report("hod-conflicts", branchIds=str(self.b1.id))), ["S3", "S5"])
        self.assertEqual(codes(self.run_report("hod-conflicts", employeeIds=str(self.S[5].id))), ["S5"])
        self.assertEqual(self.run_report("hod-conflicts", employeeIds=str(self.S[1].id))["rows"], [])

    def test_mapping_type_and_designation_filters_and_employee_without_branch(self):
        self.assertEqual(len(self.run_report("hod-mapping", employmentType="staff")["rows"]), 15)
        self.assertEqual(self.run_report("hod-mapping", employmentType="production")["rows"], [])
        self.assertEqual(self.run_report("hod-mapping", designationIds="999999")["rows"], [])
        make_emp("NB1", "Nobranch", self.sew, None)  # legacy row: no branch at all
        body = self.run_report("hod-mapping")
        row = by_key(body, "NB1")
        self.assertEqual((row["branch"], row["hodCode"]), ("No branch", "H2"))
        self.assertNotIn("NB1", codes(self.run_report("hod-mapping", self.user)))


# ═══════════════════════════════════════════════════════════════════════════
#  Recruitment
# ═══════════════════════════════════════════════════════════════════════════


class RecruitmentReportTests(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        mk = lambda n, b: Department.objects.create(name=n, branch=b)  # noqa: E731
        cls.cut, cls.sew, cls.stores, cls.dye = (mk(n, cls.b1) for n in ("CUTTING", "SEWING", "STORES", "DYEING"))
        cls.pack = mk("PACKING", cls.b2)
        cls.legacy = mk("LEGACY", None)

        # -- manpower fixtures
        n = [0]

        def staff(dept, branch, count, etype="staff", status="active"):
            for _ in range(count):
                n[0] += 1
                make_emp(f"R{n[0]}", "R", dept, branch, etype=etype, status=status)

        staff(cls.cut, cls.b1, 3)
        staff(cls.cut, cls.b1, 1, "production")
        staff(cls.cut, cls.b1, 1, "staff", "inactive")
        staff(cls.sew, cls.b1, 2)
        staff(cls.sew, cls.b1, 2, "production")
        staff(cls.pack, cls.b2, 1)
        staff(cls.dye, cls.b1, 2)
        DepartmentHeadcount.objects.create(department=cls.cut, required_count=5, notes="Peak season")
        DepartmentHeadcount.objects.create(department=cls.sew, required_count=2)
        DepartmentHeadcount.objects.create(department=cls.stores, required_count=1)
        DepartmentHeadcount.objects.create(department=cls.dye, required_count=1)

        # -- jobs and applicants
        def job(title, dept, status, when):
            j = Job.objects.create(title=title, department=dept, status=status, salary_range="15k-20k")
            Job.objects.filter(pk=j.pk).update(created_at=when)
            return j

        cls.j1 = job("Tailor", cls.cut, "open", ist(2026, 9, 10, 10))
        cls.j2 = job("Helper", cls.cut, "closed", ist(2026, 8, 1, 10))
        cls.j3 = job("Iron man", cls.sew, "open", ist(2026, 9, 19, 10))
        cls.j4 = job("Store keeper", cls.stores, "open", ist(2026, 9, 20, 9))
        cls.j5 = job("Accountant", None, "open", ist(2026, 9, 1, 10))

        def applicant(name, j, status, when, **kw):
            a = Applicant.objects.create(
                job=j, name=name, email=f"{name.split()[0].lower()}@x.in", phone="9000000000", status=status, **kw
            )
            Applicant.objects.filter(pk=a.pk).update(created_at=when)
            return a

        cls.a = {
            1: applicant("Ravi A", cls.j1, "applied", ist(2026, 9, 11, 10), experience="2 years"),
            2: applicant("Sita B", cls.j1, "attended", ist(2026, 9, 12, 10)),
            3: applicant("Uma C", cls.j1, "selected", ist(2026, 9, 13, 10)),
            4: applicant("Vikram D", cls.j1, "rejected", ist(2026, 9, 14, 10)),
            5: applicant("Wasim E", cls.j1, "hold", ist(2026, 9, 15, 10)),
            6: applicant("Old F", cls.j2, "rejected", ist(2026, 8, 5, 10)),
            7: applicant("Old G", cls.j2, "rejected", ist(2026, 8, 5, 11)),
            8: applicant("Xena F", cls.j3, "applied", ist(2026, 9, 19, 10)),
            9: applicant("Yash G", cls.j5, "selected", ist(2026, 9, 2, 10)),
            10: applicant("Zoya H", cls.j3, "applied", ist(2026, 9, 1, 0, 30)),
            11: applicant("Late I", cls.j3, "applied", ist(2026, 10, 1, 0, 30)),
        }

        # -- hiring criteria and screening candidates
        cls.rs1 = HiringRuleSet.objects.create(
            name="Cutting Master",
            department=cls.cut,
            required_skills=["Cutting", "Pattern"],
            soft_skills=["Teamwork"],
            education_qualification="ITI",
            min_experience_years=2,
            preferred_city="Tirupur",
        )
        cls.rs2 = HiringRuleSet.objects.create(name="Sewing Operator", department=cls.sew, required_skills=["Sewing"])
        cls.rs3 = HiringRuleSet.objects.create(name="Packing", department=cls.pack, is_active=False)

        def cand(
            key,
            rs,
            dept,
            name,
            status,
            score,
            created,
            source="single",
            rank=None,
            invited=None,
            when=None,
            resume=True,
        ):
            c = ScreeningCandidate.objects.create(
                rule_set=rs,
                department=dept,
                # a real stored file: 'On file' now means the file exists, not just that the column has a path
                resume_file=ContentFile(b"%PDF-1.4 resume", name="x.pdf") if resume else "",
                original_filename=f"{name}.pdf",
                source=source,
                candidate_name=name,
                phone="9111111111",
                email=f"{name.lower()}@x.in",
                city="Tirupur",
                extracted_experience_years=3.5 if key == 1 else None,
                extracted_education="ITI" if key == 1 else None,
                match_score=score,
                rank_in_batch=rank,
                status=status,
                interview_invited_at=invited,
                interview_datetime=when,
            )
            ScreeningCandidate.objects.filter(pk=c.pk).update(created_at=created)
            return c

        cls.c = {
            1: cand(1, cls.rs1, cls.cut, "Ganesh", "shortlisted", 80, ist(2026, 9, 5, 11), "bulk", 1),
            2: cand(
                2,
                cls.rs1,
                cls.cut,
                "Hari",
                "selected",
                90,
                ist(2026, 9, 6, 9),
                invited=ist(2026, 9, 8, 10),
                when=ist(2026, 9, 22, 10, 30),
            ),
            3: cand(3, cls.rs1, cls.cut, "Indu", "rejected", 40, ist(2026, 9, 6, 10), resume=False),
            4: cand(4, cls.rs1, cls.cut, "Jaya", "uploaded", None, ist(2026, 9, 7, 10)),
            5: cand(5, cls.rs2, cls.sew, "Kala", "selected", 70, ist(2026, 9, 8, 10)),
            6: cand(6, cls.rs2, cls.sew, "Lata", "shortlisted", 60, ist(2026, 9, 9, 10)),
            7: cand(7, cls.rs3, cls.pack, "Mani", "not_shortlisted", 30, ist(2026, 9, 10, 10), "bulk", 2),
            8: cand(8, cls.rs1, None, "Nila", "screened", 50, ist(2026, 9, 11, 10)),
            9: cand(
                9,
                cls.rs1,
                cls.cut,
                "Omar",
                "selected",
                85,
                ist(2026, 8, 15, 10),
                invited=ist(2026, 8, 20, 10),
                when=ist(2026, 9, 15, 10),
            ),
            10: cand(10, cls.rs2, cls.sew, "Pari", "shortlisted", 65, ist(2026, 10, 1, 0, 30)),
            11: cand(11, cls.rs2, cls.sew, "Qadir", "screened", 40, ist(2026, 9, 1, 0, 30)),
            12: cand(
                12,
                cls.rs2,
                cls.sew,
                "Ravi",
                "rejected",
                20,
                ist(2026, 9, 12, 10),
                invited=ist(2026, 9, 13, 10),
                when=ist(2026, 9, 25, 9),
            ),
            13: cand(13, cls.rs3, cls.pack, "Rani", "selected", 60, ist(2026, 9, 12, 11)),
        }
        cls.rec_user = make_user("ea_rec_b1", {"reports": "view", "recruitment": "view"}, cls.b1)
        cls.roles_user = make_user("ea_roles_b1", {"reports": "view", "recruitment.required_roles": "view"}, cls.b1)
        cls.screen_user = make_user("ea_screen_b1", {"reports": "view", "recruitment.resume_screening": "view"}, cls.b1)

    # -- manpower ----------------------------------------------------------
    def test_manpower_golden_rows(self):
        body = self.run_report("manpower-requirement")
        pick = lambda d: {
            k: by_key(body, d, "department")[k]
            for k in (  # noqa: E731
                "requiredCount",
                "currentCount",
                "vacancy",
                "surplus",
                "fillPct",
                "status",
                "openJobs",
                "inPipeline",
            )
        }
        self.assertEqual(codes(body, "department"), ["CUTTING", "DYEING", "SEWING", "STORES", "PACKING", "LEGACY"])
        self.assertEqual(
            pick("CUTTING"),
            dict(
                requiredCount=5,
                currentCount=3,
                vacancy=2,
                surplus=0,
                fillPct=60.0,
                status="Understaffed",
                openJobs=1,
                inPipeline=3,
            ),
        )
        self.assertEqual(
            pick("DYEING"),
            dict(
                requiredCount=1,
                currentCount=2,
                vacancy=0,
                surplus=1,
                fillPct=200.0,
                status="Overstaffed",
                openJobs=0,
                inPipeline=0,
            ),
        )
        self.assertEqual(
            pick("SEWING"),
            dict(
                requiredCount=2,
                currentCount=2,
                vacancy=0,
                surplus=0,
                fillPct=100.0,
                status="Fully staffed",
                openJobs=1,
                inPipeline=3,
            ),
        )
        self.assertEqual(
            pick("STORES"),
            dict(
                requiredCount=1,
                currentCount=0,
                vacancy=1,
                surplus=0,
                fillPct=0.0,
                status="Understaffed",
                openJobs=1,
                inPipeline=0,
            ),
        )
        self.assertEqual(
            pick("PACKING"),
            dict(
                requiredCount=0,
                currentCount=1,
                vacancy=0,
                surplus=0,
                fillPct=None,
                status="Not planned",
                openJobs=0,
                inPipeline=1,
            ),
        )
        self.assertEqual(by_key(body, "LEGACY", "department")["branch"], "No branch")
        self.assertEqual(by_key(body, "CUTTING", "department")["notes"], "Peak season")
        self.assertIsNotNone(by_key(body, "CUTTING", "department")["updatedOn"])
        self.assertIsNone(by_key(body, "PACKING", "department")["updatedOn"])
        self.assertEqual(
            cards(body), {"Required": 9, "Current": 8, "Vacancies": 3, "Surplus": 1, "Departments with gaps": 2}
        )
        self.assert_totals_match(body, "requiredCount", "currentCount", "vacancy", "surplus", "openJobs", "inPipeline")
        self.assertEqual(body["totals"]["inPipeline"], 7)

    def test_manpower_basis_and_filters(self):
        body = self.run_report("manpower-requirement", basis="all")
        self.assertEqual(
            by_key(body, "CUTTING", "department")["currentCount"], 4
        )  # 3 staff + 1 production, inactive excluded
        self.assertEqual(by_key(body, "SEWING", "department")["surplus"], 2)
        body = self.run_report("manpower-requirement", basis="production")
        self.assertEqual(by_key(body, "CUTTING", "department")["vacancy"], 4)
        self.assertEqual(by_key(body, "SEWING", "department")["status"], "Fully staffed")
        self.assertEqual(
            codes(self.run_report("manpower-requirement", onlyGaps="true"), "department"), ["CUTTING", "STORES"]
        )
        self.assertEqual(
            codes(self.run_report("manpower-requirement", branchIds=str(self.b2.id)), "department"), ["PACKING"]
        )
        self.assertEqual(
            codes(self.run_report("manpower-requirement", departmentIds=str(self.sew.id)), "department"), ["SEWING"]
        )

    def test_manpower_branch_isolation(self):
        body = self.run_report("manpower-requirement", self.roles_user)
        self.assertEqual(codes(body, "department"), ["CUTTING", "DYEING", "SEWING", "STORES"])
        self.assertEqual(
            self.run_report("manpower-requirement", self.roles_user, branchIds=str(self.b2.id))["rows"], []
        )
        self.assertEqual(
            self.run_report("manpower-requirement", self.roles_user, departmentIds=str(self.pack.id))["rows"], []
        )

    def test_manpower_xlsx_numbers_are_real_numbers_with_a_totals_row(self):
        ws = self.sheet("manpower-requirement")
        self.assertEqual(
            [c.value for c in ws[7]][:8],
            ["Branch", "Department", "Required", "Current", "Vacancy", "Surplus", "Fill %", "Status"],
        )
        rows = {r[1].value or r[0].value: [c.value for c in r] for r in ws.iter_rows(min_row=8)}
        self.assertEqual(rows["CUTTING"][2:8], [5, 3, 2, 0, 60.0, "Understaffed"])
        self.assertEqual((rows["CUTTING"][8], rows["CUTTING"][9], rows["CUTTING"][10]), (1, 3, "Peak season"))
        self.assertIsNone(rows["PACKING"][6])  # not planned: a blank cell, never 0 or 100
        self.assertEqual(rows["PACKING"][7], "Not planned")
        self.assertEqual(rows["TOTAL"][2:6], [9, 8, 3, 1])
        self.assertEqual((rows["TOTAL"][8], rows["TOTAL"][9]), (3, 7))

    # -- job openings ------------------------------------------------------
    def test_job_openings_golden_rows(self):
        body = self.run_report("job-openings")
        self.assertEqual(codes(body, "title"), ["Store keeper", "Iron man", "Tailor", "Accountant"])
        r = by_key(body, "Tailor", "title")
        self.assertEqual(
            (
                r["status"],
                r["postedOn"],
                r["daysOpen"],
                r["applicants"],
                r["applied"],
                r["attended"],
                r["selected"],
                r["rejected"],
                r["selectionRatePct"],
            ),
            ("Open", "2026-09-10", 10, 5, 1, 1, 1, 1, 20.0),
        )
        self.assertEqual(
            (by_key(body, "Iron man", "title")["daysOpen"], by_key(body, "Iron man", "title")["selectionRatePct"]),
            (1, 0.0),
        )
        r = by_key(body, "Store keeper", "title")
        self.assertEqual(
            (r["daysOpen"], r["applicants"], r["selectionRatePct"], r["department"], r["branch"]),
            (0, 0, None, "STORES", "Unit 1"),
        )
        r = by_key(body, "Accountant", "title")
        self.assertEqual(
            (r["department"], r["branch"], r["daysOpen"], r["selectionRatePct"]), ("Unassigned", "No branch", 19, 100.0)
        )
        self.assertEqual(cards(body), {"Open positions": 4, "Applicants": 9, "Selected": 2, "Average days open": 7.5})
        self.assert_totals_match(body, "applicants", "applied", "attended", "selected", "rejected")

    def test_job_openings_filters(self):
        body = self.run_report("job-openings", status="all")
        self.assertEqual(len(body["rows"]), 5)
        closed = by_key(body, "Helper", "title")
        self.assertEqual(
            (closed["status"], closed["daysOpen"], closed["rejected"], closed["selectionRatePct"]),
            ("Closed", None, 2, 0.0),
        )
        self.assertEqual(codes(self.run_report("job-openings", status="closed"), "title"), ["Helper"])
        self.assertEqual(
            codes(self.run_report("job-openings", status="all", postedWithinDays="5"), "title"),
            ["Store keeper", "Iron man"],
        )
        self.assertEqual(codes(self.run_report("job-openings", departmentIds=str(self.cut.id)), "title"), ["Tailor"])
        self.assertEqual(
            codes(self.run_report("job-openings", branchIds=str(self.b1.id)), "title"),
            ["Store keeper", "Iron man", "Tailor"],
        )

    def test_job_openings_branch_isolation(self):
        body = self.run_report("job-openings", self.rec_user)
        self.assertNotIn("Accountant", codes(body, "title"))  # no department = no branch = unscoped users only
        self.assertEqual(self.run_report("job-openings", self.rec_user, branchIds=str(self.b2.id))["rows"], [])

    # -- applicants --------------------------------------------------------
    def test_applicant_register_golden_rows_and_cards(self):
        body = self.run_report("applicant-register", **WINDOW)
        self.assertEqual(
            codes(body, "name"),
            ["Xena F", "Wasim E", "Vikram D", "Uma C", "Sita B", "Ravi A", "Yash G", "Zoya H"],
        )
        r = by_key(body, "Ravi A", "name")
        self.assertEqual(
            (r["appliedOn"], r["jobTitle"], r["department"], r["experience"], r["status"], r["email"]),
            ("2026-09-11", "Tailor", "CUTTING", "2 years", "Applied", "ravi@x.in"),
        )
        self.assertEqual(by_key(body, "Wasim E", "name")["status"], "Hold")
        self.assertEqual(by_key(body, "Yash G", "name")["department"], "Unassigned")
        self.assertEqual(by_key(body, "Zoya H", "name")["appliedOn"], "2026-09-01")  # 31-Aug UTC, 1-Sep IST
        self.assertNotIn("Late I", codes(body, "name"))
        self.assertEqual(cards(body), {"Applicants": 8, "Applied": 3, "Attended": 1, "Selected": 2, "Rejected": 1})
        self.assertTrue(any("Tailor 5, Iron man 2, Accountant 1" in n for n in body["notes"]))

    def test_applicant_register_filters_and_isolation(self):
        self.assertEqual(
            sorted(codes(self.run_report("applicant-register", status="applied,selected", **WINDOW), "name")),
            ["Ravi A", "Uma C", "Xena F", "Yash G", "Zoya H"],
        )
        self.assertEqual(len(self.run_report("applicant-register", jobTitle="tailor", **WINDOW)["rows"]), 5)
        self.assertEqual(
            sorted(codes(self.run_report("applicant-register", departmentIds=str(self.sew.id), **WINDOW), "name")),
            ["Xena F", "Zoya H"],
        )
        self.assertEqual(
            len(self.run_report("applicant-register", dateFrom="2026-08-01", dateTo="2026-08-31")["rows"]), 2
        )
        body = self.run_report("applicant-register", self.rec_user, **WINDOW)
        self.assertNotIn("Yash G", codes(body, "name"))
        self.assertEqual(cards(body)["Applicants"], 7)
        self.assertEqual(
            self.run_report("applicant-register", self.rec_user, branchIds=str(self.b2.id), **WINDOW)["rows"], []
        )

    # -- screening pipeline ------------------------------------------------
    def test_screening_pipeline_golden_rows_and_cards(self):
        body = self.run_report("screening-pipeline", **WINDOW)
        self.assertEqual(
            codes(body, "candidateName"),
            ["Rani", "Ravi", "Nila", "Mani", "Lata", "Kala", "Jaya", "Indu", "Hari", "Ganesh", "Qadir"],
        )
        r = by_key(body, "Ganesh", "candidateName")
        self.assertEqual(
            (
                r["status"],
                r["matchScore"],
                r["rankInBatch"],
                r["source"],
                r["experienceYears"],
                r["education"],
                r["resumeOnFile"],
                r["department"],
                r["ruleSet"],
            ),
            ("Shortlisted", 80.0, 1, "Bulk", 3.5, "ITI", "On file", "CUTTING", "Cutting Master"),
        )
        self.assertIsNone(by_key(body, "Jaya", "candidateName")["matchScore"])  # unscored stays blank, never 0
        self.assertEqual(by_key(body, "Indu", "candidateName")["resumeOnFile"], "Removed")
        self.assertEqual(by_key(body, "Mani", "candidateName")["status"], "Not shortlisted")
        self.assertEqual(by_key(body, "Hari", "candidateName")["interviewInvitedAt"], "2026-09-08 10:00")
        self.assertEqual(by_key(body, "Hari", "candidateName")["interviewDatetime"], "2026-09-22 10:30")
        self.assertEqual(by_key(body, "Nila", "candidateName")["department"], "Unassigned")
        self.assertEqual(by_key(body, "Qadir", "candidateName")["createdAt"], "2026-09-01 00:30")
        self.assertNotIn("Pari", codes(body, "candidateName"))  # IST 1-Oct 00:30 is outside September
        self.assertNotIn("Omar", codes(body, "candidateName"))
        self.assertEqual(
            cards(body),
            {
                "Candidates": 11,
                "Shortlisted": 2,
                "Selected": 3,
                "Rejected": 2,
                "Average match score": 54.0,
                "Interview invites sent": 2,
            },
        )

    def test_screening_pipeline_filters_and_isolation(self):
        self.assertEqual(
            sorted(codes(self.run_report("screening-pipeline", status="shortlisted", **WINDOW), "candidateName")),
            ["Ganesh", "Lata"],
        )
        self.assertEqual(
            sorted(codes(self.run_report("screening-pipeline", source="bulk", **WINDOW), "candidateName")),
            ["Ganesh", "Mani"],
        )
        self.assertEqual(
            sorted(codes(self.run_report("screening-pipeline", minScore="70", **WINDOW), "candidateName")),
            ["Ganesh", "Hari", "Kala"],
        )
        self.assertEqual(
            sorted(
                codes(self.run_report("screening-pipeline", departmentIds=str(self.sew.id), **WINDOW), "candidateName")
            ),
            ["Kala", "Lata", "Qadir", "Ravi"],
        )
        body = self.run_report("screening-pipeline", self.screen_user, **WINDOW)
        self.assertFalse({"Mani", "Nila", "Rani"} & set(codes(body, "candidateName")))
        self.assertEqual(cards(body)["Candidates"], 8)
        self.assertEqual(
            self.run_report("screening-pipeline", self.screen_user, branchIds=str(self.b2.id), **WINDOW)["rows"], []
        )

    def test_screening_never_exposes_resume_files(self):
        text = self.raw("screening-pipeline", **WINDOW).content.decode()
        self.assertNotIn("resumes/", text)
        self.assertNotIn("/resume", text)

    # -- interview schedule ------------------------------------------------
    def test_interview_schedule_golden_rows(self):
        body = self.run_report("interview-schedule")
        self.assertEqual(codes(body, "candidateName"), ["Omar", "Hari", "Ravi", "Kala", "Rani"])
        pick = lambda n: (
            by_key(body, n, "candidateName")["interviewDatetime"],
            by_key(body, n, "candidateName")["inviteState"],
        )  # noqa: E731
        self.assertEqual(pick("Omar"), ("2026-09-15 10:00", "Invited"))
        self.assertEqual(pick("Hari"), ("2026-09-22 10:30", "Invited"))
        self.assertEqual(pick("Ravi"), ("2026-09-25 09:00", "Invited"))  # rejected after being scheduled: still listed
        self.assertEqual(pick("Kala"), (None, "Pending invite"))
        self.assertEqual(pick("Rani"), (None, "Pending invite"))
        self.assertEqual(cards(body), {"Interviews scheduled": 3, "Invites pending": 2, "Interviews in next 7 days": 2})

    def test_interview_schedule_filters_and_isolation(self):
        names = lambda **p: codes(self.run_report("interview-schedule", **p), "candidateName")  # noqa: E731
        self.assertEqual(names(invited="pending_invite"), ["Kala", "Rani"])
        self.assertEqual(names(invited="invited"), ["Omar", "Hari", "Ravi"])
        self.assertEqual(names(when="upcoming"), ["Hari", "Ravi"])
        self.assertEqual(names(when="past"), ["Omar"])
        self.assertEqual(names(when="unscheduled"), ["Kala", "Rani"])
        self.assertEqual(names(departmentIds=str(self.pack.id)), ["Rani"])
        self.assertEqual(names(branchIds=str(self.b2.id)), ["Rani"])
        body = self.run_report("interview-schedule", self.screen_user)
        self.assertEqual(codes(body, "candidateName"), ["Omar", "Hari", "Ravi", "Kala"])
        self.assertEqual(cards(body)["Invites pending"], 1)
        self.assertEqual(self.run_report("interview-schedule", self.screen_user, branchIds=str(self.b2.id))["rows"], [])

    # -- funnel ------------------------------------------------------------
    def test_recruitment_funnel_golden_rows(self):
        body = self.run_report("recruitment-funnel", **WINDOW)
        self.assertEqual(codes(body, "department"), ["CUTTING", "PACKING", "SEWING", "Unassigned"])
        keys = (
            "candidates",
            "pending",
            "notShortlisted",
            "shortlisted",
            "selected",
            "rejected",
            "avgScore",
            "progressPct",
            "selectPct",
            "applicants",
            "applicantsSelected",
        )
        pick = lambda d: tuple(by_key(body, d, "department")[k] for k in keys)  # noqa: E731
        self.assertEqual(pick("CUTTING"), (4, 1, 0, 1, 1, 1, 70.0, 50.0, 25.0, 5, 1))
        self.assertEqual(pick("PACKING"), (2, 0, 1, 0, 1, 0, 45.0, 50.0, 50.0, 0, 0))
        self.assertEqual(pick("SEWING"), (4, 1, 0, 1, 1, 1, 47.5, 50.0, 25.0, 2, 0))
        self.assertEqual(pick("Unassigned"), (1, 1, 0, 0, 0, 0, 50.0, 0.0, 0.0, 1, 1))
        self.assertEqual(by_key(body, "Unassigned", "department")["branch"], "No branch")
        self.assertEqual(
            cards(body),
            {
                "Screened candidates": 11,
                "Shortlisted or selected": 45.5,
                "Selected": 27.3,
                "Job-board applicants": 8,
            },
        )
        self.assert_totals_match(body, "candidates", "selected", "applicants", "applicantsSelected")

    def test_recruitment_funnel_filters_and_isolation(self):
        self.assertEqual(
            codes(self.run_report("recruitment-funnel", departmentIds=str(self.cut.id), **WINDOW), "department"),
            ["CUTTING"],
        )
        self.assertEqual(
            codes(self.run_report("recruitment-funnel", branchIds=str(self.b2.id), **WINDOW), "department"), ["PACKING"]
        )
        body = self.run_report("recruitment-funnel", self.rec_user, **WINDOW)
        self.assertEqual(codes(body, "department"), ["CUTTING", "SEWING"])
        self.assertEqual(cards(body)["Screened candidates"], 8)
        self.assertEqual(
            self.run_report("recruitment-funnel", self.rec_user, branchIds=str(self.b2.id), **WINDOW)["rows"], []
        )

    # -- hiring criteria ---------------------------------------------------
    def test_hiring_criteria_golden_rows_and_isolation(self):
        body = self.run_report("hiring-criteria")
        self.assertEqual(codes(body, "name"), ["Cutting Master", "Packing", "Sewing Operator"])
        r = by_key(body, "Cutting Master", "name")
        self.assertEqual(
            (
                r["requiredSkills"],
                r["softSkills"],
                r["education"],
                r["minExperienceYears"],
                r["preferredCity"],
                r["isActive"],
                r["candidates"],
                r["branch"],
            ),
            ("Cutting, Pattern", "Teamwork", "ITI", 2.0, "Tirupur", "Active", 6, "Unit 1"),
        )
        self.assertEqual(by_key(body, "Packing", "name")["isActive"], "Inactive")
        self.assertEqual(
            (by_key(body, "Sewing Operator", "name")["candidates"], by_key(body, "Packing", "name")["candidates"]),
            (5, 2),
        )
        self.assertEqual(cards(body), {"Rule sets": 3, "Active rule sets": 2, "Candidates screened": 13})
        self.assert_totals_match(body, "candidates")
        self.assertEqual(codes(self.run_report("hiring-criteria", active="inactive"), "name"), ["Packing"])
        self.assertEqual(
            codes(self.run_report("hiring-criteria", active="active", branchIds=str(self.b1.id)), "name"),
            ["Cutting Master", "Sewing Operator"],
        )
        self.assertEqual(
            codes(self.run_report("hiring-criteria", departmentIds=str(self.sew.id)), "name"), ["Sewing Operator"]
        )
        scoped = self.run_report("hiring-criteria", self.screen_user)
        self.assertEqual(codes(scoped, "name"), ["Cutting Master", "Sewing Operator"])
        self.assertEqual(cards(scoped)["Candidates screened"], 11)
        self.assertEqual(self.run_report("hiring-criteria", self.screen_user, branchIds=str(self.b2.id))["rows"], [])

    def test_unscoped_branch_filter_on_applicants_screening_and_funnel_window(self):
        body = self.run_report("applicant-register", branchIds=str(self.b1.id), **WINDOW)
        self.assertEqual(
            sorted(codes(body, "name")),
            sorted(["Xena F", "Wasim E", "Vikram D", "Uma C", "Sita B", "Ravi A", "Zoya H"]),
        )
        self.assertEqual(codes(self.run_report("applicant-register", branchIds=str(self.b2.id), **WINDOW), "name"), [])
        body = self.run_report("screening-pipeline", branchIds=str(self.b2.id), **WINDOW)
        self.assertEqual(sorted(codes(body, "candidateName")), ["Mani", "Rani"])
        body = self.run_report("recruitment-funnel", dateFrom="2026-08-01", dateTo="2026-08-31")
        self.assertEqual(codes(body, "department"), ["CUTTING"])  # only Omar (uploaded 15-Aug)
        self.assertEqual(cards(body)["Screened candidates"], 1)
        self.assertEqual(cards(body)["Job-board applicants"], 2)  # the two closed-job rejections of 5-Aug
        self.assertEqual(by_key(body, "CUTTING", "department")["selected"], 1)

    def test_retention_note_uses_the_real_purge_window(self):
        from .screening_cleanup import RETENTION_DAYS

        body = self.run_report("screening-pipeline", **WINDOW)
        self.assertTrue(any(f"after {RETENTION_DAYS} days" in n for n in body["notes"]))


# ═══════════════════════════════════════════════════════════════════════════
#  Audit trail (admin only)
# ═══════════════════════════════════════════════════════════════════════════


class AuditReportTests(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        b1, b2 = cls.b1, cls.b2
        L = add_log
        cls.l = {
            1: L(ist(2026, 9, 1, 8, 0), "Admin One", "login", "auth", "Admin One (admin1) logged in", None, "10.0.0.1"),
            2: L(ist(2026, 9, 1, 8, 5), "system", "login_failed", "auth", "Failed login for: bob", None, "10.0.0.9"),
            3: L(
                ist(2026, 9, 1, 8, 10),
                "system",
                "login_blocked",
                "auth",
                "Locked-out login attempt for: bob",
                None,
                "10.0.0.9",
            ),
            4: L(
                ist(2026, 9, 2, 10, 0),
                "HR Two",
                "create",
                "employees",
                "Created employee 501 -Ravi Kumar",
                b1,
                "10.0.0.2",
                5,
            ),
            5: L(
                ist(2026, 9, 2, 11, 0),
                "HR Two",
                "update",
                "employees",
                "Updated employee 501 -Ravi Kumar",
                b1,
                "10.0.0.2",
                5,
            ),
            6: L(
                ist(2026, 9, 3, 12, 0),
                "HR Two",
                "delete",
                "employees",
                "Deleted employee 501 -Ravi Kumar",
                b1,
                "10.0.0.2",
                5,
            ),
            7: L(
                ist(2026, 9, 4, 9, 0),
                "HR Two",
                "create",
                "employees",
                "Bulk-imported employee 601 -Sita Devi",
                b2,
                "10.0.0.3",
                6,
            ),
            8: L(
                ist(2026, 9, 4, 9, 30),
                "Admin One",
                "update",
                "employees",
                "Bulk-updated employee 601 -changed Phone, Salary Amount",
                None,
                "10.0.0.1",
                6,
            ),
            9: L(
                ist(2026, 9, 5, 10, 0),
                "Admin One",
                "update",
                "employees",
                "Bulk enabled live location tracking for 12 employee(s)",
                None,
                "10.0.0.1",
            ),
            10: L(
                ist(2026, 9, 5, 11, 0),
                "Admin One",
                "export",
                "reports",
                "Salary Register - XLSX - 40 rows",
                None,
                "10.0.0.1",
            ),
            11: L(
                ist(2026, 9, 6, 9, 0), "Admin One", "update", "settings", "Updated company settings", None, "10.0.0.1"
            ),
            12: L(
                ist(2026, 9, 1, 0, 20), "Admin One", "login", "auth", "early boundary", None, "10.0.0.1"
            ),  # 31-Aug UTC
            13: L(
                ist(2026, 10, 1, 0, 20), "Admin One", "login", "auth", "late boundary", None, "10.0.0.1"
            ),  # 30-Sep UTC
            14: L(
                ist(2026, 9, 10, 10, 0),
                "HR Two",
                "create",
                "employees",
                "Created employee 502 -Tara S",
                b1,
                "10.0.0.2",
                7,
            ),
            15: L(ist(2026, 8, 15, 10, 0), "Admin One", "login", "auth", "August", None, "10.0.0.1"),
            16: L(
                ist(2026, 9, 11, 10, 0),
                "HR Two",
                "create",
                "employees",
                "Created employee 503 -Anna-Marie K",
                b1,
                "10.0.0.2",
                8,
            ),
        }
        cls.branch_user = make_user("ea_audit_b1", {"reports": "view", "user_management": "edit"}, cls.b1)

    # -- audit log ---------------------------------------------------------
    def test_audit_log_golden_rows_and_cards(self):
        body = self.run_report("audit-log", **WINDOW)
        self.assertEqual(len(body["rows"]), 14)
        self.assertEqual(
            [r["createdAt"] for r in body["rows"]][:3], ["2026-09-11 10:00", "2026-09-10 10:00", "2026-09-06 09:00"]
        )
        self.assertEqual(body["rows"][-1]["createdAt"], "2026-09-01 00:20")  # newest first
        r = next(x for x in body["rows"] if x["recordId"] == 5 and x["action"] == "Delete")
        self.assertEqual(
            (r["userName"], r["module"], r["description"], r["ipAddress"], r["branch"], r["userType"]),
            ("HR Two", "employees", "Deleted employee 501 -Ravi Kumar", "10.0.0.2", "Unit 1", "HR"),
        )
        self.assertIsNone(next(x for x in body["rows"] if x["description"] == "Failed login for: bob")["branch"])
        self.assertNotIn("late boundary", [r["description"] for r in body["rows"]])  # IST 1-Oct 00:20 = 30-Sep UTC
        self.assertIn("early boundary", [r["description"] for r in body["rows"]])  # IST 1-Sep 00:20 = 31-Aug UTC
        self.assertEqual(
            cards(body), {"Events": 14, "Distinct users": 3, "Failed / blocked logins": 2, "Deletes": 1, "Exports": 1}
        )

    def test_audit_log_filters(self):
        ids = lambda **p: sorted(r["description"] for r in self.run_report("audit-log", **WINDOW, **p)["rows"])  # noqa: E731
        self.assertEqual(len(ids(module="employees")), 8)
        self.assertEqual(len(ids(action="delete,export")), 2)
        self.assertEqual(len(ids(userName="hr two")), 6)
        self.assertEqual(ids(search="bob"), ["Failed login for: bob", "Locked-out login attempt for: bob"])
        self.assertEqual(len(ids(branchIds=str(self.b1.id))), 5)
        self.assertEqual(len(ids(branchIds=str(self.b2.id))), 1)
        self.assertEqual(len(ids(module="auth", action="login")), 2)
        self.assertEqual(self.run_report("audit-log", dateFrom="2026-08-01", dateTo="2026-08-31")["rowCount"], 1)

    def test_audit_log_row_limit_is_enforced_in_the_query(self):
        with mock.patch("api.reporting.runner.SCREEN_ROW_LIMIT", 5):
            body = self.run_report("audit-log", **WINDOW)
        self.assertEqual(body["rowCount"], 5)
        self.assertTrue(body["truncated"])
        self.assertEqual(cards(body)["Events"], 14)  # headline figure is over the full range, not the cut-off rows

    def test_date_range_is_capped(self):
        r = self.raw("audit-log", dateFrom="2026-01-01", dateTo="2026-09-30")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["field"], "dateTo")

    # -- employee change log -----------------------------------------------
    def test_employee_change_log_parses_descriptions(self):
        body = self.run_report("employee-change-log", **WINDOW)
        self.assertEqual(len(body["rows"]), 8)
        pick = lambda r: (r["action"], r["employeeCode"], r["employeeName"], r["fieldsChanged"])  # noqa: E731
        rows = {r["createdAt"]: r for r in body["rows"]}
        self.assertEqual(pick(rows["2026-09-11 10:00"]), ("Created", "503", "Anna-Marie K", None))
        self.assertEqual(pick(rows["2026-09-10 10:00"]), ("Created", "502", "Tara S", None))
        self.assertEqual(pick(rows["2026-09-05 10:00"]), ("Bulk change", None, None, None))
        self.assertEqual(pick(rows["2026-09-04 09:30"]), ("Bulk update", "601", None, "Phone, Salary Amount"))
        self.assertEqual(pick(rows["2026-09-04 09:00"]), ("Bulk import", "601", "Sita Devi", None))
        self.assertEqual(pick(rows["2026-09-03 12:00"]), ("Deleted", "501", "Ravi Kumar", None))
        self.assertEqual(pick(rows["2026-09-02 11:00"]), ("Updated", "501", "Ravi Kumar", None))
        self.assertEqual(pick(rows["2026-09-02 10:00"]), ("Created", "501", "Ravi Kumar", None))
        self.assertEqual(cards(body), {"Created": 3, "Updated": 1, "Deleted": 1, "Bulk imported": 1, "Bulk updated": 1})
        self.assertEqual(rows["2026-09-03 12:00"]["branch"], "Unit 1")

    def test_employee_change_log_filters(self):
        codes_of = lambda **p: [
            r["employeeCode"] for r in self.run_report("employee-change-log", **WINDOW, **p)["rows"]
        ]  # noqa: E731
        self.assertEqual(codes_of(action="delete"), ["501"])
        self.assertEqual(codes_of(employeeCode="501"), ["501", "501", "501"])
        self.assertEqual(codes_of(employeeCode="50"), [])  # exact code, not a prefix
        self.assertEqual(len(codes_of(userName="admin")), 2)
        self.assertEqual(sorted(codes_of(branchIds=str(self.b1.id))), ["501", "501", "501", "502", "503"])
        self.assertEqual(codes_of(branchIds=str(self.b2.id)), ["601"])

    # -- audit summary -----------------------------------------------------
    def test_audit_summary_by_user(self):
        body = self.run_report("audit-summary", **WINDOW)
        self.assertEqual(codes(body, "group"), ["Admin One", "HR Two", "system"])
        keys = ("total", "logins", "failedLogins", "creates", "updates", "deletes", "exports", "other")
        pick = lambda g: tuple(by_key(body, g, "group")[k] for k in keys)  # noqa: E731
        self.assertEqual(pick("Admin One"), (6, 2, 0, 0, 3, 0, 1, 0))
        self.assertEqual(pick("HR Two"), (6, 0, 0, 4, 1, 1, 0, 0))
        self.assertEqual(pick("system"), (2, 0, 2, 0, 0, 0, 0, 0))
        r = by_key(body, "Admin One", "group")
        self.assertEqual((r["firstAt"], r["lastAt"]), ("2026-09-01 00:20", "2026-09-06 09:00"))
        self.assertEqual(
            cards(body), {"Events": 14, "Most active user": "Admin One (6)", "Busiest day": "2026-09-01 (4)"}
        )
        self.assertEqual(body["totals"]["total"], 14)
        self.assert_totals_match(body, *keys)
        self.assertEqual(body["columns"][0]["label"], "User")

    def test_audit_summary_by_module_and_day(self):
        body = self.run_report("audit-summary", groupBy="module", **WINDOW)
        self.assertEqual(
            [(r["group"], r["total"]) for r in body["rows"]],
            [("employees", 8), ("auth", 4), ("reports", 1), ("settings", 1)],
        )
        self.assertEqual(by_key(body, "auth", "group")["failedLogins"], 2)
        self.assertEqual(body["columns"][0]["label"], "Module")
        body = self.run_report("audit-summary", groupBy="day", **WINDOW)
        self.assertEqual(
            [(r["group"], r["total"]) for r in body["rows"]],
            [
                ("2026-09-01", 4),
                ("2026-09-02", 2),
                ("2026-09-03", 1),
                ("2026-09-04", 2),
                ("2026-09-05", 2),
                ("2026-09-06", 1),
                ("2026-09-10", 1),
                ("2026-09-11", 1),
            ],
        )  # the 00:20 IST event is grouped on the 1st (IST), not on 31-Aug (UTC)
        self.assertEqual((body["columns"][0]["label"], body["columns"][0]["type"]), ("Day (IST)", "date"))
        body = self.run_report("audit-summary", branchIds=str(self.b1.id), **WINDOW)
        self.assertEqual(codes(body, "group"), ["HR Two"])
        self.assertEqual(cards(body)["Events"], 5)

    # -- access ------------------------------------------------------------
    def test_admin_reports_are_super_admin_only(self):
        for rid in ("audit-log", "employee-change-log", "audit-summary"):
            self.assertEqual(self.raw(rid, self.branch_user, **WINDOW).status_code, 404, rid)
            self.assertEqual(self.raw(rid, self.branch_user, fmt="xlsx", **WINDOW).status_code, 404, rid)
            ids = {x["id"] for x in self.client.get("/api/reports/catalog", **hdr(self.branch_user)).json()["reports"]}
            self.assertNotIn(rid, ids)
            ids = {x["id"] for x in self.client.get("/api/reports/catalog", **hdr(self.admin)).json()["reports"]}
            self.assertIn(rid, ids)

    def test_branch_isolation_holds_even_when_a_scope_is_pinned(self):
        # Unit 1 owns logs 4, 5, 6, 14 and 16 (all "HR Two"); Unit 2 owns log 7; the rest carry no branch.
        out = self.scoped_run("audit-log", self.b1, **WINDOW)
        self.assertEqual(len(out.rows), 5)
        self.assertEqual({r["userName"] for r in out.rows}, {"HR Two"})
        self.assertEqual({r["branch"] for r in out.rows}, {"Unit 1"})
        self.assertEqual({s["label"]: s["value"] for s in out.summary}["Events"], 5)
        # the Branch filter cannot widen the scope
        self.assertEqual(self.scoped_run("audit-log", self.b1, branchIds=str(self.b2.id), **WINDOW).rows, [])
        out = self.scoped_run("employee-change-log", self.b1, **WINDOW)
        self.assertEqual(sorted(r["employeeCode"] for r in out.rows), ["501", "501", "501", "502", "503"])
        self.assertEqual({s["label"]: s["value"] for s in out.summary}["Created"], 3)
        out = self.scoped_run("audit-summary", self.b1, **WINDOW)
        self.assertEqual([(r["group"], r["total"]) for r in out.rows], [("HR Two", 5)])
        out = self.scoped_run("audit-summary", self.b2, groupBy="module", **WINDOW)
        self.assertEqual([(r["group"], r["total"]) for r in out.rows], [("employees", 1)])


# ═══════════════════════════════════════════════════════════════════════════
#  User access (admin only)
# ═══════════════════════════════════════════════════════════════════════════


@override_settings(ADMIN_USERNAME="")
class AccessReportTests(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.r_hr = Role.objects.create(
            name="HR Manager", permissions={"employees": "edit", "attendance": "view", "reports": "view"}
        )
        cls.r_pay = Role.objects.create(
            name="Payroll Officer",
            permissions={
                "payroll": "edit",
                "salary_slip": "view",
                "employees.departments": "view",
                "bogus_key": "edit",
            },
        )
        cls.r_legacy = Role.objects.create(
            name="Legacy", permissions={"reports": {"view": True}, "users": {"read": True}}
        )

        def user(name, role=None, branch=None, login=None, **kw):
            u = HRUser.objects.create(
                username=name,
                password_hash=SECRET,
                role=role,
                branch=branch,
                full_name=name.title(),
                email=f"{name}@x.in",
                master_features={"co": True},
                **kw,
            )
            if login:
                HRUser.objects.filter(pk=u.pk).update(last_login=login)
            return u

        HRUser.objects.filter(pk=cls.admin.pk).update(last_login=datetime(2026, 9, 19, 6, 30, tzinfo=UTC))
        cls.hrm1 = user("hrm1", cls.r_hr, cls.b1, ist(2026, 9, 15, 10))
        cls.pay1 = user("pay1", cls.r_pay, None, None)
        cls.old1 = user("old1", cls.r_hr, cls.b2, datetime(2026, 7, 1, 6, 30, tzinfo=UTC))
        cls.dis1 = user("dis1", cls.r_pay, None, ist(2026, 6, 1, 10), is_active=False)
        cls.norole = user("norole", None, cls.b1, datetime(2026, 9, 18, 6, 30, tzinfo=UTC))
        cls.ghost = user("ghost", cls.r_hr, None, ist(2026, 9, 19, 10), is_hidden=True)

        def session(u, created, seen, revoked=None, ip="10.1.1.1"):
            s = LoginSession.objects.create(
                hr_user=u,
                jti=uuid.uuid4().hex,
                device_label="Chrome on Windows",
                user_agent="UA-SECRET-STRING",
                ip_address=ip,
            )
            LoginSession.objects.filter(pk=s.pk).update(created_at=created, last_seen_at=seen, revoked_at=revoked)
            return s

        cls.s1 = session(cls.hrm1, datetime(2026, 9, 20, 3, 0, tzinfo=UTC), datetime(2026, 9, 20, 5, 0, tzinfo=UTC))
        cls.s2 = session(
            cls.hrm1,
            datetime(2026, 9, 20, 2, 0, tzinfo=UTC),
            datetime(2026, 9, 20, 2, 30, tzinfo=UTC),
            datetime(2026, 9, 20, 3, 0, tzinfo=UTC),
        )
        cls.s3 = session(cls.hrm1, datetime(2026, 9, 19, 6, 0, tzinfo=UTC), datetime(2026, 9, 19, 8, 0, tzinfo=UTC))
        cls.s4 = session(
            cls.admin, datetime(2026, 9, 20, 5, 0, tzinfo=UTC), datetime(2026, 9, 20, 5, 45, tzinfo=UTC), ip="10.1.1.2"
        )
        cls.s5 = session(cls.ghost, datetime(2026, 9, 20, 4, 0, tzinfo=UTC), datetime(2026, 9, 20, 4, 10, tzinfo=UTC))
        cls.s6 = session(cls.old1, datetime(2026, 9, 10, 10, 0, tzinfo=UTC), datetime(2026, 9, 10, 10, 20, tzinfo=UTC))
        cls.s7 = session(  # 1-Oct 00:30 IST: outside September (revoked so it never counts as live "now")
            cls.old1,
            datetime(2026, 9, 30, 19, 0, tzinfo=UTC),
            datetime(2026, 9, 30, 19, 5, tzinfo=UTC),
            datetime(2026, 9, 30, 19, 10, tzinfo=UTC),
        )
        cls.s8 = session(
            cls.old1, datetime(2026, 8, 31, 19, 0, tzinfo=UTC), datetime(2026, 8, 31, 19, 5, tzinfo=UTC)
        )  # 1-Sep IST

        def attempt(username, ok, when, ip="1.1.1.1"):
            a = HrLoginAttempt.objects.create(username=username, ip_address=ip, success=ok)
            HrLoginAttempt.objects.filter(pk=a.pk).update(created_at=when)
            return a

        F_, S_ = False, True
        for i, name in enumerate(("bob", "Bob", "bob", "bob", "bob")):  # 5 failures in 8 minutes -> lockout on the 5th
            attempt(name, F_, ist(2026, 9, 10, 10, i * 2))
        attempt("bob", F_, ist(2026, 9, 10, 10, 30))  # 28 min after the 2nd failure: no new lockout
        attempt("bob", S_, ist(2026, 9, 10, 10, 35))
        attempt("carol", F_, ist(2026, 9, 10, 11, 0), "2.2.2.2")
        attempt("carol", F_, ist(2026, 9, 10, 11, 1), "2.2.2.2")
        attempt("carol", S_, ist(2026, 9, 10, 11, 2), "2.2.2.2")  # a success resets the run
        for i in range(5):
            attempt("carol", F_, ist(2026, 9, 10, 11, 3 + i), "2.2.2.2")  # five more failures -> lockout on the last
        for i in range(4):
            attempt("dave", F_, ist(2026, 9, 11, 9, i), "3.3.3.3")
        attempt("dave", F_, ist(2026, 9, 11, 16, 0), "3.3.3.4")  # 5th failure but hours later: no lockout
        attempt("erin", F_, ist(2026, 9, 1, 0, 10), "4.4.4.4")  # 31-Aug UTC, inside September in IST
        attempt("erin", F_, ist(2026, 10, 1, 0, 10), "4.4.4.4")  # 30-Sep UTC, outside
        for i in range(4):
            attempt("frank", F_, ist(2026, 8, 31, 23, 50 + i), "5.5.5.5")  # before the window: history only
        attempt("frank", F_, ist(2026, 9, 1, 0, 2), "5.5.5.5")  # completes a run that started before 1-Sep
        attempt("ghost", F_, ist(2026, 9, 12, 9, 0), "6.6.6.6")
        cls.b1_admin_like = make_user("ea_access_b1", {"reports": "view", "user_management": "edit"}, cls.b1)

    # -- hr users ----------------------------------------------------------
    ACTIVE = [
        "ea_access_b1",
        "ea_admin",
        "ea_reports_only",
        "hrm1",
        "norole",
        "old1",
        "pay1",
    ]  # ghost is hidden, dis1 disabled

    def test_hr_users_golden_rows(self):
        body = self.run_report("hr-users")
        self.assertEqual(codes(body, "username"), self.ACTIVE)
        by = {r["username"]: r for r in body["rows"]}
        r = by["ea_admin"]
        self.assertEqual(
            (r["role"], r["accessScope"], r["daysSinceLogin"], r["liveSessions"], r["flags"]),
            ("Super Admin", "Super Admin", 1, 1, None),
        )
        r = by["hrm1"]
        self.assertEqual(
            (r["role"], r["accessScope"], r["branch"], r["daysSinceLogin"], r["liveSessions"], r["flags"]),
            ("HR Manager", "Branch-limited", "Unit 1", 5, 1, None),
        )
        self.assertEqual((r["lastLogin"], r["fullName"], r["email"]), ("2026-09-15 10:00", "Hrm1", "hrm1@x.in"))
        r = by["pay1"]
        self.assertEqual(
            (r["role"], r["accessScope"], r["daysSinceLogin"], r["lastLogin"], r["flags"]),
            ("Payroll Officer", "Company-wide", None, None, "Company-wide access; Never logged in"),
        )
        r = by["old1"]
        self.assertEqual((r["daysSinceLogin"], r["flags"], r["branch"]), (81, "Dormant (30+ days)", "Unit 2"))
        r = by["norole"]
        self.assertEqual(
            (r["role"], r["accessScope"], r["flags"], r["daysSinceLogin"]),
            ("No role (no access)", "Branch-limited", "No role (no access)", 2),
        )
        self.assertEqual(by["hrm1"]["isActive"], "Active")
        self.assertEqual(body["totals"]["liveSessions"], 2)
        self.assertEqual(
            cards(body),
            {
                "Accounts": 7,
                "Active": 7,
                "Disabled": 0,
                "Super admins": 1,
                "Company-wide (non-admin)": 2,
                "Dormant / never logged in": 4,  # ea_access_b1, ea_reports_only (never), old1 (81 days), pay1 (never)
            },
        )

    def test_hr_users_filters(self):
        allb = self.run_report("hr-users", state="all")
        self.assertEqual(codes(allb, "username"), sorted([*self.ACTIVE, "dis1"]))
        self.assertEqual(by_key(allb, "dis1", "username")["isActive"], "Disabled")
        self.assertEqual((cards(allb)["Accounts"], cards(allb)["Disabled"]), (8, 1))
        self.assertEqual(cards(allb)["Dormant / never logged in"], 4)  # a disabled account is never "dormant"
        self.assertEqual(codes(self.run_report("hr-users", state="inactive"), "username"), ["dis1"])
        self.assertEqual(codes(self.run_report("hr-users", role="payroll", state="all"), "username"), ["dis1", "pay1"])
        self.assertEqual(
            codes(self.run_report("hr-users", branchIds=str(self.b1.id)), "username"),
            ["ea_access_b1", "hrm1", "norole"],
        )
        dormant = self.run_report("hr-users", dormantOnly="true")
        self.assertEqual(codes(dormant, "username"), ["ea_access_b1", "ea_reports_only", "old1", "pay1"])
        late = self.run_report("hr-users", dormantOnly="true", dormantDays="90")
        self.assertEqual(
            codes(late, "username"), ["ea_access_b1", "ea_reports_only", "pay1"]
        )  # 81 days is fine under 90
        self.assertEqual(by_key(late, "pay1", "username")["flags"], "Company-wide access; Never logged in")

    def test_hr_users_hidden_accounts_only_for_the_master_admin(self):
        self.assertNotIn("ghost", codes(self.run_report("hr-users", state="all"), "username"))
        with override_settings(ADMIN_USERNAME="ea_admin"):
            body = self.run_report("hr-users", state="all")
        self.assertIn("Hidden account", by_key(body, "ghost", "username")["flags"])
        with override_settings(ADMIN_USERNAME="somebody_else"):
            self.assertNotIn("ghost", codes(self.run_report("hr-users", state="all"), "username"))

    def test_hr_users_never_leak_secrets(self):
        for fmt in ("xlsx", "pdf"):
            content = self.raw("hr-users", fmt=fmt, state="all").content
            self.assertNotIn(b"SECRETHASH", content)
        text = self.raw("hr-users", state="all").content.decode()
        for needle in ("SECRETHASH", "master_features", "passwordHash", "password_hash"):
            self.assertNotIn(needle, text)

    # -- role access matrix ------------------------------------------------
    def test_role_matrix_levels_are_the_effective_ones(self):
        body = self.run_report("role-access-matrix")
        labels = [c["label"] for c in body["columns"]]
        self.assertEqual(labels[:3], ["Module", "Key", "Group"])
        role_cols = {c["label"]: c["key"] for c in body["columns"][3:]}
        self.assertTrue({"HR Manager", "Payroll Officer", "Legacy"} <= set(role_cols))
        lvl = lambda key, role: by_key(body, key, "moduleKey")[role_cols[role]]  # noqa: E731
        self.assertEqual(lvl("employees", "HR Manager"), "Edit")
        self.assertEqual(lvl("employees.designations", "HR Manager"), "Edit")  # inherited from the parent
        self.assertEqual(lvl("employees.departments", "Payroll Officer"), "View")  # explicit child entry
        self.assertEqual(lvl("employees.designations", "Payroll Officer"), "Hidden")
        self.assertEqual(lvl("attendance", "HR Manager"), "View")
        self.assertEqual(lvl("payroll", "Payroll Officer"), "Edit")
        self.assertEqual(lvl("salary_slip", "Payroll Officer"), "View")
        self.assertEqual(lvl("dashboard", "HR Manager"), "Hidden")  # nothing set = hidden (fail closed)
        self.assertEqual(lvl("reports", "HR Manager"), "View")
        self.assertEqual(lvl("reports", "Legacy"), "Hidden")  # legacy dict-shaped permission is not a level
        self.assertEqual(len(body["rows"]), len(all_module_keys()))
        self.assertEqual(by_key(body, "employees.departments", "moduleKey")["module"], "Employees > Departments")
        self.assertEqual(by_key(body, "employees.departments", "moduleKey")["group"], "Employees")
        self.assertEqual(
            [r["moduleKey"] for r in body["rows"]][:3], ["dashboard", "employees", "employees.departments"]
        )
        self.assertEqual(
            cards(body)["Modules"], len(MODULE_TREE) + sum(len(n.get("children", [])) for n in MODULE_TREE)
        )

    def test_role_matrix_notes_filters_and_counts(self):
        body = self.run_report("role-access-matrix")
        text = " ".join(body["notes"])
        self.assertIn("HR Manager 2", text)  # hrm1 + old1; the hidden ghost is not counted
        self.assertIn("Payroll Officer 1", text)  # dis1 is disabled
        self.assertIn("Legacy 0", text)
        self.assertIn("Payroll Officer: bogus_key", text)  # a key that is not a module
        self.assertIn("Legacy: users", text)
        self.assertIn("Legacy: reports", text)  # a known module whose value is not hidden / view / edit
        self.assertTrue(any("bypass" in n for n in body["notes"]))
        self.assertEqual(cards(body)["Active users with a role"], 2 + 1 + 0 + 1 + 1)  # + the two helper accounts
        only = self.run_report("role-access-matrix", role="payroll")
        self.assertEqual([c["label"] for c in only["columns"][3:]], ["Payroll Officer"])
        self.assertEqual((cards(only)["Roles"], cards(only)["Active users with a role"]), (1, 1))
        self.assertEqual(self.run_report("role-access-matrix", role="Legacy", hideNoAccess="true")["rows"], [])
        trimmed = self.run_report("role-access-matrix", role="Payroll", hideNoAccess="true")
        keys = {r["moduleKey"] for r in trimmed["rows"]}
        self.assertEqual(keys, {"employees.departments", "payroll", "salary_slip"})
        self.assertNotIn("dashboard", keys)

    def test_role_matrix_with_many_roles_still_exports(self):
        for i in range(25):
            Role.objects.create(name=f"Extra role {i:02d}", permissions={"employees": "view"})
        body = self.run_report("role-access-matrix")
        self.assertGreater(len(body["columns"]), 28)
        ws = self.sheet("role-access-matrix")
        self.assertEqual([c.value for c in ws[7]], [c["label"] for c in body["columns"]])
        self.assertTrue(self.raw("role-access-matrix", fmt="pdf").content.startswith(b"%PDF"))

    # -- login sessions ----------------------------------------------------
    def test_login_sessions_golden_rows(self):
        body = self.run_report("login-sessions", **WINDOW)
        self.assertEqual(len(body["rows"]), 6)
        order = [(r["username"], r["signedInAt"]) for r in body["rows"]]
        self.assertEqual(
            order,
            [
                ("ea_admin", "2026-09-20 10:30"),
                ("hrm1", "2026-09-20 08:30"),
                ("hrm1", "2026-09-20 07:30"),
                ("hrm1", "2026-09-19 11:30"),
                ("old1", "2026-09-10 15:30"),
                ("old1", "2026-09-01 00:30"),
            ],
        )
        r = body["rows"][1]
        self.assertEqual(
            (r["state"], r["duration"], r["deviceLabel"], r["ipAddress"], r["role"], r["signedOutAt"]),
            ("Live", 120, "Chrome on Windows", "10.1.1.1", "HR Manager", None),
        )
        self.assertEqual(
            (body["rows"][2]["state"], body["rows"][2]["duration"], body["rows"][2]["signedOutAt"]),
            ("Revoked", 30, "2026-09-20 08:30"),
        )
        self.assertEqual((body["rows"][3]["state"], body["rows"][3]["duration"]), ("Expired", 120))
        self.assertEqual(
            (body["rows"][0]["state"], body["rows"][0]["duration"], body["rows"][0]["role"]),
            ("Live", 45, "Super Admin"),
        )
        self.assertEqual(cards(body), {"Sessions": 6, "Distinct users": 3, "Live now": 2, "Revoked / signed out": 1})
        self.assertNotIn("ghost", codes(body, "username"))  # hidden account's session

    def test_login_sessions_filters(self):
        st = lambda **p: [r["state"] for r in self.run_report("login-sessions", **WINDOW, **p)["rows"]]  # noqa: E731
        self.assertEqual(st(state="live"), ["Live", "Live"])
        self.assertEqual(st(state="revoked"), ["Revoked"])
        self.assertEqual(st(state="expired"), ["Expired", "Expired", "Expired"])
        self.assertEqual(len(st(username="hrm")), 3)
        self.assertEqual(
            codes(self.run_report("login-sessions", branchIds=str(self.b2.id), **WINDOW), "username"), ["old1", "old1"]
        )
        self.assertEqual(
            cards(self.run_report("login-sessions", state="live", **WINDOW))["Sessions"], 6
        )  # cards ignore State

    def test_login_sessions_never_leak_tokens_or_user_agents(self):
        for fmt in ("xlsx", "pdf"):
            content = self.raw("login-sessions", fmt=fmt, **WINDOW).content
            self.assertNotIn(b"UA-SECRET-STRING", content)
            self.assertNotIn(self.s1.jti.encode(), content)
        text = self.raw("login-sessions", **WINDOW).content.decode()
        self.assertNotIn("UA-SECRET-STRING", text)
        self.assertNotIn(self.s1.jti, text)

    # -- login attempts ----------------------------------------------------
    def test_login_attempts_golden_numbers(self):
        body = self.run_report("login-attempts", **WINDOW)
        self.assertEqual(len(body["rows"]), 22)
        self.assertEqual(
            cards(body),
            {
                "Attempts": 22,
                "Failures": 20,
                "Failing usernames": 5,
                "Failing IP addresses": 6,
                "Lockouts": 3,
            },
        )
        locked = [(r["username"], r["attemptedAt"]) for r in body["rows"] if r["lockedOut"]]
        self.assertEqual(
            sorted(locked),
            sorted(
                [
                    ("bob", "2026-09-10 10:08"),
                    ("carol", "2026-09-10 11:07"),
                    ("frank", "2026-09-01 00:02"),
                ]
            ),
        )
        self.assertTrue(all(r["lockedOut"] == "Lockout triggered" for r in body["rows"] if r["lockedOut"]))
        self.assertIn("2026-09-01 00:10", codes(body, "attemptedAt"))  # erin: IST 1-Sep 00:10 = 31-Aug UTC
        self.assertNotIn("2026-10-01 00:10", codes(body, "attemptedAt"))
        outcomes = [r["success"] for r in body["rows"]]
        self.assertEqual((outcomes.count("Success"), outcomes.count("Failed")), (2, 20))

    def test_login_attempts_lockout_needs_history_before_the_window(self):
        body = self.run_report("login-attempts", username="frank", **WINDOW)
        self.assertEqual(
            [(r["attemptedAt"], r["lockedOut"]) for r in body["rows"]], [("2026-09-01 00:02", "Lockout triggered")]
        )

    def test_login_attempts_filters(self):
        rows = lambda **p: self.run_report("login-attempts", **WINDOW, **p)["rows"]  # noqa: E731
        self.assertEqual(len(rows(outcome="failed")), 20)
        self.assertEqual(len(rows(outcome="success")), 2)
        self.assertEqual(len(rows(username="bob")), 7)  # "Bob" typed differently is the same username
        failed_bob = rows(username="bob", outcome="failed")
        self.assertEqual(len(failed_bob), 6)
        self.assertEqual(
            sum(1 for r in failed_bob if r["lockedOut"]), 1
        )  # lock-out still seen when only failures are listed
        self.assertEqual(len(rows(ip="3.3.3")), 5)
        self.assertEqual(len(rows(ip="1.1.1.1")), 7)

    def test_login_attempts_hidden_accounts_and_window_cap(self):
        self.assertNotIn("ghost", codes(self.run_report("login-attempts", **WINDOW), "username"))
        with override_settings(ADMIN_USERNAME="ea_admin"):
            self.assertIn("ghost", codes(self.run_report("login-attempts", **WINDOW), "username"))
        r = self.raw("login-attempts", dateFrom="2026-08-01", dateTo="2026-09-30")
        self.assertEqual(r.status_code, 400)

    def test_branch_isolation_holds_even_when_a_scope_is_pinned(self):
        out = self.scoped_run("hr-users", self.b1, state="all")
        self.assertEqual([r["username"] for r in out.rows], ["ea_access_b1", "hrm1", "norole"])
        self.assertEqual(self.scoped_run("hr-users", self.b1, branchIds=str(self.b2.id), state="all").rows, [])
        self.assertEqual([r["username"] for r in self.scoped_run("hr-users", self.b2, state="all").rows], ["old1"])
        out = self.scoped_run("login-sessions", self.b1, **WINDOW)
        self.assertEqual([r["username"] for r in out.rows], ["hrm1", "hrm1", "hrm1"])
        self.assertEqual({s["label"]: s["value"] for s in out.summary}["Sessions"], 3)
        out = self.scoped_run("login-sessions", self.b2, **WINDOW)
        self.assertEqual([r["username"] for r in out.rows], ["old1", "old1"])

    def test_admin_access_reports_are_super_admin_only(self):
        for rid in ("hr-users", "role-access-matrix", "login-sessions", "login-attempts"):
            self.assertEqual(self.raw(rid, self.b1_admin_like, **WINDOW).status_code, 404, rid)
            self.assertEqual(self.raw(rid, self.b1_admin_like, fmt="pdf", **WINDOW).status_code, 404, rid)
            self.assertEqual(self.raw(rid, self.reports_only, **WINDOW).status_code, 404, rid)


# ═══════════════════════════════════════════════════════════════════════════
#  Cross-cutting contract: registration, permissions, exports, empty results, N+1
# ═══════════════════════════════════════════════════════════════════════════

_seq = [0]


def seed_world(n, b1, b2):
    """n units of every kind of row these reports read. Different departments, HODs, users and days per unit so
    that any per-row query shows up as growth in the query count."""
    for i in range(n):
        _seq[0] += 1
        t = _seq[0]
        branch = b1 if t % 2 else b2
        dept = Department.objects.create(name=f"DEPT{t}", branch=branch)
        desig = Designation.objects.create(title=f"Role{t}", department=dept)
        emp = make_emp(
            f"Z{t}",
            f"Emp{t}",
            dept,
            branch,
            etype="staff" if t % 3 else "production",
            designation=desig,
            password_hash=SECRET,
            last_mobile_login_at=ist(2026, 9, 1 + t % 20, 9),
        )
        PushToken.objects.create(employee=emp, token=f"tok-{t}")
        add_doc(emp, "pan_card", ist(2026, 9, 1 + t % 25, 10), f"HR {t % 3}")
        hod_emp = make_emp(f"ZH{t}", f"Hod{t}", dept, branch)
        other_emp = make_emp(f"ZO{t}", f"Oth{t}", dept, branch)
        m = DepartmentManager.objects.create(employee=hod_emp)
        m2 = DepartmentManager.objects.create(employee=other_emp)
        ManagerDepartmentAssignment.objects.create(manager=m, department=dept)
        ManagerDepartmentAssignment.objects.create(manager=m2, department=dept)  # two holders: a conflict per unit
        ManagerEmployeeAssignment.objects.create(manager=m2, employee=emp)
        DepartmentHeadcount.objects.create(department=dept, required_count=t)
        job = Job.objects.create(title=f"Job{t}", department=dept)
        Job.objects.filter(pk=job.pk).update(created_at=ist(2026, 9, 1 + t % 25, 10))
        app = Applicant.objects.create(job=job, name=f"App{t}", email="a@x.in", phone="9", status="applied")
        Applicant.objects.filter(pk=app.pk).update(created_at=ist(2026, 9, 1 + t % 25, 10))
        rs = HiringRuleSet.objects.create(name=f"RS{t}", department=dept, required_skills=["a", "b"])
        c = ScreeningCandidate.objects.create(
            rule_set=rs,
            department=dept,
            resume_file="resumes/x.pdf",
            original_filename="x.pdf",
            candidate_name=f"Cand{t}",
            status="selected",
            match_score=50 + t,
            interview_datetime=ist(2026, 9, 1 + t % 25, 10),
        )
        ScreeningCandidate.objects.filter(pk=c.pk).update(created_at=ist(2026, 9, 1 + t % 25, 10))
        role = Role.objects.create(name=f"Role {t}", permissions={"employees": "view"})
        user = HRUser.objects.create(username=f"user{t}", password_hash=SECRET, role=role, branch=branch)
        s = LoginSession.objects.create(hr_user=user, jti=uuid.uuid4().hex, device_label="Chrome", ip_address="1.2.3.4")
        LoginSession.objects.filter(pk=s.pk).update(created_at=ist(2026, 9, 1 + t % 25, 10))
        for k, (action, module, desc) in enumerate(
            (
                ("create", "employees", f"Created employee {t} -Name {t}"),
                ("login", "auth", "in"),
                ("export", "reports", "x"),
            )
        ):
            add_log(ist(2026, 9, 1 + (t + k) % 25, 10, k), f"User{t}", action, module, desc, branch)
        a = HrLoginAttempt.objects.create(username=f"user{t}", ip_address=f"9.9.9.{t % 200}", success=bool(t % 2))
        HrLoginAttempt.objects.filter(pk=a.pk).update(created_at=ist(2026, 9, 1 + t % 25, 10))


class RegistrationTests(TestCase):
    def test_every_group_report_is_registered_and_the_module_loads(self):
        registry.all_specs()
        self.assertFalse([m for m in registry.LOAD_ERRORS if "employees_admin" in m], registry.LOAD_ERRORS)
        for rid in ALL_IDS:
            self.assertIsNotNone(registry.get_spec(rid), rid)

    def test_metadata_is_sound(self):
        for rid in ALL_IDS:
            s = registry.get_spec(rid)
            self.assertEqual(s.category, "admin" if rid in ADMIN_IDS else "employees", rid)
            self.assertEqual(s.super_admin_only, rid in ADMIN_IDS, rid)
            self.assertTrue(s.description.endswith("."), rid)
            if rid in NON_ADMIN_IDS:
                self.assertTrue(s.modules, rid)
            for m in s.modules:
                self.assertIn(m, all_module_keys(), f"{rid}: {m}")
            self.assertTrue(s.icon, rid)
        fam = {}
        for rid in ALL_IDS:
            s = registry.get_spec(rid)
            if s.family:
                fam.setdefault(s.family, []).append((rid, s.variant, s.category))
        for name, members in fam.items():
            self.assertGreaterEqual(len(members), 2, name)
            self.assertEqual(len({m[2] for m in members}), 1, name)
            self.assertEqual(len({m[1] for m in members}), len(members), name)

    def test_hod_reports_use_the_user_management_module_and_the_documents_ones_the_documents_module(self):
        for rid in ("hod-directory", "hod-mapping", "hod-conflicts"):
            self.assertEqual(registry.get_spec(rid).modules, ("user_management",))
        for rid in ("document-compliance", "document-upload-log"):
            self.assertEqual(registry.get_spec(rid).modules, ("recruitment.documents",))
        self.assertEqual(registry.get_spec("mobile-app-access").modules, ("mobile_app_login",))


class ContractTests(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        seed_world(3, cls.b1, cls.b2)

    PARAMS = {
        "document-compliance": {"state": "all"},
        "hod-directory": {"state": "all"},
        "hr-users": {"state": "all"},
        "job-openings": {"status": "all"},
        "hiring-criteria": {"active": "all"},
    }
    EMPTY = {
        "document-compliance": {"employeeIds": "999999"},
        "hod-mapping": {"employeeIds": "999999"},
        "hod-conflicts": {"employeeIds": "999999"},
        "mobile-app-access": {"employeeIds": "999999"},
        "hod-directory": {"departmentIds": "999999"},
        "manpower-requirement": {"departmentIds": "999999"},
        "job-openings": {"departmentIds": "999999"},
        "interview-schedule": {"departmentIds": "999999"},
        "hiring-criteria": {"departmentIds": "999999"},
        "document-upload-log": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "applicant-register": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "screening-pipeline": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "recruitment-funnel": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "audit-log": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "employee-change-log": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "audit-summary": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "login-sessions": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "login-attempts": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        "hr-users": {"role": "no-such-role"},
        "role-access-matrix": {"role": "no-such-role", "hideNoAccess": "true"},
    }

    def params(self, rid):
        return {**WINDOW, **self.PARAMS.get(rid, {})}

    def test_role_with_only_the_reports_module_is_forbidden_not_hidden(self):
        for rid in NON_ADMIN_IDS:
            for fmt in (None, "xlsx", "pdf"):
                r = self.raw(rid, self.reports_only, fmt=fmt, **self.params(rid))
                self.assertEqual(r.status_code, 403, f"{rid} {fmt}")
                self.assertEqual(r.json()["error"], "report_forbidden", rid)

    def test_owning_module_grant_opens_the_report(self):
        for rid in NON_ADMIN_IDS:
            module = registry.get_spec(rid).modules[0]
            user = make_user(f"grant_{rid}", {"reports": "view", module: "view"})
            self.assertEqual(self.raw(rid, user, **self.params(rid)).status_code, 200, rid)
            self.assertEqual(self.raw(rid, user, fmt="pdf", **self.params(rid)).status_code, 200, rid)

    def test_parent_module_grant_cascades_to_recruitment_children(self):
        user = make_user("grant_parent", {"reports": "view", "recruitment": "view"})
        for rid in (
            "screening-pipeline",
            "interview-schedule",
            "hiring-criteria",
            "manpower-requirement",
            "document-compliance",
        ):
            self.assertEqual(self.raw(rid, user, **self.params(rid)).status_code, 200, rid)

    def test_admin_reports_are_invisible_to_everybody_but_a_super_admin(self):
        user = make_user("everything_but_super", {"reports": "edit", "user_management": "edit", "employees": "edit"})
        for rid in ADMIN_IDS:
            self.assertEqual(self.raw(rid, user, **self.params(rid)).status_code, 404, rid)
            self.assertEqual(self.raw(rid, user, fmt="xlsx", **self.params(rid)).status_code, 404, rid)
            self.assertEqual(self.raw(rid, self.admin, **self.params(rid)).status_code, 200, rid)

    def test_xlsx_and_pdf_exports_have_valid_signatures_and_round_trip(self):
        for rid in ALL_IDS:
            with self.subTest(rid):
                body = self.run_report(rid, **self.params(rid))
                pdf = self.raw(rid, fmt="pdf", **self.params(rid))
                self.assertEqual(pdf.status_code, 200)
                self.assertTrue(pdf.content.startswith(b"%PDF"))
                x = self.raw(rid, fmt="xlsx", **self.params(rid))
                self.assertEqual(x.status_code, 200)
                self.assertTrue(x.content.startswith(b"PK"))
                ws = load_workbook(io.BytesIO(x.content)).active
                self.assertEqual([c.value for c in ws[7]], [c["label"] for c in body["columns"]])
                self.assertGreater(body["rowCount"], 0)
                first = body["rows"][0]
                for idx, col in enumerate(body["columns"], start=1):
                    if col["type"] in ("text", "badge") and first.get(col["key"]) not in (None, ""):
                        self.assertEqual(ws.cell(row=8, column=idx).value, str(first[col["key"]]), col["key"])
                        break
                else:
                    self.fail("no text column with a value to compare")

    def test_empty_results_work_on_screen_and_in_both_exports(self):
        for rid in ALL_IDS:
            with self.subTest(rid):
                params = {**WINDOW, **self.EMPTY[rid]}
                body = self.run_report(rid, **params)
                self.assertEqual(body["rows"], [], rid)
                self.assertEqual(body["rowCount"], 0)
                self.assertTrue(self.raw(rid, fmt="pdf", **params).content.startswith(b"%PDF"))
                r = self.raw(rid, fmt="xlsx", **params)
                self.assertEqual(r.status_code, 200)
                ws = load_workbook(io.BytesIO(r.content)).active
                self.assertEqual([c.value for c in ws[7]], [c["label"] for c in body["columns"]])

    def test_query_count_does_not_grow_with_the_number_of_rows(self):
        def measure(rid):
            with CaptureQueriesContext(connection) as ctx:
                r = self.raw(rid, **self.params(rid))
            self.assertEqual(r.status_code, 200, f"{rid}: {r.content[:300]}")
            return len(ctx), r.json()

        small = {}
        for rid in ALL_IDS:
            measure(rid)  # warm-up (content types, settings singletons ...)
            small[rid] = measure(rid)
        seed_world(12, self.b1, self.b2)  # 3 -> 15 of everything
        for rid in ALL_IDS:
            n_small, b_small = small[rid]
            n_big, b_big = measure(rid)
            with self.subTest(rid):
                self.assertLessEqual(abs(n_big - n_small), 3, f"{rid}: {n_small} -> {n_big} queries")
                if rid == "role-access-matrix":
                    self.assertGreater(len(b_big["columns"]), len(b_small["columns"]))
                else:
                    self.assertGreater(b_big["rowCount"], b_small["rowCount"], rid)

    def test_no_report_query_loads_photos_or_employee_password_hashes(self):
        # Employee.photo_url is often a ~43 KB base64 data URI and password_hash is a secret: registers never load them.
        with CaptureQueriesContext(connection) as captured:
            for rid in ALL_IDS:
                for fmt in (None, "xlsx"):
                    self.raw(rid, fmt=fmt, **self.params(rid))
        sqls = [q["sql"] for q in captured.captured_queries]
        self.assertTrue(sqls)
        self.assertEqual([q[:100] for q in sqls if "photo_url" in q], [])
        bare_hash = re.compile(r'"employees"\."password_hash"(?!\s+(IS|=))')
        for q in sqls:
            select_list = q[: q.upper().find(" FROM ")] if " FROM " in q.upper() else q
            self.assertIsNone(bare_hash.search(select_list), q[:160])

    def test_every_report_runs_with_no_filters_at_all(self):
        for rid in ALL_IDS:
            with self.subTest(rid):
                self.assertEqual(self.raw(rid).status_code, 200)
                self.assertEqual(self.raw(rid, fmt="xlsx").status_code, 200)
                self.assertEqual(self.raw(rid, fmt="pdf").status_code, 200)

    def test_catalog_lists_them_with_filters_and_families_for_a_super_admin(self):
        body = self.client.get("/api/reports/catalog", **hdr(self.admin)).json()
        listed = {r["id"]: r for r in body["reports"]}
        for rid in ALL_IDS:
            self.assertIn(rid, listed)
            self.assertTrue(listed[rid]["filters"], rid)
        self.assertEqual(listed["role-access-matrix"]["dynamicColumns"], True)
        self.assertEqual(listed["hod-directory"]["family"], "hod")
        self.assertEqual({r["category"] for r in body["reports"] if r["id"] in ADMIN_IDS}, {"admin"})
        # a branch-scoped user is never offered the Branch filter, and never sees the admin reports
        scoped = make_user("catalog_b1", {"reports": "view", "user_management": "view"}, self.b1)
        mine = {r["id"]: r for r in self.client.get("/api/reports/catalog", **hdr(scoped)).json()["reports"]}
        self.assertFalse(set(ADMIN_IDS) & set(mine))
        self.assertIn("hod-directory", mine)
        self.assertNotIn("branch", [f["key"] for f in mine["hod-directory"]["filters"]])

    def test_a_get_never_writes(self):
        def counts():
            # every table of the app; the audit trail is the one legitimate writer (exports are logged)
            return {m.__name__: m.objects.count() for m in apps.get_app_config("api").get_models() if m is not AuditLog}

        # The export letterhead (framework branding.company) lazily creates the PayrollSettings singleton on the very
        # first export of a fresh database; that is framework behaviour, so create it up front and guard the reports.
        PayrollSettings.get()
        before = counts()
        for rid in ALL_IDS:
            for fmt in (None, "xlsx", "pdf"):
                self.assertEqual(self.raw(rid, fmt=fmt, **self.params(rid)).status_code, 200, (rid, fmt))
        self.assertEqual(counts(), before)
        # ... and every export writes exactly one audit row
        n = AuditLog.objects.count()
        self.raw("hod-directory", fmt="xlsx")
        self.assertEqual(AuditLog.objects.count(), n + 1)


class EmptyDatabaseTests(_Base):
    """A fresh installation: nothing but a super admin. Every report must still answer, and export."""

    def test_every_report_works_with_no_data_at_all(self):
        for rid in ALL_IDS:
            with self.subTest(rid):
                body = self.run_report(rid)
                self.assertIsInstance(body["rows"], list)
                self.assertTrue(self.raw(rid, fmt="pdf").content.startswith(b"%PDF"))
                self.assertTrue(self.raw(rid, fmt="xlsx").content.startswith(b"PK"))


# ═══════════════════════════════════════════════════════════════════════════
#  Adversarial review (independent reviewer): each test states the behaviour a
#  garments-company HR manager would rely on. A failing test here is a proven bug.
# ═══════════════════════════════════════════════════════════════════════════


class AdversarialReviewTests(_Base):
    def test_resume_purged_by_the_retention_job_is_reported_removed(self):
        """The report promises 'Removed' = the resume file no longer exists (rejected, or purged after the
        retention window). The purge job deletes the stored file but never clears the FileField value, so the
        column still holds a path and the report keeps saying 'On file' for a resume that is gone."""
        from .screening_cleanup import purge_expired_screening_documents

        dept = Department.objects.create(name="ADV-PURGE", branch=self.b1)
        rule = HiringRuleSet.objects.create(name="RS-adv-purge", department=dept)
        cand = ScreeningCandidate.objects.create(
            rule_set=rule,
            department=dept,
            resume_file=ContentFile(b"%PDF-1.4 resume", name="adv-old.pdf"),
            original_filename="adv-old.pdf",
            candidate_name="Old Candidate",
            status="screened",
        )
        stored_name = cand.resume_file.name
        self.assertTrue(default_storage.exists(stored_name))
        ScreeningCandidate.objects.filter(pk=cand.pk).update(created_at=NOW_UTC - timedelta(days=30))

        with mock.patch("django.utils.timezone.now", return_value=NOW_UTC):
            self.assertEqual(purge_expired_screening_documents()["purged"], 1)
        self.assertFalse(default_storage.exists(stored_name), "the purge job really removed the stored file")

        body = self.run_report("screening-pipeline", dateFrom="2026-08-01", dateTo="2026-09-30")
        row = by_key(body, "Old Candidate", "candidateName")
        self.assertEqual(row["resumeOnFile"], "Removed")

    def test_audit_log_module_filter_offers_every_module_the_application_writes(self):
        """A super admin must be able to filter the Activity Log by any module the app actually logs. Staff and
        production payroll generation is logged under module 'payroll' but the Module filter does not offer it."""
        root = Path(__file__).resolve().parent
        written = set()
        pattern = re.compile(r"""(?:log_action|_log)\(\s*request,\s*["'][a-z_]+["'],\s*["']([a-z_]+)["']""")
        for path in root.glob("*.py"):
            if path.name.startswith("tests"):
                continue
            written |= set(pattern.findall(path.read_text(encoding="utf-8")))
        offered = {value for value, _ in registry.get_spec("audit-log").filter("module").options}
        self.assertTrue({"payroll", "employees", "auth"} <= written, written)
        self.assertEqual(sorted(written - offered), [])
        self.assertEqual(self.raw("audit-log", module="payroll", **WINDOW).status_code, 200)

    def test_login_attempts_outcome_filter_is_not_blind_to_rows_behind_a_flood(self):
        """A credential-stuffing flood must not hide an older successful sign-in. The query keeps only the newest
        row_limit+100 attempts BEFORE the outcome / IP / hidden-account filters are applied, so 'Successful' shows
        nothing (and no truncation notice appears) once enough failures are newer than the success."""
        from .reporting import runner

        boss = HrLoginAttempt.objects.create(username="boss", ip_address="7.7.7.7", success=True)
        HrLoginAttempt.objects.filter(pk=boss.pk).update(created_at=ist(2026, 9, 2, 9))
        HrLoginAttempt.objects.bulk_create(
            [HrLoginAttempt(username=f"bot{i}", ip_address="8.8.8.8", success=False) for i in range(200)]
        )
        HrLoginAttempt.objects.filter(username__startswith="bot").update(created_at=ist(2026, 9, 20, 9))
        with mock.patch.object(runner, "SCREEN_ROW_LIMIT", 50):
            body = self.run_report("login-attempts", outcome="success", **WINDOW)
        self.assertEqual([r["username"] for r in body["rows"]], ["boss"])
        self.assertEqual(cards(body)["Attempts"], 1)

    def test_hod_reports_agree_on_active_employees_without_a_hod(self):
        """'Active employees with no HOD' (directory card) and 'Without a HOD (excl. HODs)' (mapping card) answer
        the same question. An INACTIVE HOD approves nothing, so his own requests still have no HOD; the directory
        card wrongly treats him as 'a HOD themselves' and under-counts."""
        dept = Department.objects.create(name="ADV-HOD", branch=self.b1)
        ex_hod = make_emp("ADV1", "Retired", dept, self.b1)
        make_emp("ADV2", "Regular", dept, self.b1)
        DepartmentManager.objects.create(employee=ex_hod, is_active=False)
        mapping = cards(self.run_report("hod-mapping", branchIds=str(self.b1.id)))
        directory = cards(self.run_report("hod-directory", state="all", branchIds=str(self.b1.id)))
        self.assertEqual(mapping["Without a HOD (excl. HODs)"], 2)
        self.assertEqual(directory["Active employees with no HOD"], mapping["Without a HOD (excl. HODs)"])

    def test_sessions_of_a_disabled_account_are_not_reported_live(self):
        """The middleware rejects every request of a disabled HR account, so its 12-hour token is dead. The
        session report must not count it as 'Live' (a security reviewer reads that card as 'can be used now')."""
        user = HRUser.objects.create(username="adv_disabled", password_hash=SECRET, is_active=False)
        session = LoginSession.objects.create(hr_user=user, jti=uuid.uuid4().hex, device_label="Chrome on Windows")
        LoginSession.objects.filter(pk=session.pk).update(
            created_at=NOW_UTC - timedelta(hours=1), last_seen_at=NOW_UTC - timedelta(minutes=5)
        )
        body = self.run_report("login-sessions", **WINDOW)
        row = by_key(body, "adv_disabled", "username")
        self.assertNotEqual(row["state"], "Live")
        self.assertEqual(cards(body)["Live now"], 0)

    # -- follow-ups on the fixes above ---------------------------------------
    def test_disabled_account_sessions_have_their_own_state_filter_and_no_live_count_on_hr_users(self):
        user = HRUser.objects.create(username="adv_off", password_hash=SECRET, is_active=False)
        session = LoginSession.objects.create(hr_user=user, jti=uuid.uuid4().hex, device_label="Edge")
        LoginSession.objects.filter(pk=session.pk).update(
            created_at=NOW_UTC - timedelta(hours=2), last_seen_at=NOW_UTC - timedelta(hours=1)
        )
        state = lambda s: [  # noqa: E731
            (r["username"], r["state"]) for r in self.run_report("login-sessions", state=s, **WINDOW)["rows"]
        ]
        self.assertEqual(state("disabled"), [("adv_off", "Account disabled")])
        self.assertNotIn("adv_off", [name for name, _ in state("live")])
        self.assertNotIn("adv_off", [name for name, _ in state("expired")])
        # an expired token stays "Expired" whether or not the account is enabled
        LoginSession.objects.filter(pk=session.pk).update(created_at=NOW_UTC - timedelta(hours=20))
        self.assertEqual(state("expired"), [("adv_off", "Expired")])
        LoginSession.objects.filter(pk=session.pk).update(created_at=NOW_UTC - timedelta(hours=2))
        row = by_key(self.run_report("hr-users", state="inactive"), "adv_off", "username")
        self.assertEqual(row["liveSessions"], 0)

    def test_login_attempts_cards_cover_every_match_and_lockouts_only_the_listed_rows(self):
        from .reporting import runner

        HrLoginAttempt.objects.bulk_create(
            [HrLoginAttempt(username="victim", ip_address="7.7.7.7", success=False) for _ in range(5)]
            + [HrLoginAttempt(username="victim", ip_address="7.7.7.7", success=True)]
            + [HrLoginAttempt(username=f"bot{i}", ip_address="8.8.8.8", success=False) for i in range(30)]
        )
        for i, a in enumerate(HrLoginAttempt.objects.filter(username="victim").order_by("id")):
            HrLoginAttempt.objects.filter(pk=a.pk).update(created_at=ist(2026, 9, 5, 9, i))
        HrLoginAttempt.objects.filter(username__startswith="bot").update(created_at=ist(2026, 9, 6, 9))
        # the lock-out is a failure row: listing only successes must not count (or show) it
        ok = self.run_report("login-attempts", outcome="success", **WINDOW)
        self.assertEqual([r["username"] for r in ok["rows"]], ["victim"])
        self.assertEqual((cards(ok)["Attempts"], cards(ok)["Lockouts"]), (1, 0))
        failed = self.run_report("login-attempts", username="victim", outcome="failed", **WINDOW)
        self.assertEqual(sum(1 for r in failed["rows"] if r["lockedOut"]), 1)
        self.assertEqual(cards(failed)["Lockouts"], 1)
        # a cut-off list: the cards still describe every matching attempt, and the report says so
        with mock.patch.object(runner, "SCREEN_ROW_LIMIT", 10):
            body = self.run_report("login-attempts", outcome="failed", **WINDOW)
        self.assertTrue(body["truncated"])
        self.assertEqual(len(body["rows"]), 10)
        self.assertEqual(cards(body)["Attempts"], 35)
        self.assertEqual(cards(body)["Failing usernames"], 31)
        self.assertTrue(any("More attempts match" in n for n in body["notes"]))

    def test_resume_still_on_disk_counts_as_on_file_and_a_missing_path_as_removed(self):
        import tempfile

        dept = Department.objects.create(name="ADV-LEGACY", branch=self.b1)
        rule = HiringRuleSet.objects.create(name="RS-adv-legacy", department=dept)
        with tempfile.TemporaryDirectory() as media:
            (Path(media) / "resumes").mkdir()
            (Path(media) / "resumes" / "legacy.pdf").write_bytes(b"%PDF-1.4 old")
            for name, path in (("Legacy Kept", "resumes/legacy.pdf"), ("Legacy Gone", "resumes/gone.pdf")):
                c = ScreeningCandidate.objects.create(
                    rule_set=rule, department=dept, resume_file=path, candidate_name=name, status="screened"
                )
                ScreeningCandidate.objects.filter(pk=c.pk).update(created_at=ist(2026, 9, 10, 9))
            with override_settings(MEDIA_ROOT=media):
                body = self.run_report("screening-pipeline", **WINDOW)
        self.assertEqual(by_key(body, "Legacy Kept", "candidateName")["resumeOnFile"], "On file")
        self.assertEqual(by_key(body, "Legacy Gone", "candidateName")["resumeOnFile"], "Removed")
