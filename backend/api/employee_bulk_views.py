"""Bulk employee upload (create) and bulk update from an Excel template."""

import io

from .audit_utils import log_action
from .auth import require_hr
from .branch_scope import get_branch_scope, scope_to_branch
from .employee_views import _assign_unit_code, _create_employee_from_data, _employee_queryset
from .models import Department
from .serializers import parse_decimal
from .view_common import _error
from datetime import date, datetime
from decimal import Decimal
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.request import Request
from rest_framework.response import Response


# --- Bulk employee upload ---
#
# Column order/text is the enforced contract with the downloaded template —
# keep this in sync with EMPLOYEE_TEMPLATE_HEADERS in
# frontend/src/pages/hr/BulkUploadEmployees.tsx if either ever changes.
EMPLOYEE_UPLOAD_HEADERS = [
    "Employee Code", "First Name", "Last Name", "Email", "Phone", "Gender",
    "Date of Birth", "Employment Type", "Department", "Designation", "Branch",
    "Salary Type", "Salary Amount", "Salary Per Shift", "Join Date",
    "Bank Name", "Bank Account", "Bank IFSC", "PF Number", "ESI Number",
    "Address", "ID Proof", "Father's Name", "Mother's Name",
    "Biometric Device ID", "Blood Group", "Emergency Contact",
]


_VALID_EMPLOYMENT_TYPES = {"staff", "production"}


_VALID_SALARY_TYPES = {"monthly", "weekly"}


_VALID_GENDERS = {"male", "female", "other"}


def _parse_date_cell(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    return raw  # unparseable -let Django's own validation reject it with its own message


def _employee_row_to_data(row: dict) -> tuple[dict, str | None]:
    """Normalize one raw Excel row (header -> cell value) into the camelCase
    dict _create_employee_from_data expects. Returns (data, error) -error is
    set when a restricted-choice column (Employment Type/Salary Type/Gender)
    holds something other than one of its known values or blank."""

    def cell(key: str) -> str:
        v = row.get(key)
        return "" if v is None else str(v).strip()

    def num_cell(key: str):
        v = row.get(key)
        if v is None or (isinstance(v, str) and v.strip() == ""):
            return None
        return v

    emp_type = cell("Employment Type").lower()
    if emp_type and emp_type not in _VALID_EMPLOYMENT_TYPES:
        return {}, f"Employment Type must be Staff or Production, got '{cell('Employment Type')}'"

    salary_type = cell("Salary Type").lower()
    if salary_type and salary_type not in _VALID_SALARY_TYPES:
        return {}, f"Salary Type must be Monthly or Weekly, got '{cell('Salary Type')}'"

    gender = cell("Gender").lower()
    if gender and gender not in _VALID_GENDERS:
        return {}, f"Gender must be Male, Female or Other, got '{cell('Gender')}'"

    return {
        "employeeCode": cell("Employee Code"),
        "firstName": cell("First Name"),
        "lastName": cell("Last Name"),
        "email": cell("Email") or None,
        "phone": cell("Phone"),
        "gender": gender or None,
        "dateOfBirth": _parse_date_cell(row.get("Date of Birth")),
        "employmentType": emp_type or None,
        "department": cell("Department") or None,
        "designation": cell("Designation") or None,
        "branch": cell("Branch") or None,
        "salaryType": salary_type or None,
        "salaryAmount": num_cell("Salary Amount"),
        "salaryPerShift": num_cell("Salary Per Shift"),
        "joinDate": _parse_date_cell(row.get("Join Date")),
        "bankName": cell("Bank Name") or None,
        "bankAccount": cell("Bank Account") or None,
        "bankIfsc": cell("Bank IFSC") or None,
        "pfNumber": cell("PF Number") or None,
        "esiNumber": cell("ESI Number") or None,
        "address": cell("Address") or None,
        "idProof": cell("ID Proof") or None,
        "fatherName": cell("Father's Name") or None,
        "motherName": cell("Mother's Name") or None,
        "biometricDeviceId": cell("Biometric Device ID") or None,
        "bloodGroup": cell("Blood Group") or None,
        "emergencyContact": cell("Emergency Contact") or None,
    }, None


@api_view(["POST"])
@parser_classes([MultiPartParser, FormParser])
@require_hr
def bulk_upload_employees(request: Request) -> Response:
    file = request.FILES.get("file")
    if not file:
        return _error("No file uploaded. Send as multipart/form-data with key 'file'.")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        return _error(f"Failed to read the Excel file: {e}")

    if not rows:
        return _error("The uploaded file is empty.")

    def _normalize_header(c) -> str:
        # The downloaded template marks required columns as "Employee Code *"
        # for the user's benefit -strip that trailing marker back off before
        # comparing, so an unmodified official template always validates.
        h = str(c).strip() if c is not None else ""
        return h[:-1].rstrip() if h.endswith("*") else h

    header_row = [_normalize_header(c) for c in rows[0]]
    while header_row and header_row[-1] == "":
        header_row.pop()

    if header_row != EMPLOYEE_UPLOAD_HEADERS:
        return Response(
            {
                "error": "invalid_template",
                "message": "Invalid template. Please upload the employee data using the official template provided by the system.",
            },
            status=400,
        )

    created = 0
    sample_skipped = 0
    errors: list[str] = []
    warnings: list[str] = []

    for idx, raw_row in enumerate(rows[1:], start=2):
        if not raw_row or all(c is None or str(c).strip() == "" for c in raw_row):
            continue  # skip fully blank rows
        first_cell = str(raw_row[0]).strip() if raw_row[0] is not None else ""
        if first_cell.upper().startswith("SAMPLE"):
            # Reference rows shipped in the downloaded template -never imported,
            # even if the user forgets to delete them before uploading.
            sample_skipped += 1
            continue
        row = dict(zip(EMPLOYEE_UPLOAD_HEADERS, raw_row))
        data, row_error = _employee_row_to_data(row)
        if row_error:
            errors.append(f"Row {idx}: {row_error}")
            continue
        emp, error, row_warnings = _create_employee_from_data(data, request, strict=False)
        if error:
            errors.append(f"Row {idx}: {error}")
            continue
        created += 1
        log_action(
            request, "create", "employees", record_id=emp.id,
            description=f"Bulk-imported employee {emp.employee_code} -{emp.first_name} {emp.last_name}",
        )
        warnings.extend(f"Row {idx}: {w}" for w in row_warnings)

    return Response(
        {
            "message": f"Imported {created} employee(s)." + (f" {len(errors)} row(s) failed." if errors else ""),
            "created": created,
            "failed": len(errors),
            "sampleRowsSkipped": sample_skipped,
            "errors": errors,
            "warnings": warnings,
        },
        status=201,
    )


# --- Bulk employee update (existing employees, matched by Employee Code) ---

# Fields the Excel updater may change, as (header, model attr) pairs. Employee
# Code is deliberately absent -it's the match key, so this flow can never
# rename it. Department/Designation/Branch are handled separately (FK
# resolution), as are the choice-validated columns.
_UPDATE_TEXT_FIELDS = {
    "First Name": "first_name",
    "Last Name": "last_name",
    "Email": "email",
    "Phone": "phone",
    "Bank Name": "bank_name",
    "Bank Account": "bank_account",
    "Bank IFSC": "bank_ifsc",
    "PF Number": "pf_number",
    "ESI Number": "esi_number",
    "Address": "address",
    "ID Proof": "id_proof",
    "Father's Name": "father_name",
    "Mother's Name": "mother_name",
    "Biometric Device ID": "biometric_device_id",
    "Blood Group": "blood_group",
    "Emergency Contact": "emergency_contact",
}


def _apply_row_updates(emp, row: dict, request: Request) -> tuple[list[str], list[str], str | None]:
    """
    Apply one Excel row's non-blank cells onto an existing Employee, writing
    only values that actually differ. Blank cells always mean "leave as is" —
    this flow exists to fill gaps and fix mistakes, so an empty cell must
    never wipe stored data. Returns (changed_field_labels, warnings, error).
    Nothing is saved here; the caller saves when changes is non-empty.
    """
    from .models import Branch, Designation as _Desig

    changed: list[str] = []
    warnings: list[str] = []

    def cell(key: str) -> str:
        v = row.get(key)
        return "" if v is None else str(v).strip()

    # Choice-validated columns -a bad value fails the whole row rather than
    # silently skipping, so typos get fixed instead of ignored.
    gender = cell("Gender").lower()
    if gender and gender not in _VALID_GENDERS:
        return [], [], f"Gender must be Male, Female or Other, got '{cell('Gender')}'"
    emp_type = cell("Employment Type").lower()
    if emp_type and emp_type not in _VALID_EMPLOYMENT_TYPES:
        return [], [], f"Employment Type must be Staff or Production, got '{cell('Employment Type')}'"
    salary_type = cell("Salary Type").lower()
    if salary_type and salary_type not in _VALID_SALARY_TYPES:
        return [], [], f"Salary Type must be Monthly or Weekly, got '{cell('Salary Type')}'"

    for header, attr in _UPDATE_TEXT_FIELDS.items():
        value = cell(header)
        if value and value != str(getattr(emp, attr) or ""):
            setattr(emp, attr, value)
            changed.append(header)

    for value, attr, header in [
        (gender, "gender", "Gender"),
        (emp_type, "employment_type", "Employment Type"),
        (salary_type, "salary_type", "Salary Type"),
    ]:
        if value and value != (getattr(emp, attr) or ""):
            setattr(emp, attr, value)
            changed.append(header)

    for header, attr in [("Date of Birth", "date_of_birth"), ("Join Date", "join_date")]:
        if cell(header):
            parsed = _parse_date_cell(row.get(header))
            if parsed and parsed != str(getattr(emp, attr) or ""):
                setattr(emp, attr, parsed)
                changed.append(header)

    for header, attr in [("Salary Amount", "salary_amount"), ("Salary Per Shift", "salary_per_shift")]:
        if cell(header):
            parsed = parse_decimal(row.get(header))
            current = getattr(emp, attr)
            if parsed is not None and (current is None or Decimal(str(current)) != Decimal(str(parsed))):
                setattr(emp, attr, parsed)
                changed.append(header)

    dept_name = cell("Department")
    if dept_name and dept_name.lower() != (emp.department.name.lower() if emp.department_id and emp.department else ""):
        dept = Department.objects.filter(name__iexact=dept_name).first()
        if dept is None:
            dept, _ = Department.objects.get_or_create(name=dept_name)
        emp.department = dept
        changed.append("Department")

    desig_title = cell("Designation")
    if desig_title and desig_title.lower() != (emp.designation.title.lower() if emp.designation_id and emp.designation else ""):
        desig_qs = _Desig.objects.filter(title__iexact=desig_title)
        desig = (desig_qs.filter(department=emp.department).first() if emp.department_id else None) or desig_qs.first()
        if desig is None:
            warnings.append(f"Designation '{desig_title}' not found -kept the current one")
        else:
            emp.designation = desig
            changed.append("Designation")

    # Branch: a branch-scoped HR user can't move employees between branches
    # (same rule as the Edit Employee form) -their rows silently keep the
    # current branch. Unscoped users match by name; unknown names warn.
    branch_name = cell("Branch")
    if branch_name and get_branch_scope(request) is None:
        current_branch_name = emp.branch.name if emp.branch_id and emp.branch else ""
        if branch_name.lower() != current_branch_name.lower():
            b = Branch.objects.filter(name__iexact=branch_name).first()
            if b is None:
                warnings.append(f"Branch '{branch_name}' not found -kept the current one")
            else:
                emp.branch_id = b.id
                # Old Unit Code described the old branch -mint a fresh one.
                emp.unit_code = _assign_unit_code(b.id)
                changed.append("Branch")

    return changed, warnings, None


@api_view(["POST"])
@parser_classes([MultiPartParser, FormParser])
@require_hr
def bulk_update_employees(request: Request) -> Response:
    """
    Companion to bulk_upload_employees for EXISTING employees: HR downloads
    the current-employees export, fills in missing/corrected cells, and
    re-uploads it here. Rows match by Employee Code; codes not in the system
    are reported (never created -that's what bulk upload is for); blank
    cells never overwrite stored data; only genuinely changed fields are
    written, and the response lists exactly what changed per employee.
    """
    file = request.FILES.get("file")
    if not file:
        return _error("No file uploaded. Send as multipart/form-data with key 'file'.")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        return _error(f"Failed to read the Excel file: {e}")

    if not rows:
        return _error("The uploaded file is empty.")

    def _normalize_header(c) -> str:
        h = str(c).strip() if c is not None else ""
        return h[:-1].rstrip() if h.endswith("*") else h

    header_row = [_normalize_header(c) for c in rows[0]]
    while header_row and header_row[-1] == "":
        header_row.pop()

    if header_row != EMPLOYEE_UPLOAD_HEADERS:
        return Response(
            {
                "error": "invalid_template",
                "message": "Invalid file. Please upload the Excel downloaded from Download Current Employees (or the official template) without changing its columns.",
            },
            status=400,
        )

    updated = 0
    unchanged = 0
    sample_skipped = 0
    not_found: list[str] = []
    errors: list[str] = []
    warnings: list[str] = []
    changes: list[str] = []

    scoped_qs = scope_to_branch(_employee_queryset(), request)

    for idx, raw_row in enumerate(rows[1:], start=2):
        if not raw_row or all(c is None or str(c).strip() == "" for c in raw_row):
            continue
        first_cell = str(raw_row[0]).strip() if raw_row[0] is not None else ""
        if first_cell.upper().startswith("SAMPLE"):
            sample_skipped += 1
            continue
        if not first_cell:
            errors.append(f"Row {idx}: Employee Code is required to match an existing employee")
            continue

        row = dict(zip(EMPLOYEE_UPLOAD_HEADERS, raw_row))
        emp = scoped_qs.filter(employee_code=first_cell).first()
        if emp is None:
            not_found.append(f"Row {idx}: no employee with code '{first_cell}' -use Bulk Upload to add new employees")
            continue

        changed, row_warnings, row_error = _apply_row_updates(emp, row, request)
        warnings.extend(f"Row {idx} ({first_cell}): {w}" for w in row_warnings)
        if row_error:
            errors.append(f"Row {idx} ({first_cell}): {row_error}")
            continue
        if not changed:
            unchanged += 1
            continue

        emp.save()
        if "Employment Type" in changed and emp.employment_type == "production":
            from .shift_views import auto_assign_production_shift
            auto_assign_production_shift(emp)
        updated += 1
        changes.append(f"{first_cell} -{emp.first_name} {emp.last_name}: {', '.join(changed)}")
        log_action(
            request, "update", "employees", record_id=emp.id,
            description=f"Bulk-updated employee {emp.employee_code} -changed {', '.join(changed)}",
        )

    return Response(
        {
            "message": f"Updated {updated} employee(s), {unchanged} already up to date."
            + (f" {len(not_found)} code(s) not found." if not_found else "")
            + (f" {len(errors)} row(s) failed." if errors else ""),
            "updated": updated,
            "unchanged": unchanged,
            "notFound": not_found,
            "failed": len(errors),
            "sampleRowsSkipped": sample_skipped,
            "errors": errors,
            "warnings": warnings,
            "changes": changes,
        }
    )
