"""
Report Center - HR Reports (category ``hr``): hr-salary-statement, the printed Salary Statement.

The report must show exactly what Payroll > Salary Statement (GET /api/payroll/salary-statement) shows, because both call
the same builder. Payrolls are made by the REAL engine (``_generate_staff_payroll``) from manual attendance day records,
for April 2026 (26 working days, like the printed statement), with "today" frozen after the month.

Run via: python manage.py test api.tests_reporting_hr_salary_statement -v 2
"""

import io
import re
from datetime import date
from decimal import Decimal
from unittest.mock import patch

from django.apps import apps
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .models import (
    Advance,
    AdvanceRepayment,
    AttendanceDayRecord,
    Branch,
    CasualLeaveRequest,
    Department,
    Designation,
    Employee,
    HRUser,
    Payroll,
    PayrollSettings,
    Role,
    SalarySlip,
)
from .payroll_views import _build_working_days, _generate_staff_payroll
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import hr_reports_salary_statement as S
from .reporting.export_xlsx import HEADER_ROW

RID = "hr-salary-statement"
MONTH, YEAR = 4, 2026
WORKING = _build_working_days(MONTH, YEAR, False, set())  # 26 Mon-Sat days
PERIOD = "2026-04"

LABELS = [
    "SNO", "ID NO", "Name", "Designation", "Basic salary", "Total", "Present", "CL", "LOP", "HD", "Total Days", "PTRL",
    "Others Allow", "E salary", "T.Advance", "Advance", "ESI", "PF", "TDS", "Others", "Lop& PSAmt", "Total Deduction",
    "NET", "Remarks",
]  # fmt: skip
KEYS = [
    "sno", "code", "name", "designation", "basic", "total", "present", "cl", "lop", "hd", "totalDays", "petrol",
    "othersAllow", "eSalary", "tAdvance", "advance", "esi", "pf", "tds", "others", "lopAmount", "totalDeduction", "net",
    "remarks",
]  # fmt: skip
PRINTED_TOTALS = [
    "petrol", "othersAllow", "eSalary", "tAdvance", "advance", "esi", "pf", "tds", "others", "totalDeduction", "net",
]  # fmt: skip
VALUE_KEYS = [k for k in KEYS if k != "sno"]

# The attendance engine only computes days that have elapsed: with "today" fixed after the month the figures do not depend
# on the date the suite runs on.
_FROZEN_TODAY = patch("api.attendance_final.ist_today", return_value=date(2026, 12, 31))


def setUpModule():
    _FROZEN_TODAY.start()


def tearDownModule():
    _FROZEN_TODAY.stop()


def hdr(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def data_rows(body) -> list[dict]:
    return [r for r in body["rows"] if not r.get("_kind")]


def by_code(body) -> dict:
    return {r["code"]: r for r in data_rows(body)}


def headings(body) -> list[str]:
    """The department heading lines: structural lines that carry a name and no money at all."""
    return [r["name"] for r in body["rows"] if r.get("_kind") == "subtotal" and r["net"] is None]


def subtotals(body) -> dict[str, dict]:
    return {
        r["name"][: -len(" Sub total")]: r
        for r in body["rows"]
        if r.get("_kind") == "subtotal" and r["name"].endswith(" Sub total")
    }


def day(emp, d, status="present", shifts="1.00"):
    return AttendanceDayRecord.objects.create(
        employee=emp,
        date=d,
        status=status,
        shifts_earned=Decimal(shifts),
        is_half_shift=(status == "half_shift"),
        source="manual",
    )


class _World(TestCase):
    """April 2026, seven staff payrolls in two units, plus staff who have none.

    Unit 1: Accounts (2316 Meena: salary 29,500 split with a Petrol part, 22 days + 2 approved Casual Leave + 2 absent; 2896
    Lokesh: a half day, an advance of 6,000 with 2,000 recovered; 30071 Sriram: 3 days, now INACTIVE) and Admin (189, 226).
    Unit 2: Cutting (3000) and its own Admin department (5000: same NAME as Unit 1's Admin, a different department).
    No payroll: 4000 (Unit 1, active, counted), 4003 (Unit 2, active, counted), 4001 (joined in May: not counted yet),
    4002 (inactive) and a production worker."""

    @classmethod
    def setUpTestData(cls):
        ps = PayrollSettings.get()
        for key, value in {
            "attendance_mode": "simple",
            "late_deduction_slabs": [],
            "permission_monthly_cap": 3,
            "compensation_feature_enabled": True,
            "company_name": "U.K TEXTILES",
        }.items():
            setattr(ps, key, value)
        ps.save()

        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.accounts = Department.objects.create(name="Accounts", branch=cls.b1)
        cls.admin_d = Department.objects.create(name="Admin", branch=cls.b1)
        cls.cutting = Department.objects.create(name="Cutting", branch=cls.b2)
        cls.admin_d2 = Department.objects.create(name="Admin", branch=cls.b2)
        cls.des_exec = Designation.objects.create(title="Accounts Executive")
        cls.des_op = Designation.objects.create(title="Operator")

        def staff(code, first, dept, branch, salary, designation=None, **kw):
            fields = {
                "status": "active",
                "salary_type": "monthly",
                "salary_amount": Decimal(salary),
                "department": dept,
                "branch": branch,
                "designation": designation or cls.des_op,
                **kw,
            }
            return Employee.objects.create(
                employee_code=code, first_name=first, last_name="Staff", employment_type="staff", **fields
            )

        cls.staff_maker = staff
        w = WORKING
        cls.meena = staff(
            "2316", "MEENATCHI", cls.accounts, cls.b1, "29500", cls.des_exec,
            salary_basic=Decimal("14750"), salary_da=Decimal("0"), salary_retaining_allowance=Decimal("0"),
            salary_other_allowance=Decimal("12750"), salary_petrol_allowance=Decimal("2000"), salary_hra=Decimal("0"),
            salary_special_allowance=Decimal("0"), salary_ca=Decimal("0"),
        )  # fmt: skip
        for d in w[:22]:
            day(cls.meena, d)
        for d in w[22:24]:
            day(cls.meena, d)  # an approved Casual Leave day is stored as a Present day
            CasualLeaveRequest.objects.create(employee=cls.meena, date=d, status="approved")
        for d in w[24:]:
            day(cls.meena, d, "absent", "0")

        cls.lokesh = staff("2896", "LOKESH", cls.accounts, cls.b1, "17500", cls.des_exec)
        for d in w[:20]:
            day(cls.lokesh, d)
        day(cls.lokesh, w[20], "half_shift", "0.50")
        day(cls.lokesh, w[21])
        CasualLeaveRequest.objects.create(employee=cls.lokesh, date=w[21], status="approved")
        for d in w[22:]:
            day(cls.lokesh, d, "absent", "0")
        advance = Advance.objects.create(
            employee=cls.lokesh, advance_type="general", amount=Decimal("6000"), status="approved",
            total_repaid=Decimal("0"), outstanding=Decimal("6000"),
        )  # fmt: skip
        AdvanceRepayment.objects.create(advance=advance, month=MONTH, year=YEAR, amount=Decimal("2000"))

        cls.sriram = staff("30071", "SRIRAM", cls.accounts, cls.b1, "50000", cls.des_exec, status="inactive")
        for d in w[:3]:
            day(cls.sriram, d)
        for d in w[3:]:
            day(cls.sriram, d, "absent", "0")

        others = [
            staff("189", "MURUGAN", cls.admin_d, cls.b1, "50000"),
            staff("226", "KRISHNAN", cls.admin_d, cls.b1, "14000"),
            staff("3000", "CHITRA", cls.cutting, cls.b2, "21000"),
            staff("5000", "RAJA", cls.admin_d2, cls.b2, "16000"),
        ]
        for emp in others:
            for d in w:
                day(emp, d)
        cls.everyone = (cls.meena, cls.lokesh, cls.sriram, *others)
        cls.payrolls = {e.employee_code: _generate_staff_payroll(e, MONTH, YEAR)["payroll"] for e in cls.everyone}

        # staff payroll has not been run for
        staff("4000", "NOPAY", cls.accounts, cls.b1, "10000")
        staff("4003", "NOPAYTWO", cls.cutting, cls.b2, "10000")
        staff("4001", "LATEJOIN", cls.accounts, cls.b1, "10000", join_date="2026-05-01")
        staff("4002", "GONE", cls.accounts, cls.b1, "10000", status="inactive")
        Employee.objects.create(
            employee_code="P1", first_name="Prod", last_name="Worker", employment_type="production",
            status="active", salary_per_shift=Decimal("500"), department=cls.accounts, branch=cls.b1,
        )  # fmt: skip

        # people
        cls.admin = HRUser.objects.create(username="ss_admin", password_hash="x", is_super_admin=True)

        def user(name, perms, **kw):
            role = Role.objects.create(name=name, permissions={"reports": "view", **perms})
            return HRUser.objects.create(username=name, password_hash="x", role=role, **kw)

        cls.payroll_user = user("ss_payroll", {"payroll": "view"})
        cls.salary_user = user("ss_salary", {"salary": "view", "salary_slip": "view"})
        cls.production_user = user("ss_production", {"production_payroll": "view"})
        cls.plain_user = user("ss_plain", {})
        cls.branch_user = user("ss_branch", {"payroll": "view"}, branch=cls.b1)

    # ── helpers ────────────────────────────────────────────────────────────
    def get(self, user=None, expect=200, **params):
        params = {"period": PERIOD, **params}
        r = self.client.get(f"/api/reports/run/{RID}", params, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, fmt, user=None, expect=200, **params):
        params = {"period": PERIOD, "fmt": fmt, **params}
        r = self.client.get(f"/api/reports/export/{RID}", params, **hdr(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r

    def statement(self, extra="", user=None):
        r = self.client.get(
            f"/api/payroll/salary-statement?month={MONTH}&year={YEAR}{extra}", **hdr(user or self.admin)
        )
        self.assertEqual(r.status_code, 200, r.content[:400])
        return r.json()

    def statement_rows(self, body=None) -> dict:
        body = body or self.statement()
        return {r["values"]["code"]: r["values"] for g in body["groups"] for r in g["rows"]}


# ═════════════════════════════════════════════════════════════════════════════
# Registration, columns and access
# ═════════════════════════════════════════════════════════════════════════════


class CatalogTests(_World):
    def test_the_report_is_registered_in_the_hr_group_with_the_printed_columns(self):
        spec = next(s for s in registry.all_specs() if s.id == RID)
        self.assertEqual(registry.LOAD_ERRORS, {})
        self.assertEqual((spec.title, spec.category), ("Salary Statement", "hr"))
        self.assertEqual([c.label for c in spec.columns], LABELS)
        self.assertEqual([c.key for c in spec.columns], KEYS)
        self.assertTrue(set(spec.modules) <= set(all_module_keys()), spec.modules)
        self.assertFalse(spec.super_admin_only)
        types = {c.key: c.type for c in spec.columns}
        self.assertEqual(types["sno"], "integer")
        for key in ("code", "name", "designation", "remarks"):
            self.assertEqual(types[key], "text", key)
        for key in ("present", "cl", "lop", "hd", "totalDays"):
            self.assertEqual(types[key], "number", key)
        for key in (
            "basic", "total", "petrol", "othersAllow", "eSalary", "tAdvance", "advance", "esi", "pf", "tds", "others",
            "lopAmount", "totalDeduction", "net",
        ):  # fmt: skip
            self.assertEqual(types[key], "currency", key)
        # exactly the printed statement's grand-total columns are added up
        self.assertEqual([c.key for c in spec.columns if c.total == "sum"], PRINTED_TOTALS)
        self.assertEqual(S.SUM_KEYS, tuple(PRINTED_TOTALS))

    def test_the_catalog_offers_it_with_its_filters(self):
        body = self.client.get("/api/reports/catalog", **hdr(self.admin)).json()
        report = next(r for r in body["reports"] if r["id"] == RID)
        self.assertEqual(report["category"], "hr")
        self.assertEqual([c["label"] for c in report["columns"]], LABELS)
        self.assertEqual(
            [(f["key"], f["kind"]) for f in report["filters"]],
            [
                ("period", "period"),
                ("branch", "branch"),
                ("department", "department"),
                ("designation", "designation"),
                ("employee", "employee"),
                ("employeeStatus", "employeeStatus"),
                ("paymentStatus", "select"),
            ],
        )
        status = next(f for f in report["filters"] if f["key"] == "employeeStatus")
        self.assertEqual(status["default"], "all")  # an employee who has since left stays on last month's statement

    def test_the_month_defaults_to_last_month(self):
        with patch("api.reporting.filters.ist_today", return_value=date(2026, 5, 10)):
            r = self.client.get(f"/api/reports/run/{RID}", **hdr(self.admin))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(data_rows(r.json())), 7)
        self.assertEqual(dict((f["label"], f["value"]) for f in r.json()["filters"])["Month"], "April 2026")

    def test_a_bad_month_is_a_400(self):
        r = self.client.get(f"/api/reports/run/{RID}", {"period": "April"}, **hdr(self.admin))
        self.assertEqual(r.status_code, 400)


class AccessTests(_World):
    def test_only_the_payroll_module_opens_it_like_the_statement_endpoint(self):
        for who in (self.payroll_user, self.admin):
            self.assertEqual(len(data_rows(self.get(user=who))), 7, who.username)
            self.assertEqual(self.export("xlsx", user=who).status_code, 200)
            self.assertEqual(
                self.client.get("/api/payroll/salary-statement?month=4&year=2026", **hdr(who)).status_code, 200
            )

    def test_a_role_without_payroll_access_is_refused_everywhere(self):
        """Salary and deductions are printed: the Salary / Salary Slip / Production Payroll modules are not enough."""
        for who in (self.salary_user, self.production_user, self.plain_user):
            catalog = self.client.get("/api/reports/catalog", **hdr(who)).json()
            self.assertNotIn(RID, [r["id"] for r in catalog["reports"]], who.username)
            r = self.client.get(f"/api/reports/run/{RID}", {"period": PERIOD}, **hdr(who))
            self.assertEqual(r.status_code, 403, who.username)
            self.assertEqual(r.json()["error"], "report_forbidden")
            for fmt in ("xlsx", "pdf"):
                r = self.client.get(f"/api/reports/export/{RID}", {"fmt": fmt, "period": PERIOD}, **hdr(who))
                self.assertEqual(r.status_code, 403, (who.username, fmt))
                self.assertEqual(r.json()["error"], "report_forbidden")

    def test_it_needs_an_hr_login(self):
        self.assertEqual(self.client.get(f"/api/reports/run/{RID}").status_code, 401)
        emp = {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'employee', 'employeeId': self.meena.id})}"}
        self.assertEqual(self.client.get(f"/api/reports/run/{RID}", **emp).status_code, 403)

    def test_a_view_only_role_can_read_it(self):
        # payroll_user holds "view" (not edit) and reads everything: the report never writes
        body = self.get(user=self.payroll_user)
        self.assertEqual(len(data_rows(body)), 7)


# ═════════════════════════════════════════════════════════════════════════════
# What the rows say
# ═════════════════════════════════════════════════════════════════════════════


class ContentTests(_World):
    def test_every_number_equals_the_payroll_statement(self):
        body = self.get()
        expected = self.statement_rows()
        got = by_code(body)
        self.assertEqual(sorted(got), sorted(expected))
        self.assertEqual(len(got), 7)
        for code, row in got.items():
            for key in VALUE_KEYS:
                self.assertEqual(row[key], expected[code][key], (code, key))

    def test_the_departments_serial_numbers_and_order_are_the_statements(self):
        body = self.get()
        statement = self.statement()
        self.assertEqual(
            headings(body), ["Accounts", "Admin (Unit 1)", "Admin (Unit 2)", "Cutting"]
        )  # two departments that share a name are told apart, as on the statement
        self.assertEqual(headings(body), [g["label"] for g in statement["groups"]])
        self.assertEqual([r["code"] for r in data_rows(body)], ["2316", "2896", "30071", "189", "226", "5000", "3000"])
        self.assertEqual([r["sno"] for r in data_rows(body)], [1, 2, 3, 4, 5, 6, 7])  # runs on across departments
        # a heading line sits right before its department's rows, the subtotal right after
        kinds = [(r.get("_kind"), r["name"]) for r in body["rows"]]
        self.assertEqual(kinds[0], ("subtotal", "Accounts"))
        self.assertEqual(kinds[4], ("subtotal", "Accounts Sub total"))
        self.assertEqual(kinds[5], ("subtotal", "Admin (Unit 1)"))
        self.assertEqual(kinds[-1], ("subtotal", "Cutting Sub total"))

    def test_a_known_row_matches_the_printed_example(self):
        row = by_code(self.get())["2316"]
        self.assertEqual(
            (row["name"], row["designation"], row["basic"], row["total"]),
            ("MEENATCHI Staff", "Accounts Executive", 14750.0, 29500.0),
        )
        self.assertEqual(
            (row["present"], row["cl"], row["lop"], row["hd"], row["totalDays"]), (22.0, 2.0, 2.0, 0.0, 24.0)
        )
        self.assertEqual((row["petrol"], row["othersAllow"]), (2000.0, 12750.0))
        lokesh = by_code(self.get())["2896"]
        self.assertEqual((lokesh["tAdvance"], lokesh["advance"], lokesh["hd"]), (6000.0, 2000.0, 1.0))
        self.assertEqual(by_code(self.get())["30071"]["remarks"], "LEFT")

    def test_department_subtotals_add_up_the_printed_columns(self):
        body = self.get()
        rows = data_rows(body)
        accounts = [r for r in rows if r["code"] in ("2316", "2896", "30071")]
        subs = subtotals(body)
        for key in PRINTED_TOTALS:
            self.assertAlmostEqual(subs["Accounts"][key], round(sum(r[key] for r in accounts), 2), 2, key)
            self.assertAlmostEqual(subs["Cutting"][key], by_code(body)["3000"][key], 2, key)
        # the money the printed sheet does not add up stays blank on the subtotal
        for key in ("basic", "total", "present", "totalDays", "lopAmount"):
            self.assertIsNone(subs["Accounts"][key], key)
        self.assertIsNone(subs["Accounts"]["sno"])
        self.assertEqual(len(subs), 4)

    def test_the_grand_total_is_the_statements_total(self):
        body = self.get()
        statement = self.statement()
        for key in PRINTED_TOTALS:
            self.assertAlmostEqual(body["totals"][key], statement["total"]["values"][key], 2, key)
            self.assertAlmostEqual(body["totals"][key], round(sum(r[key] for r in data_rows(body)), 2), 2, key)
        # subtotal lines are not counted twice
        self.assertAlmostEqual(body["totals"]["net"], round(sum(s["net"] for s in subtotals(body).values()), 2), 2)
        for key in ("basic", "total", "present", "lop", "lopAmount"):
            self.assertNotIn(key, body["totals"])
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual((cards["Employees listed"], cards["Departments"]), (7, 4))
        self.assertAlmostEqual(cards["Net pay"], body["totals"]["net"], 2)

    def test_the_notes_explain_the_columns_as_the_payroll_statement_does(self):
        body = self.get()
        statement = self.statement()
        for note in statement["notes"]:
            self.assertIn(note, body["notes"])
        text = " ".join(body["notes"])
        self.assertIn("Basic is the first salary-split portion", text)
        self.assertIn("Lop& PSAmt = Total - E salary", text)

    def test_staff_without_a_payroll_are_counted_in_the_notes_not_listed(self):
        body = self.get()
        listed = {r["code"] for r in data_rows(body)}
        self.assertFalse(listed & {"4000", "4001", "4002", "4003", "P1"})
        missing = [n for n in body["notes"] if "no payroll generated" in n]
        # 4000 (Unit 1) and 4003 (Unit 2): not 4001 (joins in May), 4002 (inactive) or the production worker
        self.assertEqual(len(missing), 1)
        self.assertTrue(missing[0].startswith("2 active staff member(s)"), missing[0])
        self.assertEqual(self.get(departmentIds=str(self.accounts.id))["notes"][-1][:2], "1 ")

    def test_a_month_with_no_payroll_lists_nobody_and_says_so(self):
        body = self.get(period="2026-03")
        self.assertEqual(body["rows"], [])
        self.assertIsNone(body["totals"])
        self.assertTrue(any("No staff payroll has been generated" in n for n in body["notes"]))

    def test_the_status_filter_defaults_to_everyone_with_a_payroll(self):
        self.assertIn("30071", by_code(self.get()))
        self.assertNotIn("30071", by_code(self.get(employeeStatus="active")))
        self.assertEqual(set(by_code(self.get(employeeStatus="inactive"))), {"30071"})
        # the 'no payroll yet' count is about ACTIVE staff whatever the filter says
        for status in ("all", "active", "inactive"):
            self.assertTrue(
                any(n.startswith("2 active staff member(s)") for n in self.get(employeeStatus=status)["notes"]), status
            )

    def test_the_scope_filters_narrow_the_statement_like_its_own_filters(self):
        dept = self.get(departmentIds=str(self.accounts.id))
        self.assertEqual([r["code"] for r in data_rows(dept)], ["2316", "2896", "30071"])
        self.assertEqual([r["sno"] for r in data_rows(dept)], [1, 2, 3])
        self.assertEqual(headings(dept), ["Accounts"])
        self.assertEqual(set(by_code(self.get(designationIds=str(self.des_op.id)))), {"189", "226", "3000", "5000"})
        self.assertEqual(list(by_code(self.get(employeeIds=str(self.lokesh.id)))), ["2896"])
        self.assertEqual(set(by_code(self.get(branchIds=str(self.b2.id)))), {"3000", "5000"})
        statement = self.statement(f"&departmentId={self.accounts.id}")
        self.assertEqual(by_code(dept)["2896"]["net"], self.statement_rows(statement)["2896"]["net"])

    def test_the_payment_status_filter(self):
        Payroll.objects.filter(pk=self.payrolls["2316"].pk).update(status=Payroll.STATUS_PAID)
        self.assertEqual(list(by_code(self.get(paymentStatus="paid"))), ["2316"])
        self.assertEqual(len(data_rows(self.get(paymentStatus="pending"))), 6)
        cards = {c["label"]: c["value"] for c in self.get()["summary"]}
        self.assertEqual((cards["Cleared"], cards["Pending"]), (1, 6))
        self.get(paymentStatus="nonsense", expect=400)


# ═════════════════════════════════════════════════════════════════════════════
# Branch isolation
# ═════════════════════════════════════════════════════════════════════════════


class BranchTests(_World):
    def test_a_branch_user_sees_only_their_own_unit_and_the_statements_numbers_for_it(self):
        body = self.get(user=self.branch_user)
        self.assertEqual([r["code"] for r in data_rows(body)], ["2316", "2896", "30071", "189", "226"])
        self.assertEqual(headings(body), ["Accounts", "Admin"])  # no 'Admin (Unit 1)': the other Admin is not visible
        expected = self.statement_rows(self.statement(user=self.branch_user))
        for code, row in by_code(body).items():
            self.assertEqual(row["net"], expected[code]["net"], code)
        self.assertEqual([r["sno"] for r in data_rows(body)], [1, 2, 3, 4, 5])

    def test_a_branch_user_cannot_widen_their_scope_with_a_parameter(self):
        for params in (
            {"branchIds": str(self.b2.id)},
            {"employeeIds": ",".join(str(e.id) for e in self.everyone if e.branch_id == self.b2.id)},
            {"departmentIds": str(self.cutting.id)},
        ):
            self.assertEqual(self.get(user=self.branch_user, **params)["rows"], [], params)

    def test_the_other_units_people_never_reach_the_notes_totals_or_exports(self):
        body = self.get(user=self.branch_user)
        # of the two staff without a payroll only Unit 1's 4000 is in this login's scope
        self.assertTrue(any(n.startswith("1 active staff member(s)") for n in body["notes"]), body["notes"])
        for key in PRINTED_TOTALS:
            self.assertAlmostEqual(body["totals"][key], round(sum(r[key] for r in data_rows(body)), 2), 2)
        ws = load_workbook(io.BytesIO(self.export("xlsx", user=self.branch_user).content)).active
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIn("MEENATCHI", cells)
        for hidden in ("CHITRA", "RAJA", "Cutting", "Unit 2"):
            self.assertNotIn(hidden, cells)

    def test_a_branch_user_gets_no_branch_filter(self):
        catalog = self.client.get("/api/reports/catalog", **hdr(self.branch_user)).json()
        report = next(r for r in catalog["reports"] if r["id"] == RID)
        self.assertNotIn("branch", [f["key"] for f in report["filters"]])


# ═════════════════════════════════════════════════════════════════════════════
# Exports
# ═════════════════════════════════════════════════════════════════════════════


class ExportTests(_World):
    def test_excel_is_the_screen_cell_for_cell(self):
        screen = self.get()
        r = self.export("xlsx")
        self.assertRegex(r["Content-Disposition"], r'filename="salary_statement_\d{8}_\d{4}\.xlsx"')
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[HEADER_ROW]][: len(LABELS)], LABELS)
        for i, row in enumerate(screen["rows"]):
            excel = [ws.cell(row=HEADER_ROW + 1 + i, column=n + 1).value for n in range(len(KEYS))]
            for n, key in enumerate(KEYS):
                want = row[key]
                got = excel[n]
                if isinstance(want, float):
                    self.assertAlmostEqual(got, want, 2, (i, key))
                else:
                    self.assertEqual(got if got != "" else None, want if want != "" else None, (i, key))
        total_row = [c.value for c in ws[HEADER_ROW + 1 + len(screen["rows"])]]
        self.assertEqual(total_row[0], "TOTAL")
        self.assertAlmostEqual(total_row[KEYS.index("net")], screen["totals"]["net"], 2)
        cells = " ".join(str(c.value) for rw in ws.iter_rows() for c in rw if c.value is not None)
        self.assertIn("April 2026", cells)  # the filter echo in the sheet header
        self.assertIn("Basic is the first salary-split portion", cells)

    def test_pdf_is_a_real_pdf(self):
        r = self.export("pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertRegex(r["Content-Disposition"], r'\.pdf"$')
        self.assertGreater(len(r.content), 2000)

    def test_the_export_is_audited(self):
        from .models import AuditLog

        before = AuditLog.objects.count()
        self.export("xlsx")
        self.assertEqual(AuditLog.objects.count(), before + 1)
        self.assertIn("Salary Statement", AuditLog.objects.order_by("-id").first().record_description)


# ═════════════════════════════════════════════════════════════════════════════
# Cost and safety
# ═════════════════════════════════════════════════════════════════════════════


class CostAndSafetyTests(_World):
    def _queries(self, **params) -> int:
        params = {"period": PERIOD, **params}
        self.client.get(f"/api/reports/run/{RID}", params, **hdr(self.admin))  # warm-up (permission caches)
        with CaptureQueriesContext(connection) as q:
            r = self.client.get(f"/api/reports/run/{RID}", params, **hdr(self.admin))
        self.assertEqual(r.status_code, 200)
        return len(q)

    def test_the_query_count_does_not_grow_with_the_number_of_employees(self):
        before = self._queries()
        for i in range(10):
            emp = self.staff_maker(
                f"7{i:03d}", f"EXTRA{i}", (self.accounts, self.admin_d, self.cutting)[i % 3], self.b1, "20000"
            )
            for d in WORKING:
                day(emp, d)
            _generate_staff_payroll(emp, MONTH, YEAR)
            if i % 2:
                self.staff_maker(f"8{i:03d}", f"NOPAY{i}", self.accounts, self.b1, "9000")
        self.assertEqual(self._queries(), before)
        self.assertEqual(len(data_rows(self.get())), 17)

    def test_running_and_exporting_never_writes(self):
        def counts():
            return {
                m.__name__: m.objects.count()
                for m in apps.get_app_config("api").get_models()
                if m.__name__ != "AuditLog"  # an export is supposed to leave an audit entry
            }

        payrolls = list(Payroll.objects.order_by("id").values())
        slips = [
            (s.id, s.net_salary, s.total_deductions, s.breakdown_details) for s in SalarySlip.objects.order_by("id")
        ]
        before = counts()
        self.get()  # the HR login row exists from here on
        with CaptureQueriesContext(connection) as q:
            self.get()
            self.get(employeeStatus="active")
            self.export("pdf")
        written = [
            x["sql"][:100]
            for x in q.captured_queries
            if re.match(r"\s*(INSERT|UPDATE|DELETE|TRUNCATE|ALTER|CREATE|DROP)\b", x["sql"], re.I)
            and "audit_log" not in x["sql"].lower()
        ]
        self.assertEqual(written, [])
        self.export("xlsx")
        self.assertEqual(counts(), before)
        self.assertEqual(list(Payroll.objects.order_by("id").values()), payrolls)
        self.assertEqual(
            [(s.id, s.net_salary, s.total_deductions, s.breakdown_details) for s in SalarySlip.objects.order_by("id")],
            slips,
        )

    def test_photo_and_password_columns_are_never_selected(self):
        with CaptureQueriesContext(connection) as q:
            self.get()
        for query in q.captured_queries:
            sql = query["sql"]
            self.assertNotIn("photo_url", sql)
            if '"hr_users"' not in sql:
                self.assertNotIn("password_hash", sql)

    def test_an_export_over_the_row_limit_is_refused_not_cut(self):
        with patch("api.reporting.runner.XLSX_ROW_LIMIT", 3):
            r = self.client.get(f"/api/reports/export/{RID}", {"fmt": "xlsx", "period": PERIOD}, **hdr(self.admin))
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.json()["error"], "too_many_rows")
