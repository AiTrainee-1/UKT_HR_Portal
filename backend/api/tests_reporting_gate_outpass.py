"""
Report Center - gate / outpass reports (group G8): outpass register, In/Out register, employee and department
summaries, not-returned, late returns, approval turnaround, pending requests, expired-unused, purpose analysis,
QR submissions, gate scan audit, gate activity and the daily gate summary.

All figures below are derived by hand from ONE fixture world (see ``build_world``): fixed IST timestamps in
September 2026 and a pinned "now" of 2026-09-15 10:00 IST, so nothing depends on the day the suite runs.

Run via: python manage.py test api.tests_reporting_gate_outpass -v 2
"""

import io
import json
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from unittest import mock

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .clock import FACTORY_TZ
from .gate_scanner_views import resolve_gate_scan
from .jwt_utils import sign_token
from .models import (
    Branch,
    Department,
    Designation,
    Employee,
    GateDevice,
    HRUser,
    OutpassGateScan,
    OutpassRecord,
    OutpassRequest,
    Role,
    TeaBreakLog,
    TeaBreakRule,
    Visitor,
    VisitorVisit,
)
from .outpass_request_views import resolve_outpass_request
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import gate_outpass_common as C

GATE_REPORT_IDS = [
    "outpass-register",
    "outpass-in-out-register",
    "outpass-employee-summary",
    "outpass-department-summary",
    "outpass-not-returned",
    "outpass-late-return",
    "outpass-approval-turnaround",
    "outpass-pending-requests",
    "outpass-expired-unused",
    "outpass-qr-submissions",
    "gate-scan-audit-log",
    "gate-activity-summary",
    "gate-daily-summary",
    "outpass-purpose-analysis",
]
# Owned by outpass_visitors AND requests; the rest by outpass_visitors alone.
REQUEST_LIFECYCLE_IDS = [
    "outpass-register",
    "outpass-in-out-register",
    "outpass-employee-summary",
    "outpass-department-summary",
    "outpass-not-returned",
    "outpass-late-return",
    "outpass-approval-turnaround",
    "outpass-pending-requests",
    "outpass-expired-unused",
    "outpass-purpose-analysis",
]
GATE_ONLY_IDS = ["outpass-qr-submissions", "gate-scan-audit-log", "gate-activity-summary", "gate-daily-summary"]

NOW = datetime(2026, 9, 15, 10, 0, tzinfo=FACTORY_TZ)  # the pinned "now" of every report
RANGE = {"dateFrom": "2026-09-01", "dateTo": "2026-09-15"}
EMPTY_RANGE = {"dateFrom": "2020-01-01", "dateTo": "2020-01-02"}
DASH = "–"  # en dash used in reviewer labels


def _day(value):
    return value.date() if hasattr(value, "date") else value


def ist(text: str) -> datetime:
    fmt = "%Y-%m-%d %H:%M:%S" if text.count(":") == 2 else "%Y-%m-%d %H:%M"
    return datetime.strptime(text, fmt).replace(tzinfo=FACTORY_TZ)


def _when(value):
    return ist(value) if isinstance(value, str) else value


def _hr_headers(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def mk_request(
    emp,
    dest,
    created,
    *,
    status="approved",
    source="manual",
    pass_type=None,
    role=None,
    by=None,
    approved=None,
    exit=None,
    exit_gate=None,
    enter=None,
    entry_gate=None,
    expected=None,
    comment=None,
    return_qr=None,
    reason=None,
):
    req = OutpassRequest.objects.create(
        employee=emp,
        destination=dest,
        reason=reason or f"reason {dest}",
        status=status,
        source=source,
        pass_type=pass_type,
        approver_role=role,
        approved_by=by,
        review_comment=comment,
        approved_at=_when(approved),
        exited_at=_when(exit),
        exit_gate=exit_gate,
        entered_at=_when(enter),
        entry_gate=entry_gate,
        expected_return_at=_when(expected),
        return_qr_generated_at=_when(return_qr),
    )
    stamp = ist(created)
    OutpassRequest.objects.filter(pk=req.pk).update(created_at=stamp, updated_at=stamp)  # auto_now_add ignores input
    return req


def mk_scan(gate, req, emp, scan_type, result, when, message="scan"):
    scan = OutpassGateScan.objects.create(
        gate=gate,
        outpass_request=req,
        employee=emp,
        scan_type=scan_type,
        result=result,
        message=message,
    )
    OutpassGateScan.objects.filter(pk=scan.pk).update(scanned_at=ist(when))
    return scan


def mk_qr(branch, emp, name, code, dest, when, source="qr"):
    rec = OutpassRecord.objects.create(
        branch=branch,
        employee=emp,
        employee_name=name,
        employee_code=code,
        destination=dest,
        source=source,
    )
    OutpassRecord.objects.filter(pk=rec.pk).update(submitted_at=ist(when))
    return rec


def mk_visit(visitor, branch, when, host=None):
    visit = VisitorVisit.objects.create(
        visitor=visitor,
        branch=branch,
        whom_to_meet="Someone",
        purpose="Meeting",
        meeting_employee=host,
    )
    VisitorVisit.objects.filter(pk=visit.pk).update(visited_at=ist(when))
    return visit


def mk_tea(emp, out_gate, out_at, in_gate=None, in_at=None):
    return TeaBreakLog.objects.create(
        employee=emp,
        out_gate=out_gate,
        out_at=ist(out_at),
        in_gate=in_gate,
        in_at=ist(in_at) if in_at else None,
    )


def mk_gate(name, branch, username, active=True, last_login=None):
    return GateDevice.objects.create(
        name=name,
        branch=branch,
        username=username,
        password_hash="x",
        login_token=f"tok-{username}",
        is_active=active,
        last_login_at=ist(last_login) if last_login else None,
    )


def mk_user(name, perms, branch=None, super_admin=False):
    role = Role.objects.create(name=name, permissions=perms)
    return HRUser.objects.create(
        username=name,
        password_hash="x",
        role=None if super_admin else role,
        branch=branch,
        is_super_admin=super_admin,
    )


def build_world(cls):
    """Branches, departments, seven employees, gates, sixteen+ passes, QR rows, scans, tea breaks and visits."""
    cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
    cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
    cls.cut = Department.objects.create(name="CUTTING", branch=cls.b1)
    cls.sew = Department.objects.create(name="SEWING", branch=cls.b2)
    cls.des = Designation.objects.create(title="Supervisor", department=cls.cut)

    def emp(code, first, last, dept=None, branch=None, **kw):
        return Employee.objects.create(
            employee_code=code, first_name=first, last_name=last, department=dept, branch=branch, **kw
        )

    cls.e1 = emp("GT001", "Asha", "Rao", cls.cut, cls.b1, designation=cls.des, phone="9000000001")
    cls.e2 = emp("GT002", "Bala", "Kumar", cls.cut, cls.b1, employment_type="production")
    cls.e3 = emp("GT003", "Chitra", "Devi", cls.sew, cls.b2)
    cls.e4 = emp("GT004", "Dinesh", "Raj", cls.cut, cls.b1, status="inactive")
    cls.e5 = emp("GT005", "Esha", "Nair", None, cls.b1, phone="9000000005")
    cls.e6 = emp("GT006", "Farid", "Khan", None, None)
    cls.e7 = emp("GT007", "Gita", "Sen", cls.sew, cls.b2)  # active, never took a pass
    Employee.objects.filter(pk=cls.e1.pk).update(reporting_manager=cls.e5)

    cls.g1 = mk_gate("Gate 1", cls.b1, "g1")
    cls.g2 = mk_gate("Gate 2", cls.b1, "g2")
    cls.g3 = mk_gate("Gate 3", cls.b2, "g3")
    cls.g4 = mk_gate("Gate 4", cls.b1, "g4", active=False, last_login="2026-09-14 08:00")
    g1, g2, g3 = cls.g1, cls.g2, cls.g3

    R = cls.R = {}
    PR, MH, SH = ("hr", "Priya HR"), ("dept_head", "Mani HOD"), ("dept_head", "Selvi HOD")

    def add(key, e, created, decider=None, **kw):
        role, by = decider or (None, None)
        R[key] = mk_request(e, key, created, role=role, by=by, **kw)

    add(
        "R01",
        cls.e1,
        "2026-09-01 09:00",
        PR,
        pass_type="official",
        approved="2026-09-01 09:10",
        exit="2026-09-01 09:20",
        exit_gate=g1,
        enter="2026-09-01 10:20",
        entry_gate=g2,
        expected="2026-09-01 10:00",
    )
    add(
        "R02",
        cls.e1,
        "2026-09-02 14:00",
        MH,
        pass_type="personal",
        approved="2026-09-02 14:30",
        exit="2026-09-02 14:35",
        exit_gate=g1,
        enter="2026-09-02 15:05",
        entry_gate=g1,
        expected="2026-09-02 15:10",
    )
    add(
        "R03",
        cls.e2,
        "2026-09-03 15:00",
        PR,
        pass_type="early_dismissal",
        approved="2026-09-03 15:05",
        exit="2026-09-03 15:10",
        exit_gate=g1,
    )
    add(
        "R04",
        cls.e2,
        "2026-09-04 11:00",
        PR,
        pass_type="official",
        approved="2026-09-04 11:05",
        exit="2026-09-04 11:10",
        exit_gate=g2,
        expected="2026-09-04 12:00",
    )
    add(
        "R05",
        cls.e3,
        "2026-09-05 10:00",
        SH,
        pass_type="personal",
        approved="2026-09-05 11:30",
        exit="2026-09-05 11:40",
        exit_gate=g3,
        enter="2026-09-05 12:00",
        entry_gate=g3,
    )
    add("R06", cls.e1, "2026-09-06 09:00", PR, status="rejected", pass_type="official", comment="Not needed")
    add("R07", cls.e1, "2026-09-14 08:00", status="pending", pass_type="personal")
    add("R08", cls.e2, "2026-09-15 09:30", status="pending", pass_type="official")
    add("R09", cls.e3, "2026-09-10 10:00", PR, pass_type="official", approved="2026-09-10 10:20")
    add("R10", cls.e1, "2026-09-11 13:00", ("system", "Rita HR"), source="on_duty", approved="2026-09-11 13:00")
    add(
        "R11",
        cls.e5,
        "2026-09-15 08:00",
        PR,
        pass_type="official",
        approved="2026-09-15 08:05",
        exit="2026-09-15 08:30",
        exit_gate=g1,
        expected="2026-09-15 09:00",
        return_qr="2026-09-15 09:30",
    )
    add(
        "R12",
        cls.e4,
        "2026-09-08 09:00",
        MH,
        pass_type="personal",
        approved="2026-09-08 09:12",
        exit="2026-09-08 09:15",
        exit_gate=g1,
        enter="2026-09-08 09:45",
        entry_gate=g1,
        expected="2026-09-08 09:30",
    )
    add(
        "R13",
        cls.e6,
        "2026-09-09 09:00",
        PR,
        pass_type="official",
        approved="2026-09-09 09:05",
        exit="2026-09-09 09:10",
        exit_gate=g1,
        enter="2026-09-09 09:40",
        entry_gate=g1,
    )
    add(
        "R14",
        cls.e1,
        "2026-08-20 09:00",
        PR,
        pass_type="official",
        approved="2026-08-20 09:05",
        exit="2026-08-20 09:10",
        exit_gate=g1,
        enter="2026-08-20 09:55",
        entry_gate=g1,
    )
    add(
        "R15",
        cls.e2,
        "2026-09-08 00:30",
        PR,
        pass_type="official",
        approved="2026-09-08 00:40",
        exit="2026-09-08 00:45",
        exit_gate=g1,
        enter="2026-09-08 01:45",
        entry_gate=g1,
    )
    add(
        "R16",
        cls.e1,
        "2026-09-08 23:40",
        MH,
        pass_type="personal",
        approved="2026-09-08 23:45",
        exit="2026-09-08 23:50",
        exit_gate=g1,
        enter="2026-09-09 00:20",
        entry_gate=g1,
    )
    add(
        "R17",
        cls.e1,
        "2026-09-12 09:00",
        PR,
        pass_type="official",
        approved="2026-09-12 09:05",
        exit="2026-09-12 09:10",
        exit_gate=g1,
        enter="2026-09-12 10:10",
        entry_gate=g1,
        expected="2026-09-12 09:40",
    )
    add(
        "R18",
        cls.e1,
        "2026-09-13 09:00",
        MH,
        pass_type="personal",
        approved="2026-09-13 09:25",
        exit="2026-09-13 09:30",
        exit_gate=g2,
        enter="2026-09-13 10:10",
        entry_gate=g2,
        expected="2026-09-13 09:50",
    )
    add("R19", cls.e3, "2026-09-15 09:20", PR, pass_type="official", approved="2026-09-15 09:30")
    OutpassRequest.objects.filter(pk=R["R10"].pk).update(reason="R10")  # on-duty passes carry reason == destination

    # Gate QR form submissions (+ one mirror copy created by an approval, which must not count as an exit).
    mk_qr(cls.b1, cls.e1, "Asha Rao", "GT001", "Q1", "2026-09-10 10:00")
    mk_qr(cls.b1, cls.e1, "Asha", "gt001", "Q2", "2026-09-11 11:00")
    mk_qr(cls.b1, None, "Unknown Person", "ZZ999", "Q3", "2026-09-12 12:00")
    mk_qr(cls.b2, cls.e3, "Chitra  Devi", "GT003", "Q4", "2026-09-13 09:00")
    mk_qr(None, None, "Nobody", "X1", "Q5", "2026-09-14 09:00")
    mk_qr(cls.b1, cls.e1, "Asha Rao", "GT001", "Q6", "2026-09-01 09:10", source="request")
    mk_qr(cls.b1, cls.e2, "Bala Kumar", "GT002", "Q7", "2026-09-08 00:30")
    mk_qr(cls.b1, cls.e1, "Asha Rao", "GT001", "Q8", "2026-08-15 10:00")

    # Scan attempts (S6 sits at 00:30 IST = the previous day in UTC).
    S = cls.S = {}
    S["S1"] = mk_scan(g1, R["R01"], cls.e1, "exit", "success", "2026-09-01 09:20", "Exit recorded via Gate 1")
    S["S2"] = mk_scan(g2, R["R01"], cls.e1, "entry", "success", "2026-09-01 10:20", "Return recorded via Gate 2")
    S["S3"] = mk_scan(
        g1, R["R01"], cls.e1, "exit", "already_scanned", "2026-09-01 09:45", "Already exited at 03:50 AM via Gate 1"
    )
    S["S11"] = mk_scan(
        g1, R["R01"], cls.e1, "entry", "already_scanned", "2026-09-01 11:00", "Already returned at 04:50 AM via Gate 2"
    )
    S["S5"] = mk_scan(
        g1, R["R06"], cls.e1, "exit", "not_approved", "2026-09-06 09:30", "This Outpass request was not approved."
    )
    S["S6"] = mk_scan(g1, None, None, "exit", "invalid_qr", "2026-09-08 00:30", "This QR code is not recognized.")
    S["S7"] = mk_scan(None, None, None, "exit", "invalid_qr", "2026-09-07 23:00", "This QR code is not recognized.")
    S["S8"] = mk_scan(g3, R["R05"], cls.e3, "exit", "success", "2026-09-05 11:40")
    S["S9"] = mk_scan(g3, R["R05"], cls.e3, "entry", "success", "2026-09-05 12:00")
    S["S4"] = mk_scan(g1, R["R09"], cls.e3, "exit", "expired", "2026-09-10 12:00", "This Outpass has expired.")
    S["S10"] = mk_scan(
        g1, R["R07"], cls.e1, "exit", "not_approved", "2026-09-14 08:30", "This Outpass request was not approved."
    )
    S["S12"] = mk_scan(g1, R["R04"], cls.e2, "exit", "success", "2026-09-04 11:10")

    # Tea breaks (rule is not stored, so the default 15 minutes applies unless a test saves one).
    cls.tea = {
        "T1": mk_tea(cls.e1, g1, "2026-09-05 11:00", g1, "2026-09-05 11:12"),
        "T2": mk_tea(cls.e1, g1, "2026-09-06 11:00", g2, "2026-09-06 11:20"),
        "T3": mk_tea(cls.e2, g2, "2026-09-07 16:00"),
        "T4": mk_tea(cls.e3, g3, "2026-09-05 10:00", g3, "2026-09-05 10:15"),
        "T5": mk_tea(cls.e2, g1, "2026-09-08 00:30", g1, "2026-09-08 00:45"),
        "T6": mk_tea(cls.e1, g1, "2026-09-09 15:00", g1, "2026-09-09 15:15:31"),
        "T7": mk_tea(cls.e1, g1, "2026-09-10 15:00", g1, "2026-09-10 15:15:29"),
        "T8": mk_tea(cls.e1, g1, "2026-09-11 15:00", g1, "2026-09-11 15:15:30"),
        "T9": mk_tea(cls.e1, g1, "2026-09-12 15:00", g1, "2026-09-12 15:16:30"),
    }

    v1 = Visitor.objects.create(name="Vis One", phone="9111111111")
    v2 = Visitor.objects.create(name="Vis Two", phone="9222222222")
    v3 = Visitor.objects.create(name="Vis Three", phone="9333333333")
    mk_visit(v1, cls.b1, "2026-09-05 10:00")
    mk_visit(v1, cls.b1, "2026-09-05 15:00")
    mk_visit(v2, cls.b2, "2026-09-05 11:00")
    mk_visit(v2, cls.b2, "2026-09-08 00:30")
    mk_visit(v3, None, "2026-09-05 12:00")

    both = {"reports": "view", "outpass_visitors": "view"}
    cls.admin = mk_user("gt_admin", {}, super_admin=True)
    cls.b1_user = mk_user("gt_b1", both, branch=cls.b1)
    cls.b2_user = mk_user("gt_b2", both, branch=cls.b2)
    cls.gate_user = mk_user("gt_gate", both)
    cls.requests_user = mk_user("gt_requests", {"reports": "view", "requests": "view"})
    cls.plain_user = mk_user("gt_plain", {"reports": "view"})


class World(TestCase):
    """Base: the shared fixture, a pinned clock and small HTTP helpers."""

    @classmethod
    def setUpTestData(cls):
        build_world(cls)

    def setUp(self):
        patcher = mock.patch("api.reporting.definitions.gate_outpass_common.now", return_value=NOW)
        patcher.start()
        self.addCleanup(patcher.stop)

    def get(self, path, user=None, **params):
        return self.client.get(path, params, **_hr_headers(user or self.admin))

    def run_report(self, rid, user=None, **params):
        r = self.get(f"/api/reports/run/{rid}", user, **params)
        self.assertEqual(r.status_code, 200, r.content[:400])
        return r.json()

    def export(self, rid, fmt, user=None, **params):
        return self.get(f"/api/reports/export/{rid}", user, fmt=fmt, **params)

    @staticmethod
    def dests(body):
        return [r["destination"] for r in body["rows"] if r.get("_kind") is None]

    @staticmethod
    def by(body, key):
        return {r[key]: r for r in body["rows"] if r.get("_kind") is None}

    @staticmethod
    def cards(body):
        return {s["label"]: s["value"] for s in body["summary"]}


ALL_REQ = {f"R{n:02d}" for n in range(1, 20)}
IN_RANGE = ALL_REQ - {"R14"}  # R14 was created in August


# ── outpass-register ────────────────────────────────────────────────────────


class RegisterTests(World):
    RID = "outpass-register"

    def test_golden_rows_summary_and_totals(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(body["rowCount"], 18)
        self.assertEqual(self.dests(body)[:4], ["R08", "R19", "R11", "R07"])  # newest request first
        rows = self.by(body, "destination")

        r1 = rows["R01"]
        for key, want in {
            "requestedAt": "2026-09-01 09:00",
            "employeeCode": "GT001",
            "employeeName": "Asha Rao",
            "department": "CUTTING",
            "designation": "Supervisor",
            "passType": "Official",
            "source": "Manual",
            "status": "Approved",
            "reviewedBy": f"HR {DASH} Priya HR",
            "approvedAt": "2026-09-01 09:10",
            "expiresAt": "2026-09-01 10:10",
            "scanStatus": "Returned",
            "exitGate": "Gate 1",
            "exitedAt": "2026-09-01 09:20",
            "entryGate": "Gate 2",
            "enteredAt": "2026-09-01 10:20",
            "outsideMinutes": 60,
            "expectedReturnAt": "2026-09-01 10:00",
            "reviewComment": None,
        }.items():
            self.assertEqual(r1[key], want, key)

        r6 = rows["R06"]  # rejected: no approval time, reviewer still named, comment kept
        self.assertEqual(
            (r6["status"], r6["approvedAt"], r6["expiresAt"], r6["scanStatus"]), ("Rejected", None, None, None)
        )
        self.assertEqual((r6["reviewedBy"], r6["reviewComment"]), (f"HR {DASH} Priya HR", "Not needed"))

        r10 = rows["R10"]  # on-duty: unspecified type, reason == destination, system approval
        self.assertEqual(
            (r10["source"], r10["passType"], r10["reason"], r10["reviewedBy"], r10["scanStatus"]),
            ("On-Duty", "Unspecified", "R10", "On-Duty approval (Rita HR)", "Expired, Not Scanned"),
        )
        self.assertEqual(rows["R05"]["reviewedBy"], f"HOD {DASH} Selvi HOD")
        self.assertEqual(rows["R16"]["outsideMinutes"], 30)
        self.assertEqual(rows["R16"]["enteredAt"], "2026-09-09 00:20")  # cross-midnight return kept as a datetime

        states = {k: rows[k]["scanStatus"] for k in ("R02", "R03", "R04", "R07", "R09", "R11", "R19")}
        self.assertEqual(
            states,
            {
                "R02": "Returned",
                "R03": "Early Dismissal",
                "R04": "Not Returned",
                "R07": None,
                "R09": "Expired, Not Scanned",
                "R11": "Awaiting Return",
                "R19": "Awaiting Exit",
            },
        )
        self.assertIsNone(rows["R04"]["outsideMinutes"])  # never a running duration for an open pass

        cards = self.cards(body)
        self.assertEqual(cards["Total requests"], 18)
        self.assertEqual((cards["Approved"], cards["Rejected"], cards["Pending"]), (15, 1, 2))
        self.assertEqual((cards["Exited"], cards["Returned"], cards["Not returned"]), (12, 9, 2))
        self.assertEqual(cards["Avg time outside"], 40)  # 360 minutes over 9 returned passes
        self.assertEqual(body["totals"]["outsideMinutes"], 360)
        joined = " ".join(body["notes"])
        self.assertIn("By pass type: Official 10, Personal 6, Early Dismissal 1, Unspecified 1.", joined)
        self.assertIn("By origin: Manual 17, On-Duty 1.", joined)

    def test_row_values_add_up_to_totals_and_cards(self):
        body = self.run_report(self.RID, **RANGE)
        minutes = [r["outsideMinutes"] for r in body["rows"] if r["outsideMinutes"] is not None]
        self.assertEqual(sum(minutes), body["totals"]["outsideMinutes"])
        self.assertEqual(len(minutes), self.cards(body)["Returned"])
        statuses = [r["status"] for r in body["rows"]]
        cards = self.cards(body)
        self.assertEqual(
            (statuses.count("Approved"), statuses.count("Rejected"), statuses.count("Pending")),
            (cards["Approved"], cards["Rejected"], cards["Pending"]),
        )

    def test_every_filter_narrows(self):
        cases = {
            "approvalStatus=pending": ({"approvalStatus": "pending"}, {"R07", "R08"}),
            "approvalStatus=rejected": ({"approvalStatus": "rejected"}, {"R06"}),
            "passType=official": (
                {"passType": "official"},
                {"R01", "R04", "R06", "R08", "R09", "R11", "R13", "R15", "R17", "R19"},
            ),
            "passType=unspecified": ({"passType": "unspecified"}, {"R10"}),
            "passType=early": ({"passType": "early_dismissal"}, {"R03"}),
            "source=on_duty": ({"source": "on_duty"}, {"R10"}),
            "reviewerRole=dept_head": ({"reviewerRole": "dept_head"}, {"R02", "R05", "R12", "R16", "R18"}),
            "reviewerRole=system": ({"reviewerRole": "system"}, {"R10"}),
            "status completed": (
                {"scanStatus": "completed"},
                {"R01", "R02", "R05", "R12", "R13", "R15", "R16", "R17", "R18"},
            ),
            "status not_returned": ({"scanStatus": "not_returned"}, {"R04"}),
            "status early": ({"scanStatus": "early_dismissal"}, {"R03"}),
            "status pending_return": ({"scanStatus": "pending_return"}, {"R11"}),
            "status expired": ({"scanStatus": "expired_unscanned"}, {"R09", "R10"}),
            "status pending_exit": ({"scanStatus": "pending_exit"}, {"R19"}),
            "status not_applicable": ({"scanStatus": "not_applicable"}, {"R06", "R07", "R08"}),
            "gate=gate 2": ({"gate": "gate 2"}, {"R01", "R04", "R18"}),  # exit OR return gate, case-insensitive
            "gate=gate 3": ({"gate": "Gate 3"}, {"R05"}),
            "department": (
                {"departmentIds": str(self.cut.id)},
                {"R01", "R02", "R03", "R04", "R06", "R07", "R08", "R10", "R12", "R15", "R16", "R17", "R18"},
            ),
            "employee": ({"employeeIds": str(self.e3.id)}, {"R05", "R09", "R19"}),
            "production": ({"employmentType": "production"}, {"R03", "R04", "R08", "R15"}),
            "designation": (
                {"designationIds": str(self.des.id)},
                {"R01", "R02", "R06", "R07", "R10", "R16", "R17", "R18"},
            ),
            "inactive employees": ({"employeeStatus": "inactive"}, {"R12"}),
            "active employees": ({"employeeStatus": "active"}, IN_RANGE - {"R12"}),
            "branch b2": ({"branchIds": str(self.b2.id)}, {"R05", "R09", "R19"}),
            "one IST day": ({"dateFrom": "2026-09-08", "dateTo": "2026-09-08"}, {"R12", "R15", "R16"}),
            "day before (UTC trap)": ({"dateFrom": "2026-09-07", "dateTo": "2026-09-07"}, set()),
            "august": ({"dateFrom": "2026-08-01", "dateTo": "2026-08-31"}, {"R14"}),
        }
        for label, (extra, want) in cases.items():
            with self.subTest(label):
                params = {**RANGE, **extra}
                body = self.run_report(self.RID, **params)
                self.assertEqual(sorted(self.dests(body)), sorted(want))

    def test_branch_scoped_user_sees_own_branch_only_and_cannot_widen(self):
        body = self.run_report(self.RID, self.b1_user, **RANGE)
        got = set(self.dests(body))
        self.assertEqual(got, IN_RANGE - {"R05", "R09", "R19", "R13"})  # E3 is in Unit 2, E6 has no branch
        for widen in (
            {"branchIds": str(self.b2.id)},
            {"employeeIds": str(self.e3.id)},
            {"departmentIds": str(self.sew.id)},
        ):
            self.assertEqual(self.run_report(self.RID, self.b1_user, **RANGE, **widen)["rows"], [], widen)
        body2 = self.run_report(self.RID, self.b2_user, **RANGE)
        self.assertEqual(set(self.dests(body2)), {"R05", "R09", "R19"})
        self.assertEqual(self.cards(body2)["Total requests"], 3)

    def test_empty_result(self):
        body = self.run_report(self.RID, **EMPTY_RANGE)
        self.assertEqual(body["rows"], [])
        self.assertEqual(self.cards(body)["Total requests"], 0)
        self.assertIsNone(self.cards(body)["Avg time outside"])
        for fmt, sig in (("xlsx", b"PK"), ("pdf", b"%PDF")):
            r = self.export(self.RID, fmt, **EMPTY_RANGE)
            self.assertEqual(r.status_code, 200, fmt)
            self.assertTrue(r.content.startswith(sig), fmt)

    def test_python_filter_and_plain_query_both_honour_the_row_limit(self):
        with mock.patch("api.reporting.runner.SCREEN_ROW_LIMIT", 5):
            plain = self.run_report(self.RID, **RANGE)
            filtered = self.run_report(self.RID, scanStatus="completed", **RANGE)
        for body in (plain, filtered):
            self.assertEqual(body["rowCount"], 5)
            self.assertTrue(body["truncated"])
            self.assertIsNone(body["totals"])
        self.assertTrue(all(r["scanStatus"] == "Returned" for r in filtered["rows"]))

    def test_xlsx_round_trip(self):
        r = self.export(self.RID, "xlsx", **RANGE)
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:5], ["Requested", "Emp Code", "Employee", "Department", "Designation"])
        self.assertEqual(ws["A8"].value, datetime(2026, 9, 15, 9, 30))  # R08, newest
        self.assertEqual(ws["C8"].value, "Bala Kumar")
        last = [c.value for c in ws[8 + 18]]
        self.assertEqual(last[0], "TOTAL")

    def test_free_text_is_flattened_and_capped(self):
        mk_request(self.e7, "line one\nline two =SUM(A1)", "2026-09-14 12:00", status="pending", reason="x" * 500)
        body = self.run_report(self.RID, employeeIds=str(self.e7.id), **RANGE)
        row = body["rows"][0]
        self.assertEqual(row["destination"], "line one line two =SUM(A1)")
        self.assertEqual(len(row["reason"]), C.MAX_TEXT)
        self.assertTrue(row["reason"].endswith("..."))
        r = self.export(self.RID, "xlsx", employeeIds=str(self.e7.id), **RANGE)
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws["H8"].data_type, "s")  # text stays text, never a formula


# ── outpass-in-out-register ─────────────────────────────────────────────────


class InOutTests(World):
    RID = "outpass-in-out-register"
    EXITED = {"R01", "R02", "R03", "R04", "R05", "R11", "R12", "R13", "R15", "R16", "R17", "R18"}

    def test_golden_rows_summary_and_notes(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(
            self.dests(body),
            [
                "R11",
                "R18",
                "R17",
                "R13",
                "R16",
                "R12",
                "R15",
                "R05",
                "R04",
                "R03",
                "R02",
                "R01",
            ],  # by exit time, newest first
        )
        rows = self.by(body, "destination")
        r16 = rows["R16"]  # exits 23:50 IST on the 8th, back 00:20 on the 9th
        self.assertEqual(
            (r16["date"], r16["exitedAt"], r16["enteredAt"], r16["outsideMinutes"], r16["state"]),
            ("2026-09-08", "2026-09-08 23:50", "2026-09-09 00:20", 30, "Returned"),
        )
        self.assertEqual(rows["R15"]["date"], "2026-09-08")  # 00:45 IST is still the 8th (UTC would say the 7th)
        self.assertEqual(
            (rows["R11"]["state"], rows["R11"]["overrunMinutes"], rows["R11"]["outsideMinutes"]),
            ("Outside Now", 60, None),
        )
        self.assertEqual((rows["R04"]["state"], rows["R04"]["overrunMinutes"]), ("Not Returned", None))
        self.assertEqual(rows["R03"]["state"], "Early Dismissal")
        self.assertEqual(
            {k: rows[k]["overrunMinutes"] for k in ("R01", "R02", "R12", "R17", "R18", "R13")},
            {"R01": 20, "R02": 0, "R12": 15, "R17": 30, "R18": 20, "R13": None},
        )
        cards = self.cards(body)
        self.assertEqual(cards["Total exits"], 12)
        self.assertEqual((cards["Returned"], cards["Outside now"], cards["Not returned (earlier days)"]), (9, 1, 1))
        self.assertEqual(
            (cards["Total time outside"], cards["Avg time outside"], cards["Longest absence"]), (360, 40, 60)
        )
        self.assertEqual(body["totals"]["outsideMinutes"], 360)
        joined = " ".join(body["notes"])
        self.assertIn("Exits per gate: Gate 1 9, Gate 2 2, Gate 3 1.", joined)
        self.assertIn("Early dismissals in this list: 1.", joined)

    def test_every_filter_narrows(self):
        cases = {
            "returned": ({"returnState": "returned"}, self.EXITED - {"R03", "R04", "R11"}),
            "outside_now": ({"returnState": "outside_now"}, {"R11"}),
            "not_returned": ({"returnState": "not_returned"}, {"R04"}),
            "early_dismissal": ({"returnState": "early_dismissal"}, {"R03"}),
            "personal": ({"passType": "personal"}, {"R02", "R05", "R12", "R16", "R18"}),
            "gate 3": ({"gate": "gate 3"}, {"R05"}),
            "department": (
                {"departmentIds": str(self.cut.id)},
                {"R01", "R02", "R03", "R04", "R12", "R15", "R16", "R17", "R18"},
            ),
            "production": ({"employmentType": "production"}, {"R03", "R04", "R15"}),
            "employee": ({"employeeIds": str(self.e1.id)}, {"R01", "R02", "R16", "R17", "R18"}),
            "exit basis": ({"dateFrom": "2026-09-09", "dateTo": "2026-09-09"}, {"R13"}),  # R16 only RETURNED on the 9th
            "cross midnight day": ({"dateFrom": "2026-09-08", "dateTo": "2026-09-08"}, {"R12", "R15", "R16"}),
            "august exit": ({"dateFrom": "2026-08-01", "dateTo": "2026-08-31"}, {"R14"}),
        }
        for label, (extra, want) in cases.items():
            with self.subTest(label):
                body = self.run_report(self.RID, **{**RANGE, **extra})
                self.assertEqual(sorted(self.dests(body)), sorted(want))

    def test_branch_isolation(self):
        got = set(self.dests(self.run_report(self.RID, self.b1_user, **RANGE)))
        self.assertEqual(got, self.EXITED - {"R05", "R13"})
        self.assertEqual(self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["rows"], [])
        self.assertEqual(set(self.dests(self.run_report(self.RID, self.b2_user, **RANGE))), {"R05"})

    def test_empty_and_exports(self):
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Date", "Emp Code", "Employee"])
        self.assertEqual(_day(ws["A8"].value), datetime(2026, 9, 15).date())  # R11
        self.assertEqual(ws["C8"].value, "Esha Nair")
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-employee-summary ────────────────────────────────────────────────


class EmployeeSummaryTests(World):
    RID = "outpass-employee-summary"

    def test_golden_rows_and_totals(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(
            [r["employeeCode"] for r in body["rows"]], ["GT001", "GT002", "GT003", "GT004", "GT005", "GT006"]
        )
        rows = self.by(body, "employeeCode")
        keys = (
            "requests",
            "approved",
            "rejected",
            "pending",
            "exited",
            "returned",
            "notReturned",
            "expiredUnused",
            "official",
            "personal",
            "earlyDismissal",
            "unspecified",
            "totalOutsideMinutes",
            "avgOutsideMinutes",
            "maxOutsideMinutes",
            "lateReturns",
            "qrSubmissions",
        )
        want = {
            "GT001": (8, 6, 1, 1, 5, 5, 0, 1, 3, 4, 0, 1, 220, 44, 60, 3, 2),
            "GT002": (4, 3, 0, 1, 3, 1, 1, 0, 3, 0, 1, 0, 60, 60, 60, 0, 1),
            "GT003": (3, 3, 0, 0, 1, 1, 0, 1, 2, 1, 0, 0, 20, 20, 20, 0, 1),
            "GT004": (1, 1, 0, 0, 1, 1, 0, 0, 0, 1, 0, 0, 30, 30, 30, 1, 0),
            "GT005": (1, 1, 0, 0, 1, 0, 1, 0, 1, 0, 0, 0, None, None, None, 1, 0),
            "GT006": (1, 1, 0, 0, 1, 1, 0, 0, 1, 0, 0, 0, 30, 30, 30, 0, 0),
        }
        for code, values in want.items():
            self.assertEqual(tuple(rows[code][k] for k in keys), values, code)
        self.assertEqual((rows["GT005"]["department"], rows["GT005"]["designation"]), ("Unassigned", None))

        totals = body["totals"]
        self.assertEqual(
            tuple(totals[k] for k in keys),
            (18, 15, 1, 2, 12, 9, 2, 2, 10, 6, 1, 1, 360, 40, 60, 5, 4),
        )
        for k in (
            "requests",
            "approved",
            "rejected",
            "pending",
            "exited",
            "returned",
            "notReturned",
            "lateReturns",
            "qrSubmissions",
        ):
            self.assertEqual(totals[k], sum(r[k] for r in body["rows"]), k)
        cards = self.cards(body)
        self.assertEqual(
            (cards["Employees with a pass"], cards["Total requests"], cards["Approved"], cards["Rejected"]),
            (6, 18, 15, 1),
        )
        self.assertEqual((cards["Total time outside"], cards["Avg per returned pass"]), (360, 40))
        notes = " ".join(body["notes"])
        self.assertIn("Most passes: GT001 Asha (8); GT002 Bala (4); GT003 Chitra (3);", notes)

    def test_qr_mirror_rows_are_never_counted(self):
        body = self.run_report(self.RID, employeeIds=str(self.e1.id), **RANGE)
        self.assertEqual(body["rows"][0]["qrSubmissions"], 2)  # Q1 + Q2; the approval copy (Q6) and August (Q8) are out

    def test_every_filter_narrows(self):
        def codes(**extra):
            return [r["employeeCode"] for r in self.run_report(self.RID, **{**RANGE, **extra})["rows"]]

        self.assertEqual(codes(minPasses="2"), ["GT001", "GT002", "GT003"])
        self.assertEqual(codes(minPasses="5"), ["GT001"])
        self.assertEqual(codes(includeZero="true"), ["GT001", "GT002", "GT003", "GT004", "GT005", "GT006", "GT007"])
        self.assertEqual(codes(includeZero="true", departmentIds=str(self.sew.id)), ["GT003", "GT007"])
        self.assertEqual(codes(passType="official"), ["GT001", "GT002", "GT003", "GT005", "GT006"])
        self.assertEqual(codes(departmentIds=str(self.sew.id)), ["GT003"])
        self.assertEqual(codes(employmentType="production"), ["GT002"])
        self.assertEqual(codes(employeeIds=f"{self.e2.id},{self.e5.id}"), ["GT002", "GT005"])
        self.assertEqual(codes(employeeStatus="inactive"), ["GT004"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), ["GT003"])
        self.assertEqual(codes(designationIds=str(self.des.id)), ["GT001"])
        self.assertEqual(
            codes(dateFrom="2026-09-08", dateTo="2026-09-08"), ["GT001", "GT002", "GT004"]
        )  # one request each that IST day
        official = self.run_report(self.RID, passType="official", **RANGE)
        self.assertEqual(self.by(official, "employeeCode")["GT001"]["requests"], 3)
        zero = self.by(self.run_report(self.RID, includeZero="true", **RANGE), "employeeCode")["GT007"]
        self.assertEqual((zero["requests"], zero["totalOutsideMinutes"], zero["avgOutsideMinutes"]), (0, None, None))

    def test_branch_isolation(self):
        body = self.run_report(self.RID, self.b1_user, **RANGE)
        self.assertEqual([r["employeeCode"] for r in body["rows"]], ["GT001", "GT002", "GT004", "GT005"])
        self.assertEqual(body["totals"]["requests"], 14)
        self.assertEqual(self.run_report(self.RID, self.b1_user, employeeIds=str(self.e3.id), **RANGE)["rows"], [])
        zero = self.run_report(self.RID, self.b1_user, includeZero="true", **RANGE)
        self.assertNotIn("GT007", [r["employeeCode"] for r in zero["rows"]])  # Unit 2 employee stays invisible

    def test_empty_and_exports(self):
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][4:7], ["Requests", "Approved", "Rejected"])
        self.assertEqual((ws["A8"].value, ws["E8"].value), ("GT001", 8))
        self.assertEqual(ws.cell(row=8 + 6, column=1).value, "TOTAL")
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-department-summary ──────────────────────────────────────────────


class DepartmentSummaryTests(World):
    RID = "outpass-department-summary"
    KEYS = (
        "headcount",
        "employeesWithPass",
        "requests",
        "approved",
        "rejected",
        "exited",
        "returned",
        "notReturned",
        "totalOutsideMinutes",
        "avgOutsideMinutes",
        "passesPer100",
    )

    def test_group_by_department(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual([r["group"] for r in body["rows"]], ["CUTTING", "SEWING", "No Department"])
        rows = self.by(body, "group")
        self.assertEqual(tuple(rows["CUTTING"][k] for k in self.KEYS), (2, 3, 13, 10, 1, 9, 7, 1, 310, 44, 650.0))
        self.assertEqual(tuple(rows["SEWING"][k] for k in self.KEYS), (2, 1, 3, 3, 0, 1, 1, 0, 20, 20, 150.0))
        self.assertEqual(tuple(rows["No Department"][k] for k in self.KEYS), (2, 2, 2, 2, 0, 2, 1, 1, 30, 30, 100.0))
        self.assertEqual(
            (rows["CUTTING"]["branch"], rows["SEWING"]["branch"], rows["No Department"]["branch"]),
            ("Unit 1", "Unit 2", None),
        )
        totals = body["totals"]
        self.assertEqual(tuple(totals[k] for k in self.KEYS), (6, 6, 18, 15, 1, 12, 9, 2, 360, 40, 300.0))
        for k in ("headcount", "requests", "approved", "exited", "returned", "notReturned"):
            self.assertEqual(totals[k], sum(r[k] for r in body["rows"]), k)
        cards = self.cards(body)
        self.assertEqual(cards["Most passes"], "CUTTING (13)")
        self.assertEqual(cards["Most time outside"], "CUTTING (5h 10m)")

    def test_group_by_branch_staff_type_and_pass_type(self):
        body = self.run_report(self.RID, groupBy="branch", **RANGE)
        rows = self.by(body, "group")
        self.assertEqual([r["group"] for r in body["rows"]], ["Unit 1", "Unit 2", "No Branch"])
        self.assertEqual(
            tuple(
                rows["Unit 1"][k]
                for k in (
                    "headcount",
                    "employeesWithPass",
                    "requests",
                    "approved",
                    "notReturned",
                    "totalOutsideMinutes",
                )
            ),
            (3, 4, 14, 11, 2, 310),
        )
        self.assertEqual((rows["Unit 2"]["requests"], rows["Unit 2"]["headcount"]), (3, 2))
        self.assertEqual((rows["No Branch"]["requests"], rows["No Branch"]["headcount"]), (1, 1))

        body = self.run_report(self.RID, groupBy="employmentType", **RANGE)
        rows = self.by(body, "group")
        self.assertEqual((rows["Staff"]["requests"], rows["Staff"]["headcount"]), (14, 5))
        self.assertEqual((rows["Production"]["requests"], rows["Production"]["headcount"]), (4, 1))

        body = self.run_report(self.RID, groupBy="passType", **RANGE)
        rows = self.by(body, "group")
        self.assertEqual(
            {k: v["requests"] for k, v in rows.items()},
            {"Official": 10, "Personal": 6, "Early Dismissal": 1, "Unspecified": 1},
        )
        self.assertTrue(all(r["headcount"] is None and r["passesPer100"] is None for r in body["rows"]))
        self.assertIsNone(body["totals"]["passesPer100"])
        # One person can hold several pass types, so the rows add up to 10 but only 6 employees took a pass.
        self.assertEqual(
            {k: v["employeesWithPass"] for k, v in rows.items()},
            {"Official": 5, "Personal": 3, "Early Dismissal": 1, "Unspecified": 1},
        )
        self.assertEqual(body["totals"]["employeesWithPass"], 6)
        self.assertEqual(self.cards(body)["Employees with a pass"], 6)

    def test_filters_and_isolation(self):
        body = self.run_report(self.RID, departmentIds=str(self.sew.id), **RANGE)
        self.assertEqual([r["group"] for r in body["rows"]], ["SEWING"])
        body = self.run_report(self.RID, employmentType="production", **RANGE)
        self.assertEqual([(r["group"], r["requests"]) for r in body["rows"]], [("CUTTING", 4)])
        body = self.run_report(self.RID, branchIds=str(self.b2.id), **RANGE)
        self.assertEqual([r["group"] for r in body["rows"]], ["SEWING"])
        body = self.run_report(self.RID, dateFrom="2026-09-08", dateTo="2026-09-08")
        self.assertEqual(self.by(body, "group")["CUTTING"]["requests"], 3)

        scoped = self.run_report(self.RID, self.b1_user, **RANGE)
        self.assertEqual([r["group"] for r in scoped["rows"]], ["CUTTING", "No Department"])
        self.assertEqual(self.by(scoped, "group")["No Department"]["headcount"], 1)  # only E5; E6 has no branch
        self.assertEqual(self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["rows"], [])

    def test_empty_and_exports(self):
        body = self.run_report(self.RID, dateFrom="2020-01-01", dateTo="2020-01-02", employmentType="production")
        self.assertEqual(
            [(r["group"], r["requests"], r["headcount"]) for r in body["rows"]], [("CUTTING", 0, 1)]
        )  # headcount stays
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Group", "Branch", "Active Headcount"])
        self.assertEqual((ws["A8"].value, ws["E8"].value), ("CUTTING", 13))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-not-returned ────────────────────────────────────────────────────


class NotReturnedTests(World):
    RID = "outpass-not-returned"

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(self.dests(body), ["R11", "R04"])  # outside now first, then earlier days
        rows = self.by(body, "destination")
        r11, r4 = rows["R11"], rows["R04"]
        self.assertEqual(
            (
                r11["employeeCode"],
                r11["phone"],
                r11["bucket"],
                r11["minutesOutside"],
                r11["overdueMinutes"],
                r11["returnQrState"],
            ),
            ("GT005", "9000000005", "Outside Now", 90, 60, "Valid"),
        )
        self.assertEqual(
            (r4["bucket"], r4["minutesOutside"], r4["overdueMinutes"], r4["returnQrState"], r4["phone"]),
            ("Not Returned", None, None, "Not generated", None),
        )
        self.assertEqual(r4["exitedAt"], "2026-09-04 11:10")
        cards = self.cards(body)
        self.assertEqual(cards["Outside now (exited today)"], 1)
        self.assertEqual(cards["Return never recorded (earlier days)"], 1)
        self.assertEqual(cards["Past expected return time"], 2)
        self.assertEqual(cards["Report as of (IST)"], "2026-09-15 10:00")
        self.assertIn("Snapshot as of 2026-09-15 10:00 IST", " ".join(body["notes"]))

    def test_filters(self):
        def dests(**extra):
            return self.dests(self.run_report(self.RID, **{**RANGE, **extra}))

        self.assertEqual(dests(includeEarlyDismissal="true"), ["R11", "R04", "R03"])
        self.assertEqual(
            self.by(self.run_report(self.RID, includeEarlyDismissal="true", **RANGE), "destination")["R03"]["bucket"],
            "Early Dismissal",
        )
        self.assertEqual(dests(minOutsideMinutes="120"), ["R04"])  # R11 has only been out 90 minutes
        self.assertEqual(dests(minOutsideMinutes="60"), ["R11", "R04"])
        self.assertEqual(dests(passType="official"), ["R11", "R04"])
        self.assertEqual(dests(passType="personal"), [])
        self.assertEqual(dests(departmentIds=str(self.cut.id)), ["R04"])
        self.assertEqual(dests(employmentType="production"), ["R04"])
        self.assertEqual(dests(employeeIds=str(self.e5.id)), ["R11"])
        self.assertEqual(dests(dateFrom="2026-09-15", dateTo="2026-09-15"), ["R11"])
        self.assertEqual(dests(dateFrom="2026-09-04", dateTo="2026-09-04"), ["R04"])

    def test_isolation_and_exports(self):
        self.assertEqual(set(self.dests(self.run_report(self.RID, self.b1_user, **RANGE))), {"R11", "R04"})
        self.assertEqual(self.run_report(self.RID, self.b2_user, **RANGE)["rows"], [])
        self.assertEqual(self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["rows"], [])
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual(ws["A8"].value, "GT005")
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-late-return ─────────────────────────────────────────────────────


class LateReturnTests(World):
    RID = "outpass-late-return"

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(self.dests(body), ["R11", "R18", "R17", "R12", "R01"])  # newest exit first
        overs = [(r["employeeCode"], r["overrunMinutes"], r["state"]) for r in body["rows"]]
        self.assertEqual(
            overs,
            [
                ("GT005", 60, "Outside Now"),
                ("GT001", 20, "Returned Late"),
                ("GT001", 30, "Returned Late"),
                ("GT004", 15, "Returned Late"),
                ("GT001", 20, "Returned Late"),
            ],
        )
        self.assertEqual(body["totals"]["overrunMinutes"], 145)
        cards = self.cards(body)
        self.assertEqual(cards["Late returns"], 5)
        self.assertEqual((cards["Average overrun"], cards["Longest overrun"]), (29, 60))
        self.assertEqual(cards["Employees with 3+ late returns"], 1)  # GT001: R01 + R17 + R18
        joined = " ".join(body["notes"])
        self.assertIn("1 pass(es) with an expected return time were scanned out on an earlier day", joined)
        self.assertIn("Outpass Not Returned", joined)

    def test_boundaries_and_filters(self):
        body = self.run_report(self.RID, minOverrunMinutes="30", **RANGE)
        self.assertEqual(
            [(r["employeeCode"], r["overrunMinutes"]) for r in body["rows"]], [("GT005", 60), ("GT001", 30)]
        )
        body = self.run_report(self.RID, minOverrunMinutes="60", **RANGE)
        self.assertEqual([r["employeeCode"] for r in body["rows"]], ["GT005"])
        self.assertEqual(
            [r["employeeCode"] for r in self.run_report(self.RID, employeeIds=str(self.e4.id), **RANGE)["rows"]],
            ["GT004"],
        )
        self.assertEqual(
            [r["employeeCode"] for r in self.run_report(self.RID, departmentIds=str(self.cut.id), **RANGE)["rows"]],
            ["GT001", "GT001", "GT004", "GT001"],
        )
        self.assertEqual(self.run_report(self.RID, employmentType="production", **RANGE)["rows"], [])
        self.assertEqual(
            len(self.run_report(self.RID, dateFrom="2026-09-12", dateTo="2026-09-13")["rows"]), 2
        )  # R17, R18
        # R02 came back 5 minutes BEFORE its expected time; a return within 30 seconds after is not "late" either.
        self.assertNotIn(
            ("GT001", 0), [(r["employeeCode"], r["overrunMinutes"]) for r in self.run_report(self.RID, **RANGE)["rows"]]
        )

    def test_isolation_and_exports(self):
        self.assertEqual(len(self.run_report(self.RID, self.b1_user, **RANGE)["rows"]), 5)
        self.assertEqual(self.run_report(self.RID, self.b2_user, **RANGE)["rows"], [])
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual(ws["A8"].value, "GT005")
        self.assertAlmostEqual(ws["I8"].value.total_seconds(), 3600, delta=1)  # "Late By" is stored as a duration
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-approval-turnaround ─────────────────────────────────────────────


class TurnaroundTests(World):
    RID = "outpass-approval-turnaround"
    KEYS = (
        "approverRole",
        "decisions",
        "approved",
        "rejected",
        "avgTurnaroundMinutes",
        "maxTurnaroundMinutes",
        "within15MinPct",
    )

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual([r["approverName"] for r in body["rows"]], ["Priya HR", "Mani HOD", "Selvi HOD"])
        rows = self.by(body, "approverName")
        self.assertEqual(tuple(rows["Priya HR"][k] for k in self.KEYS), ("HR", 10, 9, 1, 8, 20, 88.89))
        self.assertEqual(tuple(rows["Mani HOD"][k] for k in self.KEYS), ("Dept Head", 4, 4, 0, 18, 30, 50.0))
        self.assertEqual(tuple(rows["Selvi HOD"][k] for k in self.KEYS), ("Dept Head", 1, 1, 0, 90, 90, 0.0))
        totals = body["totals"]
        self.assertEqual(
            (
                totals["decisions"],
                totals["approved"],
                totals["rejected"],
                totals["avgTurnaroundMinutes"],
                totals["maxTurnaroundMinutes"],
                totals["within15MinPct"],
            ),
            (15, 14, 1, 17, 90, 71.43),
        )
        cards = self.cards(body)
        self.assertEqual(cards["Decisions"], 15)
        self.assertEqual((cards["Approved after 60+ min"], cards["Still pending"]), (1, 2))
        self.assertIn("On-Duty passes are created already approved", " ".join(body["notes"]))

    def test_filters_and_isolation(self):
        body = self.run_report(self.RID, reviewerRole="dept_head", **RANGE)
        self.assertEqual([r["approverName"] for r in body["rows"]], ["Mani HOD", "Selvi HOD"])
        self.assertEqual(body["totals"]["decisions"], 5)
        body = self.run_report(self.RID, self.b1_user, **RANGE)
        rows = self.by(body, "approverName")
        self.assertEqual(set(rows), {"Priya HR", "Mani HOD"})
        self.assertEqual(tuple(rows["Priya HR"][k] for k in self.KEYS), ("HR", 7, 6, 1, 7, 10, 100.0))
        body = self.run_report(self.RID, self.b2_user, **RANGE)
        rows = self.by(body, "approverName")
        self.assertEqual(tuple(rows["Priya HR"][k] for k in self.KEYS), ("HR", 2, 2, 0, 15, 20, 50.0))
        self.assertEqual(self.cards(body)["Still pending"], 0)
        self.assertEqual(self.run_report(self.RID, departmentIds=str(self.sew.id), **RANGE)["totals"]["decisions"], 3)
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])

    def test_exports(self):
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Approver", "Role", "Decisions"])
        self.assertEqual((ws["A8"].value, ws["C8"].value), ("Priya HR", 10))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-pending-requests ────────────────────────────────────────────────


class PendingTests(World):
    RID = "outpass-pending-requests"

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(self.dests(body), ["R07", "R08"])  # oldest first
        rows = self.by(body, "destination")
        r7, r8 = rows["R07"], rows["R08"]
        self.assertEqual(
            (r7["ageMinutes"], r7["ageBand"], r7["employeeCode"], r7["reportingManager"]),
            (1560, "Overdue", "GT001", "Esha Nair"),
        )
        self.assertEqual((r8["ageMinutes"], r8["ageBand"], r8["reportingManager"]), (30, "Under 1 Hour", None))
        cards = self.cards(body)
        self.assertEqual(
            (cards["Pending requests"], cards["Oldest waiting"], cards["Waiting over 1 day"]), (2, 1560, 1)
        )

    def test_filters(self):
        def dests(**extra):
            return self.dests(self.run_report(self.RID, **{**RANGE, **extra}))

        self.assertEqual(dests(minAge="60"), ["R07"])
        self.assertEqual(dests(minAge="30"), ["R07", "R08"])  # exactly 30 minutes counts
        self.assertEqual(dests(minAge="1440"), ["R07"])
        self.assertEqual(dests(employeeIds=str(self.e2.id)), ["R08"])
        self.assertEqual(dests(employmentType="production"), ["R08"])
        self.assertEqual(dests(departmentIds=str(self.sew.id)), [])
        body = self.run_report(self.RID, dateFrom="2026-09-15", dateTo="2026-09-15")
        self.assertEqual(self.dests(body), ["R08"])
        self.assertIn("1 more pending request(s) were created before the selected dates", " ".join(body["notes"]))

    def test_isolation_and_exports(self):
        self.assertEqual(self.dests(self.run_report(self.RID, self.b1_user, **RANGE)), ["R07", "R08"])
        self.assertEqual(self.run_report(self.RID, self.b2_user, **RANGE)["rows"], [])
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual(ws["A8"].value, datetime(2026, 9, 14, 8, 0))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-expired-unused ──────────────────────────────────────────────────


class ExpiredUnusedTests(World):
    RID = "outpass-expired-unused"

    def test_default_origin_is_manual(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(self.dests(body), ["R09"])
        self.assertEqual([r["employeeCode"] for r in body["rows"]], ["GT003"])
        row = body["rows"][0]
        self.assertEqual(
            (
                row["source"],
                row["approvedAt"],
                row["expiredAt"],
                row["reviewedBy"],
                row["reviewerRole"],
                row["passType"],
            ),
            ("Manual", "2026-09-10 10:20", "2026-09-10 11:20", f"HR {DASH} Priya HR", "HR", "Official"),
        )
        cards = self.cards(body)
        self.assertEqual((cards["Approved but never used"], cards["Approved passes in period"]), (1, 14))
        self.assertEqual(cards["Share never used"], 7.14)
        self.assertIn("that is why the default origin is Manual", " ".join(body["notes"]))

    def test_origin_role_and_window_filters(self):
        body = self.run_report(self.RID, source="all", **RANGE)
        self.assertEqual(
            [(r["employeeCode"], r["source"]) for r in body["rows"]], [("GT001", "On-Duty"), ("GT003", "Manual")]
        )
        cards = self.cards(body)
        self.assertEqual(
            (cards["Approved but never used"], cards["Approved passes in period"], cards["Share never used"]),
            (2, 15, 13.33),
        )
        self.assertIn("By origin: Manual 1, On-Duty 1.", " ".join(body["notes"]))
        body = self.run_report(self.RID, source="on_duty", **RANGE)
        self.assertEqual([r["employeeCode"] for r in body["rows"]], ["GT001"])
        self.assertEqual(self.cards(body)["Share never used"], 100.0)
        self.assertEqual(
            self.run_report(self.RID, source="all", reviewerRole="system", **RANGE)["rows"][0]["reviewerRole"], "System"
        )
        # R19 was approved 30 minutes ago: still valid, so it is never listed.
        self.assertNotIn(
            "R19", json.dumps(self.run_report(self.RID, source="all", employeeIds=str(self.e3.id), **RANGE))
        )
        self.assertEqual(
            [
                r["employeeCode"]
                for r in self.run_report(self.RID, source="all", employeeIds=str(self.e3.id), **RANGE)["rows"]
            ],
            ["GT003"],
        )
        self.assertEqual(
            self.run_report(self.RID, source="all", dateFrom="2026-09-11", dateTo="2026-09-11")["rows"][0][
                "employeeCode"
            ],
            "GT001",
        )
        self.assertEqual(
            self.run_report(self.RID, source="all", departmentIds=str(self.cut.id), **RANGE)["rows"][0]["employeeCode"],
            "GT001",
        )

    def test_isolation_and_exports(self):
        self.assertEqual(
            self.run_report(self.RID, self.b1_user, **RANGE)["rows"], []
        )  # the manual one belongs to Unit 2
        body = self.run_report(self.RID, self.b1_user, source="all", **RANGE)
        self.assertEqual([r["employeeCode"] for r in body["rows"]], ["GT001"])
        self.assertEqual(
            [r["employeeCode"] for r in self.run_report(self.RID, self.b2_user, **RANGE)["rows"]], ["GT003"]
        )
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual((ws["A8"].value, ws["G8"].value), ("GT003", datetime(2026, 9, 10, 10, 20)))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-purpose-analysis ────────────────────────────────────────────────


class PurposeTests(World):
    RID = "outpass-purpose-analysis"
    KEYS = (
        "requests",
        "approved",
        "rejected",
        "pending",
        "approvalRatePct",
        "returned",
        "avgOutsideMinutes",
        "totalOutsideMinutes",
        "sharePct",
    )

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(
            [(r["passType"], r["origin"]) for r in body["rows"]],
            [
                ("Official", "Manual"),
                ("Personal", "Manual"),
                ("Early Dismissal", "Manual"),
                ("Unspecified", "On-Duty"),
            ],
        )
        rows = self.by(body, "passType")
        self.assertEqual(tuple(rows["Official"][k] for k in self.KEYS), (10, 8, 1, 1, 88.89, 4, 53, 210, 55.56))
        self.assertEqual(tuple(rows["Personal"][k] for k in self.KEYS), (6, 5, 0, 1, 100.0, 5, 30, 150, 33.33))
        self.assertEqual(tuple(rows["Early Dismissal"][k] for k in self.KEYS), (1, 1, 0, 0, 100.0, 0, None, None, 5.56))
        self.assertEqual(tuple(rows["Unspecified"][k] for k in self.KEYS), (1, 1, 0, 0, 100.0, 0, None, None, 5.56))
        totals = body["totals"]
        self.assertEqual(
            (
                totals["requests"],
                totals["approved"],
                totals["rejected"],
                totals["pending"],
                totals["returned"],
                totals["approvalRatePct"],
                totals["totalOutsideMinutes"],
                totals["avgOutsideMinutes"],
            ),
            (18, 15, 1, 2, 9, 93.75, 360, 40),
        )
        cards = self.cards(body)
        self.assertEqual(
            (cards["Total requests"], cards["Most common pass type"], cards["Overall approval rate"]),
            (18, "Official (10)", 93.75),
        )

    def test_filters_and_isolation(self):
        body = self.run_report(self.RID, departmentIds=str(self.sew.id), **RANGE)
        self.assertEqual(body["totals"]["requests"], 3)
        self.assertEqual(self.run_report(self.RID, employmentType="production", **RANGE)["totals"]["requests"], 4)
        self.assertEqual(self.run_report(self.RID, branchIds=str(self.b2.id), **RANGE)["totals"]["requests"], 3)
        self.assertEqual(self.run_report(self.RID, dateFrom="2026-09-08", dateTo="2026-09-08")["totals"]["requests"], 3)
        self.assertEqual(self.run_report(self.RID, self.b1_user, **RANGE)["totals"]["requests"], 14)
        self.assertEqual(self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["rows"], [])
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])

    def test_exports(self):
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Pass Type", "Origin", "Requests"])
        self.assertEqual((ws["A8"].value, ws["C8"].value), ("Official", 10))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── outpass-qr-submissions ──────────────────────────────────────────────────


class QrSubmissionTests(World):
    RID = "outpass-qr-submissions"

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual([r["destination"] for r in body["rows"]], ["Q5", "Q4", "Q3", "Q2", "Q1", "Q7"])
        rows = self.by(body, "destination")
        self.assertEqual(
            {
                k: rows["Q1"][k]
                for k in (
                    "submittedAt",
                    "enteredCode",
                    "enteredName",
                    "matchedEmployee",
                    "department",
                    "designation",
                    "branch",
                    "matchStatus",
                    "source",
                )
            },
            {
                "submittedAt": "2026-09-10 10:00",
                "enteredCode": "GT001",
                "enteredName": "Asha Rao",
                "matchedEmployee": "GT001 - Asha Rao",
                "department": "CUTTING",
                "designation": "Supervisor",
                "branch": "Unit 1",
                "matchStatus": "Matched",
                "source": "QR Form",
            },
        )
        self.assertEqual(
            (rows["Q2"]["matchStatus"], rows["Q2"]["enteredCode"]), ("Name Mismatch", "gt001")
        )  # typed name "Asha"
        self.assertEqual(
            (rows["Q3"]["matchStatus"], rows["Q3"]["matchedEmployee"], rows["Q3"]["department"]),
            ("Unmatched", None, None),
        )
        self.assertEqual(rows["Q4"]["matchStatus"], "Matched")  # double space in the typed name is ignored
        self.assertEqual((rows["Q5"]["branch"], rows["Q5"]["matchStatus"]), (None, "Unmatched"))
        cards = self.cards(body)
        self.assertEqual(
            (cards["Submissions"], cards["Matched to an employee"], cards["Name mismatch"], cards["Unmatched"]),
            (6, 3, 1, 2),
        )
        joined = " ".join(body["notes"])
        self.assertIn("By QR branch: Unit 1 4, No branch 1, Unit 2 1.", joined)
        self.assertIn("By department (submissions linked to an employee): CUTTING 3, SEWING 1.", joined)
        self.assertIn("Most submissions: GT001 Asha (2)", joined)
        self.assertIn("excluded here by default", joined)

    def test_mirror_rows_and_filters(self):
        def dests(**extra):
            return sorted(r["destination"] for r in self.run_report(self.RID, **{**RANGE, **extra})["rows"])

        self.assertEqual(dests(source="all"), ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7"])
        self.assertEqual(dests(source="request"), ["Q6"])
        body = self.run_report(self.RID, source="request", **RANGE)
        self.assertEqual((body["rows"][0]["source"], body["rows"][0]["matchStatus"]), ("Approved Pass", "Matched"))
        self.assertEqual(dests(matchStatus="name_mismatch"), ["Q2"])
        self.assertEqual(dests(matchStatus="unmatched"), ["Q3", "Q5"])
        self.assertEqual(dests(matchStatus="matched"), ["Q1", "Q4", "Q7"])
        self.assertEqual(dests(departmentIds=str(self.cut.id)), ["Q1", "Q2", "Q7"])
        self.assertIn(
            "unmatched submissions are hidden",
            " ".join(self.run_report(self.RID, departmentIds=str(self.cut.id), **RANGE)["notes"]),
        )
        self.assertEqual(dests(employeeIds=str(self.e1.id)), ["Q1", "Q2"])
        self.assertEqual(dests(employmentType="production"), ["Q7"])
        self.assertEqual(dests(designationIds=str(self.des.id)), ["Q1", "Q2"])
        self.assertEqual(dests(branchIds=str(self.b2.id)), ["Q4"])  # the QR's branch, not the employee's
        self.assertEqual(dests(dateFrom="2026-09-08", dateTo="2026-09-08"), ["Q7"])  # 00:30 IST is still the 8th
        self.assertEqual(dests(dateFrom="2026-09-07", dateTo="2026-09-07"), [])
        self.assertEqual(dests(dateFrom="2026-08-01", dateTo="2026-08-31"), ["Q8"])

    def test_branch_isolation(self):
        self.assertEqual(
            sorted(r["destination"] for r in self.run_report(self.RID, self.b1_user, **RANGE)["rows"]),
            ["Q1", "Q2", "Q3", "Q7"],
        )
        self.assertEqual(self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["rows"], [])
        self.assertEqual([r["destination"] for r in self.run_report(self.RID, self.b2_user, **RANGE)["rows"]], ["Q4"])
        # a submission on a QR whose branch was deleted belongs to nobody: unscoped users only
        self.assertNotIn(
            "Q5", [r["destination"] for r in self.run_report(self.RID, self.b1_user, source="all", **RANGE)["rows"]]
        )

    def test_empty_and_exports(self):
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Submitted", "Code Entered", "Name Entered"])
        self.assertEqual((ws["A8"].value, ws["B8"].value), (datetime(2026, 9, 14, 9, 0), "X1"))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── gate-scan-audit-log ─────────────────────────────────────────────────────


class ScanAuditTests(World):
    RID = "gate-scan-audit-log"
    ALL = ["S10", "S4", "S6", "S7", "S5", "S9", "S8", "S12", "S11", "S2", "S3", "S1"]  # newest first

    def ids_of(self, body):
        by_time = {
            ist(t).strftime("%Y-%m-%d %H:%M"): k
            for k, t in {
                "S10": "2026-09-14 08:30",
                "S4": "2026-09-10 12:00",
                "S6": "2026-09-08 00:30",
                "S7": "2026-09-07 23:00",
                "S5": "2026-09-06 09:30",
                "S9": "2026-09-05 12:00",
                "S8": "2026-09-05 11:40",
                "S12": "2026-09-04 11:10",
                "S11": "2026-09-01 11:00",
                "S2": "2026-09-01 10:20",
                "S3": "2026-09-01 09:45",
                "S1": "2026-09-01 09:20",
            }.items()
        }
        return [by_time[r["scannedAt"]] for r in body["rows"]]

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(self.ids_of(body), self.ALL)
        rows = {k: r for k, r in zip(self.ids_of(body), body["rows"])}
        s1 = rows["S1"]
        self.assertEqual(
            {
                k: s1[k]
                for k in (
                    "gate",
                    "gateBranch",
                    "scanType",
                    "result",
                    "employeeCode",
                    "employeeName",
                    "department",
                    "passType",
                    "destination",
                    "detail",
                )
            },
            {
                "gate": "Gate 1",
                "gateBranch": "Unit 1",
                "scanType": "Exit",
                "result": "Success",
                "employeeCode": "GT001",
                "employeeName": "Asha Rao",
                "department": "CUTTING",
                "passType": "Official",
                "destination": "R01",
                "detail": "Exit recorded",
            },
        )
        self.assertEqual((rows["S2"]["scanType"], rows["S2"]["detail"]), ("Return", "Return recorded"))
        # already_scanned: never the stored message (its clock time is UTC); times come from the pass, in IST
        self.assertEqual(rows["S3"]["detail"], "Pass already used - exited at 01-Sep-2026 09:20 via Gate 1")
        self.assertNotIn("03:50", rows["S3"]["detail"])
        self.assertEqual(rows["S11"]["detail"], "Already returned at 01-Sep-2026 10:20 via Gate 2")
        self.assertEqual(rows["S4"]["detail"], "Pass expired at 10-Sep-2026 11:20 (valid 60 minutes from approval)")
        self.assertEqual(rows["S5"]["detail"], "Request was not approved (pending or rejected)")
        s6, s7 = rows["S6"], rows["S7"]
        self.assertEqual(
            (s6["result"], s6["employeeCode"], s6["passType"], s6["detail"]),
            ("Invalid QR", None, None, "This QR code is not recognized."),
        )
        self.assertEqual((s7["gate"], s7["gateBranch"]), ("Unknown gate", None))
        cards = self.cards(body)
        self.assertEqual((cards["Scan attempts"], cards["Successful"], cards["Denied"]), (12, 5, 7))
        self.assertEqual((cards["Denial rate"], cards["Employees with 3+ denied attempts"]), (58.33, 1))
        joined = " ".join(body["notes"])
        self.assertIn("Denied by reason: Already Scanned 2, Invalid QR 2, Not Approved 2, Expired 1.", joined)
        self.assertIn("Denial rate by gate: Unknown gate 1 of 1 (100.0%); Gate 1 6 of 8 (75.0%).", joined)

    def test_every_filter_narrows(self):
        cases = {
            "already": ({"result": "already_scanned"}, {"S3", "S11"}),
            "expired": ({"result": "expired"}, {"S4"}),
            "success": ({"result": "success"}, {"S1", "S2", "S8", "S9", "S12"}),
            "return leg": ({"scanType": "entry"}, {"S2", "S9", "S11"}),
            "gate 3": ({"gate": "gate 3"}, {"S8", "S9"}),
            "employee": ({"employeeIds": str(self.e1.id)}, {"S1", "S2", "S3", "S5", "S10", "S11"}),
            "department": ({"departmentIds": str(self.cut.id)}, {"S1", "S2", "S3", "S5", "S10", "S11", "S12"}),
            "sewing": ({"departmentIds": str(self.sew.id)}, {"S4", "S8", "S9"}),
            "production": ({"employmentType": "production"}, {"S12"}),
            "branch b2": ({"branchIds": str(self.b2.id)}, {"S4", "S8", "S9"}),  # gate OR employee is in Unit 2
            "IST day": ({"dateFrom": "2026-09-08", "dateTo": "2026-09-08"}, {"S6"}),
            "day before": ({"dateFrom": "2026-09-07", "dateTo": "2026-09-07"}, {"S7"}),
        }
        for label, (extra, want) in cases.items():
            with self.subTest(label):
                body = self.run_report(self.RID, **{**RANGE, **extra})
                self.assertEqual(set(self.ids_of(body)), want)
                self.assertEqual(len(body["rows"]), len(want))

    def test_a_stale_return_qr_keeps_its_own_message_and_leg(self):
        # A superseded return QR is logged as invalid_qr but IS tied to a pass and an employee (unlike a forged QR).
        text = "This return QR is no longer valid -a newer one was generated for this Outpass."
        mk_scan(self.g1, self.R["R01"], self.e1, "entry", "invalid_qr", "2026-09-01 12:00", text)
        body = self.run_report(self.RID, result="invalid_qr", employeeIds=str(self.e1.id), **RANGE)
        self.assertEqual(len(body["rows"]), 1)
        row = body["rows"][0]
        self.assertEqual(
            (row["scanType"], row["result"], row["employeeCode"], row["destination"], row["detail"]),
            ("Return", "Invalid QR", "GT001", "R01", text),
        )

    def test_branch_isolation(self):
        b1 = self.run_report(self.RID, self.b1_user, **RANGE)
        self.assertEqual(
            set(self.ids_of(b1)), set(self.ALL) - {"S7", "S8", "S9"}
        )  # S7: deleted gate, nobody to attribute it to
        self.assertEqual(self.cards(b1)["Scan attempts"], 9)
        widened = self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)
        self.assertEqual(set(self.ids_of(widened)), {"S4"})  # only a row that is Unit 1's own (its gate) can appear
        self.assertEqual(set(self.ids_of(self.run_report(self.RID, self.b2_user, **RANGE))), {"S4", "S8", "S9"})

    def test_empty_and_exports(self):
        self.assertEqual(self.run_report(self.RID, **EMPTY_RANGE)["rows"], [])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:3], ["Scanned", "Gate", "Gate Branch"])
        self.assertEqual((ws["A8"].value, ws["B8"].value), (datetime(2026, 9, 14, 8, 30), "Gate 1"))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── gate-activity-summary ───────────────────────────────────────────────────


class GateActivityTests(World):
    RID = "gate-activity-summary"
    KEYS = (
        "exitScans",
        "returnScans",
        "alreadyScanned",
        "expired",
        "notApproved",
        "notExited",
        "invalidQr",
        "teaOut",
        "teaIn",
        "totalScans",
        "deniedScans",
        "denialPct",
    )

    def test_golden(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(
            [r["gate"] for r in body["rows"]], ["Gate 1", "Gate 3", "Gate 2", "Unknown / removed gate", "Gate 4"]
        )
        rows = self.by(body, "gate")
        self.assertEqual(tuple(rows["Gate 1"][k] for k in self.KEYS), (2, 0, 2, 1, 2, 0, 1, 7, 6, 21, 6, 75.0))
        self.assertEqual(tuple(rows["Gate 3"][k] for k in self.KEYS), (1, 1, 0, 0, 0, 0, 0, 1, 1, 4, 0, 0.0))
        self.assertEqual(tuple(rows["Gate 2"][k] for k in self.KEYS), (0, 1, 0, 0, 0, 0, 0, 1, 1, 3, 0, 0.0))
        self.assertEqual(
            tuple(rows["Unknown / removed gate"][k] for k in self.KEYS), (0, 0, 0, 0, 0, 0, 1, 0, 0, 1, 1, 100.0)
        )
        self.assertEqual(tuple(rows["Gate 4"][k] for k in self.KEYS), (0,) * 11 + (None,))
        self.assertEqual(
            (rows["Gate 1"]["branch"], rows["Gate 1"]["isActive"], rows["Gate 4"]["isActive"]),
            ("Unit 1", "Active", "Inactive"),
        )
        self.assertEqual(rows["Gate 4"]["lastLoginAt"], "2026-09-14 08:00")
        self.assertEqual(
            (rows["Unknown / removed gate"]["branch"], rows["Unknown / removed gate"]["isActive"]), (None, None)
        )
        totals = body["totals"]
        self.assertEqual(tuple(totals[k] for k in self.KEYS), (3, 2, 2, 1, 2, 0, 2, 9, 8, 29, 7, 58.33))
        for k in self.KEYS[:-1]:
            self.assertEqual(totals[k], sum(r[k] for r in body["rows"]), k)
        cards = self.cards(body)
        self.assertEqual(
            (cards["Total scans"], cards["Outpass scan attempts"], cards["Denied"], cards["Denial rate"]),
            (29, 12, 7, 58.33),
        )
        self.assertEqual(cards["Busiest gate"], "Gate 1 (21)")

    def test_filters(self):
        def names(**extra):
            return [r["gate"] for r in self.run_report(self.RID, **{**RANGE, **extra})["rows"]]

        self.assertEqual(names(activeOnly="true"), ["Gate 1", "Gate 3", "Gate 2"])
        self.assertEqual(names(gate="gate 1"), ["Gate 1"])
        self.assertEqual(names(branchIds=str(self.b2.id)), ["Gate 3"])
        self.assertEqual(names(dateFrom="2026-09-05", dateTo="2026-09-05"), ["Gate 3", "Gate 1", "Gate 2", "Gate 4"])
        body = self.run_report(self.RID, dateFrom="2026-09-05", dateTo="2026-09-05")
        self.assertEqual([r["totalScans"] for r in body["rows"]], [4, 2, 0, 0])

    def test_isolation_and_no_secrets(self):
        self.assertEqual(
            [r["gate"] for r in self.run_report(self.RID, self.b1_user, **RANGE)["rows"]],
            ["Gate 1", "Gate 2", "Gate 4"],
        )
        self.assertEqual([r["gate"] for r in self.run_report(self.RID, self.b2_user, **RANGE)["rows"]], ["Gate 3"])
        self.assertEqual(self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["rows"], [])
        raw = self.get(f"/api/reports/run/{self.RID}", **RANGE).content.decode()
        for secret in ("tok-g1", "tok-g2", "login_token", "loginToken", "password_hash", "passwordHash", "username"):
            self.assertNotIn(secret, raw)
        cols = [c.key for c in registry.get_spec(self.RID).columns]
        self.assertFalse({"username", "loginToken", "passwordHash"} & set(cols))

    def test_empty_and_exports(self):
        body = self.run_report(self.RID, **EMPTY_RANGE)
        self.assertTrue(all(r["totalScans"] == 0 for r in body["rows"]))  # gates are still listed, all zero
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:4], ["Gate", "Branch", "Status", "Last Login"])
        self.assertEqual((ws["A8"].value, ws["N8"].value), ("Gate 1", 21))
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── gate-daily-summary ──────────────────────────────────────────────────────


class DailySummaryTests(World):
    RID = "gate-daily-summary"
    KEYS = (
        "passRequests",
        "passesApproved",
        "exits",
        "returns",
        "qrSubmissions",
        "visitors",
        "teaBreaks",
        "teaOvertime",
        "deniedScans",
    )

    def test_golden_days_and_totals(self):
        body = self.run_report(self.RID, **RANGE)
        self.assertEqual(len(body["rows"]), 15)  # every day is listed
        rows = self.by(body, "date")
        want = {
            "2026-09-05": (1, 1, 1, 1, 0, 4, 2, 0, 0),
            "2026-09-06": (1, 0, 0, 0, 0, 0, 1, 1, 1),
            "2026-09-07": (0, 0, 0, 0, 0, 0, 1, 0, 1),  # only an open tea break and a forged-QR scan at 23:00 IST
            "2026-09-08": (3, 3, 3, 2, 1, 1, 1, 0, 1),  # includes everything stamped 00:30-01:45 IST
            "2026-09-09": (1, 1, 1, 2, 0, 0, 1, 1, 0),
            "2026-09-10": (1, 1, 0, 0, 1, 0, 1, 0, 1),  # T7 took 15m29s: rounds to 15, not overtime
            "2026-09-11": (1, 1, 0, 0, 1, 0, 1, 1, 0),  # T8 took exactly 15m30s
            "2026-09-12": (1, 1, 1, 1, 1, 0, 1, 1, 0),
            "2026-09-15": (3, 2, 1, 0, 0, 0, 0, 0, 0),
        }
        for day, values in want.items():
            self.assertEqual(tuple(rows[day][k] for k in self.KEYS), values, day)
        self.assertEqual(rows["2026-09-08"]["weekday"], "Tue")
        totals = body["totals"]
        self.assertEqual(tuple(totals[k] for k in self.KEYS), (18, 15, 12, 9, 6, 5, 9, 4, 7))
        cards = self.cards(body)
        self.assertEqual(
            (cards["Pass requests"], cards["Gate exits"], cards["QR submissions"], cards["Visitors"]), (18, 12, 6, 5)
        )
        self.assertEqual((cards["Tea breaks"], cards["Denied scans"]), (9, 7))
        self.assertEqual(cards["Busiest day (most gate exits)"], "08-Sep-2026 (3 exits)")
        self.assertIn("columns are not additive", " ".join(body["notes"]))

    def test_tea_overtime_matches_the_hr_page_rule_for_any_allowance(self):
        # The HR tea-break page: overtime <=> round(taken minutes) > allowed (Python round: halves go to even).
        logs = list(TeaBreakLog.objects.filter(in_at__isnull=False))
        for allowed in (10, 14, 15, 16, 17, 20):
            TeaBreakRule.objects.update_or_create(pk=1, defaults={"allowed_minutes": allowed})
            want = sum(1 for t in logs if round((t.in_at - t.out_at).total_seconds() / 60) > allowed)
            got = self.run_report(self.RID, **RANGE)["totals"]["teaOvertime"]
            self.assertEqual(got, want, allowed)
        TeaBreakRule.objects.update_or_create(pk=1, defaults={"allowed_minutes": 16})
        note = " ".join(self.run_report(self.RID, **RANGE)["notes"])
        self.assertIn("current allowed break of 16 minutes", note)

    def test_get_never_creates_the_tea_rule(self):
        self.assertFalse(TeaBreakRule.objects.exists())
        self.run_report(self.RID, **RANGE)
        self.assertFalse(TeaBreakRule.objects.exists())

    def test_filters_and_isolation(self):
        one = self.run_report(self.RID, dateFrom="2026-09-08", dateTo="2026-09-08")
        self.assertEqual(len(one["rows"]), 1)
        self.assertEqual(tuple(one["rows"][0][k] for k in self.KEYS), (3, 3, 3, 2, 1, 1, 1, 0, 1))
        b2 = self.run_report(self.RID, branchIds=str(self.b2.id), **RANGE)["totals"]
        self.assertEqual(tuple(b2[k] for k in self.KEYS), (3, 3, 1, 1, 1, 2, 1, 0, 1))
        b1_user = self.run_report(self.RID, self.b1_user, **RANGE)["totals"]
        self.assertEqual(tuple(b1_user[k] for k in self.KEYS), (14, 11, 10, 7, 4, 2, 8, 4, 6))
        b2_user = self.run_report(self.RID, self.b2_user, **RANGE)["totals"]
        self.assertEqual(tuple(b2_user[k] for k in self.KEYS), tuple(b2[k] for k in self.KEYS))
        widened = self.run_report(self.RID, self.b1_user, branchIds=str(self.b2.id), **RANGE)["totals"]
        self.assertEqual(widened["passRequests"], 0)
        self.assertEqual(widened["visitors"], 0)

    def test_empty_and_exports(self):
        body = self.run_report(self.RID, **EMPTY_RANGE)
        self.assertEqual(len(body["rows"]), 2)
        self.assertTrue(all(all(r[k] == 0 for k in self.KEYS) for r in body["rows"]))
        self.assertIsNone(body["summary"][-1]["value"])
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **RANGE).content)).active
        self.assertEqual([c.value for c in ws[7]][:4], ["Date", "Day", "Pass Requests", "Passes Approved"])
        self.assertEqual(_day(ws["A8"].value), datetime(2026, 9, 1).date())
        self.assertEqual(ws.cell(row=8 + 15, column=1).value, "TOTAL")
        self.assertTrue(self.export(self.RID, "pdf", **RANGE).content.startswith(b"%PDF"))


# ── cross-cutting: access, catalog, exports, read-only, query counts ────────


class AccessAndCatalogTests(World):
    def test_registry_contains_exactly_the_owned_ids_with_valid_modules(self):
        specs = {s.id: s for s in registry.all_specs()}
        self.assertEqual(registry.LOAD_ERRORS.get("api.reporting.definitions.gate_outpass"), None)
        for rid in GATE_REPORT_IDS:
            self.assertIn(rid, specs)
            spec = specs[rid]
            self.assertEqual(spec.category, "gate")
            self.assertTrue(spec.modules and all(m in all_module_keys() for m in spec.modules), rid)
            self.assertFalse(spec.super_admin_only)
            self.assertTrue(spec.title and spec.description)
        family = [s for s in specs.values() if s.family == "outpass"]
        self.assertEqual(
            sorted(s.id for s in family),
            ["outpass-department-summary", "outpass-employee-summary", "outpass-in-out-register", "outpass-register"],
        )
        self.assertEqual(len({s.variant for s in family}), 4)

    def test_role_without_the_owning_module_gets_403_report_forbidden(self):
        for rid in GATE_REPORT_IDS:
            for url, extra in (("run", {}), ("export", {"fmt": "xlsx"})):
                r = self.get(f"/api/reports/{url}/{rid}", self.plain_user, **RANGE, **extra)
                self.assertEqual(r.status_code, 403, (rid, url))
                self.assertEqual(r.json()["error"], "report_forbidden", (rid, url))

    def test_request_module_opens_only_the_pass_lifecycle_reports(self):
        for rid in REQUEST_LIFECYCLE_IDS:
            self.assertEqual(self.get(f"/api/reports/run/{rid}", self.requests_user, **RANGE).status_code, 200, rid)
        for rid in GATE_ONLY_IDS:
            r = self.get(f"/api/reports/run/{rid}", self.requests_user, **RANGE)
            self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"), rid)
        for rid in GATE_REPORT_IDS:
            self.assertEqual(self.get(f"/api/reports/run/{rid}", self.gate_user, **RANGE).status_code, 200, rid)

    def test_catalog_lists_them_for_the_right_roles(self):
        def ids(user):
            return {r["id"] for r in self.get("/api/reports/catalog", user).json()["reports"]}

        self.assertTrue(set(GATE_REPORT_IDS) <= ids(self.admin))
        self.assertTrue(set(GATE_REPORT_IDS) <= ids(self.gate_user))
        self.assertFalse(set(GATE_REPORT_IDS) & ids(self.plain_user))
        self.assertEqual(set(GATE_REPORT_IDS) & ids(self.requests_user), set(REQUEST_LIFECYCLE_IDS))
        entry = next(
            r for r in self.get("/api/reports/catalog", self.b1_user).json()["reports"] if r["id"] == "outpass-register"
        )
        self.assertNotIn(
            "branch", [f["key"] for f in entry["filters"]]
        )  # branch-scoped users never see the Branch filter
        self.assertEqual(entry["family"], "outpass")

    def test_bad_filter_values_are_400(self):
        for rid, params in (
            ("outpass-register", {"approvalStatus": "maybe"}),
            ("outpass-register", {"scanStatus": "flying"}),
            ("outpass-in-out-register", {"returnState": "x"}),
            ("outpass-department-summary", {"groupBy": "planet"}),
            ("gate-scan-audit-log", {"result": "nope"}),
            ("outpass-qr-submissions", {"matchStatus": "??"}),
            ("outpass-late-return", {"minOverrunMinutes": "5"}),
            ("gate-daily-summary", {"dateFrom": "2026-09-10", "dateTo": "2026-09-01"}),
        ):
            r = self.get(f"/api/reports/run/{rid}", **params)
            self.assertEqual(r.status_code, 400, (rid, params))
            self.assertEqual(r.json()["error"], "invalid_filter")

    def test_defaults_run_without_any_parameters(self):
        for rid in GATE_REPORT_IDS:
            self.assertEqual(self.get(f"/api/reports/run/{rid}").status_code, 200, rid)


class ExportEverythingTests(World):
    def test_every_report_exports_xlsx_and_pdf_with_data(self):
        for rid in GATE_REPORT_IDS:
            spec = registry.get_spec(rid)
            x = self.export(rid, "xlsx", **RANGE)
            self.assertEqual(x.status_code, 200, rid)
            self.assertTrue(x.content.startswith(b"PK"), rid)
            ws = load_workbook(io.BytesIO(x.content)).active
            self.assertEqual([c.value for c in ws[7]], [c.label for c in spec.columns], rid)
            self.assertIsNotNone(ws["A8"].value, rid)
            p = self.export(rid, "pdf", **RANGE)
            self.assertEqual(p.status_code, 200, rid)
            self.assertTrue(p.content.startswith(b"%PDF"), rid)
            self.assertIn(".xlsx", x["Content-Disposition"])

    def test_every_report_handles_an_empty_range_in_every_format(self):
        for rid in GATE_REPORT_IDS:
            self.assertEqual(self.get(f"/api/reports/run/{rid}", **EMPTY_RANGE).status_code, 200, rid)
            for fmt, sig in (("xlsx", b"PK"), ("pdf", b"%PDF")):
                r = self.export(rid, fmt, **EMPTY_RANGE)
                self.assertEqual(r.status_code, 200, (rid, fmt))
                self.assertTrue(r.content.startswith(sig), (rid, fmt))

    def test_branch_scoped_exports_never_contain_other_branch_data(self):
        r = self.export("outpass-register", "xlsx", self.b1_user, **RANGE)
        cells = {c.value for row in load_workbook(io.BytesIO(r.content)).active.iter_rows() for c in row}
        self.assertNotIn("GT003", cells)
        self.assertNotIn("GT006", cells)
        self.assertIn("GT001", cells)
        r = self.export("outpass-qr-submissions", "xlsx", self.b1_user, source="all", **RANGE)
        cells = {c.value for row in load_workbook(io.BytesIO(r.content)).active.iter_rows() for c in row}
        self.assertNotIn("Q4", cells)
        self.assertNotIn("Q5", cells)

    def test_row_limit_is_applied_to_every_row_per_record_report(self):
        # A detail report must stop at the limit in the query/loop itself and flag the cut (never a silent truncate);
        # the summary cards then cover the listed rows only and the totals row is withheld.
        ids = (
            "outpass-register",
            "outpass-in-out-register",
            "outpass-not-returned",
            "outpass-late-return",
            "outpass-pending-requests",
            "outpass-qr-submissions",
            "gate-scan-audit-log",
        )
        with mock.patch("api.reporting.runner.SCREEN_ROW_LIMIT", 1):
            for rid in ids:
                body = self.run_report(rid, **RANGE)
                self.assertEqual(body["rowCount"], 1, rid)
                self.assertTrue(body["truncated"], rid)
                self.assertIn("cut at the row limit", " ".join(body["notes"]), rid)

    def test_an_export_over_the_row_limit_is_refused_not_truncated(self):
        one_day = {"dateFrom": "2026-09-05", "dateTo": "2026-09-05"}  # one request (R05)
        for fmt, limit_name in (("xlsx", "XLSX_ROW_LIMIT"), ("pdf", "PDF_ROW_LIMIT")):
            with mock.patch(f"api.reporting.runner.{limit_name}", 3):
                r = self.export("outpass-register", fmt, **RANGE)
                self.assertEqual(r.status_code, 413, fmt)
                self.assertEqual(r.json()["error"], "too_many_rows", fmt)
                # a narrower range under the same limit still exports
                self.assertEqual(self.export("outpass-register", fmt, **one_day).status_code, 200, fmt)

    def test_exports_are_audited(self):
        from .models import AuditLog

        before = AuditLog.objects.filter(action="export", module="reports").count()
        self.export("gate-scan-audit-log", "pdf", **RANGE)
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), before + 1)


class ReadOnlyTests(World):
    def test_running_a_report_never_writes(self):
        for rid in GATE_REPORT_IDS:
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(self.get(f"/api/reports/run/{rid}", **RANGE).status_code, 200, rid)
            writes = [
                q["sql"][:80]
                for q in ctx.captured_queries
                if q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
            ]
            self.assertEqual(writes, [], rid)
        self.assertFalse(TeaBreakRule.objects.exists())


def add_employees_with_activity(branch, dept, gate, count):
    """`count` more employees, each with a returned pass, a pending one, an unused one, an open one, a QR row,
    two scans, a tea break and a hosted visit."""
    for _ in range(count):
        i = Employee.objects.count() + 1
        emp = Employee.objects.create(
            employee_code=f"QC{i:03d}",
            first_name=f"Emp{i}",
            last_name="Q",
            department=dept,
            branch=branch,
            phone=f"90000{i:05d}",
        )
        day = f"2026-09-{(i % 9) + 1:02d}"
        req = mk_request(
            emp,
            f"D{i}",
            f"{day} 09:00",
            role="hr",
            by="Priya HR",
            pass_type="official",
            approved=f"{day} 09:05",
            exit=f"{day} 09:10",
            exit_gate=gate,
            enter=f"{day} 10:10",
            entry_gate=gate,
            expected=f"{day} 09:40",
        )
        mk_request(emp, f"P{i}", f"{day} 11:00", status="pending", pass_type="personal")
        mk_request(emp, f"X{i}", f"{day} 12:00", role="hr", by="Priya HR", approved=f"{day} 12:05")
        mk_request(
            emp,
            f"N{i}",
            f"{day} 13:00",
            role="hr",
            by="Priya HR",
            approved=f"{day} 13:05",
            exit=f"{day} 13:10",
            exit_gate=gate,
        )
        mk_qr(branch, emp, f"Emp{i} Q", emp.employee_code, f"Q{i}", f"{day} 10:00")
        mk_scan(gate, req, emp, "exit", "success", f"{day} 09:10")
        mk_scan(gate, req, emp, "exit", "already_scanned", f"{day} 09:20")
        mk_tea(emp, gate, f"{day} 15:00", gate, f"{day} 15:20")
        mk_visit(Visitor.objects.create(name=f"V{i}", phone=f"91{i:08d}"), branch, f"{day} 10:00", host=emp)


class QueryCountTests(TestCase):
    """The number of queries must not grow with the number of employees (no N+1)."""

    @classmethod
    def setUpTestData(cls):
        cls.branch = Branch.objects.create(name="Unit Q", code="UQ")
        cls.dept = Department.objects.create(name="QDEPT", branch=cls.branch)
        cls.gate = mk_gate("Gate Q", cls.branch, "gq")
        cls.admin = mk_user("qc_admin", {}, super_admin=True)
        add_employees_with_activity(cls.branch, cls.dept, cls.gate, 3)

    def add_employees(self, count):
        add_employees_with_activity(self.branch, self.dept, self.gate, count)

    def setUp(self):
        patcher = mock.patch("api.reporting.definitions.gate_outpass_common.now", return_value=NOW)
        patcher.start()
        self.addCleanup(patcher.stop)

    def counts(self):
        out = {}
        for rid in GATE_REPORT_IDS:
            with CaptureQueriesContext(connection) as ctx:
                r = self.client.get(
                    f"/api/reports/run/{rid}",
                    {"dateFrom": "2026-09-01", "dateTo": "2026-09-15"},
                    **_hr_headers(self.admin),
                )
            self.assertEqual(r.status_code, 200, rid)
            out[rid] = (len(ctx), len(r.json()["rows"]))
        return out

    def test_query_count_is_flat_from_3_to_15_employees(self):
        small = self.counts()
        self.add_employees(12)
        large = self.counts()
        for rid in GATE_REPORT_IDS:
            self.assertLessEqual(abs(large[rid][0] - small[rid][0]), 2, (rid, small[rid], large[rid]))
        # the fixture really did get bigger for the row-per-record reports
        for rid in (
            "outpass-register",
            "outpass-not-returned",
            "outpass-pending-requests",
            "gate-scan-audit-log",
            "outpass-qr-submissions",
        ):
            self.assertGreater(large[rid][1], small[rid][1], rid)


# ── the real code paths -> the reports ──────────────────────────────────────


class RealFlowTests(TestCase):
    """Produce the data with the app's own approval and gate-scan code, then read it back through the reports."""

    def setUp(self):
        self.branch = Branch.objects.create(name="Flow Unit", code="FU")
        self.dept = Department.objects.create(name="FLOWDEPT", branch=self.branch)
        self.emp = Employee.objects.create(
            employee_code="FL001",
            first_name="Flo",
            last_name="Wright",
            department=self.dept,
            branch=self.branch,
        )
        self.gate = mk_gate("Flow Gate", self.branch, "flowgate")
        self.admin = mk_user("flow_admin", {}, super_admin=True)
        patcher = mock.patch("api.reporting.definitions.gate_outpass_common.now", return_value=ist("2026-09-10 18:00"))
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_approval_exit_duplicate_scan_and_return(self):
        from django.utils import timezone

        req = OutpassRequest.objects.create(employee=self.emp, destination="Bank", reason="Cash", pass_type="official")
        resolve_outpass_request(req, "approved", "Priya HR", "hr", "ok")
        req.refresh_from_db()
        exit_qr = sign_token({"role": "outpass_pass", "requestId": req.id}, expires_in=timedelta(minutes=60))
        self.assertEqual(resolve_gate_scan(self.gate, exit_qr)[0]["result"], "success")
        self.assertEqual(resolve_gate_scan(self.gate, exit_qr)[0]["result"], "already_scanned")
        req.refresh_from_db()
        req.return_qr_generated_at = timezone.now()
        req.save(update_fields=["return_qr_generated_at"])
        return_qr = sign_token(
            {"role": "outpass_return", "requestId": req.id, "generatedAt": req.return_qr_generated_at.isoformat()},
            expires_in=timedelta(minutes=60),
        )
        self.assertEqual(resolve_gate_scan(self.gate, return_qr)[0]["result"], "success")

        # Pin every timestamp the real code stamped with "now" to a fixed morning.
        OutpassRequest.objects.filter(pk=req.pk).update(
            created_at=ist("2026-09-10 09:00"),
            approved_at=ist("2026-09-10 09:10"),
            exited_at=ist("2026-09-10 09:20"),
            return_qr_generated_at=ist("2026-09-10 10:00"),
            entered_at=ist("2026-09-10 10:35"),
        )
        for scan, when in zip(OutpassGateScan.objects.order_by("id"), ("09:20", "09:25", "10:35")):
            OutpassGateScan.objects.filter(pk=scan.pk).update(scanned_at=ist(f"2026-09-10 {when}"))
        OutpassRecord.objects.update(submitted_at=ist("2026-09-10 09:10"))

        day = {"dateFrom": "2026-09-10", "dateTo": "2026-09-10"}
        get = lambda rid, **p: self.client.get(
            f"/api/reports/run/{rid}", {**day, **p}, **_hr_headers(self.admin)
        ).json()  # noqa: E731

        reg = get("outpass-register")["rows"]
        self.assertEqual(len(reg), 1)
        self.assertEqual(
            (
                reg[0]["status"],
                reg[0]["scanStatus"],
                reg[0]["reviewedBy"],
                reg[0]["outsideMinutes"],
                reg[0]["expiresAt"],
            ),
            ("Approved", "Returned", f"HR {DASH} Priya HR", 75, "2026-09-10 10:10"),
        )
        audit = get("gate-scan-audit-log")["rows"]
        self.assertEqual(
            sorted((r["result"], r["scanType"]) for r in audit),
            [("Already Scanned", "Exit"), ("Success", "Exit"), ("Success", "Return")],
        )
        dup = next(r for r in audit if r["result"] == "Already Scanned")
        self.assertEqual(dup["detail"], "Pass already used - exited at 10-Sep-2026 09:20 via Flow Gate")
        # The approval's copy on the Outpass page is not a gate exit: hidden by default, listed on request.
        self.assertEqual(get("outpass-qr-submissions")["rows"], [])
        mirror = get("outpass-qr-submissions", source="all")["rows"]
        self.assertEqual((len(mirror), mirror[0]["source"], mirror[0]["matchStatus"]), (1, "Approved Pass", "Matched"))
        summary = get("outpass-employee-summary")["rows"][0]
        self.assertEqual(
            (summary["requests"], summary["returned"], summary["totalOutsideMinutes"], summary["qrSubmissions"]),
            (1, 1, 75, 0),
        )
        turn = get("outpass-approval-turnaround")["rows"][0]
        self.assertEqual((turn["approverName"], turn["decisions"], turn["avgTurnaroundMinutes"]), ("Priya HR", 1, 10))


# ── helper units ────────────────────────────────────────────────────────────


class HelperTests(TestCase):
    def test_pass_state_boundaries(self):
        base = dict(
            status="approved", exited_at=None, entered_at=None, return_qr_generated_at=None, pass_type="official"
        )
        approved = ist("2026-09-15 09:00")
        now = lambda hh, mm=0, ss=0: ist(f"2026-09-15 {hh:02d}:{mm:02d}:{ss:02d}")  # noqa: E731
        req = SimpleNamespace(approved_at=approved, **base)
        today = now(0, 0).date()
        self.assertEqual(C.pass_state(req, now(9, 59, 59), today), "pending_exit")
        self.assertEqual(C.pass_state(req, now(10, 0, 0), today), "expired_unscanned")  # the window closes AT 60:00
        rej = SimpleNamespace(approved_at=None, **{**base, "status": "rejected"})
        self.assertEqual(C.pass_state(rej, now(10), today), "not_applicable")
        out = SimpleNamespace(approved_at=approved, **{**base, "exited_at": ist("2026-09-15 09:10")})
        self.assertEqual(C.pass_state(out, now(10), today), "exited")
        out.return_qr_generated_at = ist("2026-09-15 09:30")
        self.assertEqual(C.pass_state(out, now(10), today), "pending_return")
        self.assertEqual(C.pass_state(out, now(10, 30), today), "return_expired")  # return QR is valid 60 minutes
        tomorrow = today + timedelta(days=1)
        self.assertEqual(C.pass_state(out, now(10), tomorrow), "not_returned")  # left on an earlier day
        out.pass_type = "early_dismissal"
        self.assertEqual(C.pass_state(out, now(10), tomorrow), "early_dismissal")
        out.entered_at = ist("2026-09-15 09:50")
        self.assertEqual(C.pass_state(out, now(10), tomorrow), "completed")

    def test_ist_day_bounds_are_closed_open_and_not_utc(self):
        ctx = SimpleNamespace(date_from=ist("2026-09-08 00:00").date(), date_to=ist("2026-09-08 00:00").date())
        start, end = C.ist_bounds(ctx)
        self.assertEqual((start, end), (ist("2026-09-08 00:00"), ist("2026-09-09 00:00")))
        # 00:30 IST on the 8th is 19:00 UTC on the 7th, yet belongs to the 8th; 23:30 IST on the 8th is 18:00 UTC on the 8th.
        self.assertTrue(start <= ist("2026-09-08 00:30") < end)
        self.assertTrue(start <= ist("2026-09-08 23:30") < end)
        self.assertFalse(start <= ist("2026-09-07 23:59:59") < end)
        self.assertFalse(start <= ist("2026-09-09 00:00") < end)

    def test_minutes_round_halves_up_and_never_go_negative(self):
        self.assertEqual(C.half_up_minutes(29), 0)
        self.assertEqual(C.half_up_minutes(30), 1)
        self.assertEqual(C.half_up_minutes(90), 2)  # 1.5 minutes -> 2
        self.assertEqual(C.half_up_minutes(150), 3)  # 2.5 -> 3 (Python's round() would give 2)
        self.assertEqual(C.half_up_minutes(-500), 0)
        self.assertEqual(C.mean_minutes([52, 53]), 53)
        self.assertIsNone(C.mean_minutes([]))

    def test_text_helpers(self):
        self.assertEqual(C.clean("  a\n\n b\t c "), "a b c")
        self.assertIsNone(C.clean("   "))
        self.assertIsNone(C.clean(None))
        self.assertEqual(len(C.clean("z" * 999)), C.MAX_TEXT)
        self.assertEqual(C.reviewer_text("hr", None), "HR")
        self.assertEqual(C.reviewer_text("dept_head", "Mani"), f"HOD {DASH} Mani")
        self.assertEqual(C.reviewer_text("system", "Rita"), "On-Duty approval (Rita)")
        self.assertEqual(C.reviewer_text(None, None), None)
        self.assertEqual(C.reviewer_text(None, "Someone"), "Someone")
        self.assertEqual(C.pass_type_label(None), "Unspecified")
        self.assertEqual(C.pass_type_label("early_dismissal"), "Early Dismissal")

    def test_overrun_only_where_measurable(self):
        exp = ist("2026-09-15 09:00")
        req = SimpleNamespace(
            expected_return_at=exp, exited_at=ist("2026-09-15 08:30"), entered_at=ist("2026-09-15 09:20")
        )
        now = ist("2026-09-15 10:00")
        self.assertEqual(C.overrun_minutes(req, "completed", now), 20)
        req.entered_at = ist("2026-09-15 08:50")
        self.assertEqual(C.overrun_minutes(req, "completed", now), 0)  # early: never negative
        req.entered_at = None
        self.assertEqual(C.overrun_minutes(req, "pending_return", now), 60)  # outside today: running, as of now
        self.assertIsNone(C.overrun_minutes(req, "not_returned", now))
        self.assertIsNone(C.overrun_minutes(req, "early_dismissal", now))
        req.expected_return_at = None
        self.assertIsNone(C.overrun_minutes(req, "pending_return", now))
        self.assertFalse(C.is_late(0))
        self.assertTrue(C.is_late(1))


# ── adversarial review (independent brute-force cross-checks) ───────────────


class AdversarialReviewTests(World):
    """Written by the reviewer: recompute each figure from the raw rows, without the reports' helpers."""

    def test_qr_report_does_not_reveal_other_branch_employees_to_a_scoped_user(self):
        # Anyone can type any employee code on the public branch QR form, so the matched employee of a Unit 2
        # code typed at a Unit 1 QR must not have its real name / department / designation shown to Unit 1 HR.
        mk_qr(self.b1, self.e3, "Whoever", "GT003", "QLEAK", "2026-09-12 10:00")
        body = self.run_report("outpass-qr-submissions", self.b1_user, **RANGE)
        row = next(r for r in body["rows"] if r["destination"] == "QLEAK")
        blob = json.dumps(row) + " ".join(body["notes"])
        self.assertNotIn("Chitra", blob)
        self.assertNotIn("SEWING", blob)
        # ...nor whether the code exists at all (Matched / Name Mismatch vs Unmatched is an oracle over every branch)
        self.assertEqual(row["matchStatus"], "Unmatched")

    def test_qr_isolation_holds_for_the_match_filter_the_cards_and_employee_filters(self):
        # The hidden other-branch match behaves as "code belongs to nobody" everywhere, not only in the row cells.
        mk_qr(self.b1, self.e3, "Chitra Devi", "GT003", "QHID", "2026-09-12 10:00")  # exact name, other branch
        mk_qr(self.b1, self.e1, "Asha Rao", "GT001", "QOWN", "2026-09-12 11:00")  # own branch, exact name
        mk_qr(self.b1, self.e2, "Wrong Name", "GT002", "QMIS", "2026-09-12 12:00")  # own branch, name differs

        def dests(user, **params):
            return self.dests(self.run_report("outpass-qr-submissions", user, **{**RANGE, **params}))

        mine = {"QHID", "QOWN", "QMIS"}
        self.assertTrue(mine <= set(dests(self.b1_user)))
        self.assertIn("QHID", dests(self.b1_user, matchStatus="unmatched"))
        self.assertNotIn("QOWN", dests(self.b1_user, matchStatus="unmatched"))
        self.assertNotIn("QHID", dests(self.b1_user, matchStatus="matched"))
        self.assertNotIn("QHID", dests(self.b1_user, matchStatus="name_mismatch"))
        self.assertIn("QOWN", dests(self.b1_user, matchStatus="matched"))
        self.assertIn("QMIS", dests(self.b1_user, matchStatus="name_mismatch"))
        # Naming another branch's department must not list who typed its members' codes.
        self.assertEqual(dests(self.b1_user, departmentIds=str(self.sew.id)), [])
        self.assertEqual(dests(self.b1_user, employeeIds=str(self.e3.id)), [])
        # Cards and notes only count / name the identified (own-branch) employees.
        body = self.run_report("outpass-qr-submissions", self.b1_user, **RANGE)
        matched_rows = sum(1 for r in body["rows"] if r["matchStatus"] == "Matched")
        self.assertEqual(self.cards(body)["Matched to an employee"], matched_rows)
        self.assertNotIn("Matched", {r["matchStatus"] for r in body["rows"] if r["destination"] == "QHID"})
        self.assertNotIn("GT003", " ".join(body["notes"]))
        # An unscoped admin still sees everything, identified.
        row = next(
            r for r in self.run_report("outpass-qr-submissions", None, **RANGE)["rows"] if r["destination"] == "QHID"
        )
        self.assertEqual(
            (row["matchStatus"], row["matchedEmployee"], row["department"]),
            ("Matched", "GT003 - Chitra Devi", "SEWING"),
        )


def _half_up(seconds):
    import math

    return max(0, math.floor(seconds / 60 + 0.5))


class RowLimitFetchTests(World):
    """README rule 8: the row limit must bound the QUERY, not only the rows kept after loading everything."""

    def instantiated(self, model, rid, **params):
        from django.db.models.signals import post_init

        seen = []

        def handler(sender, instance, **kw):
            seen.append(instance.pk)

        post_init.connect(handler, sender=model)
        try:
            with mock.patch("api.reporting.runner.SCREEN_ROW_LIMIT", 2):
                body = self.run_report(rid, **params)
        finally:
            post_init.disconnect(handler, sender=model)
        return body, len(seen)

    def test_in_out_register_does_not_load_every_exit_to_show_two_rows(self):
        for i in range(30):
            mk_request(
                self.e1,
                f"BULK{i}",
                "2026-09-02 10:00",
                approved="2026-09-02 10:05",
                exit="2026-09-02 10:10",
                exit_gate=self.g1,
                enter="2026-09-02 10:40",
                entry_gate=self.g1,
            )
        body, loaded = self.instantiated(OutpassRequest, "outpass-in-out-register", **RANGE)
        self.assertEqual(body["rowCount"], 2)
        self.assertLessEqual(loaded, 4, f"{loaded} OutpassRequest objects were built for a 2-row list")

    def test_expired_and_late_reports_do_not_load_every_candidate_to_show_two_rows(self):
        for i in range(30):
            mk_request(
                self.e1,
                f"EXP{i}",
                "2026-09-02 10:00",
                approved="2026-09-02 10:05",
                role="hr",
                by="Priya HR",
            )
            mk_request(
                self.e1,
                f"LATE{i}",
                "2026-09-03 10:00",
                approved="2026-09-03 10:05",
                exit="2026-09-03 10:10",
                exit_gate=self.g1,
                enter="2026-09-03 12:00",
                entry_gate=self.g1,
                expected="2026-09-03 11:00",
            )
        for rid in ("outpass-expired-unused", "outpass-late-return"):
            with self.subTest(rid):
                body, loaded = self.instantiated(OutpassRequest, rid, **RANGE)
                self.assertEqual(body["rowCount"], 2, rid)
                self.assertLessEqual(loaded, 4, f"{rid}: {loaded} OutpassRequest objects were built for a 2-row list")

    def test_qr_submissions_do_not_load_every_record_to_show_two_rows(self):
        for i in range(30):
            mk_qr(self.b1, self.e1, "Asha Rao", "GT001", f"QB{i}", "2026-09-02 10:00")
        body, loaded = self.instantiated(OutpassRecord, "outpass-qr-submissions", **RANGE)
        self.assertEqual(body["rowCount"], 2)
        self.assertLessEqual(loaded, 4, f"{loaded} OutpassRecord objects were built for a 2-row list")


class FilterContractTests(World):
    def test_a_blank_choice_must_not_be_labelled_all_when_the_server_turns_it_into_a_default(self):
        # The Report Center UI shows the placeholder as the blank choice of every optional select and omits blank
        # values from the query, so the server's default is what runs. "All" that runs Manual only misleads.
        cat = self.get("/api/reports/catalog").json()
        offenders = []
        for rep in cat["reports"]:
            if rep["id"] not in GATE_REPORT_IDS:
                continue
            for f in rep["filters"]:
                if f["kind"] != "select" or f.get("required") or not f.get("default"):
                    continue
                label = dict((o["value"], o["label"]) for o in f["options"]).get(f["default"], "")
                if f.get("placeholder", "").strip().lower() in ("all", "any") or not f.get("placeholder"):
                    offenders.append((rep["id"], f["key"], f["default"], f.get("placeholder"), label))
        self.assertEqual(offenders, [])

    def test_a_nul_byte_in_the_gate_filter_is_not_a_500(self):
        for rid in ("outpass-register", "outpass-in-out-register", "gate-scan-audit-log", "gate-activity-summary"):
            with self.subTest(rid):
                r = self.get(f"/api/reports/run/{rid}", gate="Gate\x00 1", **RANGE)
                self.assertIn(r.status_code, (200, 400))

    def test_a_date_range_ending_on_the_last_day_of_year_9999_is_not_a_500(self):
        for rid in GATE_REPORT_IDS:
            with self.subTest(rid):
                r = self.get(f"/api/reports/run/{rid}", dateFrom="9999-12-30", dateTo="9999-12-31")
                self.assertIn(r.status_code, (200, 400))


class RandomCrossCheckTests(TestCase):
    """Random passes (cross-midnight exits, night outings, odd statuses) recomputed from raw rows."""

    RANGE = {"dateFrom": "2026-09-01", "dateTo": "2026-09-15"}

    def setUp(self):
        import random

        patcher = mock.patch("api.reporting.definitions.gate_outpass_common.now", return_value=NOW)
        patcher.start()
        self.addCleanup(patcher.stop)
        rnd = random.Random(20260930)
        self.b = Branch.objects.create(name="RB", code="RB")
        self.d1 = Department.objects.create(name="RD1", branch=self.b)
        self.d2 = Department.objects.create(name="RD2", branch=self.b)
        self.gate = mk_gate("RG", self.b, "rg")
        self.admin = mk_user("rc_admin", {}, super_admin=True)
        self.emps = []
        for i in range(8):
            self.emps.append(
                Employee.objects.create(
                    employee_code=f"RC{i:02d}",
                    first_name=f"R{i}",
                    last_name="X",
                    department=self.d1 if i % 2 else self.d2,
                    branch=self.b,
                    status="active" if i < 6 else "resigned",
                    employment_type="production" if i % 3 == 0 else "staff",
                )
            )
        self.reqs = []
        for n in range(220):
            emp = rnd.choice(self.emps)
            created = NOW - timedelta(days=rnd.randint(0, 15), minutes=rnd.randint(0, 1439))
            status = rnd.choices(["pending", "approved", "rejected"], [2, 7, 1])[0]
            source = rnd.choices(["manual", "on_duty"], [8, 2])[0]
            ptype = rnd.choice([None, "official", "personal", "early_dismissal"]) if source == "manual" else None
            approved = exit_at = enter = expected = return_qr = None
            if status != "pending":
                approved = created + timedelta(minutes=rnd.randint(0, 120), seconds=rnd.randint(0, 59))
            if status == "approved" and rnd.random() < 0.75:
                exit_at = approved + timedelta(minutes=rnd.randint(0, 59), seconds=rnd.randint(0, 59))
                if rnd.random() < 0.7:
                    enter = exit_at + timedelta(minutes=rnd.randint(1, 400), seconds=rnd.randint(0, 59))
                elif rnd.random() < 0.5:
                    return_qr = exit_at + timedelta(minutes=rnd.randint(1, 500))
            if source == "manual" and rnd.random() < 0.6:
                expected = (exit_at or created) + timedelta(minutes=rnd.randint(-10, 200), seconds=rnd.randint(0, 59))
            if enter and enter > NOW:
                enter = None
            if exit_at and exit_at > NOW:
                exit_at = enter = None
            req = OutpassRequest.objects.create(
                employee=emp,
                destination=f"D{n}",
                reason="r",
                status=status,
                source=source,
                pass_type=ptype,
                approver_role=("hr" if rnd.random() < 0.5 else "dept_head") if status != "pending" else None,
                approved_by=rnd.choice(["Priya", "Mani"]) if status != "pending" else None,
                approved_at=approved,
                exit_gate=self.gate if exit_at else None,
                exited_at=exit_at,
                entered_at=enter,
                entry_gate=self.gate if enter else None,
                expected_return_at=expected,
                return_qr_generated_at=return_qr,
            )
            OutpassRequest.objects.filter(pk=req.pk).update(created_at=created)
            self.reqs.append(OutpassRequest.objects.get(pk=req.pk))

    def get(self, rid, **params):
        r = self.client.get(f"/api/reports/run/{rid}", {**self.RANGE, **params}, **_hr_headers(self.admin))
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    @staticmethod
    def day(dt):
        return dt.astimezone(FACTORY_TZ).date()

    def in_range(self, dt):
        return dt is not None and date(2026, 9, 1) <= self.day(dt) <= date(2026, 9, 15)

    def state(self, r):
        today = self.day(NOW)
        if r.entered_at:
            return "completed"
        if r.exited_at:
            if r.pass_type == "early_dismissal":
                return "early"
            return "not_returned" if self.day(r.exited_at) < today else "outside"
        if r.status != "approved":
            return "na"
        return "expired" if NOW >= r.approved_at + timedelta(minutes=60) else "pending_exit"

    def overrun(self, r):
        if r.expected_return_at is None or r.exited_at is None:
            return None
        if r.entered_at:
            return _half_up((r.entered_at - r.expected_return_at).total_seconds())
        if self.state(r) == "outside":
            return _half_up((NOW - r.expected_return_at).total_seconds())
        return None

    def test_register_and_purpose_and_dept_totals_agree_with_raw_rows(self):
        mine = [r for r in self.reqs if self.in_range(r.created_at)]
        reg = self.get("outpass-register")
        self.assertEqual(reg["rowCount"], len(mine))
        want_label = {
            "completed": {"Returned"},
            "early": {"Early Dismissal"},
            "not_returned": {"Not Returned"},
            "outside": {"Exited", "Awaiting Return", "Return QR Expired"},
            "na": {None},
            "expired": {"Expired, Not Scanned"},
            "pending_exit": {"Awaiting Exit"},
        }
        got = {r["destination"]: r for r in reg["rows"]}
        for r in mine:
            self.assertIn(got[r.destination]["scanStatus"], want_label[self.state(r)], (r.destination, self.state(r)))
        mins = [_half_up((r.entered_at - r.exited_at).total_seconds()) for r in mine if r.entered_at and r.exited_at]
        self.assertEqual(reg["totals"]["outsideMinutes"], sum(mins))
        purpose = self.get("outpass-purpose-analysis")
        self.assertEqual(purpose["totals"]["requests"], len(mine))
        dept = self.get("outpass-department-summary")
        self.assertEqual(dept["totals"]["requests"], len(mine))
        self.assertEqual(dept["totals"]["exited"], sum(1 for r in mine if r.exited_at))
        self.assertEqual(dept["totals"]["returned"], sum(1 for r in mine if r.entered_at))
        self.assertEqual(
            dept["totals"]["notReturned"], sum(1 for r in mine if self.state(r) in ("outside", "not_returned"))
        )

    def test_employee_summary_matches_raw_rows(self):
        mine = [r for r in self.reqs if self.in_range(r.created_at)]
        body = self.get("outpass-employee-summary")
        rows = {r["employeeCode"]: r for r in body["rows"]}
        for emp in self.emps:
            mr = [r for r in mine if r.employee_id == emp.id]
            if not mr:
                self.assertNotIn(emp.employee_code, rows)
                continue
            row = rows[emp.employee_code]
            want = {
                "requests": len(mr),
                "approved": sum(1 for r in mr if r.status == "approved"),
                "rejected": sum(1 for r in mr if r.status == "rejected"),
                "pending": sum(1 for r in mr if r.status == "pending"),
                "exited": sum(1 for r in mr if r.exited_at),
                "returned": sum(1 for r in mr if r.entered_at),
                "notReturned": sum(1 for r in mr if self.state(r) in ("outside", "not_returned")),
                "expiredUnused": sum(1 for r in mr if self.state(r) == "expired"),
                "lateReturns": sum(1 for r in mr if (self.overrun(r) or 0) >= 1),
            }
            for k, v in want.items():
                self.assertEqual(row[k], v, (emp.employee_code, k))

    def test_in_out_not_returned_and_late_match_raw_rows(self):
        exits = [r for r in self.reqs if self.in_range(r.exited_at)]
        body = self.get("outpass-in-out-register")
        self.assertEqual(sorted(r["destination"] for r in body["rows"]), sorted(r.destination for r in exits))
        nr = self.get("outpass-not-returned")
        want = [r for r in exits if not r.entered_at and r.pass_type != "early_dismissal"]
        self.assertEqual(sorted(r["destination"] for r in nr["rows"]), sorted(r.destination for r in want))
        late = self.get("outpass-late-return")
        want_late = [r for r in exits if r.expected_return_at and (self.overrun(r) or 0) >= 1]
        self.assertEqual(sorted(r["destination"] for r in late["rows"]), sorted(r.destination for r in want_late))
        self.assertEqual(sum(r["overrunMinutes"] for r in late["rows"]), sum(self.overrun(r) for r in want_late))

    def test_expired_and_turnaround_and_daily_match_raw_rows(self):
        exp = [
            r
            for r in self.reqs
            if r.status == "approved"
            and self.in_range(r.approved_at)
            and r.source == "manual"
            and not r.exited_at
            and NOW >= r.approved_at + timedelta(minutes=60)
        ]
        body = self.get("outpass-expired-unused")
        self.assertEqual(body["rowCount"], len(exp))
        tr = self.get("outpass-approval-turnaround")
        dec = [r for r in self.reqs if self.in_range(r.created_at) and r.source == "manual" and r.status != "pending"]
        self.assertEqual(tr["totals"]["decisions"], len(dec))
        apps = [r for r in dec if r.status == "approved"]
        turns = [(r.approved_at - r.created_at).total_seconds() for r in apps]
        self.assertEqual(tr["totals"]["maxTurnaroundMinutes"], _half_up(max(turns)))
        daily = self.get("gate-daily-summary")
        rows = {r["date"]: r for r in daily["rows"]}
        for d in range(1, 16):
            day = date(2026, 9, d)
            want = {
                "passRequests": sum(1 for r in self.reqs if self.day(r.created_at) == day),
                "passesApproved": sum(
                    1 for r in self.reqs if r.status == "approved" and self.day(r.approved_at) == day
                ),
                "exits": sum(1 for r in self.reqs if r.exited_at and self.day(r.exited_at) == day),
                "returns": sum(1 for r in self.reqs if r.entered_at and self.day(r.entered_at) == day),
            }
            for k, v in want.items():
                self.assertEqual(rows[day.isoformat()][k], v, (day, k))
