"""
Report Center - attendance_core reports: time-card, daily-attendance, absentee-list-daily,
attendance-muster-sheet, attendance-summary-monthly, punch-log, punch-exceptions, unmatched-punches,
manual-overrides.

Run via: python manage.py test api.tests_reporting_attendance_core -v 2

Every date is in August 2026 (a month that is already over: "today" is patched to 2026-09-05), so nothing
depends on the real clock. Shift 09:00-18:00, 15 min grace, first half ends 13:30, lunch 60 min; the
half-day cut-offs are at their defaults (13:30 / 14:30); Early-Out detection is switched ON.

Calendar used (2026): Mon 3, Tue 4, Wed 5, Thu 6 (company holiday), Fri 7, Sat 8, Sun 9, Mon 10.
"""

import dataclasses
import io
import random
from datetime import date, datetime, time, timezone
from unittest import mock

import pdfplumber

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .attendance_final import _holiday_dates_for_month, compute_range_records
from .jwt_utils import sign_token
from .models import (
    Attendance,
    AttendanceDayRecord,
    AttendanceLog,
    AttendanceOverrideRequest,
    Branch,
    CasualLeaveRequest,
    DailyShiftLog,
    Department,
    Designation,
    Employee,
    EmployeePermission,
    EmployeeShiftAssignment,
    Holiday,
    HRUser,
    LeaveRequest,
    LeaveType,
    PayrollSettings,
    ResignationRequest,
    Role,
    ShiftTemplate,
    UnmatchedPunch,
)
from .reporting import registry

IDS = [
    "time-card",
    "daily-attendance",
    "absentee-list-daily",
    "attendance-muster-sheet",
    "attendance-summary-monthly",
    "punch-log",
    "punch-exceptions",
    "unmatched-punches",
    "manual-overrides",
]
TODAY = date(2026, 9, 5)
PERIOD = {"dateFrom": "2026-08-03", "dateTo": "2026-08-10"}
SRC = "biometric:adms:SN1"


def d(day: int) -> date:
    return date(2026, 8, day)


def tm(text: str) -> time:
    parts = [int(x) for x in text.split(":")]
    return time(*parts)


def punch(emp, day: date, *times: str, source: str = SRC) -> None:
    for text in times:
        AttendanceLog.objects.create(employee=emp, date=day, punch_time=tm(text), punch_type="IN", source=source)


def hr_headers(user) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def seed_records(emp, first: date, last: date) -> None:
    """Persist the engine's verdicts the way opening the Attendance screens does (test set-up only)."""
    with mock.patch("api.attendance_final.ist_today", return_value=TODAY):
        compute_range_records(emp, first, last, PayrollSettings.get())


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        _holiday_dates_for_month.cache_clear()
        ps = PayrollSettings.get()
        ps.attendance_mode = "simple"
        ps.morning_late_in_enabled = True
        ps.evening_early_out_enabled = True
        ps.save()

        cls.b1 = Branch.objects.create(name="Unit 1", code="U1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="U2")
        cls.cutting = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.sewing = Department.objects.create(name="SEWING", branch=cls.b1)
        cls.packing = Department.objects.create(name="PACKING", branch=cls.b2)
        cls.cutter = Designation.objects.create(title="Cutter")
        cls.tailor = Designation.objects.create(title="Tailor")

        cls.shift = ShiftTemplate.objects.create(
            name="General",
            shift_type="staff",
            start_time=time(9, 0),
            end_time=time(18, 0),
            grace_period_minutes=15,
            first_half_end=time(13, 30),
            lunch_duration_minutes=60,
        )

        def mk(code, first, dept, branch, desig=None, **kw):
            emp = Employee.objects.create(
                employee_code=code,
                first_name=first,
                last_name="T",
                department=dept,
                branch=branch,
                designation=desig,
                **kw,
            )
            return emp

        cls.alice = mk("AC001", "Alice", cls.cutting, cls.b1, cls.cutter, employment_type="staff")
        cls.bob = mk("AC002", "Bob", cls.cutting, cls.b1, cls.cutter, employment_type="staff")
        cls.carol = mk("AC003", "Carol", cls.sewing, cls.b1, cls.tailor, employment_type="staff")
        cls.pat = mk("AC004", "Pat", cls.sewing, cls.b1, cls.tailor, employment_type="production")
        cls.zed = mk("AC005", "Zed", cls.packing, cls.b2, None, employment_type="staff")
        cls.leo = mk("AC006", "Leo", cls.cutting, cls.b1, cls.cutter, employment_type="staff", status="inactive")
        for emp in (cls.alice, cls.bob, cls.zed, cls.leo):
            EmployeeShiftAssignment.objects.create(employee=emp, shift=cls.shift, effective_from=date(2020, 1, 1))
        # Carol's schedule has Saturdays off (payroll's rule; the engine itself stores such a Saturday as absent)
        EmployeeShiftAssignment.objects.create(
            employee=cls.carol, shift=cls.shift, effective_from=date(2020, 1, 1), saturday_off=True
        )

        Holiday.objects.create(name="Test Holiday", date=d(6))

        # ── raw punches ──────────────────────────────────────────────────────
        punch(cls.alice, d(3), "08:55", "13:00", "13:45", "18:05")
        punch(cls.bob, d(3), "09:40", "18:10")
        punch(cls.pat, d(3), "08:30", "12:45", "13:30", "20:00")
        punch(cls.zed, d(3), "09:00", "18:00", source="geo:auto")
        punch(cls.leo, d(3), "09:00", "18:00", source="on_duty:approved")
        punch(cls.alice, d(4), "09:10")
        punch(cls.bob, d(4), "09:00", "19:30")
        punch(cls.carol, d(4), "09:00", "15:00")
        punch(cls.pat, d(4), "08:50", "20:00")
        punch(cls.alice, d(5), "09:00", "09:02", "13:00", "13:45", "18:00", "18:01")
        punch(cls.bob, d(5), "09:00", "13:00")
        punch(cls.bob, d(5), "18:00", source="manual:punch-view")
        punch(cls.carol, d(7), "15:00")
        punch(cls.bob, d(10), "10:00", "18:00", source="missing_punch:approved")

        # ── leave, permission, casual leave, legacy presence ─────────────────
        sl, _ = LeaveType.objects.get_or_create(code="SL", defaults={"name": "Sick Leave"})
        LeaveRequest.objects.create(
            employee=cls.carol,
            leave_type_ref=sl,
            type="sick",
            start_date="2026-08-05",
            end_date="2026-08-05",
            total_days=1,
            status="approved",
        )
        EmployeePermission.objects.create(
            employee=cls.bob,
            date=d(10),
            permission_time=time(9, 0),
            status="approved",
            type="morning_late_in",
            duration_minutes=60,
        )
        CasualLeaveRequest.objects.create(employee=cls.carol, date=d(10), status="approved")
        Attendance.objects.create(employee=cls.zed, date="2026-08-07", present=True)

        # ── frozen manual days (written before the range is computed, which never touches them) ──
        def frozen(emp, day, by, note, first=None, last=None):
            return AttendanceDayRecord.objects.create(
                employee=emp,
                date=day,
                status="present",
                shifts_earned="1.00",
                source="manual",
                override_by=by,
                override_note=note,
                first_punch=first,
                last_punch=last,
            )

        frozen(cls.alice, d(10), "HR Anita", "Missed punch due to power cut", time(9, 0), time(18, 0))
        frozen(cls.carol, d(10), "HOD Meena", "Casual Leave (paid) -approved")
        frozen(cls.pat, d(10), "HR Anita", "Compensation Alternative Day (paid) -redeemed")

        for emp in (cls.alice, cls.bob, cls.carol, cls.pat, cls.zed, cls.leo):
            seed_records(emp, d(3), d(10))
        AttendanceDayRecord.objects.filter(source="manual").update(
            updated_at=datetime(2026, 8, 10, 12, 0, tzinfo=timezone.utc)
        )

        # HR's Informed call on absent days
        AttendanceDayRecord.objects.filter(employee=cls.carol, date=d(3)).update(is_informed=True)
        AttendanceDayRecord.objects.filter(employee=cls.zed, date=d(5)).update(is_informed=False)

        # ── override requests (attendance date d; created / reviewed stamps are UTC and shown in IST) ──
        def utc(day, hour, minute=0):
            return datetime(2026, 8, day, hour, minute, tzinfo=timezone.utc)

        def before(status="absent", shifts="0.00"):
            return {
                "status": status,
                "isLate": False,
                "isEarlyOut": False,
                "isHalfShift": False,
                "firstPunch": None,
                "lastPunch": None,
                "shiftsEarned": shifts,
            }

        def after(status="present", shifts="1.00", first="09:00", last="18:00", note=None):
            return {
                "status": status,
                "isLate": False,
                "isEarlyOut": False,
                "isHalfShift": False,
                "firstPunch": first,
                "lastPunch": last,
                "shiftsEarned": shifts,
                "note": note,
            }

        def request(emp, day, status, by, created, **kw):
            req = AttendanceOverrideRequest.objects.create(employee=emp, date=day, status=status, requested_by=by, **kw)
            AttendanceOverrideRequest.objects.filter(pk=req.pk).update(created_at=created)
            return req

        cls.req_approved = request(
            cls.alice,
            d(10),
            "approved",
            "HR Anita",
            utc(10, 3, 30),
            previous_values=before(),
            requested_values=after(note="Missed punch due to power cut"),
            reason="Missed punch due to power cut",
            reviewed_by="HOD Ravi",
            review_comment="Verified with gate",
            reviewed_at=utc(11, 5, 0),
        )
        cls.req_pending = request(
            cls.bob,
            d(7),
            "pending",
            "HR Anita",
            utc(8, 4, 0),
            previous_values=before(),
            requested_values=after(status="half_shift", shifts="0.50", first="09:00", last=None),
            reason="Left after lunch",
        )
        cls.req_rejected = request(
            cls.zed,
            d(8),
            "rejected",
            "HR Kumar",
            utc(9, 4, 0),
            previous_values=before(),
            requested_values=after(),
            reason="Sat work",
            reviewed_by="HOD Sita",
            review_comment="No evidence",
            reviewed_at=utc(9, 9, 0),
        )
        cls.req_other_branch = request(
            cls.zed,
            d(9),
            "pending",
            "HR Kumar",
            utc(9, 5, 0),
            previous_values=before(),
            requested_values=after(),
            reason="Sunday shift",
        )

        # ── device IDs that match no employee (stamps are UTC; shown in IST) ──
        def unmatched(uid, serial, label, count, last_date, resolved=False, note=""):
            row = UnmatchedPunch.objects.create(
                device_user_id=uid,
                device_serial=serial,
                device_label=label,
                punch_count=count,
                last_punch_date=last_date,
                last_punch_time=time(9, 5, 7),
                resolved=resolved,
                resolved_note=note,
            )
            UnmatchedPunch.objects.filter(pk=row.pk).update(
                first_seen_at=utc(1, 4, 0),
                last_seen_at=datetime(last_date.year, last_date.month, last_date.day, 4, 0, tzinfo=timezone.utc),
            )
            return row

        unmatched("9001", "SN-A", "Gate A", 12, d(3))
        unmatched("9002", "", "", 5, date(2026, 9, 2))
        unmatched(cls.leo.employee_code, "SN-A", "Gate A", 3, d(3))
        unmatched("9003", "SN-B", "Gate B", 7, d(4), resolved=True, note="Test device")

        # ── users ────────────────────────────────────────────────────────────
        cls.admin = HRUser.objects.create(username="ac_admin", password_hash="x", is_super_admin=True)
        cls.role_ok = Role.objects.create(name="ac_ok", permissions={"reports": "view", "attendance": "view"})
        cls.branch_user = HRUser.objects.create(username="ac_b1", password_hash="x", role=cls.role_ok, branch=cls.b1)
        cls.no_module = HRUser.objects.create(
            username="ac_none",
            password_hash="x",
            role=Role.objects.create(name="ac_none", permissions={"reports": "view"}),
        )

    @classmethod
    def tearDownClass(cls):
        _holiday_dates_for_month.cache_clear()
        super().tearDownClass()

    def setUp(self):
        patcher = mock.patch("api.reporting.filters.ist_today", return_value=TODAY)
        patcher.start()
        self.addCleanup(patcher.stop)
        _holiday_dates_for_month.cache_clear()
        self._registered = []

    def tearDown(self):
        for rid in self._registered:
            registry._REGISTRY.pop(rid, None)

    # -- helpers ---------------------------------------------------------------
    def get(self, path, user=None, **params):
        return self.client.get(path, params, **hr_headers(user or self.admin))

    def run_report(self, rid, user=None, **params):
        r = self.get(f"/api/reports/run/{rid}", user, **params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    def export(self, rid, fmt, user=None, **params):
        return self.get(f"/api/reports/export/{rid}", user, fmt=fmt, **params)

    def codes(self, body, key="employeeCode"):
        return [r[key] for r in body["rows"] if not r.get("_kind")]

    def data_rows(self, body):
        return [r for r in body["rows"] if not r.get("_kind")]

    def row(self, body, code, date_iso=None):
        for r in body["rows"]:
            if r.get("_kind"):
                continue
            if r.get("employeeCode") == code and (date_iso is None or r.get("date") == date_iso):
                return r
        raise AssertionError(f"no row for {code} {date_iso}")


class SmokeTests(Base):
    def test_every_report_is_registered_and_well_formed(self):
        registry.all_specs()
        self.assertEqual(registry.LOAD_ERRORS, {})
        for rid in IDS:
            spec = registry.get_spec(rid)
            self.assertIsNotNone(spec, rid)
            self.assertEqual(spec.category, "attendance")
            self.assertEqual(spec.modules, ("attendance",))
            self.assertRegex(spec.id, "^[a-z0-9-]+$")

    def test_every_report_runs_and_exports_for_admin(self):
        for rid in IDS:
            body = self.run_report(rid, **PERIOD, period="2026-08")
            self.assertIn("columns", body, rid)
            for fmt, sig in (("xlsx", b"PK"), ("pdf", b"%PDF")):
                r = self.export(rid, fmt, **PERIOD, period="2026-08")
                self.assertEqual(r.status_code, 200, f"{rid} {fmt}: {r.content[:200]}")
                self.assertTrue(r.content.startswith(sig), f"{rid} {fmt}")


class Kit(Base):
    """Helpers for building extra employees and opening the exports."""

    def add_staff(self, code, first, *, dept=None, branch=None, production=False, **kw):
        emp = Employee.objects.create(
            employee_code=code,
            first_name=first,
            last_name="Q",
            department=dept or self.cutting,
            branch=branch or self.b1,
            employment_type="production" if production else "staff",
            **kw,
        )
        if not production:
            EmployeeShiftAssignment.objects.create(employee=emp, shift=self.shift, effective_from=date(2020, 1, 1))
        return emp

    def sheet(self, rid, user=None, **params):
        r = self.export(rid, "xlsx", user, **params)
        self.assertEqual(r.status_code, 200, r.content[:200])
        return load_workbook(io.BytesIO(r.content)).active

    def pdf_pages(self, rid, user=None, **params):
        r = self.export(rid, "pdf", user, **params)
        self.assertEqual(r.status_code, 200, r.content[:200])
        with pdfplumber.open(io.BytesIO(r.content)) as pdf:
            return [page.extract_text() or "" for page in pdf.pages]

    def count_queries(self, rid, **params):
        with CaptureQueriesContext(connection) as ctx:
            r = self.get(f"/api/reports/run/{rid}", **params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return len(ctx)


# ══ Time Card ═══════════════════════════════════════════════════════════════════


class TimeCardTests(Kit):
    def card(self, **params):
        both = f"{self.alice.id},{self.bob.id}"
        return self.run_report("time-card", **{**PERIOD, "employeeIds": both, **params})

    def test_golden_rows_alice(self):
        body = self.card()
        r = self.row(body, "AC001", "2026-08-03")
        self.assertEqual(
            (r["in1"], r["out1"], r["in2"], r["out2"], r["punchCount"]), ("08:55", "13:00", "13:45", "18:05", 4)
        )
        self.assertEqual(r["workedHours"], 8.42)  # (13:00-08:55) + (18:05-13:45) = 245 + 260 = 505 min
        self.assertEqual((r["status"], r["shiftsEarned"], r["shift"], r["day"]), ("Present", 1.0, "General", "Mon"))
        self.assertIsNone(r["lateMinutes"])
        r = self.row(body, "AC001", "2026-08-04")  # one lone punch before the cut-off: a half day, hours unknown
        self.assertEqual((r["in1"], r["out1"], r["workedHours"], r["punchCount"]), ("09:10", None, None, 1))
        self.assertEqual((r["status"], r["halfDay"], r["shiftsEarned"]), ("Half Day", "Morning", 0.5))
        self.assertIn("Single punch", r["remarks"])
        r = self.row(body, "AC001", "2026-08-05")  # six punches incl. two double taps: 09:00 13:00 13:45 18:00 counted
        self.assertEqual(r["punchCount"], 6)
        self.assertEqual(r["workedHours"], 8.25)  # 240 + 255 min
        self.assertIn("Punches: 09:00, 09:02, 13:00, 13:45, 18:00, 18:01", r["remarks"])
        r = self.row(body, "AC001", "2026-08-06")
        self.assertEqual(r["status"], "Holiday")
        self.assertIn("Holiday: Test Holiday", r["remarks"])
        self.assertIsNone(r["in1"])
        self.assertEqual(self.row(body, "AC001", "2026-08-07")["status"], "Absent")
        self.assertEqual(self.row(body, "AC001", "2026-08-08")["status"], "Absent")
        self.assertEqual(self.row(body, "AC001", "2026-08-09")["status"], "Weekly Off")
        r = self.row(body, "AC001", "2026-08-10")  # frozen manual day: HR typed the times, no punches behind them
        self.assertEqual((r["in1"], r["out1"], r["workedHours"], r["punchCount"]), ("09:00", "18:00", None, None))
        self.assertIn("Manual entry by HR Anita: Missed punch due to power cut", r["remarks"])

    def test_golden_rows_bob(self):
        body = self.card()
        r = self.row(body, "AC002", "2026-08-03")
        self.assertEqual(
            (r["lateMinutes"], r["workedHours"], r["status"]), (40, 7.5, "Present")
        )  # 09:40 vs 09:00; 510-60 min
        self.assertIsNone(self.row(body, "AC002", "2026-08-03")["otMinutes"])  # 18:10 is only 10 min past the end
        r = self.row(body, "AC002", "2026-08-04")
        self.assertEqual((r["otMinutes"], r["workedHours"]), (90, 9.5))  # 19:30 - 18:00; 630-60 min
        r = self.row(body, "AC002", "2026-08-05")  # 3 punches: no reliable hours
        self.assertEqual((r["punchCount"], r["workedHours"]), (3, None))
        self.assertIn("Odd punches (3)", r["remarks"])
        r = self.row(body, "AC002", "2026-08-10")  # in-cap Morning Late-In permission moves the start to 10:00
        self.assertEqual((r["permissionMinutes"], r["lateMinutes"], r["workedHours"]), (60, None, 7.0))
        self.assertIn("Permission: Morning Late-In", r["remarks"])

    def test_early_out_half_day_evening_and_production(self):
        body = self.run_report("time-card", **PERIOD, employeeIds=f"{self.carol.id},{self.pat.id}")
        r = self.row(body, "AC003", "2026-08-04")
        self.assertEqual((r["earlyOutMinutes"], r["workedHours"], r["status"]), (180, 5.0, "Present"))  # 18:00 - 15:00
        r = self.row(body, "AC003", "2026-08-07")  # only a 15:00 punch: evening half, and 360 min after the start
        self.assertEqual((r["status"], r["halfDay"], r["lateMinutes"]), ("Half Day", "Evening", 360))
        r = self.row(body, "AC003", "2026-08-05")
        self.assertEqual(r["status"], "Leave")
        self.assertIn("Approved leave (SL)", r["remarks"])
        r = self.row(body, "AC003", "2026-08-08")  # a Saturday-off Saturday is a weekly off, not an absence
        self.assertEqual(r["status"], "Weekly Off")
        r = self.row(body, "AC004", "2026-08-03")
        self.assertEqual((r["workedHours"], r["shiftsEarned"], r["shift"]), (10.75, 1.5, None))
        r = self.row(body, "AC004", "2026-08-04")  # production: late against the 08:30 reference + 10 min grace
        self.assertEqual((r["lateMinutes"], r["shiftsEarned"], r["status"]), (20, 1.25, "Present"))
        self.assertEqual(self.row(body, "AC004", "2026-08-09")["status"], "Absent")  # production works Sundays
        self.assertEqual(self.row(body, "AC004", "2026-08-06")["status"], "Holiday")

    def test_subtotals_and_totals_equal_the_sum_of_rows(self):
        body = self.card()
        keys = ["workedHours", "lateMinutes", "earlyOutMinutes", "otMinutes", "permissionMinutes", "shiftsEarned"]
        data = self.data_rows(body)
        self.assertEqual(len(data), 16)
        subs = {r["employeeName"]: r for r in body["rows"] if r.get("_kind") == "subtotal"}
        self.assertEqual(set(subs), {"AC001 total", "AC002 total"})
        for code in ("AC001", "AC002"):
            mine = [r for r in data if r["employeeCode"] == code]
            for k in keys:
                self.assertAlmostEqual(subs[f"{code} total"][k], sum(r[k] or 0 for r in mine), places=2, msg=k)
        for k in keys:
            self.assertAlmostEqual(body["totals"][k] or 0, sum(r[k] or 0 for r in data), places=2, msg=k)
        self.assertEqual(body["totals"]["workedHours"], 40.67)
        self.assertEqual(
            (body["totals"]["lateMinutes"], body["totals"]["otMinutes"], body["totals"]["permissionMinutes"]),
            (40, 90, 60),
        )
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (cards["Employees"], cards["Present days"], cards["Half days"], cards["Absent days"], cards["Late days"]),
            (2, 7, 1, 4, 1),
        )
        self.assertEqual(cards["Worked hours"], 40.67)

    def test_deduct_lunch_option(self):
        on = self.row(self.card(), "AC002", "2026-08-03")["workedHours"]
        off = self.row(self.card(deductLunch="false"), "AC002", "2026-08-03")["workedHours"]
        self.assertEqual((on, off), (7.5, 8.5))
        # a four-punch day is paired In->Out, so the lunch break is never subtracted twice
        self.assertEqual(self.row(self.card(deductLunch="false"), "AC001", "2026-08-03")["workedHours"], 8.42)

    def test_filters_narrow_the_result(self):
        def keys(**p):
            body = self.run_report("time-card", **{**PERIOD, **p})
            return {(r["employeeCode"], r["date"][-2:]) for r in self.data_rows(body)}

        self.assertEqual(keys(dayStatus="late"), {("AC002", "03"), ("AC003", "07"), ("AC004", "04")})
        self.assertEqual(keys(dayStatus="early_out"), {("AC003", "04")})
        self.assertEqual(
            keys(dayStatus="exceptions"),
            {("AC001", "04"), ("AC001", "05"), ("AC002", "05"), ("AC005", "07"), ("AC003", "07")},
        )
        self.assertEqual({c for c, _ in keys(dayStatus="half")}, {"AC001", "AC003"})
        self.assertEqual(keys(dayStatus="unprocessed"), set())
        self.assertEqual({day for _c, day in keys(dayStatus="off", employeeIds=str(self.alice.id))}, {"06", "09"})
        no_off = self.run_report("time-card", **PERIOD, employeeIds=str(self.alice.id), includeOffDays="false")
        self.assertEqual(len(self.data_rows(no_off)), 6)
        self.assertEqual({c for c, _ in keys(departmentIds=str(self.sewing.id))}, {"AC003", "AC004"})
        self.assertEqual({c for c, _ in keys(designationIds=str(self.cutter.id))}, {"AC001", "AC002", "AC006"})
        self.assertEqual({c for c, _ in keys(employmentType="production")}, {"AC004"})
        self.assertEqual({c for c, _ in keys(employeeStatus="inactive")}, {"AC006"})
        self.assertEqual({c for c, _ in keys(branchIds=str(self.b2.id))}, {"AC005"})
        self.assertEqual({c for c, _ in keys(employeeIds=f"{self.zed.id},{self.pat.id}")}, {"AC004", "AC005"})
        self.assertEqual({c for c, _ in keys()}, {"AC001", "AC002", "AC003", "AC004", "AC005", "AC006"})

    def test_inactive_employee_without_any_data_is_left_out_unless_chosen(self):
        ghost = self.add_staff("AC900", "Ghost", status="inactive")
        self.assertNotIn("AC900", self.codes(self.run_report("time-card", **PERIOD)))
        chosen = self.run_report("time-card", **PERIOD, employeeIds=str(ghost.id))
        self.assertEqual(len(self.data_rows(chosen)), 8)
        self.assertEqual({r["status"] for r in self.data_rows(chosen)}, {None})

    def test_exit_after_midnight_belongs_to_the_day_it_closes(self):
        ps = PayrollSettings.get()
        ps.last_punch_post_shift_grace_hours = 9
        ps.first_punch_pre_shift_buffer_hours = 2
        ps.save()
        eve = self.add_staff("AC910", "Eve")
        punch(eve, d(17), "09:00", "13:00", "13:45")
        punch(eve, d(18), "01:05", "09:00", "18:00")  # 01:05 is Monday's forgotten exit, stamped Tuesday
        seed_records(eve, d(17), d(18))
        body = self.run_report("time-card", dateFrom="2026-08-17", dateTo="2026-08-18", employeeIds=str(eve.id))
        mon = self.row(body, "AC910", "2026-08-17")
        self.assertEqual(
            (mon["in1"], mon["out1"], mon["in2"], mon["out2"], mon["punchCount"]),
            ("09:00", "13:00", "13:45", "01:05", 4),
        )
        self.assertEqual(mon["workedHours"], 15.33)  # 240 + (25:05 - 13:45 = 680) = 920 min - never a negative duration
        self.assertEqual((mon["status"], mon["otMinutes"]), ("Present", 425))  # 25:05 - 18:00
        self.assertIn("Exit after midnight", mon["remarks"])
        tue = self.row(body, "AC910", "2026-08-18")
        self.assertEqual((tue["in1"], tue["out1"], tue["punchCount"], tue["workedHours"]), ("09:00", "18:00", 2, 8.0))

    def test_days_outside_joining_and_leaving_dates_are_blank(self):
        fay = self.add_staff("AC920", "Fay", join_date="2026-08-05")
        ResignationRequest.objects.create(employee=fay, status="approved", last_working_date=d(6))
        gus = self.add_staff("AC921", "Gus", join_date="not a date")  # unparseable text is treated as "no date"
        for emp in (fay, gus):
            punch(emp, d(3), "09:00", "18:00")
            punch(emp, d(5), "09:00", "18:00")
            punch(emp, d(6), "09:00", "18:00")
            punch(emp, d(7), "09:00", "18:00")
            seed_records(emp, d(3), d(8))
        body = self.run_report(
            "time-card", dateFrom="2026-08-03", dateTo="2026-08-08", employeeIds=f"{fay.id},{gus.id}"
        )
        fay_days = {r["date"][-2:]: r for r in self.data_rows(body) if r["employeeCode"] == "AC920"}
        self.assertEqual(
            {k: v["status"] for k, v in fay_days.items()},
            {"03": None, "04": None, "05": "Present", "06": "Present", "07": None, "08": None},
        )
        self.assertIn("Outside joining / leaving dates", fay_days["03"]["remarks"])
        self.assertIsNone(fay_days["03"]["in1"])  # the punch on the 3rd is not shown either
        gus_days = {r["date"][-2:]: r["status"] for r in self.data_rows(body) if r["employeeCode"] == "AC921"}
        self.assertEqual(gus_days["03"], "Present")

    def test_unprocessed_and_future_days_are_blank_counted_and_nothing_is_written(self):
        before = (AttendanceDayRecord.objects.count(), DailyShiftLog.objects.count())
        body = self.run_report("time-card", dateFrom="2026-08-11", dateTo="2026-08-12", employeeIds=str(self.alice.id))
        rows = self.data_rows(body)
        self.assertEqual([r["status"] for r in rows], [None, None])
        self.assertIn("Not processed yet", rows[0]["remarks"])
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(cards["Days not processed"], 2)
        # the same days, one of them after "today": future days are not "not processed"
        body = self.run_report("time-card", dateFrom="2026-09-05", dateTo="2026-09-07", employeeIds=str(self.alice.id))
        self.assertEqual({c["label"]: c["value"] for c in body["summary"]}["Days not processed"], 1)  # 5th only
        self.assertEqual((AttendanceDayRecord.objects.count(), DailyShiftLog.objects.count()), before)

    def test_long_ranges_need_a_narrow_employee_list(self):
        r = self.get("/api/reports/run/time-card", dateFrom="2026-07-01", dateTo="2026-08-15")  # 46 days, 6 employees
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["field"], "employee")
        ok = self.get(
            "/api/reports/run/time-card", dateFrom="2026-07-01", dateTo="2026-08-15", employeeIds=str(self.alice.id)
        )
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(ok.json()["rowCount"], 47)  # 46 days + the employee's total row
        wide = self.get(
            "/api/reports/run/time-card", dateFrom="2026-01-01", dateTo="2026-08-15", employeeIds=str(self.alice.id)
        )
        self.assertEqual(wide.status_code, 400)
        self.assertEqual(wide.json()["field"], "dateTo")

    def test_pdf_is_one_page_per_employee_with_totals_and_signatures(self):
        pages = self.pdf_pages("time-card", **PERIOD, employeeIds=f"{self.alice.id},{self.bob.id}")
        self.assertEqual(len(pages), 2)
        self.assertIn("TIME CARD", pages[0])
        self.assertIn("Alice T", pages[0])
        self.assertIn("AC001", pages[0])
        self.assertIn("CUTTING", pages[0])
        self.assertIn("Cutter", pages[0])
        self.assertIn("03-Aug-26", pages[0])
        self.assertIn("Test Holiday", pages[0])
        self.assertIn("Employee signature", pages[0])
        self.assertIn("Supervisor / HOD", pages[0])
        self.assertIn("Bob T", pages[1])
        self.assertNotIn("Alice T", pages[1])
        self.assertIn("Worked hrs", pages[1])
        self.assertIn("24.00", pages[1])  # Bob's total worked hours in the footer

    def test_xlsx_round_trip(self):
        ws = self.sheet("time-card", **PERIOD, employeeIds=str(self.bob.id))
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:6], ["Emp Code", "Employee", "Department", "Date", "Day", "Shift"])
        self.assertEqual(header[11], "Worked Hrs")
        first = [c.value for c in ws[8]]
        self.assertEqual((first[0], first[1], first[4], first[6]), ("AC002", "Bob T", "Mon", time(9, 40)))
        self.assertEqual((first[11], first[12]), (7.5, 40))  # worked hours, late minutes


# ══ Daily Attendance Register ═══════════════════════════════════════════════════


class DailyAttendanceTests(Kit):
    def day(self, day=3, **params):
        """One day for the ACTIVE staff (the fixtures' left employee AC006 is exercised on its own below)."""
        iso = f"2026-08-{day:02d}"
        return self.run_report("daily-attendance", dateFrom=iso, dateTo=iso, **{"employeeStatus": "active", **params})

    def test_golden_day_and_legacy_ordering(self):
        body = self.day(3)
        self.assertEqual(self.codes(body), ["AC001", "AC002", "AC005", "AC003", "AC004"])  # department, then first name
        by = {r["employeeCode"]: r for r in body["rows"]}
        a = by["AC001"]
        self.assertEqual(
            (a["firstIn"], a["lastOut"], a["punchCount"], a["workedHours"], a["status"]),
            ("08:55", "18:05", 4, 8.42, "Present"),
        )
        self.assertEqual(
            (a["designation"], a["employeeType"], a["shift"], a["source"]), ("Cutter", "Staff", "General", "Biometric")
        )
        b = by["AC002"]
        self.assertEqual((b["lateMinutes"], b["flags"], b["workedHours"]), (40, "Late", 7.5))
        self.assertEqual((by["AC005"]["source"], by["AC005"]["department"]), ("Geo Punch", "PACKING"))
        c = by["AC003"]
        self.assertEqual(
            (c["status"], c["informed"], c["firstIn"], c["shiftsEarned"]), ("Absent", "Informed", None, 0.0)
        )
        p = by["AC004"]
        self.assertEqual(
            (p["employeeType"], p["shiftsEarned"], p["workedHours"], p["shift"]), ("Production", 1.5, 10.75, None)
        )
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (
                cards["Employee-days"],
                cards["Present"],
                cards["Half day"],
                cards["Absent"],
                cards["On leave"],
                cards["Late"],
            ),
            (5, 4, 0, 1, 0, 1),
        )
        self.assertEqual((cards["Strength %"], cards["Not processed"]), (80.0, 0))
        self.assertEqual(body["totals"]["workedHours"], 34.67)
        self.assertEqual(body["totals"]["lateMinutes"], 40)

    def test_day_with_half_day_leave_permission_and_manual_flags(self):
        body = self.run_report("daily-attendance", **PERIOD, employeeStatus="active")
        rows = {(r["employeeCode"], r["date"][-2:]): r for r in self.data_rows(body)}
        self.assertEqual(rows[("AC001", "04")]["status"], "Half Day")
        self.assertEqual(rows[("AC001", "04")]["flags"], None)
        self.assertEqual(rows[("AC002", "10")]["flags"], "Permission")
        self.assertEqual(rows[("AC001", "10")]["flags"], "HR override")
        self.assertEqual(rows[("AC003", "10")]["status"], "Casual Leave")
        self.assertEqual(rows[("AC004", "10")]["status"], "Comp Off")
        self.assertEqual(rows[("AC003", "05")]["status"], "Leave")
        self.assertEqual(rows[("AC003", "05")]["leaveType"], "SL")
        self.assertEqual(rows[("AC001", "06")]["status"], "Holiday")
        self.assertEqual(rows[("AC001", "09")]["status"], "Weekly Off")
        self.assertEqual(rows[("AC003", "08")]["status"], "Weekly Off")
        self.assertEqual(rows[("AC003", "04")]["flags"], "Early out")

    def test_multi_day_range_is_ordered_by_date_then_department(self):
        body = self.run_report("daily-attendance", dateFrom="2026-08-03", dateTo="2026-08-04", employeeStatus="active")
        dates = [r["date"] for r in self.data_rows(body)]
        self.assertEqual(dates, ["2026-08-03"] * 5 + ["2026-08-04"] * 5)
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual((cards["Present"], cards["Half day"], cards["Absent"]), (4 + 3, 1, 1 + 1))

    def test_status_flag_and_informed_filters(self):
        def keys(**p):
            return {
                (r["employeeCode"], r["date"][-2:])
                for r in self.data_rows(
                    self.run_report("daily-attendance", **{**PERIOD, "employeeStatus": "active", **p})
                )
            }

        self.assertEqual(keys(status="half"), {("AC001", "04"), ("AC003", "07")})
        self.assertEqual(keys(status="leave"), {("AC003", "05")})
        self.assertEqual({c for c, _ in keys(status="holiday")}, {"AC001", "AC002", "AC003", "AC004", "AC005"})
        self.assertEqual({day for _, day in keys(status="weekly_off")}, {"08", "09"})
        self.assertEqual(keys(status="unprocessed"), set())
        self.assertEqual(keys(flag="late"), {("AC002", "03"), ("AC003", "07"), ("AC004", "04")})
        self.assertEqual(keys(flag="early_out"), {("AC003", "04")})
        self.assertEqual(keys(flag="half_day"), {("AC001", "04"), ("AC003", "07")})
        self.assertEqual(keys(flag="permission"), {("AC002", "10")})
        self.assertEqual(keys(flag="manual"), {("AC001", "10"), ("AC003", "10"), ("AC004", "10")})
        self.assertEqual(keys(flag="compensation_day"), set())
        self.assertEqual(
            keys(flag="missing_punch"), {("AC001", "04"), ("AC002", "05"), ("AC005", "07"), ("AC003", "07")}
        )
        self.assertEqual(keys(informed="informed"), {("AC003", "03")})
        self.assertEqual(keys(informed="not_informed"), {("AC005", "05")})
        unset = keys(informed="unset")
        self.assertIn(("AC003", "05"), unset)  # an on-leave day HR has not marked
        self.assertNotIn(("AC003", "03"), unset)

    def test_scope_and_search_filters(self):
        def codes(**p):
            return set(self.codes(self.day(3, **p)))

        def by_default(**p):  # the filters exactly as the form sends them: no employeeStatus at all
            body = self.run_report("daily-attendance", dateFrom="2026-08-03", dateTo="2026-08-03", **p)
            return set(self.codes(body))

        # a register of a past day keeps people who have since left, as long as they have activity in the range
        self.assertEqual(by_default(), {"AC001", "AC002", "AC003", "AC004", "AC005", "AC006"})
        self.assertEqual(codes(), {"AC001", "AC002", "AC003", "AC004", "AC005"})  # ... "Active" narrows it
        self.assertEqual(codes(employeeStatus="all"), {"AC001", "AC002", "AC003", "AC004", "AC005", "AC006"})
        self.assertEqual(codes(employeeStatus="inactive"), {"AC006"})
        ghost = self.add_staff("AC901", "Gone", status="inactive")  # left, and nothing at all in the range
        self.assertNotIn("AC901", by_default())
        self.assertIn("AC901", by_default(employeeIds=str(ghost.id)))  # unless chosen explicitly
        self.assertEqual(codes(departmentIds=str(self.sewing.id)), {"AC003", "AC004"})
        self.assertEqual(codes(designationIds=str(self.tailor.id)), {"AC003", "AC004"})
        self.assertEqual(codes(employmentType="production"), {"AC004"})
        self.assertEqual(codes(employmentType="staff"), {"AC001", "AC002", "AC003", "AC005"})
        self.assertEqual(codes(branchIds=str(self.b2.id)), {"AC005"})
        self.assertEqual(codes(employeeIds=f"{self.alice.id},{self.bob.id}"), {"AC001", "AC002"})
        self.assertEqual(codes(search="ali"), {"AC001"})
        self.assertEqual(codes(search="AC00"), {"AC001", "AC002", "AC003", "AC004", "AC005"})

    def test_strength_counts_heads_present_or_half(self):
        cards = {c["label"]: c["value"] for c in self.day(4)["summary"]}
        # Alice half + Bob + Carol + Pat present = 4 heads; Zed absent -> 4 / 5
        self.assertEqual((cards["Present"], cards["Half day"], cards["Absent"], cards["Strength %"]), (3, 1, 1, 80.0))

    def test_unprocessed_day_is_blank_counted_and_not_computed(self):
        before = AttendanceDayRecord.objects.count()
        body = self.day(11)
        self.assertEqual({r["status"] for r in self.data_rows(body)}, {None})
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual((cards["Not processed"], cards["Present"], cards["Absent"]), (5, 0, 0))
        self.assertEqual(AttendanceDayRecord.objects.count(), before)

    def test_today_absent_is_provisional(self):
        emp = self.add_staff("AC930", "Ivy")
        seed_records(emp, d(12), d(12))
        with mock.patch("api.reporting.filters.ist_today", return_value=d(12)):
            body = self.run_report(
                "daily-attendance", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id)
            )
        row = self.row(body, "AC930")
        self.assertEqual(row["status"], "Absent")
        self.assertIn("Provisional", row["flags"])
        self.assertTrue(any("provisional" in n for n in body["notes"]))

    def test_exports(self):
        ws = self.sheet("daily-attendance", dateFrom="2026-08-03", dateTo="2026-08-03")
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:3], ["Date", "Emp Code", "Employee"])
        self.assertEqual([ws.cell(row=8, column=2).value, ws.cell(row=8, column=3).value], ["AC001", "Alice T"])
        pages = self.pdf_pages("daily-attendance", dateFrom="2026-08-03", dateTo="2026-08-03")
        self.assertIn("DAILY ATTENDANCE REGISTER", pages[0])
        self.assertIn("Alice T", pages[0])


# ══ Absentee list + Report Log parity ══════════════════════════════════════════


class AbsenteeListTests(Kit):
    def listing(self, day=3, last=None, **params):
        return self.run_report(
            "absentee-list-daily", dateFrom=f"2026-08-{day:02d}", dateTo=f"2026-08-{(last or day):02d}", **params
        )

    def test_golden_absent_list_with_informed_status(self):
        body = self.listing(3)
        self.assertEqual(len(body["rows"]), 1)
        r = body["rows"][0]
        self.assertEqual(
            (
                r["sno"],
                r["employeeCode"],
                r["employeeName"],
                r["department"],
                r["designation"],
                r["status"],
                r["informedStatus"],
            ),
            (1, "AC003", "Carol T", "SEWING", "Tailor", "Absent", "Informed"),
        )
        self.assertIn("Staff leave list for 03.08.2026", body["notes"])
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual((cards["Listed"], cards["Informed"], cards["Not Informed"], cards["Unset"]), (1, 1, 0, 0))

    def test_multi_day_numbering_restarts_each_day(self):
        body = self.listing(3, 5, show="both")
        got = [(r["date"][-2:], r["sno"], r["employeeCode"], r["informedStatus"]) for r in body["rows"]]
        self.assertEqual(
            got,
            [
                ("03", 1, "AC003", "Informed"),
                ("04", 1, "AC005", "Unset"),
                ("05", 1, "AC005", "Not Informed"),
                ("05", 2, "AC003", "Unset"),  # on approved leave: listed under "both", after PACKING
            ],
        )

    def test_show_and_informed_filters(self):
        self.assertEqual(self.codes(self.listing(5, show="leave")), ["AC003"])
        self.assertEqual(self.codes(self.listing(5, show="absent")), ["AC005"])
        self.assertEqual(self.codes(self.listing(3, 5, informed="not_informed")), ["AC005"])
        self.assertEqual(self.codes(self.listing(3, 5, informed="informed")), ["AC003"])
        self.assertEqual(sorted(self.codes(self.listing(3, 5, informed="unset"))), ["AC005"])

    def test_staff_by_default_production_optional_and_search(self):
        self.assertNotIn("AC004", self.codes(self.listing(5)))  # Pat is production
        self.assertEqual(self.codes(self.listing(5, staffType="all")), ["AC005", "AC004"])  # PACKING, then SEWING
        self.assertEqual(self.codes(self.listing(5, staffType="production")), ["AC004"])
        self.assertEqual(self.codes(self.listing(5, staffType="all", search="pat")), ["AC004"])
        self.assertEqual(self.codes(self.listing(8, departmentIds=str(self.cutting.id))), ["AC001", "AC002"])

    def test_saturday_off_and_sunday_are_not_absences(self):
        body = self.listing(8)  # Saturday: Carol's schedule has Saturdays off
        self.assertEqual(self.codes(body), ["AC001", "AC002", "AC005"])
        self.assertEqual({c["label"]: c["value"] for c in body["summary"]}["Weekly off (left out)"], 1)
        self.assertEqual(self.codes(self.listing(9)), [])  # Sunday: every staff verdict is a holiday

    def test_a_saturday_off_saturday_with_a_punch_is_still_an_absence(self):
        punch(self.carol, d(15), "14:00")  # a lone punch in the 13:30-14:30 gap: neither half attended, so absent
        seed_records(self.carol, d(15), d(15))
        self.assertEqual(AttendanceDayRecord.objects.get(employee=self.carol, date=d(15)).status, "absent")
        body = self.listing(15)
        self.assertEqual([(r["employeeCode"], r["status"]) for r in body["rows"]], [("AC003", "Absent")])
        self.assertEqual({c["label"]: c["value"] for c in body["summary"]}["Weekly off (left out)"], 0)

    def test_ordering_matches_report_log(self):
        body = self.listing(8, staffType="all")
        self.assertEqual([r["department"] for r in body["rows"]], sorted(r["department"] for r in body["rows"]))

    def test_unprocessed_days_are_counted_not_listed(self):
        body = self.listing(11)
        self.assertEqual(body["rows"], [])
        self.assertEqual({c["label"]: c["value"] for c in body["summary"]}["Not processed"], 4)  # the four active staff
        self.assertEqual(
            self.get("/api/reports/run/absentee-list-daily", dateFrom="2026-08-01", dateTo="2026-09-15").status_code,
            400,
        )

    def test_today_is_provisional(self):
        emp = self.add_staff("AC931", "Jay")
        seed_records(emp, d(12), d(12))
        with mock.patch("api.reporting.filters.ist_today", return_value=d(12)):
            body = self.run_report(
                "absentee-list-daily", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id)
            )
        self.assertEqual(self.codes(body), ["AC931"])
        self.assertTrue(any("provisional" in n for n in body["notes"]))

    def test_matches_the_report_log_daily_report(self):
        """The legacy Daily Report (GET /api/attendance/report-log?date=) lists every active staff member; its UI keeps
        the absent ones. The Report Center list must show exactly those, with the same Informed call."""
        for day in (3, 4, 5, 8):
            legacy = self.get("/api/attendance/report-log", date=f"2026-08-{day:02d}").json()
            want = []
            for row in legacy["rows"]:
                if row["status"] != "absent":
                    continue
                if day == 8 and row["employeeCode"] == "AC003":
                    continue  # the one deliberate difference: a Saturday-off Saturday is not an absence
                informed = {True: "Informed", False: "Not Informed", None: "Unset"}[row["isInformed"]]
                want.append((row["employeeCode"], informed))
            mine = self.listing(day)
            self.assertEqual([(r["employeeCode"], r["informedStatus"]) for r in mine["rows"]], want, f"Aug {day}")

    def test_exports(self):
        ws = self.sheet("absentee-list-daily", dateFrom="2026-08-05", dateTo="2026-08-05", show="both")
        header = [c.value for c in ws[7]]
        self.assertEqual(
            header,
            ["S.No", "Date", "Ticket No", "Employee Name", "Department", "Designation", "Status", "Informed Status"],
        )
        self.assertEqual([c.value for c in ws[8]][2:4], ["AC005", "Zed T"])
        self.assertIn("Not Informed", [c.value for c in ws[8]])
        pages = self.pdf_pages("absentee-list-daily", dateFrom="2026-08-05", dateTo="2026-08-05", show="both")
        self.assertIn("Zed T", pages[0])
        self.assertIn("Staff leave list for 05.08.2026", pages[0])


# ══ Muster sheet + Report Log parity ═══════════════════════════════════════════

LEGACY_CODE = {"present": "P", "absent": "A", "on_leave": "L", "holiday": "H"}


class MusterSheetTests(Kit):
    DAYS = [f"d202608{n:02d}" for n in range(3, 11)]

    def sheet_rows(self, **params):
        body = self.run_report("attendance-muster-sheet", **{**PERIOD, **params})
        rows = {r["employeeCode"]: r for r in body["rows"] if not r.get("_kind")}
        return body, rows

    def codes_of(self, row):
        return [row[k] for k in self.DAYS]

    def test_golden_codes_and_totals(self):
        body, rows = self.sheet_rows()
        self.assertEqual(list(rows), ["AC001", "AC002", "AC003", "AC005"])  # active staff, ordered by first name
        self.assertEqual(self.codes_of(rows["AC001"]), ["P", "½M", "P", "H", "A", "A", "H", "P"])
        self.assertEqual(self.codes_of(rows["AC002"]), ["P", "P", "P", "H", "A", "A", "H", "P"])
        self.assertEqual(self.codes_of(rows["AC003"]), ["A", "P", "L", "H", "½E", "A", "H", "P"])
        self.assertEqual(self.codes_of(rows["AC005"]), ["P", "A", "A", "H", "P", "A", "H", "A"])
        a = rows["AC001"]
        self.assertEqual(
            (a["present"], a["half"], a["absent"], a["leave"], a["holiday"], a["late"]), (3, 1, 2, 0, 2, 0)
        )
        self.assertEqual((a["effective"], a["shifts"]), (3.5, 3.5))  # 3 full days + half a day
        c = rows["AC003"]
        self.assertEqual(
            (c["present"], c["half"], c["absent"], c["leave"], c["holiday"], c["late"], c["effective"]),
            (2, 1, 2, 1, 2, 1, 2.5),
        )
        strength = next(r for r in body["rows"] if r.get("_kind") == "total")
        self.assertEqual(strength["employeeName"], "Strength")
        self.assertEqual([strength[k] for k in self.DAYS], [3, 3, 2, 0, 2, 0, 0, 3])  # heads present or half per date
        self.assertEqual(
            body["totals"],
            {
                "present": 11,
                "half": 2,
                "absent": 10,
                "leave": 1,
                "holiday": 8,
                "late": 2,
                "effective": 12.0,
                "shifts": 12.0,
            },
        )
        self.assertEqual([c["key"] for c in body["columns"]][:4], ["sno", "employeeCode", "employeeName", "department"])
        self.assertEqual([c["label"] for c in body["columns"]][4:6], ["3 Mo", "4 Tu"])
        self.assertEqual(
            [c["label"] for c in body["columns"]][-8:], ["P", "½", "A", "L", "H", "Late", "Eff.", "Shifts"]
        )
        self.assertEqual([r["sno"] for r in body["rows"] if not r.get("_kind")], [1, 2, 3, 4])

    def test_totals_summary_cards_and_strength_row_equal_the_rows(self):
        body, rows = self.sheet_rows(staffType="all")
        for key in ("present", "half", "absent", "leave", "holiday", "late"):
            self.assertEqual(body["totals"][key], sum(r[key] for r in rows.values()), key)
        for key in ("effective", "shifts"):
            self.assertAlmostEqual(body["totals"][key], sum(r[key] for r in rows.values()), places=2, msg=key)
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (cards["Employees"], cards["Man-days present"], cards["Absent days"]),
            (len(rows), body["totals"]["present"], body["totals"]["absent"]),
        )
        strength = next(r for r in body["rows"] if r.get("_kind") == "total")
        for key in self.DAYS:
            heads = sum(1 for r in rows.values() if r[key] in ("P", "½M", "½E", "½"))
            self.assertEqual(strength[key], heads, key)

    def test_matches_the_report_log_monthly_sheet(self):
        """Same employees, same status codes, same P / A / Eff. and the same Strength row as the legacy endpoint the
        Report Log page uses."""
        legacy = self.get("/api/attendance/report-log/sheet", **PERIOD).json()
        body, rows = self.sheet_rows()
        self.assertEqual(list(rows), [e["employeeCode"] for e in legacy["employees"]])
        for emp in legacy["employees"]:
            mine = rows[emp["employeeCode"]]
            for cell in emp["days"]:
                if cell["status"] == "half_shift":
                    want = "½M" if cell["halfDayPeriod"] == "morning" else "½E"
                else:
                    want = LEGACY_CODE.get(cell["status"])
                self.assertEqual(
                    mine["d" + cell["date"].replace("-", "")], want, f"{emp['employeeCode']} {cell['date']}"
                )
            s = emp["summary"]
            self.assertEqual(
                (mine["present"], mine["absent"], mine["effective"]),
                (s["present"], s["absent"], float(s["effectiveDays"])),
            )
            self.assertEqual(
                (mine["half"], mine["leave"], mine["holiday"], mine["late"]),
                (s["halfShift"], s["onLeave"], s["holidays"], s["late"]),
            )
            self.assertEqual(mine["shifts"], float(s["totalShifts"]))
        strength = next(r for r in body["rows"] if r.get("_kind") == "total")
        self.assertEqual([strength[k] for k in self.DAYS], legacy["strength"])
        self.assertEqual([e["employeeName"] for e in legacy["employees"]], [r["employeeName"] for r in rows.values()])

    def test_weekly_off_option_separates_sundays_and_saturday_off(self):
        body, rows = self.sheet_rows(weeklyOff="true")
        self.assertEqual(
            self.codes_of(rows["AC001"]), ["P", "½M", "P", "H", "A", "A", "WO", "P"]
        )  # Aug 6 is a holiday row
        self.assertEqual(
            self.codes_of(rows["AC003"]), ["A", "P", "L", "H", "½E", "WO", "WO", "P"]
        )  # Sat 8 is Carol's day off
        c = rows["AC003"]
        self.assertEqual((c["absent"], c["holiday"], c["weeklyOff"]), (1, 1, 2))
        self.assertEqual(rows["AC001"]["weeklyOff"], 1)
        self.assertEqual(body["totals"]["weeklyOff"], 5)  # the Sunday for four staff + Carol's Saturday
        self.assertIn("WO", [c["label"] for c in body["columns"]])

    def test_leave_codes_option(self):
        _b, rows = self.sheet_rows(showLeaveCodes="true")
        self.assertEqual(rows["AC003"]["d20260805"], "SL")
        self.assertEqual(rows["AC003"]["d20260810"], "CL")
        self.assertEqual(rows["AC001"]["d20260810"], "P")  # an HR override is still a plain present day
        self.assertEqual(self.sheet_rows()[1]["AC003"]["d20260810"], "P")

    def test_scope_filters(self):
        def codes(**p):
            return list(self.sheet_rows(**p)[1])

        self.assertEqual(codes(staffType="all"), ["AC001", "AC002", "AC003", "AC004", "AC005"])
        self.assertEqual(codes(staffType="production"), ["AC004"])
        self.assertEqual(codes(employeeStatus="all"), ["AC001", "AC002", "AC003", "AC006", "AC005"])  # by first name
        self.assertEqual(codes(employeeStatus="inactive"), ["AC006"])
        self.assertEqual(codes(departmentIds=str(self.cutting.id)), ["AC001", "AC002"])
        self.assertEqual(codes(designationIds=str(self.tailor.id)), ["AC003"])
        self.assertEqual(codes(search="zed"), ["AC005"])
        self.assertEqual(codes(branchIds=str(self.b2.id)), ["AC005"])
        self.assertEqual(codes(employeeIds=f"{self.alice.id},{self.carol.id}"), ["AC001", "AC003"])

    def test_production_rows_use_a_single_half_code_and_shift_credit(self):
        _b, rows = self.sheet_rows(staffType="production")
        pat = rows["AC004"]
        self.assertEqual(self.codes_of(pat), ["P", "P", "A", "H", "A", "A", "A", "P"])  # Sunday is a normal working day
        self.assertEqual(pat["shifts"], 3.75)  # 1.50 + 1.25 + 1.00

    def test_service_dates_can_be_blanked(self):
        fay = self.add_staff("AC940", "Faye", join_date="2026-08-05")
        for day in (3, 5, 6):
            punch(fay, d(day), "09:00", "18:00")
        seed_records(fay, d(3), d(8))
        _b, rows = self.sheet_rows(search="AC940")
        self.assertEqual(rows["AC940"]["d20260803"], "P")  # the legacy sheet shows it
        _b, rows = self.sheet_rows(search="AC940", maskService="true")
        self.assertEqual(
            (rows["AC940"]["d20260803"], rows["AC940"]["d20260804"], rows["AC940"]["d20260805"]), (None, None, "P")
        )
        self.assertEqual(rows["AC940"]["present"], 2)  # the 5th and 6th only

    def test_unprocessed_days_show_a_dash_and_are_counted(self):
        body = self.run_report("attendance-muster-sheet", dateFrom="2026-08-09", dateTo="2026-08-12")
        rows = {r["employeeCode"]: r for r in body["rows"] if not r.get("_kind")}
        self.assertEqual(
            [rows["AC001"][k] for k in ("d20260809", "d20260810", "d20260811", "d20260812")], ["H", "P", None, None]
        )
        self.assertEqual(rows["AC001"]["present"], 1)
        self.assertTrue(any("no attendance record yet" in n for n in body["notes"]))
        self.assertEqual(
            {c["label"]: c["value"] for c in body["summary"]}["Days with no record"], 8
        )  # 4 staff x 2 days

    def test_range_is_capped_at_31_days(self):
        r = self.get("/api/reports/run/attendance-muster-sheet", dateFrom="2026-08-01", dateTo="2026-09-05")
        self.assertEqual(r.status_code, 400)
        ok = self.get("/api/reports/run/attendance-muster-sheet", dateFrom="2026-08-01", dateTo="2026-08-31")
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(len([c for c in ok.json()["columns"] if c["key"].startswith("d20")]), 31)

    def test_exports(self):
        ws = self.sheet("attendance-muster-sheet", **PERIOD)
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:6], ["S.No", "Emp Code", "Employee Name", "Department", "3 Mo", "4 Tu"])
        row = [c.value for c in ws[8]]
        self.assertEqual((row[1], row[2], row[4], row[5]), ("AC001", "Alice T", "P", "½M"))
        pages = self.pdf_pages("attendance-muster-sheet", **PERIOD)
        self.assertIn("ATTENDANCE SHEET (MUSTER ROLL)", pages[0])
        self.assertIn("Strength", pages[0])
        self.assertIn("Alice T", pages[0])


# ══ Monthly attendance summary ══════════════════════════════════════════════════


class MonthlySummaryTests(Kit):
    def summary(self, **params):
        body = self.run_report("attendance-summary-monthly", period="2026-08", **params)
        return body, {r["employeeCode"]: r for r in body["rows"] if not r.get("_kind")}

    def test_golden_staff_rows(self):
        body, rows = self.summary()
        a = rows["AC001"]
        self.assertEqual(
            (
                a["workingDays"],
                a["presentDays"],
                a["halfDays"],
                a["absentDays"],
                a["leaveDays"],
                a["holidays"],
                a["weeklyOff"],
            ),
            (25, 3, 1, 2, 0, 1, 1),  # 31 days - 5 Sundays - the 6th
        )
        self.assertEqual((a["lateDays"], a["earlyOutDays"], a["effectiveDays"], a["totalShifts"]), (0, 0, 3.5, 3.5))
        self.assertEqual((a["attendancePct"], a["notProcessed"], a["employeeType"]), (14.0, 23, "Staff"))  # 3.5 / 25
        c = rows["AC003"]  # Saturdays off: 31 - 5 Sundays - 5 Saturdays - the 6th = 20 working days
        self.assertEqual(
            (
                c["workingDays"],
                c["presentDays"],
                c["halfDays"],
                c["absentDays"],
                c["leaveDays"],
                c["holidays"],
                c["weeklyOff"],
            ),
            (20, 2, 1, 1, 1, 1, 2),
        )
        self.assertEqual((c["casualLeaves"], c["permissions"], c["lateDays"], c["earlyOutDays"]), (1, 0, 1, 1))
        self.assertEqual((c["effectiveDays"], c["totalShifts"], c["attendancePct"]), (2.5, 2.5, 12.5))  # 2.5 / 20
        b = rows["AC002"]
        self.assertEqual(
            (b["presentDays"], b["absentDays"], b["lateDays"], b["permissions"], b["attendancePct"]), (4, 2, 1, 1, 16.0)
        )
        z = rows["AC005"]
        self.assertEqual((z["presentDays"], z["absentDays"], z["attendancePct"]), (2, 4, 8.0))

    def test_production_has_no_working_day_denominator(self):
        _body, rows = self.summary(employmentType="production")
        pat = rows["AC004"]
        self.assertEqual((pat["employeeType"], pat["workingDays"], pat["attendancePct"]), ("Production", None, None))
        self.assertEqual((pat["presentDays"], pat["absentDays"], pat["holidays"], pat["weeklyOff"]), (3, 4, 1, 0))
        self.assertEqual((pat["totalShifts"], pat["lateDays"]), (3.75, 1))

    def test_subtotals_totals_and_company_figures_equal_the_rows(self):
        body, rows = self.summary()
        add = [
            "workingDays",
            "presentDays",
            "halfDays",
            "absentDays",
            "leaveDays",
            "casualLeaves",
            "permissions",
            "holidays",
            "weeklyOff",
            "lateDays",
            "earlyOutDays",
            "effectiveDays",
            "totalShifts",
            "notProcessed",
        ]
        subs = {r["employeeName"]: r for r in body["rows"] if r.get("_kind") == "subtotal"}
        self.assertEqual(set(subs), {"CUTTING total", "PACKING total", "SEWING total"})
        # department attendance % is over the staff working days of the department (production has none):
        # CUTTING (3.5 + 4 + 1) / (25 x 3), SEWING Carol only 2.5 / 20, PACKING Zed only 2 / 25
        self.assertEqual(
            [subs[f"{dep} total"]["attendancePct"] for dep in ("CUTTING", "SEWING", "PACKING")], [11.3, 12.5, 8.0]
        )
        for dept, sub in (("CUTTING", subs["CUTTING total"]), ("SEWING", subs["SEWING total"])):
            mine = [r for r in rows.values() if r["department"] == dept]
            for k in add:
                self.assertAlmostEqual(sub[k], sum(r[k] or 0 for r in mine), places=2, msg=f"{dept} {k}")
        for k in add:
            self.assertAlmostEqual(body["totals"][k], sum(r[k] or 0 for r in rows.values()), places=2, msg=k)
        staff = [r for r in rows.values() if r["employeeType"] == "Staff"]
        cards = {c["label"]: c["value"] for c in body["summary"]}
        wd = sum(r["workingDays"] for r in staff)
        eff = sum(r["attendancePct"] * r["workingDays"] / 100.0 for r in staff)
        self.assertAlmostEqual(cards["Staff attendance %"], eff * 100.0 / wd, delta=0.1)
        self.assertEqual(body["totals"]["attendancePct"], cards["Staff attendance %"])
        absent = sum(r["absentDays"] for r in staff)
        self.assertAlmostEqual(cards["Staff absenteeism %"], absent * 100.0 / wd, delta=0.1)
        self.assertEqual(cards["Employees"], 6)
        self.assertEqual(cards["Days not processed"], 23 * 6)

    def test_scope_filters(self):
        def codes(**p):
            return set(self.summary(**p)[1])

        self.assertEqual(codes(departmentIds=str(self.sewing.id)), {"AC003", "AC004"})
        self.assertEqual(codes(employmentType="production"), {"AC004"})
        self.assertEqual(codes(employeeStatus="inactive"), {"AC006"})
        self.assertEqual(codes(branchIds=str(self.b2.id)), {"AC005"})
        self.assertEqual(codes(designationIds=str(self.cutter.id)), {"AC001", "AC002", "AC006"})
        self.assertEqual(codes(employeeIds=str(self.alice.id)), {"AC001"})
        self.assertEqual(codes(), {"AC001", "AC002", "AC003", "AC004", "AC005", "AC006"})
        ghost = self.add_staff("AC950", "Nobody", status="inactive")
        self.assertNotIn("AC950", codes())
        self.assertEqual(codes(employeeIds=str(ghost.id)), {"AC950"})

    def test_current_month_counts_elapsed_days_only(self):
        with mock.patch("api.reporting.filters.ist_today", return_value=d(10)):
            body = self.run_report("attendance-summary-monthly", period="2026-08", employeeIds=str(self.alice.id))
        row = self.data_rows(body)[0]
        self.assertEqual(
            row["workingDays"], 7
        )  # Aug 1-10: 8 weekdays/Saturdays minus Sundays 2 and 9, minus the 6th = 7
        self.assertEqual(row["notProcessed"], 2)  # the 1st and 2nd have no record
        self.assertTrue(any("counted up to 10-Aug-2026" in n for n in body["notes"]))

    def test_service_dates_are_respected_by_default(self):
        fay = self.add_staff("AC960", "Fern", join_date="2026-08-05")
        ResignationRequest.objects.create(employee=fay, status="approved", last_working_date=d(6))
        for day in (3, 5, 6, 7):
            punch(fay, d(day), "09:00", "18:00")
        seed_records(fay, d(3), d(8))
        _b, rows = self.summary(employeeIds=str(fay.id))
        self.assertEqual((rows["AC960"]["presentDays"], rows["AC960"]["absentDays"]), (2, 0))  # only the 5th and 6th
        _b, rows = self.summary(employeeIds=str(fay.id), maskService="false")
        self.assertEqual((rows["AC960"]["presentDays"], rows["AC960"]["absentDays"]), (4, 2))  # 4th and 8th are absent

    def test_exports(self):
        ws = self.sheet("attendance-summary-monthly", period="2026-08")
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:3], ["Emp Code", "Employee", "Department"])
        self.assertEqual(header[5], "Working Days")
        first = [c.value for c in ws[8]]
        self.assertEqual((first[0], first[1], first[5], first[6]), ("AC001", "Alice T", 25, 3))
        pages = self.pdf_pages("attendance-summary-monthly", period="2026-08")
        self.assertIn("MONTHLY ATTENDANCE SUMMARY", pages[0])
        self.assertIn("CUTTING total", "".join(pages))


# ══ Punch Log ═══════════════════════════════════════════════════════════════════


class PunchLogTests(Kit):
    def log(self, **params):
        return self.run_report("punch-log", **{"dateFrom": "2026-08-05", "dateTo": "2026-08-05", **params})

    def test_positional_direction_and_day_counts(self):
        body = self.log()
        alice = [r for r in body["rows"] if r["employeeCode"] == "AC001"]
        self.assertEqual([r["punchNo"] for r in alice], [1, 2, 3, 4, 5, 6])
        self.assertEqual([r["direction"] for r in alice], ["IN", "OUT", "IN", "OUT", "IN", "OUT"])
        self.assertEqual({r["dayPunches"] for r in alice}, {6})
        self.assertEqual(
            [r["punchTime"] for r in alice], ["09:00:00", "09:02:00", "13:00:00", "13:45:00", "18:00:00", "18:01:00"]
        )
        self.assertEqual({r["storedType"] for r in alice}, {"IN"})  # the device's own flag: everything says IN
        self.assertEqual({(r["sourceLabel"], r["device"]) for r in alice}, {("Biometric", "ADMS SN1")})
        bob = [r for r in body["rows"] if r["employeeCode"] == "AC002"]
        self.assertEqual(
            [(r["punchNo"], r["direction"], r["sourceLabel"]) for r in bob],
            [
                (1, "IN", "Biometric"),
                (2, "OUT", "Biometric"),
                (3, "IN", "Manual Punch-View"),
            ],
        )
        self.assertEqual(
            (alice[0]["date"], alice[0]["day"], alice[0]["employeeName"], alice[0]["department"]),
            ("2026-08-05", "Wed", "Alice T", "CUTTING"),
        )
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (cards["Punches"], cards["Employees"], cards["Employee-days"], cards["Days with odd punch count"]),
            (9, 2, 2, 1),
        )
        self.assertEqual((cards["Biometric punches"], cards["Other sources"]), (8, 1))

    def test_summary_over_three_days(self):
        cards = {c["label"]: c["value"] for c in self.log(dateFrom="2026-08-03", dateTo="2026-08-05")["summary"]}
        self.assertEqual(
            (cards["Punches"], cards["Employees"], cards["Employee-days"], cards["Days with odd punch count"]),
            (30, 6, 11, 2),
        )
        self.assertEqual((cards["Biometric punches"], cards["Other sources"]), (25, 5))  # 2 geo + 2 on-duty + 1 manual

    def test_source_filter_keeps_positions_of_the_whole_day(self):
        body = self.log(source="manual")
        self.assertEqual(len(body["rows"]), 1)
        r = body["rows"][0]
        self.assertEqual((r["employeeCode"], r["punchNo"], r["dayPunches"], r["direction"]), ("AC002", 3, 3, "IN"))
        wide = self.log(dateFrom="2026-08-03", dateTo="2026-08-10")
        self.assertEqual(len([1 for r in wide["rows"] if r["sourceLabel"] == "Geo Punch"]), 2)

        def n(src):
            return self.log(dateFrom="2026-08-03", dateTo="2026-08-10", source=src)["rowCount"]

        self.assertEqual((n("geo"), n("on_duty"), n("missing_punch"), n("manual")), (2, 2, 2, 1))
        self.assertEqual(n("biometric"), 26)  # 25 in the first three days + Carol's lone 15:00 punch
        self.assertEqual(
            sum(n(s) for s in ("biometric", "geo", "on_duty", "missing_punch", "manual")),
            self.log(dateFrom="2026-08-03", dateTo="2026-08-10")["rowCount"],
        )

    def test_first_and_last_punch_filters(self):
        first = self.log(dateFrom="2026-08-03", dateTo="2026-08-05", punchPosition="first")
        self.assertEqual(len(first["rows"]), 11)  # one per employee-day
        self.assertEqual({r["punchNo"] for r in first["rows"]}, {1})
        last = self.log(dateFrom="2026-08-03", dateTo="2026-08-05", punchPosition="last")
        self.assertEqual(len(last["rows"]), 11)
        self.assertTrue(all(r["punchNo"] == r["dayPunches"] for r in last["rows"]))

    def test_scope_filters(self):
        wide = dict(dateFrom="2026-08-03", dateTo="2026-08-05")

        def codes(**p):
            return {r["employeeCode"] for r in self.log(**{**wide, **p})["rows"]}

        self.assertEqual(
            codes(), {"AC001", "AC002", "AC003", "AC004", "AC005", "AC006"}
        )  # left employees keep their history
        self.assertEqual(codes(employeeStatus="active"), {"AC001", "AC002", "AC003", "AC004", "AC005"})
        self.assertEqual(codes(employeeStatus="inactive"), {"AC006"})
        self.assertEqual(codes(departmentIds=str(self.sewing.id)), {"AC003", "AC004"})
        self.assertEqual(codes(employmentType="production"), {"AC004"})
        self.assertEqual(codes(branchIds=str(self.b2.id)), {"AC005"})
        self.assertEqual(codes(designationIds=str(self.cutter.id)), {"AC001", "AC002", "AC006"})
        self.assertEqual(codes(employeeIds=str(self.pat.id)), {"AC004"})

    def test_ordering_is_person_wise(self):
        body = self.log(dateFrom="2026-08-03", dateTo="2026-08-05")
        seq = [(r["employeeName"], r["date"], r["punchTime"]) for r in body["rows"]]
        self.assertEqual(seq, sorted(seq))

    def test_row_limit_is_honoured_and_exports_are_refused_not_truncated(self):
        base = registry.get_spec("punch-log")
        small = dataclasses.replace(base, id="zz-small-punch-log", screen_limit=5, pdf_max_rows=3)
        registry.register(small)
        self._registered.append(small.id)
        with CaptureQueriesContext(connection) as q:
            r = self.get(f"/api/reports/run/{small.id}", dateFrom="2026-08-03", dateTo="2026-08-10")
        body = r.json()
        self.assertEqual((body["rowCount"], body["truncated"], body["limit"]), (5, True, 5))
        self.assertIsNone(body["totals"])
        self.assertTrue(any("LIMIT" in x["sql"] for x in q.captured_queries))  # cut in the query, not after loading
        refused = self.export(small.id, "pdf", dateFrom="2026-08-03", dateTo="2026-08-10")
        self.assertEqual(refused.status_code, 413)
        self.assertEqual(refused.json()["error"], "too_many_rows")

    def test_exports(self):
        ws = self.sheet("punch-log", dateFrom="2026-08-05", dateTo="2026-08-05")
        header = [c.value for c in ws[7]]
        self.assertEqual(
            header[:8], ["Emp Code", "Employee", "Department", "Date", "Day", "Punch No", "Time", "In / Out"]
        )
        first = [c.value for c in ws[8]]
        self.assertEqual((first[0], first[5], first[6], first[7]), ("AC001", 1, "09:00:00", "IN"))
        self.assertIn("PUNCH LOG", self.pdf_pages("punch-log", dateFrom="2026-08-05", dateTo="2026-08-05")[0])


# ══ Punch Exceptions ════════════════════════════════════════════════════════════


class PunchExceptionTests(Kit):
    def exc(self, **params):
        return self.run_report("punch-exceptions", **{**PERIOD, **params})

    def kinds(self, body):
        return sorted((r["employeeCode"], r["date"][-2:], r["exceptionType"]) for r in body["rows"])

    def test_golden_exceptions(self):
        body = self.exc()
        self.assertEqual(
            self.kinds(body),
            [
                ("AC001", "04", "Single punch"),
                ("AC001", "05", "Duplicate taps"),
                ("AC001", "05", "Many punches"),
                ("AC002", "05", "Odd punches"),
                ("AC003", "07", "Single punch"),
                ("AC005", "07", "Present without punches"),
            ],
        )
        row = next(r for r in body["rows"] if r["exceptionType"] == "Many punches")
        self.assertEqual(
            (row["punchCount"], row["punches"], row["status"], row["source"]),
            (6, "09:00, 09:02, 13:00, 13:45, 18:00, 18:01", "Present", "Biometric"),
        )
        single = self.row(body, "AC001", "2026-08-04")
        self.assertEqual((single["status"], single["punches"]), ("Half Day", "09:10"))
        self.assertIn("Half Day", single["suggestion"])
        no_punch = self.row(body, "AC005", "2026-08-07")
        self.assertEqual((no_punch["punchCount"], no_punch["punches"], no_punch["source"]), (None, None, "HR Entry"))
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (
                cards["Exceptions"],
                cards["Employees"],
                cards["Single punch"],
                cards["Odd punches"],
                cards["Many punches"],
                cards["Duplicate taps"],
            ),
            (6, 4, 2, 1, 1, 1),
        )
        self.assertEqual((cards["Present without punches"], cards["On leave / holiday"]), (1, 0))
        self.assertTrue(any("2 single-punch day(s)" in n for n in body["notes"]))
        self.assertEqual(body["rowCount"], len(body["rows"]))

    def test_filters(self):
        self.assertEqual(
            self.kinds(self.exc(exceptionType="single_punch")),
            [("AC001", "04", "Single punch"), ("AC003", "07", "Single punch")],
        )
        self.assertEqual(len(self.exc(exceptionType="duplicate_taps")["rows"]), 1)
        self.assertEqual(len(self.exc(exceptionType="present_no_punches")["rows"]), 1)
        self.assertEqual(
            {r["employeeCode"] for r in self.exc(departmentIds=str(self.cutting.id))["rows"]}, {"AC001", "AC002"}
        )
        self.assertEqual({r["employeeCode"] for r in self.exc(employeeIds=str(self.zed.id))["rows"]}, {"AC005"})
        self.assertEqual(self.exc(employmentType="production")["rows"], [])
        self.assertEqual({r["employeeCode"] for r in self.exc(branchIds=str(self.b2.id))["rows"]}, {"AC005"})
        self.assertEqual(self.exc(dateFrom="2026-08-06", dateTo="2026-08-06")["rows"], [])

    def test_punches_on_leave_and_holidays_and_night_exits(self):
        hal = self.add_staff("AC970", "Hal")
        sl = LeaveType.objects.get(code="SL")
        LeaveRequest.objects.create(
            employee=hal,
            leave_type_ref=sl,
            type="sick",
            start_date="2026-08-12",
            end_date="2026-08-12",
            total_days=1,
            status="approved",
        )
        Holiday.objects.create(name="Founders Day", date=d(13))
        _holiday_dates_for_month.cache_clear()
        punch(hal, d(12), "09:00", "18:00")
        punch(hal, d(13), "09:00", "18:00")
        eve = self.add_staff("AC971", "Evie")
        punch(eve, d(14), "09:00", "13:00", "13:45")
        punch(eve, d(15), "01:05", "09:00", "18:00")  # a night exit the engine gives to the 14th
        for emp, first, last in ((hal, d(12), d(13)), (eve, d(14), d(15))):
            seed_records(emp, first, last)
        body = self.run_report(
            "punch-exceptions", dateFrom="2026-08-12", dateTo="2026-08-15", employeeIds=f"{hal.id},{eve.id}"
        )
        self.assertEqual(self.kinds(body), [("AC970", "12", "Punch on leave"), ("AC970", "13", "Punch on holiday")])
        self.assertEqual({r["status"] for r in body["rows"]}, {"Present"})

    def test_today_is_not_judged_for_single_or_odd_punches(self):
        emp = self.add_staff("AC972", "Todd")
        punch(emp, d(12), "09:00")
        seed_records(emp, d(12), d(12))
        with mock.patch("api.reporting.filters.ist_today", return_value=d(12)):
            body = self.run_report(
                "punch-exceptions", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id)
            )
        self.assertEqual(body["rows"], [])
        body = self.run_report("punch-exceptions", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id))
        self.assertEqual([r["exceptionType"] for r in body["rows"]], ["Single punch"])

    def test_exports(self):
        ws = self.sheet("punch-exceptions", **PERIOD)
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:6], ["Emp Code", "Employee", "Department", "Date", "Day", "Exception"])
        self.assertEqual([c.value for c in ws[8]][:2], ["AC001", "Alice T"])
        self.assertIn("PUNCH EXCEPTIONS", self.pdf_pages("punch-exceptions", **PERIOD)[0])


# ══ Unmatched device IDs ════════════════════════════════════════════════════════


class UnmatchedPunchTests(Kit):
    def test_golden_rows_and_ist_times(self):
        body = self.run_report("unmatched-punches")
        rows = {r["deviceUserId"]: r for r in body["rows"]}
        self.assertEqual(set(rows), {"9001", "9002", "AC006"})  # unresolved only by default
        r = rows["9001"]
        self.assertEqual(
            (r["deviceLabel"], r["deviceSerial"], r["punchCount"], r["resolved"]), ("Gate A", "SN-A", 12, "Unresolved")
        )
        self.assertEqual(
            (r["firstSeenAt"], r["lastSeenAt"]), ("2026-08-01 09:30", "2026-08-03 09:30")
        )  # 04:00 UTC = 09:30 IST
        self.assertEqual((r["lastPunchDate"], r["lastPunchTime"]), ("2026-08-03", "09:05"))
        self.assertIn("No employee has this code", r["matchHint"])
        self.assertIsNone(rows["9002"]["deviceSerial"])  # a pull-sync row has no serial
        self.assertIn("exists but is inactive", rows["AC006"]["matchHint"])
        self.assertIn("Leo T", rows["AC006"]["matchHint"])
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (
                cards["Unresolved IDs"],
                cards["Punches discarded (approx.)"],
                cards["Punched in last 7 days"],
                cards["Shown"],
            ),
            (3, 20, 1, 3),
        )
        self.assertEqual(body["totals"]["punchCount"], 20)
        self.assertEqual([r["deviceUserId"] for r in body["rows"]], ["9002", "9001", "AC006"])  # last seen first

    def test_filters(self):
        def ids(**p):
            return {r["deviceUserId"] for r in self.run_report("unmatched-punches", **p)["rows"]}

        self.assertEqual(ids(resolved="resolved"), {"9003"})
        self.assertEqual(ids(resolved="all"), {"9001", "9002", "9003", "AC006"})
        self.assertEqual(ids(seen="7"), {"9002"})
        self.assertEqual(ids(seen="30"), {"9002"})
        self.assertEqual(ids(resolved="all", device="gate b"), {"9003"})
        self.assertEqual(ids(device="SN-A"), {"9001", "AC006"})
        self.assertEqual(ids(deviceUserId="900"), {"9001", "9002"})
        self.assertEqual(ids(resolved="resolved", device="nothing"), set())

    def test_company_wide_data_is_hidden_from_branch_restricted_users(self):
        body = self.run_report("unmatched-punches", self.branch_user)
        self.assertEqual(body["rows"], [])
        self.assertTrue(any("company-wide" in n for n in body["notes"]))

    def test_exports(self):
        ws = self.sheet("unmatched-punches")
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:4], ["Device User ID", "Device", "Device Serial", "Punches Discarded"])
        self.assertEqual({ws.cell(row=r, column=1).value for r in (8, 9, 10)}, {"9001", "9002", "AC006"})
        self.assertIn("UNMATCHED DEVICE IDS", self.pdf_pages("unmatched-punches")[0])


# ══ Manual overrides ════════════════════════════════════════════════════════════


class ManualOverrideTests(Kit):
    RANGE = {"dateFrom": "2026-08-01", "dateTo": "2026-08-31"}

    def rows(self, user=None, **params):
        body = self.run_report("manual-overrides", user, **{**self.RANGE, **params})
        return body, [(r["employeeCode"], r["date"][-2:], r["origin"], r["status"]) for r in body["rows"]]

    def test_golden_rows_newest_attendance_date_first(self):
        body, rows = self.rows()
        self.assertEqual(
            rows,
            [
                ("AC001", "10", "HR override request", "Approved"),  # the frozen day it produced is not listed twice
                ("AC003", "10", "Casual leave (approved)", "Applied"),
                ("AC004", "10", "Compensation redemption", "Applied"),
                ("AC005", "09", "HR override request", "Pending"),
                ("AC005", "08", "HR override request", "Rejected"),
                ("AC002", "07", "HR override request", "Pending"),
            ],
        )
        a = body["rows"][0]
        self.assertEqual(a["before"], "Absent; 0.00 shift")
        self.assertEqual(a["after"], "Present; 09:00 to 18:00; 1.00 shift")
        self.assertEqual(
            (a["requestedBy"], a["reviewedBy"], a["reviewComment"]), ("HR Anita", "HOD Ravi", "Verified with gate")
        )
        self.assertEqual((a["createdAt"], a["reviewedAt"]), ("2026-08-10 09:00", "2026-08-11 10:30"))  # IST
        self.assertEqual(a["reason"], "Missed punch due to power cut")
        cl = body["rows"][1]
        self.assertEqual(
            (cl["reason"], cl["reviewedBy"], cl["before"], cl["createdAt"]),
            ("Casual Leave (paid) -approved", "HOD Meena", None, "2026-08-10 17:30"),
        )
        half = body["rows"][5]
        self.assertEqual(half["after"], "Half Day; 09:00 to -; 0.50 shift")
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(
            (cards["Requests"], cards["Pending"], cards["Approved"], cards["Rejected"], cards["Days frozen as manual"]),
            (4, 2, 1, 1, 3),
        )

    def test_filters(self):
        self.assertEqual({r[0] for r in self.rows(requestStatus="pending")[1]}, {"AC005", "AC002"})
        self.assertEqual([r[3] for r in self.rows(requestStatus="pending")[1]], ["Pending", "Pending"])
        approved = self.rows(requestStatus="approved")[1]
        self.assertEqual(
            [(r[0], r[3]) for r in approved], [("AC001", "Approved"), ("AC003", "Applied"), ("AC004", "Applied")]
        )
        self.assertEqual([r[3] for r in self.rows(requestStatus="rejected")[1]], ["Rejected"])
        self.assertEqual([r[0] for r in self.rows(origin="casual_leave")[1]], ["AC003"])
        self.assertEqual([r[0] for r in self.rows(origin="compensation_redemption")[1]], ["AC004"])
        self.assertEqual(len(self.rows(origin="hr_override")[1]), 4)
        self.assertEqual([r[0] for r in self.rows(requestedBy="kumar")[1]], ["AC005", "AC005"])
        self.assertEqual({r[0] for r in self.rows(employeeIds=str(self.zed.id))[1]}, {"AC005"})
        self.assertEqual({r[0] for r in self.rows(departmentIds=str(self.sewing.id))[1]}, {"AC003", "AC004"})
        self.assertEqual(
            [r[1] for r in self.rows(dateFrom="2026-08-09", dateTo="2026-08-10")[1]], ["10", "10", "10", "09"]
        )
        self.assertEqual(self.rows(dateFrom="2020-01-01", dateTo="2020-01-31")[1], [])

    def test_branch_scoped_user_sees_only_their_branch(self):
        _body, rows = self.rows(self.branch_user)
        self.assertEqual({r[0] for r in rows}, {"AC001", "AC002", "AC003", "AC004"})
        self.assertEqual(self.rows(self.branch_user, branchIds=str(self.b2.id))[1], [])

    def test_exports(self):
        ws = self.sheet("manual-overrides", **self.RANGE)
        header = [c.value for c in ws[7]]
        self.assertEqual(header[:5], ["Emp Code", "Employee", "Department", "Attendance Date", "Origin"])
        self.assertEqual([c.value for c in ws[8]][:2], ["AC001", "Alice T"])
        self.assertIn("MANUAL ATTENDANCE OVERRIDES", self.pdf_pages("manual-overrides", **self.RANGE)[0])


# ══ Custom exports: the coloured muster roll and the one-page time card ═════════


class ExportLayoutTests(Kit):
    MONTH = {"dateFrom": "2026-08-01", "dateTo": "2026-08-31"}

    @staticmethod
    def near(colour, hex_):
        want = tuple(int(hex_[i : i + 2], 16) / 255 for i in (0, 2, 4))
        return colour is not None and len(colour) == 3 and all(abs(a - b) < 0.01 for a, b in zip(colour, want))

    def test_a_full_month_muster_roll_exports_to_pdf_on_one_page_in_the_report_log_colours(self):
        r = self.export("attendance-muster-sheet", "pdf", **self.MONTH)
        self.assertEqual(r.status_code, 200, r.content[:300])
        with pdfplumber.open(io.BytesIO(r.content)) as pdf:
            self.assertEqual(len(pdf.pages), 1)
            page = pdf.pages[0]
            text = page.extract_text() or ""
            for word in ("ATTENDANCE SHEET (MUSTER ROLL)", "Alice T", "Strength", "TOTAL", "Half day - morning worked"):
                self.assertIn(word, text)
            fills = [rect.get("non_stroking_color") for rect in page.rects]
            for name, hex_ in (
                ("present", "C6EFCE"),
                ("half morning", "FFEB9C"),
                ("half evening", "FED7AA"),
                ("absent", "FFC7CE"),
                ("holiday", "E2E8F0"),
                ("Sunday header", "FFE4E6"),
            ):
                self.assertTrue(any(self.near(f, hex_) for f in fills), f"no {name} cell ({hex_}) in the PDF")

    def test_a_week_muster_roll_pdf_uses_the_smaller_page(self):
        r = self.export("attendance-muster-sheet", "pdf", **PERIOD)
        with pdfplumber.open(io.BytesIO(r.content)) as pdf:
            self.assertEqual(len(pdf.pages), 1)
            w, h = pdf.pages[0].width, pdf.pages[0].height
        self.assertAlmostEqual(w, 841.9, delta=1)  # A4 landscape
        self.assertAlmostEqual(h, 595.3, delta=1)

    def test_a_single_day_muster_roll_exports(self):
        one = {"dateFrom": "2026-08-03", "dateTo": "2026-08-03"}
        pages = self.pdf_pages("attendance-muster-sheet", **one)
        self.assertEqual(len(pages), 1)
        self.assertIn("Alice T", pages[0])
        ws = self.sheet("attendance-muster-sheet", **one)
        self.assertEqual([c.value for c in ws[7]][4:6], ["3 Mo", "P"])
        self.assertEqual([c.value for c in ws[8]][4], "P")

    def test_muster_excel_cells_carry_the_report_log_colours(self):
        ws = self.sheet("attendance-muster-sheet", **PERIOD)
        first = list(ws[8])  # Alice: P, half morning, P, holiday, A, A, holiday, P
        self.assertEqual([c.value for c in first[4:12]], ["P", "½M", "P", "H", "A", "A", "H", "P"])
        self.assertEqual(
            [c.fill.fgColor.rgb for c in first[4:12]],
            ["00" + h for h in ("C6EFCE", "FFEB9C", "C6EFCE", "E2E8F0", "FFC7CE", "FFC7CE", "E2E8F0", "C6EFCE")],
        )
        carol = [c.fill.fgColor.rgb for c in ws[10]][4:12]  # ... half evening on Friday the 7th
        self.assertEqual(carol[4], "00FED7AA")
        header = list(ws[7])
        self.assertEqual(header[4 + 6].value, "9 Su")
        self.assertEqual(header[4 + 6].fill.fgColor.rgb, "00FFE4E6")  # the rose Sunday header
        self.assertEqual(header[4].fill.fgColor.rgb, "001E3A5F")
        self.assertEqual(ws.column_dimensions["E"].width, 5.5)
        strength = [c.value for c in ws[12]]
        self.assertEqual(strength[2], "Strength")
        self.assertEqual(strength[4:12], [3, 3, 2, 0, 2, 0, 0, 3])

    def test_muster_excel_keeps_text_that_looks_like_a_formula_as_text(self):
        evil = self.add_staff("AC992", "=SUM(A1:A9)")
        ws = self.sheet("attendance-muster-sheet", **PERIOD, employeeIds=str(evil.id))
        cell = ws.cell(row=8, column=3)
        self.assertEqual(cell.value, "=SUM(A1:A9) Q")
        self.assertEqual(cell.data_type, "s")

    def test_a_month_time_card_fits_one_page_per_employee(self):
        rng = random.Random(11)
        emps = []
        for i in range(3):
            emp = self.add_staff(f"AC97{i}", f"Month{i}")
            emps.append(emp)
            for day in range(1, 32):
                if date(2026, 8, day).weekday() == 6 or rng.random() < 0.12:
                    continue
                if rng.random() < 0.15:
                    punch(emp, d(day), "09:%02d" % rng.choice([0, 20, 40]))  # a lone punch: a remark on the row
                else:
                    punch(emp, d(day), "09:%02d" % rng.choice([0, 5, 40]), "18:%02d" % rng.choice([0, 30]))
            seed_records(emp, d(1), d(31))
        pages = self.pdf_pages("time-card", **self.MONTH, employeeIds=",".join(str(e.id) for e in emps))
        self.assertEqual(len(pages), 3)
        for page, emp in zip(pages, emps):
            self.assertIn(emp.employee_code, page)
            for word in ("TIME CARD", "01-Aug-26", "31-Aug-26", "Worked hrs", "Employee signature", "HR"):
                self.assertIn(word, page)
            self.assertEqual(sum(1 for e in emps if e.employee_code in page), 1)  # one employee per page

    def test_time_card_pdf_of_a_range_longer_than_a_month_flows_onto_more_pages(self):
        pages = self.pdf_pages("time-card", dateFrom="2026-06-01", dateTo="2026-08-31", employeeIds=str(self.alice.id))
        self.assertGreaterEqual(len(pages), 2)
        self.assertIn("Employee signature", pages[-1])
        self.assertIn("Alice T", pages[0])
        self.assertNotIn("Employee signature", pages[0])


# ══ Edge cases: missing master data, production employees ═══════════════════════


class EdgeCaseTests(Kit):
    def test_employee_without_department_designation_or_branch(self):
        nora = Employee.objects.create(employee_code="AC980", first_name="Nora", last_name="N", employment_type="staff")
        punch(nora, d(3), "09:00")  # a lone punch: half day, and a punch exception
        punch(nora, d(4), "09:00", "18:00")
        seed_records(nora, d(3), d(4))
        one = {"employeeIds": str(nora.id)}

        body = self.run_report("time-card", **PERIOD, **one)
        r = self.row(body, "AC980", "2026-08-03")
        self.assertEqual(
            (r["department"], r["shift"], r["status"], r["lateMinutes"]), ("Unassigned", None, "Half Day", None)
        )
        self.assertIn("No shift assigned", r["remarks"])
        self.assertEqual(
            self.row(body, "AC980", "2026-08-04")["workedHours"], 9.0
        )  # no shift, so no lunch break to deduct

        body = self.run_report("daily-attendance", dateFrom="2026-08-04", dateTo="2026-08-04", **one)
        r = self.row(body, "AC980")
        self.assertEqual(
            (r["department"], r["designation"], r["shift"], r["status"]), ("Unassigned", None, None, "Present")
        )

        body = self.run_report("attendance-muster-sheet", **PERIOD, **one)
        r = self.row(body, "AC980")
        self.assertEqual((r["department"], r["d20260803"], r["d20260804"]), (None, "½M", "P"))

        body = self.run_report("attendance-summary-monthly", period="2026-08", **one)
        r = self.row(body, "AC980")
        self.assertEqual(
            (r["department"], r["designation"], r["presentDays"], r["halfDays"]), ("Unassigned", None, 1, 1)
        )
        self.assertEqual([x["employeeName"] for x in body["rows"] if x.get("_kind")], ["Unassigned total"])

        body = self.run_report("punch-log", dateFrom="2026-08-03", dateTo="2026-08-04", **one)
        self.assertEqual({r["department"] for r in body["rows"]}, {"Unassigned"})
        body = self.run_report("punch-exceptions", **PERIOD, **one)
        self.assertEqual([r["exceptionType"] for r in body["rows"]], ["Single punch"])

        # the unassigned people sort after every department; a branch-restricted user never sees a branchless person
        everyone = self.run_report("attendance-summary-monthly", period="2026-08")
        subs = [x["employeeName"] for x in everyone["rows"] if x.get("_kind")]
        self.assertEqual(subs[-1], "Unassigned total")
        for rid, params in (
            ("time-card", PERIOD),
            ("daily-attendance", PERIOD),
            ("attendance-muster-sheet", PERIOD),
            ("attendance-summary-monthly", {"period": "2026-08"}),
            ("punch-log", PERIOD),
            ("punch-exceptions", PERIOD),
        ):
            mine = self.run_report(rid, self.branch_user, **params)
            self.assertNotIn("AC980", [r.get("employeeCode") for r in mine["rows"]], rid)
            self.assertIn("AC980", [r.get("employeeCode") for r in self.run_report(rid, **params)["rows"]], rid)

    def test_production_employees_ignore_leave_requests(self):
        sl = LeaveType.objects.get(code="SL")
        LeaveRequest.objects.create(
            employee=self.pat,
            leave_type_ref=sl,
            type="sick",
            start_date="2026-08-04",
            end_date="2026-08-04",
            total_days=1,
            status="approved",
        )  # Pat punched that day; the engine gives production no leave, so this request is neither status nor exception
        body = self.run_report("time-card", **PERIOD, employeeIds=str(self.pat.id))
        r = self.row(body, "AC004", "2026-08-04")
        self.assertEqual(r["status"], "Present")
        self.assertNotIn("leave", (r["remarks"] or "").lower())
        exc = self.run_report("punch-exceptions", **PERIOD, employeeIds=str(self.pat.id))
        self.assertEqual(exc["rows"], [])
        daily = self.run_report(
            "daily-attendance", dateFrom="2026-08-04", dateTo="2026-08-04", employeeIds=str(self.pat.id)
        )
        self.assertEqual((self.row(daily, "AC004")["status"], self.row(daily, "AC004")["leaveType"]), ("Present", None))

    def test_strict_mode_days_read_the_same_and_the_reports_write_nothing(self):
        ps = PayrollSettings.get()
        ps.attendance_mode = "strict"
        ps.save()
        sam = self.add_staff("AC995", "Sam")
        punch(sam, d(12), "09:40", "13:00", "13:50", "18:05")
        seed_records(sam, d(12), d(12))  # a strict-mode compute writes DailyShiftLog: set-up only
        self.assertEqual(AttendanceDayRecord.objects.get(employee=sam, date=d(12)).computed_mode, "strict")
        before = (AttendanceDayRecord.objects.count(), DailyShiftLog.objects.count())
        body = self.run_report("time-card", dateFrom="2026-08-12", dateTo="2026-08-13", employeeIds=str(sam.id))
        r = self.row(body, "AC995", "2026-08-12")
        self.assertEqual((r["status"], r["lateMinutes"], r["workedHours"], r["punchCount"]), ("Present", 40, 7.58, 4))
        self.assertIsNone(self.row(body, "AC995", "2026-08-13")["status"])  # unprocessed: still not computed
        self.assertEqual(
            self.row(
                self.run_report(
                    "daily-attendance", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(sam.id)
                ),
                "AC995",
            )["status"],
            "Present",
        )
        self.export("time-card", "pdf", dateFrom="2026-08-12", dateTo="2026-08-13", employeeIds=str(sam.id))
        self.assertEqual((AttendanceDayRecord.objects.count(), DailyShiftLog.objects.count()), before)

    def test_overtime_note_follows_the_settings_switch(self):
        body = self.run_report("time-card", **PERIOD, employeeIds=str(self.bob.id))
        self.assertTrue(any("Overtime detection is switched off" in n for n in body["notes"]))
        ps = PayrollSettings.get()
        ps.ot_detection_enabled = True
        ps.compensation_feature_enabled = True
        ps.save()
        body = self.run_report("time-card", **PERIOD, employeeIds=str(self.bob.id))
        self.assertFalse(any("Overtime detection is switched off" in n for n in body["notes"]))


# ══ Cross-cutting: access, branch isolation, no writes, empty, query counts ═════


class CrossCuttingTests(Kit):
    ALL_PARAMS = {**PERIOD, "period": "2026-08"}

    def test_reports_role_without_the_attendance_module_gets_report_forbidden(self):
        for rid in IDS:
            for verb, extra in (("run", {}), ("export", {"fmt": "xlsx"}), ("export", {"fmt": "pdf"})):
                r = self.get(f"/api/reports/{verb}/{rid}", self.no_module, **self.ALL_PARAMS, **extra)
                self.assertEqual(r.status_code, 403, (rid, verb))
                self.assertEqual(r.json()["error"], "report_forbidden", (rid, verb))
        listed = {x["id"] for x in self.get("/api/reports/catalog", self.no_module).json()["reports"]}
        self.assertFalse(listed & set(IDS))
        listed = {x["id"] for x in self.get("/api/reports/catalog", self.branch_user).json()["reports"]}
        self.assertTrue(set(IDS) <= listed)
        self.assertEqual(self.get("/api/reports/run/time-card", self.branch_user, **PERIOD).status_code, 200)

    def test_modules_are_real_permission_keys(self):
        from .permission_registry import all_module_keys

        for rid in IDS:
            for m in registry.get_spec(rid).modules:
                self.assertIn(m, all_module_keys())

    def test_branch_scoped_user_sees_only_their_branch_and_cannot_widen_it(self):
        for rid in IDS[:5] + ["punch-log", "punch-exceptions", "manual-overrides"]:
            base = {**self.ALL_PARAMS, "dateTo": "2026-08-10", "employeeStatus": "all", "staffType": "all"}
            everyone = self.run_report(rid, **base)
            self.assertIn(
                "AC005", {r.get("employeeCode") for r in everyone["rows"]}, rid
            )  # Zed (Unit 2) shows for the admin
            mine = self.run_report(rid, self.branch_user, **base)
            codes = {r.get("employeeCode") for r in mine["rows"] if r.get("employeeCode")}
            self.assertTrue(codes, rid)
            self.assertNotIn("AC005", codes, rid)
            other = self.run_report(rid, self.branch_user, **base, branchIds=str(self.b2.id))
            self.assertEqual([r for r in other["rows"] if r.get("employeeCode")], [], rid)
            zed = self.run_report(rid, self.branch_user, **base, employeeIds=str(self.zed.id))
            self.assertEqual([r for r in zed["rows"] if r.get("employeeCode")], [], rid)

    def test_reading_a_report_never_writes(self):
        counts = lambda: (  # noqa: E731
            AttendanceDayRecord.objects.count(),
            DailyShiftLog.objects.count(),
            AttendanceLog.objects.count(),
            Attendance.objects.count(),
            AttendanceOverrideRequest.objects.count(),
            UnmatchedPunch.objects.count(),
        )
        before = counts()
        # Aug 11 and 12 have no records: a computing report would create them
        params = {"dateFrom": "2026-08-03", "dateTo": "2026-08-12", "period": "2026-08", "staffType": "all"}
        for rid in IDS:
            self.run_report(rid, **params)
            self.assertEqual(self.export(rid, "xlsx", **params).status_code, 200, rid)
            self.assertEqual(self.export(rid, "pdf", **params).status_code, 200, rid)
        self.assertEqual(counts(), before)

    def test_empty_results_work_on_screen_and_in_both_exports(self):
        nobody = {"employeeIds": "9999999"}
        for rid in IDS[:7]:
            body = self.run_report(rid, **self.ALL_PARAMS, **nobody)
            self.assertEqual([r for r in body["rows"] if r.get("employeeCode")], [], rid)
            for fmt, sig in (("xlsx", b"PK"), ("pdf", b"%PDF")):
                r = self.export(rid, fmt, **self.ALL_PARAMS, **nobody)
                self.assertEqual(r.status_code, 200, (rid, fmt))
                self.assertTrue(r.content.startswith(sig), (rid, fmt))
        empties = {
            "unmatched-punches": {"resolved": "resolved", "device": "no such device"},
            "manual-overrides": {"dateFrom": "2020-01-01", "dateTo": "2020-01-31"},
        }
        for rid, extra in empties.items():
            self.assertEqual(self.run_report(rid, **extra)["rows"], [], rid)
            for fmt in ("xlsx", "pdf"):
                self.assertEqual(self.export(rid, fmt, **extra).status_code, 200, (rid, fmt))
        ws = self.sheet("time-card", **PERIOD, **nobody)
        self.assertEqual(ws.cell(row=7, column=1).value, "Emp Code")

    def test_every_export_is_a_valid_file_with_the_spec_headers(self):
        for rid in IDS:
            spec = registry.get_spec(rid)
            r = self.export(rid, "xlsx", **self.ALL_PARAMS)
            ws = load_workbook(io.BytesIO(r.content)).active
            self.assertEqual(ws.cell(row=3, column=1).value, spec.title.upper(), rid)
            if spec.columns:
                self.assertEqual(ws.cell(row=7, column=1).value, spec.columns[0].label, rid)
            self.assertIn("attachment", r["Content-Disposition"])
            pdf = self.export(rid, "pdf", **self.ALL_PARAMS)
            self.assertTrue(pdf.content.startswith(b"%PDF"), rid)
            self.assertIn("application/pdf", pdf["Content-Type"])

    def test_query_counts_do_not_grow_with_the_number_of_employees(self):
        dept = Department.objects.create(name="QDEPT", branch=self.b1)

        def grow(prefix, count):
            for i in range(count):
                emp = self.add_staff(f"{prefix}{i:02d}", f"Q{prefix}{i}", dept=dept)
                punch(emp, d(17), "09:00", "13:00", "13:45", "18:00")
                punch(emp, d(18), "09:30", "18:30")
                punch(emp, d(19), "09:00")
                seed_records(emp, d(17), d(19))
                EmployeePermission.objects.create(
                    employee=emp, date=d(18), status="approved", type="morning_late_in", duration_minutes=60
                )

        grow("QA", 3)
        params = {
            "departmentIds": str(dept.id),
            "dateFrom": "2026-08-17",
            "dateTo": "2026-08-19",
            "period": "2026-08",
            "staffType": "all",
        }
        ids = [
            "time-card",
            "daily-attendance",
            "absentee-list-daily",
            "attendance-muster-sheet",
            "attendance-summary-monthly",
            "punch-log",
            "punch-exceptions",
            "manual-overrides",
        ]
        for rid in ids:
            self.count_queries(rid, **params)  # warm up
        small = {rid: self.count_queries(rid, **params) for rid in ids}
        grow("QB", 12)
        for rid in ids:
            self.assertLessEqual(
                abs(self.count_queries(rid, **params) - small[rid]), 1, f"{rid}: query count grew with employees"
            )
        # a report over the 15 employees really does contain them
        self.assertGreaterEqual(len(self.run_report("time-card", **params)["rows"]), 15 * 3)

    def test_unmatched_and_override_query_counts_are_constant(self):
        runs = (("unmatched-punches", {"resolved": "all"}), ("manual-overrides", ManualOverrideTests.RANGE))
        small = {}
        for rid, params in runs:
            self.count_queries(rid, **params)  # warm up
            small[rid] = self.count_queries(rid, **params)
        for i in range(12):
            UnmatchedPunch.objects.create(device_user_id=f"77{i:02d}", punch_count=1, last_punch_date=d(4))
            AttendanceOverrideRequest.objects.create(
                employee=self.alice,
                date=d(11 + i % 5),
                status="pending",
                requested_by="HR Anita",
                previous_values={"status": "absent"},
                requested_values={"status": "present"},
            )
        for rid, params in runs:
            self.assertLessEqual(abs(self.count_queries(rid, **params) - small[rid]), 1, rid)
            self.assertGreater(len(self.run_report(rid, **params)["rows"]), 10, rid)


# ══ Adversarial review: each test states the behaviour a payroll / HR manager relies on ═══════


class AdversarialReviewTests(Kit):
    def test_time_card_never_silently_drops_employees_when_a_filter_shrinks_the_rows(self):
        # The employee list is cut to what the row limit could hold BEFORE the day filters remove rows, so a
        # roster bigger than the limit loses its tail without any "truncated" notice.
        with mock.patch("api.reporting.runner.limit_for", return_value=30):
            body = self.run_report("time-card", **PERIOD, dayStatus="absent")
        codes = {r["employeeCode"] for r in self.data_rows(body)}
        # Carol (AC003) and Pat (AC004) are absent on several days of the period
        missing = {"AC003", "AC004"} - codes
        self.assertTrue(body["truncated"] or not missing, f"employees {sorted(missing)} silently left out")

    def test_time_card_dormant_leavers_do_not_use_up_the_employee_budget(self):
        # Default filters (Employee status = All): people who left long ago and have no activity are dropped AFTER the
        # employee list was cut to the row budget, so the active tail of a big roster is lost without a notice.
        for i in range(1, 6):
            self.add_staff(f"AC000{i}", f"Gone{i}", status="inactive")  # CUTTING, sorts before AC001
        with mock.patch("api.reporting.runner.limit_for", return_value=30):
            body = self.run_report("time-card", **PERIOD)
        codes = {r["employeeCode"] for r in self.data_rows(body)}
        active = {"AC001", "AC002", "AC003", "AC004", "AC005"}
        self.assertTrue(body["truncated"] or active <= codes, f"active employees missing: {sorted(active - codes)}")

    def test_a_lone_punch_after_shift_end_is_not_overtime(self):
        # One punch at 19:30: the engine stores no last punch (single punch), so it can never detect overtime;
        # the time card must not invent 90 OT minutes from a punch that may just as well be the arrival.
        emp = self.add_staff("AC800", "Lone")
        punch(emp, d(12), "19:30")
        seed_records(emp, d(12), d(12))
        body = self.run_report("time-card", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id))
        r = self.row(body, "AC800", "2026-08-12")
        self.assertIsNone(r["out1"])
        self.assertIsNone(r["otMinutes"])

    def _leave_type_cl_employee(self):
        cl, _ = LeaveType.objects.get_or_create(code="CL", defaults={"name": "Casual Leave"})
        emp = self.add_staff("AC801", "Cel")
        LeaveRequest.objects.create(
            employee=emp,
            leave_type_ref=cl,
            type="casual",
            start_date="2026-08-12",
            end_date="2026-08-12",
            total_days=1,
            status="approved",
        )
        seed_records(emp, d(12), d(12))
        self.assertEqual(AttendanceDayRecord.objects.get(employee=emp, date=d(12)).status, "on_leave")
        return emp

    def test_muster_leave_type_cl_is_not_confused_with_paid_casual_leave(self):
        # LeaveType "CL" (Casual Leave) is a seeded default: an approved one is status on_leave (L column, unpaid),
        # while a paid CasualLeaveRequest day is a PRESENT day. Both must not share the cell code "CL".
        self._leave_type_cl_employee()
        body = self.run_report(
            "attendance-muster-sheet",
            dateFrom="2026-08-10",
            dateTo="2026-08-12",
            showLeaveCodes="true",
            staffType="all",
        )
        rows = {r["employeeCode"]: r for r in body["rows"] if not r.get("_kind")}
        paid = rows["AC003"]["d20260810"]  # Carol's approved paid casual leave
        on_leave = rows["AC801"]["d20260812"]  # an approved full-day leave of type CL
        self.assertEqual(paid, "CL")
        self.assertNotEqual(on_leave, paid)

    def test_muster_excel_does_not_paint_a_leave_day_in_the_present_colour(self):
        emp = self._leave_type_cl_employee()
        ws = self.sheet(
            "attendance-muster-sheet",
            dateFrom="2026-08-12",
            dateTo="2026-08-12",
            showLeaveCodes="true",
            employeeIds=str(emp.id),
        )
        cell = ws.cell(row=8, column=5)
        self.assertEqual(ws.cell(row=8, column=2).value, "AC801")
        self.assertNotEqual(cell.fill.fgColor.rgb, "00C6EFCE", f"leave day {cell.value!r} painted as present")

    def test_reports_say_when_late_or_early_out_detection_is_switched_off(self):
        ps = PayrollSettings.get()
        ps.morning_late_in_enabled = False
        ps.evening_early_out_enabled = False
        ps.save()
        for rid, params in (
            ("attendance-summary-monthly", {"period": "2026-08"}),
            ("daily-attendance", PERIOD),
            ("attendance-muster-sheet", PERIOD),
        ):
            with self.subTest(report=rid):
                body = self.run_report(rid, **params)
                text = " ".join(body["notes"]).lower()
                self.assertTrue(
                    "switched off" in text or "disabled" in text or "not enabled" in text,
                    f"{rid}: a Late / Early-Out column of zeros with no note that the detection is off",
                )

    def test_a_day_with_punches_is_not_shown_as_a_plain_absence_without_a_stale_warning(self):
        # The verdict was stored at 09:00, before the device synced this employee's punches. Both facts are on the
        # row (status Absent, In 09:00 / Out 18:00, 8 worked hours) and nothing tells the reader it is stale.
        emp = self.add_staff("AC804", "Late")
        seed_records(emp, d(12), d(12))
        self.assertEqual(AttendanceDayRecord.objects.get(employee=emp, date=d(12)).status, "absent")
        punch(emp, d(12), "09:00", "18:00")
        body = self.run_report("time-card", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id))
        r = self.row(body, "AC804", "2026-08-12")
        self.assertEqual((r["in1"], r["out1"], r["workedHours"]), ("09:00", "18:00", 8.0))
        remarks = (r["remarks"] or "").lower()
        warned = any(w in remarks for w in ("stale", "out of date", "punches after", "recompute", "not up to date"))
        self.assertTrue(r["status"] != "Absent" or warned, f"Absent with two punches and no warning: {r['remarks']!r}")
        daily = self.run_report("daily-attendance", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id))
        d_row = self.row(daily, "AC804")
        flags = (d_row["flags"] or "").lower()
        self.assertTrue(
            d_row["status"] != "Absent" or any(w in flags for w in ("stale", "out of date", "recompute")),
            f"daily register: Absent with punches and no warning: {d_row['flags']!r}",
        )

    def test_absentee_list_does_not_call_someone_absent_whose_leave_was_approved_after_the_day_was_stored(self):
        emp = self.add_staff("AC805", "Leaver")
        seed_records(emp, d(12), d(12))  # stored as absent
        sl = LeaveType.objects.get(code="SL")
        LeaveRequest.objects.create(
            employee=emp,
            leave_type_ref=sl,
            type="sick",
            start_date="2026-08-12",
            end_date="2026-08-12",
            total_days=1,
            status="approved",
        )
        body = self.run_report("absentee-list-daily", dateFrom="2026-08-12", dateTo="2026-08-12")
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertTrue(
            "AC805" not in rows or rows["AC805"]["status"] != "Absent",
            "a person on approved leave is listed as plain Absent for the morning call",
        )

    def test_absentee_list_saturday_off_saturday_with_only_a_night_exit_is_not_an_absence(self):
        ps = PayrollSettings.get()
        ps.last_punch_post_shift_grace_hours = 9
        ps.first_punch_pre_shift_buffer_hours = 2
        ps.save()
        sat = self.add_staff("AC802", "Satoff")
        EmployeeShiftAssignment.objects.filter(employee=sat).update(saturday_off=True)
        punch(sat, d(14), "09:00", "13:00", "13:45")  # Friday, exit forgotten ...
        punch(sat, d(15), "01:05")  # ... made after midnight: stamped Saturday, but it is Friday's exit
        seed_records(sat, d(14), d(15))
        tc = self.run_report("time-card", dateFrom="2026-08-15", dateTo="2026-08-15", employeeIds=str(sat.id))
        self.assertEqual(self.row(tc, "AC802")["status"], "Weekly Off")
        body = self.run_report("absentee-list-daily", dateFrom="2026-08-15", dateTo="2026-08-15")
        self.assertNotIn("AC802", self.codes(body))

    def test_a_filter_echo_never_names_an_employee_outside_the_users_branch(self):
        body = self.run_report("time-card", self.branch_user, **PERIOD, employeeIds=str(self.zed.id))
        echo = " ".join(f"{f['label']}: {f['value']}" for f in body["filters"])
        self.assertEqual(self.data_rows(body), [])
        self.assertNotIn("AC005", echo)
        self.assertNotIn("Zed", echo)

    def test_summary_cards_agree_with_the_totals_row(self):
        emp = self.add_staff("AC803", "Round")
        for day in (12, 13, 14):
            punch(emp, d(day), "09:00", "16:35")  # 455 min = 7.583 h each; rows show 7.58
        seed_records(emp, d(12), d(14))
        body = self.run_report(
            "time-card", dateFrom="2026-08-12", dateTo="2026-08-14", employeeIds=str(emp.id), deductLunch="false"
        )
        cards = {c["label"]: c["value"] for c in body["summary"]}
        self.assertEqual(cards["Worked hours"], body["totals"]["workedHours"])


# ══ Fixes that came out of the adversarial review ═══════════════════════════════


class ReviewFixTests(Kit):
    def test_time_card_reads_on_past_the_first_block_and_flags_truncation_only_when_rows_are_cut(self):
        with mock.patch("api.reporting.runner.limit_for", return_value=30):
            body = self.run_report("time-card", **PERIOD)  # 6 employees x (8 days + total) = 54 rows > 30
            self.assertTrue(body["truncated"])
            self.assertEqual(
                len(self.data_rows(body)), 30
            )  # the limit counts data rows; each employee's total line rides along
            self.assertIsNone(body["totals"])
        with mock.patch("api.reporting.runner.limit_for", return_value=100):
            body = self.run_report("time-card", **PERIOD)
            self.assertFalse(body["truncated"])
            self.assertEqual(len(self.data_rows(body)), 48)
        # the same 30-row limit with a filter that leaves few rows: every employee with a matching day is listed
        with mock.patch("api.reporting.runner.limit_for", return_value=30):
            body = self.run_report("time-card", **PERIOD, dayStatus="half")
        self.assertEqual({r["employeeCode"] for r in self.data_rows(body)}, {"AC001", "AC003"})

    def test_a_fresh_day_carries_no_out_of_date_flag(self):
        body = self.run_report("time-card", **PERIOD, employeeIds=str(self.alice.id))
        self.assertTrue(all("out of date" not in (r["remarks"] or "").lower() for r in self.data_rows(body)))
        daily = self.run_report("daily-attendance", **PERIOD)
        self.assertTrue(all("out of date" not in (r["flags"] or "").lower() for r in self.data_rows(daily)))

    def test_leave_approved_after_the_day_was_stored_is_flagged_and_listed_as_leave(self):
        emp = self.add_staff("AC806", "Later")
        seed_records(emp, d(12), d(12))
        LeaveRequest.objects.create(
            employee=emp,
            leave_type_ref=LeaveType.objects.get(code="SL"),
            type="sick",
            start_date="2026-08-12",
            end_date="2026-08-12",
            total_days=1,
            status="approved",
        )
        daily = self.run_report("daily-attendance", dateFrom="2026-08-12", dateTo="2026-08-12", employeeIds=str(emp.id))
        row = self.row(daily, "AC806")
        self.assertEqual(row["status"], "Absent")  # the stored verdict is shown as it is ...
        self.assertIn("leave approved after", row["flags"])  # ... with the reason it cannot be trusted
        one = {"dateFrom": "2026-08-12", "dateTo": "2026-08-12", "employeeIds": str(emp.id)}
        self.assertEqual(self.codes(self.run_report("absentee-list-daily", **one)), [])
        listed = self.run_report("absentee-list-daily", **one, show="leave")
        self.assertEqual([(r["employeeCode"], r["status"]) for r in listed["rows"]], [("AC806", "Leave")])
        self.assertTrue(any("approved since" in n for n in listed["notes"]))

    def test_leave_type_code_that_clashes_with_a_sheet_code_is_prefixed(self):
        cl, _ = LeaveType.objects.get_or_create(code="CL", defaults={"name": "Casual Leave"})
        emp = self.add_staff("AC801", "Cel")
        LeaveRequest.objects.create(
            employee=emp,
            leave_type_ref=cl,
            type="casual",
            start_date="2026-08-12",
            end_date="2026-08-12",
            total_days=1,
            status="approved",
        )
        seed_records(emp, d(12), d(12))
        body = self.run_report(
            "attendance-muster-sheet",
            dateFrom="2026-08-12",
            dateTo="2026-08-12",
            showLeaveCodes="true",
            staffType="all",
        )
        rows = {r["employeeCode"]: r for r in body["rows"] if not r.get("_kind")}
        self.assertEqual(rows["AC801"]["d20260812"], "LCL")
        self.assertTrue(any("LCL" in n for n in body["notes"]))

    def _night_exit_saturday_off(self):
        ps = PayrollSettings.get()
        ps.last_punch_post_shift_grace_hours = 9
        ps.first_punch_pre_shift_buffer_hours = 2
        ps.save()
        sat = self.add_staff("AC807", "Nite")
        EmployeeShiftAssignment.objects.filter(employee=sat).update(saturday_off=True)
        punch(sat, d(14), "09:00", "13:00", "13:45")
        punch(sat, d(15), "01:05")  # Friday's forgotten exit, stamped after midnight
        seed_records(sat, d(14), d(15))
        return sat

    def test_muster_and_summary_count_a_night_exit_saturday_off_as_a_weekly_off_not_an_absence(self):
        sat = self._night_exit_saturday_off()
        one = {"employeeIds": str(sat.id), "staffType": "all"}
        muster = self.run_report(
            "attendance-muster-sheet", dateFrom="2026-08-14", dateTo="2026-08-15", weeklyOff="true", **one
        )
        row = self.row(muster, "AC807")
        self.assertEqual((row["d20260815"], row["weeklyOff"]), ("WO", 1))
        summary = self.row(self.run_report("attendance-summary-monthly", period="2026-08", **one), "AC807")
        self.assertEqual((summary["weeklyOff"], summary["absentDays"]), (1, 0))

    def test_a_saturday_with_a_real_punch_is_still_judged_by_the_engine_not_a_weekly_off(self):
        sat = self.add_staff("AC808", "Real")
        EmployeeShiftAssignment.objects.filter(employee=sat).update(saturday_off=True)
        punch(sat, d(15), "14:00")  # a lone punch inside the half-day gap: the engine stores it as absent
        seed_records(sat, d(15), d(15))
        self.assertEqual(AttendanceDayRecord.objects.get(employee=sat, date=d(15)).status, "absent")
        body = self.run_report(
            "absentee-list-daily", dateFrom="2026-08-15", dateTo="2026-08-15", employeeIds=str(sat.id)
        )
        self.assertEqual(self.codes(body), ["AC808"])
