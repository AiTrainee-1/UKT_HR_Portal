"""
Report Center, group G7 (requests / permission / on-duty / missing punch): permission-register,
permission-monthly-counts, permission-excess-salary-impact, on-duty-register, on-duty-employee-summary,
on-duty-punch-verification-log, missing-punch-register, missing-punch-monthly-counts, absence-around-holidays and
leave-attendance-conflicts.

Every date is fixed (March 2026 unless stated) so nothing depends on the day the suite runs.

Run via: python manage.py test api.tests_reporting_requests_permission_duty -v 2
"""

import io
from datetime import date, datetime, time, timedelta, timezone as dt_timezone
from decimal import Decimal

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .jwt_utils import sign_token
from .models import (
    AttendanceDayRecord,
    AttendanceLog,
    Branch,
    CasualLeaveRequest,
    Department,
    DepartmentManager,
    Designation,
    Employee,
    EmployeePermission,
    EmployeeShiftAssignment,
    Holiday,
    HRUser,
    LeaveRequest,
    LeaveType,
    ManagerDepartmentAssignment,
    MissingPunchRequest,
    OnDutyPunchVerification,
    OnDutySession,
    OutpassRequest,
    PayrollSettings,
    Role,
    SalarySlip,
    ShiftTemplate,
)
from .payroll_views import _generate_staff_payroll
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import requests_permission_duty_checks as checks
from .reporting.definitions import requests_permission_duty_permissions as perms_mod

UTC = dt_timezone.utc

DOMAIN_MODULES = {
    "reports": "view",
    "requests": "view",
    "geo_attendance": "view",
    "missing_punch": "view",
    "attendance": "view",
    "leave": "view",
    "casual_leave": "view",
    "payroll": "view",
}
GROUP_IDS = (
    "permission-register", "permission-monthly-counts", "permission-excess-salary-impact",
    "on-duty-register", "on-duty-employee-summary", "on-duty-punch-verification-log",
    "missing-punch-register", "missing-punch-monthly-counts", "absence-around-holidays", "leave-attendance-conflicts",
)
NO_ONE = "999999999"  # a valid id that matches no employee -> an empty report


def _headers(user):
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def utc(y, m, d, hh=0, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=UTC)


class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.settings_row = PayrollSettings.get()
        cls.b1 = Branch.objects.create(name="Unit 1", code="PDU1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="PDU2")
        cls.cutting = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.finishing = Department.objects.create(name="FINISHING", branch=cls.b1)
        cls.sewing = Department.objects.create(name="SEWING", branch=cls.b2)
        cls.operator = Designation.objects.create(title="Operator")

        mk = cls.mk
        cls.a = mk("PD_A", "Asha", cls.cutting, cls.b1)
        cls.b = mk("PD_B", "Bala", cls.finishing, cls.b1)
        cls.c = mk("PD_C", "Chitra", cls.sewing, cls.b2)
        cls.p = mk("PD_P", "Pavan", cls.cutting, cls.b1, employment_type="production")
        cls.x = mk("PD_X", "Xavier", None, cls.b1, desig=False)
        cls.gone = mk("PD_G", "Gopal", cls.cutting, cls.b1, status="inactive")

        cls.admin = HRUser.objects.create(username="pd_admin", password_hash="x", is_super_admin=True)
        cls.b1_user = HRUser.objects.create(
            username="pd_b1", password_hash="x", branch=cls.b1,
            role=Role.objects.create(name="pd_b1_role", permissions=dict(DOMAIN_MODULES)),
        )
        cls.only_reports = HRUser.objects.create(
            username="pd_plain", password_hash="x", role=Role.objects.create(name="pd_plain", permissions={"reports": "view"}),
        )

    @classmethod
    def mk(cls, code, first, dept, branch, desig=True, **kw):
        kw.setdefault("employment_type", "staff")
        kw.setdefault("salary_amount", Decimal("26000"))
        return Employee.objects.create(
            employee_code=code, first_name=first, last_name="Kumar", department=dept, branch=branch,
            designation=cls.operator if desig else None, **kw,
        )

    def hr(self, name, perms, **kw):
        return HRUser.objects.create(
            username=name, password_hash="x", role=Role.objects.create(name=name, permissions=perms), **kw
        )

    # ── http helpers ──
    def get(self, path, user=None, **params):
        return self.client.get(path, params, **_headers(user or self.admin))

    def run_report(self, rid, user=None, **params):
        r = self.get(f"/api/reports/run/{rid}", user, **params)
        self.assertEqual(r.status_code, 200, r.content[:400])
        return r.json()

    def export(self, rid, fmt, user=None, **params):
        return self.get(f"/api/reports/export/{rid}", user, fmt=fmt, **params)

    def codes(self, body, key="employeeCode"):
        return [r[key] for r in body["rows"] if r.get("_kind") != "subtotal"]

    def data_rows(self, body):
        return [r for r in body["rows"] if r.get("_kind") != "subtotal"]

    def assert_totals_consistent(self, body):
        """The totals row equals the sum of the data rows; every subtotal equals the sum of its group."""
        data = self.data_rows(body)
        for col in body["columns"]:
            if col["total"] == "sum" and body["totals"] is not None:
                want = sum(r[col["key"]] for r in data if isinstance(r.get(col["key"]), (int, float)))
                self.assertAlmostEqual(body["totals"][col["key"]] or 0, want, places=2, msg=col["key"])
        group: list = []
        for r in body["rows"]:
            if r.get("_kind") == "subtotal":
                for col in body["columns"]:
                    if col["total"] == "sum":
                        want = sum(g[col["key"]] for g in group if isinstance(g.get(col["key"]), (int, float)))
                        got = r.get(col["key"])
                        self.assertAlmostEqual(got or 0, want, places=2, msg=f"subtotal {col['key']}")
                group = []
            else:
                group.append(r)

    def summary(self, body):
        return {s["label"]: s["value"] for s in body["summary"]}

    def query_count(self, rid, **params):
        with CaptureQueriesContext(connection) as ctx:
            body = self.run_report(rid, **params)
        return len(ctx), body

    def assert_constant_queries(self, rid, params, grow, slack=3):
        """The query count must not depend on the row count (no N+1)."""
        self.run_report(rid, **params)  # warm any per-process caches
        small, body_small = self.query_count(rid, **params)
        grow()
        big, body_big = self.query_count(rid, **params)
        self.assertGreater(len(self.data_rows(body_big)), len(self.data_rows(body_small)))
        self.assertLessEqual(big - small, slack, f"{rid}: {small} -> {big} queries")

    def many(self, n, prefix="PD_M", dept=None, branch=None, **kw):
        return [
            self.mk(f"{prefix}{i:02d}", f"Many{i:02d}", dept or self.cutting, branch or self.b1, **kw)
            for i in range(n)
        ]

    # ── contract shared by every report ──
    def check_exports(self, rid, params, first_code, header_first="Emp Code"):
        """xlsx round-trips (header + a data cell), pdf is a pdf, and an empty result works in all three formats."""
        r = self.export(rid, "xlsx", **params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertTrue(r.content.startswith(b"PK"))
        ws = load_workbook(io.BytesIO(r.content)).active
        header_row = next(i for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15), 1) if row[0].value == header_first)
        headers = [c.value for c in ws[header_row]]
        self.assertEqual(headers[0], header_first)
        self.assertIn(first_code, [ws.cell(row=header_row + 1, column=i).value for i in range(1, len(headers) + 1)])
        p = self.export(rid, "pdf", **params)
        self.assertEqual(p.status_code, 200)
        self.assertTrue(p.content.startswith(b"%PDF"))
        empty = dict(params, employeeIds=NO_ONE)
        body = self.run_report(rid, **empty)
        self.assertEqual(self.data_rows(body), [])
        self.assertEqual(self.export(rid, "xlsx", **empty).status_code, 200)
        self.assertEqual(self.export(rid, "pdf", **empty).status_code, 200)

    def check_gating(self, rid, params, owner_perms):
        """A role with Reports but without the owning module gets report_forbidden; with it, the report opens."""
        plain = self.hr(f"gate_none_{rid}", {"reports": "view"})
        for kind in ("run", "export"):
            r = self.get(f"/api/reports/{kind}/{rid}", plain, fmt="xlsx", **params)
            self.assertEqual(r.status_code, 403, kind)
            self.assertEqual(r.json()["error"], "report_forbidden", kind)
        owner = self.hr(f"gate_own_{rid}", {"reports": "view", **owner_perms})
        self.assertEqual(self.get(f"/api/reports/run/{rid}", owner, **params).status_code, 200)

    def check_branch_isolation(self, rid, params, own_codes, other_code, other_id):
        """A branch-scoped user only sees their branch and cannot widen it through the filters."""
        body = self.run_report(rid, self.b1_user, **params)
        self.assertEqual(sorted(set(self.codes(body))), sorted(own_codes))
        self.assertNotIn(other_code, self.codes(body))
        wide = self.run_report(rid, self.b1_user, branchIds=str(self.b2.id), **params)
        self.assertEqual(self.data_rows(wide), [])
        wide = self.run_report(rid, self.b1_user, employeeIds=str(other_id), **params)
        self.assertEqual(self.data_rows(wide), [])
        xl = self.export(rid, "xlsx", self.b1_user, **params)
        cells = [c.value for row in load_workbook(io.BytesIO(xl.content)).active.iter_rows() for c in row]
        self.assertNotIn(other_code, cells)


# ════════════════════════════════════════════════════════════════════════════
#  Registry / catalog
# ════════════════════════════════════════════════════════════════════════════


class RegistryTests(_Base):
    def test_all_group_reports_are_registered_with_real_modules_and_families(self):
        specs = {s.id: s for s in registry.all_specs()}
        for rid in GROUP_IDS:
            self.assertIn(rid, specs)
            self.assertEqual(specs[rid].category, "leave")
            self.assertTrue(specs[rid].modules, rid)
            for m in specs[rid].modules:
                self.assertIn(m, all_module_keys())
        fam = {rid: (specs[rid].family, specs[rid].variant) for rid in GROUP_IDS}
        self.assertEqual({fam[r][0] for r in ("permission-register", "permission-monthly-counts", "permission-excess-salary-impact")}, {"permissions"})
        self.assertEqual({fam[r][0] for r in ("on-duty-register", "on-duty-employee-summary", "on-duty-punch-verification-log")}, {"on-duty"})
        self.assertEqual({fam[r][0] for r in ("missing-punch-register", "missing-punch-monthly-counts")}, {"missing-punch"})
        self.assertEqual(len({v for _f, v in fam.values() if v}), len([v for _f, v in fam.values() if v]))
        self.assertEqual(registry.LOAD_ERRORS, {})

    def test_conflict_type_modules_are_real_and_permission_types_match_the_model(self):
        for m in checks.CONFLICT_MODULE.values():
            self.assertIn(m, all_module_keys())
        self.assertEqual(
            dict(perms_mod.TYPE_OPTIONS[:3]), dict(EmployeePermission.TYPE_LABELS),
        )

    def test_catalog_lists_the_reports_for_a_role_with_the_modules(self):
        body = self.get("/api/reports/catalog", self.b1_user).json()
        ids = {r["id"] for r in body["reports"]}
        self.assertTrue(set(GROUP_IDS) <= ids)
        plain = self.get("/api/reports/catalog", self.only_reports).json()
        self.assertFalse(set(GROUP_IDS) & {r["id"] for r in plain["reports"]})

    def test_reports_never_write(self):
        Holiday.objects.create(name="H", date=date(2026, 3, 5))
        EmployeePermission.objects.create(employee=self.a, date=date(2026, 3, 2), type="morning_late_in", status="approved")
        PayrollSettings.objects.all().delete()
        before = (AttendanceDayRecord.objects.count(), AttendanceLog.objects.count(), PayrollSettings.objects.count())
        for rid in GROUP_IDS:
            self.assertEqual(self.get(f"/api/reports/run/{rid}", dateFrom="2026-03-01", dateTo="2026-03-31", period="2026-03").status_code, 200, rid)
        self.assertEqual(before, (AttendanceDayRecord.objects.count(), AttendanceLog.objects.count(), PayrollSettings.objects.count()))
        self.assertEqual(PayrollSettings.objects.count(), 0)  # reading the settings never creates the singleton row
        PayrollSettings.get()  # the (framework) letterhead creates it on export; the attendance tables must still not change
        for rid in GROUP_IDS:
            self.assertEqual(self.export(rid, "xlsx", dateFrom="2026-03-01", dateTo="2026-03-31", period="2026-03").status_code, 200, rid)
        self.assertEqual(before[:2], (AttendanceDayRecord.objects.count(), AttendanceLog.objects.count()))


# ════════════════════════════════════════════════════════════════════════════
#  Permissions
# ════════════════════════════════════════════════════════════════════════════

MARCH = dict(dateFrom="2026-03-01", dateTo="2026-03-31")


def perm(emp, d, typ, status="approved", minutes=60, at=None, by=None, role=None, comment=None):
    return EmployeePermission.objects.create(
        employee=emp, date=d, type=typ, status=status, duration_minutes=minutes, permission_time=at,
        approved_by=by, approver_role=role, hr_comment=comment, reason=f"reason {d.isoformat()}",
    )


class _PermissionFixture(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()

        d = lambda day: date(2026, 3, day)  # noqa: E731
        cls.P1 = perm(cls.a, d(2), "morning_late_in", at=time(9, 30), by="Ravi", role="hr")
        cls.P2 = perm(cls.a, d(3), "Early Out", at=time(17, 15), by="Meena", role="dept_head")
        cls.P3 = perm(cls.a, d(4), "middle_permission", minutes=90, at=time(11, 0))
        cls.P4 = perm(cls.a, d(5), "Late In", at=time(9, 30))
        cls.P5 = perm(cls.a, d(6), "morning_late_in", status="pending")
        cls.P6 = perm(cls.a, d(9), "middle_permission", status="rejected", comment="Not needed")
        cls.P7 = perm(cls.a, d(10), None, at=time(12, 0))
        cls.Q1 = perm(cls.b, d(11), "morning_late_in", minutes=None, at=time(9, 20))
        cls.Q2 = perm(cls.b, d(12), "evening_early_out", status="pending")
        cls.R1 = perm(cls.c, d(12), "morning_late_in")
        cls.R2 = perm(cls.c, d(13), "middle_permission")
        cls.S1 = perm(cls.p, d(16), "morning_late_in", at=time(23, 30))
        cls.G1 = perm(cls.gone, d(17), "morning_late_in")
        cls.X1 = perm(cls.x, d(18), "middle_permission")
        EmployeePermission.objects.filter(pk=cls.P1.pk).update(created_at=utc(2026, 3, 1, 4, 30))  # 10:00 IST
        AttendanceDayRecord.objects.create(employee=cls.a, date=d(2), status="present", morning_permission_applied=True)
        AttendanceDayRecord.objects.create(employee=cls.a, date=d(4), status="present")
        AttendanceDayRecord.objects.create(employee=cls.a, date=d(5), status="present", morning_permission_excess=True)


class PermissionRegisterTests(_PermissionFixture):
    RID = "permission-register"

    def test_golden_rows_cap_position_and_day_effect(self):
        body = self.run_report(self.RID, **MARCH)
        self.assertEqual(body["rowCount"], 14)
        rows = body["rows"]
        self.assertEqual(
            [(r["employeeCode"], r["date"]) for r in rows],
            [("PD_A", "2026-03-02"), ("PD_A", "2026-03-03"), ("PD_A", "2026-03-04"), ("PD_A", "2026-03-05"),
             ("PD_A", "2026-03-06"), ("PD_A", "2026-03-09"), ("PD_A", "2026-03-10"), ("PD_B", "2026-03-11"),
             ("PD_B", "2026-03-12"), ("PD_C", "2026-03-12"), ("PD_C", "2026-03-13"), ("PD_P", "2026-03-16"),
             ("PD_G", "2026-03-17"), ("PD_X", "2026-03-18")],
        )
        by_date = {(r["employeeCode"], r["date"]): r for r in rows}
        p1 = by_date[("PD_A", "2026-03-02")]
        self.assertEqual(p1["typeLabel"], "Morning Late-In")
        self.assertEqual((p1["permissionTime"], p1["permissionEnd"], p1["durationMinutes"]), ("09:30", "10:30", 60))
        self.assertEqual((p1["status"], p1["capStatus"], p1["monthSeq"]), ("Approved", "Within cap", 1))
        self.assertEqual(p1["decidedBy"], "Ravi (HR)")
        self.assertEqual(p1["appliedOn"], "2026-03-01 10:00")  # stored 04:30 UTC, shown in IST
        self.assertEqual(p1["dayEffect"], "Late-in boundary moved +60 min")
        self.assertEqual(p1["department"], "CUTTING")
        self.assertEqual(p1["branch"], "Unit 1")
        p2 = by_date[("PD_A", "2026-03-03")]  # legacy spelling "Early Out"
        self.assertEqual((p2["typeLabel"], p2["capStatus"], p2["monthSeq"]), ("Evening Early-Out", "Within cap", 2))
        self.assertEqual(p2["decidedBy"], "Meena (HOD)")
        self.assertIsNone(p2["dayEffect"])  # no attendance record for that day yet
        p3 = by_date[("PD_A", "2026-03-04")]
        self.assertEqual((p3["durationMinutes"], p3["permissionEnd"], p3["monthSeq"]), (90, "12:30", 3))
        self.assertEqual(p3["dayEffect"], "Not reflected on the day record")
        p4 = by_date[("PD_A", "2026-03-05")]  # legacy "Late In", 4th approved -> excess
        self.assertEqual((p4["typeLabel"], p4["capStatus"], p4["monthSeq"]), ("Morning Late-In", "Excess", 4))
        self.assertEqual(p4["dayEffect"], "Excess: late-in not protected")
        p5, p6 = by_date[("PD_A", "2026-03-06")], by_date[("PD_A", "2026-03-09")]
        self.assertEqual((p5["status"], p5["capStatus"], p5["monthSeq"], p5["dayEffect"]), ("Pending", None, None, "Pending - no effect yet"))
        self.assertEqual((p6["status"], p6["hrComment"], p6["dayEffect"]), ("Rejected", "Not needed", "Rejected - no effect"))
        p7 = by_date[("PD_A", "2026-03-10")]  # untyped legacy row still counts toward the cap
        self.assertEqual((p7["typeLabel"], p7["capStatus"], p7["monthSeq"]), ("Unclassified", "Excess", 5))
        self.assertEqual(p7["dayEffect"], "Counts toward the cap; type unknown")
        q1 = by_date[("PD_B", "2026-03-11")]  # no stored duration -> the fixed 60 minutes
        self.assertEqual((q1["durationMinutes"], q1["permissionEnd"], q1["capStatus"], q1["monthSeq"]), (60, "10:20", "Within cap", 1))
        s1 = by_date[("PD_P", "2026-03-16")]  # production; 23:30 + 60 min crosses midnight
        self.assertEqual((s1["permissionTime"], s1["permissionEnd"]), ("23:30", "00:30"))
        self.assertEqual(s1["dayEffect"], "None (production: permissions have no effect)")
        x1 = by_date[("PD_X", "2026-03-18")]  # no department
        self.assertEqual(x1["department"], "Unassigned")

    def test_summary_and_notes(self):
        body = self.run_report(self.RID, **MARCH)
        s = self.summary(body)
        self.assertEqual(s["Total requests"], 14)
        self.assertEqual((s["Approved"], s["Pending"], s["Rejected"]), (11, 2, 1))
        self.assertEqual((s["Within cap"], s["Excess"], s["Employees with excess"]), (9, 2, 1))
        self.assertEqual(s["Approved hours"], 11.5)  # 330 + 60 + 120 + 60 + 60 + 60 minutes
        notes = " ".join(body["notes"])
        self.assertIn("cap in force: 3", notes)
        self.assertIn("Unclassified", notes)
        self.assertIn("60 minutes", notes)  # the NULL-duration row
        self.assertIn("Morning Late-In 7", notes)

    def test_filters_narrow_the_result(self):
        n = lambda **kw: len(self.run_report(self.RID, **{**MARCH, **kw})["rows"])  # noqa: E731
        self.assertEqual(n(permissionType="morning_late_in"), 7)  # incl. the legacy "Late In" row
        self.assertEqual(n(permissionType="middle_permission"), 4)
        self.assertEqual(n(permissionType="evening_early_out"), 2)
        self.assertEqual(n(permissionType="unclassified"), 1)
        self.assertEqual(n(status="pending"), 2)
        self.assertEqual(n(status="rejected"), 1)
        self.assertEqual(n(status="approved"), 11)
        self.assertEqual(n(capStatus="excess"), 2)
        self.assertEqual(n(capStatus="within_cap"), 9)
        self.assertEqual(n(approverRole="hr"), 1)
        self.assertEqual(n(approverRole="dept_head"), 1)
        self.assertEqual(n(departmentIds=str(self.finishing.id)), 2)
        self.assertEqual(n(employmentType="production"), 1)
        self.assertEqual(n(employmentType="staff"), 13)
        self.assertEqual(n(employeeIds=f"{self.a.id},{self.c.id}"), 9)
        self.assertEqual(n(employeeStatus="active"), 13)
        self.assertEqual(n(employeeStatus="inactive"), 1)
        self.assertEqual(n(branchIds=str(self.b2.id)), 2)
        self.assertEqual(n(designationIds=str(self.operator.id)), 13)  # PD_X has no designation
        self.assertEqual(len(self.run_report(self.RID, dateFrom="2026-03-05", dateTo="2026-03-06")["rows"]), 2)
        self.assertEqual(len(self.run_report(self.RID, dateFrom="2026-04-01", dateTo="2026-04-30")["rows"]), 0)

    def test_the_cap_comes_from_the_live_setting(self):
        row = lambda: {r["date"]: r for r in self.run_report(self.RID, employeeIds=str(self.a.id), **MARCH)["rows"]}  # noqa: E731
        self.assertEqual(row()["2026-03-05"]["capStatus"], "Excess")
        ps = PayrollSettings.get()
        ps.permission_monthly_cap = 4
        ps.save()
        after = row()
        self.assertEqual(after["2026-03-05"]["capStatus"], "Within cap")
        self.assertEqual(after["2026-03-10"]["capStatus"], "Excess")

    def test_cap_status_agrees_with_the_engine_helper(self):
        from .attendance_final import permission_cap_status

        body = self.run_report(self.RID, employeeIds=str(self.a.id), **MARCH)
        engine = {p.date.isoformat(): permission_cap_status(p, 3) for p in EmployeePermission.objects.filter(employee=self.a)}
        label = {"within_cap": "Within cap", "excess": "Excess", "not_applicable": None}
        for r in body["rows"]:
            self.assertEqual(r["capStatus"], label[engine[r["date"]]], r["date"])
            if r["monthSeq"] is not None:
                self.assertEqual(r["monthSeq"] <= 3, r["capStatus"] == "Within cap")

    def test_row_order_is_deterministic(self):
        first = self.run_report(self.RID, **MARCH)["rows"]
        second = self.run_report(self.RID, **MARCH)["rows"]
        self.assertEqual(first, second)

    def test_exports_empty_and_totals(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        body = self.run_report(self.RID, **MARCH)
        self.assertIsNone(body["totals"])  # no additive column here: minutes across statuses would mislead

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B", "PD_G", "PD_P", "PD_X"], "PD_C", self.c.id)

    def test_needs_the_requests_module(self):
        self.check_gating(self.RID, MARCH, {"requests": "view"})

    def test_query_count_is_flat(self):
        def grow():
            for i, e in enumerate(self.many(12)):
                perm(e, date(2026, 3, 20), "morning_late_in", at=time(9, 30))
                AttendanceDayRecord.objects.create(employee=e, date=date(2026, 3, 20), status="present", morning_permission_applied=True)

        for i, e in enumerate(self.many(2, prefix="PD_S")):
            perm(e, date(2026, 3, 20), "morning_late_in")
        self.assert_constant_queries(self.RID, MARCH, grow)


class PermissionCountsTests(_PermissionFixture):
    RID = "permission-monthly-counts"

    def rows_by_code(self, **params):
        body = self.run_report(self.RID, **{**MARCH, **params})
        return {r["employeeCode"]: r for r in self.data_rows(body)}, body

    def test_golden_counts_per_employee_month(self):
        rows, body = self.rows_by_code()
        self.assertEqual(sorted(rows), ["PD_A", "PD_B", "PD_C", "PD_G", "PD_P", "PD_X"])
        a = rows["PD_A"]
        self.assertEqual(a["month"], "2026-03")
        self.assertEqual((a["morningLateIn"], a["eveningEarlyOut"], a["middleOneHour"], a["unclassified"]), (2, 1, 1, 1))
        self.assertEqual((a["totalApproved"], a["pending"], a["rejected"]), (5, 1, 1))
        self.assertEqual((a["monthlyCap"], a["excess"], a["capUsedPct"]), (3, 2, 166.7))
        self.assertEqual((a["approvedMinutes"], a["lastPermissionDate"]), (330, "2026-03-10"))
        b = rows["PD_B"]
        self.assertEqual((b["totalApproved"], b["pending"], b["excess"], b["capUsedPct"], b["approvedMinutes"]), (1, 1, 0, 33.3, 60))
        c = rows["PD_C"]
        self.assertEqual((c["morningLateIn"], c["middleOneHour"], c["totalApproved"], c["capUsedPct"], c["approvedMinutes"]), (1, 1, 2, 66.7, 120))
        self.assertEqual(rows["PD_X"]["department"], "Unassigned")

    def test_summary_totals_and_subtotals_add_up(self):
        rows, body = self.rows_by_code()
        s = self.summary(body)
        self.assertEqual((s["Approved permissions"], s["Excess permissions"], s["Employees over the cap"]), (11, 2, 1))
        self.assertEqual((s["Approved hours"], s["Pending requests"]), (11.5, 2))
        self.assert_totals_consistent(body)
        self.assertEqual(body["totals"]["totalApproved"], 11)
        self.assertEqual(body["totals"]["excess"], 2)
        subtotals = {r["employeeName"]: r for r in body["rows"] if r.get("_kind") == "subtotal"}
        self.assertEqual(sorted(subtotals), ["CUTTING (Unit 1) total", "FINISHING (Unit 1) total", "SEWING (Unit 2) total", "Unassigned (Unit 1) total"])
        self.assertEqual(subtotals["CUTTING (Unit 1) total"]["totalApproved"], 7)  # A 5 + G 1 + P 1

    def test_filters(self):
        self.assertEqual(sorted(self.rows_by_code(onlyExcess="true")[0]), ["PD_A"])
        self.assertEqual(sorted(self.rows_by_code(employmentType="production")[0]), ["PD_P"])
        self.assertEqual(sorted(self.rows_by_code(departmentIds=str(self.sewing.id))[0]), ["PD_C"])
        self.assertEqual(sorted(self.rows_by_code(employeeStatus="active")[0]), ["PD_A", "PD_B", "PD_C", "PD_P", "PD_X"])
        self.assertEqual(sorted(self.rows_by_code(employeeIds=str(self.b.id))[0]), ["PD_B"])
        self.assertEqual(sorted(self.rows_by_code(branchIds=str(self.b2.id))[0]), ["PD_C"])

    def test_range_spanning_months_gives_one_row_per_month_and_is_widened_to_whole_months(self):
        perm(self.a, date(2026, 4, 7), "middle_permission")
        rows = self.data_rows(self.run_report(self.RID, dateFrom="2026-03-20", dateTo="2026-04-10", employeeIds=str(self.a.id)))
        self.assertEqual([(r["month"], r["totalApproved"]) for r in rows], [("2026-03", 5), ("2026-04", 1)])  # March counted whole
        body = self.run_report(self.RID, dateFrom="2026-03-20", dateTo="2026-04-10", employeeIds=str(self.a.id))
        self.assertIn("widened to whole months (2026-03-01 to 2026-04-30)", " ".join(body["notes"]))

    def test_excess_uses_the_engines_late_pool_formula(self):
        from .attendance_final import late_pool_summary

        rows, _ = self.rows_by_code()
        ps = PayrollSettings.get()
        for code, r in rows.items():
            self.assertEqual(r["excess"], late_pool_summary((), r["totalApproved"], ps)["excess_permissions"], code)

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"requests": "edit"})

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B", "PD_G", "PD_P", "PD_X"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for e in self.many(12):
                perm(e, date(2026, 3, 20), "morning_late_in")
                perm(e, date(2026, 3, 21), "middle_permission", status="pending")

        for e in self.many(2, prefix="PD_S"):
            perm(e, date(2026, 3, 20), "morning_late_in")
        self.assert_constant_queries(self.RID, MARCH, grow)


def impact_perm(emp, day, status="approved"):
    return EmployeePermission.objects.create(
        employee=emp, date=date(2026, 3, day), type="morning_late_in", status=status, duration_minutes=60,
    )


def impact_punch(emp, day, first, last=time(18, 0)):
    for t in (first, last):
        AttendanceLog.objects.create(employee=emp, date=date(2026, 3, day), punch_time=t, punch_type="IN", source="test")


class SalaryImpactTests(_Base):
    RID = "permission-excess-salary-impact"
    PARAMS = {"period": "2026-03"}

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        ps = PayrollSettings.get()
        ps.attendance_mode = "simple"
        ps.morning_late_in_enabled = True
        ps.evening_early_out_enabled = False
        ps.permission_monthly_cap = 3
        ps.late_free_allowance = 3
        ps.save()
        shift = ShiftTemplate.objects.create(
            name="PD Shift", shift_type="staff", start_time=time(9, 0), end_time=time(18, 0), grace_period_minutes=15,
            first_half_end=time(13, 30), lunch_duration_minutes=60, lunch_grace_minutes=10,
        )
        cls.d = cls.mk("PD_D", "Devi", cls.cutting, cls.b1)
        for emp in (cls.a, cls.b, cls.d):
            EmployeeShiftAssignment.objects.create(employee=emp, shift=shift, effective_from=date(2020, 1, 1))

        perm, punch = impact_perm, impact_punch

        # A: 5 approved permissions (3 in cap + 2 excess) and 5 plain lates -> pool 5 + 2 = 7, billable 4 -> 0.25 shift.
        for day in (2, 3, 4, 5, 6):
            perm(cls.a, day)
            punch(cls.a, day, time(10, 10))
        for day in (9, 10, 11, 12, 13):
            punch(cls.a, day, time(9, 40))
        # B: one in-cap permission, no lateness.  D: three in-cap permissions; a 4th is approved AFTER the payslip.
        perm(cls.b, 2)
        punch(cls.b, 2, time(10, 10))
        for day in (2, 3, 4):
            perm(cls.d, day)
            punch(cls.d, day, time(10, 10))
        cls.slips = {e.employee_code: _generate_staff_payroll(e, 3, 2026)["slip"] for e in (cls.a, cls.b, cls.d)}
        perm(cls.d, 5)
        # C (branch 2): four approved permissions and no payslip.  P (production): one, never listed.
        for day in (2, 3, 4, 5):
            perm(cls.c, day)
        EmployeePermission.objects.create(employee=cls.p, date=date(2026, 3, 2), type="morning_late_in", status="approved")

    def rows(self, **params):
        body = self.run_report(self.RID, **{**self.PARAMS, **params})
        return {r["employeeCode"]: r for r in self.data_rows(body)}, body

    def test_the_generated_slip_is_what_this_test_assumes(self):
        late = self.slips["PD_A"].breakdown_details["deductions"]["lateSummary"]
        self.assertEqual(
            (late["lateInCount"], late["excessPermissionCount"], late["totalLateCount"], late["billableLateCount"], late["shiftDeductions"]),
            (5, 2, 7, 4, 0.25),
        )
        self.assertEqual(self.slips["PD_A"].breakdown_details["earnings"]["dailyRate"], 1000.0)
        self.assertEqual(self.slips["PD_A"].other_deductions, Decimal("250.00"))

    def test_golden_rows(self):
        rows, body = self.rows()
        self.assertEqual(sorted(rows), ["PD_A", "PD_B", "PD_C", "PD_D"])  # not the production employee
        a = rows["PD_A"]
        self.assertEqual((a["approvedPermissions"], a["monthlyCap"], a["excessPermissions"]), (5, 3, 2))
        self.assertEqual((a["lateInCount"], a["earlyOutCount"], a["totalPool"], a["freeAllowance"], a["billableLate"]), (5, 0, 7, 3, 4))
        self.assertEqual((a["shiftDeductions"], a["dailyRate"], a["latePenalty"]), (0.25, 1000.0, 250.0))
        # Without the two excess permissions the pool would be 5 -> billable 2 -> below the first slab -> no deduction.
        self.assertEqual(a["excessCost"], 250.0)
        self.assertEqual(a["slipStatus"], "Payslip generated")
        b = rows["PD_B"]
        self.assertEqual((b["approvedPermissions"], b["excessPermissions"], b["totalPool"], b["latePenalty"], b["excessCost"]), (1, 0, 0, 0.0, 0.0))
        d = rows["PD_D"]  # a 4th permission was approved after the payslip was generated
        self.assertEqual((d["approvedPermissions"], d["excessPermissions"], d["slipStatus"]), (4, 0, "Changed since payslip"))
        c = rows["PD_C"]  # no payslip: live counts only, the pool is unknown (dash), not zero
        self.assertEqual((c["approvedPermissions"], c["excessPermissions"], c["slipStatus"]), (4, 1, "No payslip"))
        for key in ("lateInCount", "totalPool", "billableLate", "shiftDeductions", "dailyRate", "latePenalty", "excessCost"):
            self.assertIsNone(c[key], key)

    def test_summary_totals_and_subtotals(self):
        rows, body = self.rows()
        s = self.summary(body)
        self.assertEqual(s["Late-pool salary deducted"], 250.0)
        self.assertEqual(s["Salary caused by excess permissions"], 250.0)
        self.assertEqual((s["Employees penalised"], s["Excess permissions"], s["Billable occurrences"], s["Changed since payslip"]), (1, 3, 4, 1))
        self.assert_totals_consistent(body)
        self.assertEqual(body["totals"]["latePenalty"], 250.0)
        self.assertEqual(body["totals"]["approvedPermissions"], 14)
        labels = [r["employeeName"] for r in body["rows"] if r.get("_kind") == "subtotal"]
        self.assertEqual(labels, ["CUTTING (Unit 1) total", "FINISHING (Unit 1) total", "SEWING (Unit 2) total"])

    def test_show_filter_and_scope_filters(self):
        self.assertEqual(sorted(self.rows(show="excess")[0]), ["PD_A", "PD_C"])
        self.assertEqual(sorted(self.rows(show="penalised")[0]), ["PD_A"])
        self.assertEqual(sorted(self.rows(departmentIds=str(self.finishing.id))[0]), ["PD_B"])
        self.assertEqual(sorted(self.rows(branchIds=str(self.b2.id))[0]), ["PD_C"])
        self.assertEqual(sorted(self.rows(employeeIds=str(self.d.id))[0]), ["PD_D"])
        self.assertEqual(self.rows(period="2026-04")[0], {})

    def test_the_excess_cost_is_blank_when_the_slabs_changed_since_the_payslip(self):
        ps = PayrollSettings.get()
        ps.late_deduction_slabs = [{"fromLates": 3, "deductionShifts": 1.0}]
        ps.save()
        rows, _ = self.rows()
        self.assertEqual(rows["PD_A"]["latePenalty"], 250.0)  # still what was paid
        self.assertIsNone(rows["PD_A"]["excessCost"])  # re-pricing with the new slabs would not reproduce the payslip

    def test_a_payslip_without_late_data_is_reported_as_such(self):
        SalarySlip.objects.filter(pk=self.slips["PD_B"].pk).update(breakdown_details={"type": "staff"})
        rows, _ = self.rows()
        self.assertEqual(rows["PD_B"]["slipStatus"], "Payslip without late data")
        self.assertIsNone(rows["PD_B"]["latePenalty"])

    def test_production_slips_and_other_months_are_ignored(self):
        SalarySlip.objects.filter(pk=self.slips["PD_D"].pk).update(month=2)
        rows, _ = self.rows()
        self.assertEqual(rows["PD_D"]["slipStatus"], "No payslip")

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, self.PARAMS, "PD_A")
        r = self.get(f"/api/reports/run/{self.RID}", self.hr("gate_requests_only", {"reports": "view", "requests": "view"}), **self.PARAMS)
        self.assertEqual((r.status_code, r.json()["error"]), (403, "report_forbidden"))  # salary data needs a payroll module
        for perms in ({"payroll": "view"}, {"salary": "view"}, {"salary_slip": "view"}):
            self.assertEqual(self.get(f"/api/reports/run/{self.RID}", self.hr(f"g_{next(iter(perms))}", {"reports": "view", **perms}), **self.PARAMS).status_code, 200)

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, self.PARAMS, ["PD_A", "PD_B", "PD_D"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for e in self.many(12):
                impact_perm(e, 20)

        for e in self.many(2, prefix="PD_S"):
            impact_perm(e, 20)
        self.assert_constant_queries(self.RID, self.PARAMS, grow)


# ════════════════════════════════════════════════════════════════════════════
#  On-duty
# ════════════════════════════════════════════════════════════════════════════


def od_session(emp, dest, status, created, **kw):
    s = OnDutySession.objects.create(employee=emp, destination=dest, status=status, branch=kw.pop("branch", emp.branch), **kw)
    OnDutySession.objects.filter(pk=s.pk).update(created_at=created)
    s.refresh_from_db()
    return s


def od_punch(s, n, t, status="approved", mocked=False, acc=None, photo="on_duty_punch_verifications/2026/03/p.jpg",
             day=None, lat="11.104512", lng="77.341234", **kw):
    return OnDutyPunchVerification.objects.create(
        session=s, employee=s.employee, punch_date=day or date(2026, 3, 10), punch_time=t,
        punch_type="IN" if n % 2 else "OUT", punch_number=n, latitude=Decimal(lat), longitude=Decimal(lng),
        accuracy_m=acc, is_mocked=mocked, photo=photo, status=status, **kw,
    )


class _OnDutyFixture(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()

        session, punch = od_session, od_punch
        # S1: 01:30 IST on 10 Mar (= 20:00 UTC on the 9th). HOD then HR; completed by the employee.
        cls.S1 = session(cls.a, "Coimbatore mill", "completed", utc(2026, 3, 9, 20, 0),
                         hod_reviewed_by="Meena K", hod_reviewed_at=utc(2026, 3, 9, 21, 0), hod_review_comment="ok",
                         hr_reviewed_by="Ravi", hr_reviewed_at=utc(2026, 3, 9, 23, 0), hr_review_comment="fine",
                         completion_reason="manual", started_at=utc(2026, 3, 9, 23, 0))
        punch(cls.S1, 1, time(8, 30), acc=10)
        punch(cls.S1, 2, time(12, 30), acc=20, mocked=True)
        punch(cls.S1, 3, time(13, 30), acc=30)
        punch(cls.S1, 4, time(17, 30), acc=40)
        cls.S2 = session(cls.a, "Erode", "rejected", utc(2026, 3, 11, 5, 0),
                         hod_reviewed_by="Meena K", hod_reviewed_at=utc(2026, 3, 11, 7, 0))
        punch(cls.S2, 1, time(9, 0), status="rejected", photo="", day=date(2026, 3, 11),
              hr_review_comment="Voided automatically -the On-Duty request was rejected by HOD.")
        cls.S3 = session(cls.b, "Chennai", "pending_hr", utc(2026, 3, 12, 4, 0),
                         hod_reviewed_by="Meena K", hod_reviewed_at=utc(2026, 3, 12, 6, 0),
                         employee_ended_at=utc(2026, 3, 12, 11, 0))
        punch(cls.S3, 1, time(9, 0), status="pending", acc=100, day=date(2026, 3, 12))
        punch(cls.S3, 2, time(9, 30), status="pending", day=date(2026, 3, 12))
        cls.S4 = session(cls.c, "Salem", "active", utc(2026, 3, 13, 4, 0),
                         hr_reviewed_by="Ravi", hr_reviewed_at=utc(2026, 3, 13, 9, 0))
        cls.S5 = session(cls.p, "Tiruppur", "pending_hod", utc(2026, 3, 14, 4, 0), branch=None)
        OutpassRequest.objects.create(
            employee=cls.a, destination="Coimbatore mill", reason="on duty", status="approved", source="on_duty",
            on_duty_session=cls.S1,
        )


class OnDutyRegisterTests(_OnDutyFixture):
    RID = "on-duty-register"

    def test_golden_rows(self):
        body = self.run_report(self.RID, **MARCH)
        self.assertEqual([r["destination"] for r in body["rows"]], ["Coimbatore mill", "Erode", "Chennai", "Salem", "Tiruppur"])
        s1, s2, s3, s4, s5 = body["rows"]
        self.assertEqual(s1["requestedAt"], "2026-03-10 01:30")  # IST, although stored on the 9th in UTC
        self.assertEqual((s1["status"], s1["decisionPath"], s1["hodBy"], s1["hrBy"]), ("Completed", "HOD, then HR", "Meena K", "Ravi"))
        self.assertEqual((s1["hodAt"], s1["hrAt"], s1["decisionHours"]), ("2026-03-10 02:30", "2026-03-10 04:30", 3.0))
        self.assertEqual(s1["comments"], "HOD: ok; HR: fine")
        self.assertEqual((s1["punchesTotal"], s1["punchesApproved"], s1["punchesPending"], s1["punchesRejected"]), (4, 4, 0, 0))
        self.assertEqual((s1["firstPunch"], s1["lastPunch"], s1["mockedPunches"]), ("08:30", "17:30", 1))
        self.assertEqual((s1["completion"], s1["outpass"], s1["branch"]), ("Employee marked done", "Issued", "Unit 1"))
        self.assertEqual((s2["status"], s2["decisionPath"], s2["decisionHours"]), ("Rejected", "Rejected by HOD", 2.0))
        self.assertEqual((s2["punchesTotal"], s2["punchesRejected"], s2["firstPunch"], s2["lastPunch"]), (1, 1, None, None))
        self.assertEqual((s3["status"], s3["decisionPath"], s3["decisionHours"]), ("Pending", "Awaiting HR (HOD approved)", None))
        self.assertEqual((s3["punchesPending"], s3["firstPunch"], s3["lastPunch"], s3["completion"]), (2, "09:00", "09:30", "Employee done (awaiting HR)"))
        self.assertEqual((s4["status"], s4["decisionPath"], s4["decisionHours"], s4["hodBy"], s4["punchesTotal"]), ("Active", "HR directly", 5.0, None, 0))
        self.assertEqual((s4["branch"], s4["department"]), ("Unit 2", "SEWING"))
        self.assertEqual((s5["status"], s5["decisionPath"], s5["branch"], s5["outpass"]), ("Pending", "Awaiting HOD", None, None))

    def test_summary_totals_and_notes(self):
        body = self.run_report(self.RID, **MARCH)
        s = self.summary(body)
        self.assertEqual((s["Sessions"], s["Pending decision"], s["Approved by HR"], s["Rejected"]), (5, 2, 2, 1))
        self.assertEqual((s["Punches captured"], s["Sessions with mock GPS"], s["Employees on duty"]), (7, 1, 4))
        self.assertEqual(s["Avg request-to-decision (h)"], 3.33)  # (3 + 2 + 5) / 3
        self.assertEqual(body["totals"]["punchesTotal"], 7)
        self.assertEqual(body["totals"]["mockedPunches"], 1)
        self.assert_totals_consistent(body)
        self.assertIn("IST date the on-duty request was submitted", " ".join(body["notes"]))

    def test_filters(self):
        dest = lambda **kw: [r["destination"] for r in self.run_report(self.RID, **{**MARCH, **kw})["rows"]]  # noqa: E731
        self.assertEqual(dest(status="pending_hod"), ["Tiruppur"])
        self.assertEqual(dest(status="pending_hr"), ["Chennai"])
        self.assertEqual(dest(status="completed"), ["Coimbatore mill"])
        self.assertEqual(dest(status="rejected"), ["Erode"])
        self.assertEqual(dest(decisionPath="hod_then_hr"), ["Coimbatore mill"])
        self.assertEqual(dest(decisionPath="hod_rejected"), ["Erode"])
        self.assertEqual(dest(decisionPath="hr_direct"), ["Salem"])
        self.assertEqual(dest(decisionPath="awaiting"), ["Chennai", "Tiruppur"])
        self.assertEqual(dest(completion="manual"), ["Coimbatore mill"])
        self.assertEqual(dest(completion="not_completed"), ["Erode", "Chennai", "Salem", "Tiruppur"])
        self.assertEqual(dest(destination="ERODE"), ["Erode"])
        self.assertEqual(dest(mockedOnly="true"), ["Coimbatore mill"])
        self.assertEqual(dest(departmentIds=str(self.finishing.id)), ["Chennai"])
        self.assertEqual(dest(employmentType="production"), ["Tiruppur"])
        self.assertEqual(dest(employeeIds=str(self.a.id)), ["Coimbatore mill", "Erode"])
        self.assertEqual(dest(branchIds=str(self.b2.id)), ["Salem"])
        self.assertEqual(dest(employeeStatus="inactive"), [])

    def test_the_date_is_the_ist_day_not_the_utc_day(self):
        one = lambda d: [r["destination"] for r in self.run_report(self.RID, dateFrom=d, dateTo=d)["rows"]]  # noqa: E731
        self.assertEqual(one("2026-03-10"), ["Coimbatore mill"])  # 01:30 IST on the 10th
        self.assertEqual(one("2026-03-09"), [])  # a UTC-date lookup would have filed it here
        self.assertEqual(one("2026-03-11"), ["Erode"])

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"geo_attendance": "view"})

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B", "PD_P"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for i, e in enumerate(self.many(12)):
                s = od_session(e, f"Dest {i}", "active", utc(2026, 3, 20, 4, 0), hr_reviewed_at=utc(2026, 3, 20, 6, 0))
                od_punch(s, 1, time(9, 0), day=date(2026, 3, 20))

        for i, e in enumerate(self.many(2, prefix="PD_S")):
            s = od_session(e, f"Dest {i}", "active", utc(2026, 3, 20, 4, 0), hr_reviewed_at=utc(2026, 3, 20, 6, 0))
            od_punch(s, 1, time(9, 0), day=date(2026, 3, 20))
        self.assert_constant_queries(self.RID, MARCH, grow)


class OnDutySummaryTests(_OnDutyFixture):
    RID = "on-duty-employee-summary"

    def test_golden_rows_subtotals_and_summary(self):
        body = self.run_report(self.RID, **MARCH)
        rows = {r["employeeCode"]: r for r in self.data_rows(body)}
        self.assertEqual(sorted(rows), ["PD_A", "PD_B", "PD_C", "PD_P"])
        a = rows["PD_A"]
        self.assertEqual((a["sessions"], a["daysOnDuty"], a["approved"], a["rejected"], a["pending"]), (2, 1, 1, 1, 0))
        self.assertEqual((a["punchesTotal"], a["punchesRejected"], a["mockedPunches"]), (5, 1, 1))
        self.assertEqual((a["topDestination"], a["lastOnDuty"]), ("Coimbatore mill", "2026-03-11"))  # tie -> alphabetical
        b = rows["PD_B"]
        self.assertEqual((b["sessions"], b["daysOnDuty"], b["approved"], b["pending"], b["punchesTotal"]), (1, 0, 0, 1, 2))
        c = rows["PD_C"]
        self.assertEqual((c["sessions"], c["daysOnDuty"], c["approved"], c["punchesTotal"]), (1, 1, 1, 0))
        s = self.summary(body)
        self.assertEqual((s["Employees with on-duty"], s["Sessions"], s["Approved on-duty days"], s["Mock-GPS punches"]), (4, 5, 2, 1))
        self.assertEqual(s["Most sessions"], "Asha Kumar (2)")
        self.assert_totals_consistent(body)
        self.assertEqual(body["totals"]["sessions"], 5)
        labels = [r["employeeName"] for r in body["rows"] if r.get("_kind") == "subtotal"]
        self.assertEqual(labels, ["CUTTING (Unit 1) total", "FINISHING (Unit 1) total", "SEWING (Unit 2) total"])

    def test_approved_days_count_distinct_dates(self):
        s = od_session(self.a, "Palladam", "active", utc(2026, 3, 10, 6, 0))  # same IST day as S1
        self.assertIsNotNone(s.pk)
        a = next(r for r in self.data_rows(self.run_report(self.RID, **MARCH)) if r["employeeCode"] == "PD_A")
        self.assertEqual((a["sessions"], a["approved"], a["daysOnDuty"]), (3, 2, 1))

    def test_filters(self):
        codes = lambda **kw: self.codes(self.run_report(self.RID, **{**MARCH, **kw}))  # noqa: E731
        self.assertEqual(codes(status="rejected"), ["PD_A"])
        self.assertEqual(codes(mockedOnly="true"), ["PD_A"])
        self.assertEqual(codes(decisionPath="awaiting"), ["PD_P", "PD_B"])  # ordered by department
        self.assertEqual(codes(employmentType="staff"), ["PD_A", "PD_B", "PD_C"])
        self.assertEqual(codes(destination="salem"), ["PD_C"])
        self.assertEqual(codes(dateFrom="2026-03-12", dateTo="2026-03-12"), ["PD_B"])

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"geo_attendance": "edit"})

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B", "PD_P"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for i, e in enumerate(self.many(12)):
                s = od_session(e, f"Dest {i}", "active", utc(2026, 3, 20, 4, 0))
                od_punch(s, 1, time(9, 0), day=date(2026, 3, 20))

        for e in self.many(2, prefix="PD_S"):
            od_session(e, "X", "active", utc(2026, 3, 20, 4, 0))
        self.assert_constant_queries(self.RID, MARCH, grow)


class OnDutyPunchLogTests(_OnDutyFixture):
    RID = "on-duty-punch-verification-log"

    def test_golden_rows_and_no_photo_leak(self):
        body = self.run_report(self.RID, **MARCH)
        self.assertEqual(body["rowCount"], 7)
        keys = [c["key"] for c in body["columns"]]
        self.assertNotIn("photo", keys)
        self.assertTrue(all("photo" not in k.lower() or k == "hasPhoto" for k in keys))
        r = body["rows"][1]  # 12:30 mocked punch of S1
        self.assertEqual((r["punchDate"], r["punchTime"], r["punchNumber"], r["punchType"]), ("2026-03-10", "12:30", 2, "OUT"))
        self.assertEqual((r["latitude"], r["longitude"], r["accuracyM"]), ("11.104512", "77.341234", 20.0))
        self.assertEqual((r["isMocked"], r["hasPhoto"], r["status"], r["sessionStatus"]), ("Mocked", "Yes", "Approved", "Completed"))
        self.assertEqual(r["mapLink"], "https://www.google.com/maps?q=11.104512,77.341234")
        self.assertEqual(r["destination"], "Coimbatore mill")
        voided = next(x for x in body["rows"] if x["destination"] == "Erode")
        self.assertEqual((voided["hasPhoto"], voided["status"], voided["accuracyM"]), ("No", "Rejected", None))
        self.assertTrue(voided["hrReviewComment"].startswith("Voided automatically"))
        self.assertEqual(self.summary(body), {
            "Punches": 7, "Approved": 4, "Pending": 2, "Rejected": 1, "Mock-GPS punches": 1, "Average accuracy (m)": 40.0,
        })
        self.assertEqual(body["totals"]["accuracyM"], 40.0)

    def test_filters(self):
        n = lambda **kw: len(self.run_report(self.RID, **{**MARCH, **kw})["rows"])  # noqa: E731
        self.assertEqual(n(status="pending"), 2)
        self.assertEqual(n(status="approved"), 4)
        self.assertEqual(n(status="rejected"), 1)
        self.assertEqual(n(mockedOnly="true"), 1)
        self.assertEqual(n(accuracyWorse="50"), 1)  # only the 100 m fix
        self.assertEqual(n(departmentIds=str(self.finishing.id)), 2)
        self.assertEqual(n(employeeIds=str(self.a.id)), 5)
        self.assertEqual(len(self.run_report(self.RID, dateFrom="2026-03-10", dateTo="2026-03-10")["rows"]), 4)
        self.assertEqual(len(self.run_report(self.RID, dateFrom="2026-03-13", dateTo="2026-03-31")["rows"]), 0)

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"geo_attendance": "view"})

    def test_pdf_and_xlsx_carry_no_photo_reference(self):
        cells = [c.value for row in load_workbook(io.BytesIO(self.export(self.RID, "xlsx", **MARCH).content)).active.iter_rows() for c in row]
        self.assertFalse(any("on_duty_punch_verifications" in str(c) or "/photo" in str(c) for c in cells))

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for i, e in enumerate(self.many(12)):
                s = od_session(e, f"Dest {i}", "active", utc(2026, 3, 20, 4, 0))
                od_punch(s, 1, time(9, 0), day=date(2026, 3, 20))

        for e in self.many(2, prefix="PD_S"):
            od_punch(od_session(e, "X", "active", utc(2026, 3, 20, 4, 0)), 1, time(9, 0), day=date(2026, 3, 20))
        self.assert_constant_queries(self.RID, MARCH, grow)


# ════════════════════════════════════════════════════════════════════════════
#  Missing punch
# ════════════════════════════════════════════════════════════════════════════


def mp_req(emp, day, at, typ, slot, status, created, **kw):
    r = MissingPunchRequest.objects.create(
        employee=emp, date=date(2026, 3, day), punch_time=at, punch_type=typ, punch_slot=slot, status=status,
        reason=f"forgot {day}", **kw,
    )
    MissingPunchRequest.objects.filter(pk=r.pk).update(created_at=created)
    r.refresh_from_db()
    return r


class _MissingPunchFixture(_Base):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()

        req = mp_req
        # An active HOD for CUTTING (with the missing-punch right) and one for SEWING (without it).
        cls.hod = cls.mk("PD_HOD", "Hari", cls.cutting, cls.b1)
        m = DepartmentManager.objects.create(employee=cls.hod, can_approve_missing_punch=True)
        ManagerDepartmentAssignment.objects.create(manager=m, department=cls.cutting)
        cls.hod2 = cls.mk("PD_HOD2", "Hema", cls.sewing, cls.b2)
        m2 = DepartmentManager.objects.create(employee=cls.hod2, can_approve_missing_punch=False)
        ManagerDepartmentAssignment.objects.create(manager=m2, department=cls.sewing)

        cls.M1 = req(cls.a, 2, time(9, 0), "IN", "morning_in", "approved", utc(2026, 3, 3, 5, 0),
                     hod_reviewed_by="Hari Kumar", hod_reviewed_at=utc(2026, 3, 3, 6, 0),
                     hr_reviewed_by="Ravi", hr_reviewed_at=utc(2026, 3, 4, 7, 0), hr_review_comment="ok")
        cls.M2 = req(cls.a, 3, time(13, 0), "OUT", "lunch_out", "approved", utc(2026, 3, 3, 5, 0),
                     hod_reviewed_at=utc(2026, 3, 3, 6, 0), hr_reviewed_at=utc(2026, 3, 3, 8, 0))
        cls.M3 = req(cls.a, 4, time(18, 0), "OUT", "evening_out", "approved", utc(2026, 3, 8, 5, 0),
                     hod_reviewed_at=utc(2026, 3, 8, 6, 0), hr_reviewed_at=utc(2026, 3, 8, 8, 0))
        cls.M4 = req(cls.a, 5, time(9, 5), "IN", None, "rejected", utc(2026, 3, 6, 5, 0),
                     hod_reviewed_by="Hari Kumar", hod_reviewed_at=utc(2026, 3, 6, 6, 30), hod_review_comment="not sure")
        cls.M5 = req(cls.a, 6, time(12, 30), "OUT", "lunch_out", "rejected", utc(2026, 3, 9, 5, 0),
                     hod_reviewed_at=utc(2026, 3, 9, 6, 0), hr_reviewed_at=utc(2026, 3, 9, 7, 0))
        cls.M6 = req(cls.a, 9, time(9, 0), "IN", "morning_in", "pending_hod", utc(2026, 3, 10, 5, 0))
        cls.M7 = req(cls.a, 10, time(9, 0), "IN", "morning_in", "pending_hr", utc(2026, 3, 11, 5, 0),
                     hod_reviewed_at=utc(2026, 3, 11, 6, 0))
        cls.M8 = req(cls.b, 11, time(9, 0), "IN", "morning_in", "pending_hod", utc(2026, 3, 12, 5, 0))  # no HOD
        cls.M9 = req(cls.c, 12, time(9, 0), "IN", "morning_in", "pending_hod", utc(2026, 3, 13, 5, 0))  # HOD lacks the right
        cls.M10 = req(cls.x, 13, time(17, 0), "OUT", "evening_out", "pending_hod", utc(2026, 3, 14, 5, 0))  # no dept, no HOD
        AttendanceLog.objects.create(employee=cls.a, date=date(2026, 3, 2), punch_time=time(9, 0), punch_type="IN", source="missing_punch:approved")
        AttendanceLog.objects.create(employee=cls.a, date=date(2026, 3, 3), punch_time=time(13, 0), punch_type="OUT", source="biometric:essl")


class MissingPunchRegisterTests(_MissingPunchFixture):
    RID = "missing-punch-register"

    def test_golden_rows(self):
        body = self.run_report(self.RID, **MARCH)
        self.assertEqual(body["rowCount"], 10)
        by = {r["reason"]: r for r in body["rows"]}
        m1 = by["forgot 2"]
        self.assertEqual((m1["date"], m1["punchSlot"], m1["punchType"], m1["punchTime"]), ("2026-03-02", "Morning Check-In", "IN", "09:00"))
        self.assertEqual((m1["appliedOn"], m1["lagDays"], m1["status"]), ("2026-03-03 10:30", 1, "Approved"))
        self.assertEqual((m1["hodBy"], m1["hodAt"], m1["hrBy"], m1["hrAt"], m1["comments"]), ("Hari Kumar", "2026-03-03 11:30", "Ravi", "2026-03-04 12:30", "HR: ok"))
        self.assertEqual((m1["turnaroundHours"], m1["punchWritten"], m1["pendingWith"]), (26.0, "Yes - added by approval", None))
        m2 = by["forgot 3"]
        self.assertEqual((m2["turnaroundHours"], m2["punchWritten"]), (3.0, "Yes - already recorded"))
        m3 = by["forgot 4"]
        self.assertEqual((m3["lagDays"], m3["punchWritten"]), (4, "Missing from punch log"))
        m4 = by["forgot 5"]  # rejected by the HOD: the HR columns stay empty
        self.assertEqual((m4["status"], m4["punchSlot"], m4["hrBy"], m4["hrAt"], m4["turnaroundHours"], m4["punchWritten"]), ("Rejected", "Not specified", None, None, 1.5, None))
        self.assertEqual(m4["comments"], "HOD: not sure")
        m5 = by["forgot 6"]
        self.assertEqual((m5["status"], m5["hrAt"], m5["turnaroundHours"]), ("Rejected", "2026-03-09 12:30", 2.0))
        self.assertEqual((by["forgot 9"]["status"], by["forgot 9"]["pendingWith"], by["forgot 9"]["turnaroundHours"]), ("Pending", "HOD: Hari Kumar", None))
        self.assertEqual(by["forgot 10"]["pendingWith"], "HR")
        self.assertEqual(by["forgot 11"]["pendingWith"], "No HOD assigned - HR cannot act")
        self.assertEqual(by["forgot 12"]["pendingWith"], "HOD Hema Kumar: no Missing Punch approval right - stuck")
        self.assertEqual((by["forgot 13"]["pendingWith"], by["forgot 13"]["department"]), ("No HOD assigned - HR cannot act", "Unassigned"))

    def test_summary_and_notes(self):
        body = self.run_report(self.RID, **MARCH)
        s = self.summary(body)
        self.assertEqual((s["Requests"], s["Approved"], s["Rejected by HOD"], s["Rejected by HR"], s["Pending"]), (10, 3, 1, 1, 5))
        self.assertEqual((s["Stuck (no HOD / no right)"], s["Approved but not in log"]), (3, 1))
        # decided requests: M1 26h, M2 3h, M3 3h, M4 1.5h, M5 2h -> 35.5 / 5
        self.assertEqual(s["Avg turnaround (h)"], 7.1)
        notes = " ".join(body["notes"])
        self.assertIn("Requests by slot", notes)
        self.assertIn("Morning Check-In 5", notes)

    def test_filters(self):
        reasons = lambda **kw: sorted(r["reason"] for r in self.run_report(self.RID, **{**MARCH, **kw})["rows"])  # noqa: E731
        self.assertEqual(len(reasons(status="pending")), 5)  # HOD and HR stage together
        self.assertEqual(reasons(status="pending_hr"), ["forgot 10"])
        self.assertEqual(len(reasons(status="pending_hod")), 4)
        self.assertEqual(len(reasons(status="approved")), 3)
        self.assertEqual(reasons(punchSlot="unspecified"), ["forgot 5"])
        self.assertEqual(len(reasons(punchSlot="morning_in")), 5)
        self.assertEqual(reasons(punchSlot="evening_out"), ["forgot 13", "forgot 4"])
        self.assertEqual(reasons(rejectedAt="hod"), ["forgot 5"])
        self.assertEqual(reasons(rejectedAt="hr"), ["forgot 6"])
        self.assertEqual(reasons(departmentIds=str(self.finishing.id)), ["forgot 11"])
        self.assertEqual(reasons(branchIds=str(self.b2.id)), ["forgot 12"])
        self.assertEqual(reasons(employeeIds=str(self.x.id)), ["forgot 13"])
        self.assertEqual(reasons(employeeStatus="inactive"), [])
        self.assertEqual(reasons(dateFrom="2026-03-04", dateTo="2026-03-05"), ["forgot 4", "forgot 5"])

    def test_the_hod_stage_follows_the_one_hod_rule(self):
        # Give PD_A an individual HOD (the SEWING HOD): it beats the department HOD.
        from .models import ManagerEmployeeAssignment

        ManagerEmployeeAssignment.objects.create(manager=self.hod2.manager_profile, employee=self.a)
        row = next(r for r in self.run_report(self.RID, **MARCH)["rows"] if r["reason"] == "forgot 9")
        self.assertEqual(row["pendingWith"], "HOD Hema Kumar: no Missing Punch approval right - stuck")

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"missing_punch": "view"})

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B", "PD_X"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for i, e in enumerate(self.many(12)):
                mp_req(e, 20, time(9, 0), "IN", "morning_in", "pending_hod", utc(2026, 3, 21, 5, 0))
                mp_req(e, 21, time(9, 0), "IN", "morning_in", "approved", utc(2026, 3, 22, 5, 0), hr_reviewed_at=utc(2026, 3, 22, 8, 0))

        for e in self.many(2, prefix="PD_S"):
            mp_req(e, 20, time(9, 0), "IN", "morning_in", "pending_hod", utc(2026, 3, 21, 5, 0))
        self.assert_constant_queries(self.RID, MARCH, grow)


class MissingPunchCountsTests(_MissingPunchFixture):
    RID = "missing-punch-monthly-counts"

    def test_golden_counts(self):
        body = self.run_report(self.RID, **MARCH)
        rows = {r["employeeCode"]: r for r in self.data_rows(body)}
        self.assertEqual(sorted(rows), ["PD_A", "PD_B", "PD_C", "PD_X"])
        a = rows["PD_A"]
        self.assertEqual(a["month"], "2026-03")
        self.assertEqual((a["requests"], a["approved"], a["rejected"], a["pending"]), (7, 3, 2, 2))
        self.assertEqual((a["morningIn"], a["lunchOut"], a["lunchIn"], a["eveningOut"], a["unspecified"]), (3, 2, 0, 1, 1))
        # days late: M1 1, M2 0, M3 4, M4 1, M5 3, M6 1, M7 1 -> 11 / 7
        self.assertEqual(a["avgLagDays"], 1.6)
        self.assertEqual(rows["PD_X"]["department"], "Unassigned")
        s = self.summary(body)
        self.assertEqual((s["Requests"], s["Employees with 3 or more"], s["Pending"]), (10, 1, 5))
        self.assertEqual(s["Approval rate"], 60.0)  # 3 approved of 5 decided
        self.assert_totals_consistent(body)
        self.assertEqual(body["totals"]["requests"], 10)

    def test_filters_and_month_bucketing(self):
        MissingPunchRequest.objects.create(
            employee=self.a, date=date(2026, 4, 2), punch_time=time(9, 0), punch_type="IN", punch_slot=None, status="approved", reason="apr"
        )
        rows = self.data_rows(self.run_report(self.RID, dateFrom="2026-03-01", dateTo="2026-04-30", employeeIds=str(self.a.id)))
        self.assertEqual([(r["month"], r["requests"], r["unspecified"]) for r in rows], [("2026-03", 7, 1), ("2026-04", 1, 1)])
        codes = lambda **kw: self.codes(self.run_report(self.RID, **{**MARCH, **kw}))  # noqa: E731
        self.assertEqual(codes(minRequests="3"), ["PD_A"])
        self.assertEqual(codes(minRequests="2"), ["PD_A"])
        self.assertEqual(codes(minRequests="5"), ["PD_A"])
        self.assertEqual(codes(departmentIds=str(self.sewing.id)), ["PD_C"])
        self.assertEqual(codes(employmentType="production"), [])

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"missing_punch": "edit"})

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, MARCH, ["PD_A", "PD_B", "PD_X"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for e in self.many(12):
                mp_req(e, 20, time(9, 0), "IN", "morning_in", "approved", utc(2026, 3, 21, 5, 0))

        for e in self.many(2, prefix="PD_S"):
            mp_req(e, 20, time(9, 0), "IN", "morning_in", "approved", utc(2026, 3, 21, 5, 0))
        self.assert_constant_queries(self.RID, MARCH, grow)


# ════════════════════════════════════════════════════════════════════════════
#  Absence around holidays
# ════════════════════════════════════════════════════════════════════════════

JAN = dict(dateFrom="2026-01-01", dateTo="2026-01-31")


def day_rec(emp, day, status, month=1):
    return AttendanceDayRecord.objects.create(employee=emp, date=date(2026, month, day), status=status)


class AbsenceAroundHolidaysTests(_Base):
    RID = "absence-around-holidays"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        Holiday.objects.create(name="Pongal", date=date(2026, 1, 14))
        Holiday.objects.create(name="Thiruvalluvar Day", date=date(2026, 1, 15))
        cls.sat_off = cls.mk("PD_SAT", "Sathya", cls.cutting, cls.b1)
        shift = ShiftTemplate.objects.create(name="PD Sat shift", shift_type="staff", start_time=time(9, 0), end_time=time(18, 0))
        EmployeeShiftAssignment.objects.create(employee=cls.sat_off, shift=shift, effective_from=date(2020, 1, 1), saturday_off=True)
        cls.joiner = cls.mk("PD_JN", "Jaya", cls.cutting, cls.b1, join_date="12/01/2026")  # dd/mm/yyyy, like real data
        rec = day_rec
        # PD_A: four sandwiches around Sundays 4, 11, 18 and the Pongal block (14-15).
        for day, status in ((3, "present"), (5, "absent"), (10, "absent"), (12, "absent"), (13, "absent"), (16, "absent"),
                            (17, "present"), (19, "on_leave")):
            rec(cls.a, day, status)
        # PD_SAT has Saturday off: an 'absent' Saturday record must never make the Sunday look sandwiched.
        for day, status in ((2, "absent"), (3, "absent"), (5, "present"), (9, "present"), (10, "absent"), (12, "present")):
            rec(cls.sat_off, day, status)
        # Production works Sundays: only the holidays are off days.
        for day, status in ((4, "absent"), (13, "absent"), (16, "present")):
            rec(cls.p, day, status)
        # Joined on 12 Jan: earlier 'absent' verdicts are the engine's artefact and must be masked.
        for day, status in ((3, "absent"), (5, "absent"), (10, "absent"), (13, "present"), (16, "present")):
            rec(cls.joiner, day, status)
        for day, status in ((13, "absent"), (16, "absent")):
            rec(cls.c, day, status)

    def rows(self, **params):
        body = self.run_report(self.RID, **{**JAN, **params})
        return body["rows"], body

    def test_golden_rows(self):
        rows, body = self.rows()
        got = [(r["employeeCode"], r["offDay"], r["offDayName"], r["offDays"], r["dayBefore"], r["dayAfter"], r["pattern"]) for r in rows]
        self.assertEqual(got, [
            ("PD_A", "2026-01-04", "Sunday", 1, "Present", "Absent", "After only"),
            ("PD_A", "2026-01-11", "Sunday", 1, "Absent", "Absent", "Before and after"),
            ("PD_A", "2026-01-14", "Pongal + Thiruvalluvar Day", 2, "Absent", "Absent", "Before and after"),
            ("PD_A", "2026-01-18", "Sunday", 1, "Present", "On leave", "After only"),
            ("PD_C", "2026-01-14", "Pongal + Thiruvalluvar Day", 2, "Absent", "Absent", "Before and after"),
            ("PD_P", "2026-01-14", "Pongal + Thiruvalluvar Day", 2, "Absent", "Present", "Before only"),
            ("PD_SAT", "2026-01-03", "Saturday (weekly off) + Sunday", 2, "Absent", "Present", "Before only"),
        ])
        first = rows[1]
        self.assertEqual((first["beforeDate"], first["afterDate"]), ("2026-01-10", "2026-01-12"))
        block = rows[2]
        self.assertEqual((block["beforeDate"], block["afterDate"]), ("2026-01-13", "2026-01-16"))
        s = self.summary(body)
        self.assertEqual((s["Sandwich absences"], s["Away before and after"], s["Employees affected"], s["Employees with 3 or more"]), (7, 3, 4, 1))

    def test_joining_date_masks_the_days_before_joining(self):
        rows, _ = self.rows()
        self.assertNotIn("PD_JN", [r["employeeCode"] for r in rows])

    def test_weekly_off_saturday_is_not_treated_as_a_working_day(self):
        rows, _ = self.rows(employeeIds=str(self.sat_off.id))
        self.assertEqual([r["offDay"] for r in rows], ["2026-01-03"])  # nothing around Sunday 11 Jan

    def test_filters(self):
        codes = lambda **kw: [(r["employeeCode"], r["offDay"]) for r in self.rows(**kw)[0]]  # noqa: E731
        self.assertEqual(len(codes(pattern="both")), 3)
        self.assertEqual(codes(pattern="before"), [("PD_P", "2026-01-14"), ("PD_SAT", "2026-01-03")])
        self.assertEqual(codes(pattern="after"), [("PD_A", "2026-01-04"), ("PD_A", "2026-01-18")])
        self.assertEqual(len(codes(includeLeave="false")), 6)  # the on-leave Monday no longer counts
        self.assertEqual(codes(employmentType="production"), [("PD_P", "2026-01-14")])
        self.assertEqual(codes(departmentIds=str(self.sewing.id)), [("PD_C", "2026-01-14")])
        self.assertEqual(codes(branchIds=str(self.b2.id)), [("PD_C", "2026-01-14")])
        self.assertEqual(codes(employeeIds=str(self.a.id), dateFrom="2026-01-15", dateTo="2026-01-15"), [("PD_A", "2026-01-14")])  # block found from its 2nd day
        self.assertEqual(codes(employeeIds=str(self.a.id), dateFrom="2026-01-12", dateTo="2026-01-13"), [])
        self.assertEqual(codes(dateFrom="2026-02-01", dateTo="2026-02-28"), [])

    def test_a_block_running_past_the_range_edge_is_found_whole(self):
        # Sunday 25 Jan is followed by a company holiday on Monday 26 Jan: one 2-day block, day after = Tuesday 27th.
        Holiday.objects.create(name="Republic Day", date=date(2026, 1, 26))
        day_rec(self.a, 23, "absent")
        day_rec(self.a, 27, "absent")
        day_rec(self.a, 24, "present")
        rows = [r for r in self.rows(employeeIds=str(self.a.id), dateFrom="2026-01-25", dateTo="2026-01-25")[0]]
        self.assertEqual([(r["offDay"], r["offDayName"], r["offDays"], r["dayAfter"]) for r in rows], [("2026-01-25", "Sunday + Republic Day", 2, "Absent")])

    def test_days_from_today_on_are_never_flagged(self):
        from .clock import ist_today

        today = ist_today()
        # A Sunday whose day-before is today: the record is 'in progress' and must be ignored.
        sunday = today + timedelta(days=(6 - today.weekday()) % 7 + 7)
        AttendanceDayRecord.objects.create(employee=self.a, date=sunday - timedelta(days=1), status="absent")
        AttendanceDayRecord.objects.create(employee=self.a, date=sunday + timedelta(days=1), status="absent")
        body = self.run_report(self.RID, dateFrom=sunday.isoformat(), dateTo=sunday.isoformat(), employeeIds=str(self.a.id))
        self.assertEqual(body["rows"], [])

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, JAN, "PD_A")
        self.check_gating(self.RID, JAN, {"attendance": "view"})

    def test_branch_isolation(self):
        self.check_branch_isolation(self.RID, JAN, ["PD_A", "PD_P", "PD_SAT"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for e in self.many(12):
                day_rec(e, 13, "absent")
                day_rec(e, 16, "absent")

        for e in self.many(2, prefix="PD_S"):
            day_rec(e, 13, "absent")
            day_rec(e, 16, "absent")
        self.assert_constant_queries(self.RID, JAN, grow)


# ════════════════════════════════════════════════════════════════════════════
#  Leave / request vs attendance conflicts
# ════════════════════════════════════════════════════════════════════════════


def log_punch(emp, day, *times):
    for t in times:
        AttendanceLog.objects.create(employee=emp, date=date(2026, 3, day), punch_time=t, punch_type="IN", source="biometric:test")


class ConflictTests(_Base):
    RID = "leave-attendance-conflicts"

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()

        punch = log_punch
        casual = LeaveType.objects.create(name="Casual Leave", code="PDCL")
        # 1. leave but punched: PD_A approved 16-18 Mar, punched on the 17th only.
        LeaveRequest.objects.create(employee=cls.a, leave_type_ref=casual, type="casual", start_date="2026-03-16", end_date="2026-03-18", status="approved")
        punch(cls.a, 17, time(9, 0), time(18, 0))
        LeaveRequest.objects.create(employee=cls.b, type="casual", start_date="2026-03-16", end_date="2026-03-16", status="approved", is_half_day=True, half_day_slot="morning", total_days=Decimal("0.5"))
        punch(cls.b, 16, time(9, 0))  # half-day leave: not a conflict
        LeaveRequest.objects.create(employee=cls.b, type="casual", start_date="2026-03-17", end_date="2026-03-17", status="pending")
        punch(cls.b, 17, time(9, 0))  # pending leave: not a conflict
        LeaveRequest.objects.create(employee=cls.p, type="casual", start_date="2026-03-17", end_date="2026-03-17", status="approved")
        punch(cls.p, 17, time(9, 0))  # production: leave has no effect
        LeaveRequest.objects.create(employee=cls.x, type="casual", start_date="soon", end_date="later", status="approved")  # unparseable
        AttendanceDayRecord.objects.create(employee=cls.a, date=date(2026, 3, 17), status="present")
        # 2. casual leave rejected but worked.
        CasualLeaveRequest.objects.create(employee=cls.a, date=date(2026, 3, 19), status="rejected")
        punch(cls.a, 19, time(9, 0), time(17, 0), time(17, 30))
        CasualLeaveRequest.objects.create(employee=cls.b, date=date(2026, 3, 19), status="approved")
        punch(cls.b, 19, time(9, 0))
        CasualLeaveRequest.objects.create(employee=cls.b, date=date(2026, 3, 20), status="rejected")  # no punches -> fine
        AttendanceDayRecord.objects.create(employee=cls.a, date=date(2026, 3, 19), status="on_leave", source="manual")
        # 3. on-duty approved without any approved punch.
        cls.od_bad = OnDutySession.objects.create(employee=cls.a, destination="Karur", status="active", branch=cls.b1)
        OnDutySession.objects.filter(pk=cls.od_bad.pk).update(created_at=utc(2026, 3, 21, 4, 0))
        ok = OnDutySession.objects.create(employee=cls.b, destination="Trichy", status="completed", branch=cls.b1)
        OnDutySession.objects.filter(pk=ok.pk).update(created_at=utc(2026, 3, 21, 4, 0))
        OnDutyPunchVerification.objects.create(
            session=ok, employee=cls.b, punch_date=date(2026, 3, 21), punch_time=time(9, 0), punch_type="IN", punch_number=1,
            latitude=Decimal("10.1"), longitude=Decimal("77.1"), photo="x.jpg", status="approved",
        )
        pend = OnDutySession.objects.create(employee=cls.a, destination="Pending place", status="pending_hr", branch=cls.b1)
        OnDutySession.objects.filter(pk=pend.pk).update(created_at=utc(2026, 3, 21, 4, 0))
        # 4. approved missing punch that is not in the log.
        MissingPunchRequest.objects.create(employee=cls.a, date=date(2026, 3, 23), punch_time=time(9, 0), punch_type="IN", status="approved", reason="r")
        MissingPunchRequest.objects.create(employee=cls.b, date=date(2026, 3, 23), punch_time=time(9, 0), punch_type="IN", status="approved", reason="r")
        punch(cls.b, 23, time(9, 0))
        MissingPunchRequest.objects.create(employee=cls.a, date=date(2026, 3, 24), punch_time=time(9, 0), punch_type="IN", status="pending_hr", reason="r")
        # 5. approved permission not reflected on a worked, auto-computed day.
        for day, typ in ((24, "morning_late_in"), (25, "evening_early_out"), (26, "middle_permission")):
            EmployeePermission.objects.create(employee=cls.a, date=date(2026, 3, day), type=typ, status="approved")
            AttendanceDayRecord.objects.create(employee=cls.a, date=date(2026, 3, day), status="present", source="auto")
        EmployeePermission.objects.create(employee=cls.b, date=date(2026, 3, 24), type="morning_late_in", status="approved")
        AttendanceDayRecord.objects.create(employee=cls.b, date=date(2026, 3, 24), status="present", morning_permission_applied=True)
        EmployeePermission.objects.create(employee=cls.b, date=date(2026, 3, 25), type="morning_late_in", status="approved")  # no record
        EmployeePermission.objects.create(employee=cls.b, date=date(2026, 3, 26), type="middle_permission", status="approved")
        AttendanceDayRecord.objects.create(employee=cls.b, date=date(2026, 3, 26), status="present", source="manual")
        EmployeePermission.objects.create(employee=cls.b, date=date(2026, 3, 27), type=None, status="approved")  # untyped: cannot verify
        AttendanceDayRecord.objects.create(employee=cls.b, date=date(2026, 3, 27), status="present", source="auto")
        EmployeePermission.objects.create(employee=cls.p, date=date(2026, 3, 24), type="morning_late_in", status="approved")  # production
        AttendanceDayRecord.objects.create(employee=cls.p, date=date(2026, 3, 24), status="present", source="auto")

    def rows(self, user=None, **params):
        body = self.run_report(self.RID, user, **{**MARCH, **params})
        return body["rows"], body

    def test_golden_rows(self):
        rows, body = self.rows()
        got = [(r["employeeCode"], r["date"], r["conflictType"], r["requestRef"].split(" #")[0], r["attendanceStatus"]) for r in rows]
        self.assertEqual(got, [
            ("PD_A", "2026-03-17", "Approved leave, but punched", "Leave", "Present"),
            ("PD_A", "2026-03-19", "Casual leave rejected, but worked", "Casual leave", "On leave"),
            ("PD_A", "2026-03-21", "On-duty approved, no approved punches", "On-duty", None),
            ("PD_A", "2026-03-23", "Missing punch approved, not in the log", "Missing punch", None),
            ("PD_A", "2026-03-24", "Permission approved, not reflected", "Permission", "Present"),
            ("PD_A", "2026-03-25", "Permission approved, not reflected", "Permission", "Present"),
            ("PD_A", "2026-03-26", "Permission approved, not reflected", "Permission", "Present"),
        ])
        self.assertIn("Approved Casual Leave (16-Mar-2026 to 18-Mar-2026) but 2 punch(es)", rows[0]["detail"])
        self.assertIn("3 punch(es)", rows[1]["detail"])
        self.assertIn("Karur", rows[2]["detail"])
        self.assertIn("09:00", rows[3]["detail"])
        self.assertIn("possibly stale", rows[4]["detail"])
        s = self.summary(body)
        self.assertEqual((s["Conflicts"], s["Employees affected"]), (7, 1))
        for label in checks.CONFLICT_LABELS.values():
            self.assertIn(label, s)
        self.assertEqual(s["Permission approved, not reflected"], 3)
        self.assertEqual(s["Approved leave, but punched"], 1)

    def test_conflict_type_filter(self):
        only = lambda t: [(r["employeeCode"], r["date"]) for r in self.rows(conflictType=t)[0]]  # noqa: E731
        self.assertEqual(only("leave_but_punched"), [("PD_A", "2026-03-17")])
        self.assertEqual(only("cl_rejected_but_worked"), [("PD_A", "2026-03-19")])
        self.assertEqual(only("on_duty_no_punches"), [("PD_A", "2026-03-21")])
        self.assertEqual(only("missing_punch_not_written"), [("PD_A", "2026-03-23")])
        self.assertEqual(len(only("permission_no_effect")), 3)

    def test_other_filters(self):
        codes = lambda **kw: [r["employeeCode"] for r in self.rows(**kw)[0]]  # noqa: E731
        self.assertEqual(len(codes(employeeIds=str(self.a.id))), 7)
        self.assertEqual(codes(employeeIds=str(self.b.id)), [])
        self.assertEqual(codes(departmentIds=str(self.finishing.id)), [])
        self.assertEqual(len(codes(employmentType="staff")), 7)
        self.assertEqual(codes(employmentType="production"), [])
        self.assertEqual(codes(employeeStatus="inactive"), [])
        self.assertEqual(len(codes(dateFrom="2026-03-19", dateTo="2026-03-21")), 2)
        self.assertEqual(codes(branchIds=str(self.b2.id)), [])

    def test_a_role_only_sees_conflicts_of_requests_it_can_open(self):
        leave_only = self.hr("pd_leave_only", {"reports": "view", "attendance": "view", "leave": "view"})
        rows, body = self.rows(user=leave_only)
        self.assertEqual([r["conflictType"] for r in rows], ["Approved leave, but punched"])
        self.assertIn("Not shown, because your role cannot open the underlying requests", " ".join(body["notes"]))
        self.assertIn("Permission approved, not reflected", " ".join(body["notes"]))
        forbidden = self.rows(user=leave_only, conflictType="on_duty_no_punches")[0]
        self.assertEqual(forbidden, [])  # asking for a type the role cannot open returns nothing, not an error
        full = self.hr("pd_all", {"reports": "view", **{k: "view" for k in ("attendance", "leave", "casual_leave", "geo_attendance", "missing_punch", "requests")}})
        self.assertEqual(len(self.rows(user=full)[0]), 7)

    def test_totals_and_summary_agree_with_the_rows(self):
        rows, body = self.rows()
        self.assertEqual(sum(v for k, v in self.summary(body).items() if k in checks.CONFLICT_LABELS.values()), len(rows))
        self.assertIsNone(body["totals"])

    def test_exports_empty_and_gating(self):
        self.check_exports(self.RID, MARCH, "PD_A")
        self.check_gating(self.RID, MARCH, {"attendance": "view"})

    def test_branch_isolation(self):
        # Put a conflict on PD_C (branch 2): the branch-1 user must never see it.
        CasualLeaveRequest.objects.create(employee=self.c, date=date(2026, 3, 19), status="rejected")
        log_punch(self.c, 19, time(9, 0))
        rows, _ = self.rows()
        self.assertIn("PD_C", [r["employeeCode"] for r in rows])
        self.check_branch_isolation(self.RID, MARCH, ["PD_A"], "PD_C", self.c.id)

    def test_query_count_is_flat(self):
        def grow():
            for e in self.many(12):
                CasualLeaveRequest.objects.create(employee=e, date=date(2026, 3, 19), status="rejected")
                log_punch(e, 19, time(9, 0))
                MissingPunchRequest.objects.create(employee=e, date=date(2026, 3, 23), punch_time=time(9, 0), punch_type="IN", status="approved", reason="r")

        for e in self.many(2, prefix="PD_S"):
            CasualLeaveRequest.objects.create(employee=e, date=date(2026, 3, 19), status="rejected")
            log_punch(e, 19, time(9, 0))
        self.assert_constant_queries(self.RID, MARCH, grow)
