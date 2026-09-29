"""
Report Center - Gate & Visitors group G9: visitor and tea-break reports, employee time away, device register.

Run via: python manage.py test api.tests_reporting_gate_visitor_tea -v 2

All dates are fixed and the "current instant" the reports use for open tea breaks is frozen (NOW below), so nothing
depends on the day the suite runs. Fixture map (times are IST unless stated):

Employees (all "T00x"): T001 Anita Kumar CUTTING/Operator/Unit 1 staff; T002 Bala Raj CUTTING/Helper/Unit 1 production;
T003 Chitra Devi SEWING/Operator/Unit 1 staff; T004 Dinesh Babu PACKING/Operator/Unit 2 staff; T005 Esther Mary (resigned,
no department/designation) Unit 1 production; T006 Farook Ali CUTTING/Helper, no branch. Tea allowance = 15 min.

Tea breaks (A..M), out -> in  (gate out/in):
  A T001 09-05 10:00 -> 10:10 (G1/G1)   10 min  on time
  B T001 09-05 15:00:00 -> 15:15:30     15.5 -> 16 (half-even) overtime, 1 over
  C T002 09-05 10:05 -> 10:20           15 min  on time
  D T002 09-05 15:10:00 -> 15:26:30     16.5 -> 16 (half-even, SQL ROUND would say 17) overtime, 1 over  (G1/G2)
  E T001 09-06 23:50 -> 09-07 00:10     20 min  overtime, 5 over, crosses midnight                        (G1/G2)
  F T003 09-07 01:30 -> 01:40           10 min  on time  (UTC date is 09-06!)                             (G2/G2)
  G T003 09-08 09:00 -> 11:30           150 min overtime, 135 over, "Long break"                          (G2/G2)
  H T004 09-05 10:00 -> 10:12           12 min  on time  (Unit 2)                                         (G3/G3)
  I T004 09-05 16:00 -> 16:30           30 min  overtime, 15 over                                         (G3/G3)
  J T002 open since 09-15 13:00         150 min open -> not returned                                      (G1)
  K T004 open since 09-15 15:00         30 min open -> in progress                                        (G3)
  L T001 open since 09-14 20:00         19.5 h open -> not returned + "Left open"                         (G1)
  M T005 09-05 11:00 -> 11:14           14 min  on time  (resigned employee)                              (G2/G2)
NOW = 15-Sep-2026 15:30 IST.
Totals (all): 13 breaks, 10 completed, 5 on time, 5 overtime, 2 not returned, 1 in progress, 2 long/left-open,
293 minutes, 157 minutes over allowance, longest 150.

Visits (visitor, branch, time, host):
  v1 Ravi Kumar  Unit 1 09-02 10:00 host T001 (linked)   email +4s, WhatsApp +40s
  v2 Ravi Kumar  Unit 1 09-09 02:30 free text "anita kumar "  (UTC date is 09-08!)
  v3 Sita Devi   Unit 1 09-09 11:00 host T002 (linked)   email +10s
  v4 Mohan       Unit 2 09-09 15:00 free text "ANITA KUMAR"
  v5 Sita Devi   Unit 1 09-10 09:00 host T001 (linked)   WhatsApp +6s
  v6 Ravi Kumar  Unit 2 09-12 12:00 host T004 (linked)   nothing sent
  v7 Mohan       no branch 09-12 13:00 free text "Mr X"
"""

import io
import json
from collections import Counter
from datetime import date, datetime, timedelta
from datetime import timezone as dt_timezone
from unittest import mock

from django.db import connection
from django.db.models import Max, Sum
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .clock import FACTORY_TZ
from .jwt_utils import sign_token
from .models import (
    AuditLog, Branch, Department, Designation, Employee, GateDevice, HRUser, OutpassGateScan, OutpassRecord,
    OutpassRequest, ReceptionDevice, Role, TeaBreakLog, TeaBreakRule, Visitor, VisitorVisit,
)
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import gate_visitor_tea_common as C
from .reporting.runner import run_report
from .tea_break_views import _log_json

NOW = datetime(2026, 9, 15, 10, 0, tzinfo=dt_timezone.utc)  # 15:30 IST on 15 Sep 2026
SEPT = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30"}
FULL_AADHAAR_1 = "123456789012"
FULL_AADHAAR_2 = "987654321098"

MY_IDS = [
    "visitor-register", "visitor-host-summary", "visitor-frequency", "visitor-daily-summary", "visitor-notification-delivery",
    "tea-break-register", "tea-break-employee-summary", "tea-break-department-summary", "tea-break-exceptions",
    "tea-break-daily-trend", "tea-break-hourly-distribution", "tea-break-frequency", "employee-time-away-summary",
    "gate-device-register",
]


def ist(month, day, hh=0, mm=0, ss=0):
    return datetime(2026, month, day, hh, mm, ss, tzinfo=FACTORY_TZ)


def headers(user):
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


def summary_of(body):
    return {s["label"]: s["value"] for s in body["summary"]}


class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.b1 = Branch.objects.create(name="Unit 1", code="GU1")
        cls.b2 = Branch.objects.create(name="Unit 2", code="GU2")
        cls.b_empty = Branch.objects.create(name="Unit 3", code="GU3")
        cls.d_cut = Department.objects.create(name="CUTTING", branch=cls.b1)
        cls.d_sew = Department.objects.create(name="SEWING", branch=cls.b1)
        cls.d_pack = Department.objects.create(name="PACKING", branch=cls.b2)
        cls.des_op = Designation.objects.create(title="Operator")
        cls.des_help = Designation.objects.create(title="Helper")

        def emp(code, first, last, dept, desig, branch, etype, status="active", **kw):
            return Employee.objects.create(
                employee_code=code, first_name=first, last_name=last, department=dept, designation=desig,
                branch=branch, employment_type=etype, status=status, **kw,
            )

        cls.e1 = emp("T001", "Anita", "Kumar", cls.d_cut, cls.des_op, cls.b1, "staff", email="anita@example.com", phone="9111111111")
        cls.e2 = emp("T002", "Bala", "Raj", cls.d_cut, cls.des_help, cls.b1, "production", phone="9222222222")
        cls.e3 = emp("T003", "Chitra", "Devi", cls.d_sew, cls.des_op, cls.b1, "staff")
        cls.e4 = emp("T004", "Dinesh", "Babu", cls.d_pack, cls.des_op, cls.b2, "staff", email="")
        cls.e5 = emp("T005", "Esther", "Mary", None, None, cls.b1, "production", status="resigned")
        cls.e6 = emp("T006", "Farook", "Ali", cls.d_cut, cls.des_help, None, "staff")

        role_ok = Role.objects.create(name="gvt_ok", permissions={"reports": "view", "outpass_visitors": "view"})
        role_plain = Role.objects.create(name="gvt_plain", permissions={"reports": "view"})
        mk = lambda name, **kw: HRUser.objects.create(username=name, password_hash="x", **kw)  # noqa: E731
        cls.admin = mk("gvt_admin", is_super_admin=True)
        cls.b1_user = mk("gvt_b1", role=role_ok, branch=cls.b1)
        cls.b2_user = mk("gvt_b2", role=role_ok, branch=cls.b2)
        cls.open_user = mk("gvt_open", role=role_ok)  # role has the module, no branch -> unscoped
        cls.plain_user = mk("gvt_plain", role=role_plain)

        rule = TeaBreakRule.objects.create(pk=1, allowed_minutes=15)
        TeaBreakRule.objects.filter(pk=1).update(updated_at=ist(9, 3, 12, 0))
        cls.rule = rule

        # ── gates / desks
        def gate(name, branch, user, active=True, last=None):
            return GateDevice.objects.create(
                name=name, branch=branch, username=user, password_hash=f"hash-secret-{user}", login_token=f"tok-{user}",
                is_active=active, created_by="Admin", last_login_at=last,
            )

        cls.g1 = gate("Gate 1", cls.b1, "gate1", True, ist(9, 10, 13, 30))
        cls.g2 = gate("Gate 2", cls.b1, "gate2", False)
        cls.g3 = gate("Gate 3", cls.b2, "gate3", True, ist(9, 1, 9, 0))
        cls.r1 = ReceptionDevice.objects.create(
            name="Reception 1", branch=cls.b1, username="rec1", password_hash="hash-secret-rec1", login_token="tok-rec1",
            is_active=True, created_by="Admin", last_login_at=ist(9, 11, 9, 0),
        )
        cls.r2 = ReceptionDevice.objects.create(
            name="Reception 2", branch=cls.b2, username="rec2", password_hash="hash-secret-rec2", login_token="tok-rec2",
            is_active=False, created_by="Admin",
        )

        # ── tea breaks
        def tea(emp_, out, back=None, og=None, ig=None):
            return TeaBreakLog.objects.create(employee=emp_, out_gate=og, out_at=out, in_gate=ig, in_at=back)

        g1, g2, g3 = cls.g1, cls.g2, cls.g3
        cls.tA = tea(cls.e1, ist(9, 5, 10, 0), ist(9, 5, 10, 10), g1, g1)
        cls.tB = tea(cls.e1, ist(9, 5, 15, 0), ist(9, 5, 15, 15, 30), g1, g1)
        cls.tC = tea(cls.e2, ist(9, 5, 10, 5), ist(9, 5, 10, 20), g1, g1)
        cls.tD = tea(cls.e2, ist(9, 5, 15, 10), ist(9, 5, 15, 26, 30), g1, g2)
        cls.tE = tea(cls.e1, ist(9, 6, 23, 50), ist(9, 7, 0, 10), g1, g2)
        cls.tF = tea(cls.e3, ist(9, 7, 1, 30), ist(9, 7, 1, 40), g2, g2)
        cls.tG = tea(cls.e3, ist(9, 8, 9, 0), ist(9, 8, 11, 30), g2, g2)
        cls.tH = tea(cls.e4, ist(9, 5, 10, 0), ist(9, 5, 10, 12), g3, g3)
        cls.tI = tea(cls.e4, ist(9, 5, 16, 0), ist(9, 5, 16, 30), g3, g3)
        cls.tJ = tea(cls.e2, ist(9, 15, 13, 0), None, g1)
        cls.tK = tea(cls.e4, ist(9, 15, 15, 0), None, g3)
        cls.tL = tea(cls.e1, ist(9, 14, 20, 0), None, g1)
        cls.tM = tea(cls.e5, ist(9, 5, 11, 0), ist(9, 5, 11, 14), g2, g2)

        # ── visitors
        cls.v1 = Visitor.objects.create(name="Ravi Kumar", phone="9000000001", aadhaar_number=FULL_AADHAAR_1)
        cls.v2 = Visitor.objects.create(name="Sita Devi", phone="9000000002", aadhaar_number=FULL_AADHAAR_2)
        cls.v3 = Visitor.objects.create(name="Mohan", phone="9000000003", aadhaar_number=None)

        def visit(visitor, branch, at, whom, purpose, why=None, host=None, email=None, wa=None):
            v = VisitorVisit.objects.create(
                visitor=visitor, branch=branch, why_came=why, whom_to_meet=whom, purpose=purpose, meeting_employee=host,
                notified_email_at=(at + timedelta(seconds=email)) if email is not None else None,
                notified_whatsapp_at=(at + timedelta(seconds=wa)) if wa is not None else None,
            )
            VisitorVisit.objects.filter(pk=v.pk).update(visited_at=at)  # auto_now_add ignores a value passed to create()
            return v

        cls.vis1 = visit(cls.v1, cls.b1, ist(9, 2, 10, 0), "Anita Kumar", "Fabric samples", "Supplier", cls.e1, 4, 40)
        cls.vis2 = visit(cls.v1, cls.b1, ist(9, 9, 2, 30), "anita kumar ", "Follow up")
        cls.vis3 = visit(cls.v2, cls.b1, ist(9, 9, 11, 0), "Bala Raj", "Interview", "Interview", cls.e2, 10, None)
        cls.vis4 = visit(cls.v3, cls.b2, ist(9, 9, 15, 0), "ANITA KUMAR", "Parcel", "Delivery")
        cls.vis5 = visit(cls.v2, cls.b1, ist(9, 10, 9, 0), "Anita Kumar", "Meeting", None, cls.e1, None, 6)
        cls.vis6 = visit(cls.v1, cls.b2, ist(9, 12, 12, 0), "Dinesh Babu", "Audit", None, cls.e4)
        cls.vis7 = visit(cls.v3, None, ist(9, 12, 13, 0), "Mr X", "Vendor")

        # ── outpass (for employee-time-away-summary and the device scan counts)
        def op(emp_, out, back=None, status="approved"):
            return OutpassRequest.objects.create(
                employee=emp_, destination="Bank", reason="Personal", status=status, approved_at=out - timedelta(minutes=10),
                exit_gate=cls.g1, exited_at=out, entry_gate=cls.g1 if back else None, entered_at=back,
            )

        op(cls.e1, ist(9, 5, 11, 0), ist(9, 5, 12, 30))  # 90 min
        op(cls.e1, ist(9, 6, 14, 0), ist(9, 6, 14, 46, 30))  # 46.5 -> 46 (half-even)
        op(cls.e1, ist(9, 7, 9, 0))  # exited, never returned: counted as an exit, no minutes
        op(cls.e2, ist(9, 8, 10, 0), ist(9, 8, 10, 20))  # 20 min
        op(cls.e4, ist(9, 5, 10, 0), ist(9, 5, 10, 30))  # 30 min (Unit 2)
        OutpassRequest.objects.create(employee=cls.e3, destination="Bank", reason="x", status="pending")  # never exited

        def qr(emp_, branch, at, source="qr", name="x", code="x"):
            rec = OutpassRecord.objects.create(
                branch=branch, employee=emp_, employee_name=name, employee_code=code, destination="Market", source=source,
            )
            OutpassRecord.objects.filter(pk=rec.pk).update(submitted_at=at)

        qr(cls.e1, cls.b1, ist(9, 5, 12, 0))
        qr(cls.e1, cls.b1, ist(9, 6, 12, 0))
        qr(None, cls.b1, ist(9, 6, 13, 0), name="Stranger", code="ZZ99")  # unmatched: not attributable to an employee
        qr(cls.e4, cls.b2, ist(9, 5, 12, 0))
        qr(cls.e1, cls.b1, ist(9, 5, 13, 0), source="request")  # mirror row of an approved pass: never an exit

        def scan(gate_, scan_type, result, at):
            s = OutpassGateScan.objects.create(gate=gate_, scan_type=scan_type, result=result, message="m")
            OutpassGateScan.objects.filter(pk=s.pk).update(scanned_at=at)

        scan(cls.g1, "exit", "success", ist(9, 5, 11, 0))
        scan(cls.g1, "entry", "success", ist(9, 5, 12, 30))
        scan(cls.g1, "exit", "expired", ist(9, 6, 15, 0))
        scan(cls.g1, "exit", "success", ist(8, 30, 10, 0))  # before the period
        scan(cls.g3, "exit", "not_approved", ist(9, 5, 10, 0))

    def setUp(self):
        patcher = mock.patch.object(C, "now_utc", return_value=NOW)
        patcher.start()
        self.addCleanup(patcher.stop)

    # ── helpers
    def get_report(self, rid, params=None, user=None, expect=200):
        r = self.client.get(f"/api/reports/run/{rid}", params or {}, **headers(user or self.admin))
        self.assertEqual(r.status_code, expect, r.content[:400])
        return r.json()

    def export(self, rid, fmt, params=None, user=None):
        return self.client.get(f"/api/reports/export/{rid}", {"fmt": fmt, **(params or {})}, **headers(user or self.admin))

    def rows(self, rid, params=None, user=None):
        return self.get_report(rid, params, user)["rows"]


# ═════════════════════════════════════════════════════════════════════════════
# Visitors
# ═════════════════════════════════════════════════════════════════════════════


class VisitorRegisterTests(_Base):
    RID = "visitor-register"

    def test_golden_rows_summary_and_notes(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual(
            [r["visitedAt"] for r in rows],
            ["2026-09-02 10:00", "2026-09-09 02:30", "2026-09-09 11:00", "2026-09-09 15:00",
             "2026-09-10 09:00", "2026-09-12 12:00", "2026-09-12 13:00"],
        )
        self.assertEqual([r["visitNo"] for r in rows], [1, 2, 1, 1, 2, 3, 2])
        self.assertEqual([r["visitorType"] for r in rows], ["First visit", "Repeat", "First visit", "First visit", "Repeat", "Repeat", "Repeat"])
        r0 = rows[0]
        self.assertEqual(r0["visitorName"], "Ravi Kumar")
        self.assertEqual(r0["phone"], "9000000001")
        self.assertEqual(r0["aadhaarMasked"], "XXXX XXXX 9012")
        self.assertEqual((r0["whyCame"], r0["whomToMeet"], r0["purpose"], r0["branch"]), ("Supplier", "Anita Kumar", "Fabric samples", "Unit 1"))
        self.assertEqual((r0["hostCode"], r0["hostName"], r0["hostDepartment"]), ("T001", "Anita Kumar", "CUTTING"))
        self.assertEqual((r0["emailNotified"], r0["whatsappNotified"]), ("Sent", "Sent"))
        # free-text host: no employee cells and notification does not apply (dash), never "Not sent"
        r1 = rows[1]
        self.assertEqual((r1["hostCode"], r1["hostName"], r1["hostDepartment"]), (None, None, None))
        self.assertEqual((r1["emailNotified"], r1["whatsappNotified"]), (None, None))
        self.assertIsNone(r1["whyCame"])
        # linked host, nothing delivered -> "Not sent"; only one channel -> mixed
        self.assertEqual((rows[5]["emailNotified"], rows[5]["whatsappNotified"]), ("Not sent", "Not sent"))
        self.assertEqual((rows[2]["emailNotified"], rows[2]["whatsappNotified"]), ("Sent", "Not sent"))
        self.assertEqual(rows[6]["branch"], None)  # visit whose branch link is gone
        self.assertEqual(rows[4]["aadhaarMasked"], "XXXX XXXX 1098")
        self.assertIsNone(rows[3]["aadhaarMasked"])  # visitor with no Aadhaar on file

        s = summary_of(body)
        self.assertEqual(s["Visits"], 7)
        self.assertEqual(s["Unique visitors"], 3)
        self.assertEqual(s["First-time visits"], 3)
        self.assertEqual(s["Repeat visits"], 4)
        self.assertEqual(s["With a linked host employee"], 4)
        self.assertEqual(s["Host emailed"], 2)
        self.assertEqual(s["Host sent WhatsApp"], 2)
        self.assertEqual(s["Visits"], len(rows))
        notes = " ".join(body["notes"])
        self.assertIn("Visits by branch: Unit 1 4; Unit 2 2; No branch 1.", notes)
        self.assertIn("check-out", notes)

    def test_aadhaar_is_never_emitted_in_full(self):
        body = self.get_report(self.RID, SEPT)
        text = json.dumps(body)
        for full in (FULL_AADHAAR_1, FULL_AADHAAR_2):
            self.assertNotIn(full, text)
            self.assertNotIn(full[:8], text)  # not even a long prefix
        self.assertIn("XXXX XXXX 9012", text)
        # everything the PDF prints comes from RunOutput; assert on that too (rows, summary cards, notes, filters)
        class _Request:  # what run_report reads from a request: the branch scope and the HR identity
            hr_branch_id = None
            jwt_user = {"hrUserId": self.admin.id}

        req = _Request()
        for purpose in ("pdf", "xlsx"):
            out = run_report(req, registry.get_spec(self.RID), SEPT, purpose=purpose)
            blob = json.dumps(out.payload())
            self.assertNotIn(FULL_AADHAAR_1, blob)
            self.assertNotIn(FULL_AADHAAR_2, blob)
        r = self.export(self.RID, "xlsx", SEPT)
        ws = load_workbook(io.BytesIO(r.content)).active
        cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertFalse([c for c in cells if FULL_AADHAAR_1 in c or FULL_AADHAAR_2 in c])
        self.assertIn("XXXX XXXX 9012", cells)
        pdf = self.export(self.RID, "pdf", SEPT)
        self.assertEqual(pdf.status_code, 200)
        self.assertNotIn(FULL_AADHAAR_1.encode(), pdf.content)

    def test_short_or_untrimmed_aadhaar_is_handled(self):
        Visitor.objects.filter(pk=self.v3.pk).update(aadhaar_number="  4444 5555 6666  ")
        Visitor.objects.filter(pk=self.v2.pk).update(aadhaar_number="12")
        rows = self.rows(self.RID, SEPT)
        self.assertEqual(rows[3]["aadhaarMasked"], "XXXX XXXX 6666")
        self.assertIsNone(rows[2]["aadhaarMasked"])  # fewer than four digits: nothing is guessed

    def test_ist_day_boundaries(self):
        one_day = lambda d: self.rows(self.RID, {"dateFrom": d, "dateTo": d})  # noqa: E731
        # v2 checked in at 02:30 IST on 09-09 (21:00 UTC on 09-08): it belongs to 09-09
        self.assertEqual([r["visitedAt"] for r in one_day("2026-09-09")], ["2026-09-09 02:30", "2026-09-09 11:00", "2026-09-09 15:00"])
        self.assertEqual(one_day("2026-09-08"), [])

    def test_filters_narrow(self):
        def at(**p):
            return [r["visitedAt"][5:] for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(at(branchIds=self.b2.id), ["09-09 15:00", "09-12 12:00"])
        self.assertEqual(at(hostLinked="linked_employee"), ["09-02 10:00", "09-09 11:00", "09-10 09:00", "09-12 12:00"])
        self.assertEqual(at(hostLinked="free_text"), ["09-09 02:30", "09-09 15:00", "09-12 13:00"])
        self.assertEqual(at(departmentIds=self.d_cut.id), ["09-02 10:00", "09-09 11:00", "09-10 09:00"])
        self.assertEqual(at(employeeIds=self.e1.id), ["09-02 10:00", "09-10 09:00"])
        self.assertEqual(at(visitorType="first_visit"), ["09-02 10:00", "09-09 11:00", "09-09 15:00"])
        self.assertEqual(at(visitorType="repeat"), ["09-09 02:30", "09-10 09:00", "09-12 12:00", "09-12 13:00"])
        self.assertEqual(at(notification="email_sent"), ["09-02 10:00", "09-09 11:00"])
        self.assertEqual(at(notification="whatsapp_sent"), ["09-02 10:00", "09-10 09:00"])
        self.assertEqual(at(notification="either"), ["09-02 10:00", "09-09 11:00", "09-10 09:00"])
        self.assertEqual(at(notification="neither"), ["09-09 02:30", "09-09 15:00", "09-12 12:00", "09-12 13:00"])
        self.assertEqual(at(q="0002"), ["09-09 11:00", "09-10 09:00"])
        self.assertEqual(at(q="mohan"), ["09-09 15:00", "09-12 13:00"])
        self.assertEqual(at(dateFrom="2026-09-10", dateTo="2026-09-12"), ["09-10 09:00", "09-12 12:00", "09-12 13:00"])
        # a first-visit filter must not renumber: v5 is Sita's 2nd visit even when v3 is filtered out
        self.assertEqual([r["visitNo"] for r in self.rows(self.RID, {**SEPT, "employeeIds": self.e1.id})], [1, 2])
        # a filtered summary follows the same filter as the rows
        s = summary_of(self.get_report(self.RID, {**SEPT, "hostLinked": "free_text"}))
        self.assertEqual((s["Visits"], s["With a linked host employee"]), (3, 0))

    def test_branch_isolation_and_visit_numbers_are_scoped(self):
        b1 = self.get_report(self.RID, SEPT, self.b1_user)
        self.assertEqual([r["visitedAt"][5:] for r in b1["rows"]], ["09-02 10:00", "09-09 02:30", "09-09 11:00", "09-10 09:00"])
        self.assertEqual([r["visitNo"] for r in b1["rows"]], [1, 2, 1, 2])
        self.assertEqual(summary_of(b1)["Visits"], 4)
        self.assertEqual(summary_of(b1)["Unique visitors"], 2)
        # cannot be widened with filters
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual(self.rows(self.RID, {**SEPT, "employeeIds": self.e4.id}, self.b1_user), [])
        b2 = self.get_report(self.RID, SEPT, self.b2_user)
        self.assertEqual([r["visitedAt"][5:] for r in b2["rows"]], ["09-09 15:00", "09-12 12:00"])
        # Ravi's third visit overall is his FIRST one at Unit 2: a Unit 2 user must not learn about other branches' visits
        self.assertEqual([r["visitNo"] for r in b2["rows"]], [1, 1])
        # the visit with no branch is visible to unscoped users only
        self.assertEqual(len(self.rows(self.RID, SEPT, self.open_user)), 7)

    def test_totals_none_and_sum_of_rows(self):
        body = self.get_report(self.RID, SEPT)
        self.assertIsNone(body["totals"])


class VisitorHostSummaryTests(_Base):
    RID = "visitor-host-summary"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual(len(rows), 5)
        key = lambda r: (r["hostType"], r["hostCode"], r["hostName"].lower())  # noqa: E731
        self.assertEqual(
            [key(r) for r in rows],
            [("Employee", "T001", "anita kumar"), ("Free text", None, "anita kumar"), ("Employee", "T002", "bala raj"),
             ("Employee", "T004", "dinesh babu"), ("Free text", None, "mr x")],
        )
        e1 = rows[0]
        self.assertEqual((e1["department"], e1["visits"], e1["uniqueVisitors"]), ("CUTTING", 2, 2))
        self.assertEqual((e1["firstVisitAt"], e1["lastVisitAt"]), ("2026-09-02 10:00", "2026-09-10 09:00"))
        self.assertEqual((e1["emailSent"], e1["whatsappSent"]), (1, 2))
        free = rows[1]  # "anita kumar " (Unit 1) and "ANITA KUMAR" (Unit 2) are one free-text host
        self.assertEqual((free["department"], free["visits"], free["uniqueVisitors"]), (None, 2, 2))
        self.assertEqual((free["emailSent"], free["whatsappSent"]), (0, 0))
        e2 = rows[2]
        self.assertEqual((e2["visits"], e2["uniqueVisitors"], e2["emailSent"], e2["whatsappSent"]), (1, 1, 1, 0))
        self.assertEqual(body["totals"]["visits"], 7)
        self.assertEqual(body["totals"]["emailSent"], 2)
        self.assertEqual(body["totals"]["whatsappSent"], 2)
        self.assertIsNone(body["totals"].get("uniqueVisitors"))  # not additive across hosts
        s = summary_of(body)
        self.assertEqual((s["Visits"], s["Hosts"]), (7, 5))
        self.assertEqual(s["Visits"], sum(r["visits"] for r in rows))
        self.assertEqual(s["Visits with no linked employee"], round(100 * 3 / 7, 1))
        self.assertTrue(s["Most visited host"].endswith("(2)"))

    def test_no_phone_or_aadhaar_columns(self):
        body = self.get_report(self.RID, SEPT)
        text = json.dumps(body)
        self.assertNotIn("9000000001", text)
        self.assertNotIn("XXXX", text)
        self.assertNotIn("phone", [c["key"] for c in body["columns"]])

    def test_filters(self):
        def hosts(**p):
            return sorted(r["hostName"].lower() for r in self.rows(self.RID, {**SEPT, **p}))

        self.assertEqual(hosts(hostLinked="free_text"), ["anita kumar", "mr x"])
        self.assertEqual(hosts(hostLinked="linked_employee"), ["anita kumar", "bala raj", "dinesh babu"])
        self.assertEqual(hosts(departmentIds=self.d_cut.id), ["anita kumar", "bala raj"])  # free text cannot be matched
        self.assertEqual(hosts(employeeIds=self.e2.id), ["bala raj"])
        self.assertEqual(hosts(branchIds=self.b2.id), ["anita kumar", "dinesh babu"])
        self.assertEqual(hosts(dateFrom="2026-09-11", dateTo="2026-09-30"), ["dinesh babu", "mr x"])

    def test_branch_isolation(self):
        rows = self.rows(self.RID, SEPT, self.b1_user)
        self.assertEqual(sorted((r["hostType"], r["visits"]) for r in rows), [("Employee", 1), ("Employee", 2), ("Free text", 1)])
        self.assertEqual(sum(r["visits"] for r in rows), 4)
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        b2 = self.rows(self.RID, SEPT, self.b2_user)
        self.assertEqual(sum(r["visits"] for r in b2), 2)


class VisitorFrequencyTests(_Base):
    RID = "visitor-frequency"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual([r["visitorName"] for r in rows], ["Ravi Kumar", "Mohan", "Sita Devi"])
        v1 = rows[0]
        self.assertEqual(v1["phone"], "9000000001")
        self.assertEqual(v1["aadhaarMasked"], "XXXX XXXX 9012")
        self.assertEqual(v1["firstSeenAt"], "2026-09-02 10:00")
        self.assertEqual((v1["visitsInPeriod"], v1["totalVisitsEver"]), (3, 3))
        self.assertEqual(v1["lastVisitAt"], "2026-09-12 12:00")
        self.assertEqual(v1["distinctHosts"], 3)  # T001, "anita kumar " (free text), T004
        self.assertEqual((v1["lastHost"], v1["lastPurpose"]), ("Dinesh Babu", "Audit"))
        v3 = rows[1]
        self.assertEqual((v3["visitsInPeriod"], v3["distinctHosts"], v3["lastHost"]), (2, 2, "Mr X"))
        self.assertIsNone(v3["aadhaarMasked"])
        v2 = rows[2]
        self.assertEqual((v2["visitsInPeriod"], v2["distinctHosts"], v2["firstSeenAt"]), (2, 2, "2026-09-09 11:00"))
        self.assertEqual((v2["lastHost"], v2["lastPurpose"]), ("Anita Kumar", "Meeting"))
        self.assertEqual(body["totals"]["visitsInPeriod"], 7)
        s = summary_of(body)
        self.assertEqual((s["Unique visitors"], s["New in period"], s["Returning"]), (3, 3, 0))
        self.assertEqual(s["Most frequent visitor"], "Ravi Kumar (3)")
        self.assertNotIn(FULL_AADHAAR_1, json.dumps(body))

    def test_total_counts_history_outside_the_period(self):
        rows = self.rows(self.RID, {"dateFrom": "2026-09-09", "dateTo": "2026-09-12"})
        v1 = next(r for r in rows if r["visitorName"] == "Ravi Kumar")
        self.assertEqual((v1["visitsInPeriod"], v1["totalVisitsEver"]), (2, 3))
        self.assertEqual(v1["firstSeenAt"], "2026-09-02 10:00")

    def test_filters(self):
        def names(**p):
            return [r["visitorName"] for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(names(minVisits=3), ["Ravi Kumar"])
        self.assertEqual(names(minVisits=2), ["Ravi Kumar", "Mohan", "Sita Devi"])
        window = {"dateFrom": "2026-09-09", "dateTo": "2026-09-12"}
        got = lambda **p: [r["visitorName"] for r in self.rows(self.RID, {**window, **p})]  # noqa: E731
        self.assertEqual(got(visitorType="returning"), ["Ravi Kumar"])  # first seen 09-02, before the window
        self.assertEqual(got(visitorType="new_in_period"), ["Mohan", "Sita Devi"])
        self.assertEqual(names(q="sita"), ["Sita Devi"])
        self.assertEqual(names(q="9000000003"), ["Mohan"])
        self.assertEqual(names(branchIds=self.b2.id), ["Mohan", "Ravi Kumar"])  # one visit each at Unit 2: name breaks the tie

    def test_branch_isolation(self):
        b1 = self.rows(self.RID, SEPT, self.b1_user)
        self.assertEqual([(r["visitorName"], r["visitsInPeriod"], r["totalVisitsEver"]) for r in b1], [("Ravi Kumar", 2, 2), ("Sita Devi", 2, 2)])
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        b2 = self.rows(self.RID, SEPT, self.b2_user)
        # Ravi has 3 visits overall but a Unit 2 user only counts (and dates) the one at Unit 2
        ravi = next(r for r in b2 if r["visitorName"] == "Ravi Kumar")
        self.assertEqual((ravi["totalVisitsEver"], ravi["firstSeenAt"]), (1, "2026-09-12 12:00"))


class VisitorDailySummaryTests(_Base):
    RID = "visitor-daily-summary"
    WINDOW = {"dateFrom": "2026-09-08", "dateTo": "2026-09-12"}

    def test_golden_calendar_is_filled_and_days_are_ist(self):
        body = self.get_report(self.RID, self.WINDOW)
        rows = body["rows"]
        self.assertEqual([r["date"] for r in rows], ["2026-09-08", "2026-09-09", "2026-09-10", "2026-09-11", "2026-09-12"])
        self.assertEqual([r["weekday"] for r in rows], ["Tue", "Wed", "Thu", "Fri", "Sat"])
        self.assertEqual([r["visits"] for r in rows], [0, 3, 1, 0, 2])  # v2 (02:30 IST) is on the 9th, not the 8th
        d9 = rows[1]
        self.assertEqual((d9["uniqueVisitors"], d9["firstTime"], d9["repeat"], d9["hostLinked"]), (3, 2, 1, 1))
        self.assertEqual(d9["peakHour"], "02:00-02:59")  # three single check-ins: earliest hour wins the tie
        self.assertEqual((rows[2]["firstTime"], rows[2]["repeat"], rows[2]["peakHour"]), (0, 1, "09:00-09:59"))
        self.assertEqual((rows[3]["visits"], rows[3]["uniqueVisitors"], rows[3]["peakHour"]), (0, 0, None))
        d12 = rows[4]
        self.assertEqual((d12["uniqueVisitors"], d12["firstTime"], d12["repeat"], d12["hostLinked"]), (2, 0, 2, 1))
        self.assertEqual(body["totals"]["visits"], 6)
        self.assertEqual(body["totals"]["firstTime"] + body["totals"]["repeat"], 6)
        self.assertEqual(body["totals"]["hostLinked"], 3)
        s = summary_of(body)
        self.assertEqual(s["Visits"], sum(r["visits"] for r in rows))
        self.assertEqual(s["Unique visitors in period"], 3)
        self.assertEqual(s["Daily average (calendar days)"], 1.2)
        self.assertEqual(s["Busiest day"], "2026-09-09 (3)")

    def test_filters(self):
        free = self.rows(self.RID, {**self.WINDOW, "hostLinked": "free_text"})
        self.assertEqual([r["visits"] for r in free], [0, 2, 0, 0, 1])
        linked = self.rows(self.RID, {**self.WINDOW, "hostLinked": "linked_employee"})
        self.assertEqual([r["visits"] for r in linked], [0, 1, 1, 0, 1])
        b2 = self.rows(self.RID, {**self.WINDOW, "branchIds": self.b2.id})
        self.assertEqual([r["visits"] for r in b2], [0, 1, 0, 0, 1])

    def test_branch_isolation_and_scoped_first_visit(self):
        rows = self.rows(self.RID, self.WINDOW, self.b1_user)
        self.assertEqual([r["visits"] for r in rows], [0, 2, 1, 0, 0])
        self.assertEqual((rows[1]["firstTime"], rows[1]["repeat"]), (1, 1))  # v3 is Sita's first, v2 is Ravi's 2nd at Unit 1
        self.assertEqual(self.rows(self.RID, {**self.WINDOW, "branchIds": self.b2.id}, self.b1_user)[1]["visits"], 0)


class VisitorNotificationTests(_Base):
    RID = "visitor-notification-delivery"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual([r["visitedAt"][5:] for r in rows], ["09-02 10:00", "09-09 11:00", "09-10 09:00", "09-12 12:00"])
        r0, r1, r2, r3 = rows
        self.assertEqual((r0["hostName"], r0["hostDepartment"], r0["hostHasEmail"], r0["hostHasPhone"]), ("Anita Kumar", "CUTTING", "Yes", "Yes"))
        self.assertEqual((r0["emailDelaySeconds"], r0["whatsappDelaySeconds"]), (4, 40))
        self.assertEqual(r0["emailSentAt"], "2026-09-02 10:00")
        self.assertEqual((r1["hostHasEmail"], r1["hostHasPhone"]), ("No", "Yes"))
        self.assertEqual((r1["emailDelaySeconds"], r1["whatsappSentAt"], r1["whatsappDelaySeconds"]), (10, None, None))
        self.assertEqual((r2["emailSentAt"], r2["whatsappDelaySeconds"]), (None, 6))
        self.assertEqual((r3["hostHasEmail"], r3["hostHasPhone"], r3["emailDelaySeconds"]), ("No", "No", None))  # "" email and NULL phone
        s = summary_of(body)
        self.assertEqual(s["Visits with a linked host"], 4)
        self.assertEqual((s["Email delivered"], s["WhatsApp delivered"]), (50.0, 50.0))
        self.assertEqual((s["Avg email delay (s)"], s["Avg WhatsApp delay (s)"]), (7, 23))
        self.assertEqual((s["Hosts with no email on file"], s["Hosts with no phone on file"]), (2, 1))
        self.assertEqual(s["Visits with a linked host"], len(rows))
        text = json.dumps(body)
        self.assertNotIn("anita@example.com", text)  # presence only: contact details are not printed
        self.assertNotIn("9111111111", text)

    def test_free_text_hosts_are_not_listed(self):
        names = {r["visitorName"] for r in self.rows(self.RID, SEPT)}
        self.assertEqual(names, {"Ravi Kumar", "Sita Devi"})
        self.assertEqual(len(self.rows(self.RID, {"dateFrom": "2026-09-09", "dateTo": "2026-09-09"})), 1)  # only v3 has a host that day

    def test_filters(self):
        def at(**p):
            return [r["visitedAt"][5:10] + " " + r["visitorName"] for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(at(deliveryState="both_sent"), ["09-02 Ravi Kumar"])
        self.assertEqual(at(deliveryState="email_only"), ["09-09 Sita Devi"])
        self.assertEqual(at(deliveryState="whatsapp_only"), ["09-10 Sita Devi"])
        self.assertEqual(at(deliveryState="none_sent"), ["09-12 Ravi Kumar"])
        self.assertEqual(at(departmentIds=self.d_cut.id), ["09-02 Ravi Kumar", "09-09 Sita Devi", "09-10 Sita Devi"])
        self.assertEqual(at(employeeIds=self.e4.id), ["09-12 Ravi Kumar"])
        self.assertEqual(at(branchIds=self.b2.id), ["09-12 Ravi Kumar"])

    def test_branch_isolation(self):
        self.assertEqual(len(self.rows(self.RID, SEPT, self.b1_user)), 3)
        self.assertEqual(len(self.rows(self.RID, SEPT, self.b2_user)), 1)
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])


# ═════════════════════════════════════════════════════════════════════════════
# Tea breaks
# ═════════════════════════════════════════════════════════════════════════════


class TeaRegisterTests(_Base):
    RID = "tea-break-register"
    DAY = {"dateFrom": "2026-09-05", "dateTo": "2026-09-05"}

    def test_golden_day(self):
        body = self.get_report(self.RID, self.DAY)
        rows = body["rows"]
        self.assertEqual([r["employeeCode"] for r in rows], ["T001", "T004", "T002", "T005", "T001", "T002", "T004"])
        self.assertEqual([r["outAt"] for r in rows], ["10:00", "10:00", "10:05", "11:00", "15:00", "15:10", "16:00"])
        a = rows[0]
        self.assertEqual(
            (a["date"], a["employeeName"], a["department"], a["designation"], a["employmentType"]),
            ("2026-09-05", "Anita Kumar", "CUTTING", "Operator", "Staff"),
        )
        self.assertEqual((a["outGate"], a["inAt"], a["inGate"]), ("Gate 1", "10:10", "Gate 1"))
        self.assertEqual((a["takenMinutes"], a["allowedMinutes"], a["overMinutes"], a["remark"], a["flag"]), (10, 15, 0, "On time", None))
        m = rows[3]  # resigned employee with no department / designation
        self.assertEqual((m["department"], m["designation"], m["employmentType"], m["takenMinutes"]), ("Unassigned", None, "Production", 14))
        b, d = rows[4], rows[5]
        self.assertEqual((b["takenMinutes"], b["overMinutes"], b["remark"]), (16, 1, "Overtime"))  # 15 min 30 s -> 16
        self.assertEqual((d["inAt"], d["takenMinutes"], d["overMinutes"], d["inGate"]), ("15:26", 16, 1, "Gate 2"))  # 16.5 -> 16 (half-even)
        self.assertEqual(rows[6]["overMinutes"], 15)
        self.assertEqual(body["totals"]["takenMinutes"], 113)
        self.assertEqual(body["totals"]["overMinutes"], 17)
        s = summary_of(body)
        self.assertEqual((s["Allowed minutes"], s["Breaks"], s["Completed"], s["On time"], s["Overtime"]), (15, 7, 7, 4, 3))
        self.assertEqual((s["Not returned"], s["In progress"], s["Long or left open"], s["Employees"]), (0, 0, 0, 4))
        self.assertEqual(s["Total time taken"], 113)
        self.assertEqual(s["Minutes over allowance"], 17)
        self.assertEqual((s["Longest break (min)"], s["Avg minutes / break"]), (30, round(113 / 7, 2)))
        # summary agrees with the rows / totals row
        self.assertEqual(s["Total time taken"], body["totals"]["takenMinutes"])
        self.assertEqual(s["Breaks"], len(rows))
        self.assertEqual(s["Minutes over allowance"], body["totals"]["overMinutes"])
        notes = " ".join(body["notes"])
        self.assertIn("Allowed minutes used: 15 (rule last changed 03-Sep-2026)", notes)
        self.assertIn("Informational only", notes)

    def test_cross_midnight_break_belongs_to_the_day_it_started(self):
        rows = self.rows(self.RID, {"dateFrom": "2026-09-06", "dateTo": "2026-09-06"})
        self.assertEqual(len(rows), 1)
        e = rows[0]
        self.assertEqual((e["date"], e["outAt"], e["inAt"], e["takenMinutes"], e["overMinutes"]), ("2026-09-06", "23:50", "00:10 (+1d)", 20, 5))
        self.assertEqual(self.rows(self.RID, {"dateFrom": "2026-09-07", "dateTo": "2026-09-07"})[0]["employeeCode"], "T003")

    def test_night_shift_break_is_bucketed_by_ist_not_utc(self):
        # F started 01:30 IST on 09-07 = 20:00 UTC on 09-06
        on_7th = self.rows(self.RID, {"dateFrom": "2026-09-07", "dateTo": "2026-09-07"})
        self.assertEqual([(r["employeeCode"], r["outAt"], r["date"]) for r in on_7th], [("T003", "01:30", "2026-09-07")])
        self.assertNotIn("T003", [r["employeeCode"] for r in self.rows(self.RID, {"dateFrom": "2026-09-06", "dateTo": "2026-09-06"})])

    def test_open_breaks_have_no_duration_and_are_marked(self):
        rows = self.rows(self.RID, {"dateFrom": "2026-09-14", "dateTo": "2026-09-15"})
        self.assertEqual([(r["employeeCode"], r["remark"], r["flag"]) for r in rows], [
            ("T001", "Not returned", "Left open"),  # since 20:00 the day before: 19.5 h
            ("T002", "Not returned", None),  # 150 min
            ("T004", "In progress", None),  # 30 min
        ])
        for r in rows:
            self.assertIsNone(r["takenMinutes"])
            self.assertIsNone(r["overMinutes"])
            self.assertIsNone(r["inAt"])
        s = summary_of(self.get_report(self.RID, {"dateFrom": "2026-09-14", "dateTo": "2026-09-15"}))
        self.assertEqual((s["Breaks"], s["Completed"], s["Not returned"], s["In progress"], s["Long or left open"]), (3, 0, 2, 1, 1))
        self.assertIsNone(s["Avg minutes / break"])  # nothing completed: a dash, not 0
        self.assertIsNone(s["Longest break (min)"])

    def test_long_break_is_flagged(self):
        rows = self.rows(self.RID, {"dateFrom": "2026-09-08", "dateTo": "2026-09-08"})
        g = rows[0]
        self.assertEqual((g["takenMinutes"], g["overMinutes"], g["remark"], g["flag"]), (150, 135, "Overtime", "Long break"))

    def test_filters(self):
        def codes(**p):
            return sorted(r["employeeCode"] for r in self.rows(self.RID, {**SEPT, **p}))

        self.assertEqual(len(self.rows(self.RID, SEPT)), 13)
        self.assertEqual(self.rows(self.RID, {"dateFrom": "2026-09-05", "dateTo": "2026-09-07"})[0]["date"], "2026-09-05")
        self.assertEqual(codes(remark="on_time"), ["T001", "T002", "T003", "T004", "T005"])  # A C F H M
        self.assertEqual(codes(remark="overtime"), ["T001", "T001", "T002", "T003", "T004"])  # B E D G I
        self.assertEqual(codes(remark="not_returned"), ["T001", "T002"])  # L J
        self.assertEqual(codes(remark="in_progress"), ["T004"])  # K
        self.assertEqual(sorted(r["flag"] for r in self.rows(self.RID, {**SEPT, "longBreaks": "only"})), ["Left open", "Long break"])
        self.assertEqual(len(self.rows(self.RID, {**SEPT, "longBreaks": "exclude"})), 11)
        self.assertEqual(codes(departmentIds=self.d_cut.id), ["T001"] * 4 + ["T002"] * 3)
        self.assertEqual(codes(employmentType="production"), ["T002"] * 3 + ["T005"])
        self.assertEqual(len(self.rows(self.RID, {**SEPT, "designationIds": self.des_op.id})), 9)  # E1 4 + E3 2 + E4 3
        self.assertEqual(codes(employeeIds=f"{self.e3.id},{self.e5.id}"), ["T003", "T003", "T005"])
        self.assertEqual(codes(employeeStatus="inactive"), ["T005"])  # "resigned" counts as inactive
        self.assertEqual(len(self.rows(self.RID, {**SEPT, "employeeStatus": "active"})), 12)
        self.assertEqual(len(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id})), 3)

    def test_ist_day_edges_2330_and_0030(self):
        TeaBreakLog.objects.create(employee=self.e1, out_at=ist(9, 25, 23, 30), in_at=ist(9, 25, 23, 35))  # 18:00 UTC on the 25th
        TeaBreakLog.objects.create(employee=self.e1, out_at=ist(9, 26, 0, 30), in_at=ist(9, 26, 0, 35))  # 19:00 UTC on the 25th
        d25 = self.rows(self.RID, {"dateFrom": "2026-09-25", "dateTo": "2026-09-25"})
        d26 = self.rows(self.RID, {"dateFrom": "2026-09-26", "dateTo": "2026-09-26"})
        self.assertEqual([(r["date"], r["outAt"]) for r in d25], [("2026-09-25", "23:30")])
        self.assertEqual([(r["date"], r["outAt"]) for r in d26], [("2026-09-26", "00:30")])

    def test_branch_isolation(self):
        rows = self.rows(self.RID, SEPT, self.b1_user)
        self.assertEqual(len(rows), 10)
        self.assertNotIn("T004", [r["employeeCode"] for r in rows])
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual(self.rows(self.RID, {**SEPT, "employeeIds": self.e4.id}, self.b1_user), [])
        self.assertEqual(len(self.rows(self.RID, SEPT, self.b2_user)), 3)
        self.assertEqual(summary_of(self.get_report(self.RID, SEPT, self.b1_user))["Breaks"], 10)

    def test_get_never_creates_the_rule_row(self):
        TeaBreakRule.objects.all().delete()
        body = self.get_report(self.RID, self.DAY)
        self.assertEqual(TeaBreakRule.objects.count(), 0)
        self.assertEqual(summary_of(body)["Allowed minutes"], 15)
        self.assertIn("default; no rule saved yet", " ".join(body["notes"]))

    def test_rule_change_recomputes_history_and_says_so(self):
        TeaBreakRule.objects.filter(pk=1).update(allowed_minutes=30)
        body = self.get_report(self.RID, self.DAY)
        self.assertEqual(summary_of(body)["Overtime"], 0)
        self.assertEqual([r["allowedMinutes"] for r in body["rows"]], [30] * 7)
        self.assertIn("applied to every break in the period as it stands today", " ".join(body["notes"]))

    def test_default_range_is_today(self):
        self.assertEqual(registry.get_spec(self.RID).filter("dateRange").default, "today")
        with mock.patch("api.reporting.filters.ist_today", return_value=date(2026, 9, 5)):
            self.assertEqual(len(self.get_report(self.RID, {})["rows"]), 7)
        with mock.patch("api.reporting.filters.ist_today", return_value=date(2026, 9, 20)):
            self.assertEqual(self.get_report(self.RID, {})["rows"], [])

    def test_matches_the_hr_tea_break_page_row_for_row(self):
        counts = Counter()
        for log in TeaBreakLog.objects.select_related("employee__department"):
            page = _log_json(log, 15, NOW)
            counts[page["remark"]] += 1
        s = summary_of(self.get_report(self.RID, SEPT))
        self.assertEqual(
            (s["On time"], s["Overtime"], s["Not returned"], s["In progress"]),
            (counts["on_time"], counts["overtime"], counts["not_returned"], counts["in_progress"]),
        )
        self.assertEqual(sum(counts.values()), s["Breaks"])


class TeaEmployeeSummaryTests(_Base):
    RID = "tea-break-employee-summary"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual([r["employeeCode"] for r in body["rows"]], ["T001", "T002", "T003", "T004", "T005"])
        e1 = rows["T001"]
        self.assertEqual((e1["breaks"], e1["completed"], e1["onTime"], e1["overtime"], e1["notReturned"], e1["inProgress"]), (4, 3, 1, 2, 1, 0))
        self.assertEqual((e1["longBreaks"], e1["totalMinutes"], e1["avgMinutes"], e1["maxMinutes"]), (1, 46, 15.33, 20))
        self.assertEqual((e1["overMinutesTotal"], e1["overtimePct"], e1["daysWithBreaks"]), (6, 66.7, 3))
        e2 = rows["T002"]
        self.assertEqual((e2["breaks"], e2["overtime"], e2["notReturned"], e2["totalMinutes"], e2["avgMinutes"]), (3, 1, 1, 31, 15.5))
        self.assertEqual((e2["overMinutesTotal"], e2["daysWithBreaks"], e2["longBreaks"]), (1, 2, 0))
        e3 = rows["T003"]
        self.assertEqual((e3["totalMinutes"], e3["avgMinutes"], e3["maxMinutes"], e3["overMinutesTotal"], e3["longBreaks"]), (160, 80.0, 150, 135, 1))
        self.assertEqual(e3["daysWithBreaks"], 2)  # 09-07 (IST) and 09-08
        e4 = rows["T004"]
        self.assertEqual((e4["inProgress"], e4["totalMinutes"], e4["avgMinutes"], e4["overMinutesTotal"]), (1, 42, 21.0, 15))
        e5 = rows["T005"]
        self.assertEqual((e5["department"], e5["designation"], e5["overtime"], e5["overtimePct"], e5["avgMinutes"]), ("Unassigned", None, 0, 0.0, 14.0))
        t = body["totals"]
        self.assertEqual((t["breaks"], t["completed"], t["onTime"], t["overtime"], t["notReturned"], t["inProgress"]), (13, 10, 5, 5, 2, 1))
        self.assertEqual((t["longBreaks"], t["totalMinutes"], t["overMinutesTotal"], t["daysWithBreaks"]), (2, 293, 157, 10))
        s = summary_of(body)
        self.assertEqual((s["Employees with breaks"], s["Breaks"], s["Overtime breaks"], s["Long or left open"]), (5, 13, 5, 2))
        self.assertEqual(s["Avg minutes / break"], 29.3)
        self.assertEqual(s["Most overtime breaks"], "T001 Anita Kumar (2)")
        self.assertEqual(s["Most total time"], "T003 Chitra Devi (160 min)")
        self.assertEqual(s["Breaks"], t["breaks"])

    def test_filters(self):
        def codes(**p):
            return [r["employeeCode"] for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(codes(onlyOvertime="true"), ["T001", "T002", "T003", "T004"])
        self.assertEqual(codes(departmentIds=self.d_cut.id), ["T001", "T002"])
        self.assertEqual(codes(employmentType="production"), ["T002", "T005"])
        self.assertEqual(codes(designationIds=self.des_help.id), ["T002"])
        self.assertEqual(codes(employeeIds=self.e3.id), ["T003"])
        self.assertEqual(codes(employeeStatus="active"), ["T001", "T002", "T003", "T004"])
        self.assertEqual(codes(branchIds=self.b2.id), ["T004"])
        self.assertEqual(codes(dateFrom="2026-09-06", dateTo="2026-09-07"), ["T001", "T003"])

    def test_exclude_suspect_removes_long_and_left_open_breaks(self):
        body = self.get_report(self.RID, {**SEPT, "excludeSuspect": "true"})
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["T001"]["breaks"], 3)  # L (left open) gone
        self.assertEqual(rows["T003"]["breaks"], 1)  # G (150 min) gone
        self.assertEqual(rows["T003"]["totalMinutes"], 10)
        self.assertEqual(body["totals"]["breaks"], 11)
        self.assertEqual(body["totals"]["longBreaks"], 0)
        self.assertIn("left out of this report", " ".join(body["notes"]))

    def test_branch_isolation(self):
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, SEPT, self.b1_user)], ["T001", "T002", "T003", "T005"])
        self.assertEqual(self.rows(self.RID, {**SEPT, "employeeIds": self.e4.id}, self.b1_user), [])
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, SEPT, self.b2_user)], ["T004"])


class TeaDepartmentSummaryTests(_Base):
    RID = "tea-break-department-summary"

    def test_golden_by_department_default_active_only(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual([(r["group"], r["branch"]) for r in rows], [("CUTTING", "Unit 1"), ("PACKING", "Unit 2"), ("SEWING", "Unit 1")])
        cut, pack, sew = rows
        self.assertEqual((cut["headcount"], cut["employeesOnBreak"], cut["participationPct"]), (3, 2, 66.7))  # T006 never scans
        self.assertEqual((cut["breaks"], cut["totalMinutes"], cut["avgMinutes"]), (7, 77, 15.4))
        self.assertEqual((cut["overtime"], cut["overtimePct"], cut["notReturned"], cut["longBreaks"]), (3, 60.0, 2, 1))
        self.assertEqual(cut["avgBreaksPerEmployeeDay"], 1.4)  # 7 breaks over 5 employee-days
        self.assertEqual((sew["headcount"], sew["employeesOnBreak"], sew["participationPct"], sew["breaks"]), (1, 1, 100.0, 2))
        self.assertEqual((sew["totalMinutes"], sew["avgMinutes"], sew["overtime"], sew["longBreaks"]), (160, 80.0, 1, 1))
        self.assertEqual((pack["totalMinutes"], pack["avgMinutes"], pack["avgBreaksPerEmployeeDay"]), (42, 21.0, 1.5))
        t = body["totals"]
        self.assertEqual((t["headcount"], t["employeesOnBreak"], t["breaks"], t["totalMinutes"], t["overtime"]), (5, 4, 12, 279, 5))
        s = summary_of(body)
        self.assertEqual((s["Breaks"], s["Headcount"]), (12, 5))
        self.assertEqual(s["Participation"], 80.0)
        self.assertEqual(s["Overtime %"], round(100 * 5 / 9, 1))  # 9 completed breaks among active employees
        self.assertEqual(s["Highest overtime %"], "CUTTING (60.0%)")
        self.assertEqual(s["Breaks"], t["breaks"])

    def test_all_statuses_adds_no_department_group(self):
        rows = self.rows(self.RID, {**SEPT, "employeeStatus": "all"})
        self.assertEqual([r["group"] for r in rows], ["CUTTING", "No department", "PACKING", "SEWING"])
        nd = rows[1]
        self.assertEqual((nd["branch"], nd["headcount"], nd["employeesOnBreak"], nd["breaks"], nd["overtime"], nd["overtimePct"]), (None, 1, 1, 1, 0, 0.0))

    def test_group_by_branch_and_employment_type(self):
        body = self.get_report(self.RID, {**SEPT, "employeeStatus": "all", "groupBy": "branch"})
        self.assertEqual(body["columns"][0]["label"], "Branch")
        self.assertNotIn("branch", [c["key"] for c in body["columns"]])
        rows = {r["group"]: r for r in body["rows"]}
        self.assertEqual(list(rows), ["No branch", "Unit 1", "Unit 2"])
        self.assertEqual((rows["No branch"]["headcount"], rows["No branch"]["breaks"], rows["No branch"]["participationPct"]), (1, 0, 0.0))
        self.assertIsNone(rows["No branch"]["avgMinutes"])  # nothing to average: a dash
        self.assertEqual((rows["Unit 1"]["headcount"], rows["Unit 1"]["breaks"]), (4, 10))
        self.assertEqual((rows["Unit 2"]["headcount"], rows["Unit 2"]["breaks"]), (1, 3))
        body = self.get_report(self.RID, {**SEPT, "employeeStatus": "all", "groupBy": "employmentType"})
        self.assertEqual(body["columns"][0]["label"], "Employee type")
        rows = {r["group"]: r for r in body["rows"]}
        self.assertEqual((rows["Staff"]["headcount"], rows["Staff"]["breaks"]), (4, 9))
        self.assertEqual((rows["Production"]["headcount"], rows["Production"]["breaks"]), (2, 4))

    def test_filters_and_exclude_suspect(self):
        rows = self.rows(self.RID, {**SEPT, "departmentIds": self.d_sew.id})
        self.assertEqual([r["group"] for r in rows], ["SEWING"])
        rows = self.rows(self.RID, {**SEPT, "employmentType": "production"})
        self.assertEqual([(r["group"], r["headcount"], r["breaks"]) for r in rows], [("CUTTING", 1, 3)])
        body = self.get_report(self.RID, {**SEPT, "excludeSuspect": "true"})
        self.assertEqual(body["totals"]["breaks"], 10)
        self.assertEqual(body["totals"]["longBreaks"], 0)

    def test_branch_isolation(self):
        rows = self.rows(self.RID, SEPT, self.b1_user)
        self.assertEqual([(r["group"], r["headcount"]) for r in rows], [("CUTTING", 2), ("SEWING", 1)])  # T006 has no branch
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual([r["group"] for r in self.rows(self.RID, SEPT, self.b2_user)], ["PACKING"])


class TeaExceptionsTests(_Base):
    RID = "tea-break-exceptions"

    def test_golden_worst_first(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual([(r["employeeCode"], r["outAt"]) for r in rows], [
            ("T003", "09:00"), ("T004", "16:00"), ("T001", "23:50"), ("T002", "15:10"), ("T001", "15:00"),  # overtime, longest first
            ("T001", "20:00"), ("T002", "13:00"),  # never closed last, oldest first
        ])
        g = rows[0]
        self.assertEqual(
            (g["date"], g["inAt"], g["takenMinutes"], g["allowedMinutes"], g["overMinutes"], g["remark"], g["flag"]),
            ("2026-09-08", "11:30", 150, 15, 135, "Overtime", "Long break"),
        )
        self.assertEqual((g["overtimeCountInPeriod"], g["outGate"], g["inGate"]), (1, "Gate 2", "Gate 2"))
        self.assertEqual((rows[3]["takenMinutes"], rows[4]["takenMinutes"]), (16, 16))  # 16m30s ranks above 15m30s though both show 16
        self.assertEqual(rows[2]["inAt"], "00:10 (+1d)")
        l_row = rows[5]
        self.assertEqual((l_row["date"], l_row["inAt"], l_row["takenMinutes"], l_row["overMinutes"]), ("2026-09-14", None, None, None))
        self.assertEqual((l_row["remark"], l_row["flag"], l_row["overtimeCountInPeriod"], l_row["inGate"]), ("Not returned", "Left open", 2, None))
        self.assertEqual(rows[6]["flag"], None)
        self.assertEqual(body["totals"]["takenMinutes"], 232)
        self.assertEqual(body["totals"]["overMinutes"], 157)
        s = summary_of(body)
        self.assertEqual((s["Overtime breaks"], s["Minutes over allowance"], s["Not returned"]), (5, 157, 2))
        self.assertEqual(s["Likely missed scans (long or left open)"], 2)
        self.assertEqual(s["Employees with 3+ overtime breaks (whole period)"], 0)
        self.assertEqual(s["Minutes over allowance"], body["totals"]["overMinutes"])

    def test_filters(self):
        def key(**p):
            return [(r["employeeCode"], r["outAt"]) for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(len(key(exceptionType="overtime")), 5)
        self.assertEqual(key(exceptionType="not_returned"), [("T001", "20:00"), ("T002", "13:00")])
        # rounded minutes >= allowed + N: N=10 -> >= 25 (G 150, I 30); N=5 -> >= 20 (adds E, exactly 20)
        self.assertEqual(key(exceptionType="overtime", minOverMinutes=10), [("T003", "09:00"), ("T004", "16:00")])
        self.assertEqual(key(exceptionType="overtime", minOverMinutes=5), [("T003", "09:00"), ("T004", "16:00"), ("T001", "23:50")])
        # open breaks are not affected by the minimum-over filter
        self.assertEqual(len(key(minOverMinutes=30)), 1 + 2)
        self.assertEqual(key(employmentType="production"), [("T002", "15:10"), ("T002", "13:00")])  # D overtime, J open; M is on time
        self.assertEqual(key(departmentIds=self.d_sew.id), [("T003", "09:00")])
        self.assertEqual(len(key(employeeIds=self.e1.id)), 3)
        self.assertEqual(len(key(dateFrom="2026-09-14", dateTo="2026-09-30")), 2)

    def test_repeat_offenders_counted_over_the_whole_period(self):
        for i in range(3):  # T002 gets three more overtime breaks on another day -> 4 in total
            TeaBreakLog.objects.create(employee=self.e2, out_at=ist(9, 21, 9 + i, 0), in_at=ist(9, 21, 9 + i, 20))
        body = self.get_report(self.RID, {**SEPT, "employeeIds": self.e2.id})
        self.assertEqual(summary_of(body)["Employees with 3+ overtime breaks (whole period)"], 1)
        self.assertEqual({r["overtimeCountInPeriod"] for r in body["rows"]}, {4})

    def test_branch_isolation(self):
        rows = self.rows(self.RID, SEPT, self.b1_user)
        self.assertNotIn("T004", [r["employeeCode"] for r in rows])
        self.assertEqual(len(rows), 6)
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual(len(self.rows(self.RID, SEPT, self.b2_user)), 1)


class TeaDailyTrendTests(_Base):
    RID = "tea-break-daily-trend"
    WINDOW = {"dateFrom": "2026-09-01", "dateTo": "2026-09-10"}

    def test_golden_calendar_filled_and_ist_days(self):
        body = self.get_report(self.RID, self.WINDOW)
        rows = {r["date"]: r for r in body["rows"]}
        self.assertEqual(len(body["rows"]), 10)
        d5 = rows["2026-09-05"]
        self.assertEqual((d5["weekday"], d5["breaks"], d5["employees"], d5["totalMinutes"], d5["avgMinutes"]), ("Sat", 7, 4, 113, round(113 / 7, 2)))
        self.assertEqual((d5["overtime"], d5["overtimePct"], d5["notReturned"], d5["peakHour"]), (3, 42.9, 0, "10:00-10:59"))
        d6 = rows["2026-09-06"]
        self.assertEqual((d6["breaks"], d6["totalMinutes"], d6["overtime"], d6["overtimePct"], d6["peakHour"]), (1, 20, 1, 100.0, "23:00-23:59"))
        d7 = rows["2026-09-07"]  # the 01:30 IST night-shift break is on the 7th (it is 20:00 UTC on the 6th)
        self.assertEqual((d7["breaks"], d7["totalMinutes"], d7["overtimePct"], d7["peakHour"]), (1, 10, 0.0, "01:00-01:59"))
        self.assertEqual((rows["2026-09-08"]["totalMinutes"], rows["2026-09-08"]["overtime"]), (150, 1))
        empty = rows["2026-09-01"]
        self.assertEqual((empty["breaks"], empty["employees"], empty["totalMinutes"], empty["overtime"]), (0, 0, 0, 0))
        self.assertEqual((empty["avgMinutes"], empty["overtimePct"], empty["peakHour"]), (None, None, None))
        self.assertEqual(body["totals"]["breaks"], 10)
        self.assertEqual(body["totals"]["totalMinutes"], 293)
        self.assertEqual(body["totals"]["overtime"], 5)
        s = summary_of(body)
        self.assertEqual((s["Breaks"], s["Avg breaks / calendar day"], s["Overtime breaks"]), (10, 1.0, 5))
        self.assertEqual(s["Worst day (overtime %)"], "2026-09-06 (100.0%)")  # tie with 09-08 -> earlier day
        self.assertEqual(s["Breaks"], body["totals"]["breaks"])

    def test_filters_exclude_suspect_and_branch_isolation(self):
        body = self.get_report(self.RID, {**self.WINDOW, "excludeSuspect": "true"})
        self.assertEqual(body["totals"]["breaks"], 9)  # G (150 min) gone
        self.assertEqual({r["date"]: r["breaks"] for r in body["rows"]}["2026-09-08"], 0)
        prod = self.get_report(self.RID, {**self.WINDOW, "employmentType": "production"})
        self.assertEqual(prod["totals"]["breaks"], 3)  # C, D, M
        self.assertEqual(self.get_report(self.RID, {**self.WINDOW, "departmentIds": self.d_pack.id})["totals"]["breaks"], 2)
        b1 = self.get_report(self.RID, self.WINDOW, self.b1_user)
        self.assertEqual(b1["totals"]["breaks"], 8)  # H and I (Unit 2) hidden
        self.assertEqual(self.get_report(self.RID, {**self.WINDOW, "branchIds": self.b2.id}, self.b1_user)["totals"]["breaks"], 0)
        self.assertEqual(self.get_report(self.RID, self.WINDOW, self.b2_user)["totals"]["breaks"], 2)


class TeaHourlyTests(_Base):
    RID = "tea-break-hourly-distribution"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = {r["hourBand"]: r for r in body["rows"]}
        self.assertEqual(len(body["rows"]), 23)  # 01:00 .. 23:00, contiguous
        self.assertEqual(body["rows"][0]["hourBand"], "01:00-01:59")
        h1 = rows["01:00-01:59"]  # night shift: 01:30 IST, not folded into the previous UTC day/hour
        self.assertEqual((h1["breaks"], h1["sharePct"], h1["avgMinutes"], h1["overtime"]), (1, 7.7, 10.0, 0))
        h10 = rows["10:00-10:59"]
        self.assertEqual((h10["breaks"], h10["sharePct"], h10["avgMinutes"], h10["overtime"]), (3, 23.1, 12.33, 0))
        h15 = rows["15:00-15:59"]  # B, D and the open K (open breaks count as breaks but not in the average)
        self.assertEqual((h15["breaks"], h15["avgMinutes"], h15["overtime"]), (3, 16.0, 2))
        quiet = rows["02:00-02:59"]
        self.assertEqual((quiet["breaks"], quiet["sharePct"], quiet["avgMinutes"], quiet["overtime"]), (0, 0.0, None, 0))
        self.assertEqual(body["totals"]["breaks"], 13)
        self.assertEqual(body["totals"]["overtime"], 5)
        s = summary_of(body)
        self.assertEqual((s["Breaks"], s["Peak hour"], s["Breaks in peak hour"]), (13, "10:00-10:59", 3))  # 10 and 15 tie: earlier
        self.assertEqual(s["Breaks"], body["totals"]["breaks"])

    def test_filters_and_isolation(self):
        body = self.get_report(self.RID, {**SEPT, "excludeSuspect": "true"})
        self.assertEqual(body["totals"]["breaks"], 11)
        self.assertEqual(self.get_report(self.RID, {**SEPT, "employmentType": "production"})["totals"]["breaks"], 4)
        self.assertEqual(self.get_report(self.RID, SEPT, self.b1_user)["totals"]["breaks"], 10)
        self.assertEqual(self.get_report(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user)["rows"], [])
        self.assertEqual(self.get_report(self.RID, {"dateFrom": "2025-01-01", "dateTo": "2025-01-31"})["rows"], [])


class TeaFrequencyTests(_Base):
    RID = "tea-break-frequency"

    def test_golden_more_than_one_break_a_day(self):
        body = self.get_report(self.RID, {**SEPT, "minBreaksPerDay": 1})
        rows = body["rows"]
        self.assertEqual([(r["date"], r["employeeCode"]) for r in rows], [("2026-09-05", "T001"), ("2026-09-05", "T002"), ("2026-09-05", "T004")])
        a, b, c = rows
        self.assertEqual((a["breaksThatDay"], a["totalMinutes"], a["overtime"], a["stillOpen"], a["firstOutAt"], a["lastInAt"]), (2, 26, 1, 0, "10:00", "15:15"))
        self.assertEqual((b["totalMinutes"], b["firstOutAt"], b["lastInAt"]), (31, "10:05", "15:26"))
        self.assertEqual((c["totalMinutes"], c["overtime"], c["lastInAt"]), (42, 1, "16:30"))
        self.assertEqual(body["totals"]["breaksThatDay"], 6)
        s = summary_of(body)
        self.assertEqual((s["Employee-days flagged"], s["Employees flagged"], s["Highest breaks in a day"]), (3, 3, 2))
        self.assertEqual(s["Most often flagged"], "T001 Anita Kumar (1 day)")

    def test_default_threshold_and_open_breaks(self):
        self.assertEqual(self.rows(self.RID, SEPT), [])  # nobody had MORE than 2 in a day
        for i, (h, m) in enumerate([(9, 0), (12, 0), (15, 0)]):
            TeaBreakLog.objects.create(employee=self.e3, out_at=ist(9, 20, h, m), in_at=ist(9, 20, h, m + 5) if i < 2 else None)
        rows = self.rows(self.RID, SEPT)
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r["date"], r["employeeCode"], r["breaksThatDay"], r["totalMinutes"], r["stillOpen"]), ("2026-09-20", "T003", 3, 10, 1))
        self.assertEqual((r["firstOutAt"], r["lastInAt"]), ("09:00", "12:05"))
        self.assertEqual(self.rows(self.RID, {**SEPT, "minBreaksPerDay": 3}), [])

    def test_filters_and_isolation(self):
        p = {**SEPT, "minBreaksPerDay": 1}
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, {**p, "departmentIds": self.d_cut.id})], ["T001", "T002"])
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, {**p, "employeeIds": self.e4.id})], ["T004"])
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, p, self.b1_user)], ["T001", "T002"])
        self.assertEqual(self.rows(self.RID, {**p, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, p, self.b2_user)], ["T004"])


# ═════════════════════════════════════════════════════════════════════════════
# Overview
# ═════════════════════════════════════════════════════════════════════════════


class TimeAwayTests(_Base):
    RID = "employee-time-away-summary"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(list(rows), ["T001", "T002", "T003", "T004", "T005"])
        e1 = rows["T001"]
        self.assertEqual((e1["outpassExits"], e1["outpassMinutes"], e1["qrExits"]), (3, 136, 2))  # 90 + 46 (46.5 half-even); 3rd never returned
        self.assertEqual((e1["teaBreaks"], e1["teaMinutes"], e1["teaOvertime"], e1["visitorsHosted"], e1["totalAwayMinutes"]), (4, 46, 2, 2, 182))
        e2 = rows["T002"]
        self.assertEqual((e2["outpassExits"], e2["outpassMinutes"], e2["qrExits"], e2["teaBreaks"], e2["teaMinutes"]), (1, 20, 0, 3, 31))
        self.assertEqual((e2["teaOvertime"], e2["visitorsHosted"], e2["totalAwayMinutes"]), (1, 1, 51))
        e3 = rows["T003"]
        self.assertEqual((e3["outpassExits"], e3["teaMinutes"], e3["visitorsHosted"], e3["totalAwayMinutes"]), (0, 160, 0, 160))  # pending pass is not an exit
        e4 = rows["T004"]
        self.assertEqual((e4["outpassExits"], e4["outpassMinutes"], e4["qrExits"], e4["teaMinutes"], e4["totalAwayMinutes"]), (1, 30, 1, 42, 72))
        self.assertEqual((e4["visitorsHosted"], e4["teaBreaks"]), (1, 3))
        e5 = rows["T005"]
        self.assertEqual((e5["department"], e5["totalAwayMinutes"], e5["outpassExits"]), ("Unassigned", 14, 0))
        t = body["totals"]
        self.assertEqual((t["outpassExits"], t["outpassMinutes"], t["qrExits"], t["teaBreaks"], t["teaMinutes"]), (5, 186, 3, 13, 293))
        self.assertEqual((t["teaOvertime"], t["visitorsHosted"], t["totalAwayMinutes"]), (5, 4, 186 + 293))
        s = summary_of(body)
        self.assertEqual(s["Employees with gate activity"], 5)
        self.assertEqual((s["Outpass exits"], s["Tea breaks"], s["Visitors hosted"]), (5, 13, 4))
        self.assertEqual(s["Time away (outpass + tea)"], t["totalAwayMinutes"])
        self.assertEqual(s["Most time away"], "T001 Anita Kumar")

    def test_unmatched_qr_rows_and_mirror_rows_are_not_counted(self):
        e1 = next(r for r in self.rows(self.RID, SEPT) if r["employeeCode"] == "T001")
        self.assertEqual(e1["qrExits"], 2)  # not 3 (mirror row) and the unmatched "Stranger" belongs to nobody
        self.assertNotIn("ZZ99", json.dumps(self.rows(self.RID, SEPT)))

    def test_filters(self):
        def codes(**p):
            return [r["employeeCode"] for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(codes(departmentIds=self.d_cut.id), ["T001", "T002"])
        self.assertEqual(codes(employmentType="production"), ["T002", "T005"])
        self.assertEqual(codes(employeeStatus="active"), ["T001", "T002", "T003", "T004"])
        self.assertEqual(codes(employeeIds=self.e5.id), ["T005"])
        self.assertEqual(codes(dateFrom="2026-09-08", dateTo="2026-09-08"), ["T002", "T003"])  # T002 pass out 10:00, T003 break 09:00
        body = self.get_report(self.RID, {**SEPT, "excludeSuspect": "true"})
        rows = {r["employeeCode"]: r for r in body["rows"]}
        self.assertEqual(rows["T003"]["teaBreaks"], 1)  # the 150-minute break is left out

    def test_branch_isolation(self):
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, SEPT, self.b1_user)], ["T001", "T002", "T003", "T005"])
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        self.assertEqual(self.rows(self.RID, {**SEPT, "employeeIds": self.e4.id}, self.b1_user), [])
        self.assertEqual([r["employeeCode"] for r in self.rows(self.RID, SEPT, self.b2_user)], ["T004"])


class DeviceRegisterTests(_Base):
    RID = "gate-device-register"

    def test_golden(self):
        body = self.get_report(self.RID, SEPT)
        rows = body["rows"]
        self.assertEqual([r["name"] for r in rows], ["Gate 1", "Gate 2", "Gate 3", "Reception 1", "Reception 2"])
        g1, g2, g3, r1, r2 = rows
        self.assertEqual((g1["deviceType"], g1["branch"], g1["username"], g1["status"], g1["createdBy"]), ("Gate scanner", "Unit 1", "gate1", "Active", "Admin"))
        self.assertEqual(g1["lastLoginAt"], "2026-09-10 13:30")
        self.assertEqual((g1["outpassScans"], g1["teaScans"]), (3, 10))  # 4th scan is in August; tea: 7 OUT + 3 IN
        self.assertEqual((g2["status"], g2["lastLoginAt"], g2["outpassScans"], g2["teaScans"]), ("Deactivated", None, 0, 8))
        self.assertEqual((g3["outpassScans"], g3["teaScans"]), (1, 5))
        self.assertEqual((r1["deviceType"], r1["outpassScans"], r1["teaScans"]), ("Reception desk", None, None))  # desks do not scan
        self.assertEqual((r2["status"], r2["lastLoginAt"]), ("Deactivated", None))
        self.assertEqual(body["totals"]["outpassScans"], 4)
        self.assertEqual(body["totals"]["teaScans"], 23)
        s = summary_of(body)
        self.assertEqual((s["Devices"], s["Active"], s["Deactivated"], s["Never logged in"]), (5, 3, 2, 2))
        self.assertEqual((s["Outpass scans in period"], s["Tea-break scans in period"]), (4, 23))
        self.assertEqual(s["Outpass scans in period"], body["totals"]["outpassScans"])

    def test_secrets_never_appear(self):
        body = self.get_report(self.RID, SEPT)
        text = json.dumps(body)
        for secret in ("hash-secret", "tok-gate1", "tok-rec1", "password_hash", "login_token"):
            self.assertNotIn(secret, text)
        ws = load_workbook(io.BytesIO(self.export(self.RID, "xlsx", SEPT).content)).active
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertNotIn("hash-secret", cells)
        self.assertNotIn("tok-", cells)

    def test_filters(self):
        def names(**p):
            return [r["name"] for r in self.rows(self.RID, {**SEPT, **p})]

        self.assertEqual(names(deviceType="gate"), ["Gate 1", "Gate 2", "Gate 3"])
        self.assertEqual(names(deviceType="reception"), ["Reception 1", "Reception 2"])
        self.assertEqual(names(deviceStatus="active"), ["Gate 1", "Gate 3", "Reception 1"])
        self.assertEqual(names(deviceStatus="deactivated"), ["Gate 2", "Reception 2"])
        self.assertEqual(names(branchIds=self.b2.id), ["Gate 3", "Reception 2"])
        self.assertEqual(names(branchIds=self.b_empty.id), [])
        window = self.rows(self.RID, {"dateFrom": "2026-08-30", "dateTo": "2026-08-30"})
        self.assertEqual(window[0]["outpassScans"], 1)  # the August scan

    def test_branch_isolation(self):
        self.assertEqual([r["name"] for r in self.rows(self.RID, SEPT, self.b1_user)], ["Gate 1", "Gate 2", "Reception 1"])
        self.assertEqual([r["name"] for r in self.rows(self.RID, SEPT, self.b2_user)], ["Gate 3", "Reception 2"])
        self.assertEqual(self.rows(self.RID, {**SEPT, "branchIds": self.b2.id}, self.b1_user), [])
        b1 = self.get_report(self.RID, SEPT, self.b1_user)
        self.assertEqual(summary_of(b1)["Outpass scans in period"], 3)  # Gate 3's denied scan is not visible to Unit 1
        self.assertEqual(summary_of(b1)["Tea-break scans in period"], 18)


# ═════════════════════════════════════════════════════════════════════════════
# SQL <-> Python parity of the tea-break rules
# ═════════════════════════════════════════════════════════════════════════════


class TeaRuleParityTests(_Base):
    """The summaries aggregate in SQL and the registers compute per row in Python; both must equal the HR page."""

    DURATIONS = [
        0, 29, 30, 31, 59, 60, 89, 90, 91, 149, 150, 151, 14 * 60 + 29, 14 * 60 + 30, 14 * 60 + 31, 15 * 60 + 29, 15 * 60 + 30,
        15 * 60 + 31, 16 * 60 + 29, 16 * 60 + 30, 16 * 60 + 31, 17 * 60 + 30, 59 * 60 + 30, 60 * 60, 60 * 60 + 29, 60 * 60 + 30,
        60 * 60 + 31, 61 * 60 + 30, 200 * 60,
    ]
    OPEN_AGES = [1800, 3599, 3600, 3629, 3630, 3631, 3660, 12 * 3600 - 1, 12 * 3600, 12 * 3600 + 1, 13 * 3600]

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.par = Employee.objects.create(employee_code="PAR001", first_name="Par", last_name="Ity")
        base = ist(9, 20, 0, 0)
        for i, secs in enumerate(cls.DURATIONS):
            out = base + timedelta(minutes=i)
            TeaBreakLog.objects.create(employee=cls.par, out_at=out, in_at=out + timedelta(seconds=secs))
        for age in cls.OPEN_AGES:
            TeaBreakLog.objects.create(employee=cls.par, out_at=NOW - timedelta(seconds=age))

    def logs(self):
        return TeaBreakLog.objects.filter(employee=self.par)

    def test_python_rules_equal_the_hr_page(self):
        for allowed in (15, 16):
            for log in self.logs().select_related("employee__department"):
                page = _log_json(log, allowed, NOW)
                taken, remark = C.tea_metrics(log.out_at, log.in_at, allowed, NOW)
                self.assertEqual((taken, remark), (page["takenMinutes"], page["remark"]), (allowed, log.out_at, log.in_at))

    def test_sql_filters_equal_python_remarks(self):
        for allowed in (15, 16):
            py = {log.id: C.tea_metrics(log.out_at, log.in_at, allowed, NOW)[1] for log in self.logs()}
            for remark in ("on_time", "overtime", "not_returned", "in_progress"):
                sql_ids = set(self.logs().filter(C.remark_q(remark, allowed, NOW)).values_list("id", flat=True))
                self.assertEqual(sql_ids, {i for i, r in py.items() if r == remark}, (allowed, remark))

    def test_sql_minutes_round_half_even_like_python(self):
        done = [log for log in self.logs() if log.in_at]
        expected = [C.tea_metrics(log.out_at, log.in_at, 15, NOW)[0] for log in done]
        agg = self.logs().aggregate(s=Sum(C.break_minutes(), filter=C.completed_q()), m=Max(C.break_minutes(), filter=C.completed_q()))
        self.assertEqual(agg["s"], sum(expected))
        self.assertEqual(agg["m"], max(expected))
        # every duration one by one - including the exact half minutes (15m30s -> 16, 16m30s -> 16, 60m30s -> 60), where
        # a plain SQL ROUND would round them all up
        halves = 0
        for log in done:
            secs = (log.in_at - log.out_at).total_seconds()
            got = self.logs().filter(pk=log.pk).aggregate(m=Max(C.break_minutes()))["m"]
            self.assertEqual(got, round(secs / 60), secs)
            halves += secs % 60 == 30
        self.assertGreaterEqual(halves, 8)

    def test_suspect_partition_matches_flags(self):
        every = set(self.logs().values_list("id", flat=True))
        suspect = set(self.logs().filter(C.suspect_q(NOW)).values_list("id", flat=True))
        clean = set(self.logs().filter(C.not_suspect_q(NOW)).values_list("id", flat=True))
        self.assertEqual(suspect | clean, every)
        self.assertFalse(suspect & clean)
        for log in self.logs():
            taken, _ = C.tea_metrics(log.out_at, log.in_at, 15, NOW)
            self.assertEqual(C.tea_flag(log.out_at, log.in_at, taken, NOW) is not None, log.id in suspect, (log.out_at, log.in_at))

    def test_register_summary_uses_the_same_numbers(self):
        p = {"dateFrom": "2026-09-01", "dateTo": "2026-09-30", "employeeIds": self.par.id}
        body = self.get_report("tea-break-register", p)
        s = summary_of(body)
        py = [(log, *C.tea_metrics(log.out_at, log.in_at, 15, NOW)) for log in self.logs()]
        self.assertEqual(s["Overtime"], sum(1 for _, _, r in py if r == "overtime"))
        self.assertEqual(s["Not returned"], sum(1 for _, _, r in py if r == "not_returned"))
        self.assertEqual(s["Total time taken"], sum(t for log, t, _ in py if log.in_at))
        self.assertEqual(sum(1 for r in body["rows"] if r["remark"] == "Overtime"), s["Overtime"])
        self.assertEqual(body["totals"]["takenMinutes"], s["Total time taken"])
        self.assertEqual(s["Breaks"], len(py))


# ═════════════════════════════════════════════════════════════════════════════
# Cross-cutting contract: metadata, exports, empty results, permissions, query counts
# ═════════════════════════════════════════════════════════════════════════════


class ContractTests(_Base):
    def params_for(self, rid, empty=False):
        if not empty:
            # nobody had MORE than the default 2 breaks in a day, so ask for more-than-1 to get rows
            return {**SEPT, "minBreaksPerDay": 1} if rid == "tea-break-frequency" else dict(SEPT)
        p = {"dateFrom": "2025-01-01", "dateTo": "2025-01-31"}
        if rid == "gate-device-register":
            p["branchIds"] = self.b_empty.id  # a device list is not date-based
        return p

    def test_registered_metadata(self):
        registered = {s.id: s for s in registry.all_specs()}
        for rid in MY_IDS:
            self.assertIn(rid, registered)
            spec = registered[rid]
            self.assertEqual(spec.category, "gate", rid)
            self.assertEqual(spec.modules, ("outpass_visitors",), rid)
            self.assertFalse(spec.super_admin_only)
            for m in spec.modules:
                self.assertIn(m, all_module_keys())
            self.assertTrue(spec.description and spec.title)
        fam = lambda name: sorted((s.variant, s.id) for s in registered.values() if s.family == name)  # noqa: E731
        self.assertEqual(fam("visitors"), [("By host", "visitor-host-summary"), ("Daily", "visitor-daily-summary"), ("Register", "visitor-register")])
        self.assertEqual(fam("tea-break"), [("By department", "tea-break-department-summary"), ("By employee", "tea-break-employee-summary"), ("Register", "tea-break-register")])
        self.assertEqual(registry.LOAD_ERRORS, {})

    def test_every_report_runs_and_exports_with_valid_signatures(self):
        for rid in MY_IDS:
            p = self.params_for(rid)
            body = self.get_report(rid, p)
            xlsx = self.export(rid, "xlsx", p)
            self.assertEqual(xlsx.status_code, 200, rid)
            self.assertEqual(xlsx.content[:2], b"PK", rid)
            pdf = self.export(rid, "pdf", p)
            self.assertEqual(pdf.status_code, 200, rid)
            self.assertTrue(pdf.content.startswith(b"%PDF"), rid)
            # xlsx round trip: header row and the first text cell
            ws = load_workbook(io.BytesIO(xlsx.content)).active
            self.assertEqual([c.value for c in ws[7]][: len(body["columns"])], [c["label"] for c in body["columns"]], rid)
            self.assertTrue(body["rows"], rid)
            first = body["rows"][0]
            idx, col = next((i, c) for i, c in enumerate(body["columns"]) if c["type"] == "text" and first.get(c["key"]))
            self.assertEqual(ws.cell(row=8, column=idx + 1).value, first[col["key"]], rid)

    def test_empty_results_work_on_screen_and_in_both_exports(self):
        for rid in MY_IDS:
            p = self.params_for(rid, empty=True)
            body = self.get_report(rid, p)
            if rid in ("visitor-daily-summary", "tea-break-daily-trend", "tea-break-department-summary"):
                # calendar / department listings keep their zero rows (a department with no scans is the finding)
                key = "visits" if rid == "visitor-daily-summary" else "breaks"
                self.assertTrue(all(r[key] == 0 for r in body["rows"]), rid)
            else:
                self.assertEqual(body["rows"], [], rid)
            for fmt, sig in (("xlsx", b"PK"), ("pdf", b"%PDF")):
                r = self.export(rid, fmt, p)
                self.assertEqual(r.status_code, 200, (rid, fmt))
                self.assertTrue(r.content.startswith(sig), (rid, fmt))

    def test_totals_row_equals_sum_of_rows_for_every_additive_column(self):
        for rid in MY_IDS:
            body = self.get_report(rid, self.params_for(rid))
            wanted = [c for c in body["columns"] if c["total"] == "sum"]
            for c in wanted:
                vals = [r[c["key"]] for r in body["rows"] if r.get("_kind") not in ("subtotal", "total") and isinstance(r.get(c["key"]), (int, float))]
                self.assertEqual(body["totals"][c["key"]], round(sum(vals), 2) if c["type"] not in ("integer", "minutes", "duration") else sum(vals), (rid, c["key"]))

    def test_role_without_the_owning_module_gets_403_report_forbidden(self):
        for rid in MY_IDS:
            for url, extra in (("run", {}), ("export", {"fmt": "xlsx"})):
                r = self.client.get(f"/api/reports/{url}/{rid}", {**SEPT, **extra}, **headers(self.plain_user))
                self.assertEqual(r.status_code, 403, (rid, url))
                self.assertEqual(r.json()["error"], "report_forbidden", (rid, url))
        catalog = self.client.get("/api/reports/catalog", **headers(self.plain_user)).json()
        self.assertFalse({r["id"] for r in catalog["reports"]} & set(MY_IDS))
        allowed = self.client.get("/api/reports/catalog", **headers(self.b1_user)).json()
        self.assertTrue(set(MY_IDS) <= {r["id"] for r in allowed["reports"]})

    def test_branch_scoped_user_gets_no_branch_filter_in_the_catalog(self):
        catalog = self.client.get("/api/reports/catalog", **headers(self.b1_user)).json()
        for r in catalog["reports"]:
            if r["id"] in MY_IDS:
                self.assertNotIn("branch", [f["key"] for f in r["filters"]], r["id"])

    def _scale(self, start, stop):
        depts = [self.d_cut, self.d_sew, self.d_pack]
        for i in range(start, stop):
            e = Employee.objects.create(
                employee_code=f"S{i:03d}", first_name=f"Scale{i}", last_name="Emp", department=depts[i % 3],
                designation=self.des_op if i % 2 else self.des_help, branch=self.b1, employment_type="staff" if i % 2 else "production",
            )
            TeaBreakLog.objects.create(employee=e, out_gate=self.g1, in_gate=self.g2, out_at=ist(9, 10, 9, i % 50), in_at=ist(9, 10, 9, 20 + i % 50))
            TeaBreakLog.objects.create(employee=e, out_gate=self.g1, out_at=ist(9, 15, 8, i % 50))
            v = Visitor.objects.create(name=f"Scale Visitor {i}", phone=f"8{i:09d}", aadhaar_number=f"5555666677{i:02d}")
            vv = VisitorVisit.objects.create(
                visitor=v, branch=self.b1, whom_to_meet=f"Host {i}", purpose="x", meeting_employee=e,
                notified_email_at=NOW if i % 2 else None,
            )
            VisitorVisit.objects.filter(pk=vv.pk).update(visited_at=ist(9, 10, 10, i % 50))
            OutpassRequest.objects.create(
                employee=e, destination="d", reason="r", status="approved", exit_gate=self.g1, exited_at=ist(9, 11, 10, 0),
                entry_gate=self.g1, entered_at=ist(9, 11, 10, 20 + i % 30),
            )
            rec = OutpassRecord.objects.create(branch=self.b1, employee=e, employee_name="n", employee_code="c", destination="d")
            OutpassRecord.objects.filter(pk=rec.pk).update(submitted_at=ist(9, 12, 10, 0))

    def _queries(self, rid):
        with CaptureQueriesContext(connection) as cap:
            r = self.client.get(f"/api/reports/run/{rid}", SEPT, **headers(self.admin))
        self.assertEqual(r.status_code, 200, rid)
        return len(cap)

    def test_query_count_does_not_grow_with_the_number_of_employees(self):
        self._scale(0, 3)
        small = {rid: self._queries(rid) for rid in MY_IDS}
        self._scale(3, 15)
        for rid in MY_IDS:
            self.assertLessEqual(abs(self._queries(rid) - small[rid]), 2, (rid, small[rid]))

    def test_row_limit_is_applied_in_the_query_and_exports_are_refused_not_cut(self):
        with mock.patch("api.reporting.runner.SCREEN_ROW_LIMIT", 3), mock.patch("api.reporting.runner.XLSX_ROW_LIMIT", 3):
            with CaptureQueriesContext(connection) as cap:
                body = self.get_report("visitor-register", SEPT)
            self.assertEqual((body["rowCount"], body["truncated"], body["limit"]), (3, True, 3))
            self.assertTrue(any("LIMIT 4" in q["sql"] for q in cap.captured_queries))  # limit + 1, in SQL
            self.assertEqual(summary_of(body)["Visits"], 7)  # the summary still covers the whole filtered set
            body = self.get_report("tea-break-register", SEPT)
            self.assertEqual((body["rowCount"], body["truncated"]), (3, True))
            self.assertEqual(summary_of(body)["Breaks"], 13)
            self.assertIsNone(body["totals"])  # a sum over a cut-off list would under-state the truth
            r = self.export("tea-break-register", "xlsx", SEPT)
            self.assertEqual(r.status_code, 413)
            self.assertEqual(r.json()["error"], "too_many_rows")

    def test_screen_runs_never_write(self):
        TeaBreakRule.objects.all().delete()
        models = (TeaBreakRule, TeaBreakLog, VisitorVisit, Visitor, GateDevice, ReceptionDevice, OutpassRequest, OutpassRecord, OutpassGateScan, Employee, AuditLog)
        before = {m.__name__: m.objects.count() for m in models}
        for rid in MY_IDS:
            self.get_report(rid, SEPT)
        self.assertEqual({m.__name__: m.objects.count() for m in models}, before)
        self.assertEqual(TeaBreakRule.objects.count(), 0)

    def test_pdf_and_screen_never_receive_free_text_as_markup_or_formulas(self):
        Visitor.objects.filter(pk=self.v1.pk).update(name='=HYPERLINK("http://evil","x")')
        VisitorVisit.objects.filter(pk=self.vis1.pk).update(purpose="<b>bold</b> & more")
        self.assertEqual(self.export("visitor-register", "pdf", SEPT).status_code, 200)
        ws = load_workbook(io.BytesIO(self.export("visitor-register", "xlsx", SEPT).content)).active
        cell = next(c for row in ws.iter_rows(min_row=8, max_row=8) for c in row if str(c.value).startswith("=HYPERLINK"))
        self.assertEqual(cell.data_type, "s")  # stored as text, never evaluated

    def test_unauthenticated_and_employee_tokens_are_refused(self):
        self.assertEqual(self.client.get("/api/reports/run/visitor-register", SEPT).status_code, 401)
        emp = {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'employee', 'employeeId': self.e1.id})}"}
        self.assertEqual(self.client.get("/api/reports/run/tea-break-register", SEPT, **emp).status_code, 403)
