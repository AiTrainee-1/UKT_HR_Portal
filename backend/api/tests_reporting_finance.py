"""
Report Center - finance reports: advance-ledger, advance-recovery-schedule, bonus-register,
increment-history, promotion-history.

Fixtures are fixed-date and "today" is pinned to 2026-09-15 (IST), so every expected figure below is
hand-derived and nothing depends on the day the suite runs. Where a report describes rows the app itself
creates (advances, payroll, bonus, increments, promotions) a test drives the real endpoint / engine and
checks that the report agrees with it.

Run via: python manage.py test api.tests_reporting_finance -v 2
"""

import io
import itertools
import json
from datetime import date, datetime, timezone as dt_timezone
from decimal import Decimal
from unittest import mock

from django.db import connection
from django.test import SimpleTestCase, TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from .clock import FACTORY_TZ
from .jwt_utils import sign_token
from .models import (
    Advance,
    AdvanceRepayment,
    AttendanceDayRecord,
    Bonus,
    Branch,
    Department,
    Designation,
    Employee,
    HRUser,
    PayrollSettings,
    Promotion,
    Role,
    SalaryIncrement,
    SalarySlip,
)
from .payroll_views import _build_working_days, _generate_production_payroll, _generate_staff_payroll
from .permission_registry import all_module_keys
from .reporting import registry
from .reporting.definitions import finance_common as C
from .reporting.formatting import fmt_dt

UTC = dt_timezone.utc
TODAY = date(2026, 9, 15)

ADV_LEDGER = "advance-ledger"
ADV_SCHEDULE = "advance-recovery-schedule"
BONUS = "bonus-register"
INCREMENT = "increment-history"
PROMOTION = "promotion-history"
ALL_IDS = (ADV_LEDGER, ADV_SCHEDULE, BONUS, INCREMENT, PROMOTION)
OWNING_MODULE = {
    ADV_LEDGER: "settlement",
    ADV_SCHEDULE: "settlement",
    BONUS: "bonus",
    INCREMENT: "increment",
    PROMOTION: "promotion",
}
# Params that make each report return the fixture data (the schedule needs an explicit month range).
FULL_PARAMS = {
    ADV_LEDGER: {},
    ADV_SCHEDULE: {"dateFrom": "2026-06-01", "dateTo": "2026-09-30"},
    BONUS: {},
    INCREMENT: {},
    PROMOTION: {},
}


def ist(y, m, d, h=10, mi=0):
    return datetime(y, m, d, h, mi, tzinfo=FACTORY_TZ)


def utc(y, m, d, h=0, mi=0):
    return datetime(y, m, d, h, mi, tzinfo=UTC)


def headers(user: HRUser) -> dict:
    return {"HTTP_AUTHORIZATION": f"Bearer {sign_token({'role': 'hr', 'hrUserId': user.id})}"}


# ── fixtures ─────────────────────────────────────────────────────────────────


def make_org(cls):
    cls.b1 = Branch.objects.create(name="Unit 1", code="FN1")
    cls.b2 = Branch.objects.create(name="Unit 2", code="FN2")
    cls.cutting = Department.objects.create(name="CUTTING", branch=cls.b1)
    cls.packing = Department.objects.create(name="PACKING", branch=cls.b1)
    cls.sewing = Department.objects.create(name="SEWING", branch=cls.b2)
    cls.d_op = Designation.objects.create(title="Operator")
    cls.d_sup = Designation.objects.create(title="Supervisor")
    cls.d_mgr = Designation.objects.create(title="Manager")

    def mk(code, first, dept, branch, desig, **kw):
        return Employee.objects.create(
            employee_code=code,
            first_name=first,
            last_name="Fin",
            department=dept,
            designation=desig,
            branch=branch,
            **kw,
        )

    cls.e1 = mk("F1", "Asha", cls.cutting, cls.b1, cls.d_op, employment_type="staff")
    cls.e2 = mk("F2", "Bala", cls.cutting, cls.b1, cls.d_op, employment_type="production")
    cls.e3 = mk("F10", "Chitra", cls.sewing, cls.b2, cls.d_op, employment_type="staff")
    cls.e4 = mk("F3", "Dev", cls.packing, cls.b1, cls.d_op, employment_type="staff", status="inactive")
    cls.e5 = mk("F4", "Esha", None, None, None, employment_type="staff")  # NULL department / designation / branch
    cls.e_none = mk("F99", "Nobody", cls.cutting, cls.b1, cls.d_op, employment_type="staff")  # no finance data at all

    cls.admin = HRUser.objects.create(username="fin_admin", password_hash="x", is_super_admin=True)

    def user(name, perms, **kw):
        return HRUser.objects.create(
            username=name, password_hash="x", role=Role.objects.create(name=name, permissions=perms), **kw
        )

    everything = {"reports": "view", "settlement": "view", "bonus": "view", "increment": "view", "promotion": "view"}
    cls.b1_user = user("fin_b1", everything, branch=cls.b1)
    cls.reports_only = user("fin_reports_only", {"reports": "view"})
    cls.by_module = {m: user(f"fin_only_{m}", {"reports": "view", m: "view"}) for m in set(OWNING_MODULE.values())}


def make_advance(
    emp,
    kind,
    amount,
    status,
    created,
    *,
    approved_at=None,
    approved_by=None,
    emi="0",
    months=None,
    purpose=None,
    start=None,
    stored_repaid="0",
    stored_outstanding=None,
    reps=(),
):
    """An Advance row as the app leaves it (stored balances included) plus its instalment rows
    ``reps = [(year, month, amount, processed), ...]``."""
    a = Advance.objects.create(
        employee=emp,
        advance_type=kind,
        amount=Decimal(amount),
        status=status,
        approved_by=approved_by,
        approved_at=approved_at,
        emi_amount=Decimal(emi),
        repayment_months=months,
        purpose=purpose,
        repayment_start_month=start[0] if start else None,
        repayment_start_year=start[1] if start else None,
        total_repaid=Decimal(stored_repaid),
        outstanding=Decimal(stored_outstanding if stored_outstanding is not None else amount),
    )
    Advance.objects.filter(pk=a.pk).update(created_at=created)
    for y, m, amt, done in reps:
        AdvanceRepayment.objects.create(advance=a, year=y, month=m, amount=Decimal(amt), is_processed=done)
    return a


def make_advances(cls):
    """Eight advances covering every status / type / overdue / mismatch case. 'Today' is 2026-09-15."""
    cls.a1 = make_advance(  # term loan, 2 of 4 instalments deducted, August still open = overdue
        cls.e1,
        "term",
        "12000",
        "approved",
        ist(2026, 5, 25, 10),
        approved_at=utc(2026, 5, 28, 20, 0),
        approved_by="Payroll Head",
        emi="3000",
        months=4,
        purpose="House repair",
        start=(6, 2026),
        stored_repaid="6000",
        stored_outstanding="6000",
        reps=[(2026, 6, 3000, True), (2026, 7, 3000, True), (2026, 8, 3000, False), (2026, 9, 3000, False)],
    )
    cls.a2 = make_advance(  # general, requested 00:30 IST on 1 Sep (= 31 Aug in UTC)
        cls.e2,
        "general",
        "5000",
        "approved",
        ist(2026, 9, 1, 0, 30),
        approved_at=utc(2026, 9, 2, 4, 0),
        approved_by="Payroll Head",
        purpose="Medical",
        stored_repaid="0",
        stored_outstanding="5000",
        reps=[(2026, 9, 5000, False)],
    )
    cls.a3 = make_advance(  # fully recovered and closed
        cls.e3,
        "general",
        "2000",
        "closed",
        ist(2026, 7, 20, 11),
        approved_at=utc(2026, 7, 21, 5, 0),
        approved_by="Payroll Head",
        stored_repaid="2000",
        stored_outstanding="0",
        reps=[(2026, 8, 2000, True)],
    )
    cls.a4 = make_advance(  # still awaiting approval (stored outstanding = amount, as the app writes it)
        cls.e1,
        "general",
        "1000",
        "pending",
        ist(2026, 9, 10, 15),
        stored_outstanding="1000",
    )
    cls.a5 = make_advance(  # employee has left; August instalment never deducted
        cls.e4,
        "general",
        "4000",
        "approved",
        ist(2026, 7, 30, 9),
        approved_at=utc(2026, 7, 31, 5, 0),
        approved_by="Payroll Head",
        stored_repaid="0",
        stored_outstanding="4000",
        reps=[(2026, 8, 4000, False)],
    )
    cls.a6 = make_advance(  # rejected term loan
        cls.e2,
        "term",
        "900",
        "rejected",
        ist(2026, 8, 5, 9),
        approved_by="HR Admin",
        stored_outstanding="900",
    )
    cls.a7 = make_advance(  # instalment deducted but the stored balance was never updated
        cls.e5,
        "general",
        "700",
        "approved",
        ist(2026, 9, 3, 9),
        approved_at=utc(2026, 9, 3, 6, 0),
        approved_by="Payroll Head",
        stored_repaid="0",
        stored_outstanding="700",
        reps=[(2026, 9, 700, True)],
    )
    cls.a8 = make_advance(  # closed by hand with nothing recovered
        cls.e3,
        "general",
        "1500",
        "closed",
        ist(2026, 6, 1, 9),
        approved_at=utc(2026, 6, 2, 5, 0),
        approved_by="Payroll Head",
        stored_repaid="0",
        stored_outstanding="1500",
    )


def make_bonus_rows(cls):
    def mk(emp, fy, records, base, pct, amount, status, when, by="HR Admin", notes=None):
        b = Bonus.objects.create(
            employee=emp,
            financial_year=fy,
            records_considered=records,
            calculation_base=Decimal(base),
            bonus_percent_applied=Decimal(pct),
            bonus_amount=Decimal(amount),
            status=status,
            computed_by=by,
            notes=notes,
        )
        Bonus.objects.filter(pk=b.pk).update(created_at=when)
        return b

    cls.bn1 = mk(cls.e1, "2025-26", 12, "84000.00", "8.33", "6997.20", "approved", ist(2026, 4, 20, 10))
    cls.bn2 = mk(
        cls.e2, "2025-26", 50, "30000.00", "8.33", "2499.00", "paid", utc(2026, 4, 21, 19, 0), notes="Weekly slips"
    )
    cls.bn3 = mk(cls.e3, "2025-26", 8, "56000.00", "8.33", "4664.80", "calculated", ist(2026, 4, 20, 10))
    cls.bn4 = mk(cls.e1, "2024-25", 12, "84000.00", "8.33", "6997.20", "paid", ist(2025, 4, 18, 10))
    cls.bn5 = mk(cls.e4, "2025-26", 12, "60000.00", "8.33", "4998.00", "calculated", ist(2026, 4, 20, 10))


def make_increments(cls):
    def mk(emp, prev, new, pct, eff, by="HR Admin", when=None, notes=None):
        i = SalaryIncrement.objects.create(
            employee=emp,
            previous_salary=Decimal(prev),
            new_salary=Decimal(new),
            percent=Decimal(pct),
            effective_date=eff,
            added_by=by,
            notes=notes,
        )
        SalaryIncrement.objects.filter(pk=i.pk).update(created_at=when or ist(2026, 1, 1, 10))
        return i

    cls.i1 = mk(cls.e1, "20000", "22000", "10.00", date(2026, 4, 1), when=utc(2026, 3, 30, 22, 0), notes="Annual")
    cls.i2 = mk(cls.e1, "22000", "23100", "5.00", date(2026, 9, 1))
    cls.i3 = mk(cls.e3, "15000", "16500", "10.00", date(2025, 4, 1))
    cls.i4 = mk(cls.e4, "12000", "12600", "5.00", date(2026, 7, 15))
    cls.i5 = mk(cls.e5, "30000", "33000", "10.00", date(2026, 10, 1), by=None)  # future-dated, no recorder


def make_promotions(cls):
    def mk(emp, pd, pg, nd, ng, eff, by="HR Admin", when=None, notes=None):
        p = Promotion.objects.create(
            employee=emp,
            previous_department=pd,
            previous_designation=pg,
            new_department=nd,
            new_designation=ng,
            effective_date=eff,
            promoted_by=by,
            notes=notes,
        )
        Promotion.objects.filter(pk=p.pk).update(created_at=when or ist(2026, 1, 1, 10))
        return p

    cls.p1 = mk(
        cls.e1,
        cls.cutting,
        cls.d_op,
        cls.cutting,
        cls.d_sup,
        date(2026, 4, 1),
        when=utc(2026, 3, 31, 19, 0),
        notes="Line supervisor",
    )
    cls.p2 = mk(cls.e1, cls.cutting, cls.d_sup, cls.packing, cls.d_mgr, date(2026, 9, 1))
    cls.p3 = mk(cls.e3, cls.sewing, cls.d_op, cls.sewing, cls.d_sup, date(2025, 6, 1))
    cls.p4 = mk(cls.e4, cls.packing, cls.d_op, cls.cutting, cls.d_op, date(2026, 7, 1))
    cls.p5 = mk(cls.e5, None, None, cls.sewing, None, date(2026, 8, 10), by=None)  # previous department was deleted


class _Case(TestCase):
    """Org + users, 'today' pinned, and request helpers."""

    TODAY = TODAY

    @classmethod
    def setUpTestData(cls):
        make_org(cls)

    def setUp(self):
        patcher = mock.patch("api.reporting.filters.ist_today", return_value=self.TODAY)
        patcher.start()
        self.addCleanup(patcher.stop)

    def get(self, kind, rid, user=None, **params):
        return self.client.get(f"/api/reports/{kind}/{rid}", params, **headers(user or self.admin))

    def run_report(self, rid, user=None, **params):
        r = self.get("run", rid, user, **params)
        self.assertEqual(r.status_code, 200, r.content[:400])
        return r.json()

    def export(self, rid, fmt, user=None, **params):
        return self.get("export", rid, user, fmt=fmt, **params)

    @staticmethod
    def data(body):
        return [r for r in body["rows"] if "_kind" not in r]

    @staticmethod
    def subtotals(body):
        return [r for r in body["rows"] if r.get("_kind") == "subtotal"]

    @staticmethod
    def cards(body):
        return {c["label"]: c["value"] for c in body["summary"]}

    @staticmethod
    def col_sum(body, key):
        return round(sum(r[key] for r in _Case.data(body) if r.get(key) is not None), 2)

    def row(self, body, code, amount=None, key="amount"):
        found = [r for r in self.data(body) if r["employeeCode"] == code and (amount is None or r.get(key) == amount)]
        self.assertEqual(len(found), 1, f"{code}/{amount}: {[r['employeeCode'] for r in self.data(body)]}")
        return found[0]

    def codes(self, body):
        return [r["employeeCode"] for r in self.data(body)]


# ═════════════════════════════════════════════════════════════════════════════
#  Pure helpers
# ═════════════════════════════════════════════════════════════════════════════


class FinanceCommonTests(SimpleTestCase):
    def test_window_bounds_for_every_token(self):
        t = date(2026, 9, 15)
        expect = {
            "thisMonth": (date(2026, 9, 1), date(2026, 9, 30)),
            "lastMonth": (date(2026, 8, 1), date(2026, 8, 31)),
            "last3": (date(2026, 7, 1), date(2026, 9, 30)),
            "last6": (date(2026, 4, 1), date(2026, 9, 30)),
            "last12": (date(2025, 10, 1), date(2026, 9, 30)),
            "thisFY": (date(2026, 4, 1), date(2027, 3, 31)),
            "lastFY": (date(2025, 4, 1), date(2026, 3, 31)),
            "thisYear": (date(2026, 1, 1), date(2026, 12, 31)),
            "lastYear": (date(2025, 1, 1), date(2025, 12, 31)),
        }
        self.assertEqual({v for v, _ in C.WINDOW_OPTIONS}, set(expect))
        for token, bounds in expect.items():
            self.assertEqual(C.window_bounds(token, t), bounds, token)
        for empty in (None, "", "nonsense"):
            self.assertEqual(C.window_bounds(empty, t), (None, None))

    def test_windows_cross_year_and_financial_year_boundaries(self):
        self.assertEqual(C.window_bounds("lastMonth", date(2026, 1, 5)), (date(2025, 12, 1), date(2025, 12, 31)))
        self.assertEqual(C.window_bounds("last3", date(2026, 2, 10)), (date(2025, 12, 1), date(2026, 2, 28)))
        self.assertEqual(C.window_bounds("thisFY", date(2026, 2, 10)), (date(2025, 4, 1), date(2026, 3, 31)))
        self.assertEqual(C.window_bounds("lastFY", date(2026, 2, 10)), (date(2024, 4, 1), date(2025, 3, 31)))
        self.assertEqual(C.window_bounds("thisFY", date(2026, 4, 1)), (date(2026, 4, 1), date(2027, 3, 31)))
        self.assertEqual(C.window_bounds("thisFY", date(2026, 3, 31)), (date(2025, 4, 1), date(2026, 3, 31)))

    def test_fy_bounds_and_options(self):
        self.assertEqual(C.fy_bounds("2025-26"), (date(2025, 4, 1), date(2026, 3, 31)))
        self.assertIsNone(C.fy_bounds("junk"))
        self.assertIsNone(C.fy_bounds("2025-2026"))
        values = [v for v, _ in C.FY_OPTIONS]
        self.assertIn("2025-26", values)
        self.assertEqual(values, sorted(values, reverse=True))  # newest first
        self.assertEqual(C.add_months(date(2026, 1, 1), -1), date(2025, 12, 1))
        self.assertEqual(C.add_months(date(2026, 11, 1), 3), date(2027, 2, 1))

    def test_natural_key_orders_numeric_looking_codes(self):
        codes = ["10", "2", "1", "A1", "30005", "73", "F10", "F2"]
        self.assertEqual(sorted(codes, key=C.natural_key), ["1", "2", "10", "73", "30005", "A1", "F2", "F10"])
        self.assertEqual(C.natural_key(None), [""])

    def test_ist_date_iso_uses_the_factory_day(self):
        self.assertEqual(C.ist_date_iso(utc(2026, 5, 28, 20, 0)), "2026-05-29")  # 01:30 IST next day
        self.assertEqual(C.ist_date_iso(utc(2026, 5, 28, 18, 29)), "2026-05-28")
        self.assertIsNone(C.ist_date_iso(None))

    def test_group_subtotals_sums_and_averages(self):
        rows = [
            {"g": "A", "amt": 10.0, "pct": 10.0},
            {"g": "A", "amt": 5.0, "pct": 5.0},
            {"g": "B", "amt": 1.0, "pct": None},
        ]
        out = C.group_subtotals(rows, lambda r: r["g"], sums=["amt"], avgs=["pct"], label_key="g")
        self.assertEqual([r.get("_kind") for r in out], [None, None, "subtotal", None, "subtotal"])
        self.assertEqual((out[2]["amt"], out[2]["pct"], out[2]["g"]), (15.0, 7.5, "A total"))
        self.assertEqual((out[4]["amt"], out[4]["pct"]), (1.0, None))


# ═════════════════════════════════════════════════════════════════════════════
#  advance-ledger
# ═════════════════════════════════════════════════════════════════════════════


class AdvanceLedgerTests(_Case):
    RID = ADV_LEDGER

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        make_advances(cls)

    def test_rows_are_in_natural_employee_order(self):
        body = self.run_report(self.RID)
        self.assertEqual(
            [(r["employeeCode"], r["amount"]) for r in self.data(body)],
            [
                ("F1", 12000.0),
                ("F1", 1000.0),
                ("F2", 900.0),
                ("F2", 5000.0),
                ("F3", 4000.0),
                ("F4", 700.0),
                ("F10", 1500.0),
                ("F10", 2000.0),
            ],  # F10 after F4, not after F1
        )

    def test_term_loan_with_an_overdue_instalment(self):
        r = self.row(self.run_report(self.RID), "F1", 12000.0)
        self.assertEqual(r["employeeName"], "Asha Fin")
        self.assertEqual((r["department"], r["empStatus"]), ("CUTTING", "Active"))
        self.assertEqual((r["advanceType"], r["status"]), ("Term Loan", "Approved"))
        self.assertEqual(r["requestedOn"], "2026-05-25")
        self.assertEqual(r["approvedOn"], "2026-05-29")  # approved 20:00 UTC on the 28th = 01:30 IST on the 29th
        self.assertEqual((r["approvedBy"], r["purpose"]), ("Payroll Head", "House repair"))
        self.assertEqual((r["instalments"], r["emi"]), ("2 of 4", 3000.0))
        self.assertEqual((r["repaid"], r["outstanding"], r["repaidPct"]), (6000.0, 6000.0, 50.0))
        self.assertEqual(r["nextDue"], "Aug 2026")
        # August is overdue; September (the current month) is merely due
        self.assertEqual((r["overdueInstalments"], r["overdueAmount"]), (1, 3000.0))

    def test_general_advance_due_this_month_is_not_overdue(self):
        r = self.row(self.run_report(self.RID), "F2", 5000.0)
        self.assertEqual(r["requestedOn"], "2026-09-01")  # 00:30 IST = 31 Aug UTC: must land on the IST day
        self.assertEqual(r["approvedOn"], "2026-09-02")
        self.assertEqual((r["instalments"], r["emi"]), ("0 of 1", None))
        self.assertEqual((r["repaid"], r["outstanding"], r["repaidPct"]), (0.0, 5000.0, 0.0))
        self.assertEqual((r["nextDue"], r["overdueInstalments"], r["overdueAmount"]), ("Sep 2026", 0, 0.0))

    def test_pending_and_rejected_requests_carry_no_balance(self):
        body = self.run_report(self.RID)
        for code, amount, status in (("F1", 1000.0, "Pending"), ("F2", 900.0, "Rejected")):
            r = self.row(body, code, amount)
            self.assertEqual(r["status"], status)
            # the stored outstanding equals the amount for every status - the report must not repeat that
            for key in (
                "repaid",
                "outstanding",
                "repaidPct",
                "instalments",
                "nextDue",
                "overdueInstalments",
                "overdueAmount",
            ):
                self.assertIsNone(r[key], f"{status}/{key}")
        self.assertIsNone(self.row(body, "F1", 1000.0)["approvedOn"])

    def test_closed_advances(self):
        body = self.run_report(self.RID)
        done = self.row(body, "F10", 2000.0)
        self.assertEqual(
            (done["status"], done["repaid"], done["outstanding"], done["repaidPct"]), ("Closed", 2000.0, 0.0, 100.0)
        )
        self.assertEqual((done["instalments"], done["nextDue"], done["overdueInstalments"]), ("1 of 1", None, None))
        manual = self.row(body, "F10", 1500.0)  # closed by hand, nothing recovered
        self.assertEqual(
            (manual["repaid"], manual["outstanding"], manual["repaidPct"], manual["instalments"]), (0.0, 0.0, 0.0, None)
        )

    def test_employee_who_left_and_employee_without_department(self):
        body = self.run_report(self.RID)
        left = self.row(body, "F3")
        self.assertEqual((left["empStatus"], left["department"]), ("Inactive", "PACKING"))
        self.assertEqual((left["outstanding"], left["overdueInstalments"], left["overdueAmount"]), (4000.0, 1, 4000.0))
        nobody = self.row(body, "F4")  # NULL department / designation / branch: listed for an unscoped admin
        self.assertEqual(nobody["department"], "Unassigned")

    def test_resigned_status_counts_as_left_too(self):
        # Employee.status is free text (active | inactive | resigned): anything but "active" is an ex-employee
        Employee.objects.filter(pk=self.e4.pk).update(status="resigned")
        body = self.run_report(self.RID)
        self.assertEqual(self.row(body, "F3")["empStatus"], "Resigned")
        self.assertEqual(self.cards(body)["Outstanding with ex-employees"], 4000.0)
        self.assertEqual(self.codes(self.run_report(self.RID, employeeStatus="inactive")), ["F3"])
        sched = self.run_report(ADV_SCHEDULE, dateFrom="2026-08-01", dateTo="2026-08-31")
        self.assertEqual(self.cards(sched)["Not yet recovered - ex-employees"], 4000.0)

    def test_deducted_instalment_wins_over_a_stale_stored_balance(self):
        r = self.row(self.run_report(self.RID), "F4")
        # stored total_repaid is 0 / outstanding 700, but payroll did deduct the 700 instalment
        self.assertEqual(
            (r["repaid"], r["outstanding"], r["repaidPct"], r["instalments"]), (700.0, 0.0, 100.0, "1 of 1")
        )
        self.assertEqual((r["nextDue"], r["overdueInstalments"]), (None, 0))

    def test_totals_summary_and_reconciliation(self):
        body = self.run_report(self.RID)
        self.assertEqual(
            body["totals"], {"repaid": 8700.0, "outstanding": 15000.0, "overdueInstalments": 2, "overdueAmount": 7000.0}
        )
        for key, total in body["totals"].items():
            self.assertEqual(self.col_sum(body, key), total, key)
        cards = self.cards(body)
        self.assertEqual(cards["Advances listed"], 8)
        self.assertEqual(cards["Sanctioned (approved + closed)"], 25200.0)
        self.assertEqual(cards["Recovered"], 8700.0)
        self.assertEqual(cards["Outstanding"], 15000.0)
        self.assertEqual(cards["Overdue instalments"], 2)
        self.assertEqual(cards["Overdue amount"], 7000.0)
        self.assertEqual(cards["Awaiting approval"], 1)
        self.assertEqual(cards["Outstanding with ex-employees"], 4000.0)
        # sanctioned = recovered + outstanding + what closed-by-hand advances never recovered (1,500)
        self.assertEqual(cards["Sanctioned (approved + closed)"], cards["Recovered"] + cards["Outstanding"] + 1500.0)

    def test_notes_flag_data_problems(self):
        notes = " ".join(self.run_report(self.RID)["notes"])
        self.assertIn("1 advance(s) store a 'total repaid' that differs", notes)
        self.assertIn("1 closed advance(s) were closed with less recovered than sanctioned", notes)
        self.assertIn("Rs. 1,500.00", notes)
        self.assertIn("September 2026", notes)  # the month overdue is measured against

    def test_status_mix_note_counts_what_is_listed(self):
        self.assertEqual(
            self.run_report(self.RID)["notes"][0], "Status mix: 4 Approved, 2 Closed, 1 Pending, 1 Rejected."
        )
        self.assertEqual(
            self.run_report(self.RID, status="pending,closed")["notes"][0], "Status mix: 2 Closed, 1 Pending."
        )
        empty = self.run_report(self.RID, employeeIds=str(self.e_none.id))
        self.assertFalse(any(n.startswith("Status mix") for n in empty["notes"]))

    def test_status_and_type_filters(self):
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, status="approved"))), ["F1", "F2", "F3", "F4"])
        body = self.run_report(self.RID, status="pending,rejected")
        self.assertEqual(
            [(r["employeeCode"], r["status"]) for r in self.data(body)], [("F1", "Pending"), ("F2", "Rejected")]
        )
        self.assertEqual(
            sorted(r["amount"] for r in self.data(self.run_report(self.RID, advanceType="term"))), [900.0, 12000.0]
        )
        self.assertEqual(len(self.data(self.run_report(self.RID, advanceType="general"))), 6)
        self.assertEqual(self.run_report(self.RID, status="closed")["totals"]["repaid"], 2000.0)

    def test_requested_window_uses_ist_days(self):
        def amounts(token):
            return sorted(r["amount"] for r in self.data(self.run_report(self.RID, requested=token)))

        self.assertEqual(amounts("thisMonth"), [700.0, 1000.0, 5000.0])  # incl. 1 Sep 00:30 IST (= 31 Aug UTC)
        self.assertEqual(amounts("lastMonth"), [900.0])
        self.assertEqual(amounts("last3"), [700.0, 900.0, 1000.0, 2000.0, 4000.0, 5000.0])
        self.assertEqual(len(amounts("thisFY")), 8)
        self.assertEqual(amounts("lastFY"), [])
        self.assertEqual(len(amounts("thisYear")), 8)
        body = self.run_report(self.RID, requested="lastMonth")
        self.assertIn("Requested", [f["label"] for f in body["filters"]])
        self.assertTrue(any("Only advances requested from 01-Aug-2026 to 31-Aug-2026" in n for n in body["notes"]))

    def test_window_edges_are_the_ist_midnights(self):
        make_advance(self.e1, "general", "111", "pending", ist(2026, 9, 30, 23, 59))  # last minute of September, IST
        make_advance(self.e1, "general", "222", "pending", ist(2026, 10, 1, 0, 1))  # first minutes of October, IST
        make_advance(self.e1, "general", "333", "pending", ist(2026, 8, 31, 23, 59))  # last minute of August, IST

        def amounts(token):
            return sorted(r["amount"] for r in self.data(self.run_report(self.RID, requested=token)))

        self.assertEqual(amounts("thisMonth"), [111.0, 700.0, 1000.0, 5000.0])
        self.assertEqual(amounts("lastMonth"), [333.0, 900.0])

    def test_overdue_only(self):
        body = self.run_report(self.RID, overdueOnly="true")
        self.assertEqual(sorted(self.codes(body)), ["F1", "F3"])
        self.assertEqual(body["totals"]["overdueInstalments"], 2)
        # a pending / rejected request can never be overdue, even with an old unprocessed row
        make_advance(self.e2, "general", "300", "pending", ist(2026, 1, 5), reps=[(2026, 2, 300, False)])
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, overdueOnly="true"))), ["F1", "F3"])

    def test_employee_scope_filters(self):
        self.assertEqual(
            sorted(self.codes(self.run_report(self.RID, departmentIds=str(self.cutting.id)))), ["F1", "F1", "F2", "F2"]
        )
        self.assertEqual(self.codes(self.run_report(self.RID, employmentType="production")), ["F2", "F2"])
        self.assertEqual(self.codes(self.run_report(self.RID, employeeIds=str(self.e4.id))), ["F3"])
        self.assertEqual(sorted(set(self.codes(self.run_report(self.RID, branchIds=str(self.b2.id))))), ["F10"])
        self.assertEqual(self.run_report(self.RID, designationIds=str(self.d_mgr.id))["rows"], [])

    def test_designation_filter(self):
        body = self.run_report(self.RID, designationIds=str(self.d_op.id))
        self.assertEqual(len(self.data(body)), 7)  # everyone except F4, who has no designation
        self.assertNotIn("F4", self.codes(body))
        Employee.objects.filter(pk=self.e2.pk).update(designation=self.d_sup)
        self.assertEqual(self.codes(self.run_report(self.RID, designationIds=str(self.d_sup.id))), ["F2", "F2"])

    def test_default_request_echoes_no_filters(self):
        body = self.run_report(self.RID)
        self.assertEqual(body["filters"], [])  # exports print "All records", not a stray "Order by"

    def test_approved_advance_without_a_schedule_is_flagged(self):
        # a term loan approved without an EMI amount gets no instalment rows, so payroll can never recover it
        make_advance(self.e1, "term", "8000", "approved", ist(2026, 9, 5), approved_at=utc(2026, 9, 6, 5, 0), emi="0")
        body = self.run_report(self.RID)
        r = self.row(body, "F1", 8000.0)
        self.assertEqual(
            (r["instalments"], r["nextDue"], r["overdueInstalments"], r["outstanding"]), (None, None, 0, 8000.0)
        )
        self.assertTrue(any("1 approved advance(s) have no instalment schedule" in n for n in body["notes"]))
        base = " ".join(self.run_report(self.RID, status="closed")["notes"])
        self.assertNotIn("no instalment schedule", base)

    def test_schedule_that_covers_less_than_the_sanctioned_amount_is_flagged(self):
        # 5,000 approved as a 3 x 1,000 plan: 2,000 will never be scheduled for recovery
        make_advance(
            self.e1,
            "term",
            "5000",
            "approved",
            ist(2026, 9, 5),
            approved_at=utc(2026, 9, 6, 5, 0),
            emi="1000",
            months=3,
            start=(9, 2026),
            reps=[(2026, 9, 1000, False), (2026, 10, 1000, False), (2026, 11, 1000, False)],
        )
        body = self.run_report(self.RID)
        r = self.row(body, "F1", 5000.0)
        self.assertEqual((r["instalments"], r["outstanding"], r["nextDue"]), ("0 of 3", 5000.0, "Sep 2026"))
        notes = " ".join(body["notes"])
        self.assertIn(
            "1 approved advance(s) have an instalment schedule that adds up to less than the sanctioned", notes
        )
        self.assertIn("Rs. 2,000.00 is not scheduled for recovery", notes)
        # the fixture plans are complete (a1: 4 x 3,000 = 12,000): without the new term loan there is no such note
        self.assertNotIn("adds up to less", " ".join(self.run_report(self.RID, advanceType="general")["notes"]))
        self.assertNotIn("adds up to less", " ".join(self.run_report(self.RID, status="closed,pending")["notes"]))

    def test_employee_status_filter_keeps_leavers_reachable(self):
        self.assertEqual(self.codes(self.run_report(self.RID, employeeStatus="inactive")), ["F3"])
        self.assertNotIn("F3", self.codes(self.run_report(self.RID, employeeStatus="active")))
        self.assertIn("F3", self.codes(self.run_report(self.RID)))  # default: everyone

    def test_sort_by_outstanding_puts_unknown_last(self):
        body = self.run_report(self.RID, sortBy="outstanding")
        self.assertEqual(
            [r["outstanding"] for r in self.data(body)], [6000.0, 5000.0, 4000.0, 0.0, 0.0, 0.0, None, None]
        )

    def test_sort_by_newest_request_first(self):
        body = self.run_report(self.RID, sortBy="newest")
        self.assertEqual(
            [r["requestedOn"] for r in self.data(body)],
            sorted((r["requestedOn"] for r in self.data(body)), reverse=True),
        )
        self.assertEqual(self.data(body)[0]["requestedOn"], "2026-09-10")

    def test_department_order_adds_subtotals_that_add_up(self):
        body = self.run_report(self.RID, sortBy="department")
        subs = self.subtotals(body)
        self.assertEqual(
            [s["employeeName"] for s in subs], ["CUTTING total", "PACKING total", "SEWING total", "Unassigned total"]
        )
        cutting = subs[0]
        self.assertEqual(
            (cutting["repaid"], cutting["outstanding"], cutting["overdueAmount"]), (6000.0, 11000.0, 3000.0)
        )
        for key in ("repaid", "outstanding", "overdueInstalments", "overdueAmount"):
            self.assertEqual(
                round(sum(s[key] for s in subs), 2), body["totals"][key], key
            )  # subtotals never double count
        # a subtotal row follows each department's rows
        self.assertEqual([r.get("_kind") for r in body["rows"]].count("subtotal"), 4)
        self.assertEqual(len(self.data(body)), 8)

    def test_no_subtotals_for_the_default_order(self):
        self.assertEqual(self.subtotals(self.run_report(self.RID)), [])

    def test_invalid_filters_are_rejected(self):
        for params in ({"requested": "someday"}, {"status": "lost"}, {"advanceType": "gift"}, {"sortBy": "random"}):
            r = self.get("run", self.RID, **params)
            self.assertEqual(r.status_code, 400, params)
            self.assertEqual(r.json()["error"], "invalid_filter")

    def test_branch_scoped_user_cannot_widen_the_scope(self):
        body = self.run_report(self.RID, user=self.b1_user)
        # Unit 1 = F1, F2, F3 (leaver); F10 is Unit 2, F4 has no branch at all
        self.assertEqual(sorted(set(self.codes(body))), ["F1", "F2", "F3"])
        self.assertEqual(len(self.data(body)), 5)
        self.assertEqual(body["totals"]["outstanding"], 15000.0)  # A1 6000 + A2 5000 + A5 4000
        self.assertEqual(self.cards(body)["Sanctioned (approved + closed)"], 21000.0)
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e3.id))["rows"], [])
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e5.id))["rows"], [])
        self.assertNotIn("F10", self.codes(self.run_report(self.RID, user=self.b1_user, sortBy="department")))


# ═════════════════════════════════════════════════════════════════════════════
#  advance-recovery-schedule
# ═════════════════════════════════════════════════════════════════════════════


class AdvanceScheduleTests(_Case):
    RID = ADV_SCHEDULE

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        make_advances(cls)

    def rows_for(self, dateFrom, dateTo, **params):
        return self.run_report(self.RID, dateFrom=dateFrom, dateTo=dateTo, **params)

    def key(self, r):
        return (r["employeeCode"], r["dueMonth"], r["status"])

    def test_this_month_golden(self):
        body = self.rows_for("2026-09-01", "2026-09-30")
        self.assertEqual(
            sorted(self.key(r) for r in self.data(body)),
            [("F1", "Sep 2026", "Pending"), ("F2", "Sep 2026", "Pending"), ("F4", "Sep 2026", "Deducted")],
        )
        self.assertEqual(body["totals"], {"amount": 8700.0})
        cards = self.cards(body)
        self.assertEqual(
            (
                cards["Instalments"],
                cards["Total scheduled"],
                cards["Deducted"],
                cards["Pending"],
                cards["Overdue amount"],
            ),
            (3, 8700.0, 700.0, 8000.0, 0.0),
        )
        self.assertEqual(cards["Overdue instalments"], 0)

    def test_instalment_position_balance_and_lateness(self):
        body = self.rows_for("2026-08-01", "2026-09-30")
        f1 = {r["dueMonth"]: r for r in self.data(body) if r["employeeCode"] == "F1"}
        self.assertEqual(sorted(f1), ["Aug 2026", "Sep 2026"])
        self.assertEqual((f1["Aug 2026"]["instalmentNo"], f1["Aug 2026"]["balanceAfter"]), ("3 of 4", 3000.0))
        self.assertEqual((f1["Aug 2026"]["status"], f1["Aug 2026"]["monthsLate"]), ("Overdue", 1))
        self.assertEqual((f1["Sep 2026"]["instalmentNo"], f1["Sep 2026"]["balanceAfter"]), ("4 of 4", 0.0))
        self.assertEqual((f1["Sep 2026"]["status"], f1["Sep 2026"]["monthsLate"]), ("Pending", None))
        self.assertEqual(
            (f1["Aug 2026"]["advanceAmount"], f1["Aug 2026"]["advanceType"], f1["Aug 2026"]["advanceStatus"]),
            (12000.0, "Term Loan", "Approved"),
        )
        self.assertEqual(f1["Aug 2026"]["paymentMethod"], "Payroll")
        deducted = self.row(body, "F10")
        self.assertEqual(
            (deducted["status"], deducted["amount"], deducted["advanceStatus"]), ("Deducted", 2000.0, "Closed")
        )

    def test_position_counts_instalments_outside_the_filter(self):
        # only September is requested, but it is still instalment 4 of 4 of the plan
        body = self.rows_for("2026-09-01", "2026-09-30", advanceType="term")
        self.assertEqual([(r["instalmentNo"], r["dueMonth"]) for r in self.data(body)], [("4 of 4", "Sep 2026")])

    def test_august_to_september_totals(self):
        body = self.rows_for("2026-08-01", "2026-09-30")
        self.assertEqual(len(self.data(body)), 6)
        self.assertEqual(body["totals"]["amount"], 17700.0)
        cards = self.cards(body)
        self.assertEqual((cards["Deducted"], cards["Pending"], cards["Overdue amount"]), (2700.0, 8000.0, 7000.0))
        self.assertEqual(cards["Deducted"] + cards["Pending"] + cards["Overdue amount"], cards["Total scheduled"])
        self.assertEqual(cards["Overdue instalments"], 2)
        self.assertEqual(cards["Not yet recovered - ex-employees"], 4000.0)
        self.assertTrue(
            any(
                "1 pending/overdue instalment(s) belong to employees who are no longer active" in n
                for n in body["notes"]
            )
        )

    def test_whole_history_range(self):
        body = self.rows_for("2026-06-01", "2026-09-30")
        self.assertEqual(len(self.data(body)), 8)
        self.assertEqual(body["totals"]["amount"], 23700.0)  # 12000 + 5000 + 2000 + 4000 + 700
        # advances with no instalment rows (pending 1,000, rejected 900, closed-by-hand 1,500) contribute nothing
        self.assertEqual(
            sorted({r["advanceAmount"] for r in self.data(body)}), [700.0, 2000.0, 4000.0, 5000.0, 12000.0]
        )

    def test_instalment_status_filter(self):
        rng = ("2026-06-01", "2026-09-30")
        self.assertEqual(
            sorted(self.key(r) for r in self.data(self.rows_for(*rng, instalmentStatus="overdue"))),
            [("F1", "Aug 2026", "Overdue"), ("F3", "Aug 2026", "Overdue")],
        )
        self.assertEqual(
            sorted(self.key(r) for r in self.data(self.rows_for(*rng, instalmentStatus="deducted"))),
            [
                ("F1", "Jul 2026", "Deducted"),
                ("F1", "Jun 2026", "Deducted"),
                ("F10", "Aug 2026", "Deducted"),
                ("F4", "Sep 2026", "Deducted"),
            ],
        )
        self.assertEqual(
            sorted(self.key(r) for r in self.data(self.rows_for(*rng, instalmentStatus="pending"))),
            [("F1", "Sep 2026", "Pending"), ("F2", "Sep 2026", "Pending")],
        )
        both = self.rows_for(*rng, instalmentStatus="pending,overdue")
        self.assertEqual(len(self.data(both)), 4)
        self.assertEqual(both["totals"]["amount"], 15000.0)  # 3000 + 4000 + 3000 + 5000

    def test_type_and_scope_filters(self):
        rng = ("2026-06-01", "2026-09-30")
        self.assertEqual({r["employeeCode"] for r in self.data(self.rows_for(*rng, advanceType="term"))}, {"F1"})
        self.assertEqual(len(self.data(self.rows_for(*rng, advanceType="term"))), 4)
        self.assertEqual(
            {r["employeeCode"] for r in self.data(self.rows_for(*rng, departmentIds=str(self.sewing.id)))}, {"F10"}
        )
        self.assertEqual(
            {r["employeeCode"] for r in self.data(self.rows_for(*rng, employmentType="production"))}, {"F2"}
        )
        self.assertEqual(
            {r["employeeCode"] for r in self.data(self.rows_for(*rng, employeeIds=str(self.e5.id)))}, {"F4"}
        )
        self.assertEqual({r["employeeCode"] for r in self.data(self.rows_for(*rng, employeeStatus="inactive"))}, {"F3"})
        self.assertEqual(self.rows_for(*rng, branchIds=str(self.b2.id))["rows"][0]["employeeCode"], "F10")

    def test_designation_filter(self):
        body = self.rows_for("2026-06-01", "2026-09-30", designationIds=str(self.d_op.id))
        self.assertEqual(len(self.data(body)), 7)  # F4's September instalment has no designation to match
        self.assertNotIn("F4", {r["employeeCode"] for r in self.data(body)})

    def test_closed_advance_with_open_instalments_is_flagged(self):
        # payroll never checks the advance status, so this instalment WOULD still be deducted
        make_advance(self.e1, "general", "900", "closed", ist(2026, 8, 1), reps=[(2026, 9, 900, False)])
        body = self.rows_for("2026-09-01", "2026-09-30")
        r = self.row(body, "F1", 900.0)
        self.assertEqual((r["status"], r["advanceStatus"]), ("Pending", "Closed"))
        self.assertTrue(any("belong to advances already marked Closed" in n for n in body["notes"]))
        self.assertFalse(any("Closed - payroll" in n for n in self.rows_for("2026-08-01", "2026-08-31")["notes"]))

    def test_pending_advances_are_hidden_but_a_real_deduction_never_is(self):
        make_advance(
            self.e1, "general", "800", "pending", ist(2026, 9, 1), reps=[(2026, 9, 800, False)]
        )  # never deducted
        make_advance(
            self.e2, "general", "600", "rejected", ist(2026, 9, 1), stored_repaid="600", reps=[(2026, 9, 600, True)]
        )
        body = self.rows_for("2026-09-01", "2026-09-30")
        amounts = sorted(r["amount"] for r in self.data(body))
        self.assertEqual(
            amounts, [600.0, 700.0, 3000.0, 5000.0]
        )  # 800 (pending, unprocessed) hidden, 600 (deducted) shown
        rejected = self.row(body, "F2", 600.0)
        self.assertEqual((rejected["status"], rejected["advanceStatus"]), ("Deducted", "Rejected"))

    def test_schedule_that_falls_short_of_the_sanctioned_amount_is_flagged(self):
        make_advance(
            self.e1,
            "term",
            "5000",
            "approved",
            ist(2026, 9, 5),
            approved_at=utc(2026, 9, 6, 5, 0),
            emi="1000",
            months=3,
            start=(9, 2026),
            reps=[(2026, 9, 1000, False), (2026, 10, 1000, False), (2026, 11, 1000, False)],
        )
        body = self.rows_for("2026-09-01", "2026-09-30", employeeIds=str(self.e1.id))
        r = self.row(body, "F1", 1000.0)
        self.assertEqual((r["instalmentNo"], r["balanceAfter"]), ("1 of 3", 4000.0))
        self.assertTrue(
            any("1 approved advance(s) listed here have a schedule that adds up to less" in n for n in body["notes"])
        )
        last = self.row(self.rows_for("2026-11-01", "2026-11-30"), "F1", 1000.0)
        self.assertEqual((last["instalmentNo"], last["balanceAfter"]), ("3 of 3", 2000.0))  # 2,000 never scheduled
        # complete plans (the base fixtures) carry no such note
        clean = self.rows_for("2026-06-01", "2026-08-31")
        self.assertFalse(any("adds up to less" in n for n in clean["notes"]))

    def test_department_subtotals_add_up(self):
        body = self.rows_for("2026-08-01", "2026-09-30")
        subs = self.subtotals(body)
        self.assertEqual(
            [s["employeeName"] for s in subs], ["CUTTING total", "PACKING total", "SEWING total", "Unassigned total"]
        )
        self.assertEqual([s["amount"] for s in subs], [11000.0, 4000.0, 2000.0, 700.0])
        self.assertEqual(round(sum(s["amount"] for s in subs), 2), body["totals"]["amount"])

    def test_single_department_has_no_subtotals(self):
        body = self.rows_for("2026-08-01", "2026-09-30", departmentIds=str(self.cutting.id))
        self.assertEqual(self.subtotals(body), [])

    def test_range_validation(self):
        self.assertEqual(self.get("run", self.RID, dateFrom="2026-09-30", dateTo="2026-09-01").status_code, 400)
        r = self.get("run", self.RID, dateFrom="2020-01-01", dateTo="2026-09-30")
        self.assertEqual((r.status_code, r.json()["field"]), (400, "dateTo"))  # wider than ~3 years
        self.assertEqual(self.get("run", self.RID, instalmentStatus="lost").status_code, 400)

    def test_range_is_matched_by_month_not_by_day(self):
        body = self.rows_for("2026-09-20", "2026-09-25")  # mid-month: still the whole of September
        self.assertEqual(len(self.data(body)), 3)

    def test_default_range_is_the_current_month(self):
        body = self.run_report(self.RID)
        self.assertEqual(len(self.data(body)), 3)

    def test_branch_isolation(self):
        body = self.run_report(self.RID, user=self.b1_user, **FULL_PARAMS[self.RID])
        self.assertEqual({r["employeeCode"] for r in self.data(body)}, {"F1", "F2", "F3"})
        self.assertEqual(body["totals"]["amount"], 12000.0 + 5000.0 + 4000.0)
        params = dict(FULL_PARAMS[self.RID])
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, branchIds=str(self.b2.id), **params)["rows"], [])
        self.assertEqual(
            self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e3.id), **params)["rows"], []
        )
        self.assertEqual(
            self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e5.id), **params)["rows"], []
        )


# ═════════════════════════════════════════════════════════════════════════════
#  The report agrees with the real advance + payroll flow
# ═════════════════════════════════════════════════════════════════════════════


class AdvanceEngineAgreementTests(TestCase):
    """Create the advance through the API, deduct it with the real payroll engine, then read it back."""

    def setUp(self):
        self.today = date(2026, 3, 10)
        patcher = mock.patch("api.reporting.filters.ist_today", side_effect=lambda: self.today)
        patcher.start()
        self.addCleanup(patcher.stop)
        wa = mock.patch("api.settlement_views.whatsapp_approvals.notify_decision")
        wa.start()
        self.addCleanup(wa.stop)

        ps = PayrollSettings.get()
        ps.staff_payroll_rules_enabled = False
        ps.late_deduction_slabs = []
        ps.attendance_mode = "simple"
        ps.save()
        dept = Department.objects.create(name="ENGINE")
        self.emp = Employee.objects.create(
            employee_code="EA1",
            first_name="Eng",
            last_name="Agree",
            employment_type="staff",
            status="active",
            salary_type="monthly",
            salary_amount=Decimal("24000.00"),
            department=dept,
        )
        for d in _build_working_days(2, 2026, False, set()):
            AttendanceDayRecord.objects.create(
                employee=self.emp,
                date=d,
                status="present",
                shifts_earned=Decimal("1.00"),
                is_half_shift=False,
                source="manual",
            )
        self.admin = HRUser.objects.create(username="eng_admin", password_hash="x", is_super_admin=True)
        self.hdr = headers(self.admin)

    def post(self, path, body):
        return self.client.post(path, data=json.dumps(body), content_type="application/json", **self.hdr)

    def put(self, path, body):
        return self.client.put(path, data=json.dumps(body), content_type="application/json", **self.hdr)

    def report(self, rid, **params):
        r = self.client.get(f"/api/reports/run/{rid}", params, **self.hdr)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    def test_ledger_and_schedule_match_the_engine(self):
        created = self.post(
            "/api/advances",
            {
                "employeeId": self.emp.id,
                "amount": 5000,
                "advanceType": "term",
                "emiAmount": 2000,
                "repaymentMonths": 3,
                "repaymentStartMonth": 2,
                "repaymentStartYear": 2026,
                "purpose": "Wedding",
            },
        )
        self.assertEqual(created.status_code, 201)
        adv_id = created.json()["id"]

        # before approval: a pending request with no balance
        row = self.report("advance-ledger")["rows"][0]
        self.assertEqual((row["status"], row["outstanding"], row["repaid"]), ("Pending", None, None))

        approved = self.put(f"/api/advances/{adv_id}", {"status": "approved", "approvedBy": "HR Head"})
        self.assertEqual(approved.status_code, 200)
        _generate_staff_payroll(self.emp, 2, 2026)
        slip = SalarySlip.objects.get(employee=self.emp, month=2, year=2026)
        self.assertEqual(slip.advance_deduction, Decimal("2000.00"))

        adv = Advance.objects.get(pk=adv_id)
        row = self.report("advance-ledger")["rows"][0]
        self.assertEqual((row["status"], row["advanceType"], row["approvedBy"]), ("Approved", "Term Loan", "HR Head"))
        self.assertEqual(row["repaid"], float(adv.total_repaid))  # same figure the app keeps...
        self.assertEqual(row["repaid"], float(slip.advance_deduction))  # ...and the slip deducted
        self.assertEqual((row["repaid"], row["outstanding"], row["repaidPct"]), (2000.0, 3000.0, 40.0))
        self.assertEqual((row["instalments"], row["emi"], row["nextDue"]), ("1 of 3", 2000.0, "Mar 2026"))
        self.assertEqual((row["overdueInstalments"], row["overdueAmount"]), (0, 0.0))
        body = self.report("advance-ledger")
        self.assertFalse(any("differs" in n for n in body["notes"]))  # stored balance and instalments agree

        sched = self.report("advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-04-30")
        self.assertEqual(
            [
                (r["dueMonth"], r["amount"], r["status"], r["instalmentNo"], r["balanceAfter"])
                for r in sched["rows"]
                if "_kind" not in r
            ],
            [
                ("Feb 2026", 2000.0, "Deducted", "1 of 3", 3000.0),
                ("Mar 2026", 2000.0, "Pending", "2 of 3", 1000.0),
                ("Apr 2026", 1000.0, "Pending", "3 of 3", 0.0),
            ],
        )
        feb = sched["rows"][0]
        self.assertEqual(
            feb["amount"], float(slip.advance_deduction)
        )  # the schedule's "deducted" is the slip's deduction

        # two months later payroll never ran for March / April: both are overdue, and the ledger says so
        self.today = date(2026, 5, 10)
        late = self.report("advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-04-30")
        self.assertEqual(
            [(r["dueMonth"], r["status"], r["monthsLate"]) for r in late["rows"] if "_kind" not in r],
            [("Feb 2026", "Deducted", None), ("Mar 2026", "Overdue", 2), ("Apr 2026", "Overdue", 1)],
        )
        row = self.report("advance-ledger")["rows"][0]
        self.assertEqual((row["overdueInstalments"], row["overdueAmount"], row["nextDue"]), (2, 3000.0, "Mar 2026"))

    def test_final_instalment_closes_the_advance_in_the_report(self):
        adv_id = self.post(
            "/api/advances",
            {
                "employeeId": self.emp.id,
                "amount": 2000,
                "advanceType": "general",
                "repaymentStartMonth": 2,
                "repaymentStartYear": 2026,
            },
        ).json()["id"]
        self.put(f"/api/advances/{adv_id}", {"status": "approved", "approvedBy": "HR Head"})
        _generate_staff_payroll(self.emp, 2, 2026)
        self.assertEqual(Advance.objects.get(pk=adv_id).status, "closed")
        row = self.report("advance-ledger")["rows"][0]
        self.assertEqual(
            (row["status"], row["repaid"], row["outstanding"], row["repaidPct"]), ("Closed", 2000.0, 0.0, 100.0)
        )
        self.assertEqual((row["overdueInstalments"], row["nextDue"]), (None, None))
        sched = self.report("advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-02-28")
        self.assertEqual(
            [(r["status"], r["advanceStatus"], r["balanceAfter"]) for r in sched["rows"]], [("Deducted", "Closed", 0.0)]
        )

    # ── a term loan followed through real monthly payroll runs ────────────────────────────────────
    def present_all_month(self, month):
        for d in _build_working_days(month, 2026, False, set()):
            AttendanceDayRecord.objects.create(
                employee=self.emp,
                date=d,
                status="present",
                shifts_earned=Decimal("1.00"),
                is_half_shift=False,
                source="manual",
            )

    def approve_term_loan(self, amount=3000, emi=1000, months=3):
        adv_id = self.post(
            "/api/advances",
            {
                "employeeId": self.emp.id,
                "amount": amount,
                "advanceType": "term",
                "emiAmount": emi,
                "repaymentMonths": months,
                "repaymentStartMonth": 2,
                "repaymentStartYear": 2026,
            },
        ).json()["id"]
        approved = self.put(f"/api/advances/{adv_id}", {"status": "approved", "approvedBy": "HR Head"})
        self.assertEqual(approved.status_code, 200)
        return adv_id

    @staticmethod
    def cards_of(body):
        return {c["label"]: c["value"] for c in body["summary"]}

    def test_a_term_loan_walks_to_closed_over_three_real_payroll_runs(self):
        self.approve_term_loan()
        self.present_all_month(3)
        self.present_all_month(4)
        expected = [
            # month generated, "today", recovered, outstanding, pct, instalments, next due, status
            (2, date(2026, 2, 28), 1000.0, 2000.0, 33.33, "1 of 3", "Mar 2026", "Approved"),
            (3, date(2026, 3, 31), 2000.0, 1000.0, 66.67, "2 of 3", "Apr 2026", "Approved"),
            (4, date(2026, 4, 30), 3000.0, 0.0, 100.0, "3 of 3", None, "Closed"),
        ]
        deducted = 0.0
        for month, today, repaid, outstanding, pct, inst, nxt, status in expected:
            _generate_staff_payroll(self.emp, month, 2026)
            self.today = today
            deducted += float(SalarySlip.objects.get(employee=self.emp, month=month, year=2026).advance_deduction)
            row = self.report("advance-ledger")["rows"][0]
            self.assertEqual(
                (
                    row["repaid"],
                    row["outstanding"],
                    row["repaidPct"],
                    row["instalments"],
                    row["nextDue"],
                    row["status"],
                ),
                (repaid, outstanding, pct, inst, nxt, status),
                month,
            )
            self.assertEqual(row["repaid"], deducted, month)  # exactly what the slips deducted so far
            self.assertEqual(row["overdueInstalments"], 0 if status == "Approved" else None, month)
        sched = self.report("advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-04-30")
        self.assertEqual(
            [
                (r["dueMonth"], r["status"], r["instalmentNo"], r["balanceAfter"])
                for r in sched["rows"]
                if "_kind" not in r
            ],
            [
                ("Feb 2026", "Deducted", "1 of 3", 2000.0),
                ("Mar 2026", "Deducted", "2 of 3", 1000.0),
                ("Apr 2026", "Deducted", "3 of 3", 0.0),
            ],
        )
        self.assertEqual(sched["totals"]["amount"], 3000.0)
        self.assertEqual(self.cards_of(sched)["Overdue amount"], 0.0)

    def test_regenerating_a_month_cannot_make_the_recovery_disappear(self):
        # Payroll regeneration is known to drop the slip's advance figure while the instalment stays processed.
        # The report reads the instalment rows, so what has been recovered is stable whatever regeneration does.
        self.approve_term_loan(amount=2000, emi=1000, months=2)
        _generate_staff_payroll(self.emp, 2, 2026)
        before = self.report("advance-ledger")["rows"][0]
        self.assertEqual((before["repaid"], before["outstanding"]), (1000.0, 1000.0))
        _generate_staff_payroll(self.emp, 2, 2026)
        after = self.report("advance-ledger")["rows"][0]
        for key in ("repaid", "outstanding", "repaidPct", "instalments", "nextDue", "status"):
            self.assertEqual(after[key], before[key], key)
        sched = self.report("advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-02-28")
        self.assertEqual(
            [(r["status"], r["amount"]) for r in sched["rows"] if "_kind" not in r], [("Deducted", 1000.0)]
        )

    def test_a_month_payroll_did_not_pick_up_stays_overdue(self):
        self.approve_term_loan(amount=2000, emi=1000, months=2)
        # payroll is only ever run for March: the February instalment is left behind
        self.present_all_month(3)
        _generate_staff_payroll(self.emp, 3, 2026)
        self.today = date(2026, 3, 31)
        sched = self.report("advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-03-31")
        rows = {r["dueMonth"]: r for r in sched["rows"] if "_kind" not in r}
        # payroll takes each instalment in its OWN month: March's run found the March row and left February's
        self.assertEqual((rows["Feb 2026"]["status"], rows["Feb 2026"]["monthsLate"]), ("Overdue", 1))
        self.assertEqual(rows["Mar 2026"]["status"], "Deducted")
        row = self.report("advance-ledger")["rows"][0]
        self.assertEqual((row["repaid"], row["outstanding"], row["instalments"]), (1000.0, 1000.0, "1 of 2"))
        self.assertEqual((row["overdueInstalments"], row["overdueAmount"], row["nextDue"]), (1, 1000.0, "Feb 2026"))
        self.assertEqual(len(self.report("advance-ledger", overdueOnly="true")["rows"]), 1)

    def test_production_advance_is_recovered_in_the_month_the_period_ends(self):
        ps = PayrollSettings.get()
        ps.prod_payroll_rules_enabled = False
        ps.prod_pf_ef_enabled = False
        ps.prod_late_detection_enabled = False
        ps.save()
        worker = Employee.objects.create(
            employee_code="EP1",
            first_name="Prod",
            last_name="Agree",
            employment_type="production",
            status="active",
            salary_per_shift=Decimal("500.00"),
            department=self.emp.department,
        )
        for i in range(10):  # 2 - 11 Feb
            AttendanceDayRecord.objects.create(
                employee=worker,
                date=date(2026, 2, 2 + i),
                status="present",
                shifts_earned=Decimal("1.00"),
                is_half_shift=False,
                source="manual",
            )
        adv_id = self.post(
            "/api/advances",
            {
                "employeeId": worker.id,
                "amount": 1000,
                "advanceType": "general",
                "repaymentStartMonth": 2,
                "repaymentStartYear": 2026,
            },
        ).json()["id"]
        self.put(f"/api/advances/{adv_id}", {"status": "approved", "approvedBy": "HR Head"})
        _generate_production_payroll(worker, date(2026, 2, 2), date(2026, 2, 11))
        slip = SalarySlip.objects.get(employee=worker)
        self.assertEqual(slip.advance_deduction, Decimal("1000.00"))
        self.today = date(2026, 2, 28)
        ledger = self.report("advance-ledger")
        row = next(r for r in ledger["rows"] if r["employeeCode"] == "EP1")
        self.assertEqual((row["status"], row["repaid"], row["outstanding"]), ("Closed", 1000.0, 0.0))
        self.assertEqual(row["repaid"], float(slip.advance_deduction))
        sched = self.report(
            "advance-recovery-schedule", dateFrom="2026-02-01", dateTo="2026-02-28", employmentType="production"
        )
        self.assertEqual(
            [(r["employeeCode"], r["dueMonth"], r["status"], r["amount"]) for r in sched["rows"] if "_kind" not in r],
            [("EP1", "Feb 2026", "Deducted", 1000.0)],
        )


# ═════════════════════════════════════════════════════════════════════════════
#  bonus-register
# ═════════════════════════════════════════════════════════════════════════════


class BonusRegisterTests(_Case):
    RID = BONUS

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        make_bonus_rows(cls)

    def test_all_years_golden_totals_and_year_subtotals(self):
        body = self.run_report(self.RID)
        rows = self.data(body)
        self.assertEqual(len(rows), 5)
        self.assertEqual(body["totals"], {"calculationBase": 314000.0, "bonusAmount": 26156.2})
        # newest year first, then department, then natural employee code
        self.assertEqual(
            [(r["financialYear"], r["employeeCode"]) for r in rows],
            [("2025-26", "F1"), ("2025-26", "F2"), ("2025-26", "F3"), ("2025-26", "F10"), ("2024-25", "F1")],
        )
        subs = self.subtotals(body)
        self.assertEqual([s["employeeName"] for s in subs], ["FY 2025-26 total", "FY 2024-25 total"])
        self.assertEqual(
            [(s["calculationBase"], s["bonusAmount"]) for s in subs], [(230000.0, 19159.0), (84000.0, 6997.2)]
        )
        self.assertEqual(round(sum(s["bonusAmount"] for s in subs), 2), body["totals"]["bonusAmount"])

    def test_row_values(self):
        body = self.run_report(self.RID, financialYear="2025-26")
        r = self.row(body, "F1", 6997.2, key="bonusAmount")
        self.assertEqual(r["employeeName"], "Asha Fin")
        self.assertEqual((r["department"], r["employmentType"], r["financialYear"]), ("CUTTING", "Staff", "2025-26"))
        self.assertEqual((r["recordsConsidered"], r["calculationBase"], r["bonusPercent"]), (12, 84000.0, 8.33))
        self.assertEqual((r["status"], r["computedBy"], r["notes"]), ("Approved", "HR Admin", None))
        self.assertEqual(r["createdOn"], "2026-04-20")
        prod = self.row(body, "F2")
        self.assertEqual(
            (prod["employmentType"], prod["status"], prod["notes"], prod["recordsConsidered"]),
            ("Production", "Paid", "Weekly slips", 50),
        )
        self.assertEqual(prod["createdOn"], "2026-04-22")  # 19:00 UTC on the 21st = 00:30 IST on the 22nd
        leaver = self.row(body, "F3")
        self.assertEqual(
            (leaver["department"], leaver["status"]), ("PACKING", "Calculated")
        )  # the employee has left; row stays

    def test_single_year_gets_department_subtotals(self):
        body = self.run_report(self.RID, financialYear="2025-26")
        self.assertEqual(len(self.data(body)), 4)
        self.assertEqual(body["totals"], {"calculationBase": 230000.0, "bonusAmount": 19159.0})
        subs = self.subtotals(body)
        self.assertEqual([s["employeeName"] for s in subs], ["CUTTING total", "PACKING total", "SEWING total"])
        self.assertEqual([s["bonusAmount"] for s in subs], [9496.2, 4998.0, 4664.8])
        self.assertEqual(round(sum(s["bonusAmount"] for s in subs), 2), 19159.0)

    def test_summary_cards_and_status_split(self):
        cards = self.cards(self.run_report(self.RID))
        self.assertEqual(cards["Employees"], 4)  # F1 appears in two years
        self.assertEqual(cards["Total bonus"], 26156.2)
        self.assertEqual(cards["Average bonus"], 5231.24)
        self.assertEqual(
            (cards["Calculated (not approved)"], cards["Approved (not paid)"], cards["Paid"]), (9662.8, 6997.2, 9496.2)
        )
        self.assertAlmostEqual(
            cards["Calculated (not approved)"] + cards["Approved (not paid)"] + cards["Paid"],
            cards["Total bonus"],
            places=2,
        )

    def test_filters_narrow(self):
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, financialYear="2024-25"))), ["F1"])
        both = self.run_report(self.RID, financialYear="2024-25,2025-26")
        self.assertEqual(len(self.data(both)), 5)
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, status="paid"))), ["F1", "F2"])
        self.assertEqual(
            sorted(self.codes(self.run_report(self.RID, status="calculated,approved"))), ["F1", "F10", "F3"]
        )
        self.assertEqual(self.codes(self.run_report(self.RID, employmentType="production")), ["F2"])
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, departmentIds=str(self.sewing.id)))), ["F10"])
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, employeeIds=str(self.e1.id)))), ["F1", "F1"])
        self.assertEqual(self.run_report(self.RID, financialYear="2023-24")["rows"], [])
        self.assertEqual(self.get("run", self.RID, financialYear="20xx").status_code, 400)
        self.assertEqual(self.get("run", self.RID, status="lost").status_code, 400)

    def test_designation_and_branch_filters(self):
        Employee.objects.filter(pk=self.e2.pk).update(designation=self.d_sup)
        self.assertEqual(self.codes(self.run_report(self.RID, designationIds=str(self.d_sup.id))), ["F2"])
        self.assertEqual(len(self.data(self.run_report(self.RID, designationIds=str(self.d_op.id)))), 4)
        self.assertEqual(self.codes(self.run_report(self.RID, branchIds=str(self.b2.id))), ["F10"])
        self.assertEqual(
            sorted(set(self.codes(self.run_report(self.RID, branchIds=str(self.b1.id))))), ["F1", "F2", "F3"]
        )

    def test_default_request_echoes_no_filters(self):
        self.assertEqual(self.run_report(self.RID)["filters"], [])

    def test_data_quality_note_counts_short_staff_years(self):
        notes = " ".join(self.run_report(self.RID)["notes"])
        self.assertIn(
            "1 staff row(s) rest on fewer than 12 payroll months", notes
        )  # F10 has 8; production F2 (50) is exempt
        self.assertNotIn("staff row(s)", " ".join(self.run_report(self.RID, employmentType="production")["notes"]))

    def test_branch_isolation(self):
        body = self.run_report(self.RID, user=self.b1_user)
        self.assertEqual(sorted(self.codes(body)), ["F1", "F1", "F2", "F3"])  # F10 is Unit 2
        self.assertEqual(body["totals"]["bonusAmount"], round(6997.2 + 2499.0 + 6997.2 + 4998.0, 2))
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e3.id))["rows"], [])

    def test_agrees_with_the_real_bonus_generation(self):
        """Real payroll slip -> real /api/bonus/generate -> the register shows exactly what the Bonus page stored."""
        Bonus.objects.all().delete()
        ps = PayrollSettings.get()
        ps.staff_payroll_rules_enabled = False
        ps.late_deduction_slabs = []
        ps.attendance_mode = "simple"
        ps.save()
        emp = Employee.objects.create(
            employee_code="BG1",
            first_name="Bonus",
            last_name="Real",
            employment_type="staff",
            status="active",
            salary_type="monthly",
            salary_amount=Decimal("14400.00"),
            department=self.cutting,
            branch=self.b1,
        )
        for d in _build_working_days(2, 2026, False, set()):
            AttendanceDayRecord.objects.create(
                employee=emp,
                date=d,
                status="present",
                shifts_earned=Decimal("1.00"),
                is_half_shift=False,
                source="manual",
            )
        _generate_staff_payroll(emp, 2, 2026)  # basic = 50% of 14,400 = 7,200 (> the 7,000 wage ceiling)
        hdr = headers(self.admin)
        r = self.client.post(
            "/api/bonus/generate", data=json.dumps({"financialYear": "2025-26"}), content_type="application/json", **hdr
        )
        self.assertEqual(r.status_code, 200)
        stored = Bonus.objects.get(employee=emp)

        row = self.data(self.run_report(self.RID, financialYear="2025-26"))[0]
        self.assertEqual((row["employeeCode"], row["financialYear"], row["status"]), ("BG1", "2025-26", "Calculated"))
        self.assertEqual((row["recordsConsidered"], row["calculationBase"]), (1, 7000.0))  # min(7,200, 7,000)
        self.assertEqual((row["bonusPercent"], row["bonusAmount"]), (8.33, 583.1))  # 7,000 x 8.33%
        self.assertEqual(row["bonusAmount"], float(stored.bonus_amount))
        self.assertEqual(row["computedBy"], stored.computed_by or None)

        r = self.client.patch(
            f"/api/bonus/{stored.id}", data=json.dumps({"status": "paid"}), content_type="application/json", **hdr
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.data(self.run_report(self.RID, financialYear="2025-26"))[0]["status"], "Paid")


# ═════════════════════════════════════════════════════════════════════════════
#  increment-history
# ═════════════════════════════════════════════════════════════════════════════


class IncrementHistoryTests(_Case):
    RID = INCREMENT

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        make_increments(cls)

    def effective(self, body):
        return [(r["employeeCode"], r["effectiveDate"]) for r in self.data(body)]

    def test_all_time_golden_and_default_order(self):
        body = self.run_report(self.RID)
        self.assertEqual(
            self.effective(body),
            [
                ("F4", "2026-10-01"),
                ("F1", "2026-09-01"),
                ("F3", "2026-07-15"),
                ("F1", "2026-04-01"),
                ("F10", "2025-04-01"),
            ],
        )
        self.assertEqual(
            body["totals"], {"incrementAmount": 8200.0, "percent": 8.0}
        )  # 2000+1100+1500+600+3000; mean of 10,5,10,5,10
        self.assertEqual(self.subtotals(body), [])

    def test_row_values(self):
        body = self.run_report(self.RID)
        r = self.row(body, "F1", 2000.0, key="incrementAmount")
        self.assertEqual((r["employeeName"], r["department"], r["designation"]), ("Asha Fin", "CUTTING", "Operator"))
        self.assertEqual(
            (r["previousSalary"], r["newSalary"], r["incrementAmount"], r["percent"]), (20000.0, 22000.0, 2000.0, 10.0)
        )
        self.assertEqual((r["addedBy"], r["notes"]), ("HR Admin", "Annual"))
        self.assertEqual(r["recordedOn"], "2026-03-31 03:30")  # 22:00 UTC on the 30th = 03:30 IST on the 31st
        nobody = self.row(body, "F4")
        self.assertEqual((nobody["department"], nobody["designation"], nobody["addedBy"]), ("Unassigned", None, None))
        self.assertEqual(self.row(body, "F3")["department"], "PACKING")  # employee has left: still listed

    def test_summary_cards(self):
        cards = self.cards(self.run_report(self.RID))
        self.assertEqual(cards["Increments"], 5)
        self.assertEqual(cards["Employees"], 4)
        self.assertEqual(cards["Monthly uplift"], 8200.0)
        self.assertEqual(cards["Annual uplift (x12)"], 98400.0)
        self.assertEqual(cards["Average increment (%)"], 8.0)
        self.assertEqual(cards["Highest increment (%)"], 10.0)

    def test_relative_windows_filter_on_effective_date(self):
        def eff(token):
            return sorted(r["effectiveDate"] for r in self.data(self.run_report(self.RID, effective=token)))

        self.assertEqual(eff("thisMonth"), ["2026-09-01"])
        self.assertEqual(eff("lastMonth"), [])
        self.assertEqual(eff("last3"), ["2026-07-15", "2026-09-01"])
        self.assertEqual(eff("thisFY"), ["2026-04-01", "2026-07-15", "2026-09-01", "2026-10-01"])
        self.assertEqual(eff("lastFY"), ["2025-04-01"])
        self.assertEqual(eff("thisYear"), ["2026-04-01", "2026-07-15", "2026-09-01", "2026-10-01"])
        self.assertEqual(eff("lastYear"), ["2025-04-01"])

    def test_financial_year_filter_and_its_combination_with_the_window(self):
        def eff(**p):
            return sorted(r["effectiveDate"] for r in self.data(self.run_report(self.RID, **p)))

        self.assertEqual(eff(financialYear="2025-26"), ["2025-04-01"])
        self.assertEqual(eff(financialYear="2026-27"), ["2026-04-01", "2026-07-15", "2026-09-01", "2026-10-01"])
        self.assertEqual(eff(financialYear="2026-27", effective="thisMonth"), ["2026-09-01"])  # both must hold
        # FY 2025-26 (to 31 Mar 2026) intersected with calendar 2026 = Jan-Mar 2026: no increment falls there
        self.assertEqual(eff(financialYear="2025-26", effective="thisYear"), [])
        self.assertEqual(eff(financialYear="2024-25"), [])

    def test_window_is_described_in_notes_and_header(self):
        body = self.run_report(self.RID, effective="last3")
        self.assertTrue(any("Only increments effective from 01-Jul-2026 to 30-Sep-2026" in n for n in body["notes"]))
        self.assertIn(
            ("Effective", "Last 3 months (incl. this month)"), [(f["label"], f["value"]) for f in body["filters"]]
        )

    def test_future_dated_increment_is_flagged(self):
        notes = " ".join(self.run_report(self.RID)["notes"])
        self.assertIn("1 increment(s) are dated after today but are already applied", notes)
        self.assertNotIn("dated after today", " ".join(self.run_report(self.RID, effective="thisMonth")["notes"]))

    def test_employee_scope_filters(self):
        self.assertEqual(
            sorted(self.codes(self.run_report(self.RID, departmentIds=str(self.cutting.id)))), ["F1", "F1"]
        )
        self.assertEqual(self.codes(self.run_report(self.RID, employeeIds=str(self.e3.id))), ["F10"])
        self.assertEqual(self.codes(self.run_report(self.RID, designationIds=str(self.d_op.id)))[:1], ["F1"])
        self.assertEqual(len(self.data(self.run_report(self.RID, designationIds=str(self.d_op.id)))), 4)  # F4 has none
        self.assertEqual(self.run_report(self.RID, designationIds=str(self.d_mgr.id))["rows"], [])
        self.assertEqual(
            self.run_report(self.RID, employmentType="production")["rows"], []
        )  # production pay is never recorded
        self.assertEqual(sorted(set(self.codes(self.run_report(self.RID, branchIds=str(self.b1.id))))), ["F1", "F3"])

    def test_department_order_gives_subtotals_with_average_percent(self):
        body = self.run_report(self.RID, sortBy="department")
        subs = self.subtotals(body)
        self.assertEqual(
            [s["employeeName"] for s in subs], ["CUTTING total", "PACKING total", "SEWING total", "Unassigned total"]
        )
        self.assertEqual(
            [(s["incrementAmount"], s["percent"]) for s in subs],
            [(3100.0, 7.5), (600.0, 5.0), (1500.0, 10.0), (3000.0, 10.0)],
        )
        self.assertEqual(round(sum(s["incrementAmount"] for s in subs), 2), body["totals"]["incrementAmount"])
        self.assertEqual(body["totals"]["percent"], 8.0)  # the grand average ignores subtotal rows
        # inside a department the newest revision comes first
        cut = [r["effectiveDate"] for r in self.data(body) if r["department"] == "CUTTING"]
        self.assertEqual(cut, ["2026-09-01", "2026-04-01"])

    def test_default_request_echoes_no_filters(self):
        self.assertEqual(self.run_report(self.RID)["filters"], [])

    def test_employee_order(self):
        body = self.run_report(self.RID, sortBy="employee")
        self.assertEqual(self.codes(body), ["F1", "F1", "F3", "F4", "F10"])
        self.assertEqual([r["effectiveDate"] for r in self.data(body)][:2], ["2026-09-01", "2026-04-01"])

    def test_invalid_filters(self):
        for params in ({"effective": "soon"}, {"financialYear": "FY26"}, {"sortBy": "salary"}):
            self.assertEqual(self.get("run", self.RID, **params).status_code, 400, params)

    def test_branch_isolation(self):
        body = self.run_report(self.RID, user=self.b1_user)
        self.assertEqual(sorted(self.codes(body)), ["F1", "F1", "F3"])  # F10 = Unit 2, F4 has no branch
        self.assertEqual(body["totals"]["incrementAmount"], 3700.0)
        self.assertEqual(self.cards(body)["Increments"], 3)
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e5.id))["rows"], [])

    def test_agrees_with_the_real_increment_endpoint(self):
        SalaryIncrement.objects.all().delete()
        emp = Employee.objects.create(
            employee_code="IN1",
            first_name="Inc",
            last_name="Real",
            employment_type="staff",
            status="active",
            salary_type="monthly",
            salary_amount=Decimal("20000.00"),
            department=self.cutting,
            designation=self.d_op,
            branch=self.b1,
        )
        r = self.client.post(
            "/api/increments",
            data=json.dumps(
                {"employeeId": emp.id, "percent": 10, "effectiveDate": "2026-04-01", "notes": "Annual review"}
            ),
            content_type="application/json",
            **headers(self.admin),
        )
        self.assertEqual(r.status_code, 201)
        stored = SalaryIncrement.objects.get(employee=emp)
        emp.refresh_from_db()
        row = self.data(self.run_report(self.RID))[0]
        self.assertEqual(
            (row["previousSalary"], row["newSalary"], row["incrementAmount"], row["percent"]),
            (20000.0, 22000.0, 2000.0, 10.0),
        )
        self.assertEqual(row["newSalary"], float(emp.salary_amount))  # the salary the app now pays
        self.assertEqual(
            (row["effectiveDate"], row["notes"], row["addedBy"]),
            ("2026-04-01", "Annual review", stored.added_by or None),
        )
        self.assertEqual(row["recordedOn"], fmt_dt(stored.created_at))


# ═════════════════════════════════════════════════════════════════════════════
#  promotion-history
# ═════════════════════════════════════════════════════════════════════════════


class PromotionHistoryTests(_Case):
    RID = PROMOTION

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        make_promotions(cls)

    def by_change(self, **params):
        body = self.run_report(self.RID, **params)
        return [(r["employeeCode"], r["change"]) for r in self.data(body)]

    def test_all_time_golden_and_order(self):
        body = self.run_report(self.RID)
        self.assertEqual(
            [(r["employeeCode"], r["effectiveDate"], r["change"]) for r in self.data(body)],
            [
                ("F1", "2026-09-01", "Both"),
                ("F4", "2026-08-10", "Department"),
                ("F3", "2026-07-01", "Department"),
                ("F1", "2026-04-01", "Designation"),
                ("F10", "2025-06-01", "Designation"),
            ],
        )
        self.assertIsNone(body["totals"])  # nothing to add up

    def test_row_values(self):
        body = self.run_report(self.RID)
        both = self.row(body, "F1", "2026-09-01", key="effectiveDate")
        self.assertEqual(both["employeeName"], "Asha Fin")
        self.assertEqual((both["previousDesignation"], both["newDesignation"]), ("Supervisor", "Manager"))
        self.assertEqual((both["previousDepartment"], both["newDepartment"]), ("CUTTING", "PACKING"))
        self.assertEqual((both["promotedBy"], both["notes"]), ("HR Admin", None))
        first = self.row(body, "F1", "2026-04-01", key="effectiveDate")
        self.assertEqual((first["previousDesignation"], first["newDesignation"]), ("Operator", "Supervisor"))
        self.assertEqual(
            (first["previousDepartment"], first["newDepartment"], first["notes"]),
            ("CUTTING", "CUTTING", "Line supervisor"),
        )
        self.assertEqual(first["recordedOn"], "2026-04-01 00:30")  # 19:00 UTC on 31 Mar = 00:30 IST on 1 Apr
        deleted = self.row(body, "F4")
        # the previous department was deleted (NULL): shown as a dash, and the move still counts as a department change
        self.assertEqual(
            (deleted["previousDepartment"], deleted["newDepartment"], deleted["change"]), (None, "SEWING", "Department")
        )
        self.assertEqual(
            (deleted["previousDesignation"], deleted["newDesignation"], deleted["promotedBy"]), (None, None, None)
        )

    def test_summary_and_notes(self):
        body = self.run_report(self.RID)
        cards = self.cards(body)
        self.assertEqual(cards["Promotions"], 5)
        self.assertEqual(cards["Employees promoted"], 4)
        self.assertEqual((cards["Designation changed"], cards["Department changed"]), (3, 3))
        notes = " ".join(body["notes"])
        self.assertIn("Promoted into: Supervisor (2), Manager (1).", notes)

    def test_change_type_filter(self):
        self.assertEqual(self.by_change(changeType="designation"), [("F1", "Designation"), ("F10", "Designation")])
        self.assertEqual(self.by_change(changeType="department"), [("F4", "Department"), ("F3", "Department")])
        self.assertEqual(self.by_change(changeType="both"), [("F1", "Both")])
        self.assertEqual(self.get("run", self.RID, changeType="pay").status_code, 400)

    def test_department_filter_matches_previous_or_new_department(self):
        # PACKING: F1 moved INTO it, F3 moved OUT of it
        got = self.by_change(departmentIds=str(self.packing.id))
        self.assertEqual(sorted(got), [("F1", "Both"), ("F3", "Department")])
        # SEWING: F10 promoted inside it, F4 moved into it
        self.assertEqual(
            sorted(self.by_change(departmentIds=str(self.sewing.id))), [("F10", "Designation"), ("F4", "Department")]
        )
        self.assertEqual(
            sorted(self.by_change(departmentIds=f"{self.packing.id},{self.sewing.id}")),
            [("F1", "Both"), ("F10", "Designation"), ("F3", "Department"), ("F4", "Department")],
        )

    def test_department_and_designation_filters_say_what_they_match(self):
        catalog = self.client.get("/api/reports/catalog", **headers(self.admin)).json()
        spec = next(r for r in catalog["reports"] if r["id"] == PROMOTION)
        labels = {f["key"]: f["label"] for f in spec["filters"]}
        self.assertEqual(labels["department"], "Department (moved from or to)")
        self.assertEqual(labels["designation"], "Designation (moved from or to)")
        body = self.run_report(self.RID, departmentIds=str(self.packing.id))
        self.assertIn(("Department (moved from or to)", "PACKING"), [(f["label"], f["value"]) for f in body["filters"]])

    def test_designation_filter_matches_previous_or_new_designation(self):
        self.assertEqual(self.by_change(designationIds=str(self.d_mgr.id)), [("F1", "Both")])
        # Supervisor: promoted into (F1 Apr, F10) or out of (F1 Sep)
        self.assertEqual(
            sorted(self.by_change(designationIds=str(self.d_sup.id))),
            [("F1", "Both"), ("F1", "Designation"), ("F10", "Designation")],
        )

    def test_windows_and_financial_year(self):
        def eff(**p):
            return sorted(r["effectiveDate"] for r in self.data(self.run_report(self.RID, **p)))

        self.assertEqual(eff(effective="thisMonth"), ["2026-09-01"])
        self.assertEqual(eff(effective="last3"), ["2026-07-01", "2026-08-10", "2026-09-01"])
        self.assertEqual(eff(effective="thisFY"), ["2026-04-01", "2026-07-01", "2026-08-10", "2026-09-01"])
        self.assertEqual(eff(effective="lastFY"), ["2025-06-01"])
        self.assertEqual(eff(financialYear="2025-26"), ["2025-06-01"])
        self.assertEqual(eff(financialYear="2026-27", effective="lastMonth"), ["2026-08-10"])
        self.assertEqual(eff(financialYear="2026-27", changeType="designation"), ["2026-04-01"])
        self.assertTrue(
            any(
                "Only promotions effective from" in n for n in self.run_report(self.RID, effective="thisMonth")["notes"]
            )
        )

    def test_employee_filters(self):
        self.assertEqual(self.codes(self.run_report(self.RID, employeeIds=str(self.e3.id))), ["F10"])
        self.assertEqual(
            sorted(self.codes(self.run_report(self.RID, employmentType="staff"))), ["F1", "F1", "F10", "F3", "F4"]
        )
        self.assertEqual(self.run_report(self.RID, employmentType="production")["rows"], [])
        self.assertEqual(sorted(self.codes(self.run_report(self.RID, branchIds=str(self.b2.id)))), ["F10"])

    def test_branch_isolation(self):
        body = self.run_report(self.RID, user=self.b1_user)
        self.assertEqual(sorted(self.codes(body)), ["F1", "F1", "F3"])  # F10 = Unit 2, F4 has no branch
        self.assertEqual(self.cards(body)["Promotions"], 3)
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, branchIds=str(self.b2.id))["rows"], [])
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, employeeIds=str(self.e3.id))["rows"], [])
        # the department filter does not let a branch user reach another branch's promotions
        self.assertEqual(self.run_report(self.RID, user=self.b1_user, departmentIds=str(self.sewing.id))["rows"], [])

    def test_agrees_with_the_real_promotion_endpoint(self):
        Promotion.objects.all().delete()
        emp = Employee.objects.create(
            employee_code="PR1",
            first_name="Pro",
            last_name="Real",
            employment_type="staff",
            status="active",
            department=self.cutting,
            designation=self.d_op,
            branch=self.b1,
        )
        r = self.client.post(
            "/api/promotions",
            data=json.dumps(
                {
                    "employeeId": emp.id,
                    "newDesignationId": self.d_sup.id,
                    "newDepartmentId": self.packing.id,
                    "effectiveDate": "2026-09-01",
                    "notes": "Shift in-charge",
                }
            ),
            content_type="application/json",
            **headers(self.admin),
        )
        self.assertEqual(r.status_code, 201)
        stored = Promotion.objects.get(employee=emp)
        row = self.data(self.run_report(self.RID))[0]
        self.assertEqual((row["previousDesignation"], row["newDesignation"]), ("Operator", "Supervisor"))
        self.assertEqual(
            (row["previousDepartment"], row["newDepartment"], row["change"]), ("CUTTING", "PACKING", "Both")
        )
        self.assertEqual(
            (row["effectiveDate"], row["notes"], row["promotedBy"]),
            ("2026-09-01", "Shift in-charge", stored.promoted_by or None),
        )
        emp.refresh_from_db()
        self.assertEqual(emp.designation_id, self.d_sup.id)


# ═════════════════════════════════════════════════════════════════════════════
#  Contract: registration, permissions, exports, empty results, limits, N+1
# ═════════════════════════════════════════════════════════════════════════════


class FinanceContractTests(_Case):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        make_advances(cls)
        make_bonus_rows(cls)
        make_increments(cls)
        make_promotions(cls)

    def test_reports_are_registered_correctly(self):
        specs = {s.id: s for s in registry.all_specs()}
        self.assertEqual(registry.LOAD_ERRORS, {})
        for rid in ALL_IDS:
            s = specs[rid]
            self.assertEqual(s.category, "finance", rid)
            self.assertEqual(s.modules, (OWNING_MODULE[rid],), rid)
            self.assertFalse(s.super_admin_only, rid)
            self.assertTrue(s.title and s.description and s.icon, rid)
            self.assertTrue(s.columns, rid)
            for m in s.modules:
                self.assertIn(m, all_module_keys())
            keys = [c.key for c in s.columns]
            self.assertEqual(len(keys), len(set(keys)), rid)
            self.assertIsNone(s.family, rid)  # these five are stand-alone reports, not variants of one subject
            self.assertRegex(s.id, r"^[a-z0-9-]+$")

    def test_catalog_lists_them_for_the_owning_module_only(self):
        def ids(user):
            body = self.client.get("/api/reports/catalog", **headers(user)).json()
            return {r["id"] for r in body["reports"]} & set(ALL_IDS)

        self.assertEqual(ids(self.admin), set(ALL_IDS))
        self.assertEqual(ids(self.b1_user), set(ALL_IDS))
        self.assertEqual(ids(self.reports_only), set())
        self.assertEqual(ids(self.by_module["settlement"]), {ADV_LEDGER, ADV_SCHEDULE})
        self.assertEqual(ids(self.by_module["bonus"]), {BONUS})
        self.assertEqual(ids(self.by_module["increment"]), {INCREMENT})
        self.assertEqual(ids(self.by_module["promotion"]), {PROMOTION})

    def test_reports_permission_alone_gets_report_forbidden(self):
        for rid in ALL_IDS:
            for kind, extra in (("run", {}), ("export", {"fmt": "xlsx"})):
                r = self.client.get(
                    f"/api/reports/{kind}/{rid}", {**FULL_PARAMS[rid], **extra}, **headers(self.reports_only)
                )
                self.assertEqual(r.status_code, 403, (rid, kind))
                self.assertEqual(r.json()["error"], "report_forbidden", (rid, kind))

    def test_only_the_owning_module_opens_a_report(self):
        for rid in ALL_IDS:
            for module, user in self.by_module.items():
                r = self.client.get(f"/api/reports/run/{rid}", FULL_PARAMS[rid], **headers(user))
                expected = 200 if module == OWNING_MODULE[rid] else 403
                self.assertEqual(r.status_code, expected, (rid, module))
                if expected == 403:
                    self.assertEqual(r.json()["error"], "report_forbidden")

    def test_denied_export_writes_no_audit_row(self):
        from .models import AuditLog

        before = AuditLog.objects.filter(action="export", module="reports").count()
        self.export(BONUS, "pdf", user=self.by_module["settlement"])
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), before)
        self.export(BONUS, "pdf", user=self.by_module["bonus"])
        self.assertEqual(AuditLog.objects.filter(action="export", module="reports").count(), before + 1)

    def test_screen_totals_equal_the_sum_of_rows_for_every_report(self):
        for rid in ALL_IDS:
            body = self.run_report(rid, **FULL_PARAMS[rid])
            self.assertTrue(self.data(body), rid)
            for col in body["columns"]:
                if col["total"] != "sum":
                    continue
                self.assertEqual(body["totals"][col["key"]], self.col_sum(body, col["key"]), (rid, col["key"]))

    def test_xlsx_round_trip_and_signature(self):
        first = {ADV_LEDGER: "F1", ADV_SCHEDULE: "F1", BONUS: "F1", INCREMENT: "F4", PROMOTION: "F1"}
        for rid in ALL_IDS:
            r = self.export(rid, "xlsx", **FULL_PARAMS[rid])
            self.assertEqual(r.status_code, 200, rid)
            self.assertTrue(r.content.startswith(b"PK"), rid)
            self.assertIn("spreadsheetml", r["Content-Type"])
            ws = load_workbook(io.BytesIO(r.content)).active
            header = [c.value for c in ws[7]]
            self.assertEqual(header[0], "Emp Code", rid)
            self.assertEqual(header[1], "Employee", rid)
            self.assertEqual(ws["A8"].value, first[rid], rid)
            self.assertEqual(ws["B8"].value, "Asha Fin" if first[rid] == "F1" else "Esha Fin", rid)

    def test_xlsx_carries_the_money_and_the_total_row(self):
        ws = load_workbook(io.BytesIO(self.export(ADV_LEDGER, "xlsx").content)).active
        header = [c.value for c in ws[7]]
        first = [c.value for c in ws[8]]
        self.assertEqual(first[header.index("Recovered")], 6000.0)
        self.assertEqual(first[header.index("Outstanding")], 6000.0)
        self.assertEqual(first[header.index("Instalments paid")], "2 of 4")
        self.assertEqual(first[header.index("Type")], "Term Loan")
        total = [c.value for c in ws[8 + 8]]  # 8 advances
        self.assertEqual(total[0], "TOTAL")
        self.assertEqual(total[header.index("Recovered")], 8700.0)
        self.assertEqual(total[header.index("Outstanding")], 15000.0)
        notes = [ws.cell(row=r, column=1).value for r in range(17, ws.max_row + 1) if ws.cell(row=r, column=1).value]
        self.assertTrue(any("Recovered = instalments already deducted" in n for n in notes))

    def test_xlsx_bonus_percent_and_promotion_blank_cells(self):
        ws = load_workbook(io.BytesIO(self.export(BONUS, "xlsx", financialYear="2025-26").content)).active
        header = [c.value for c in ws[7]]
        self.assertEqual(ws.cell(row=8, column=header.index("Bonus %") + 1).value, 8.33)
        ws = load_workbook(io.BytesIO(self.export(PROMOTION, "xlsx", changeType="department").content)).active
        header = [c.value for c in ws[7]]
        row = {h: c.value for h, c in zip(header, ws[8])}
        self.assertEqual(row["Emp Code"], "F4")
        self.assertIsNone(row["Previous department"])  # an unknown value stays an empty cell, never a 0

    def test_xlsx_cells_are_typed_dates_and_numbers(self):
        ws = load_workbook(io.BytesIO(self.export(INCREMENT, "xlsx").content)).active
        header = [c.value for c in ws[7]]
        row = {h: c.value for h, c in zip(header, ws[8])}
        self.assertEqual(row["Emp Code"], "F4")
        self.assertEqual(row["Effective date"].date(), date(2026, 10, 1))  # a real date cell, not text
        self.assertEqual((row["Previous salary"], row["New salary"], row["Increment"]), (30000.0, 33000.0, 3000.0))
        self.assertEqual(row["Increment %"], 10.0)
        self.assertIsNone(row["Recorded by"])  # nobody recorded it: empty cell, not "None" or 0
        ws = load_workbook(io.BytesIO(self.export(ADV_SCHEDULE, "xlsx", **FULL_PARAMS[ADV_SCHEDULE]).content)).active
        header = [c.value for c in ws[7]]
        row = {h: c.value for h, c in zip(header, ws[8])}
        self.assertEqual((row["Emp Code"], row["Due month"], row["Status"]), ("F1", "Jun 2026", "Deducted"))
        self.assertEqual(
            (row["Instalment"], row["Instalment amount"], row["Balance after"]), ("1 of 4", 3000.0, 9000.0)
        )

    def test_free_text_written_by_users_can_never_become_a_formula(self):
        make_advance(self.e2, "general", "100", "pending", ist(2026, 9, 12), purpose='=HYPERLINK("http://x","y")')
        ws = load_workbook(io.BytesIO(self.export(ADV_LEDGER, "xlsx").content)).active
        header = [c.value for c in ws[7]]
        col = header.index("Purpose") + 1
        cells = [ws.cell(row=r, column=col) for r in range(8, ws.max_row + 1)]
        risky = [c for c in cells if isinstance(c.value, str) and "HYPERLINK" in c.value]
        self.assertEqual(len(risky), 1)
        self.assertNotEqual(risky[0].data_type, "f")

    def test_pdf_text_carries_the_title_the_figures_and_the_total(self):
        import pdfplumber

        def pdf_text(rid, **params):
            with pdfplumber.open(io.BytesIO(self.export(rid, "pdf", **params).content)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)

        text = pdf_text(ADV_LEDGER)
        self.assertIn("ADVANCE & LOAN LEDGER", text.upper())
        for expected in (
            "Asha Fin",
            "12,000.00",
            "TOTAL",
            "15,000.00",
            "8,700.00",
            "Recovered = instalments already deducted",
        ):
            self.assertIn(expected, text)
        text = pdf_text(BONUS, financialYear="2025-26")
        self.assertIn("STATUTORY BONUS REGISTER", text.upper())
        for expected in ("Financial year: FY 2025-26", "19,159.00", "CUTTING total", "Weekly slips"):
            self.assertIn(expected, text)
        text = pdf_text(PROMOTION)
        self.assertIn("Line supervisor", text)
        self.assertIn("Promoted into: Supervisor (2), Manager (1).", text)

    def test_pdf_export_signature(self):
        for rid in ALL_IDS:
            r = self.export(rid, "pdf", **FULL_PARAMS[rid])
            self.assertEqual(r.status_code, 200, rid)
            self.assertTrue(r.content.startswith(b"%PDF"), rid)
            self.assertIn(b"%%EOF", r.content[-64:], rid)
            self.assertIn("application/pdf", r["Content-Type"])

    def test_empty_result_works_on_screen_and_in_both_exports(self):
        for rid in ALL_IDS:
            params = {**FULL_PARAMS[rid], "employeeIds": str(self.e_none.id)}
            body = self.run_report(rid, **params)
            self.assertEqual(body["rows"], [], rid)
            self.assertEqual(body["rowCount"], 0)
            self.assertTrue(body["summary"], rid)  # cards still render (zeros, not errors)
            for fmt, sig in (("xlsx", b"PK"), ("pdf", b"%PDF")):
                r = self.export(rid, fmt, **params)
                self.assertEqual(r.status_code, 200, (rid, fmt))
                self.assertTrue(r.content.startswith(sig), (rid, fmt))
            ws = load_workbook(io.BytesIO(self.export(rid, "xlsx", **params).content)).active
            self.assertEqual(ws["A7"].value, "Emp Code")

    def test_empty_summary_cards_have_no_fake_averages(self):
        cards = self.cards(self.run_report(BONUS, employeeIds=str(self.e_none.id)))
        self.assertIsNone(cards["Average bonus"])
        self.assertEqual(cards["Total bonus"], 0.0)
        cards = self.cards(self.run_report(INCREMENT, employeeIds=str(self.e_none.id)))
        self.assertIsNone(cards["Average increment (%)"])
        self.assertIsNone(cards["Highest increment (%)"])

    def test_row_limit_is_honoured_in_the_query_and_flagged(self):
        with mock.patch("api.reporting.runner.SCREEN_ROW_LIMIT", 2):
            for rid in ALL_IDS:
                body = self.run_report(rid, **FULL_PARAMS[rid])
                self.assertTrue(body["truncated"], rid)
                self.assertEqual(body["rowCount"], 2, rid)
                self.assertIsNone(body["totals"], rid)  # a sum over a cut-off list would under-state the truth
        with mock.patch("api.reporting.runner.XLSX_ROW_LIMIT", 2):
            r = self.export(ADV_LEDGER, "xlsx")
            self.assertEqual(r.status_code, 413)
            self.assertEqual(r.json()["error"], "too_many_rows")

    def test_row_limit_is_applied_in_the_query_itself(self):
        from types import SimpleNamespace

        from .reporting.filters import ReportContext, parse_params

        for rid in ALL_IDS:
            spec = registry.get_spec(rid)
            ctx = ReportContext(
                request=SimpleNamespace(), spec=spec, params=parse_params(spec, FULL_PARAMS[rid]), row_limit=3
            )
            self.assertLessEqual(len(spec.run(ctx).rows), 3, rid)  # never loads the whole table for a capped run

    def test_reports_never_write(self):
        models = (Advance, AdvanceRepayment, Bonus, SalaryIncrement, Promotion, AttendanceDayRecord, SalarySlip)
        before = {m.__name__: list(m.objects.order_by("pk").values()) for m in models}
        for rid in ALL_IDS:
            self.run_report(rid, **FULL_PARAMS[rid])
            self.export(rid, "xlsx", **FULL_PARAMS[rid])
        after = {m.__name__: list(m.objects.order_by("pk").values()) for m in models}
        self.assertEqual(before, after)


_SEQ = itertools.count(1)


class QueryCountTests(TestCase):
    """3 employees vs 15 employees must cost (almost) the same number of queries: no N+1."""

    @classmethod
    def setUpTestData(cls):
        make_org(cls)

    def setUp(self):
        patcher = mock.patch("api.reporting.filters.ist_today", return_value=TODAY)
        patcher.start()
        self.addCleanup(patcher.stop)

    def seed(self, n):
        for _ in range(n):
            i = next(_SEQ)
            dept = (self.cutting, self.packing, self.sewing)[i % 3]
            emp = Employee.objects.create(
                employee_code=f"Q{i}",
                first_name=f"Bulk{i}",
                last_name="Fin",
                department=dept,
                designation=(self.d_op, self.d_sup, None)[i % 3],
                branch=self.b1,
                employment_type="staff",
            )
            make_advance(
                emp,
                "term",
                "3000",
                "approved",
                ist(2026, 8, 1),
                approved_at=utc(2026, 8, 2),
                approved_by="HR",
                emi="1000",
                months=3,
                start=(8, 2026),
                reps=[(2026, 8, 1000, True), (2026, 9, 1000, False), (2026, 10, 1000, False)],
            )
            make_advance(emp, "general", "500", "pending", ist(2026, 9, 2))
            Bonus.objects.create(
                employee=emp,
                financial_year="2025-26",
                records_considered=12,
                calculation_base=Decimal("1000"),
                bonus_percent_applied=Decimal("8.33"),
                bonus_amount=Decimal("83.30"),
                status="calculated",
            )
            SalaryIncrement.objects.create(
                employee=emp,
                previous_salary=Decimal("10000"),
                new_salary=Decimal("11000"),
                percent=Decimal("10"),
                effective_date=date(2026, 7, 1),
            )
            Promotion.objects.create(
                employee=emp,
                previous_department=self.packing,
                new_department=dept,
                previous_designation=self.d_op,
                new_designation=self.d_sup,
                effective_date=date(2026, 7, 1),
            )

    def count(self, rid):
        params = dict(FULL_PARAMS[rid])
        self.client.get(f"/api/reports/run/{rid}", params, **headers(self.admin))  # warm-up (settings singleton etc.)
        with CaptureQueriesContext(connection) as ctx:
            r = self.client.get(f"/api/reports/run/{rid}", params, **headers(self.admin))
        self.assertEqual(r.status_code, 200)
        return len(ctx), len([x for x in r.json()["rows"] if "_kind" not in x])

    def test_query_count_is_flat(self):
        self.seed(3)
        small = {rid: self.count(rid) for rid in ALL_IDS}
        self.seed(12)
        big = {rid: self.count(rid) for rid in ALL_IDS}
        for rid in ALL_IDS:
            (q_small, n_small), (q_big, n_big) = small[rid], big[rid]
            self.assertGreaterEqual(n_big, n_small + 12, rid)  # the data really did grow
            self.assertLessEqual(
                q_big - q_small, 1, f"{rid}: {q_small} queries for {n_small} rows vs {q_big} for {n_big}"
            )
            self.assertLessEqual(q_big, 8, f"{rid}: {q_big} queries")

    def test_query_count_is_flat_for_the_department_ordered_ledger(self):
        self.seed(3)
        params = {"sortBy": "department"}
        self.client.get(f"/api/reports/run/{ADV_LEDGER}", params, **headers(self.admin))
        with CaptureQueriesContext(connection) as small:
            self.client.get(f"/api/reports/run/{ADV_LEDGER}", params, **headers(self.admin))
        self.seed(12)
        with CaptureQueriesContext(connection) as big:
            self.client.get(f"/api/reports/run/{ADV_LEDGER}", params, **headers(self.admin))
        self.assertLessEqual(len(big) - len(small), 1)
